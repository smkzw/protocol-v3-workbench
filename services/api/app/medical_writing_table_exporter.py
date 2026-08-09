from __future__ import annotations

import hashlib
import io
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Callable, Dict, List, Mapping, Optional, Sequence, Tuple

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor, Twips
from lxml import etree

from packages.contracts.workbench_contracts import (
    StructuredTable,
    StructuredTableCell,
    StructuredTableNote,
    StructuredTableRow,
)


_FIXED_PACKAGE_TIMESTAMP = (2000, 1, 1, 0, 0, 0)
_FIXED_CORE_TIMESTAMP = datetime(2000, 1, 1, tzinfo=timezone.utc)
_WORDPROCESSINGML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_DRAWINGML_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_EXPORT_EAST_ASIA_FONT = "宋体"
_EXPORT_LATIN_FONT = "Times New Roman"
_EXPORT_SVG_FONT_STACK = "'Times New Roman', SimSun, 'Songti SC', serif"
_MIN_COLUMN_WIDTH_TWIPS = 120
_SUPPORTED_SPLIT_STRATEGIES = {
    "repeat_header_keep_rows_intact",
    "repeat_header_allow_row_split",
}


class StructuredTableDocxExportError(ValueError):
    """Raised when a structured table cannot be represented as a regular Word grid."""


@dataclass(frozen=True)
class StructuredTableDocxExportResult:
    content: bytes
    metadata: Mapping[str, object]


@dataclass(frozen=True)
class _PlacedCell:
    cell: StructuredTableCell
    row_index: int
    column_index: int


@dataclass(frozen=True)
class _ExportPlan:
    rows: Sequence[StructuredTableRow]
    column_ids: Sequence[str]
    placed_cells: Sequence[_PlacedCell]
    column_widths_twips: Sequence[int]
    orientation: str
    repeat_header_rows: int
    split_strategy: str
    page_width_twips: int
    page_height_twips: int
    margins_twips: Mapping[str, int]


def export_structured_table_docx(
    table: StructuredTable,
) -> StructuredTableDocxExportResult:
    """Render one StructuredTable as deterministic, independently readable DOCX bytes."""

    plan = _build_export_plan(table)
    document = Document()
    _configure_document(document, table, plan)
    _render_title(document, table)
    word_table = _render_table(document, table, plan)
    _render_notes(document, table.notes)

    raw = io.BytesIO()
    document.save(raw)
    content = _normalize_docx_package(raw.getvalue())
    metadata: Dict[str, object] = {
        "schema_version": "structured_table_docx_export_v1",
        "table_id": table.table_id,
        "row_count": len(plan.rows),
        "column_count": len(plan.column_ids),
        "orientation": plan.orientation,
        "page_width_twips": plan.page_width_twips,
        "page_height_twips": plan.page_height_twips,
        "margins_twips": dict(plan.margins_twips),
        "column_widths_twips": list(plan.column_widths_twips),
        "repeat_header_rows": plan.repeat_header_rows,
        "long_table_split_strategy": plan.split_strategy,
        "merged_ranges": [
            _merge_range(placed)
            for placed in plan.placed_cells
            if placed.cell.row_span > 1 or placed.cell.column_span > 1
        ],
        "note_markers": [marker for marker, _ in _ordered_notes(table.notes)],
        "docx_sha256": hashlib.sha256(content).hexdigest(),
        "word_table_rows": len(word_table.rows),
        "word_table_columns": len(plan.column_ids),
    }
    return StructuredTableDocxExportResult(content=content, metadata=metadata)


def _build_export_plan(table: StructuredTable) -> _ExportPlan:
    if not table.columns:
        raise StructuredTableDocxExportError(
            "structured table must contain at least one column"
        )
    if not table.rows:
        raise StructuredTableDocxExportError(
            "structured table must contain at least one row"
        )

    columns = sorted(table.columns, key=lambda item: (item.order, item.column_id))
    rows = sorted(table.rows, key=lambda item: (item.order, item.row_id))
    _require_unique_orders([column.order for column in columns], "column")
    _require_unique_orders([row.order for row in rows], "row")
    column_ids = [column.column_id for column in columns]
    column_index_by_id = {
        column_id: index for index, column_id in enumerate(column_ids)
    }

    occupancy: List[List[Optional[str]]] = [[None for _ in columns] for _ in rows]
    placed_cells: List[_PlacedCell] = []
    for row_index, row in enumerate(rows):
        for cell in sorted(
            row.cells,
            key=lambda item: (column_index_by_id[item.column_id], item.cell_id),
        ):
            if _is_covered_merge_continuation(cell):
                continue
            column_index = column_index_by_id[cell.column_id]
            row_end = row_index + cell.row_span
            column_end = column_index + cell.column_span
            if row_end > len(rows) or column_end > len(columns):
                raise StructuredTableDocxExportError(
                    f"cell {cell.cell_id} merge range exceeds the table grid"
                )
            for occupied_row in range(row_index, row_end):
                for occupied_column in range(column_index, column_end):
                    existing = occupancy[occupied_row][occupied_column]
                    if existing is not None:
                        raise StructuredTableDocxExportError(
                            f"cell {cell.cell_id} overlaps cell {existing} at "
                            f"row {occupied_row}, column {occupied_column}"
                        )
                    occupancy[occupied_row][occupied_column] = cell.cell_id
            placed_cells.append(
                _PlacedCell(
                    cell=cell,
                    row_index=row_index,
                    column_index=column_index,
                )
            )

    holes = [
        (row_index, column_index)
        for row_index, row in enumerate(occupancy)
        for column_index, cell_id in enumerate(row)
        if cell_id is None
    ]
    if holes:
        preview = ", ".join(f"({row}, {column})" for row, column in holes[:5])
        raise StructuredTableDocxExportError(
            f"structured table has uncovered grid positions: {preview}"
        )

    layout = dict(table.word_layout)
    orientation = str(layout.get("orientation", "auto")).strip().lower()
    if orientation == "auto":
        orientation = "landscape" if len(columns) >= 8 else "portrait"
    if orientation not in {"portrait", "landscape"}:
        raise StructuredTableDocxExportError(
            "word_layout.orientation must be portrait, landscape, or auto"
        )

    page_width_twips, page_height_twips = _page_dimensions(orientation)
    margins_twips = _margins(layout)
    available_width = page_width_twips - margins_twips["left"] - margins_twips["right"]
    if available_width <= 0:
        raise StructuredTableDocxExportError("page margins leave no usable table width")
    column_widths = _column_widths(columns, available_width, layout)

    repeat_header_rows = _repeat_header_rows(table, layout)
    split_strategy = str(
        layout.get("long_table_split_strategy", "repeat_header_keep_rows_intact")
    ).strip()
    if split_strategy not in _SUPPORTED_SPLIT_STRATEGIES:
        choices = ", ".join(sorted(_SUPPORTED_SPLIT_STRATEGIES))
        raise StructuredTableDocxExportError(
            f"unsupported long table split strategy; expected one of: {choices}"
        )

    return _ExportPlan(
        rows=rows,
        column_ids=column_ids,
        placed_cells=placed_cells,
        column_widths_twips=column_widths,
        orientation=orientation,
        repeat_header_rows=repeat_header_rows,
        split_strategy=split_strategy,
        page_width_twips=page_width_twips,
        page_height_twips=page_height_twips,
        margins_twips=margins_twips,
    )


def _is_covered_merge_continuation(cell: StructuredTableCell) -> bool:
    value = cell.semantic_value
    if not isinstance(value, dict):
        return False
    metadata = value.get("_medical_writing_table")
    return isinstance(metadata, dict) and bool(metadata.get("hidden"))


def _configure_document(
    document: Document, table: StructuredTable, plan: _ExportPlan
) -> None:
    section = document.sections[0]
    section.orientation = (
        WD_ORIENT.LANDSCAPE if plan.orientation == "landscape" else WD_ORIENT.PORTRAIT
    )
    section.page_width = Twips(plan.page_width_twips)
    section.page_height = Twips(plan.page_height_twips)
    section.left_margin = Twips(plan.margins_twips["left"])
    section.right_margin = Twips(plan.margins_twips["right"])
    section.top_margin = Twips(plan.margins_twips["top"])
    section.bottom_margin = Twips(plan.margins_twips["bottom"])

    properties = document.core_properties
    properties.title = table.title or table.table_id
    properties.subject = f"Structured table export: {table.domain.value}"
    properties.identifier = table.table_id
    properties.author = "CMS AI医学经理工作台"
    properties.last_modified_by = "CMS AI医学经理工作台"
    properties.created = _FIXED_CORE_TIMESTAMP
    properties.modified = _FIXED_CORE_TIMESTAMP


def _render_title(document: Document, table: StructuredTable) -> None:
    if table.role.value == "layout" or not table.title.strip():
        return
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(table.title.strip())
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_run_fonts(run, east_asia="宋体", latin="Times New Roman")


def _render_table(
    document: Document,
    table: StructuredTable,
    plan: _ExportPlan,
    *,
    rich_text_run_renderer: Optional[
        Callable[
            [object, str, Sequence[Mapping[str, object]], Mapping[str, object]], object
        ]
    ] = None,
):
    layout = dict(table.word_layout)
    word_table = document.add_table(rows=len(plan.rows), cols=len(plan.column_ids))
    table_style = layout.get("table_style", "Table Grid")
    word_table.style = (
        table_style if isinstance(table_style, str) and table_style else None
    )
    word_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    word_table.autofit = False
    table_width_type, table_width_value = _table_width_spec(
        layout,
        sum(plan.column_widths_twips),
    )
    cell_widths = (
        _scale_widths(plan.column_widths_twips, table_width_value)
        if table_width_type == "pct"
        else list(plan.column_widths_twips)
    )
    _set_fixed_table_layout(
        word_table,
        table_width_value,
        width_type=table_width_type,
    )
    _set_grid_column_widths(word_table, plan.column_widths_twips)
    if table.role.value == "layout":
        _set_table_borders_none(word_table)
    elif layout.get("explicit_black_borders") is True:
        _set_table_borders_single_black(word_table)

    row_min_heights = _row_min_heights(layout, len(plan.rows))
    for row_index, row in enumerate(word_table.rows):
        is_header = row_index < plan.repeat_header_rows
        is_section = _normalized_role(plan.rows[row_index].style_role) == "section"
        if is_header:
            _set_repeat_header(row)
        if (
            plan.split_strategy == "repeat_header_keep_rows_intact"
            or is_header
            or is_section
        ):
            _prevent_row_split(row)
        if row_index in row_min_heights:
            _set_row_min_height(row, row_min_heights[row_index])
        for column_index, width in enumerate(cell_widths):
            _set_cell_width(
                row.cells[column_index],
                width,
                width_type=table_width_type,
            )

    merged_cells: Dict[str, object] = {}
    for placed in plan.placed_cells:
        cell = placed.cell
        start = word_table.cell(placed.row_index, placed.column_index)
        if cell.row_span > 1 or cell.column_span > 1:
            end = word_table.cell(
                placed.row_index + cell.row_span - 1,
                placed.column_index + cell.column_span - 1,
            )
            start = start.merge(end)
        merged_cells[cell.cell_id] = start
        width = sum(
            cell_widths[placed.column_index : placed.column_index + cell.column_span]
        )
        _set_cell_width(start, width, width_type=table_width_type)

    font_size = _font_size(table, len(plan.column_ids))
    for placed in plan.placed_cells:
        source_row = plan.rows[placed.row_index]
        word_cell = merged_cells[placed.cell.cell_id]
        cell_role = _normalized_role(placed.cell.style_role)
        row_role = _normalized_role(source_row.style_role)
        role = row_role if cell_role == "body" and row_role != "body" else cell_role
        is_label = str(placed.cell.style_role or "").strip().lower() == "label" or (
            placed.column_index == 0 and role == "body"
        )
        _write_cell(
            word_cell,
            placed.cell,
            role=role,
            font_size=font_size,
            first_column=placed.column_index == 0,
            is_label=is_label,
            layout=layout,
            rich_text_run_renderer=rich_text_run_renderer,
        )
    return word_table


def _set_table_borders_none(word_table) -> None:
    properties = word_table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "nil")
        borders.append(border)
    _insert_before_first(
        properties,
        borders,
        (
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
            "w:tblPrChange",
        ),
    )


def _set_table_borders_single_black(word_table) -> None:
    properties = word_table._tbl.tblPr
    existing = properties.find(qn("w:tblBorders"))
    if existing is not None:
        properties.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement(f"w:{edge}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "auto")
        borders.append(border)
    _insert_before_first(
        properties,
        borders,
        (
            "w:shd",
            "w:tblLayout",
            "w:tblCellMar",
            "w:tblLook",
            "w:tblCaption",
            "w:tblDescription",
            "w:tblPrChange",
        ),
    )


def _render_notes(document: Document, notes: Sequence[StructuredTableNote]) -> None:
    ordered = _ordered_notes(notes)
    if not ordered:
        return
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run("附注")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_run_fonts(run, east_asia="宋体", latin="Times New Roman")

    for marker, note in ordered:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Mm(4)
        paragraph.paragraph_format.first_line_indent = Mm(-4)
        paragraph.paragraph_format.space_after = Pt(1.5)
        prefix = f"{marker} " if marker else ""
        run = paragraph.add_run(f"{prefix}{note.text.strip()}")
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(0, 0, 0)
        _set_run_fonts(run, east_asia="宋体", latin="Times New Roman")


def _write_cell(
    word_cell,
    cell: StructuredTableCell,
    role: str,
    font_size: float,
    first_column: bool,
    is_label: bool,
    layout: Mapping[str, object],
    rich_text_run_renderer=None,
) -> None:
    if isinstance(cell.rich_text, dict):
        _write_rich_cell(
            word_cell,
            cell.rich_text,
            role=role,
            font_size=font_size,
            first_column=first_column,
            is_label=is_label,
            layout=layout,
            rich_text_run_renderer=rich_text_run_renderer,
        )
        return
    word_cell.text = ""
    paragraph = word_cell.paragraphs[0]
    _format_cell_paragraph(
        paragraph,
        {},
        role=role,
        font_size=font_size,
        first_column=first_column,
        is_label=is_label,
        layout=layout,
    )
    run = paragraph.add_run()
    lines = str(cell.text or "").split("\n")
    for index, line in enumerate(lines):
        if index:
            run.add_break()
        run.add_text(line)
    run.bold = role in {"header", "section"} or (
        is_label and layout.get("bold_labels") is True
    )
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_run_fonts(run, east_asia="宋体", latin="Times New Roman")
    _apply_cell_vertical_alignment(
        word_cell,
        role=role,
        is_label=is_label,
        layout=layout,
    )
    _set_cell_shading(
        word_cell,
        _cell_fill(layout, role),
    )
    _apply_cell_margins(word_cell, layout)


_RICH_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_RICH_HIGHLIGHTS = {
    "#FFFF00": WD_COLOR_INDEX.YELLOW,
    "#00FF00": WD_COLOR_INDEX.BRIGHT_GREEN,
    "#00FFFF": WD_COLOR_INDEX.TURQUOISE,
    "#FF00FF": WD_COLOR_INDEX.PINK,
}
_RICH_PRESETS = {
    "heading_1": (16.0, "宋体", True),
    "heading_2": (14.0, "宋体", True),
    "heading_3": (12.0, "宋体", True),
    "heading_4": (10.5, "宋体", True),
    "body": (10.5, "宋体", False),
    "note": (9.0, "宋体", False),
}


def _write_rich_cell(
    word_cell,
    root: Mapping[str, object],
    *,
    role: str,
    font_size: float,
    first_column: bool,
    is_label: bool,
    layout: Mapping[str, object],
    rich_text_run_renderer=None,
) -> None:
    blocks = _rich_cell_blocks(root)
    word_cell.text = ""
    if not blocks:
        blocks = [({"type": "paragraph", "content": []}, "", None)]
    numbering_cache: dict[int, int] = {}
    for index, (node, _, list_spec) in enumerate(blocks):
        paragraph = word_cell.paragraphs[0] if index == 0 else word_cell.add_paragraph()
        _format_rich_cell_paragraph(
            paragraph,
            node,
            role=role,
            font_size=font_size,
            first_column=first_column,
            is_label=is_label,
            layout=layout,
        )
        if list_spec is not None:
            _apply_word_list_numbering(
                paragraph,
                list_spec,
                numbering_cache=numbering_cache,
            )
        content = node.get("content", [])
        if not isinstance(content, list):
            raise StructuredTableDocxExportError(
                "table cell rich_text content must be a list"
            )
        for child in content:
            if not isinstance(child, Mapping):
                raise StructuredTableDocxExportError(
                    "table cell rich_text contains an invalid inline node"
                )
            child_type = str(child.get("type") or "")
            if child_type == "hardBreak":
                paragraph.add_run().add_break()
                continue
            if child_type != "text":
                raise StructuredTableDocxExportError(
                    f"unsupported table cell rich_text inline node: {child_type or '<empty>'}"
                )
            marks = child.get("marks", [])
            if not isinstance(marks, list) or any(
                not isinstance(mark, Mapping) for mark in marks
            ):
                raise StructuredTableDocxExportError(
                    "table cell rich_text marks are invalid"
                )
            text = str(child.get("text") or "")
            preset_name = str((node.get("attrs") or {}).get("stylePreset") or "body")
            if rich_text_run_renderer is not None:
                rich_text_run_renderer(
                    paragraph,
                    text,
                    marks,
                    _rich_cell_run_preset(
                        role=role,
                        font_size=font_size,
                        preset_name=preset_name,
                        layout=layout,
                    ),
                )
            else:
                _add_rich_cell_run(
                    paragraph,
                    text,
                    marks,
                    role=role,
                    font_size=font_size,
                    preset_name=preset_name,
                    layout=layout,
                )
    _apply_cell_vertical_alignment(
        word_cell,
        role=role,
        is_label=is_label,
        layout=layout,
    )
    _set_cell_shading(
        word_cell,
        _cell_fill(layout, role),
    )
    _apply_cell_margins(word_cell, layout)


def _rich_cell_blocks(
    root: Mapping[str, object],
) -> list[tuple[Mapping[str, object], str, Optional[Mapping[str, object]]]]:
    root_type = str(root.get("type") or "")
    nodes = root.get("content", []) if root_type == "doc" else [root]
    if not isinstance(nodes, list):
        raise StructuredTableDocxExportError(
            "table cell rich_text root content must be a list"
        )
    blocks: list[tuple[Mapping[str, object], str, Optional[Mapping[str, object]]]] = []
    list_sequence = 0
    for node in nodes:
        if not isinstance(node, Mapping):
            raise StructuredTableDocxExportError(
                "table cell rich_text contains an invalid block"
            )
        node_type = str(node.get("type") or "")
        if node_type in {"paragraph", "heading"}:
            blocks.append((node, "", None))
            continue
        if node_type not in {"bulletList", "orderedList"}:
            raise StructuredTableDocxExportError(
                f"unsupported table cell rich_text block: {node_type or '<empty>'}"
            )
        items = node.get("content", [])
        if not isinstance(items, list):
            raise StructuredTableDocxExportError(
                "table cell rich_text list content must be a list"
            )
        list_sequence += 1
        list_attrs = node.get("attrs") or {}
        if not isinstance(list_attrs, Mapping):
            raise StructuredTableDocxExportError(
                "table cell rich_text list attributes are invalid"
            )
        for item in items:
            children = item.get("content", []) if isinstance(item, Mapping) else []
            if not isinstance(children, list):
                raise StructuredTableDocxExportError(
                    "table cell rich_text list item is invalid"
                )
            for child_index, child in enumerate(children):
                if not isinstance(child, Mapping) or child.get("type") not in {
                    "paragraph",
                    "heading",
                }:
                    raise StructuredTableDocxExportError(
                        "nested table cell rich_text lists are not supported"
                    )
                blocks.append(
                    (
                        child,
                        "",
                        {
                            "list_type": node_type,
                            "sequence": list_sequence,
                            "format": str(list_attrs.get("numberingFormat") or ""),
                        }
                        if child_index == 0
                        else None,
                    )
                )
    return blocks


def _format_rich_cell_paragraph(
    paragraph,
    node: Mapping[str, object],
    *,
    role: str,
    font_size: float,
    first_column: bool,
    is_label: bool,
    layout: Mapping[str, object],
) -> None:
    attrs = node.get("attrs") or {}
    if not isinstance(attrs, Mapping):
        raise StructuredTableDocxExportError(
            "table cell rich_text paragraph attributes are invalid"
        )
    _format_cell_paragraph(
        paragraph,
        attrs,
        role=role,
        font_size=font_size,
        first_column=first_column,
        is_label=is_label,
        layout=layout,
    )
    paragraph.paragraph_format.left_indent = Pt(
        font_size * float(attrs.get("leftIndentChars") or 0)
    )
    paragraph.paragraph_format.right_indent = Pt(
        font_size * float(attrs.get("rightIndentChars") or 0)
    )
    paragraph.paragraph_format.first_line_indent = Pt(
        font_size * float(attrs.get("firstLineIndentChars") or 0)
    )


def _add_rich_cell_run(
    paragraph,
    text: str,
    marks: Sequence[Mapping[str, object]],
    *,
    role: str,
    font_size: float,
    preset_name: str,
    layout: Mapping[str, object],
) -> None:
    preset_size, preset_font, preset_bold = _rich_cell_preset_values(
        role=role,
        font_size=font_size,
        preset_name=preset_name,
        layout=layout,
    )
    run = paragraph.add_run(text)
    run.bold = preset_bold or role in {"header", "section"}
    run.font.size = Pt(preset_size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    _set_run_fonts(run, east_asia=preset_font, latin="Times New Roman")
    for mark in marks:
        mark_type = str(mark.get("type") or "")
        attrs = mark.get("attrs") or {}
        if mark_type == "bold":
            run.bold = True
        elif mark_type == "italic":
            run.italic = True
        elif mark_type == "underline":
            run.underline = True
        elif mark_type == "superscript":
            run.font.superscript = True
        elif mark_type == "subscript":
            run.font.subscript = True
        elif mark_type == "textStyle":
            if not isinstance(attrs, Mapping):
                raise StructuredTableDocxExportError(
                    "table cell textStyle attributes are invalid"
                )
            family = str(attrs.get("fontFamily") or "").strip()
            if family:
                _set_run_fonts(
                    run,
                    east_asia="宋体",
                    latin="Times New Roman",
                )
            size = str(attrs.get("fontSize") or "").strip()
            if size:
                match = re.fullmatch(r"([0-9]+(?:\.[0-9]+)?)pt", size)
                if not match:
                    raise StructuredTableDocxExportError(
                        "table cell font size must use pt units"
                    )
                run.font.size = Pt(float(match.group(1)))
            color = str(attrs.get("color") or "").strip().upper()
            if color:
                if not re.fullmatch(r"#[0-9A-F]{6}", color):
                    raise StructuredTableDocxExportError(
                        "table cell text color must be #RRGGBB"
                    )
                run.font.color.rgb = RGBColor.from_string(color[1:])
        elif mark_type == "highlight":
            color = (
                str(attrs.get("color") or "").upper()
                if isinstance(attrs, Mapping)
                else ""
            )
            if color not in _RICH_HIGHLIGHTS:
                raise StructuredTableDocxExportError(
                    "unsupported table cell highlight color"
                )
            run.font.highlight_color = _RICH_HIGHLIGHTS[color]
        else:
            raise StructuredTableDocxExportError(
                f"unsupported table cell rich_text mark: {mark_type or '<empty>'}"
            )


def _rich_cell_run_preset(
    *,
    role: str,
    font_size: float,
    preset_name: str,
    layout: Mapping[str, object],
) -> Mapping[str, object]:
    preset_size, preset_font, preset_bold = _rich_cell_preset_values(
        role=role,
        font_size=font_size,
        preset_name=preset_name,
        layout=layout,
    )
    return {
        "font_size_pt": preset_size,
        "east_asia": preset_font,
        "latin": "Times New Roman",
        "bold": preset_bold or role in {"header", "section"},
    }


def _rich_cell_preset_values(
    *,
    role: str,
    font_size: float,
    preset_name: str,
    layout: Mapping[str, object],
) -> tuple[float, str, bool]:
    if preset_name == "body" and layout.get("use_table_font_size_for_body") is True:
        return font_size, "宋体", role in {"header", "section"}
    return _RICH_PRESETS.get(
        preset_name,
        (font_size, "宋体", role in {"header", "section"}),
    )


def _format_cell_paragraph(
    paragraph,
    attrs: Mapping[str, object],
    *,
    role: str,
    font_size: float,
    first_column: bool,
    is_label: bool,
    layout: Mapping[str, object],
) -> None:
    if "spacingBeforePt" in attrs:
        paragraph.paragraph_format.space_before = Pt(float(attrs["spacingBeforePt"]))
    elif "paragraph_spacing_before_twips" in layout:
        paragraph.paragraph_format.space_before = Twips(
            _layout_nonnegative_int(layout, "paragraph_spacing_before_twips")
        )
    else:
        paragraph.paragraph_format.space_before = Pt(0)

    if "spacingAfterPt" in attrs:
        paragraph.paragraph_format.space_after = Pt(float(attrs["spacingAfterPt"]))
    elif "paragraph_spacing_after_twips" in layout:
        paragraph.paragraph_format.space_after = Twips(
            _layout_nonnegative_int(layout, "paragraph_spacing_after_twips")
        )
    else:
        paragraph.paragraph_format.space_after = Pt(0)

    if "lineHeight" in attrs:
        paragraph.paragraph_format.line_spacing = float(attrs["lineHeight"])
    elif "paragraph_line_twips" in layout:
        paragraph.paragraph_format.line_spacing = Twips(
            _layout_positive_int(layout, "paragraph_line_twips")
        )
    else:
        paragraph.paragraph_format.line_spacing = 1

    alignment = str(attrs.get("textAlign") or "").strip().lower()
    if not alignment:
        configured = layout.get("paragraph_alignment_by_role")
        if configured is not None and not isinstance(configured, Mapping):
            raise StructuredTableDocxExportError(
                "word_layout.paragraph_alignment_by_role must be a mapping"
            )
        alignment_key = "label" if is_label else role
        alignment = str((configured or {}).get(alignment_key) or "").strip().lower()
    if alignment:
        if alignment not in _RICH_ALIGNMENTS:
            raise StructuredTableDocxExportError(
                f"unsupported table cell paragraph alignment: {alignment}"
            )
        paragraph.alignment = _RICH_ALIGNMENTS[alignment]
    else:
        paragraph.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT
            if first_column and role == "body"
            else WD_ALIGN_PARAGRAPH.CENTER
        )


def _layout_nonnegative_int(layout: Mapping[str, object], key: str) -> int:
    value = layout.get(key)
    if isinstance(value, bool) or not isinstance(value, int) or value < 0:
        raise StructuredTableDocxExportError(
            f"word_layout.{key} must be a non-negative integer"
        )
    return value


def _layout_positive_int(layout: Mapping[str, object], key: str) -> int:
    value = layout.get(key)
    if isinstance(value, bool) or not isinstance(value, int) or value <= 0:
        raise StructuredTableDocxExportError(
            f"word_layout.{key} must be a positive integer"
        )
    return value


def _apply_cell_vertical_alignment(
    word_cell,
    *,
    role: str,
    is_label: bool,
    layout: Mapping[str, object],
) -> None:
    configured = layout.get("vertical_alignment_by_role")
    if configured is None:
        word_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        return
    if not isinstance(configured, Mapping):
        raise StructuredTableDocxExportError(
            "word_layout.vertical_alignment_by_role must be a mapping"
        )
    alignment_key = "label" if is_label else role
    value = str(configured.get(alignment_key, "inherit") or "inherit").strip().lower()
    properties = word_cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:vAlign"))
    if value == "inherit":
        if existing is not None:
            properties.remove(existing)
        return
    choices = {
        "top": WD_CELL_VERTICAL_ALIGNMENT.TOP,
        "center": WD_CELL_VERTICAL_ALIGNMENT.CENTER,
        "bottom": WD_CELL_VERTICAL_ALIGNMENT.BOTTOM,
    }
    if value not in choices:
        raise StructuredTableDocxExportError(
            f"unsupported table cell vertical alignment: {value}"
        )
    word_cell.vertical_alignment = choices[value]


def _cell_fill(layout: Mapping[str, object], role: str) -> str:
    configured = layout.get("cell_fill_by_role")
    if configured is None:
        return {"header": "FFFFFF", "section": "EEECE1"}.get(role, "FFFFFF")
    if not isinstance(configured, Mapping):
        raise StructuredTableDocxExportError(
            "word_layout.cell_fill_by_role must be a mapping"
        )
    fill = str(configured.get(role, configured.get("body", "FFFFFF"))).strip().upper()
    if not re.fullmatch(r"[0-9A-F]{6}", fill):
        raise StructuredTableDocxExportError(
            "word_layout.cell_fill_by_role values must be six-digit RGB colors"
        )
    return fill


def _apply_cell_margins(word_cell, layout: Mapping[str, object]) -> None:
    if "cell_margins_twips" not in layout:
        _set_cell_margins(word_cell, top=45, start=55, bottom=45, end=55)
        return
    configured = layout.get("cell_margins_twips")
    properties = word_cell._tc.get_or_add_tcPr()
    existing = properties.find(qn("w:tcMar"))
    if configured is None:
        if existing is not None:
            properties.remove(existing)
        return
    if not isinstance(configured, Mapping):
        raise StructuredTableDocxExportError(
            "word_layout.cell_margins_twips must be a mapping or null"
        )
    values = {}
    for side in ("top", "start", "bottom", "end"):
        value = configured.get(side, 0)
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise StructuredTableDocxExportError(f"invalid table cell {side} margin")
        values[side] = value
    _set_cell_margins(word_cell, **values)


def _apply_word_list_numbering(
    paragraph,
    list_spec: Mapping[str, object],
    *,
    numbering_cache: dict[int, int],
) -> None:
    sequence = list_spec.get("sequence")
    if isinstance(sequence, bool) or not isinstance(sequence, int) or sequence < 1:
        raise StructuredTableDocxExportError("table cell list sequence is invalid")
    list_type = str(list_spec.get("list_type") or "").strip()
    format_name = str(list_spec.get("format") or "").strip()
    if sequence not in numbering_cache:
        numbering_cache[sequence] = _create_word_numbering_instance(
            paragraph,
            list_type=list_type,
            format_name=format_name,
        )
    properties = paragraph._p.get_or_add_pPr()
    existing = properties.find(qn("w:numPr"))
    if existing is not None:
        properties.remove(existing)
    num_properties = OxmlElement("w:numPr")
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    num_id = OxmlElement("w:numId")
    num_id.set(qn("w:val"), str(numbering_cache[sequence]))
    num_properties.extend((level, num_id))
    _insert_before_first(
        properties,
        num_properties,
        (
            "w:suppressLineNumbers",
            "w:pBdr",
            "w:shd",
            "w:tabs",
            "w:suppressAutoHyphens",
            "w:kinsoku",
            "w:wordWrap",
            "w:overflowPunct",
            "w:topLinePunct",
            "w:autoSpaceDE",
            "w:autoSpaceDN",
            "w:bidi",
            "w:adjustRightInd",
            "w:snapToGrid",
            "w:spacing",
            "w:ind",
            "w:contextualSpacing",
            "w:mirrorIndents",
            "w:suppressOverlap",
            "w:jc",
            "w:textDirection",
            "w:textAlignment",
            "w:textboxTightWrap",
            "w:outlineLvl",
            "w:divId",
            "w:cnfStyle",
            "w:rPr",
            "w:sectPr",
            "w:pPrChange",
        ),
    )


def _create_word_numbering_instance(
    paragraph,
    *,
    list_type: str,
    format_name: str,
) -> int:
    definitions = {
        ("bulletList", ""): ("bullet", "•", "tab", 360, 360),
        ("orderedList", ""): ("decimal", "%1.", "space", 420, 420),
        ("orderedList", "decimal_dot"): ("decimal", "%1.", "space", 420, 420),
        ("orderedList", "decimal_half_paren"): (
            "decimal",
            "%1)",
            "space",
            246,
            244,
        ),
        ("orderedList", "decimal_fullwidth_paren"): (
            "decimal",
            "%1）",
            "nothing",
            0,
            0,
        ),
    }
    definition = definitions.get((list_type, format_name))
    if definition is None:
        raise StructuredTableDocxExportError(
            f"unsupported table cell list format: {list_type}/{format_name}"
        )
    num_format, level_text, suffix, left, hanging = definition
    numbering = paragraph.part.numbering_part.element
    abstract_ids = [
        int(item.get(qn("w:abstractNumId")))
        for item in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [int(item.get(qn("w:numId"))) for item in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi_level = OxmlElement("w:multiLevelType")
    multi_level.set(qn("w:val"), "singleLevel")
    abstract.append(multi_level)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    number_format = OxmlElement("w:numFmt")
    number_format.set(qn("w:val"), num_format)
    suffix_element = OxmlElement("w:suff")
    suffix_element.set(qn("w:val"), suffix)
    level_text_element = OxmlElement("w:lvlText")
    level_text_element.set(qn("w:val"), level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.extend(
        (
            start,
            number_format,
            suffix_element,
            level_text_element,
            level_justification,
        )
    )
    if left or hanging:
        paragraph_properties = OxmlElement("w:pPr")
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), str(left))
        indentation.set(qn("w:hanging"), str(hanging))
        paragraph_properties.append(indentation)
        level.append(paragraph_properties)
    if list_type == "bulletList":
        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Symbol")
        fonts.set(qn("w:hAnsi"), "Symbol")
        run_properties.append(fonts)
        level.append(run_properties)
    abstract.append(level)
    first_num = numbering.find(qn("w:num"))
    if first_num is None:
        numbering.append(abstract)
    else:
        numbering.insert(numbering.index(first_num), abstract)

    instance = OxmlElement("w:num")
    instance.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    instance.append(abstract_reference)
    numbering.append(instance)
    return num_id


def _ordered_notes(
    notes: Sequence[StructuredTableNote],
) -> List[Tuple[str, StructuredTableNote]]:
    indexed = list(enumerate(notes))
    ordered = sorted(
        indexed,
        key=lambda item: (_marker_sort_key(item[1].marker), item[0]),
    )
    return [
        (note.marker.strip() or str(index + 1), note)
        for index, (_, note) in enumerate(ordered)
    ]


def _marker_sort_key(marker: str) -> Tuple[object, ...]:
    value = marker.strip()
    if not value:
        return (3,)
    if value.isdigit():
        try:
            number = int(value)
        except ValueError:
            try:
                number = int(
                    "".join(str(unicodedata.digit(character)) for character in value)
                )
            except (TypeError, ValueError):
                number = None
        if number is not None:
            return (0, number)
    if re.fullmatch(r"[A-Za-z]", value):
        return (1, value.lower())
    parts = re.split(r"(\d+)", value.casefold())
    natural = tuple(int(part) if part.isdigit() else part for part in parts)
    return (2, *natural)


def _repeat_header_rows(table: StructuredTable, layout: Mapping[str, object]) -> int:
    configured = layout.get("repeat_header_rows", table.header_row_count)
    if isinstance(configured, bool):
        count = table.header_row_count if configured else 0
    elif isinstance(configured, int):
        count = configured
    else:
        raise StructuredTableDocxExportError(
            "word_layout.repeat_header_rows must be a boolean or integer"
        )
    if count < 0 or count > len(table.rows):
        raise StructuredTableDocxExportError(
            "repeat header row count exceeds table rows"
        )
    return count


def _margins(layout: Mapping[str, object]) -> Dict[str, int]:
    default = {"left": 720, "right": 720, "top": 720, "bottom": 720}
    configured = layout.get("margins_twips", {})
    if configured is None:
        return default
    if not isinstance(configured, Mapping):
        raise StructuredTableDocxExportError(
            "word_layout.margins_twips must be a mapping"
        )
    result = dict(default)
    for side in result:
        if side not in configured:
            continue
        value = configured[side]
        if isinstance(value, bool) or not isinstance(value, int) or value < 0:
            raise StructuredTableDocxExportError(f"invalid {side} page margin")
        result[side] = value
    return result


def _table_width_spec(
    layout: Mapping[str, object],
    resolved_width_twips: int,
) -> tuple[str, int]:
    width_type = str(layout.get("table_width_type") or "dxa").strip().lower()
    if width_type not in {"dxa", "pct"}:
        raise StructuredTableDocxExportError(
            "word_layout.table_width_type must be dxa or pct"
        )
    if width_type == "dxa":
        return width_type, resolved_width_twips
    value = layout.get("table_width_value", 5000)
    if isinstance(value, bool) or not isinstance(value, int) or not 1 <= value <= 5000:
        raise StructuredTableDocxExportError(
            "word_layout.table_width_value must be an integer between 1 and 5000"
        )
    return width_type, value


def _row_min_heights(
    layout: Mapping[str, object],
    row_count: int,
) -> Dict[int, int]:
    configured = layout.get("row_min_heights_twips", {})
    if configured is None:
        return {}
    if not isinstance(configured, Mapping):
        raise StructuredTableDocxExportError(
            "word_layout.row_min_heights_twips must be a mapping"
        )
    result: Dict[int, int] = {}
    for raw_index, raw_height in configured.items():
        try:
            index = int(raw_index)
        except (TypeError, ValueError) as exc:
            raise StructuredTableDocxExportError(
                "word_layout row height index must be an integer"
            ) from exc
        if index < 0 or index >= row_count:
            raise StructuredTableDocxExportError(
                "word_layout row height index exceeds table rows"
            )
        if (
            isinstance(raw_height, bool)
            or not isinstance(raw_height, int)
            or raw_height <= 0
        ):
            raise StructuredTableDocxExportError(
                "word_layout row height must be a positive integer"
            )
        result[index] = raw_height
    return result


def _column_widths(
    columns, available_width: int, layout: Mapping[str, object]
) -> List[int]:
    specified = [column.width_twips for column in columns]
    missing_count = sum(width is None for width in specified)
    specified_total = sum(width or 0 for width in specified)
    widths: List[int]
    if missing_count:
        remaining = available_width - specified_total
        fallback = max(
            _MIN_COLUMN_WIDTH_TWIPS,
            remaining // missing_count
            if remaining > 0
            else available_width // len(columns),
        )
        widths = [int(width or fallback) for width in specified]
    else:
        widths = [int(width) for width in specified if width is not None]

    fit_to_page = layout.get("fit_to_page", True)
    if not isinstance(fit_to_page, bool):
        raise StructuredTableDocxExportError("word_layout.fit_to_page must be boolean")
    if fit_to_page and sum(widths) > available_width:
        widths = _scale_widths(widths, available_width)
    if any(width < _MIN_COLUMN_WIDTH_TWIPS for width in widths):
        raise StructuredTableDocxExportError("resolved column width is too small")
    return widths


def _scale_widths(widths: Sequence[int], target: int) -> List[int]:
    total = sum(widths)
    scaled = [
        max(_MIN_COLUMN_WIDTH_TWIPS, int(width * target / total)) for width in widths
    ]
    difference = target - sum(scaled)
    cursor = len(scaled) - 1
    while difference and cursor >= 0:
        candidate = scaled[cursor] + difference
        if candidate >= _MIN_COLUMN_WIDTH_TWIPS:
            scaled[cursor] = candidate
            difference = 0
        else:
            difference += scaled[cursor] - _MIN_COLUMN_WIDTH_TWIPS
            scaled[cursor] = _MIN_COLUMN_WIDTH_TWIPS
            cursor -= 1
    if difference:
        raise StructuredTableDocxExportError(
            "column widths cannot fit within page margins"
        )
    return scaled


def _page_dimensions(orientation: str) -> Tuple[int, int]:
    portrait_width = int(Mm(210).twips)
    portrait_height = int(Mm(297).twips)
    if orientation == "landscape":
        return portrait_height, portrait_width
    return portrait_width, portrait_height


def _font_size(table: StructuredTable, column_count: int) -> float:
    configured = table.word_layout.get("font_size_pt")
    if configured is None:
        return 7.0 if column_count >= 12 else 9.0
    if isinstance(configured, bool) or not isinstance(configured, (int, float)):
        raise StructuredTableDocxExportError("word_layout.font_size_pt must be numeric")
    if configured < 5 or configured > 14:
        raise StructuredTableDocxExportError(
            "word_layout.font_size_pt must be between 5 and 14"
        )
    return float(configured)


def _require_unique_orders(values: Sequence[int], label: str) -> None:
    if len(values) != len(set(values)):
        raise StructuredTableDocxExportError(
            f"structured table {label} orders must be unique"
        )


def _normalized_role(role: str) -> str:
    normalized = str(role or "body").strip().lower()
    if normalized in {"header", "table_header", "column_header"}:
        return "header"
    if normalized in {"section", "section_header", "group", "group_header"}:
        return "section"
    return "body"


def _merge_range(placed: _PlacedCell) -> Dict[str, object]:
    return {
        "cell_id": placed.cell.cell_id,
        "start_row": placed.row_index,
        "start_column": placed.column_index,
        "end_row": placed.row_index + placed.cell.row_span - 1,
        "end_column": placed.column_index + placed.cell.column_span - 1,
    }


def _set_run_fonts(run, east_asia: str, latin: str) -> None:
    run.font.name = latin
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:eastAsia"), east_asia)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:cs"), latin)


def _set_fixed_table_layout(
    word_table,
    width_value: int,
    *,
    width_type: str = "dxa",
) -> None:
    properties = word_table._tbl.tblPr
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        _insert_before_first(
            properties,
            layout,
            (
                "w:tblCellMar",
                "w:tblLook",
                "w:tblCaption",
                "w:tblDescription",
                "w:tblPrChange",
            ),
        )
    layout.set(qn("w:type"), "fixed")
    width = properties.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        _insert_before_first(
            properties,
            width,
            (
                "w:jc",
                "w:tblCellSpacing",
                "w:tblInd",
                "w:tblBorders",
                "w:shd",
                "w:tblLayout",
                "w:tblCellMar",
                "w:tblLook",
                "w:tblCaption",
                "w:tblDescription",
                "w:tblPrChange",
            ),
        )
    width.set(qn("w:type"), width_type)
    width.set(qn("w:w"), str(width_value))


def _set_grid_column_widths(word_table, widths: Sequence[int]) -> None:
    grid_columns = word_table._tbl.tblGrid.findall(qn("w:gridCol"))
    for grid_column, width in zip(grid_columns, widths):
        grid_column.set(qn("w:w"), str(width))


def _set_cell_width(
    word_cell,
    width_value: int,
    *,
    width_type: str = "dxa",
) -> None:
    properties = word_cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        _insert_before_first(
            properties,
            width,
            (
                "w:gridSpan",
                "w:hMerge",
                "w:vMerge",
                "w:tcBorders",
                "w:shd",
                "w:noWrap",
                "w:tcMar",
                "w:textDirection",
                "w:tcFitText",
                "w:vAlign",
                "w:hideMark",
                "w:headers",
                "w:cellIns",
                "w:cellDel",
                "w:cellMerge",
                "w:tcPrChange",
            ),
        )
    width.set(qn("w:type"), width_type)
    width.set(qn("w:w"), str(width_value))


def _set_repeat_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    header = properties.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        _insert_before_first(
            properties,
            header,
            (
                "w:tblCellSpacing",
                "w:jc",
                "w:hidden",
                "w:ins",
                "w:del",
                "w:trPrChange",
            ),
        )
    # tblHeader and cantSplit use the presence-only CT_OnOff type.
    header.attrib.pop(qn("w:val"), None)


def _prevent_row_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = properties.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        _insert_before_first(
            properties,
            cant_split,
            (
                "w:trHeight",
                "w:tblHeader",
                "w:tblCellSpacing",
                "w:jc",
                "w:hidden",
                "w:ins",
                "w:del",
                "w:trPrChange",
            ),
        )
    cant_split.attrib.pop(qn("w:val"), None)


def _set_row_min_height(row, height_twips: int) -> None:
    properties = row._tr.get_or_add_trPr()
    height = properties.find(qn("w:trHeight"))
    if height is None:
        height = OxmlElement("w:trHeight")
        _insert_before_first(
            properties,
            height,
            (
                "w:tblHeader",
                "w:tblCellSpacing",
                "w:jc",
                "w:hidden",
                "w:ins",
                "w:del",
                "w:trPrChange",
            ),
        )
    height.set(qn("w:val"), str(height_twips))
    height.attrib.pop(qn("w:hRule"), None)


def _set_cell_shading(word_cell, fill: str) -> None:
    properties = word_cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        _insert_before_first(
            properties,
            shading,
            (
                "w:noWrap",
                "w:tcMar",
                "w:textDirection",
                "w:tcFitText",
                "w:vAlign",
                "w:hideMark",
                "w:headers",
                "w:cellIns",
                "w:cellDel",
                "w:cellMerge",
                "w:tcPrChange",
            ),
        )
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)


def _set_cell_margins(word_cell, top: int, start: int, bottom: int, end: int) -> None:
    properties = word_cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        _insert_before_first(
            properties,
            margins,
            (
                "w:textDirection",
                "w:tcFitText",
                "w:vAlign",
                "w:hideMark",
                "w:headers",
                "w:cellIns",
                "w:cellDel",
                "w:cellMerge",
                "w:tcPrChange",
            ),
        )
    for name, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        margin = margins.find(qn(f"w:{name}"))
        if margin is None:
            margin = OxmlElement(f"w:{name}")
            margins.append(margin)
        margin.set(qn("w:w"), str(value))
        margin.set(qn("w:type"), "dxa")


def _insert_before_first(parent, element, following_tags: Sequence[str]) -> None:
    following = {qn(tag) for tag in following_tags}
    for index, child in enumerate(parent):
        if child.tag in following:
            parent.insert(index, element)
            return
    parent.append(element)


def _normalize_docx_package(content: bytes) -> bytes:
    source = io.BytesIO(content)
    output = io.BytesIO()
    with (
        zipfile.ZipFile(source, "r") as archive,
        zipfile.ZipFile(
            output,
            "w",
            compression=zipfile.ZIP_DEFLATED,
            compresslevel=9,
        ) as normalized,
    ):
        for source_info in sorted(archive.infolist(), key=lambda item: item.filename):
            info = zipfile.ZipInfo(source_info.filename, _FIXED_PACKAGE_TIMESTAMP)
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = source_info.external_attr
            info.create_system = source_info.create_system
            payload = archive.read(source_info.filename)
            if source_info.filename.startswith(
                "word/"
            ) and source_info.filename.endswith(".xml"):
                payload = _normalize_word_xml_fonts(payload)
            elif source_info.filename.startswith(
                "word/media/"
            ) and source_info.filename.lower().endswith(".svg"):
                payload = _normalize_svg_fonts(payload)
            normalized.writestr(info, payload)
    return output.getvalue()


def _normalize_word_xml_fonts(payload: bytes) -> bytes:
    """Apply the production bilingual font contract without changing content."""

    if (
        _WORDPROCESSINGML_NS.encode() not in payload
        and _DRAWINGML_NS.encode() not in payload
    ):
        return payload
    try:
        root = etree.fromstring(payload)
    except etree.XMLSyntaxError as exc:
        raise StructuredTableDocxExportError(
            "DOCX contains malformed Word XML during font normalization"
        ) from exc
    namespaces = {
        "w": _WORDPROCESSINGML_NS,
        "a": _DRAWINGML_NS,
    }
    w_rpr = f"{{{_WORDPROCESSINGML_NS}}}rPr"
    w_rfonts = f"{{{_WORDPROCESSINGML_NS}}}rFonts"

    for run in root.xpath(".//w:r", namespaces=namespaces):
        run_properties = run.find(w_rpr)
        if run_properties is None:
            run_properties = etree.Element(w_rpr)
            run.insert(0, run_properties)
        fonts = run_properties.find(w_rfonts)
        if fonts is None:
            fonts = etree.Element(w_rfonts)
            run_properties.insert(0, fonts)

    for run_properties in root.xpath(".//w:rPr", namespaces=namespaces):
        fonts = run_properties.find(w_rfonts)
        if fonts is None:
            fonts = etree.Element(w_rfonts)
            run_properties.insert(0, fonts)
        fonts.set(f"{{{_WORDPROCESSINGML_NS}}}eastAsia", _EXPORT_EAST_ASIA_FONT)
        fonts.set(f"{{{_WORDPROCESSINGML_NS}}}ascii", _EXPORT_LATIN_FONT)
        fonts.set(f"{{{_WORDPROCESSINGML_NS}}}hAnsi", _EXPORT_LATIN_FONT)
        fonts.set(f"{{{_WORDPROCESSINGML_NS}}}cs", _EXPORT_LATIN_FONT)

    for element in root.xpath(".//a:latin", namespaces=namespaces):
        element.set("typeface", _EXPORT_LATIN_FONT)
    for element in root.xpath(".//a:ea", namespaces=namespaces):
        element.set("typeface", _EXPORT_EAST_ASIA_FONT)
    for element in root.xpath(".//a:cs", namespaces=namespaces):
        element.set("typeface", _EXPORT_LATIN_FONT)

    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=payload.lstrip().startswith(b"<?xml"),
        standalone=True,
    )


def _normalize_svg_fonts(payload: bytes) -> bytes:
    try:
        root = etree.fromstring(payload)
    except etree.XMLSyntaxError as exc:
        raise StructuredTableDocxExportError(
            "DOCX contains malformed SVG during font normalization"
        ) from exc
    for element in root.xpath(".//*[local-name()='text']"):
        element.set("font-family", _EXPORT_SVG_FONT_STACK)
    return etree.tostring(
        root,
        encoding="UTF-8",
        xml_declaration=payload.lstrip().startswith(b"<?xml"),
    )
