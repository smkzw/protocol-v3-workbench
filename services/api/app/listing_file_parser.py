from __future__ import annotations

import csv
import io
import re
import zipfile
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Set, Tuple

import openpyxl
import xlrd
from openpyxl.worksheet.filters import AutoFilter

from packages.contracts.workbench_contracts import ListingSheetPayload


SUPPORTED_EXTENSIONS = {".csv", ".xls", ".xlsx", ".xlsm"}
LISTING_PARSER_VERSION = "listing_file_parser_v3_ooxml_metadata_recovery"
INVALID_AUTOFILTER_WARNING = (
    "worksheet_autofilter_ignored: invalid range metadata ignored; cell values preserved"
)
AUTOFILTER_ELEMENT_PATTERN = re.compile(
    rb"<autoFilter\b[^>]*/>|<autoFilter\b[^>]*>.*?</autoFilter>",
    re.DOTALL,
)
AUTOFILTER_REF_PATTERN = re.compile(rb"""\bref=(["'])(.*?)\1""")

COMMON_CHINESE_HEADER_ALIASES = {
    "项目编号": "STUDYID",
    "表单编号": "FORMOID",
    "受试者编号": "SUBJID",
    "受试者筛选号": "SUBJID",
    "受试者": "SUBJID",
    "姓名缩写": "SUBJINI",
    "受试者状态": "SUBJSTA",
    "试验中心编号": "SITEID",
    "中心编号": "SITEID",
    "试验中心名称": "SITENM",
    "研究中心": "SITENM",
    "中心名称": "SITENM",
    "数据节": "VISIT",
    "访视名称": "VISIT",
    "访视号": "VISTREP",
    "Instance顺序号": "VISTREP",
    "数据块": "FORMNM",
    "数据页": "FORMNM",
    "页面名称": "FORMNM",
    "页面号": "FORMREP",
    "最后修改时间": "PAGELMDT",
    "页面最近修改时间": "PAGELMDT",
    "行号": "RECREP",
    "记录号": "RECREP",
    "随机号": "RANDNO",
    "随机时间": "RANDDTC",
    "研究分组": "ARM",
    "性别": "SEX",
}

HEADER_CODE_PATTERN = re.compile(r"[（(]([A-Za-z][A-Za-z0-9_]{1,31})[）)]")
ASCII_CODE_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9_]{1,31}$")


def _clean_header(value: Any, index: int) -> str:
    text = "" if value is None else str(value).strip()
    if not text:
        return f"UNNAMED_{index + 1}"
    if text in COMMON_CHINESE_HEADER_ALIASES:
        return COMMON_CHINESE_HEADER_ALIASES[text]
    code_match = HEADER_CODE_PATTERN.search(text)
    if code_match:
        return code_match.group(1).upper()
    return text


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _dedupe_headers(headers: List[str]) -> List[str]:
    seen: Dict[str, int] = {}
    deduped: List[str] = []
    for header in headers:
        seen[header] = seen.get(header, 0) + 1
        deduped.append(header if seen[header] == 1 else f"{header}__{seen[header]}")
    return deduped


def _is_ascii_code(value: Any) -> bool:
    text = "" if value is None else str(value).strip()
    return bool(ASCII_CODE_PATTERN.match(text)) and text == text.upper()


def _non_empty_count(values: Any) -> int:
    return sum(1 for value in values if _cell_text(value))


def _looks_like_variable_header(values: Any) -> bool:
    non_empty = _non_empty_count(values)
    if non_empty < 3:
        return False
    ascii_codes = sum(1 for value in values if _is_ascii_code(value))
    return ascii_codes >= 3 and ascii_codes / max(non_empty, 1) >= 0.6


def _looks_like_header(values: Any) -> bool:
    headers = [_cell_text(value) for value in values]
    if _non_empty_count(headers) < 2:
        return False
    markers = {
        "受试者",
        "受试者编号",
        "受试者筛选号",
        "中心编号",
        "研究中心",
        "试验中心编号",
        "项目编号",
        "表单编号",
        "访视名称",
        "数据节",
        "SUBJID",
        "SITEID",
        "VISIT",
    }
    return any(header in markers or HEADER_CODE_PATTERN.search(header) for header in headers)


def _header_and_data_rows(raw_rows: List[List[Any]]) -> tuple[List[str], List[List[Any]], int]:
    header_index = None
    scan_limit = min(25, len(raw_rows))
    for index in range(scan_limit):
        if _looks_like_header(raw_rows[index]):
            header_index = index
            break
    if header_index is None:
        header_index = 0

    header_values = raw_rows[header_index]
    data_start = header_index + 1
    if data_start < len(raw_rows) and _looks_like_variable_header(raw_rows[data_start]):
        header_values = [
            raw_rows[data_start][index] if index < len(raw_rows[data_start]) and _cell_text(raw_rows[data_start][index]) else value
            for index, value in enumerate(header_values)
        ]
        data_start += 1

    headers = _dedupe_headers([_clean_header(value, index) for index, value in enumerate(header_values)])
    return headers, raw_rows[data_start:], data_start


def _build_rows(
    headers: List[str],
    data_rows: List[List[Any]],
    data_start: int,
) -> tuple[List[Dict[str, str]], List[int]]:
    rows: List[Dict[str, str]] = []
    row_numbers: List[int] = []
    for offset, values in enumerate(data_rows):
        row = {
            headers[index]: _cell_text(values[index] if index < len(values) else "")
            for index in range(len(headers))
        }
        if any(value for value in row.values()):
            rows.append(row)
            row_numbers.append(data_start + offset + 1)
    return rows, row_numbers


def _has_usable_headers(headers: List[str]) -> bool:
    return sum(1 for header in headers if not header.startswith("UNNAMED_")) >= 2


def parse_listing_file(filename: str, content: bytes) -> List[ListingSheetPayload]:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported listing file type: {suffix or 'unknown'}")
    if suffix == ".csv":
        return [_parse_csv_listing(filename, content)]
    if suffix == ".xls":
        return _parse_xls_listing(content)
    return _parse_xlsx_listing(content)


def _parse_csv_listing(filename: str, content: bytes) -> ListingSheetPayload:
    text = content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(text))
    raw_rows = [list(row) for row in reader]
    if not raw_rows:
        raise ValueError("csv listing has no header row")
    headers, data_rows, data_start = _header_and_data_rows(raw_rows)
    rows, row_numbers = _build_rows(headers, data_rows, data_start)
    sheet_name = Path(filename or "CSV").stem or "CSV"
    return ListingSheetPayload(
        sheet_name=sheet_name,
        headers=headers,
        rows=rows,
        row_numbers=row_numbers,
    )


def _parse_xlsx_listing(content: bytes) -> List[ListingSheetPayload]:
    readable_content, invalid_autofilter_paths = _ignore_invalid_autofilter_metadata(content)
    workbook = openpyxl.load_workbook(
        io.BytesIO(readable_content),
        read_only=True,
        data_only=True,
    )
    try:
        sheets: List[ListingSheetPayload] = []
        for worksheet in workbook.worksheets:
            parser_warnings: List[str] = []
            if worksheet._worksheet_path in invalid_autofilter_paths:
                parser_warnings.append(INVALID_AUTOFILTER_WARNING)
            if (
                worksheet.calculate_dimension() in {"A1", "A1:A1"}
                and hasattr(worksheet, "reset_dimensions")
            ):
                worksheet.reset_dimensions()
                parser_warnings.append(
                    "worksheet_dimension_reset: declared A1:A1; parsed from worksheet XML bounds"
                )
            raw_rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
            if not raw_rows:
                continue
            headers, data_rows, data_start = _header_and_data_rows(raw_rows)
            rows, row_numbers = _build_rows(headers, data_rows, data_start)
            if rows or _has_usable_headers(headers):
                sheets.append(
                    ListingSheetPayload(
                        sheet_name=worksheet.title,
                        headers=headers,
                        rows=rows,
                        row_numbers=row_numbers,
                        parser_warnings=parser_warnings,
                    )
                )
        if not sheets:
            raise ValueError("xlsx listing has no non-empty sheets")
        return sheets
    finally:
        workbook.close()


def _ignore_invalid_autofilter_metadata(content: bytes) -> Tuple[bytes, Set[str]]:
    """Remove unreadable worksheet filter metadata from an in-memory OOXML copy."""

    affected_paths: Set[str] = set()
    output = io.BytesIO()
    changed = False
    with zipfile.ZipFile(io.BytesIO(content), "r") as source:
        with zipfile.ZipFile(output, "w") as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename.startswith("xl/worksheets/") and entry.filename.endswith(".xml"):
                    worksheet_changed = False

                    def replace_autofilter(match: re.Match[bytes]) -> bytes:
                        nonlocal worksheet_changed
                        ref_match = AUTOFILTER_REF_PATTERN.search(match.group(0))
                        if ref_match is None:
                            return match.group(0)
                        ref = ref_match.group(2).decode("utf-8", errors="replace")
                        try:
                            AutoFilter(ref=ref)
                        except ValueError:
                            worksheet_changed = True
                            return b""
                        return match.group(0)

                    payload = AUTOFILTER_ELEMENT_PATTERN.sub(replace_autofilter, payload)
                    if worksheet_changed:
                        changed = True
                        affected_paths.add(entry.filename)
                target.writestr(entry, payload)
    return (output.getvalue(), affected_paths) if changed else (content, affected_paths)


def _parse_xls_listing(content: bytes) -> List[ListingSheetPayload]:
    workbook = xlrd.open_workbook(file_contents=content, on_demand=True)
    try:
        sheets: List[ListingSheetPayload] = []
        for sheet_name in workbook.sheet_names():
            worksheet = workbook.sheet_by_name(sheet_name)
            raw_rows = [
                [_xls_cell_text(workbook, worksheet, row_index, col_index) for col_index in range(worksheet.ncols)]
                for row_index in range(worksheet.nrows)
            ]
            if not raw_rows:
                continue
            headers, data_rows, data_start = _header_and_data_rows(raw_rows)
            rows, row_numbers = _build_rows(headers, data_rows, data_start)
            if rows or _has_usable_headers(headers):
                sheets.append(
                    ListingSheetPayload(
                        sheet_name=sheet_name,
                        headers=headers,
                        rows=rows,
                        row_numbers=row_numbers,
                    )
                )
        if not sheets:
            raise ValueError("xls listing has no non-empty sheets")
        return sheets
    finally:
        workbook.release_resources()


def _xls_cell_text(workbook: xlrd.book.Book, worksheet: xlrd.sheet.Sheet, row_index: int, col_index: int) -> str:
    cell = worksheet.cell(row_index, col_index)
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            dt = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode)
            if dt.hour == 0 and dt.minute == 0 and dt.second == 0:
                return dt.date().isoformat()
            return dt.strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return str(cell.value).strip()
    return _cell_text(cell.value)
