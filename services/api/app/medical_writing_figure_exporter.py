from __future__ import annotations

import io
import re
import unicodedata
from typing import Optional, Union
from xml.etree import ElementTree

from docx.document import Document as DocxDocument
from docx.image.image import Image
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shape import InlineShape
from docx.shared import Length
from docx.text.paragraph import Paragraph
from lxml import etree


SVG_CONTENT_TYPE = "image/svg+xml"
SVG_NAMESPACE = "http://www.w3.org/2000/svg"
ASVG_NAMESPACE = "http://schemas.microsoft.com/office/drawing/2016/SVG/main"
ASVG_EXTENSION_URI = "{96DAC541-7B7A-43D3-8B79-37D633B846F1}"

_FORBIDDEN_SVG_ELEMENTS = frozenset({"foreignobject", "script"})
_DTD_OR_ENTITY_RE = re.compile(r"<!\s*(?:DOCTYPE|ENTITY)\b", re.IGNORECASE)
_CSS_IMPORT_RE = re.compile(r"@import\b", re.IGNORECASE)
_CSS_URL_RE = re.compile(r"url\(\s*([^)]*?)\s*\)", re.IGNORECASE)


class MedicalWritingFigureExportError(ValueError):
    """Raised when a medical-writing figure cannot be embedded safely."""


def add_svg_figure_with_png_fallback(
    target: Union[DocxDocument, Paragraph],
    *,
    svg_bytes: bytes,
    png_bytes: bytes,
    figure_id: str,
    alt_text: str,
    width: Optional[Union[int, Length]] = None,
) -> InlineShape:
    """Add an embedded SVG figure with a PNG fallback to a DOCX story.

    ``target`` may be a document, in which case a paragraph is appended, or an
    existing paragraph. ``width`` uses python-docx EMU/Length units. All caller
    input is validated before the document is changed; unsafe SVG never degrades
    to a PNG-only figure.
    """

    story_part = _story_part_for(target)
    safe_figure_id = _validated_text(figure_id, field="figure_id", limit=255)
    safe_alt_text = _validated_text(alt_text, field="alt_text", limit=1024)
    safe_width = _validated_width(width)
    safe_svg = _validated_svg(svg_bytes)
    safe_png = _validated_png(png_bytes)

    existing_relationship_ids = set(story_part.rels)
    created_paragraph = None
    created_run = None
    try:
        if isinstance(target, DocxDocument):
            paragraph = target.add_paragraph()
            created_paragraph = paragraph
        else:
            paragraph = target

        created_run = paragraph.add_run()
        inline_shape = created_run.add_picture(io.BytesIO(safe_png), width=safe_width)

        svg_partname = story_part.package.next_partname("/word/media/image%d.svg")
        svg_part = Part(
            svg_partname,
            SVG_CONTENT_TYPE,
            safe_svg,
            story_part.package,
        )
        svg_relationship_id = story_part.relate_to(svg_part, RT.IMAGE)
        _apply_svg_extension(
            inline_shape,
            svg_relationship_id=svg_relationship_id,
            figure_id=safe_figure_id,
            alt_text=safe_alt_text,
        )
        return inline_shape
    except Exception as exc:
        _rollback_insertion(
            story_part=story_part,
            existing_relationship_ids=existing_relationship_ids,
            paragraph=created_paragraph,
            run=created_run,
        )
        raise MedicalWritingFigureExportError(
            "failed to embed the SVG figure and PNG fallback"
        ) from exc


def _story_part_for(target: Union[DocxDocument, Paragraph]):
    if isinstance(target, DocxDocument):
        story_part = target.part
    elif isinstance(target, Paragraph):
        story_part = target.part
    else:
        raise MedicalWritingFigureExportError(
            "target must be a python-docx Document or Paragraph"
        )
    if story_part.package is None:
        raise MedicalWritingFigureExportError("target is not attached to a DOCX package")
    return story_part


def _validated_text(value: str, *, field: str, limit: int) -> str:
    if not isinstance(value, str):
        raise MedicalWritingFigureExportError(f"{field} must be text")
    normalized = unicodedata.normalize("NFC", value.strip())
    if not normalized:
        raise MedicalWritingFigureExportError(f"{field} must not be empty")
    if len(normalized) > limit:
        raise MedicalWritingFigureExportError(
            f"{field} must contain no more than {limit} characters"
        )
    if any(unicodedata.category(character) in {"Cc", "Cs"} for character in normalized):
        raise MedicalWritingFigureExportError(
            f"{field} contains unsupported control characters"
        )
    return normalized


def _validated_width(width: Optional[Union[int, Length]]) -> Optional[int]:
    if width is None:
        return None
    if isinstance(width, bool) or not isinstance(width, int) or width <= 0:
        raise MedicalWritingFigureExportError(
            "width must be a positive python-docx Length/EMU value"
        )
    return int(width)


def _validated_png(png_bytes: bytes) -> bytes:
    if not isinstance(png_bytes, bytes) or not png_bytes:
        raise MedicalWritingFigureExportError("png_bytes must be non-empty bytes")
    try:
        image = Image.from_blob(png_bytes)
    except Exception as exc:
        raise MedicalWritingFigureExportError("png_bytes is not a readable PNG") from exc
    if image.content_type != "image/png":
        raise MedicalWritingFigureExportError("png_bytes must contain a PNG image")
    return png_bytes


def _validated_svg(svg_bytes: bytes) -> bytes:
    if not isinstance(svg_bytes, bytes) or not svg_bytes:
        raise MedicalWritingFigureExportError("svg_bytes must be non-empty bytes")
    try:
        svg_text = svg_bytes.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise MedicalWritingFigureExportError("svg_bytes must be UTF-8 encoded") from exc
    if _DTD_OR_ENTITY_RE.search(svg_text):
        raise MedicalWritingFigureExportError("SVG document types and entities are forbidden")
    try:
        root = ElementTree.fromstring(svg_bytes)
    except ElementTree.ParseError as exc:
        raise MedicalWritingFigureExportError("svg_bytes is not well-formed XML") from exc

    namespace, local_name = _expanded_name(root.tag)
    if namespace != SVG_NAMESPACE or local_name.casefold() != "svg":
        raise MedicalWritingFigureExportError("SVG root must use the standard SVG namespace")

    for element in root.iter():
        _, element_name = _expanded_name(element.tag)
        element_name = element_name.casefold()
        if element_name in _FORBIDDEN_SVG_ELEMENTS:
            raise MedicalWritingFigureExportError(
                f"SVG element '{element_name}' is forbidden"
            )

        for attribute_name, raw_value in element.attrib.items():
            _, local_attribute = _expanded_name(attribute_name)
            local_attribute = local_attribute.casefold()
            value = raw_value.strip()
            if local_attribute.startswith("on"):
                raise MedicalWritingFigureExportError("SVG event handlers are forbidden")
            if local_attribute == "base":
                raise MedicalWritingFigureExportError("SVG base URIs are forbidden")
            if local_attribute in {"href", "src"} and not _is_local_fragment(value):
                raise MedicalWritingFigureExportError("external SVG references are forbidden")
            if _contains_nonlocal_css_url(value):
                raise MedicalWritingFigureExportError("external SVG references are forbidden")

        if element_name == "style":
            style_text = "".join(element.itertext())
            if _CSS_IMPORT_RE.search(style_text) or _contains_nonlocal_css_url(style_text):
                raise MedicalWritingFigureExportError("external SVG styles are forbidden")
    return svg_bytes


def _expanded_name(name: str) -> tuple[str, str]:
    if name.startswith("{") and "}" in name:
        namespace, local_name = name[1:].split("}", 1)
        return namespace, local_name
    return "", name


def _is_local_fragment(value: str) -> bool:
    return len(value) > 1 and value.startswith("#")


def _contains_nonlocal_css_url(value: str) -> bool:
    for match in _CSS_URL_RE.finditer(value):
        reference = match.group(1).strip().strip("\"'").strip()
        if not _is_local_fragment(reference):
            return True
    return False


def _apply_svg_extension(
    inline_shape: InlineShape,
    *,
    svg_relationship_id: str,
    figure_id: str,
    alt_text: str,
) -> None:
    inline = inline_shape._inline
    inline.docPr.set("name", figure_id)
    inline.docPr.set("title", figure_id)
    inline.docPr.set("descr", alt_text)

    non_visual_properties = inline.xpath(".//pic:cNvPr")[0]
    non_visual_properties.set("name", figure_id)
    non_visual_properties.set("title", figure_id)
    non_visual_properties.set("descr", alt_text)

    fallback_blip = inline.xpath(".//a:blip")[0]
    extension_list = OxmlElement("a:extLst")
    extension = OxmlElement("a:ext")
    extension.set("uri", ASVG_EXTENSION_URI)
    svg_blip = etree.Element(
        etree.QName(ASVG_NAMESPACE, "svgBlip"),
        nsmap={"asvg": ASVG_NAMESPACE},
    )
    svg_blip.set(qn("r:embed"), svg_relationship_id)
    extension.append(svg_blip)
    extension_list.append(extension)
    fallback_blip.append(extension_list)


def _rollback_insertion(
    *,
    story_part,
    existing_relationship_ids: set[str],
    paragraph: Optional[Paragraph],
    run,
) -> None:
    if paragraph is not None:
        paragraph._element.getparent().remove(paragraph._element)
    elif run is not None and run._element.getparent() is not None:
        run._element.getparent().remove(run._element)
    for relationship_id in set(story_part.rels) - existing_relationship_ids:
        story_part.drop_rel(relationship_id)
