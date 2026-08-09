from __future__ import annotations

import base64
import binascii
import copy
import difflib
import hashlib
import io
import json
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import date
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Mapping, Optional, Sequence
from urllib.parse import urlsplit, urlunsplit

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from lxml import etree

from packages.contracts.workbench_contracts import (
    ApprovalState,
    MedicalWritingDocumentPreview,
    MedicalWritingLiteratureLibrary,
    MedicalWritingPreviewBlock,
    MedicalWritingPreviewPage,
    MedicalWritingProjectReference,
    MedicalWritingWordVerificationReceipt,
    ProtocolDocument,
    StructuredTable,
    StructuredTableCell,
    StructuredTableDomain,
    StructuredTableNote,
)

from . import medical_writing_table_exporter as table_docx
from .medical_writing_greenfield import (
    _greenfield_document_control_table,
    _initial_glossary_rows,
)
from .medical_writing_tables import MedicalWritingTableService
from .medical_writing_figure_exporter import (
    MedicalWritingFigureExportError,
    add_svg_figure_with_png_fallback,
)
from .medical_writing_instrument_appendix import (
    MedicalWritingInstrumentAppendixError,
    validate_instrument_appendix_block,
)
from .medical_writing_style_profile import MedicalWritingStyleProfileService
from .medical_writing_protocol_template import (
    REPEATABLE_OBJECTIVE_NODE_IDS,
    materialize_repeatable_objective_heading,
)
from .medical_writing_source_reference_reindex import (
    SourceReferenceScanError,
    SourceReferenceTransformError,
    apply_source_reference_reindex,
    build_source_reference_manifest,
    decide_source_reference_reindex,
)
from .medical_writing_legacy_reference_index import (
    LegacyReferenceEntry,
    LegacyReferenceIndex,
    build_legacy_reference_index,
    is_reference_heading,
    is_superscript_reference_marker,
    parse_legacy_reference_marker,
)


class MedicalWritingDocumentExportMode(str, Enum):
    DRAFT_PREVIEW = "draft_preview"
    APPROVED_FINAL = "approved_final"


class MedicalWritingDocumentDocxExportError(ValueError):
    """Raised when a medical-writing document is not safe to export."""


@dataclass(frozen=True)
class MedicalWritingDocumentDocxExportResult:
    content: bytes
    metadata: Mapping[str, object]


@dataclass(frozen=True)
class _FrontMatterContext:
    block_ids: frozenset[str]
    document_title: str
    study_phase: str
    investigational_product: str
    protocol_date: str
    sponsor: str


_CMS_SPONSOR_NAME = "深圳市康哲生物科技有限公司"


def export_source_preserving_medical_writing_document_docx(
    source_path: Path | str,
    baseline_document: ProtocolDocument,
    document: ProtocolDocument,
    *,
    mode: MedicalWritingDocumentExportMode | str,
    approval_reference: Optional[str] = None,
    literature_library: Optional[MedicalWritingLiteratureLibrary] = None,
) -> MedicalWritingDocumentDocxExportResult:
    """Export an imported protocol without rebuilding its Word package.

    Sections, headers, footers, numbering, transparent layout tables, fields,
    drawings, and pagination remain source-authored. The exported copy receives
    only the mandatory bilingual font normalization (宋体 for East Asian text,
    Times New Roman for Latin text and digits) plus source-located content
    patches; the original DOCX remains byte-for-byte read-only. Source-located
    incremental patching is handled separately from greenfield template
    rendering.
    """

    export_mode = _parse_mode(mode)
    approval_ref = str(approval_reference or "").strip()
    _validate_export_boundary(document, export_mode, approval_ref)
    if (
        baseline_document.project_id != document.project_id
        or baseline_document.document_id != document.document_id
    ):
        raise MedicalWritingDocumentDocxExportError(
            "source-preserving export baseline does not match the assembled document"
        )
    source = Path(source_path).expanduser().resolve(strict=True)
    if source.suffix.lower() != ".docx":
        raise MedicalWritingDocumentDocxExportError(
            "source-preserving export requires an original DOCX source"
        )

    changed_blocks = _changed_source_blocks(baseline_document, document)
    has_managed_citations = _document_has_managed_citations(document)
    managed_citation_plan = (
        _build_citation_plan(
            document,
            ordered_blocks=_ordered_blocks(document),
            literature_library=literature_library,
        )
        if has_managed_citations
        else None
    )
    patched_block_ids = list(changed_blocks)
    if has_managed_citations:
        legacy_index = build_legacy_reference_index(document)
        for block in _all_document_blocks(document):
            block_id = str(block.get("block_id") or "")
            if (
                block_id
                and block_id not in patched_block_ids
                and any(_block_citation_reference_ids(block, legacy_index))
            ):
                patched_block_ids.append(block_id)
    source_content = source.read_bytes()
    content = (
        _patch_imported_docx_package(
            source_content,
            baseline_document,
            document,
            patched_block_ids,
            literature_library=literature_library,
            citation_plan=managed_citation_plan,
        )
        if patched_block_ids
        else source_content
    )
    content, reference_reindex_metadata = _apply_source_reference_reindex_for_export(
        content,
        source_content=source_content,
        managed_citation_plan=managed_citation_plan,
        export_mode=export_mode,
    )
    content = table_docx._normalize_docx_package(content)
    if not content.startswith(b"PK"):
        raise MedicalWritingDocumentDocxExportError(
            "original protocol source is not a valid OOXML package"
        )
    blocks = [block for section in document.sections for block in section.content_blocks]
    return MedicalWritingDocumentDocxExportResult(
        content=content,
        metadata={
            "schema_version": "medical_writing_source_preserving_export_v1",
            "document_id": document.document_id,
            "project_id": document.project_id,
            "protocol_id": document.protocol_id,
            "protocol_version": document.version,
            "mode": export_mode.value,
            "approval_reference": (
                approval_ref
                if export_mode == MedicalWritingDocumentExportMode.APPROVED_FINAL
                else ""
            ),
            "source_snapshot_sha256": _document_snapshot_digest(document),
            "block_count": len(blocks),
            "paragraph_count": sum(
                str(block.get("block_type") or "") in {"paragraph", "heading"}
                for block in blocks
            ),
            "table_count": sum(block.get("block_type") == "table" for block in blocks),
            "figure_count": sum(block.get("block_type") == "figure" for block in blocks),
            "docx_sha256": hashlib.sha256(content).hexdigest(),
            "source_docx_sha256": hashlib.sha256(source_content).hexdigest(),
            "source_preservation_mode": (
                "source_locator_patch_font_normalized"
                if patched_block_ids
                else "original_package_font_normalized"
            ),
            "changed_source_block_count": len(changed_blocks),
            "citation_rewritten_block_count": (
                len(set(patched_block_ids).difference(changed_blocks))
                if has_managed_citations
                else 0
            ),
            "literature_reference_count": (
                len(literature_library.references) if literature_library else 0
            ),
            "source_reference_reindex": reference_reindex_metadata,
        },
    )


def _apply_source_reference_reindex_for_export(
    content: bytes,
    *,
    source_content: bytes,
    managed_citation_plan: Optional[_CitationPlan],
    export_mode: MedicalWritingDocumentExportMode = MedicalWritingDocumentExportMode.DRAFT_PREVIEW,
) -> tuple[bytes, Mapping[str, object]]:
    if managed_citation_plan is not None:
        legacy_mapping = []
        for source_number, reference_id in sorted(
            managed_citation_plan.legacy_number_to_reference_id.items()
        ):
            canonical_id = managed_citation_plan.aliases_by_reference_id[reference_id]
            legacy_mapping.append(
                {
                    "source_number": source_number,
                    "target_number": managed_citation_plan.numbers_by_reference_id[
                        canonical_id
                    ],
                }
            )
        metadata = {
            "status": "applied",
            "action": "managed_citation_plan",
            "changed": True,
            "issue_codes": list(managed_citation_plan.issue_codes),
            "mapping": legacy_mapping,
            "source_digest": managed_citation_plan.legacy_source_digest,
            "initial_source_digest": managed_citation_plan.legacy_source_digest,
            "citation_count": len(managed_citation_plan.cited_reference_ids),
            "reference_count": len(managed_citation_plan.ordered_references),
            "failure_code": "",
            "failure_reason": "",
            "user_action_required": False,
            "warning_message": "",
        }
        return content, metadata

    initial_source_digest = ""
    try:
        initial_source_digest = build_source_reference_manifest(
            source_content
        ).source_digest
        manifest = build_source_reference_manifest(content)
    except SourceReferenceScanError as exc:
        metadata = {
            "status": "blocked",
            "action": "block",
            "changed": False,
            "issue_codes": ["source_reference_scan_failed"],
            "mapping": [],
            "source_digest": "",
            "initial_source_digest": initial_source_digest,
            "citation_count": 0,
            "reference_count": 0,
            "failure_code": "source_reference_scan_failed",
            "failure_reason": str(exc),
            "user_action_required": True,
            "warning_message": "文献引用扫描失败，请检查源文档",
        }
        if export_mode == MedicalWritingDocumentExportMode.APPROVED_FINAL:
            raise MedicalWritingDocumentDocxExportError(
                f"approved_final export blocked: source reference scan failed"
            )
        return content, metadata

    if not manifest.reference_groups and not manifest.citations:
        metadata = {
            "status": "not_applicable",
            "action": "not_applicable",
            "changed": False,
            "issue_codes": sorted({issue.code for issue in manifest.issues}),
            "mapping": [],
            "source_digest": manifest.source_digest,
            "initial_source_digest": initial_source_digest,
            "citation_count": 0,
            "reference_count": 0,
            "failure_code": "",
            "failure_reason": "",
            "user_action_required": False,
            "warning_message": "",
        }
        return content, metadata

    decision = decide_source_reference_reindex(manifest)

    action_to_status = {
        "apply": "applied",
        "preserve_with_notice": "preserved",
        "block": "blocked",
    }
    status = action_to_status.get(decision.action, "blocked")
    user_action_required = status == "blocked"

    warning_message = ""
    if status == "blocked":
        external_manager_issues = [
            issue
            for issue in decision.issues
            if issue.code
            in {
                "citation_manager_field_preserved",
                "citation_manager_metadata_conflict",
            }
        ]
        if external_manager_issues:
            warning_message = (
                "请在 Word 的临时工作副本中使用对应文献管理器刷新或重新绑定引用，"
                "保存后重新导入。"
            )
        else:
            blocking_issues = [
                issue for issue in decision.issues if issue.blocking
            ]
            if blocking_issues:
                warning_message = f"存在{len(blocking_issues)}个阻碍项，需处理后方可导出终稿"
            else:
                warning_message = "文献引用状态受阻，请处理后重新导入。"

    metadata = {
        "status": status,
        "action": decision.action,
        "changed": False,
        "issue_codes": sorted({issue.code for issue in decision.issues}),
        "mapping": [
            {
                "source_number": source_number,
                "target_number": target_number,
            }
            for source_number, target_number in decision.number_mapping
        ],
        "source_digest": manifest.source_digest,
        "initial_source_digest": initial_source_digest,
        "citation_count": len(manifest.citations),
        "reference_count": len(manifest.reference_groups),
        "failure_code": "",
        "failure_reason": "",
        "user_action_required": user_action_required,
        "warning_message": warning_message,
    }
    if (
        status == "blocked"
        and export_mode == MedicalWritingDocumentExportMode.APPROVED_FINAL
    ):
        raise MedicalWritingDocumentDocxExportError(
            f"approved_final export blocked: source reference status is {status}"
        )

    if decision.action != "apply":
        return content, metadata

    try:
        transformed = apply_source_reference_reindex(content, manifest, decision)
    except SourceReferenceTransformError as exc:
        metadata.update(
            {
                "status": "blocked",
                "changed": False,
                "failure_code": "source_reference_transform_failed",
                "failure_reason": str(exc),
                "user_action_required": True,
                "warning_message": f"引用转换失败：{str(exc)}",
            }
        )
        if export_mode == MedicalWritingDocumentExportMode.APPROVED_FINAL:
            raise MedicalWritingDocumentDocxExportError(
                f"approved_final export blocked: transform failed"
            )
        return content, metadata

    metadata.update(
        {
            "status": "applied",
            "changed": transformed != content,
            "user_action_required": False,
            "warning_message": "",
        }
    )
    return transformed, metadata


def _patch_imported_docx_package(
    source_content: bytes,
    baseline_document: ProtocolDocument,
    document: ProtocolDocument,
    changed_block_ids: Sequence[str],
    *,
    literature_library: Optional[MedicalWritingLiteratureLibrary],
    citation_plan: Optional[_CitationPlan],
) -> bytes:
    baseline_blocks = {
        str(block.get("block_id") or ""): block
        for section in baseline_document.sections
        for block in section.content_blocks
    }
    current_blocks = {
        str(block.get("block_id") or ""): block
        for section in document.sections
        for block in section.content_blocks
    }
    current_sections_by_block = {
        str(block.get("block_id") or ""): section
        for section in document.sections
        for block in section.content_blocks
    }
    has_managed_citations = any(
        True
        for section in document.sections
        for block in section.content_blocks
        for _ in _block_managed_citation_reference_ids(block)
    )
    if has_managed_citations and citation_plan is None:
        raise MedicalWritingDocumentDocxExportError(
            "managed citations require a unified citation plan"
        )
    try:
        with zipfile.ZipFile(io.BytesIO(source_content)) as source_zip:
            document_xml = source_zip.read("word/document.xml")
            root = etree.fromstring(document_xml)
            body = root.find(qn("w:body"))
            if body is None:
                raise MedicalWritingDocumentDocxExportError(
                    "original DOCX has no Word document body"
                )
            body_children = list(body)
            insertion_anchors: Dict[str, object] = {}
            package_replacements: Dict[str, bytes] = {}
            package_new_parts: Dict[str, tuple[zipfile.ZipInfo, bytes]] = {}
            index_plan = _build_document_index_plan(
                document,
                _ordered_blocks(document),
            )
            _ensure_imported_source_index_bookmarks(
                root,
                body_children,
                current_blocks,
                index_plan,
            )
            for block_id in changed_block_ids:
                baseline_block = baseline_blocks.get(block_id)
                current_block = current_blocks.get(block_id)
                if current_block is None:
                    raise MedicalWritingDocumentDocxExportError(
                        "source-preserving export cannot remove source blocks"
                    )
                if baseline_block is None:
                    section = current_sections_by_block[block_id]
                    anchor = insertion_anchors.get(section.section_id)
                    if anchor is None:
                        source_body_orders = [
                            block.get("body_order")
                            for block in section.content_blocks
                            if block.get("source_kind") == "original_protocol_docx"
                            and isinstance(block.get("body_order"), int)
                            and not isinstance(block.get("body_order"), bool)
                        ]
                        if not source_body_orders:
                            raise MedicalWritingDocumentDocxExportError(
                                "generated imported-document block has no source insertion anchor"
                            )
                        anchor_order = max(source_body_orders)
                        if anchor_order < 0 or anchor_order >= len(body_children):
                            raise MedicalWritingDocumentDocxExportError(
                                "generated imported-document insertion anchor is invalid"
                            )
                        anchor = body_children[anchor_order]
                    index_entry = index_plan.by_block_id.get(block_id)
                    package_seed = _compose_source_package(
                        source_content,
                        package_replacements,
                        package_new_parts,
                    )
                    inserted, replacements, new_parts = _render_generated_import_block(
                        package_seed,
                        current_block,
                        index_entry=index_entry,
                        index_plan=index_plan,
                    )
                    package_replacements.update(replacements)
                    duplicate_new_parts = set(package_new_parts).intersection(new_parts)
                    if duplicate_new_parts:
                        raise MedicalWritingDocumentDocxExportError(
                            "generated imported-document media names conflict"
                        )
                    package_new_parts.update(new_parts)
                    if not inserted:
                        raise MedicalWritingDocumentDocxExportError(
                            "generated imported-document block rendered no Word content"
                        )
                    for element in inserted:
                        anchor.addnext(element)
                        anchor = element
                    insertion_anchors[section.section_id] = anchor
                    continue
                if current_block.get("source_kind") != "original_protocol_docx":
                    raise MedicalWritingDocumentDocxExportError(
                        "generated imported-document objects require governed insertion"
                    )
                body_order = current_block.get("body_order")
                if (
                    not isinstance(body_order, int)
                    or isinstance(body_order, bool)
                    or body_order < 0
                    or body_order >= len(body_children)
                ):
                    raise MedicalWritingDocumentDocxExportError(
                        f"source block {block_id} has an invalid body locator"
                    )
                block_type = str(current_block.get("block_type") or "")
                target = body_children[body_order]
                if block_type in {"paragraph", "heading"}:
                    if target.tag != qn("w:p"):
                        raise MedicalWritingDocumentDocxExportError(
                            f"source paragraph locator no longer resolves: {block_id}"
                        )
                    _patch_word_paragraph(
                        target,
                        baseline_block,
                        current_block,
                        citation_plan=citation_plan,
                        index_plan=index_plan,
                    )
                elif block_type == "table":
                    if target.tag != qn("w:tbl"):
                        raise MedicalWritingDocumentDocxExportError(
                            f"source table locator no longer resolves: {block_id}"
                        )
                    _patch_word_table(
                        target,
                        baseline_block,
                        current_block,
                        citation_plan=citation_plan,
                        index_plan=index_plan,
                    )
                elif block_type == "figure":
                    if current_block != baseline_block:
                        raise MedicalWritingDocumentDocxExportError(
                            "source-linked figures are immutable in imported protocols"
                        )
                else:
                    raise MedicalWritingDocumentDocxExportError(
                        f"unsupported source-preserving block type: {block_type or '<empty>'}"
                    )

            if citation_plan is not None:
                _replace_imported_reference_list(
                    body,
                    body_children,
                    source_content,
                    document,
                    citation_plan,
                )

            patched_xml = etree.tostring(
                root,
                xml_declaration=True,
                encoding="UTF-8",
                standalone=True,
            )
            output_buffer = io.BytesIO()
            with zipfile.ZipFile(output_buffer, "w") as output_zip:
                output_zip.comment = source_zip.comment
                for item in source_zip.infolist():
                    output_zip.writestr(
                        item,
                        patched_xml
                        if item.filename == "word/document.xml"
                        else package_replacements.get(
                            item.filename,
                            source_zip.read(item.filename),
                        ),
                    )
                for item, payload in package_new_parts.values():
                    output_zip.writestr(item, payload)
            return output_buffer.getvalue()
    except zipfile.BadZipFile as exc:
        raise MedicalWritingDocumentDocxExportError(
            "original protocol source is not a valid DOCX package"
        ) from exc


def _render_generated_import_block(
    source_content: bytes,
    block: Mapping[str, object],
    *,
    index_entry: Optional[_IndexEntry],
    index_plan: _DocumentIndexPlan,
) -> tuple[
    List[object],
    Dict[str, bytes],
    Dict[str, tuple[zipfile.ZipInfo, bytes]],
]:
    source_kind = str(block.get("source_kind") or "")
    block_type = str(block.get("block_type") or "")
    if source_kind == "medical_writing_template" and block_type == "table":
        if index_entry is None or index_entry.kind != "table":
            raise MedicalWritingDocumentDocxExportError(
                "generated imported-document table has no governed index entry"
            )
        scratch = Document(io.BytesIO(source_content))
        body = scratch.element.body
        existing_count = len(body)
        table = MedicalWritingTableService().from_table_block(block)
        plan = table_docx._build_export_plan(table)
        _render_indexed_caption(scratch, index_entry)
        table_docx._render_table(scratch, table, plan)
        table_docx._render_notes(scratch, _deduplicated_notes(table.notes))
        appended = list(body)[existing_count - 1 : -1]
        return [copy.deepcopy(child) for child in appended], {}, {}
    if (
        source_kind == "medical_writing_intervention_rules"
        and block_type == "paragraph"
    ):
        return _render_source_patch_paragraphs(block, index_plan=index_plan), {}, {}
    if (
        source_kind == "medical_writing_study_schema"
        and block_type == "figure"
    ):
        if index_entry is None or index_entry.kind != "figure":
            raise MedicalWritingDocumentDocxExportError(
                "generated imported-document figure has no governed index entry"
            )
        return _render_generated_import_figure(
            source_content,
            block,
            index_entry=index_entry,
        )
    raise MedicalWritingDocumentDocxExportError(
        "unsupported generated block in imported protocol export"
    )


def _ensure_imported_source_index_bookmarks(
    root,
    body_children: Sequence[object],
    current_blocks: Mapping[str, Mapping[str, object]],
    index_plan: _DocumentIndexPlan,
) -> None:
    used_ids = {
        int(value)
        for node in root.iter(qn("w:bookmarkStart"))
        if (value := node.get(qn("w:id"))) is not None and value.isdigit()
    }
    next_id = max(used_ids, default=0) + 1
    existing_names = {
        str(node.get(qn("w:name")) or "")
        for node in root.iter(qn("w:bookmarkStart"))
    }
    for entry in index_plan.tables:
        if entry.bookmark_name in existing_names:
            continue
        block = current_blocks.get(entry.block_id)
        if not block or block.get("source_kind") != "original_protocol_docx":
            continue
        caption = block.get("table_caption")
        body_order = caption.get("body_order") if isinstance(caption, Mapping) else None
        if (
            not isinstance(body_order, int)
            or isinstance(body_order, bool)
            or body_order < 0
            or body_order >= len(body_children)
        ):
            continue
        paragraph = body_children[body_order]
        if paragraph.tag != qn("w:p"):
            continue
        children = list(paragraph)
        for begin_index, child in enumerate(children):
            field_chars = child.findall(".//" + qn("w:fldChar"))
            if not any(
                field.get(qn("w:fldCharType")) == "begin"
                for field in field_chars
            ):
                continue
            instruction_parts: List[str] = []
            end_index: Optional[int] = None
            for candidate_index in range(begin_index, len(children)):
                candidate = children[candidate_index]
                instruction_parts.extend(
                    str(node.text or "")
                    for node in candidate.findall(".//" + qn("w:instrText"))
                )
                if any(
                    field.get(qn("w:fldCharType")) == "end"
                    for field in candidate.findall(".//" + qn("w:fldChar"))
                ):
                    end_index = candidate_index
                    break
            instruction = "".join(instruction_parts)
            if end_index is None or not re.search(r"\bSEQ\s+表\b", instruction):
                continue
            bookmark_start = OxmlElement("w:bookmarkStart")
            bookmark_start.set(qn("w:id"), str(next_id))
            bookmark_start.set(qn("w:name"), entry.bookmark_name)
            bookmark_end = OxmlElement("w:bookmarkEnd")
            bookmark_end.set(qn("w:id"), str(next_id))
            paragraph.insert(begin_index, bookmark_start)
            paragraph.insert(end_index + 2, bookmark_end)
            existing_names.add(entry.bookmark_name)
            next_id += 1
            break


def _compose_source_package(
    source_content: bytes,
    replacements: Mapping[str, bytes],
    new_parts: Mapping[str, tuple[zipfile.ZipInfo, bytes]],
) -> bytes:
    if not replacements and not new_parts:
        return source_content
    output_buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(source_content)) as source_zip, zipfile.ZipFile(
        output_buffer,
        "w",
    ) as output_zip:
        output_zip.comment = source_zip.comment
        for item in source_zip.infolist():
            output_zip.writestr(
                item,
                replacements.get(item.filename, source_zip.read(item.filename)),
            )
        for item, payload in new_parts.values():
            output_zip.writestr(item, payload)
    return output_buffer.getvalue()


def _render_generated_import_figure(
    source_content: bytes,
    block: Mapping[str, object],
    *,
    index_entry: _IndexEntry,
) -> tuple[
    List[object],
    Dict[str, bytes],
    Dict[str, tuple[zipfile.ZipInfo, bytes]],
]:
    scratch = Document(io.BytesIO(source_content))
    source_body = list(scratch.element.body)
    if (
        not source_body
        or source_body[-1].tag != qn("w:sectPr")
        or sum(child.tag == qn("w:sectPr") for child in source_body) != 1
    ):
        raise MedicalWritingDocumentDocxExportError(
            "generated figure source body must end in exactly one sectPr"
        )
    source_body_count = len(source_body) - 1
    _render_study_schema_figure(
        scratch,
        block,
        figure_number=index_entry.number,
        bookmark_name=index_entry.bookmark_name,
    )
    rendered_buffer = io.BytesIO()
    scratch.save(rendered_buffer)
    rendered_content = rendered_buffer.getvalue()
    allowed_replacements = {
        "[Content_Types].xml",
        "word/_rels/document.xml.rels",
    }
    replacements: Dict[str, bytes] = {}
    new_parts: Dict[str, tuple[zipfile.ZipInfo, bytes]] = {}
    with zipfile.ZipFile(io.BytesIO(source_content)) as source_zip, zipfile.ZipFile(
        io.BytesIO(rendered_content)
    ) as rendered_zip:
        source_names = set(source_zip.namelist())
        for item in rendered_zip.infolist():
            name = item.filename
            payload = rendered_zip.read(name)
            if name == "word/document.xml":
                continue
            if name not in source_names:
                if not name.startswith("word/media/"):
                    raise MedicalWritingDocumentDocxExportError(
                        f"generated figure introduced an unexpected package part: {name}"
                    )
                new_parts[name] = (item, payload)
                continue
            if payload == source_zip.read(name):
                continue
            if _xml_payloads_are_equivalent(source_zip.read(name), payload):
                continue
            if name not in allowed_replacements:
                raise MedicalWritingDocumentDocxExportError(
                    f"generated figure changed an unrelated source package part: {name}"
                )
            replacements[name] = payload

        rendered_root = etree.fromstring(rendered_zip.read("word/document.xml"))
        rendered_body = rendered_root.find(qn("w:body"))
        if rendered_body is None:
            raise MedicalWritingDocumentDocxExportError(
                "generated figure package has no Word body"
            )
        rendered_children = list(rendered_body)
        if (
            not rendered_children
            or rendered_children[-1].tag != qn("w:sectPr")
            or sum(
                child.tag == qn("w:sectPr")
                for child in rendered_children
            )
            != 1
        ):
            raise MedicalWritingDocumentDocxExportError(
                "generated figure output body must end in exactly one sectPr"
            )
        appended = rendered_children[:-1][source_body_count:]
        if len(appended) != 1 or appended[0].tag != qn("w:p"):
            raise MedicalWritingDocumentDocxExportError(
                "generated figure must render exactly one atomic image/caption paragraph"
            )
        if not new_parts or not replacements.keys() >= allowed_replacements:
            raise MedicalWritingDocumentDocxExportError(
                "generated figure package is missing media or relationship updates"
            )
        return [copy.deepcopy(child) for child in appended], replacements, new_parts


def _xml_payloads_are_equivalent(left: bytes, right: bytes) -> bool:
    try:
        left_root = etree.fromstring(left)
        right_root = etree.fromstring(right)
    except etree.XMLSyntaxError:
        return False
    return etree.tostring(left_root, method="c14n") == etree.tostring(
        right_root,
        method="c14n",
    )


def _patch_word_paragraph(
    paragraph_element,
    baseline_block: Mapping[str, object],
    block: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan],
    index_plan: _DocumentIndexPlan,
) -> None:
    if _can_patch_source_text_in_place(
        baseline_block,
        block,
        citation_plan=citation_plan,
    ) and _patch_word_text_nodes_in_place(
        paragraph_element,
        str(baseline_block.get("text") or ""),
        str(block.get("text") or ""),
    ):
        return
    rendered = _render_source_patch_paragraphs(
        block,
        citation_plan=citation_plan,
        index_plan=index_plan,
    )
    if len(rendered) != 1:
        raise MedicalWritingDocumentDocxExportError(
            "one source paragraph cannot be expanded into multiple paragraphs during patch export"
        )
    replacement = rendered[0]
    for child in list(paragraph_element):
        if child.tag != qn("w:pPr"):
            paragraph_element.remove(child)
    for child in replacement:
        if child.tag != qn("w:pPr"):
            paragraph_element.append(copy.deepcopy(child))


def _can_patch_source_text_in_place(
    baseline_block: Mapping[str, object],
    block: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan],
) -> bool:
    baseline_rich_text = baseline_block.get("rich_text")
    current_rich_text = block.get("rich_text")
    if not isinstance(baseline_rich_text, Mapping) or not isinstance(
        current_rich_text, Mapping
    ):
        return citation_plan is None
    if citation_plan is not None or _rich_text_contains_governed_mark(current_rich_text):
        return False
    return _rich_text_format_signature(baseline_rich_text) == _rich_text_format_signature(
        current_rich_text
    )


def _rich_text_format_signature(node: Mapping[str, object]) -> object:
    node_type = str(node.get("type") or "")
    marks = node.get("marks")
    normalized_marks = (
        tuple(
            sorted(
                (
                    str(mark.get("type") or ""),
                    json.dumps(mark.get("attrs") or {}, ensure_ascii=False, sort_keys=True),
                )
                for mark in marks
                if isinstance(mark, Mapping)
            )
        )
        if isinstance(marks, list)
        else ()
    )
    attrs = node.get("attrs")
    normalized_attrs = json.dumps(
        attrs if isinstance(attrs, Mapping) else {},
        ensure_ascii=False,
        sort_keys=True,
    )
    content = node.get("content")
    children = (
        tuple(
            _rich_text_format_signature(child)
            for child in content
            if isinstance(child, Mapping)
        )
        if isinstance(content, list)
        else ()
    )
    return node_type, normalized_attrs, normalized_marks, children


def _rich_text_contains_governed_mark(node: Mapping[str, object]) -> bool:
    marks = node.get("marks")
    if isinstance(marks, list) and any(
        isinstance(mark, Mapping)
        and mark.get("type") in {"citation", "crossReference"}
        for mark in marks
    ):
        return True
    content = node.get("content")
    return isinstance(content, list) and any(
        isinstance(child, Mapping) and _rich_text_contains_governed_mark(child)
        for child in content
    )


def _patch_word_text_nodes_in_place(
    paragraph_element,
    source_text: str,
    target_text: str,
) -> bool:
    text_nodes = paragraph_element.findall(".//" + qn("w:t"))
    if "".join(str(node.text or "") for node in text_nodes) != source_text:
        return False
    if not text_nodes:
        return not target_text
    offsets: List[tuple[int, int]] = []
    cursor = 0
    for node in text_nodes:
        value = str(node.text or "")
        offsets.append((cursor, cursor + len(value)))
        cursor += len(value)
    projected = ["" for _ in text_nodes]

    def node_at(position: int, *, prefer_previous: bool = False) -> int:
        if position >= len(source_text):
            return len(text_nodes) - 1
        if prefer_previous and position > 0:
            position -= 1
        for index, (start, end) in enumerate(offsets):
            if start <= position < end:
                return index
        return 0

    matcher = difflib.SequenceMatcher(a=source_text, b=target_text, autojunk=False)
    for opcode, source_start, source_end, target_start, target_end in matcher.get_opcodes():
        replacement = target_text[target_start:target_end]
        if opcode == "equal":
            for index, (node_start, node_end) in enumerate(offsets):
                overlap_start = max(source_start, node_start)
                overlap_end = min(source_end, node_end)
                if overlap_start < overlap_end:
                    relative_start = overlap_start - source_start
                    relative_end = overlap_end - source_start
                    projected[index] += replacement[relative_start:relative_end]
            continue
        if opcode in {"replace", "insert"} and replacement:
            projected[node_at(source_start, prefer_previous=opcode == "insert")] += replacement

    for node, value in zip(text_nodes, projected):
        node.text = value
        if value.startswith(" ") or value.endswith(" "):
            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        else:
            node.attrib.pop("{http://www.w3.org/XML/1998/namespace}space", None)
    return True


def _patch_word_table(
    table_element,
    baseline_block: Mapping[str, object],
    current_block: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan],
    index_plan: _DocumentIndexPlan,
) -> None:
    baseline_rows = baseline_block.get("rows")
    current_rows = current_block.get("rows")
    if not isinstance(baseline_rows, list) or not isinstance(current_rows, list):
        raise MedicalWritingDocumentDocxExportError(
            "source-linked table rows are unavailable for patch export"
        )
    word_rows = table_element.findall(qn("w:tr"))
    if len(word_rows) != len(current_rows) or len(baseline_rows) != len(current_rows):
        raise MedicalWritingDocumentDocxExportError(
            "source-preserving export cannot change imported table row structure"
        )
    for row_index, (baseline_row, current_row, word_row) in enumerate(
        zip(baseline_rows, current_rows, word_rows)
    ):
        word_cells = word_row.findall(qn("w:tc"))
        if (
            not isinstance(baseline_row, list)
            or not isinstance(current_row, list)
            or len(word_cells) != len(current_row)
            or len(baseline_row) != len(current_row)
        ):
            raise MedicalWritingDocumentDocxExportError(
                f"source-preserving export cannot change imported table row {row_index}"
            )
        for baseline_cell, current_cell, word_cell in zip(
            baseline_row, current_row, word_cells
        ):
            if not isinstance(baseline_cell, Mapping) or not isinstance(
                current_cell, Mapping
            ):
                raise MedicalWritingDocumentDocxExportError(
                    "source-linked table cell payload is invalid"
                )
            if (
                baseline_cell.get("text") == current_cell.get("text")
                and baseline_cell.get("rich_text") == current_cell.get("rich_text")
            ):
                continue
            _patch_word_table_cell(
                word_cell,
                current_cell,
                citation_plan=citation_plan,
                index_plan=index_plan,
            )


def _patch_word_table_cell(
    cell_element,
    cell: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan],
    index_plan: _DocumentIndexPlan,
) -> None:
    block = {
        "block_id": str(cell.get("cell_id") or "source-table-cell"),
        "block_type": "paragraph",
        "text": str(cell.get("text") or ""),
        "rich_text": cell.get("rich_text"),
    }
    rendered = _render_source_patch_paragraphs(
        block,
        citation_plan=citation_plan,
        index_plan=index_plan,
    )
    existing = cell_element.findall(qn("w:p"))
    for index, replacement in enumerate(rendered):
        if index < len(existing):
            target = existing[index]
            for child in list(target):
                if child.tag != qn("w:pPr"):
                    target.remove(child)
            for child in replacement:
                if child.tag != qn("w:pPr"):
                    target.append(copy.deepcopy(child))
        else:
            cell_element.append(copy.deepcopy(replacement))
    for paragraph in existing[len(rendered) :]:
        cell_element.remove(paragraph)
    if not rendered:
        cell_element.append(OxmlElement("w:p"))


def _render_source_patch_paragraphs(
    block: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan] = None,
    index_plan: Optional[_DocumentIndexPlan] = None,
) -> List[object]:
    scratch = Document()
    body = scratch.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    rich_text = block.get("rich_text")
    if isinstance(rich_text, Mapping):
        if _rich_text_plain_text(rich_text) != str(block.get("text") or ""):
            raise MedicalWritingDocumentDocxExportError(
                f"block {block.get('block_id', '<unknown>')} rich text does not match its text projection"
            )
        _render_rich_text_root(
            scratch,
            rich_text,
            block,
            citation_plan=citation_plan,
            index_plan=index_plan,
        )
    else:
        paragraph = scratch.add_paragraph()
        paragraph.add_run(str(block.get("text") or ""))
    return [
        copy.deepcopy(child)
        for child in body
        if child.tag == qn("w:p")
    ]


def _replace_imported_reference_list(
    body,
    original_body_children: Sequence[object],
    source_content: bytes,
    document: ProtocolDocument,
    citation_plan: _CitationPlan,
) -> None:
    section_indices = sorted(_reference_section_indices(document))
    if len(section_indices) != 1:
        raise MedicalWritingDocumentDocxExportError(
            "managed citations require exactly one imported reference section"
        )
    reference_section = document.sections[section_indices[0]]
    heading_block = next(
        (
            block
            for block in reference_section.content_blocks
            if block.get("block_type") == "heading"
            and isinstance(block.get("body_order"), int)
        ),
        None,
    )
    if heading_block is None:
        raise MedicalWritingDocumentDocxExportError(
            "imported reference section has no stable heading anchor"
        )
    heading_order = int(heading_block["body_order"])
    if heading_order < 0 or heading_order >= len(original_body_children):
        raise MedicalWritingDocumentDocxExportError(
            "imported reference heading locator is invalid"
        )
    anchor = original_body_children[heading_order]

    legacy_orders = sorted(
        {
            int(block["body_order"])
            for section in document.sections
            for block in section.content_blocks
            if str(block.get("block_id") or "")
            in citation_plan.legacy_reference_block_ids
            and isinstance(block.get("body_order"), int)
        },
        reverse=True,
    )
    for body_order in legacy_orders:
        if body_order < 0 or body_order >= len(original_body_children):
            raise MedicalWritingDocumentDocxExportError(
                "imported reference entry locator is invalid"
            )
        target = original_body_children[body_order]
        if target.getparent() is body:
            body.remove(target)

    scratch = Document(io.BytesIO(source_content))
    scratch_body = scratch.element.body
    existing_count = len(scratch_body)
    reserved_bookmark_ids = {
        int(value)
        for node in body.getroottree().getroot().iter(qn("w:bookmarkStart"))
        if (value := node.get(qn("w:id"))) is not None and value.isdigit()
    }
    _render_reference_list(
        scratch,
        citation_plan,
        include_heading=False,
        style_presets=_resolve_style_profile(document).style_presets,
        reserved_bookmark_ids=reserved_bookmark_ids,
    )
    appended = list(scratch_body)[existing_count - 1 : -1]
    if not appended:
        raise MedicalWritingDocumentDocxExportError(
            "managed reference list rendered no Word paragraphs"
        )
    for paragraph in appended:
        inserted = copy.deepcopy(paragraph)
        anchor.addnext(inserted)
        anchor = inserted


def _changed_source_blocks(
    baseline_document: ProtocolDocument,
    document: ProtocolDocument,
) -> List[str]:
    baseline_sections = {
        section.section_id: section for section in baseline_document.sections
    }
    changed: List[str] = []
    for section in document.sections:
        baseline_section = baseline_sections.get(section.section_id)
        if baseline_section is None:
            changed.extend(
                str(block.get("block_id") or section.section_id)
                for block in section.content_blocks
            )
            continue
        baseline_blocks = {
            str(block.get("block_id") or ""): block
            for block in baseline_section.content_blocks
        }
        for block in section.content_blocks:
            block_id = str(block.get("block_id") or "")
            baseline_block = baseline_blocks.get(block_id)
            if baseline_block is None or block != baseline_block:
                changed.append(block_id or section.section_id)
        current_ids = {
            str(block.get("block_id") or "") for block in section.content_blocks
        }
        changed.extend(
            block_id
            for block_id in baseline_blocks
            if block_id not in current_ids
        )
    current_section_ids = {section.section_id for section in document.sections}
    changed.extend(
        section_id
        for section_id in baseline_sections
        if section_id not in current_section_ids
    )
    return list(dict.fromkeys(changed))


def _all_document_blocks(document: ProtocolDocument) -> List[Mapping[str, object]]:
    return [
        block
        for section in document.sections
        for block in section.content_blocks
    ]


def _document_has_managed_citations(document: ProtocolDocument) -> bool:
    return any(
        True
        for block in _all_document_blocks(document)
        for _ in _block_managed_citation_reference_ids(block)
    )


@dataclass(frozen=True)
class _CitationReference:
    reference_id: str
    legacy: Optional[LegacyReferenceEntry] = None
    managed: Optional[MedicalWritingProjectReference] = None


@dataclass(frozen=True)
class _CitationPlan:
    style: str
    ordered_references: Sequence[_CitationReference]
    numbers_by_reference_id: Mapping[str, int]
    bookmarks_by_reference_id: Mapping[str, str]
    aliases_by_reference_id: Mapping[str, str]
    legacy_number_to_reference_id: Mapping[int, str]
    legacy_reference_block_ids: frozenset[str]
    cited_reference_ids: Sequence[str]
    uncited_reference_ids: Sequence[str]
    legacy_source_digest: str
    issue_codes: Sequence[str]


@dataclass(frozen=True)
class _ResolvedStyleProfile:
    body_layout: tuple[object, ...]
    style_presets: Mapping[str, Mapping[str, object]]
    style_profile_id: str
    style_profile_version: str
    definition_sha256: str


@dataclass(frozen=True)
class _IndexEntry:
    kind: str
    object_id: str
    block_id: str
    number: int
    title: str
    bookmark_name: str


@dataclass(frozen=True)
class _DocumentIndexPlan:
    tables: Sequence[_IndexEntry]
    figures: Sequence[_IndexEntry]
    source_index_section_indices: frozenset[int]
    injection_section_index: int

    @property
    def by_block_id(self) -> Mapping[str, _IndexEntry]:
        return {
            entry.block_id: entry
            for entry in (*self.tables, *self.figures)
        }

    @property
    def by_object_key(self) -> Mapping[tuple[str, str], _IndexEntry]:
        return {
            (entry.kind, entry.object_id): entry
            for entry in (*self.tables, *self.figures)
        }


@dataclass(frozen=True)
class _VisualBlockLayout:
    target_layout: tuple[object, ...]
    effective_width_inches: float
    orientation: str
    context_indices: tuple[int, ...] = ()


@dataclass(frozen=True)
class _VisualLayoutPlan:
    target_layout_by_index: Mapping[int, tuple[object, ...]]
    visual_blocks_by_index: Mapping[int, _VisualBlockLayout]
    keep_with_next_indices: frozenset[int]
    compact_context_indices: frozenset[int]


_DEFAULT_HEADER_DISTANCE_TWIPS = 720
_DEFAULT_FOOTER_DISTANCE_TWIPS = 720
_HEADER_CONTENT_HEIGHT_PT = 31.5
_HEADER_BODY_GAP_PT = 8.0
_STUDY_SCHEMA_MIN_LANDSCAPE_WIDTH_INCHES = 8.0
_STUDY_SCHEMA_CAPTION_HEIGHT_PT = 12.0
_STUDY_SCHEMA_SECTION_BREAK_RESERVE_PT = 18.0
_STUDY_SCHEMA_RENDERER_ROUNDING_RESERVE_PT = 6.0
_SAFE_HEADER_TOP_MARGIN_TWIPS = _DEFAULT_HEADER_DISTANCE_TWIPS + round(
    (_HEADER_CONTENT_HEIGHT_PT + _HEADER_BODY_GAP_PT) * 20
)
_BODY_LAYOUT = (
    "portrait",
    table_docx._page_dimensions("portrait"),
    (
        720,
        720,
        _SAFE_HEADER_TOP_MARGIN_TWIPS,
        720,
        _DEFAULT_HEADER_DISTANCE_TWIPS,
        _DEFAULT_FOOTER_DISTANCE_TWIPS,
    ),
)
_APPENDIX_HEADER_DISTANCE_TWIPS = 720
_APPENDIX_FOOTER_DISTANCE_TWIPS = 720
_APPENDIX_HEADER_CONTENT_HEIGHT_PT = 31.5
_APPENDIX_FOOTER_CONTENT_HEIGHT_PT = 42.0
_APPENDIX_HEADER_FOOTER_GAP_PT = 6.0
_APPENDIX_TITLE_FONT_SIZE_PT = 10.5
_APPENDIX_TITLE_LINE_HEIGHT = 1.0
_APPENDIX_TITLE_SPACE_AFTER_PT = 4.0
_APPENDIX_IMAGE_LINE_BOX_RESERVE_PT = 7.0
_APPENDIX_SECTION_BREAK_RESERVE_PT = 18.0
# LibreOffice with real CJK fallback fonts allocates a substantially taller
# line/grid box around inline drawings than its font-less headless path. Keep
# one extra inch so title + drawing remains atomic in both LibreOffice and Word.
_APPENDIX_RENDERER_ROUNDING_RESERVE_PT = 86.0
_FINAL_APPROVAL_STATES = {
    ApprovalState.MEDICALLY_APPROVED,
    ApprovalState.LOCKED_FOR_SUBMISSION,
}
_MARKDOWN_SEPARATOR_RE = re.compile(
    r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$"
)
_HEX_COLOR_RE = re.compile(r"^#[0-9A-Fa-f]{6}$")
_FONT_SIZE_RE = re.compile(r"^(\d+(?:\.\d+)?)pt$")
_PARAGRAPH_ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_HIGHLIGHT_COLORS = {
    "#FFFF00": WD_COLOR_INDEX.YELLOW,
    "#00FF00": WD_COLOR_INDEX.BRIGHT_GREEN,
    "#00FFFF": WD_COLOR_INDEX.TURQUOISE,
    "#FF00FF": WD_COLOR_INDEX.PINK,
    "#0000FF": WD_COLOR_INDEX.BLUE,
    "#FF0000": WD_COLOR_INDEX.RED,
    "#000080": WD_COLOR_INDEX.DARK_BLUE,
    "#008080": WD_COLOR_INDEX.TEAL,
    "#008000": WD_COLOR_INDEX.GREEN,
    "#800080": WD_COLOR_INDEX.VIOLET,
    "#800000": WD_COLOR_INDEX.DARK_RED,
    "#808000": WD_COLOR_INDEX.DARK_YELLOW,
    "#808080": WD_COLOR_INDEX.GRAY_50,
    "#C0C0C0": WD_COLOR_INDEX.GRAY_25,
    "#000000": WD_COLOR_INDEX.BLACK,
}
_STYLE_PRESETS = {
    "heading_1": {
        "heading_level": 1,
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 16.0,
        "bold": True,
        "line_height": 1.0,
        "spacing_before_pt": 18.0,
        "spacing_after_pt": 8.0,
        "first_line_indent_chars": 0.0,
        "keep_with_next": True,
    },
    "heading_2": {
        "heading_level": 2,
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 14.0,
        "bold": True,
        "line_height": 1.0,
        "spacing_before_pt": 14.0,
        "spacing_after_pt": 6.0,
        "first_line_indent_chars": 0.0,
        "keep_with_next": True,
    },
    "heading_3": {
        "heading_level": 3,
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 12.0,
        "bold": True,
        "line_height": 1.0,
        "spacing_before_pt": 10.0,
        "spacing_after_pt": 4.0,
        "first_line_indent_chars": 0.0,
        "keep_with_next": True,
    },
    "heading_4": {
        "heading_level": 4,
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 10.5,
        "bold": True,
        "line_height": 1.0,
        "spacing_before_pt": 8.0,
        "spacing_after_pt": 3.0,
        "first_line_indent_chars": 0.0,
        "keep_with_next": True,
    },
    "body": {
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 10.5,
        "bold": False,
        "line_height": 1.5,
        "spacing_before_pt": 0.0,
        "spacing_after_pt": 4.0,
        "first_line_indent_chars": 2.0,
        "keep_with_next": False,
    },
    "note": {
        "east_asia": "宋体",
        "latin": "Times New Roman",
        "font_size_pt": 9.0,
        "bold": False,
        "line_height": 1.0,
        "spacing_before_pt": 2.0,
        "spacing_after_pt": 2.0,
        "first_line_indent_chars": 0.0,
        "keep_with_next": False,
    },
}

_INDEX_SECTION_HEADINGS = {
    "目录",
    "目录、表目录和图目录",
    "表格列表",
    "表目录",
    "图目录",
    "插图目录",
}
_INDEX_TEMPLATE_NODE_IDS = {"cms_indexes"}
_PRE_INDEX_TEMPLATE_NODE_IDS = {
    "ich_m11_front_matter",
    "cms_front_matter",
    "cms_confidentiality",
    "cms_signatures",
    "cms_version_history",
}
_NON_CATALOG_TABLE_LABELS = {
    "概要",
    "方案概要",
    "方案摘要",
    "缩略语",
    "缩略语表",
    "缩略语列表和术语定义",
    "术语定义",
    "方案修订记录",
    "文件版本历史",
}


def _mm_to_twips(value: object) -> int:
    try:
        millimetres = float(value)
    except (TypeError, ValueError) as exc:
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing style profile contains an invalid page measurement"
        ) from exc
    return round(millimetres * 1440 / 25.4)


def _safe_top_margin_twips(top_margin: int, header_distance: int) -> int:
    return max(
        top_margin,
        header_distance + round(
            (_HEADER_CONTENT_HEIGHT_PT + _HEADER_BODY_GAP_PT) * 20
        ),
    )


def _resolve_style_profile(document: ProtocolDocument) -> _ResolvedStyleProfile:
    binding = (
        document.style_profile_id.strip(),
        document.style_profile_version.strip(),
        document.style_profile_definition_sha256.strip(),
    )
    if not any(binding):
        return _ResolvedStyleProfile(
            body_layout=_BODY_LAYOUT,
            style_presets=_STYLE_PRESETS,
            style_profile_id="",
            style_profile_version="",
            definition_sha256="",
        )
    if not all(binding):
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing document has an incomplete style-profile binding"
        )
    try:
        definition = MedicalWritingStyleProfileService().definition(
            document.style_profile_id,
            document.style_profile_version,
        )
    except KeyError as exc:
        raise MedicalWritingDocumentDocxExportError(str(exc)) from exc
    if definition.definition_sha256 != document.style_profile_definition_sha256:
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing document style-profile hash does not match the registry"
        )
    layout = definition.page_layout
    header_distance = _mm_to_twips(layout.get("header_mm", 12.7))
    body_layout = (
        "portrait",
        table_docx._page_dimensions("portrait"),
        (
            _mm_to_twips(layout["left_mm"]),
            _mm_to_twips(layout["right_mm"]),
            _safe_top_margin_twips(
                _mm_to_twips(layout["top_mm"]),
                header_distance,
            ),
            _mm_to_twips(layout["bottom_mm"]),
            header_distance,
            _mm_to_twips(layout.get("footer_mm", 12.7)),
        ),
    )
    return _ResolvedStyleProfile(
        body_layout=body_layout,
        style_presets=definition.style_presets,
        style_profile_id=definition.style_profile_id,
        style_profile_version=definition.style_profile_version,
        definition_sha256=definition.definition_sha256,
    )


def export_medical_writing_document_docx(
    document: ProtocolDocument,
    *,
    mode: MedicalWritingDocumentExportMode | str,
    approval_reference: Optional[str] = None,
    literature_library: Optional[MedicalWritingLiteratureLibrary] = None,
    front_matter_overrides: Optional[Mapping[str, str]] = None,
    plan_consumption_helper: Any = None,
) -> MedicalWritingDocumentDocxExportResult:
    """Render one immutable ProtocolDocument snapshot as a complete Word document.

    Blocks from all sections are assembled by their original ``body_order``. Table
    blocks are converted to ``StructuredTable`` and rendered into this document;
    no independent table DOCX package is embedded or copied into the result.

    When *plan_consumption_helper* is injected, the export consumes one confirmed
    current plan revision (``docx_toc`` projection) and fails closed if the plan
    is missing, unconfirmed, stale, or has unresolved blocking drivers.
    """

    export_mode = _parse_mode(mode)
    approval_ref = str(approval_reference or "").strip()
    _validate_export_boundary(document, export_mode, approval_ref)
    # Pin export to one confirmed plan revision for TOC/applicability.
    # Fails closed on missing/unconfirmed/stale/unresolved plans.
    design_projection = None
    confirmed_plan = None
    if plan_consumption_helper is not None:
        plan_state, current_definition, design_projection = (
            plan_consumption_helper.require_confirmed_design_projection(
                project_id=document.project_id,
                projection_kind="docx_toc",
            )
        )
        confirmed_plan = plan_state.plan
        if document.source_study_definition_id and (
            document.source_study_definition_id,
            document.source_study_definition_revision,
            document.source_study_definition_sha256,
        ) != (
            current_definition.definition_id,
            current_definition.revision,
            current_definition.state_sha256,
        ):
            raise MedicalWritingDocumentDocxExportError(
                "protocol document StudyDefinition binding does not match the "
                "confirmed design projection"
            )
    source_digest = medical_writing_document_export_snapshot_digest(
        document,
        front_matter_overrides=front_matter_overrides,
    )
    style_profile = _resolve_style_profile(document)
    front_matter = _front_matter_context(
        document,
        overrides=front_matter_overrides,
    )
    ordered_blocks = _materialize_legacy_document_control_blocks(
        document,
        _ordered_blocks(document),
        front_matter=front_matter,
    )
    index_plan = _build_document_index_plan(document, ordered_blocks)
    visual_layout_plan = _build_visual_layout_plan(
        ordered_blocks,
        style_profile.body_layout,
        style_presets=style_profile.style_presets,
    )
    citation_plan = _build_citation_plan(
        document,
        ordered_blocks=ordered_blocks,
        literature_library=literature_library,
    )

    output = Document()
    _configure_default_styles(
        output,
        style_presets=style_profile.style_presets,
        enable_company_numbering=bool(style_profile.style_profile_id),
    )
    _configure_core_properties(output, document, export_mode, approval_ref)
    _configure_header(output.sections[0], document, export_mode, approval_ref)
    _configure_footer(output.sections[0], document, export_mode)
    _apply_body_layout(output.sections[0], style_profile.body_layout)

    table_service = MedicalWritingTableService()
    current_layout = style_profile.body_layout
    has_content = False
    last_paragraph_text = ""
    table_metadata: List[Dict[str, object]] = []
    figure_metadata: List[Dict[str, object]] = []
    appendix_metadata: List[Dict[str, object]] = []
    block_count = 0
    paragraph_count = 0
    pending_rendered_notes: List[str] = []
    pending_table_caption: Optional[Mapping[str, object]] = None
    indexes_rendered = False
    skipped_source_index_blocks = 0
    index_entry_by_block_id = index_plan.by_block_id

    reference_section_indices = _reference_section_indices(document)
    reference_insertion_index = _reference_insertion_index(
        ordered_blocks,
        reference_section_indices,
    )
    references_rendered = False
    objective_instance_plan = _repeatable_objective_instance_plan(document)
    if front_matter is not None:
        paragraph_count += _render_protocol_cover(
            output,
            document=document,
            context=front_matter,
        )
        has_content = True
    for ordered_index, (section_index, _, source_block) in enumerate(ordered_blocks):
        planned_layout_changed = False
        if (
            front_matter is not None
            and document.sections[section_index].template_node_id
            == "cms_confidentiality"
        ):
            # The company cover already contains the authoritative
            # confidentiality statement. Rendering the canonical document
            # control node again created an empty duplicate page.
            block_count += 1
            continue
        if (
            front_matter is not None
            and str(source_block.get("block_id") or "") in front_matter.block_ids
        ):
            block_count += 1
            continue
        if section_index in index_plan.source_index_section_indices:
            if not indexes_rendered:
                body_layout_already_active = current_layout == style_profile.body_layout
                current_layout = _ensure_body_section(
                    output,
                    current_layout=current_layout,
                    has_content=has_content,
                    document=document,
                    mode=export_mode,
                    approval_reference=approval_ref,
                    body_layout=style_profile.body_layout,
                )
                paragraph_count += _render_document_indexes(
                    output,
                    index_plan,
                    has_content=has_content and body_layout_already_active,
                )
                has_content = True
                indexes_rendered = True
            skipped_source_index_blocks += 1
            block_count += 1
            continue
        if (
            not indexes_rendered
            and section_index >= index_plan.injection_section_index
        ):
            body_layout_already_active = current_layout == style_profile.body_layout
            current_layout = _ensure_body_section(
                output,
                current_layout=current_layout,
                has_content=has_content,
                document=document,
                mode=export_mode,
                approval_reference=approval_ref,
                body_layout=style_profile.body_layout,
            )
            paragraph_count += _render_document_indexes(
                output,
                index_plan,
                has_content=has_content and body_layout_already_active,
            )
            has_content = True
            indexes_rendered = True
        block = _project_repeatable_objective_heading(
            document.sections[section_index],
            source_block,
            objective_instance_plan,
        )
        block_type = str(block.get("block_type") or "").strip().lower()
        native_heading_number: tuple[int, ...] = ()
        if style_profile.style_profile_id and block_type == "heading":
            block, native_heading_number = _project_native_heading_number(
                document.sections[section_index],
                block,
            )
        planned_layout = visual_layout_plan.target_layout_by_index.get(ordered_index)
        if planned_layout is not None:
            planned_layout_changed = current_layout != planned_layout
            current_layout = _ensure_figure_section(
                output,
                planned_layout,
                current_layout=current_layout,
                has_content=has_content,
                document=document,
                mode=export_mode,
                approval_reference=approval_ref,
            )
        if (
            citation_plan
            and str(block.get("block_id") or "")
            in citation_plan.legacy_reference_block_ids
        ):
            block_count += 1
            if ordered_index == reference_insertion_index:
                current_layout = _ensure_body_section(
                    output,
                    current_layout=current_layout,
                    has_content=has_content,
                    document=document,
                    mode=export_mode,
                    approval_reference=approval_ref,
                    body_layout=style_profile.body_layout,
                )
                paragraph_count += _render_reference_list(
                    output,
                    citation_plan,
                    include_heading=not _reference_heading_is_rendered(
                        document,
                        reference_section_indices,
                    ),
                    style_presets=style_profile.style_presets,
                )
                has_content = True
                references_rendered = True
            continue
        if block_type in {"paragraph", "heading"}:
            text = str(block.get("text") or "")
            if _looks_like_markdown_table(text):
                raise MedicalWritingDocumentDocxExportError(
                    f"block {block.get('block_id', '<unknown>')} contains an unrendered "
                    "Markdown table; convert it to a structured table before DOCX export"
                )
            note_text = _normalized_note_text(text)
            if pending_rendered_notes and (
                not note_text or note_text in pending_rendered_notes
            ):
                if note_text:
                    pending_rendered_notes.remove(note_text)
                block_count += 1
                continue
            pending_rendered_notes.clear()
            next_block = (
                ordered_blocks[ordered_index + 1][2]
                if ordered_index + 1 < len(ordered_blocks)
                else None
            )
            if _is_table_caption(block, next_block):
                pending_table_caption = block
                block_count += 1
                continue
            if planned_layout is None:
                current_layout = _ensure_body_section(
                    output,
                    current_layout=current_layout,
                    has_content=has_content,
                    document=document,
                    mode=export_mode,
                    approval_reference=approval_ref,
                    body_layout=style_profile.body_layout,
                )
            paragraph_start = len(output.paragraphs)
            paragraph_count += _render_paragraph(
                output,
                block,
                block_type,
                citation_plan=citation_plan,
                index_plan=index_plan,
                style_presets=style_profile.style_presets,
            )
            if native_heading_number:
                rendered_heading = next(
                    (
                        paragraph
                        for paragraph in output.paragraphs[paragraph_start:]
                        if paragraph.style.name.startswith("Heading ")
                    ),
                    None,
                )
                if rendered_heading is None:
                    raise MedicalWritingDocumentDocxExportError(
                        f"block {block.get('block_id', '<unknown>')} did not render an editable Heading paragraph"
                    )
                _apply_native_heading_number(rendered_heading, native_heading_number)
            if ordered_index in visual_layout_plan.keep_with_next_indices:
                for paragraph in output.paragraphs[paragraph_start:]:
                    paragraph.paragraph_format.keep_together = True
                    paragraph.paragraph_format.keep_with_next = True
                    if (
                        ordered_index
                        in visual_layout_plan.compact_context_indices
                    ):
                        paragraph.paragraph_format.line_spacing = 1.0
                        paragraph.paragraph_format.space_before = Pt(0)
                        paragraph.paragraph_format.space_after = Pt(0)
            has_content = True
            last_paragraph_text = text
        elif block_type == "table":
            try:
                structured = table_service.from_table_block(block)
                try:
                    plan = table_docx._build_export_plan(structured)
                except table_docx.StructuredTableDocxExportError as exc:
                    preserved_holes = _preserved_source_grid_holes(block)
                    if (
                        "uncovered grid positions" not in str(exc)
                        or (
                            preserved_holes is None
                            and not _is_original_source_version_zero(block)
                        )
                    ):
                        raise
                    structured = _fill_source_grid_holes(
                        structured,
                        expected_holes=preserved_holes,
                    )
                    plan = table_docx._build_export_plan(structured)
            except ValueError as exc:
                raise MedicalWritingDocumentDocxExportError(
                    f"table block {block.get('block_id', '<unknown>')} cannot be rendered: {exc}"
                ) from exc
            current_layout = _ensure_table_section(
                output,
                plan,
                current_layout=current_layout,
                has_content=has_content,
                document=document,
                mode=export_mode,
                approval_reference=approval_ref,
            )
            caption_text = ""
            table_index_entry = index_entry_by_block_id.get(
                str(block.get("block_id") or "")
            )
            if pending_table_caption is not None:
                caption_text = str(pending_table_caption.get("text") or "")
                if table_index_entry is None:
                    paragraph_count += _render_paragraph(
                        output,
                        pending_table_caption,
                        str(pending_table_caption.get("block_type") or "paragraph"),
                        citation_plan=citation_plan,
                        index_plan=index_plan,
                        style_presets=style_profile.style_presets,
                    )
                pending_table_caption = None
            if table_index_entry is not None:
                _render_indexed_caption(output, table_index_entry)
                paragraph_count += 1
            elif structured.title.strip() and _normalized_visible_text(
                structured.title
            ) not in {
                _normalized_visible_text(last_paragraph_text),
                _normalized_visible_text(caption_text),
            }:
                table_docx._render_title(output, structured)
                paragraph_count += 1
            word_table = table_docx._render_table(
                output,
                structured,
                plan,
                rich_text_run_renderer=(
                    lambda paragraph, text, marks, preset: _render_text_run(
                        paragraph,
                        text,
                        marks,
                        preset,
                        citation_plan=citation_plan,
                        index_plan=index_plan,
                    )
                ),
            )
            nested_synopsis_table_count = _render_protocol_synopsis_nested_groups(
                word_table,
                structured,
                plan,
                rich_text_run_renderer=(
                    lambda paragraph, text, marks, preset: _render_text_run(
                        paragraph,
                        text,
                        marks,
                        preset,
                        citation_plan=citation_plan,
                        index_plan=index_plan,
                    )
                ),
            )
            export_notes = _deduplicated_notes(structured.notes)
            table_docx._render_notes(output, export_notes)
            pending_rendered_notes = [
                _normalized_note_text(note.text)
                for note in export_notes
                if _normalized_note_text(note.text)
            ]
            has_content = True
            last_paragraph_text = ""
            table_metadata.append(
                {
                    "block_id": structured.block_id,
                    "table_id": structured.table_id,
                    "body_order": _body_order(block),
                    "row_count": len(plan.rows),
                    "column_count": len(plan.column_ids),
                    "orientation": plan.orientation,
                    "repeat_header_rows": plan.repeat_header_rows,
                    "merged_ranges": [
                        table_docx._merge_range(placed)
                        for placed in plan.placed_cells
                        if placed.cell.row_span > 1 or placed.cell.column_span > 1
                    ],
                    "note_markers": [
                        marker for marker, _ in table_docx._ordered_notes(export_notes)
                    ],
                    "word_table_rows": len(word_table.rows),
                    "nested_synopsis_table_count": nested_synopsis_table_count,
                    "indexed": table_index_entry is not None,
                    "table_number": (
                        table_index_entry.number if table_index_entry else None
                    ),
                    "caption": (
                        f"表 {table_index_entry.number} {table_index_entry.title}"
                        if table_index_entry
                        else structured.title.strip()
                    ),
                    "bookmark_name": (
                        table_index_entry.bookmark_name if table_index_entry else ""
                    ),
                }
            )
        elif block_type == "figure":
            figure_index_entry = index_entry_by_block_id.get(
                str(block.get("block_id") or "")
            )
            if figure_index_entry is None:
                raise MedicalWritingDocumentDocxExportError(
                    "indexed figure or image is missing from the document index"
                )
            if block.get("figure_kind") == "source_docx_image":
                current_layout = _ensure_body_section(
                    output,
                    current_layout=current_layout,
                    has_content=has_content,
                    document=document,
                    mode=export_mode,
                    approval_reference=approval_ref,
                    body_layout=style_profile.body_layout,
                )
                figure_metadata.append(
                    _render_source_docx_image(
                        output,
                        block,
                        entry=figure_index_entry,
                    )
                )
                paragraph_count += 2
                has_content = True
                last_paragraph_text = ""
                block_count += 1
                continue
            figure_layout = visual_layout_plan.visual_blocks_by_index[ordered_index]
            if figure_index_entry is None or figure_index_entry.kind != "figure":
                raise MedicalWritingDocumentDocxExportError(
                    "governed study-schema figure is missing from the document index"
                )
            figure = _render_study_schema_figure(
                output,
                block,
                figure_number=figure_index_entry.number,
                bookmark_name=figure_index_entry.bookmark_name,
                effective_width_inches=figure_layout.effective_width_inches,
                orientation=figure_layout.orientation,
            )
            figure_metadata.append(figure)
            paragraph_count += 1
            has_content = True
            last_paragraph_text = ""
        elif block_type == "appendix_image":
            page_layout = visual_layout_plan.visual_blocks_by_index[ordered_index]
            previous_is_appendix = bool(
                ordered_index > 0
                and str(
                    ordered_blocks[ordered_index - 1][2].get("block_type") or ""
                ).strip().lower()
                == "appendix_image"
            )
            appendix_metadata.append(
                _render_instrument_appendix_page(
                    output,
                    block,
                    effective_width_inches=page_layout.effective_width_inches,
                    orientation=page_layout.orientation,
                    page_break_before=(
                        has_content
                        and not previous_is_appendix
                        and not page_layout.context_indices
                        and not planned_layout_changed
                        and not _document_ends_with_page_break(output)
                    ),
                    page_break_after=ordered_index + 1 < len(ordered_blocks),
                )
            )
            paragraph_count += 2
            has_content = True
            last_paragraph_text = ""
        else:
            raise MedicalWritingDocumentDocxExportError(
                f"unsupported medical-writing block type: {block_type or '<empty>'}"
            )
        block_count += 1
        if citation_plan and ordered_index == reference_insertion_index:
            current_layout = _ensure_body_section(
                output,
                current_layout=current_layout,
                has_content=has_content,
                document=document,
                mode=export_mode,
                approval_reference=approval_ref,
                body_layout=style_profile.body_layout,
            )
            paragraph_count += _render_reference_list(
                output,
                citation_plan,
                include_heading=not _reference_heading_is_rendered(
                    document,
                    reference_section_indices,
                ),
                style_presets=style_profile.style_presets,
            )
            has_content = True
            references_rendered = True

    if citation_plan and not references_rendered:
        current_layout = _ensure_body_section(
            output,
            current_layout=current_layout,
            has_content=has_content,
            document=document,
            mode=export_mode,
            approval_reference=approval_ref,
            body_layout=style_profile.body_layout,
        )
        paragraph_count += _render_reference_list(
            output,
            citation_plan,
            include_heading=not _reference_heading_is_rendered(
                document,
                reference_section_indices,
            ),
            style_presets=style_profile.style_presets,
        )
        has_content = True

    if not indexes_rendered:
        current_layout = _ensure_body_section(
            output,
            current_layout=current_layout,
            has_content=has_content,
            document=document,
            mode=export_mode,
            approval_reference=approval_ref,
            body_layout=style_profile.body_layout,
        )
        paragraph_count += _render_document_indexes(
            output,
            index_plan,
            has_content=has_content,
        )
        has_content = True
        indexes_rendered = True

    if block_count == 0:
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing document contains no exportable content blocks"
        )
    if (
        medical_writing_document_export_snapshot_digest(
            document,
            front_matter_overrides=front_matter_overrides,
        )
        != source_digest
    ):
        raise RuntimeError(
            "medical-writing export mutated its source document snapshot"
        )

    _set_update_fields_on_open(output)
    _set_full_fidelity_compatibility_mode(output)
    raw = io.BytesIO()
    output.save(raw)
    content = table_docx._normalize_docx_package(raw.getvalue())
    metadata: Dict[str, object] = {
        "schema_version": "medical_writing_document_docx_export_v1",
        "document_id": document.document_id,
        "project_id": document.project_id,
        "protocol_id": document.protocol_id,
        "protocol_version": document.version,
        "mode": export_mode.value,
        "approval_reference": approval_ref
        if export_mode == MedicalWritingDocumentExportMode.APPROVED_FINAL
        else "",
        "source_snapshot_sha256": source_digest,
        "block_count": block_count,
        "paragraph_count": paragraph_count,
        "table_count": len(table_metadata),
        "tables": table_metadata,
        "figure_count": len(figure_metadata),
        "figures": figure_metadata,
        "appendix_image_page_count": len(appendix_metadata),
        "appendix_image_pages": appendix_metadata,
        "index_schema_version": "medical_writing_document_index_v1",
        "toc_field": True,
        "table_index_field": bool(index_plan.tables),
        "figure_index_field": bool(index_plan.figures),
        "indexed_table_count": len(index_plan.tables),
        "indexed_figure_count": len(index_plan.figures),
        "skipped_source_index_block_count": skipped_source_index_blocks,
        "docx_sha256": hashlib.sha256(content).hexdigest(),
        "style_profile_id": style_profile.style_profile_id,
        "style_profile_version": style_profile.style_profile_version,
        "style_profile_definition_sha256": style_profile.definition_sha256,
        "corpus_snapshot_id": document.corpus_snapshot_id,
        "corpus_snapshot_version": document.corpus_snapshot_version,
        "corpus_snapshot_sha256": document.corpus_snapshot_sha256,
    }
    if design_projection is not None and confirmed_plan is not None:
        metadata.update(
            {
                "protocol_assembly_plan_id": confirmed_plan.plan_id,
                "protocol_assembly_plan_revision": confirmed_plan.revision,
                "normalized_design_definition_id": (
                    design_projection.source_definition_id
                ),
                "normalized_design_definition_revision": (
                    design_projection.source_definition_revision
                ),
                "normalized_design_definition_sha256": (
                    design_projection.source_definition_sha256
                ),
            }
        )
    if citation_plan:
        metadata.update(
            {
                "citation_style": citation_plan.style,
                "citation_count": len(citation_plan.cited_reference_ids),
                "reference_count": len(citation_plan.numbers_by_reference_id),
                "cited_reference_ids": list(citation_plan.cited_reference_ids),
                "reference_ids": [
                    reference.reference_id for reference in citation_plan.ordered_references
                ],
                "uncited_reference_ids": list(citation_plan.uncited_reference_ids),
                "legacy_reference_count": sum(
                    1 for reference in citation_plan.ordered_references if reference.legacy
                ),
                "legacy_reference_source_sha256": citation_plan.legacy_source_digest,
                "reference_index_issue_codes": list(citation_plan.issue_codes),
            }
        )
    return MedicalWritingDocumentDocxExportResult(content=content, metadata=metadata)


def _parse_mode(
    mode: MedicalWritingDocumentExportMode | str,
) -> MedicalWritingDocumentExportMode:
    try:
        return MedicalWritingDocumentExportMode(mode)
    except ValueError as exc:
        choices = ", ".join(item.value for item in MedicalWritingDocumentExportMode)
        raise MedicalWritingDocumentDocxExportError(
            f"unsupported export mode; expected one of: {choices}"
        ) from exc


def _validate_export_boundary(
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
) -> None:
    corpus_binding = (
        document.corpus_snapshot_id.strip(),
        document.corpus_snapshot_version.strip(),
        document.corpus_snapshot_sha256.strip(),
    )
    if any(corpus_binding) and not all(corpus_binding):
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing document has an incomplete corpus-snapshot binding"
        )
    if corpus_binding[2] and not re.fullmatch(r"[0-9a-f]{64}", corpus_binding[2]):
        raise MedicalWritingDocumentDocxExportError(
            "medical-writing document corpus-snapshot hash must be lowercase SHA-256"
        )
    if mode != MedicalWritingDocumentExportMode.APPROVED_FINAL:
        return
    if not approval_reference:
        raise MedicalWritingDocumentDocxExportError(
            "approved final export requires an approval reference"
        )
    unapproved = [
        section.section_id
        for section in document.sections
        if section.approval_state not in _FINAL_APPROVAL_STATES
    ]
    if unapproved:
        preview = ", ".join(unapproved[:5])
        raise MedicalWritingDocumentDocxExportError(
            f"approved final export contains unapproved sections: {preview}"
        )


def _ordered_blocks(document: ProtocolDocument):
    indexed = []
    for section_index, section in enumerate(document.sections):
        for block_index, block in enumerate(section.content_blocks):
            indexed.append((section_index, block_index, block))
    return sorted(
        indexed,
        key=lambda item: (
            _body_order(item[2]) is None,
            _body_order(item[2]) if _body_order(item[2]) is not None else item[0],
            item[0],
            item[1],
        ),
    )


def _materialize_legacy_document_control_blocks(
    document: ProtocolDocument,
    ordered_blocks,
    *,
    front_matter: Optional[_FrontMatterContext],
):
    """Upgrade blank pre-contract document-control sections during export.

    Greenfield documents persisted before the structured document-control
    contract contain a heading plus an empty paragraph for signatures, version
    history, and glossary. Re-exporting those records must not recreate blank
    pages. The upgrade is export-local and uses only persisted document facts.
    """

    if front_matter is None:
        return ordered_blocks
    target_sections = {
        section_index: section
        for section_index, section in enumerate(document.sections)
        if section.template_node_id
        in {"cms_signatures", "cms_version_history", "cms_glossary"}
    }
    if not target_sections:
        return ordered_blocks
    section_has_table = {
        section_index: any(
            str(block.get("block_type") or "").strip().lower() == "table"
            for block in section.content_blocks
        )
        for section_index, section in target_sections.items()
    }
    upgraded_sections: set[int] = set()
    upgraded = []
    for section_index, block_index, block in ordered_blocks:
        section = target_sections.get(section_index)
        if (
            section is None
            or section_has_table[section_index]
            or section_index in upgraded_sections
            or str(block.get("block_type") or "").strip().lower() != "paragraph"
            or str(block.get("text") or "").strip()
        ):
            upgraded.append((section_index, block_index, block))
            continue
        body_order = _body_order(block)
        common = {
            "section_id": section.section_id,
            "document_id": document.document_id,
            "body_order": body_order if body_order is not None else block_index,
            "source_fact_ids": list(block.get("source_fact_ids") or []),
        }
        if section.template_node_id == "cms_signatures":
            replacement = _greenfield_document_control_table(
                **common,
                title="签字页",
                domain=StructuredTableDomain.GENERIC,
                columns=[
                    ("签署角色", "signature_role"),
                    ("姓名/职务", "signatory_identity"),
                    ("签名", "signature"),
                    ("日期", "signature_date"),
                ],
                rows=[
                    ("申办者代表", "", "", ""),
                    ("主要研究者", "", "", ""),
                ],
            )
        elif section.template_node_id == "cms_version_history":
            replacement = _greenfield_document_control_table(
                **common,
                title="研究方案版本更新记录",
                domain=StructuredTableDomain.VERSION_HISTORY,
                columns=[
                    ("版本号", "version_label"),
                    ("版本日期", "version_date"),
                    ("变更范围", "change_scope"),
                    ("变更说明", "change_summary"),
                    ("变更理由", "change_reason"),
                    ("批准状态", "approval_status"),
                ],
                rows=[
                    (
                        document.version,
                        front_matter.protocol_date,
                        "",
                        "",
                        "",
                        "",
                    )
                ],
            )
        else:
            glossary_rows = _initial_glossary_rows(front_matter.document_title)
            if not glossary_rows:
                upgraded.append((section_index, block_index, block))
                continue
            replacement = _greenfield_document_control_table(
                **common,
                title="缩略语与术语定义",
                domain=StructuredTableDomain.GENERIC,
                columns=[
                    ("缩略语", "abbreviation"),
                    ("中文全称", "definition_zh"),
                ],
                rows=glossary_rows,
            )
        upgraded.append((section_index, block_index, replacement))
        upgraded_sections.add(section_index)
    return upgraded


def document_index_catalog(document: ProtocolDocument) -> Dict[str, object]:
    """Return the deterministic table/figure targets used by DOCX indexes."""

    plan = _build_document_index_plan(document, _ordered_blocks(document))
    return {
        "schema_version": "medical_writing_document_index_v1",
        "document_id": document.document_id,
        "project_id": document.project_id,
        "tables": [_index_entry_payload(entry) for entry in plan.tables],
        "figures": [_index_entry_payload(entry) for entry in plan.figures],
    }


def _build_document_index_plan(document: ProtocolDocument, ordered_blocks) -> _DocumentIndexPlan:
    source_index_sections = frozenset(
        index
        for index, section in enumerate(document.sections)
        if _is_document_index_section(section)
    )
    injection_section_index = (
        min(source_index_sections)
        if source_index_sections
        else _default_index_injection_section_index(document)
    )
    tables: List[_IndexEntry] = []
    figures: List[_IndexEntry] = []
    for section_index, _, block in ordered_blocks:
        if section_index in source_index_sections:
            continue
        block_type = str(block.get("block_type") or "").strip().lower()
        block_id = str(block.get("block_id") or "").strip()
        if not block_id:
            continue
        if block_type == "table" and _table_is_catalog_content(
            document.sections[section_index], block
        ):
            object_id = str(block.get("table_id") or "").strip()
            title = _catalog_table_title(str(block.get("title") or ""))
            if not object_id or not title:
                continue
            tables.append(
                _IndexEntry(
                    kind="table",
                    object_id=object_id,
                    block_id=block_id,
                    number=len(tables) + 1,
                    title=title,
                    bookmark_name=_stable_object_bookmark("table", object_id),
                )
            )
        elif block_type == "figure":
            object_id = str(block.get("figure_id") or "").strip()
            raw_title = str(block.get("title") or "").strip()
            if (
                block.get("figure_kind") == "source_docx_image"
                and block.get("semantic_role") == "table_image"
            ):
                title = _catalog_table_title(raw_title)
                if object_id and title:
                    tables.append(
                        _IndexEntry(
                            kind="table",
                            object_id=object_id,
                            block_id=block_id,
                            number=len(tables) + 1,
                            title=title,
                            bookmark_name=_stable_object_bookmark("table", object_id),
                        )
                    )
            elif object_id and raw_title:
                figures.append(
                    _IndexEntry(
                        kind="figure",
                        object_id=object_id,
                        block_id=block_id,
                        number=len(figures) + 1,
                        title=re.sub(
                            r"^(?:图|附图)\s*\d+\s*[\.．、:]?\s*",
                            "",
                            raw_title,
                        ).strip(),
                        bookmark_name=_stable_object_bookmark("figure", object_id),
                    )
                )
    return _DocumentIndexPlan(
        tables=tuple(tables),
        figures=tuple(figures),
        source_index_section_indices=source_index_sections,
        injection_section_index=injection_section_index,
    )


def _normalized_index_heading(value: object) -> str:
    return re.sub(r"\s+", "", str(value or "").strip())


def _is_document_index_section(section) -> bool:
    return (
        str(getattr(section, "template_node_id", "") or "")
        in _INDEX_TEMPLATE_NODE_IDS
        or "document_index_editor"
        in set(getattr(section, "interaction_types", ()) or ())
        or _normalized_index_heading(section.heading) in _INDEX_SECTION_HEADINGS
    )


def _default_index_injection_section_index(document: ProtocolDocument) -> int:
    has_front_matter_boundary = False
    for index, section in enumerate(document.sections):
        template_node_id = str(getattr(section, "template_node_id", "") or "")
        node_kind = str(getattr(section, "node_kind", "") or "")
        if (
            template_node_id in _PRE_INDEX_TEMPLATE_NODE_IDS
            or node_kind in {"front_matter", "document_control"}
        ):
            has_front_matter_boundary = True
            continue
        if has_front_matter_boundary:
            return index
        return 0
    return len(document.sections)


def _table_is_catalog_content(section, block: Mapping[str, object]) -> bool:
    structure = block.get("structured_table")
    role = str(structure.get("role") or "") if isinstance(structure, Mapping) else ""
    if role in {"layout", "protocol_synopsis", "document_control"}:
        return False
    title = re.sub(r"\s+", "", str(block.get("title") or "").strip())
    section_heading = re.sub(r"\s+", "", str(section.heading or "").strip())
    if not title:
        return False
    if str(block.get("source_kind") or "") == "medical_writing_template":
        return str(block.get("template_id") or "") != "version_history"
    if title in _NON_CATALOG_TABLE_LABELS or section_heading in _NON_CATALOG_TABLE_LABELS:
        return False
    if re.match(r"^附表\s*\d+", str(block.get("title") or "").strip()):
        return False
    if re.match(r"^表\s*\d+", str(block.get("title") or "").strip()):
        return True
    return False


def _catalog_table_title(value: str) -> str:
    title = value.strip()
    title = re.sub(r"^表\s*\d+\s*[\.．、:]?\s*", "", title)
    return title.strip()


def _stable_object_bookmark(kind: str, object_id: str) -> str:
    prefix = "MWTAB" if kind == "table" else "MWFIG"
    digest = hashlib.sha256(object_id.encode("utf-8")).hexdigest()[:24]
    return f"_{prefix}_{digest}"


def _index_entry_payload(entry: _IndexEntry) -> Dict[str, object]:
    return {
        "kind": entry.kind,
        "object_id": entry.object_id,
        "block_id": entry.block_id,
        "number": entry.number,
        "title": entry.title,
        "bookmark_name": entry.bookmark_name,
        "display_label": f"{'表' if entry.kind == 'table' else '图'} {entry.number} {entry.title}",
    }


def _repeatable_objective_instance_plan(
    document: ProtocolDocument,
) -> Mapping[str, tuple[int, int]]:
    grouped: Dict[str, List[str]] = {}
    for section in document.sections:
        if section.template_node_id in REPEATABLE_OBJECTIVE_NODE_IDS:
            grouped.setdefault(section.template_node_id, []).append(section.section_id)
    return {
        section_id: (index, len(section_ids))
        for section_ids in grouped.values()
        for index, section_id in enumerate(section_ids, start=1)
    }


def _project_repeatable_objective_heading(
    section,
    block: Mapping[str, object],
    instance_plan: Mapping[str, tuple[int, int]],
) -> Mapping[str, object]:
    if str(block.get("block_type") or "").strip().lower() != "heading":
        return block
    if section.template_node_id not in REPEATABLE_OBJECTIVE_NODE_IDS:
        return block
    text = str(block.get("text") or "")
    if "<#>" not in text:
        return block
    instance_index, instance_count = instance_plan.get(section.section_id, (1, 1))
    projected = dict(block)
    projected["text"] = materialize_repeatable_objective_heading(
        text,
        template_node_id=section.template_node_id,
        instance_index=instance_index,
        instance_count=instance_count,
    )
    projected.pop("rich_text", None)
    return projected


def _project_native_heading_number(
    section,
    block: Mapping[str, object],
) -> tuple[Mapping[str, object], tuple[int, ...]]:
    """Render the authoritative number as stable Heading text.

    Word and LibreOffice do not agree when an isolated child heading starts a
    multilevel list without a visible ancestor in the same section.  Keep the
    paragraph as Heading 1-4 for navigation/TOC, but suppress inherited list
    numbering and include the complete authoritative number in its visible
    text.  This prevents both duplicate and truncated numbers across section
    breaks.
    """

    if block.get("include_in_toc") is False or block.get("suppress_numbering") is True:
        return block, ()
    number = str(section.section_number or "").strip()
    if not re.fullmatch(r"[1-9]\d*(?:\.[1-9]\d*){0,3}", number):
        return block, ()
    parts = tuple(int(part) for part in number.split("."))
    outline = block.get("outline_level")
    heading_level = outline + 1 if isinstance(outline, int) and 0 <= outline <= 3 else 1
    rich_text = block.get("rich_text")
    if isinstance(rich_text, Mapping) and str(rich_text.get("type") or "") == "heading":
        attrs = rich_text.get("attrs")
        rich_level = attrs.get("level") if isinstance(attrs, Mapping) else None
        if isinstance(rich_level, int) and not isinstance(rich_level, bool):
            heading_level = rich_level
    if heading_level != len(parts):
        return block, ()

    text = str(block.get("text") or "")
    prefix_match = re.match(
        rf"^\s*{re.escape(number)}(?=\s|[\u3001：:])(?:\s|[\u3001：:])*",
        text,
    )
    visible_title = text[prefix_match.end() :] if prefix_match else text
    section_heading = str(section.heading or "")
    section_prefix = re.match(
        rf"^\s*{re.escape(number)}(?=\s|[\u3001：:])(?:\s|[\u3001：:])*",
        section_heading,
    )
    canonical_heading = (
        section_heading[section_prefix.end() :] if section_prefix else section_heading
    )
    if (
        not visible_title.strip()
        or _normalized_visible_text(visible_title)
        != _normalized_visible_text(canonical_heading)
    ):
        return block, ()

    projected = dict(block)
    projected["text"] = (
        text.strip()
        if prefix_match is not None
        else f"{number} {visible_title.strip()}"
    )
    projected["suppress_numbering"] = True
    # Heading marks do not carry semantic content; the paragraph preset
    # preserves the editable Heading style and avoids a second rich-text path
    # reintroducing inherited numbering.
    projected.pop("rich_text", None)
    return projected, ()


def _strip_rich_text_prefix(
    root: Mapping[str, object],
    prefix: str,
) -> Mapping[str, object]:
    remaining = prefix

    def project(node: Mapping[str, object]) -> Mapping[str, object]:
        nonlocal remaining
        result = dict(node)
        if str(node.get("type") or "") == "text" and remaining:
            value = str(node.get("text") or "")
            consumed = min(len(value), len(remaining))
            if value[:consumed] != remaining[:consumed]:
                raise MedicalWritingDocumentDocxExportError(
                    "rich-text heading prefix does not match its text projection"
                )
            result["text"] = value[consumed:]
            remaining = remaining[consumed:]
            return result
        content = node.get("content")
        if isinstance(content, list):
            result["content"] = [
                project(child) if isinstance(child, Mapping) else child
                for child in content
            ]
        return result

    projected = project(root)
    if remaining:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text heading does not contain its visible numbering prefix"
        )
    return projected


def _build_citation_plan(
    document: ProtocolDocument,
    *,
    ordered_blocks,
    literature_library: Optional[MedicalWritingLiteratureLibrary],
) -> Optional[_CitationPlan]:
    legacy_index = build_legacy_reference_index(document)
    # A superscript number alone is not enough to prove a bibliography binding.
    # Activate legacy enforcement only when export would rebuild a unified list.
    has_managed_citations = any(
        True
        for _, _, block in ordered_blocks
        for _ in _block_managed_citation_reference_ids(block)
    )
    blocking_issues = [
        issue
        for issue in legacy_index.issues
        if issue.blocking
        and (
            legacy_index.entries
            or (issue.code == "unparsed_entry" and has_managed_citations)
        )
    ]
    if blocking_issues:
        details = "; ".join(
            f"{issue.code}:{issue.source_number or issue.block_id}"
            for issue in blocking_issues[:10]
        )
        raise MedicalWritingDocumentDocxExportError(
            "imported reference index contains unresolved blocking issues: " + details
        )

    style = "gbt_7714_2015_numeric"
    managed_by_id: Dict[str, MedicalWritingProjectReference] = {}
    if literature_library is not None:
        if literature_library.project_id != document.project_id:
            raise MedicalWritingDocumentDocxExportError(
                "project literature library does not belong to the exported document"
            )
        if literature_library.citation_style != "gbt_7714_2015_numeric":
            raise MedicalWritingDocumentDocxExportError(
                "unsupported citation style for DOCX export: "
                f"{literature_library.citation_style}"
            )
        style = literature_library.citation_style
        for reference in literature_library.references:
            if reference.reference_id in managed_by_id:
                raise MedicalWritingDocumentDocxExportError(
                    "project literature library contains duplicate reference id: "
                    f"{reference.reference_id}"
                )
            managed_by_id[reference.reference_id] = reference

    legacy_by_id = {entry.reference_id: entry for entry in legacy_index.entries}
    aliases: Dict[str, str] = {reference_id: reference_id for reference_id in legacy_by_id}
    for reference_id, reference in managed_by_id.items():
        if reference.project_id != document.project_id:
            raise MedicalWritingDocumentDocxExportError(
                f"project reference belongs to another project: {reference.reference_id}"
            )
        matching_legacy_ids = [
            entry.reference_id
            for entry in legacy_index.entries
            if _managed_matches_legacy(reference, entry)
        ]
        if len(matching_legacy_ids) > 1:
            raise MedicalWritingDocumentDocxExportError(
                "managed reference matches multiple imported references: "
                f"{reference_id}"
            )
        aliases[reference_id] = (
            matching_legacy_ids[0] if matching_legacy_ids else reference_id
        )

    legacy_block_ids = frozenset(
        source.block_id
        for entry in legacy_index.entries
        for source in entry.sources
        if source.block_id
    )
    ordered_reference_ids: List[str] = []
    seen_reference_ids: set[str] = set()
    cited_reference_ids: List[str] = []
    missing_managed_ids: List[str] = []
    for _, _, block in ordered_blocks:
        if str(block.get("block_id") or "") in legacy_block_ids:
            continue
        for raw_reference_id in _block_citation_reference_ids(block, legacy_index):
            if raw_reference_id not in aliases:
                missing_managed_ids.append(raw_reference_id)
                continue
            reference_id = aliases[raw_reference_id]
            if reference_id in seen_reference_ids:
                continue
            seen_reference_ids.add(reference_id)
            ordered_reference_ids.append(reference_id)
            cited_reference_ids.append(reference_id)
    if missing_managed_ids:
        raise MedicalWritingDocumentDocxExportError(
            "cited reference does not exist in the project literature library: "
            + ", ".join(dict.fromkeys(missing_managed_ids))
        )

    uncited_reference_ids: List[str] = []
    for entry in sorted(legacy_index.entries, key=lambda item: item.source_number):
        if entry.reference_id in seen_reference_ids:
            continue
        seen_reference_ids.add(entry.reference_id)
        ordered_reference_ids.append(entry.reference_id)
        uncited_reference_ids.append(entry.reference_id)

    if not ordered_reference_ids:
        return None

    ordered_references: List[_CitationReference] = []
    for reference_id in ordered_reference_ids:
        legacy = legacy_by_id.get(reference_id)
        if legacy is not None:
            ordered_references.append(
                _CitationReference(reference_id=reference_id, legacy=legacy)
            )
            continue
        managed = managed_by_id.get(reference_id)
        if managed is None:
            raise MedicalWritingDocumentDocxExportError(
                f"canonical reference is unavailable for export: {reference_id}"
            )
        if managed.validation_status == "needs_review":
            raise MedicalWritingDocumentDocxExportError(
                f"cited reference requires review before export: {reference_id}"
            )
        ordered_references.append(
            _CitationReference(reference_id=reference_id, managed=managed)
        )

    numbers = {
        reference_id: index
        for index, reference_id in enumerate(ordered_reference_ids, start=1)
    }
    bookmarks = {
        reference_id: "_MWREF_"
        + hashlib.sha256(reference_id.encode("utf-8")).hexdigest()[:16]
        for reference_id in ordered_reference_ids
    }
    return _CitationPlan(
        style=style,
        ordered_references=ordered_references,
        numbers_by_reference_id=numbers,
        bookmarks_by_reference_id=bookmarks,
        aliases_by_reference_id=aliases,
        legacy_number_to_reference_id=legacy_index.number_to_reference_id,
        legacy_reference_block_ids=legacy_block_ids,
        cited_reference_ids=tuple(cited_reference_ids),
        uncited_reference_ids=tuple(uncited_reference_ids),
        legacy_source_digest=legacy_index.source_digest,
        issue_codes=tuple(issue.code for issue in legacy_index.issues),
    )


def _block_citation_reference_ids(
    block: Mapping[str, object],
    legacy_index: LegacyReferenceIndex,
):
    rich_text = block.get("rich_text")
    if isinstance(rich_text, Mapping):
        yield from _unified_citation_reference_ids(rich_text, legacy_index)
    rows = block.get("rows")
    if not isinstance(rows, list):
        return
    for row in rows:
        if not isinstance(row, list):
            continue
        for cell in row:
            if not isinstance(cell, Mapping):
                continue
            cell_rich_text = cell.get("rich_text")
            if isinstance(cell_rich_text, Mapping):
                yield from _unified_citation_reference_ids(cell_rich_text, legacy_index)


def _block_managed_citation_reference_ids(block: Mapping[str, object]):
    rich_text = block.get("rich_text")
    if isinstance(rich_text, Mapping):
        yield from _citation_reference_ids(rich_text)
    rows = block.get("rows")
    if not isinstance(rows, list):
        return
    for row in rows:
        if not isinstance(row, list):
            continue
        for cell in row:
            if not isinstance(cell, Mapping):
                continue
            cell_rich_text = cell.get("rich_text")
            if isinstance(cell_rich_text, Mapping):
                yield from _citation_reference_ids(cell_rich_text)


def _unified_citation_reference_ids(
    node: Mapping[str, object],
    legacy_index: LegacyReferenceIndex,
    depth: int = 0,
):
    if depth > 12:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text nesting exceeds 12 levels"
        )
    if node.get("type") == "text":
        marks = node.get("marks", [])
        if not isinstance(marks, list) or any(
            not isinstance(mark, Mapping) for mark in marks
        ):
            raise MedicalWritingDocumentDocxExportError("rich-text marks are invalid")
        citation_marks = [mark for mark in marks if mark.get("type") == "citation"]
        if len(citation_marks) > 1:
            raise MedicalWritingDocumentDocxExportError(
                "rich-text text node contains duplicate citation marks"
            )
        if citation_marks:
            yield from _citation_reference_ids_from_mark(citation_marks[0])
            return
        if is_superscript_reference_marker(node):
            for source_number in parse_legacy_reference_marker(str(node.get("text") or "")):
                reference_id = legacy_index.number_to_reference_id.get(source_number)
                if reference_id:
                    yield reference_id
        return
    content = node.get("content", [])
    if not isinstance(content, list) or any(
        not isinstance(child, Mapping) for child in content
    ):
        raise MedicalWritingDocumentDocxExportError("rich-text content is invalid")
    for child in content:
        yield from _unified_citation_reference_ids(child, legacy_index, depth + 1)


def _citation_reference_ids(node: Mapping[str, object], depth: int = 0):
    if depth > 12:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text nesting exceeds 12 levels"
        )
    if node.get("type") == "text":
        marks = node.get("marks", [])
        if not isinstance(marks, list) or any(
            not isinstance(mark, Mapping) for mark in marks
        ):
            raise MedicalWritingDocumentDocxExportError("rich-text marks are invalid")
        citation_marks = [mark for mark in marks if mark.get("type") == "citation"]
        if len(citation_marks) > 1:
            raise MedicalWritingDocumentDocxExportError(
                "rich-text text node contains duplicate citation marks"
            )
        if citation_marks:
            yield from _citation_reference_ids_from_mark(citation_marks[0])
        return
    content = node.get("content", [])
    if not isinstance(content, list) or any(
        not isinstance(child, Mapping) for child in content
    ):
        raise MedicalWritingDocumentDocxExportError("rich-text content is invalid")
    for child in content:
        yield from _citation_reference_ids(child, depth + 1)


def _citation_reference_ids_from_mark(mark: Mapping[str, object]) -> List[str]:
    attrs = mark.get("attrs")
    if not isinstance(attrs, Mapping):
        raise MedicalWritingDocumentDocxExportError(
            "rich-text citation attributes are invalid"
        )
    if set(attrs) == {"referenceId"}:
        values = [attrs.get("referenceId")]
    elif set(attrs) == {"referenceIds"} and isinstance(attrs.get("referenceIds"), list):
        values = attrs["referenceIds"]
    else:
        raise MedicalWritingDocumentDocxExportError(
            "unsupported rich-text citation attributes: "
            + ", ".join(sorted(attrs))
        )
    if not 1 <= len(values) <= 50:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text citation must contain between 1 and 50 reference ids"
        )
    reference_ids = []
    for value in values:
        if not isinstance(value, str) or not value.strip():
            raise MedicalWritingDocumentDocxExportError(
                "rich-text citation referenceId is required"
            )
        reference_id = value.strip()
        if len(reference_id) > 200:
            raise MedicalWritingDocumentDocxExportError(
                "rich-text citation referenceId is too long"
            )
        if reference_id in reference_ids:
            raise MedicalWritingDocumentDocxExportError(
                "rich-text citation contains duplicate reference ids"
            )
        reference_ids.append(reference_id)
    if not reference_ids:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text citation referenceId is required"
        )
    return reference_ids


def _reference_section_indices(document: ProtocolDocument) -> set[int]:
    indices = set()
    for index, section in enumerate(document.sections):
        if is_reference_heading(section.heading):
            indices.add(index)
            continue
        if any(
            block.get("block_type") == "heading"
            and is_reference_heading(str(block.get("text") or ""))
            for block in section.content_blocks
        ):
            indices.add(index)
    return indices

def _managed_matches_legacy(
    managed: MedicalWritingProjectReference,
    legacy: LegacyReferenceEntry,
) -> bool:
    managed_values = {
        "doi": managed.doi.strip().casefold(),
        "pmid": re.sub(r"\D", "", managed.pmid),
        "url": _normalize_reference_url(managed.url),
    }
    for field_name in ("doi", "pmid", "url"):
        managed_value = managed_values[field_name]
        legacy_value = getattr(legacy, field_name)
        if managed_value and legacy_value:
            return managed_value == legacy_value
    managed_title = re.sub(
        r"[^\w]+",
        "",
        unicodedata.normalize("NFKC", managed.title).casefold(),
        flags=re.UNICODE,
    )
    return bool(
        managed_title
        and legacy.normalized_title
        and managed.year.strip()
        and legacy.year
        and managed_title == legacy.normalized_title
        and managed.year.strip() == legacy.year
    )


def _normalize_reference_url(value: str) -> str:
    parsed = urlsplit(value.strip())
    host = (parsed.hostname or "").casefold()
    if not host:
        return ""
    if host.startswith("www."):
        host = host[4:]
    path = re.sub(r"/{2,}", "/", parsed.path or "/").rstrip("/") or "/"
    return urlunsplit(("https", host, path, parsed.query, ""))


def _reference_insertion_index(ordered_blocks, section_indices: set[int]) -> Optional[int]:
    positions = [
        ordered_index
        for ordered_index, (section_index, _, _) in enumerate(ordered_blocks)
        if section_index in section_indices
    ]
    return max(positions) if positions else None


def _reference_heading_is_rendered(
    document: ProtocolDocument,
    section_indices: set[int],
) -> bool:
    return any(
        block.get("block_type") == "heading"
        and is_reference_heading(str(block.get("text") or ""))
        for section_index in section_indices
        for block in document.sections[section_index].content_blocks
    )


def _body_order(block: Mapping[str, object]) -> Optional[int]:
    value = block.get("body_order")
    return value if isinstance(value, int) and not isinstance(value, bool) else None


def _render_document_indexes(
    output: Document,
    plan: _DocumentIndexPlan,
    *,
    has_content: bool,
) -> int:
    index_specs = [
        ("目录", ' TOC \\o "1-4" \\h \\z \\u ', "更新域后显示目录"),
    ]
    if plan.figures:
        index_specs.append(
            ("图目录", ' TOC \\h \\z \\c "图" ', "更新域后显示图目录")
        )
    if plan.tables:
        index_specs.append(
            ("表目录", ' TOC \\h \\z \\c "表" ', "更新域后显示表目录")
        )
    rendered = 0
    if has_content:
        output.add_page_break()
        rendered += 1
    for index, (label, instruction, fallback) in enumerate(index_specs):
        if index:
            output.add_page_break()
            rendered += 1
        heading = output.add_paragraph()
        if "TOC Heading" in output.styles:
            heading.style = "TOC Heading"
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_after = Pt(12)
        heading_run = heading.add_run(label)
        heading_run.bold = True
        heading_run.font.size = Pt(16)
        heading_run.font.color.rgb = RGBColor(0, 0, 0)
        table_docx._set_run_fonts(
            heading_run,
            east_asia="宋体",
            latin="Times New Roman",
        )
        field_paragraph = output.add_paragraph()
        _append_word_field(
            field_paragraph,
            instruction,
            fallback,
            dirty=True,
        )
        rendered += 2
    output.add_page_break()
    return rendered + 1


def _render_indexed_caption(output: Document, entry: _IndexEntry) -> None:
    paragraph = output.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_after = Pt(6)
    prefix = paragraph.add_run(f"{'表' if entry.kind == 'table' else '图'} ")
    prefix.bold = True
    prefix.font.size = Pt(10.5)
    table_docx._set_run_fonts(prefix, east_asia="宋体", latin="Times New Roman")
    bookmark_id = str(7000 + entry.number + (0 if entry.kind == "table" else 1000))
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), bookmark_id)
    bookmark_start.set(qn("w:name"), entry.bookmark_name)
    paragraph._p.append(bookmark_start)
    _append_word_field(
        paragraph,
        f" SEQ {'表' if entry.kind == 'table' else '图'} \\* ARABIC ",
        str(entry.number),
        dirty=True,
        bold=True,
        font_size_pt=10.5,
    )
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), bookmark_id)
    paragraph._p.append(bookmark_end)
    title_run = paragraph.add_run(f" {entry.title}")
    title_run.bold = True
    title_run.font.size = Pt(10.5)
    table_docx._set_run_fonts(title_run, east_asia="宋体", latin="Times New Roman")


def _render_protocol_synopsis_nested_groups(
    word_table,
    structured: StructuredTable,
    plan,
    *,
    rich_text_run_renderer,
) -> int:
    """Project declared synopsis groups as editable Word tables inside the outer table."""

    layout = dict(structured.word_layout)
    raw_groups = layout.get("nested_groups")
    if structured.role.value != "protocol_synopsis" or not raw_groups:
        return 0
    if not isinstance(raw_groups, list):
        raise MedicalWritingDocumentDocxExportError(
            "protocol synopsis word_layout.nested_groups must be a list"
        )
    if len(plan.column_ids) < 3:
        raise MedicalWritingDocumentDocxExportError(
            "protocol synopsis nested groups require a label column and at least "
            "two content columns"
        )

    placed_by_coordinate = {
        (placed.row_index, placed.column_index): placed
        for placed in plan.placed_cells
    }
    content_column_indexes = list(range(1, len(plan.column_ids)))
    rendered_ranges: set[int] = set()
    rendered_count = 0
    font_size = table_docx._font_size(structured, len(plan.column_ids))

    for raw_group in raw_groups:
        if not isinstance(raw_group, Mapping):
            raise MedicalWritingDocumentDocxExportError(
                "protocol synopsis nested group must be an object"
            )
        group_id = str(raw_group.get("group_id") or "").strip()
        label = str(raw_group.get("label") or "").strip()
        declared_columns = raw_group.get("columns")
        if not group_id or not label:
            raise MedicalWritingDocumentDocxExportError(
                "protocol synopsis nested group requires group_id and label"
            )
        if not isinstance(declared_columns, list) or len(declared_columns) != len(
            content_column_indexes
        ):
            raise MedicalWritingDocumentDocxExportError(
                f"protocol synopsis nested group {group_id} column count does not "
                "match the outer content columns"
            )

        candidates = []
        for placed in plan.placed_cells:
            if placed.column_index != 0 or placed.cell.row_span < 2:
                continue
            semantic_value = placed.cell.semantic_value
            semantic_group_id = (
                str(semantic_value.get("nested_group_id") or "").strip()
                if isinstance(semantic_value, Mapping)
                else ""
            )
            if semantic_group_id == group_id or (
                not semantic_group_id
                and _normalized_visible_text(placed.cell.text)
                == _normalized_visible_text(label)
            ):
                candidates.append(placed)
        if len(candidates) != 1:
            raise MedicalWritingDocumentDocxExportError(
                f"protocol synopsis nested group {group_id} must resolve to exactly "
                "one row-spanning label cell"
            )

        group_cell = candidates[0]
        start_row = group_cell.row_index
        end_row = start_row + group_cell.cell.row_span - 1
        if end_row >= len(plan.rows):
            raise MedicalWritingDocumentDocxExportError(
                f"protocol synopsis nested group {group_id} exceeds the table rows"
            )
        occupied_rows = set(range(start_row, end_row + 1))
        if rendered_ranges.intersection(occupied_rows):
            raise MedicalWritingDocumentDocxExportError(
                f"protocol synopsis nested group {group_id} overlaps another group"
            )
        for row_index in occupied_rows:
            for column_index in content_column_indexes:
                if (row_index, column_index) not in placed_by_coordinate:
                    raise MedicalWritingDocumentDocxExportError(
                        f"protocol synopsis nested group {group_id} has an uncovered "
                        "inner grid position"
                    )

        merged_cell = word_table.cell(start_row, content_column_indexes[0]).merge(
            word_table.cell(end_row, content_column_indexes[-1])
        )
        merged_cell._tc.clear_content()
        table_docx._apply_cell_margins(
            merged_cell,
            {
                "cell_margins_twips": {
                    "top": 0,
                    "start": 0,
                    "bottom": 0,
                    "end": 0,
                }
            },
        )
        table_docx._set_cell_shading(merged_cell, "FFFFFF")
        merged_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP

        nested_widths = [
            plan.column_widths_twips[column_index]
            for column_index in content_column_indexes
        ]
        nested_width = sum(nested_widths)
        nested_table = merged_cell.add_table(
            rows=len(occupied_rows),
            cols=len(content_column_indexes),
        )
        nested_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        nested_table.autofit = False
        table_docx._set_fixed_table_layout(
            nested_table,
            nested_width,
            width_type="dxa",
        )
        table_docx._set_grid_column_widths(nested_table, nested_widths)
        table_docx._set_table_borders_single_black(nested_table)

        # The outer merged row can be taller than one page. Keeping it atomic
        # makes Word continue the row above the normal top margin.
        for source_row_index in occupied_rows:
            row_properties = word_table.rows[source_row_index]._tr.get_or_add_trPr()
            cant_split = row_properties.find(qn("w:cantSplit"))
            if cant_split is not None:
                row_properties.remove(cant_split)

        for nested_row_index, source_row_index in enumerate(
            range(start_row, end_row + 1)
        ):
            source_row = plan.rows[source_row_index]
            row_role = table_docx._normalized_role(source_row.style_role)
            nested_row = nested_table.rows[nested_row_index]
            table_docx._prevent_row_split(nested_row)
            for nested_column_index, source_column_index in enumerate(
                content_column_indexes
            ):
                placed = placed_by_coordinate[
                    (source_row_index, source_column_index)
                ]
                source_cell = placed.cell
                cell_role = table_docx._normalized_role(source_cell.style_role)
                role = (
                    row_role
                    if cell_role == "body" and row_role != "body"
                    else cell_role
                )
                word_cell = nested_row.cells[nested_column_index]
                table_docx._set_cell_width(
                    word_cell,
                    nested_widths[nested_column_index],
                    width_type="dxa",
                )
                table_docx._write_cell(
                    word_cell,
                    source_cell,
                    role=role,
                    font_size=font_size,
                    first_column=False,
                    is_label=False,
                    layout=layout,
                    rich_text_run_renderer=rich_text_run_renderer,
                )
                if role != "body":
                    continue
                numbering_cache: dict[int, int] = {}
                numbering_format = (
                    "decimal_half_paren"
                    if nested_column_index == 0
                    else "decimal_fullwidth_paren"
                )
                for paragraph in word_cell.paragraphs:
                    if (
                        not paragraph.text.strip()
                        or paragraph._p.xpath("./w:pPr/w:numPr")
                    ):
                        continue
                    table_docx._apply_word_list_numbering(
                        paragraph,
                        {
                            "list_type": "orderedList",
                            "sequence": 1,
                            "format": numbering_format,
                        },
                        numbering_cache=numbering_cache,
                    )
        rendered_ranges.update(occupied_rows)
        rendered_count += 1
    return rendered_count


def _append_word_field(
    paragraph,
    instruction_text: str,
    fallback_text: str,
    *,
    dirty: bool,
    bold: bool = False,
    font_size_pt: float = 10.5,
):
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    if dirty:
        field_begin.set(qn("w:dirty"), "true")
    begin_run = OxmlElement("w:r")
    begin_run.append(field_begin)
    paragraph._p.append(begin_run)

    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = instruction_text
    instruction_run = OxmlElement("w:r")
    instruction_run.append(instruction)
    paragraph._p.append(instruction_run)

    field_separator = OxmlElement("w:fldChar")
    field_separator.set(qn("w:fldCharType"), "separate")
    separator_run = OxmlElement("w:r")
    separator_run.append(field_separator)
    paragraph._p.append(separator_run)

    fallback = paragraph.add_run(fallback_text)
    fallback.bold = bold
    fallback.font.size = Pt(font_size_pt)
    table_docx._set_run_fonts(fallback, east_asia="宋体", latin="Times New Roman")

    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    end_run = OxmlElement("w:r")
    end_run.append(field_end)
    paragraph._p.append(end_run)
    return fallback


def _set_update_fields_on_open(output: Document) -> None:
    settings = output.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        following = {
            qn(tag)
            for tag in (
                "w:hdrShapeDefaults",
                "w:footnotePr",
                "w:endnotePr",
                "w:compat",
                "w:docVars",
                "w:rsids",
                "m:mathPr",
                "w:uiCompat97To2003",
                "w:attachedSchema",
                "w:themeFontLang",
                "w:clrSchemeMapping",
                "w:doNotIncludeSubdocsInStats",
                "w:doNotAutoCompressPictures",
                "w:forceUpgrade",
                "w:captions",
                "w:readModeInkLockDown",
                "w:smartTagType",
                "w:schemaLibrary",
                "w:shapeDefaults",
                "w:doNotEmbedSmartTags",
                "w:decimalSymbol",
                "w:listSeparator",
            )
        }
        for index, child in enumerate(settings):
            if child.tag in following:
                settings.insert(index, update_fields)
                break
        else:
            settings.append(update_fields)
    update_fields.set(qn("w:val"), "1")


def _set_full_fidelity_compatibility_mode(output: Document) -> None:
    """Open greenfield DOCX files with Word's full feature set enabled."""

    settings = output.settings._element
    compatibility = settings.find(qn("w:compat"))
    if compatibility is None:
        compatibility = OxmlElement("w:compat")
        settings.append(compatibility)
    mode = next(
        (
            item
            for item in compatibility.findall(qn("w:compatSetting"))
            if item.get(qn("w:name")) == "compatibilityMode"
            and item.get(qn("w:uri")) == "http://schemas.microsoft.com/office/word"
        ),
        None,
    )
    if mode is None:
        mode = OxmlElement("w:compatSetting")
        mode.set(qn("w:name"), "compatibilityMode")
        mode.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        compatibility.insert(0, mode)
    # 15 is Word 2013+ full-fidelity mode. The greenfield package is generated
    # by this service, so there is no legacy layout contract to preserve.
    mode.set(qn("w:val"), "15")


def _render_study_schema_figure(
    output: Document,
    block: Mapping[str, object],
    *,
    figure_number: int,
    bookmark_name: str,
    effective_width_inches: Optional[float] = None,
    orientation: str = "portrait",
) -> Dict[str, object]:
    if block.get("figure_kind") != "study_schema":
        raise MedicalWritingDocumentDocxExportError(
            "only governed study-schema figure blocks can be exported"
        )
    if block.get("source_kind") != "medical_writing_study_schema":
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure provenance is missing"
        )
    svg = str(block.get("svg") or "")
    svg_sha256 = str(block.get("svg_sha256") or "")
    if not svg or hashlib.sha256(svg.encode("utf-8")).hexdigest() != svg_sha256:
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure SVG payload does not match its hash"
        )
    try:
        png = base64.b64decode(str(block.get("png_base64") or ""), validate=True)
    except (ValueError, binascii.Error) as exc:
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure PNG fallback is not valid base64"
        ) from exc
    png_sha256 = str(block.get("png_sha256") or "")
    if not png.startswith(b"\x89PNG\r\n\x1a\n") or hashlib.sha256(png).hexdigest() != png_sha256:
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure PNG fallback does not match its hash"
        )
    figure_id = str(block.get("figure_id") or "").strip()
    title = str(block.get("title") or "").strip()
    alt_text = str(block.get("alt_text") or title).strip()
    if not figure_id or not title or not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]{0,39}", bookmark_name):
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure identity, title, or bookmark is invalid"
        )
    width_inches = block.get("width_inches", 6.45)
    if (
        isinstance(width_inches, bool)
        or not isinstance(width_inches, (int, float))
        or not 2.0 <= float(width_inches) <= 6.6
    ):
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure width is outside the body-page range"
        )

    rendered_width_inches = (
        float(effective_width_inches)
        if effective_width_inches is not None
        else float(width_inches)
    )
    image_paragraph = output.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_together = True
    image_paragraph.paragraph_format.keep_with_next = False
    image_paragraph.paragraph_format.line_spacing = 1.0
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(0)
    try:
        add_svg_figure_with_png_fallback(
            image_paragraph,
            svg_bytes=svg.encode("utf-8"),
            png_bytes=png,
            figure_id=figure_id,
            alt_text=alt_text,
            width=Inches(rendered_width_inches),
        )
    except MedicalWritingFigureExportError as exc:
        raise MedicalWritingDocumentDocxExportError(
            f"study-schema figure {figure_id} cannot be embedded: {exc}"
        ) from exc

    # The caption is the second line of the same keepLines paragraph as the
    # inline figure. Word may break a keepNext chain when the image consumes
    # nearly all remaining page height; one atomic paragraph makes the image
    # and caption structurally inseparable from the following section break.
    image_paragraph.add_run().add_break()
    caption = image_paragraph
    prefix = caption.add_run("图 ")
    table_docx._set_run_fonts(prefix, east_asia="宋体", latin="Times New Roman")
    bookmark_id = str(5000 + figure_number)
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), bookmark_id)
    bookmark_start.set(qn("w:name"), bookmark_name)
    caption._p.append(bookmark_start)
    _append_word_field(
        caption,
        " SEQ 图 \\* ARABIC ",
        str(figure_number),
        dirty=True,
    )
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), bookmark_id)
    caption._p.append(bookmark_end)
    title_run = caption.add_run(f" {title}")
    table_docx._set_run_fonts(title_run, east_asia="宋体", latin="Times New Roman")

    return {
        "block_id": str(block.get("block_id") or ""),
        "figure_id": figure_id,
        "figure_number": figure_number,
        "caption": f"图 {figure_number} {title}",
        "bookmark_name": bookmark_name,
        "body_order": _body_order(block),
        "schema_id": str(block.get("schema_id") or ""),
        "schema_revision": block.get("schema_revision"),
        "schema_state_sha256": str(block.get("schema_state_sha256") or ""),
        "layout_revision": block.get("layout_revision"),
        "svg_sha256": svg_sha256,
        "png_sha256": png_sha256,
        "requested_width_inches": float(width_inches),
        "width_inches": rendered_width_inches,
        "orientation": orientation,
        "vector_primary": True,
        "png_fallback": True,
    }


def _render_source_docx_image(
    output: Document,
    block: Mapping[str, object],
    *,
    entry: _IndexEntry,
) -> Dict[str, object]:
    if block.get("source_kind") != "original_protocol_docx":
        raise MedicalWritingDocumentDocxExportError(
            "source DOCX image provenance is missing"
        )
    try:
        image_bytes = base64.b64decode(
            str(block.get("image_base64") or ""),
            validate=True,
        )
    except (ValueError, binascii.Error) as exc:
        raise MedicalWritingDocumentDocxExportError(
            "source DOCX image payload is not valid base64"
        ) from exc
    image_sha256 = str(block.get("image_sha256") or "")
    media_type = str(block.get("media_type") or "")
    if (
        not image_bytes
        or hashlib.sha256(image_bytes).hexdigest() != image_sha256
        or media_type not in {"image/png", "image/jpeg"}
    ):
        raise MedicalWritingDocumentDocxExportError(
            "source DOCX image payload, hash, or media type is invalid"
        )
    if media_type == "image/png" and not image_bytes.startswith(b"\x89PNG\r\n\x1a\n"):
        raise MedicalWritingDocumentDocxExportError(
            "source DOCX PNG signature is invalid"
        )
    if media_type == "image/jpeg" and not image_bytes.startswith(b"\xff\xd8"):
        raise MedicalWritingDocumentDocxExportError(
            "source DOCX JPEG signature is invalid"
        )

    _render_indexed_caption(output, entry)
    image_paragraph = output.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    width_emu = block.get("width_emu")
    width_inches = (
        float(width_emu) / 914400
        if isinstance(width_emu, int) and not isinstance(width_emu, bool) and width_emu > 0
        else 6.0
    )
    width_inches = min(6.45, max(1.0, width_inches))
    try:
        image_paragraph.add_run().add_picture(
            io.BytesIO(image_bytes),
            width=Inches(width_inches),
        )
    except Exception as exc:
        raise MedicalWritingDocumentDocxExportError(
            f"source DOCX image {entry.object_id} cannot be embedded"
        ) from exc
    return {
        "block_id": str(block.get("block_id") or ""),
        "figure_id": entry.object_id,
        "figure_kind": "source_docx_image",
        "indexed_as": entry.kind,
        "number": entry.number,
        "title": entry.title,
        "bookmark_name": entry.bookmark_name,
        "media_type": media_type,
        "image_sha256": image_sha256,
        "width_inches": width_inches,
        "source_locator": str(block.get("source_locator") or ""),
    }


def _render_instrument_appendix_page(
    output: Document,
    block: Mapping[str, object],
    *,
    effective_width_inches: float,
    orientation: str,
    page_break_before: bool,
    page_break_after: bool,
) -> Dict[str, object]:
    try:
        validate_instrument_appendix_block(dict(block))
    except MedicalWritingInstrumentAppendixError as exc:
        raise MedicalWritingDocumentDocxExportError(str(exc)) from exc
    try:
        image_bytes = base64.b64decode(
            str(block.get("image_base64") or ""),
            validate=True,
        )
    except (ValueError, binascii.Error) as exc:
        raise MedicalWritingDocumentDocxExportError(
            "instrument appendix PNG is not valid base64"
        ) from exc

    title = str(block.get("title") or "").strip()
    page_number = int(block.get("page_number") or 0)
    page_count = int(block.get("page_count") or 0)
    title_paragraph = output.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.page_break_before = page_break_before
    title_paragraph.paragraph_format.keep_with_next = True
    title_paragraph.paragraph_format.keep_together = True
    title_paragraph.paragraph_format.line_spacing = _APPENDIX_TITLE_LINE_HEIGHT
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(
        _APPENDIX_TITLE_SPACE_AFTER_PT
    )
    title_run = title_paragraph.add_run(
        f"{title}（原始附件第{page_number}/{page_count}页）"
    )
    title_run.bold = True
    title_run.font.size = Pt(_APPENDIX_TITLE_FONT_SIZE_PT)
    table_docx._set_run_fonts(
        title_run,
        east_asia="宋体",
        latin="Times New Roman",
    )

    image_paragraph = output.add_paragraph()
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.keep_together = True
    image_paragraph.paragraph_format.space_before = Pt(0)
    image_paragraph.paragraph_format.space_after = Pt(0)
    try:
        image_run = image_paragraph.add_run()
        picture = image_run.add_picture(
            io.BytesIO(image_bytes),
            width=Inches(effective_width_inches),
        )
    except Exception as exc:
        raise MedicalWritingDocumentDocxExportError(
            f"instrument appendix page {page_number} cannot be embedded"
        ) from exc
    alt_text = str(block.get("alt_text") or title).strip()
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", title)
    if page_break_after:
        image_run.add_break(WD_BREAK.PAGE)

    return {
        "block_id": str(block.get("block_id") or ""),
        "instrument_id": str(block.get("instrument_id") or ""),
        "title": title,
        "page_number": page_number,
        "page_count": page_count,
        "render_dpi": int(block.get("render_dpi") or 0),
        "media_type": "image/png",
        "image_sha256": str(block.get("image_sha256") or ""),
        "source_pdf_sha256": str(block.get("source_pdf_sha256") or ""),
        "source_filename": str(block.get("source_filename") or ""),
        "width_inches": effective_width_inches,
        "orientation": orientation,
        "page_break_after": page_break_after,
        "indexed": False,
    }


def _build_visual_layout_plan(
    ordered_blocks,
    body_layout,
    *,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
) -> _VisualLayoutPlan:
    target_layout_by_index: Dict[int, tuple[object, ...]] = {}
    visual_blocks_by_index: Dict[int, _VisualBlockLayout] = {}
    keep_with_next_indices: set[int] = set()
    compact_context_indices: set[int] = set()
    for ordered_index, (_, _, block) in enumerate(ordered_blocks):
        block_type = str(block.get("block_type") or "").strip().lower()
        is_study_schema = (
            block_type == "figure" and block.get("figure_kind") == "study_schema"
        )
        is_appendix_page = block_type == "appendix_image"
        if not is_study_schema and not is_appendix_page:
            continue

        context_indices = (
            _leading_visual_context_indices(ordered_blocks, ordered_index)
            if is_study_schema or int(block.get("page_number") or 0) == 1
            else ()
        )
        context_height_inches = sum(
            _estimated_visual_context_height_inches(
                ordered_blocks[index][2],
                style_presets=style_presets,
            )
            for index in context_indices
        )
        if is_study_schema:
            target_layout, effective_width_inches, orientation = (
                _study_schema_figure_word_layout(
                    block,
                    body_layout,
                    context_height_inches=context_height_inches,
                )
            )
        else:
            target_layout, effective_width_inches, orientation = (
                _instrument_appendix_page_word_layout(
                    block,
                    body_layout,
                    context_height_inches=context_height_inches,
                )
            )
        layout = _VisualBlockLayout(
            target_layout=target_layout,
            effective_width_inches=effective_width_inches,
            orientation=orientation,
            context_indices=context_indices,
        )
        visual_blocks_by_index[ordered_index] = layout
        target_layout_by_index[ordered_index] = target_layout
        for context_index in context_indices:
            target_layout_by_index[context_index] = target_layout
            keep_with_next_indices.add(context_index)
            if is_study_schema:
                compact_context_indices.add(context_index)

    return _VisualLayoutPlan(
        target_layout_by_index=target_layout_by_index,
        visual_blocks_by_index=visual_blocks_by_index,
        keep_with_next_indices=frozenset(keep_with_next_indices),
        compact_context_indices=frozenset(compact_context_indices),
    )


def _leading_visual_context_indices(ordered_blocks, visual_index: int) -> tuple[int, ...]:
    section_index = ordered_blocks[visual_index][0]
    context: List[int] = []
    cursor = visual_index - 1
    while cursor >= 0 and len(context) < 3:
        candidate_section_index, _, candidate = ordered_blocks[cursor]
        if candidate_section_index != section_index:
            break
        block_type = str(candidate.get("block_type") or "").strip().lower()
        if block_type not in {"heading", "paragraph"}:
            break
        context.append(cursor)
        if block_type == "heading":
            break
        cursor -= 1
    return tuple(reversed(context))


def _estimated_visual_context_height_inches(
    block: Mapping[str, object],
    *,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
) -> float:
    block_type = str(block.get("block_type") or "").strip().lower()
    outline = block.get("outline_level")
    level = outline + 1 if isinstance(outline, int) and 0 <= outline <= 8 else 1
    preset_name = f"heading_{min(level, 4)}" if block_type == "heading" else "body"
    preset = style_presets[preset_name]
    text_lines = str(block.get("text") or "").splitlines() or [""]
    line_count = sum(max(1, (len(line.strip()) + 59) // 60) for line in text_lines)
    points = (
        float(preset["spacing_before_pt"])
        + float(preset["spacing_after_pt"])
        + line_count
        * float(preset["font_size_pt"])
        * float(preset["line_height"])
        + 6.0
    )
    return points / 72.0


def _instrument_appendix_page_word_layout(
    block,
    body_layout,
    *,
    context_height_inches: float = 0.0,
):
    pixel_width = int(block.get("pixel_width") or 0)
    pixel_height = int(block.get("pixel_height") or 0)
    if pixel_width <= 0 or pixel_height <= 0:
        raise MedicalWritingDocumentDocxExportError(
            "instrument appendix page pixel dimensions are invalid"
        )
    use_landscape = pixel_width > pixel_height
    if use_landscape:
        page_width_twips, page_height_twips = table_docx._page_dimensions(
            "landscape"
        )
        margins = (
            720,
            720,
            720,
            720,
            _APPENDIX_HEADER_DISTANCE_TWIPS,
            _APPENDIX_FOOTER_DISTANCE_TWIPS,
        )
        target_layout = (
            "landscape",
            (page_width_twips, page_height_twips),
            margins,
        )
        orientation = "landscape"
    else:
        _, page_dimensions, body_margins = body_layout
        margins = (
            *body_margins[:4],
            (
                body_margins[4]
                if len(body_margins) >= 6
                else _APPENDIX_HEADER_DISTANCE_TWIPS
            ),
            (
                body_margins[5]
                if len(body_margins) >= 6
                else _APPENDIX_FOOTER_DISTANCE_TWIPS
            ),
        )
        target_layout = ("portrait", page_dimensions, margins)
        orientation = "portrait"
        _, (page_width_twips, page_height_twips), margins = target_layout
    if use_landscape:
        _, (page_width_twips, page_height_twips), margins = target_layout
    left, right, top, bottom = margins[:4]
    available_width_inches = (
        page_width_twips - left - right
    ) / 1440
    if context_height_inches < 0:
        raise MedicalWritingDocumentDocxExportError(
            "instrument appendix context height is invalid"
        )
    header_distance = margins[4]
    footer_distance = margins[5]
    header_boundary_inches = max(
        top / 1440,
        header_distance / 1440
        + (
            _APPENDIX_HEADER_CONTENT_HEIGHT_PT
            + _APPENDIX_HEADER_FOOTER_GAP_PT
        )
        / 72,
    )
    footer_boundary_inches = max(
        bottom / 1440,
        footer_distance / 1440
        + (
            _APPENDIX_FOOTER_CONTENT_HEIGHT_PT
            + _APPENDIX_HEADER_FOOTER_GAP_PT
        )
        / 72,
    )
    title_text = (
        f"{str(block.get('title') or '').strip()}"
        f"（原始附件第{int(block.get('page_number') or 0)}/"
        f"{int(block.get('page_count') or 0)}页）"
    )
    title_line_count = _appendix_title_line_count(
        title_text,
        available_width_inches=available_width_inches,
    )
    non_image_height_inches = (
        title_line_count
        * _APPENDIX_TITLE_FONT_SIZE_PT
        * _APPENDIX_TITLE_LINE_HEIGHT
        + _APPENDIX_TITLE_SPACE_AFTER_PT
        + _APPENDIX_IMAGE_LINE_BOX_RESERVE_PT
        + _APPENDIX_SECTION_BREAK_RESERVE_PT
        + _APPENDIX_RENDERER_ROUNDING_RESERVE_PT
    ) / 72
    available_height_inches = (
        page_height_twips / 1440
        - header_boundary_inches
        - footer_boundary_inches
        - context_height_inches
        - non_image_height_inches
    )
    if available_height_inches <= 1.0:
        raise MedicalWritingDocumentDocxExportError(
            "instrument appendix page has insufficient vertical layout space"
        )
    aspect_ratio = pixel_width / pixel_height
    width_by_height = available_height_inches * aspect_ratio
    effective_width_inches = max(
        1.0,
        min(available_width_inches, width_by_height),
    )
    return target_layout, effective_width_inches, orientation


def _appendix_title_line_count(
    text: str,
    *,
    available_width_inches: float,
) -> int:
    if available_width_inches <= 0:
        raise MedicalWritingDocumentDocxExportError(
            "instrument appendix title has no horizontal layout space"
        )
    character_units = sum(
        1.0 if unicodedata.east_asian_width(character) in {"W", "F", "A"} else 0.55
        for character in text
    )
    units_per_line = max(
        1.0,
        available_width_inches * 72 / _APPENDIX_TITLE_FONT_SIZE_PT,
    )
    return max(1, math.ceil(character_units / units_per_line))


def _study_schema_figure_word_layout(
    block,
    body_layout,
    *,
    context_height_inches: float = 0.0,
):
    svg = str(block.get("svg") or "")
    width_match = re.search(r'<svg\b[^>]*\bwidth="([0-9]+(?:\.[0-9]+)?)"', svg)
    height_match = re.search(r'<svg\b[^>]*\bheight="([0-9]+(?:\.[0-9]+)?)"', svg)
    canvas_width = float(width_match.group(1)) if width_match else 0.0
    canvas_height = float(height_match.group(1)) if height_match else 0.0
    node_count = svg.count('data-node-id="')
    prefer_landscape = canvas_width >= 1600 or (
        node_count >= 7
        and canvas_height > 0
        and canvas_width / canvas_height >= 1.4
    )
    if not prefer_landscape:
        target_layout = body_layout
        orientation = "portrait"
        requested_width_inches = float(block.get("width_inches", 6.45))
    else:
        page_width_twips, page_height_twips = table_docx._page_dimensions("landscape")
        margins = (
            720,
            720,
            _SAFE_HEADER_TOP_MARGIN_TWIPS,
            720,
            _DEFAULT_HEADER_DISTANCE_TWIPS,
            _DEFAULT_FOOTER_DISTANCE_TWIPS,
        )
        target_layout = (
            "landscape",
            (page_width_twips, page_height_twips),
            margins,
        )
        orientation = "landscape"
        requested_width_inches = 10.4
    _, (page_width_twips, page_height_twips), margins = target_layout
    available_width_inches = (page_width_twips - margins[0] - margins[1]) / 1440
    available_height_inches = (page_height_twips - margins[2] - margins[3]) / 1440
    # Keep the leading context, image, generated caption, and following
    # section-break paragraph on one page. The compact figure paragraphs use
    # zero spacing, so only the caption line, break paragraph, and a small
    # renderer rounding allowance need to be reserved below the image.
    non_image_reserve_inches = (
        _STUDY_SCHEMA_CAPTION_HEIGHT_PT
        + _STUDY_SCHEMA_SECTION_BREAK_RESERVE_PT
        + _STUDY_SCHEMA_RENDERER_ROUNDING_RESERVE_PT
    ) / 72
    figure_height_inches = max(
        1.0,
        available_height_inches
        - context_height_inches
        - non_image_reserve_inches,
    )
    width_by_height = (
        figure_height_inches * canvas_width / canvas_height
        if canvas_width > 0 and canvas_height > 0
        else available_width_inches
    )
    effective_width_inches = min(
        requested_width_inches,
        available_width_inches,
        width_by_height,
    )
    if (
        orientation == "landscape"
        and effective_width_inches < _STUDY_SCHEMA_MIN_LANDSCAPE_WIDTH_INCHES
    ):
        raise MedicalWritingDocumentDocxExportError(
            "landscape study-schema figure cannot retain its grouped context "
            "at the minimum 8.0-inch width"
        )
    if effective_width_inches < 2.0:
        raise MedicalWritingDocumentDocxExportError(
            "study-schema figure has insufficient grouped layout space"
        )
    return target_layout, effective_width_inches, orientation


def _document_ends_with_page_break(output: Document) -> bool:
    paragraphs = output.paragraphs
    return bool(
        paragraphs
        and paragraphs[-1]._p.xpath('.//w:br[@w:type="page"]')
    )


def _ensure_body_section(
    output: Document,
    *,
    current_layout,
    has_content: bool,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
    body_layout,
):
    if current_layout == body_layout:
        return current_layout
    section = (
        output.add_section(WD_SECTION.NEW_PAGE) if has_content else output.sections[-1]
    )
    _apply_body_layout(section, body_layout)
    _configure_header(section, document, mode, approval_reference)
    # Footer inherits from the initial section so PAGE/NUMPAGES fields
    # count continuously across all sections, matching company templates
    # where landscape or alternate-orientation sections omit footerReference.
    section.footer.is_linked_to_previous = True
    return body_layout


def _ensure_figure_section(
    output: Document,
    target_layout,
    *,
    current_layout,
    has_content: bool,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
):
    if current_layout == target_layout:
        return current_layout
    section = (
        output.add_section(WD_SECTION.NEW_PAGE) if has_content else output.sections[-1]
    )
    orientation, (width, height), margins = target_layout
    section.orientation = (
        WD_ORIENT.LANDSCAPE if orientation == "landscape" else WD_ORIENT.PORTRAIT
    )
    section.page_width = Twips(width)
    section.page_height = Twips(height)
    section.left_margin = Twips(margins[0])
    section.right_margin = Twips(margins[1])
    section.top_margin = Twips(margins[2])
    section.bottom_margin = Twips(margins[3])
    if len(margins) >= 6:
        section.header_distance = Twips(margins[4])
        section.footer_distance = Twips(margins[5])
    _configure_header(section, document, mode, approval_reference)
    section.footer.is_linked_to_previous = True
    return target_layout


def _ensure_table_section(
    output: Document,
    plan,
    *,
    current_layout,
    has_content: bool,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
):
    target_layout = (
        plan.orientation,
        (plan.page_width_twips, plan.page_height_twips),
        (
            plan.margins_twips["left"],
            plan.margins_twips["right"],
            plan.margins_twips["top"],
            plan.margins_twips["bottom"],
        ),
    )
    if current_layout == target_layout:
        return current_layout
    section = (
        output.add_section(WD_SECTION.NEW_PAGE) if has_content else output.sections[-1]
    )
    _apply_table_layout(section, plan)
    _configure_header(section, document, mode, approval_reference)
    section.footer.is_linked_to_previous = True
    return target_layout


def _apply_body_layout(section, body_layout) -> None:
    width, height = body_layout[1]
    margins = body_layout[2]
    left, right, top, bottom = margins[:4]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Twips(width)
    section.page_height = Twips(height)
    section.left_margin = Twips(left)
    section.right_margin = Twips(right)
    section.top_margin = Twips(top)
    section.bottom_margin = Twips(bottom)
    if len(margins) >= 6:
        section.header_distance = Twips(margins[4])
        section.footer_distance = Twips(margins[5])


def _apply_table_layout(section, plan) -> None:
    section.orientation = (
        WD_ORIENT.LANDSCAPE if plan.orientation == "landscape" else WD_ORIENT.PORTRAIT
    )
    section.page_width = Twips(plan.page_width_twips)
    section.page_height = Twips(plan.page_height_twips)
    section.left_margin = Twips(plan.margins_twips["left"])
    section.right_margin = Twips(plan.margins_twips["right"])
    header_distance = int(
        getattr(section.header_distance, "twips", 0)
        or _DEFAULT_HEADER_DISTANCE_TWIPS
    )
    section.header_distance = Twips(header_distance)
    section.top_margin = Twips(
        _safe_top_margin_twips(
            plan.margins_twips["top"],
            header_distance,
        )
    )
    section.bottom_margin = Twips(plan.margins_twips["bottom"])


def _render_paragraph(
    output: Document,
    block: Mapping[str, object],
    block_type: str,
    *,
    citation_plan: Optional[_CitationPlan] = None,
    index_plan: Optional[_DocumentIndexPlan] = None,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
) -> int:
    text = str(block.get("text") or "")
    rich_text = block.get("rich_text")
    if isinstance(rich_text, Mapping):
        if _rich_text_plain_text(rich_text) != text:
            raise MedicalWritingDocumentDocxExportError(
                f"block {block.get('block_id', '<unknown>')} rich text does not match its text projection"
            )
        return _render_rich_text_root(
            output,
            rich_text,
            block,
            citation_plan=citation_plan,
            index_plan=index_plan,
            style_presets=style_presets,
        )

    outline = block.get("outline_level")
    level = outline + 1 if isinstance(outline, int) and 0 <= outline <= 8 else 1
    preset_name = f"heading_{min(level, 4)}" if block_type == "heading" else "body"
    text_segments = [text]
    if block_type == "paragraph" and ("\n" in text or "\r" in text):
        text_segments = [
            segment.strip()
            for segment in text.splitlines()
            if segment.strip()
        ]
    if not text_segments:
        text_segments = [""]

    for segment_index, segment in enumerate(text_segments):
        paragraph = output.add_paragraph()
        _apply_paragraph_preset(
            paragraph,
            preset_name,
            {},
            style_presets=style_presets,
        )
        if block_type == "heading" and block.get("include_in_toc") is False:
            if "CMS Document Control Heading" in output.styles:
                paragraph.style = "CMS Document Control Heading"
        if (
            block_type == "heading"
            and block.get("page_break_before") is True
            and segment_index == 0
        ):
            paragraph.paragraph_format.page_break_before = True
        if block_type == "heading" and block.get("suppress_numbering") is True:
            _suppress_paragraph_numbering(paragraph)
        _render_text_run(
            paragraph,
            segment,
            [],
            style_presets[preset_name],
            citation_plan=citation_plan,
            index_plan=index_plan,
        )
    return len(text_segments)


def _suppress_paragraph_numbering(paragraph) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    numbering_properties = paragraph_properties.find(qn("w:numPr"))
    if numbering_properties is None:
        numbering_properties = OxmlElement("w:numPr")
        table_docx._insert_before_first(
            paragraph_properties,
            numbering_properties,
            (
                "w:suppressLineNumbers",
                "w:pBdr",
                "w:shd",
                "w:tabs",
                "w:suppressAutoHyphens",
                "w:kinsoku",
                "w:wordWrap",
                "w:overflowPunct",
                "w:topLinePunct",
                "w:autoSpaceDE",
                "w:autoSpaceDN",
                "w:bidi",
                "w:adjustRightInd",
                "w:snapToGrid",
                "w:spacing",
                "w:ind",
                "w:contextualSpacing",
                "w:mirrorIndents",
                "w:suppressOverlap",
                "w:jc",
                "w:textDirection",
                "w:textAlignment",
                "w:textboxTightWrap",
                "w:outlineLvl",
                "w:divId",
                "w:cnfStyle",
                "w:rPr",
                "w:sectPr",
                "w:pPrChange",
            ),
        )
    level = numbering_properties.find(qn("w:ilvl"))
    if level is None:
        level = OxmlElement("w:ilvl")
        numbering_properties.append(level)
    level.set(qn("w:val"), "0")
    numbering_id = numbering_properties.find(qn("w:numId"))
    if numbering_id is None:
        numbering_id = OxmlElement("w:numId")
        numbering_properties.append(numbering_id)
    numbering_id.set(qn("w:val"), "0")


def _rich_text_plain_text(node: Mapping[str, object], depth: int = 0) -> str:
    if depth > 12:
        raise MedicalWritingDocumentDocxExportError("rich-text nesting exceeds 12 levels")
    node_type = str(node.get("type") or "")
    if node_type == "text":
        value = node.get("text")
        if not isinstance(value, str):
            raise MedicalWritingDocumentDocxExportError("rich-text text node has no text value")
        return value
    if node_type == "hardBreak":
        return "\n"
    if node_type not in {
        "doc",
        "paragraph",
        "heading",
        "bulletList",
        "orderedList",
        "listItem",
    }:
        raise MedicalWritingDocumentDocxExportError(
            f"unsupported rich-text node: {node_type or '<empty>'}"
        )
    content = node.get("content", [])
    if not isinstance(content, list) or any(not isinstance(child, Mapping) for child in content):
        raise MedicalWritingDocumentDocxExportError("rich-text content is invalid")
    child_text = [_rich_text_plain_text(child, depth + 1) for child in content]
    return "\n".join(child_text) if node_type == "doc" else "".join(child_text)


def _render_rich_text_root(
    output: Document,
    root: Mapping[str, object],
    block: Mapping[str, object],
    *,
    citation_plan: Optional[_CitationPlan] = None,
    index_plan: Optional[_DocumentIndexPlan] = None,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
) -> int:
    root_type = str(root.get("type") or "")
    if root_type == "doc":
        children = root.get("content", [])
        if not isinstance(children, list) or not children:
            raise MedicalWritingDocumentDocxExportError("rich-text document has no block content")
        rendered = 0
        for child in children:
            if not isinstance(child, Mapping) or child.get("type") == "doc":
                raise MedicalWritingDocumentDocxExportError("rich-text document contains an invalid block")
            rendered += _render_rich_text_root(
                output,
                child,
                block,
                citation_plan=citation_plan,
                index_plan=index_plan,
                style_presets=style_presets,
            )
        return rendered
    if root_type in {"paragraph", "heading"}:
        _render_rich_text_paragraph(
            output,
            root,
            block=block,
            citation_plan=citation_plan,
            index_plan=index_plan,
            style_presets=style_presets,
        )
        return 1
    if root_type not in {"bulletList", "orderedList"}:
        raise MedicalWritingDocumentDocxExportError(
            f"unsupported rich-text root node: {root_type or '<empty>'}"
        )
    list_style = "List Bullet" if root_type == "bulletList" else "List Number"
    rendered = 0
    items = root.get("content", [])
    if not isinstance(items, list):
        raise MedicalWritingDocumentDocxExportError("rich-text list content is invalid")
    for item in items:
        if not isinstance(item, Mapping) or item.get("type") != "listItem":
            raise MedicalWritingDocumentDocxExportError("rich-text list contains a non-list-item node")
        children = item.get("content", [])
        if not isinstance(children, list) or not children:
            raise MedicalWritingDocumentDocxExportError("rich-text list item has no paragraph")
        for child in children:
            if not isinstance(child, Mapping) or child.get("type") not in {"paragraph", "heading"}:
                raise MedicalWritingDocumentDocxExportError(
                    "rich-text nested lists are not supported for DOCX export"
                )
            paragraph = _render_rich_text_paragraph(
                output,
                child,
                block=block,
                citation_plan=citation_plan,
                index_plan=index_plan,
                style_presets=style_presets,
            )
            if list_style in output.styles:
                paragraph.style = list_style
            rendered += 1
    return rendered


def _render_rich_text_paragraph(
    output: Document,
    node: Mapping[str, object],
    *,
    block: Mapping[str, object],
    citation_plan: Optional[_CitationPlan] = None,
    index_plan: Optional[_DocumentIndexPlan] = None,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
):
    node_type = str(node.get("type") or "")
    attrs = node.get("attrs") or {}
    if not isinstance(attrs, Mapping):
        raise MedicalWritingDocumentDocxExportError("rich-text paragraph attributes are invalid")
    if node_type == "heading":
        level = attrs.get("level")
        if not isinstance(level, int) or isinstance(level, bool) or level not in range(1, 7):
            outline = block.get("outline_level")
            level = outline + 1 if isinstance(outline, int) and 0 <= outline <= 5 else 1
        default_preset = f"heading_{min(level, 4)}"
    else:
        default_preset = "body"
    preset_name = str(attrs.get("stylePreset") or default_preset)
    if preset_name not in style_presets:
        raise MedicalWritingDocumentDocxExportError(
            f"unsupported rich-text style preset: {preset_name}"
        )
    paragraph = output.add_paragraph()
    preset = _apply_paragraph_preset(
        paragraph,
        preset_name,
        attrs,
        style_presets=style_presets,
    )
    content = node.get("content", [])
    if not isinstance(content, list):
        raise MedicalWritingDocumentDocxExportError("rich-text paragraph content is invalid")
    for child in content:
        if not isinstance(child, Mapping):
            raise MedicalWritingDocumentDocxExportError("rich-text inline node is invalid")
        child_type = child.get("type")
        if child_type == "hardBreak":
            paragraph.add_run().add_break()
            continue
        if child_type != "text":
            raise MedicalWritingDocumentDocxExportError(
                f"unsupported rich-text inline node: {child_type or '<empty>'}"
            )
        marks = child.get("marks", [])
        if not isinstance(marks, list) or any(not isinstance(mark, Mapping) for mark in marks):
            raise MedicalWritingDocumentDocxExportError("rich-text marks are invalid")
        _render_text_run(
            paragraph,
            str(child.get("text") or ""),
            marks,
            preset,
            citation_plan=citation_plan,
            index_plan=index_plan,
        )
    return paragraph


def _apply_paragraph_preset(
    paragraph,
    preset_name: str,
    attrs: Mapping[str, object],
    *,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
):
    preset = dict(style_presets[preset_name])
    heading_level = preset.get("heading_level")
    if heading_level:
        style_name = f"Heading {heading_level}"
        if style_name in paragraph.part.document.styles:
            paragraph.style = style_name
    alignment = attrs.get("textAlign")
    if alignment is not None:
        if alignment not in _PARAGRAPH_ALIGNMENTS:
            raise MedicalWritingDocumentDocxExportError(
                f"unsupported rich-text paragraph alignment: {alignment}"
            )
        paragraph.alignment = _PARAGRAPH_ALIGNMENTS[str(alignment)]
    paragraph_format = paragraph.paragraph_format
    paragraph_format.keep_with_next = bool(preset.get("keep_with_next"))
    paragraph_format.line_spacing = _numeric_attr(
        attrs,
        "lineHeight",
        float(preset["line_height"]),
        minimum=0.8,
        maximum=3.0,
    )
    paragraph_format.space_before = Pt(_numeric_attr(
        attrs,
        "spacingBeforePt",
        float(preset["spacing_before_pt"]),
        minimum=0.0,
        maximum=72.0,
    ))
    paragraph_format.space_after = Pt(_numeric_attr(
        attrs,
        "spacingAfterPt",
        float(preset["spacing_after_pt"]),
        minimum=0.0,
        maximum=72.0,
    ))
    base_font_size = float(preset["font_size_pt"])
    paragraph_format.left_indent = Pt(base_font_size * _numeric_attr(
        attrs, "leftIndentChars", 0.0, minimum=0.0, maximum=20.0
    ))
    paragraph_format.right_indent = Pt(base_font_size * _numeric_attr(
        attrs, "rightIndentChars", 0.0, minimum=0.0, maximum=20.0
    ))
    first_line_indent_chars = _numeric_attr(
        attrs,
        "firstLineIndentChars",
        float(preset["first_line_indent_chars"]),
        minimum=-10.0,
        maximum=20.0,
    )
    paragraph_format.first_line_indent = Pt(
        base_font_size * first_line_indent_chars
    )
    _set_character_first_line_indent(paragraph, first_line_indent_chars)
    return preset


def _set_character_first_line_indent(paragraph, indent_chars: float) -> None:
    """Persist Word's character-relative indent beside python-docx's twip value."""
    paragraph_properties = paragraph._p.get_or_add_pPr()
    indentation = paragraph_properties.find(qn("w:ind"))
    if indentation is None:
        indentation = OxmlElement("w:ind")
        paragraph_properties.append(indentation)
    indentation.attrib.pop(qn("w:firstLineChars"), None)
    indentation.attrib.pop(qn("w:hangingChars"), None)
    if indent_chars >= 0:
        indentation.set(qn("w:firstLineChars"), str(round(indent_chars * 100)))
    else:
        indentation.set(qn("w:hangingChars"), str(round(abs(indent_chars) * 100)))


def _numeric_attr(
    attrs: Mapping[str, object],
    key: str,
    default: float,
    *,
    minimum: float,
    maximum: float,
) -> float:
    value = attrs.get(key)
    if value is None:
        return default
    if not isinstance(value, (int, float)) or isinstance(value, bool):
        raise MedicalWritingDocumentDocxExportError(f"rich-text {key} must be numeric")
    numeric = float(value)
    if numeric < minimum or numeric > maximum:
        raise MedicalWritingDocumentDocxExportError(
            f"rich-text {key} must be between {minimum:g} and {maximum:g}"
        )
    return numeric


def _render_text_run(
    paragraph,
    text: str,
    marks: Sequence[Mapping[str, object]],
    preset,
    *,
    citation_plan: Optional[_CitationPlan] = None,
    index_plan: Optional[_DocumentIndexPlan] = None,
):
    citation_marks = [mark for mark in marks if mark.get("type") == "citation"]
    if len(citation_marks) > 1:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text text node contains duplicate citation marks"
        )
    cross_reference_marks = [
        mark for mark in marks if mark.get("type") == "crossReference"
    ]
    if len(cross_reference_marks) > 1:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text text node contains duplicate cross-reference marks"
        )
    if citation_marks and cross_reference_marks:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text text node cannot be both a citation and a cross-reference"
        )
    if cross_reference_marks:
        target_kind, target_id = _cross_reference_target_from_mark(
            cross_reference_marks[0]
        )
        entry = (
            index_plan.by_object_key.get((target_kind, target_id))
            if index_plan is not None
            else None
        )
        if entry is None:
            raise MedicalWritingDocumentDocxExportError(
                f"cross-reference target is not available for rendering: {target_kind}/{target_id}"
            )
        visible_marks = [
            mark for mark in marks if mark.get("type") != "crossReference"
        ]
        prefix = "表 " if target_kind == "table" else "图 "
        _add_formatted_run(paragraph, prefix, visible_marks, preset)
        return _append_word_field(
            paragraph,
            f" REF {entry.bookmark_name} \\h ",
            str(entry.number),
            dirty=True,
            bold=bool(preset.get("bold")),
            font_size_pt=float(preset["font_size_pt"]),
        )
    citation_reference_ids: List[str] = []
    rendered_text = text
    if citation_marks:
        raw_reference_ids = _citation_reference_ids_from_mark(citation_marks[0])
        citation_reference_ids = list(dict.fromkeys(
            citation_plan.aliases_by_reference_id.get(reference_id, reference_id)
            if citation_plan
            else reference_id
            for reference_id in raw_reference_ids
        ))
    elif citation_plan and is_superscript_reference_marker(
        {"type": "text", "text": text, "marks": list(marks)}
    ):
        source_numbers = parse_legacy_reference_marker(text)
        citation_reference_ids = list(
            dict.fromkeys(
                citation_plan.legacy_number_to_reference_id[number]
                for number in source_numbers
                if number in citation_plan.legacy_number_to_reference_id
            )
        )

    if citation_reference_ids:
        if citation_plan is None or any(
            reference_id not in citation_plan.numbers_by_reference_id
            for reference_id in citation_reference_ids
        ):
            raise MedicalWritingDocumentDocxExportError(
                "cited reference is not available for rendering: "
                + ", ".join(citation_reference_ids)
            )
        citation_numbers = [
            citation_plan.numbers_by_reference_id[reference_id]
            for reference_id in citation_reference_ids
        ]
        if len(citation_reference_ids) > 1:
            last_run = _add_formatted_run(paragraph, "[", marks, preset)
            for index, (reference_id, number) in enumerate(
                zip(citation_reference_ids, citation_numbers)
            ):
                if index:
                    last_run = _add_formatted_run(paragraph, ",", marks, preset)
                last_run = _add_formatted_run(paragraph, str(number), marks, preset)
                _wrap_run_in_hyperlink(
                    paragraph,
                    last_run,
                    citation_plan.bookmarks_by_reference_id[reference_id],
                )
            return _add_formatted_run(paragraph, "]", marks, preset)
        rendered_text = f"[{citation_numbers[0]}]"

    run = _add_formatted_run(paragraph, rendered_text, marks, preset)
    if citation_reference_ids:
        _wrap_run_in_hyperlink(
            paragraph,
            run,
            citation_plan.bookmarks_by_reference_id[citation_reference_ids[0]],
        )
    return run


def _cross_reference_target_from_mark(
    mark: Mapping[str, object],
) -> tuple[str, str]:
    attrs = mark.get("attrs")
    if not isinstance(attrs, Mapping) or set(attrs) != {"targetKind", "targetId"}:
        raise MedicalWritingDocumentDocxExportError(
            "rich-text cross-reference attributes are invalid"
        )
    target_kind = str(attrs.get("targetKind") or "").strip()
    target_id = str(attrs.get("targetId") or "").strip()
    if target_kind not in {"table", "figure"} or not re.fullmatch(
        r"[A-Za-z][A-Za-z0-9_-]{0,127}", target_id
    ):
        raise MedicalWritingDocumentDocxExportError(
            "rich-text cross-reference target is invalid"
        )
    return target_kind, target_id


def _add_formatted_run(
    paragraph,
    text: str,
    marks: Sequence[Mapping[str, object]],
    preset,
):
    run = paragraph.add_run(text)
    run.bold = bool(preset.get("bold"))
    run.font.size = Pt(float(preset["font_size_pt"]))
    table_docx._set_run_fonts(
        run,
        east_asia=str(preset["east_asia"]),
        latin=str(preset["latin"]),
    )
    run.font.color.rgb = RGBColor(0, 0, 0)
    seen_marks = set()
    for mark in marks:
        mark_type = str(mark.get("type") or "")
        if mark_type in seen_marks:
            raise MedicalWritingDocumentDocxExportError(
                f"duplicate rich-text mark: {mark_type or '<empty>'}"
            )
        seen_marks.add(mark_type)
        if mark_type == "bold":
            run.bold = True
        elif mark_type == "italic":
            run.italic = True
        elif mark_type == "underline":
            run.underline = True
        elif mark_type == "superscript":
            run.font.superscript = True
        elif mark_type == "subscript":
            run.font.subscript = True
        elif mark_type == "textStyle":
            _apply_text_style_mark(run, mark)
        elif mark_type == "highlight":
            attrs = mark.get("attrs") or {}
            if not isinstance(attrs, Mapping):
                raise MedicalWritingDocumentDocxExportError("rich-text highlight attributes are invalid")
            color = str(attrs.get("color") or "").upper()
            if color not in _HIGHLIGHT_COLORS:
                raise MedicalWritingDocumentDocxExportError(
                    f"unsupported rich-text highlight color: {color or '<empty>'}"
                )
            run.font.highlight_color = _HIGHLIGHT_COLORS[color]
        elif mark_type == "citation":
            run.font.superscript = True
        else:
            raise MedicalWritingDocumentDocxExportError(
                f"unsupported rich-text mark: {mark_type or '<empty>'}"
            )
    return run


def _wrap_run_in_hyperlink(paragraph, run, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    hyperlink.set(qn("w:history"), "1")
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def _render_reference_list(
    output: Document,
    citation_plan: _CitationPlan,
    *,
    include_heading: bool,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
    reserved_bookmark_ids: Sequence[int] = (),
) -> int:
    rendered = 0
    used_bookmark_ids = set(reserved_bookmark_ids)
    used_bookmark_ids.update({
        int(value)
        for node in output.element.iter(qn("w:bookmarkStart"))
        if (value := node.get(qn("w:id"))) is not None and value.isdigit()
    })
    next_bookmark_id = max(used_bookmark_ids, default=0) + 1
    if include_heading:
        heading = output.add_paragraph()
        preset = _apply_paragraph_preset(
            heading,
            "heading_1",
            {},
            style_presets=style_presets,
        )
        _render_text_run(heading, "参考文献", [], preset)
        rendered += 1

    for reference in citation_plan.ordered_references:
        number = citation_plan.numbers_by_reference_id[reference.reference_id]
        paragraph = output.add_paragraph()
        preset = _apply_paragraph_preset(
            paragraph,
            "body",
            {
                "firstLineIndentChars": -2,
                "leftIndentChars": 2,
            },
            style_presets=style_presets,
        )
        bookmark_id = str(next_bookmark_id)
        next_bookmark_id += 1
        bookmark_start = OxmlElement("w:bookmarkStart")
        bookmark_start.set(qn("w:id"), bookmark_id)
        bookmark_start.set(
            qn("w:name"),
            citation_plan.bookmarks_by_reference_id[reference.reference_id],
        )
        paragraph._p.append(bookmark_start)
        _render_text_run(
            paragraph,
            _format_citation_reference(number, reference),
            [],
            preset,
        )
        bookmark_end = OxmlElement("w:bookmarkEnd")
        bookmark_end.set(qn("w:id"), bookmark_id)
        paragraph._p.append(bookmark_end)
        rendered += 1
    return rendered


def _format_citation_reference(number: int, reference: _CitationReference) -> str:
    if reference.legacy is not None:
        raw_text = reference.legacy.raw_text.strip()
        body = re.sub(
            r"^\s*(?:\[\s*\d{1,4}\s*\]|\d{1,4}[.、])\s*",
            "",
            raw_text,
            count=1,
        )
        return f"[{number}] {body}"
    if reference.managed is None:
        raise MedicalWritingDocumentDocxExportError(
            f"reference has no renderable source: {reference.reference_id}"
        )
    return _format_gbt_7714_2015_reference(number, reference.managed)


def _format_gbt_7714_2015_reference(
    number: int,
    reference: MedicalWritingProjectReference,
) -> str:
    authors = [item.strip() for item in reference.authors if item.strip()]
    if not authors:
        author_text = "佚名"
    elif len(authors) <= 3:
        author_text = ", ".join(authors)
    else:
        suffix = "等" if any(re.search(r"[\u4e00-\u9fff]", item) for item in authors[:3]) else "et al"
        author_text = f"{', '.join(authors[:3])}, {suffix}"

    title = reference.title.strip() or "题名缺失"
    journal = reference.journal.strip()
    if journal:
        source = f"{title}[J]. {journal}"
        publication_parts: List[str] = []
        if reference.year.strip():
            publication_parts.append(reference.year.strip())
        volume_issue = reference.volume.strip()
        if reference.issue.strip():
            volume_issue += f"({reference.issue.strip()})"
        if volume_issue:
            publication_parts.append(volume_issue)
        publication = ", ".join(publication_parts)
        if reference.pages.strip():
            publication = (
                f"{publication}: {reference.pages.strip()}"
                if publication
                else reference.pages.strip()
            )
        if publication:
            source += f", {publication}"
        source += "."
    else:
        source = f"{title}[EB/OL]."

    identifiers: List[str] = []
    if reference.doi.strip():
        identifiers.append(f"DOI: {reference.doi.strip()}")
    elif reference.url.strip():
        identifiers.append(reference.url.strip())
    if identifiers:
        source += " " + ". ".join(identifiers) + "."
    return f"[{number}] {author_text}. {source}"


def _apply_text_style_mark(run, mark: Mapping[str, object]) -> None:
    attrs = mark.get("attrs") or {}
    if not isinstance(attrs, Mapping):
        raise MedicalWritingDocumentDocxExportError("rich-text textStyle attributes are invalid")
    unknown = set(attrs) - {"fontFamily", "fontSize", "color"}
    if unknown:
        raise MedicalWritingDocumentDocxExportError(
            f"unsupported rich-text textStyle attributes: {', '.join(sorted(unknown))}"
        )
    font_family = attrs.get("fontFamily")
    if font_family:
        family = str(font_family).strip()
        if not family or len(family) > 80:
            raise MedicalWritingDocumentDocxExportError("rich-text font family is invalid")
        # The editor may retain a requested family for on-screen authoring,
        # but the regulatory DOCX export has one fixed bilingual font pair.
        table_docx._set_run_fonts(
            run,
            east_asia="宋体",
            latin="Times New Roman",
        )
    font_size = attrs.get("fontSize")
    if font_size:
        match = _FONT_SIZE_RE.fullmatch(str(font_size).strip())
        if not match:
            raise MedicalWritingDocumentDocxExportError("rich-text font size must use pt units")
        points = float(match.group(1))
        if points < 6 or points > 72:
            raise MedicalWritingDocumentDocxExportError("rich-text font size is outside 6-72 pt")
        run.font.size = Pt(points)
    color = attrs.get("color")
    if color:
        value = str(color).strip().upper()
        if not _HEX_COLOR_RE.fullmatch(value):
            raise MedicalWritingDocumentDocxExportError("rich-text text color must be #RRGGBB")
        run.font.color.rgb = RGBColor.from_string(value[1:])


def _front_matter_context(
    document: ProtocolDocument,
    *,
    overrides: Optional[Mapping[str, str]] = None,
) -> Optional[_FrontMatterContext]:
    for section in document.sections:
        if (
            str(getattr(section, "node_kind", "") or "") != "front_matter"
            and str(getattr(section, "template_node_id", "") or "")
            != "ich_m11_front_matter"
        ):
            continue
        if not any(
            str(block.get("source_kind") or "").startswith("greenfield_")
            for block in section.content_blocks
        ):
            continue
        values: Dict[str, str] = {}
        block_ids: set[str] = set()
        for block in section.content_blocks:
            block_id = str(block.get("block_id") or "").strip()
            if block_id:
                block_ids.add(block_id)
            if str(block.get("block_type") or "") != "table":
                continue
            structure = block.get("structured_table")
            if not isinstance(structure, Mapping):
                continue
            rows = structure.get("rows")
            if not isinstance(rows, list):
                continue
            for row in rows:
                if not isinstance(row, Mapping):
                    continue
                cells = row.get("cells")
                if not isinstance(cells, list):
                    continue
                texts = [
                    str(cell.get("text") or "").strip()
                    for cell in cells
                    if isinstance(cell, Mapping)
                ]
                label = str(row.get("label") or "").strip()
                if not label and texts:
                    label = texts[0]
                value = texts[1] if len(texts) >= 2 else ""
                if label:
                    values[label] = value
        authoritative = {
            key: str(value or "").strip()
            for key, value in (overrides or {}).items()
            if key
            in {
                "document_title",
                "study_phase",
                "investigational_product",
                "protocol_date",
                "sponsor",
            }
            and str(value or "").strip()
        }
        return _FrontMatterContext(
            block_ids=frozenset(block_ids),
            document_title=authoritative.get(
                "document_title",
                values.get("方案标题", "").strip(),
            ),
            study_phase=authoritative.get(
                "study_phase",
                values.get("研究分期", "").strip(),
            ),
            investigational_product=authoritative.get(
                "investigational_product",
                values.get("研究药物", "").strip(),
            ),
            protocol_date=authoritative.get(
                "protocol_date",
                values.get("版本日期", "").strip(),
            ),
            sponsor=authoritative.get(
                "sponsor",
                values.get("申办者", "").strip() or _CMS_SPONSOR_NAME,
            ),
        )
    return None


def _render_protocol_cover(
    output: Document,
    *,
    document: ProtocolDocument,
    context: _FrontMatterContext,
) -> int:
    """Render the editable company cover instead of a numbered front-matter chapter."""

    rendered = 0
    spacer = output.add_paragraph()
    spacer.paragraph_format.space_after = Pt(18)
    rendered += 1

    cover_label = output.add_paragraph()
    if "Subtitle" in output.styles:
        cover_label.style = "Subtitle"
    cover_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_label.paragraph_format.space_before = Pt(8)
    cover_label.paragraph_format.space_after = Pt(36)
    cover_label_run = cover_label.add_run("研究方案")
    cover_label_run.bold = True
    cover_label_run.italic = False
    cover_label_run.font.size = Pt(15)
    cover_label_run.font.color.rgb = RGBColor(0, 0, 0)
    table_docx._set_run_fonts(
        cover_label_run,
        east_asia="宋体",
        latin="Times New Roman",
    )
    rendered += 1

    title = output.add_paragraph()
    if "Title" in output.styles:
        title.style = "Title"
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(12)
    title.paragraph_format.space_after = Pt(72)
    title_run = title.add_run(
        context.document_title
        or f"{document.protocol_id}临床试验方案"
    )
    title_run.bold = True
    title_run.font.size = Pt(14)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    table_docx._set_run_fonts(
        title_run,
        east_asia="宋体",
        latin="Times New Roman",
    )
    rendered += 1

    metadata_rows = [
        ("研究药物：", context.investigational_product),
        ("研究阶段：", context.study_phase),
        ("方案/研究编号：", document.protocol_id),
        ("版本号：", document.version),
        (
            "版本日期：",
            (context.protocol_date or "").strip()
            or (getattr(document, "protocol_date", "") or "").strip()
            or date.today().isoformat(),
        ),
        ("申办者：", context.sponsor),
    ]
    metadata = output.add_table(rows=len(metadata_rows), cols=2)
    metadata.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata.autofit = False
    table_docx._set_table_borders_none(metadata)
    for row, (label, value) in zip(metadata.rows, metadata_rows):
        row.cells[0].width = Inches(1.55)
        row.cells[1].width = Inches(4.45)
        for cell, text in zip(row.cells, (label, value)):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(1.5)
            paragraph.paragraph_format.space_after = Pt(1.5)
            paragraph.paragraph_format.line_spacing = 1.5
            run = paragraph.add_run(text)
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0, 0, 0)
            table_docx._set_run_fonts(
                run,
                east_asia="宋体",
                latin="Times New Roman",
            )
    rendered += len(metadata_rows)

    confidentiality = output.add_paragraph()
    confidentiality.paragraph_format.space_before = Pt(84)
    confidentiality.paragraph_format.line_spacing = 1.5
    confidentiality.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    label = confidentiality.add_run("保密声明：")
    label.bold = True
    body = confidentiality.add_run(
        "本文件中的信息保密，所有权归申办者。根据中国法律法规要求，"
        "本文件仅提供给研究者、伦理委员会、监管管理部门等相关机构审阅。"
        "除为研究实施所必需并获得申办者书面授权外，任何个人不得分发、"
        "复制、发表或披露本文件的任何内容。"
    )
    for run in (label, body):
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0, 0, 0)
        table_docx._set_run_fonts(
            run,
            east_asia="宋体",
            latin="Times New Roman",
        )
    rendered += 1
    return rendered


def _configure_core_properties(
    output: Document,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
) -> None:
    properties = output.core_properties
    properties.title = f"{document.protocol_id} {document.version}".strip()
    properties.subject = (
        "医学写作批准终稿"
        if mode == MedicalWritingDocumentExportMode.APPROVED_FINAL
        else "医学写作草稿预览"
    )
    properties.identifier = document.document_id
    properties.author = "CMS AI医学经理工作台"
    properties.last_modified_by = "CMS AI医学经理工作台"
    properties.comments = approval_reference
    properties.created = table_docx._FIXED_CORE_TIMESTAMP
    properties.modified = table_docx._FIXED_CORE_TIMESTAMP


def _configure_default_styles(
    output: Document,
    *,
    style_presets: Mapping[str, Mapping[str, object]] = _STYLE_PRESETS,
    enable_company_numbering: bool = False,
) -> None:
    """Set company-standard default fonts on the Normal style.

    Company templates (D005, D017, RUX, D001) all use 宋体 (SimSun) for
    Chinese body text and Times New Roman for Latin text.  python-docx
    defaults to Calibri, which does not match the company formatting
    authority for greenfield documents rendered from scratch.
    """
    style = output.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(10.5)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    style.font.color.rgb = RGBColor(0, 0, 0)

    for style_name, size, bold in (
        ("Title", 14.0, True),
        ("Subtitle", 14.0, True),
        ("TOC Heading", 16.0, True),
    ):
        if style_name not in output.styles:
            continue
        target = output.styles[style_name]
        _configure_word_style_font(
            target,
            size_pt=size,
            bold=bold,
            color="000000",
        )
        paragraph_properties = target.element.find(qn("w:pPr"))
        if paragraph_properties is not None:
            paragraph_border = paragraph_properties.find(qn("w:pBdr"))
            if paragraph_border is not None:
                paragraph_properties.remove(paragraph_border)
        if style_name == "TOC Heading":
            # Word's built-in TOC Heading is based on Heading 1. Once Heading 1
            # is bound to the protocol multilevel list, that inheritance makes
            # "目录" consume chapter number 1. Index titles are display-only and
            # must remain outside the protocol outline numbering sequence.
            target.base_style = output.styles["Normal"]

    try:
        document_control_heading = output.styles["CMS Document Control Heading"]
    except KeyError:
        document_control_heading = output.styles.add_style(
            "CMS Document Control Heading",
            WD_STYLE_TYPE.PARAGRAPH,
        )
    _configure_word_style_font(
        document_control_heading,
        size_pt=14.0,
        bold=True,
        color="000000",
    )
    document_control_heading.paragraph_format.keep_with_next = True
    document_control_heading.paragraph_format.line_spacing = 1.25
    document_control_heading.paragraph_format.space_before = Pt(12)
    document_control_heading.paragraph_format.space_after = Pt(6)
    document_control_heading.quick_style = True
    document_control_heading.hidden = False

    for level in range(1, 5):
        style_name = f"Heading {level}"
        if style_name not in output.styles:
            continue
        preset = style_presets[f"heading_{level}"]
        target = output.styles[style_name]
        _configure_word_style_font(
            target,
            size_pt=float(preset["font_size_pt"]),
            bold=bool(preset["bold"]),
            color="000000",
        )
        target.paragraph_format.keep_with_next = True
        target.paragraph_format.line_spacing = float(preset["line_height"])
        target.paragraph_format.space_before = Pt(
            float(preset["spacing_before_pt"])
        )
        target.paragraph_format.space_after = Pt(
            float(preset["spacing_after_pt"])
        )
        target.quick_style = True
        target.hidden = False

    for level, left_indent_pt in ((1, 0.0), (2, 21.0), (3, 42.0), (4, 63.0)):
        style_name = f"TOC {level}"
        try:
            toc_style = output.styles[style_name]
        except KeyError:
            toc_style = output.styles.add_style(
                style_name,
                WD_STYLE_TYPE.PARAGRAPH,
            )
        _configure_word_style_font(
            toc_style,
            size_pt=10.5,
            bold=False,
            color="000000",
        )
        toc_style.paragraph_format.left_indent = Pt(left_indent_pt)
        toc_style.paragraph_format.space_before = Pt(0)
        toc_style.paragraph_format.space_after = Pt(0)
        toc_style.paragraph_format.line_spacing = 1.15

    if enable_company_numbering:
        _configure_company_multilevel_heading_numbering(output)


def _configure_word_style_font(
    style,
    *,
    size_pt: float,
    bold: bool,
    color: str,
) -> None:
    style.font.name = "Times New Roman"
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:eastAsia"), "宋体")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    color_element = rpr.find(qn("w:color"))
    if color_element is None:
        color_element = OxmlElement("w:color")
        rpr.append(color_element)
    color_element.set(qn("w:val"), color)


def _configure_company_multilevel_heading_numbering(output: Document) -> int:
    """Bind Heading 1-4 to one editable Word multilevel list."""

    numbering = output.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if str(element.get(qn("w:abstractNumId")) or "").isdigit()
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
        if str(element.get(qn("w:numId")) or "").isdigit()
    ]
    abstract_num_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_num_id))
    numbering_signature = f"{0x617A5780 + abstract_num_id:08X}"[-8:]
    numbering_session_id = OxmlElement("w:nsid")
    numbering_session_id.set(qn("w:val"), numbering_signature)
    abstract.append(numbering_session_id)
    multi_level_type = OxmlElement("w:multiLevelType")
    multi_level_type.set(qn("w:val"), "multilevel")
    abstract.append(multi_level_type)
    numbering_template = OxmlElement("w:tmpl")
    numbering_template.set(qn("w:val"), numbering_signature)
    abstract.append(numbering_template)

    for level in range(4):
        style = output.styles[f"Heading {level + 1}"]
        item = OxmlElement("w:lvl")
        item.set(qn("w:ilvl"), str(level))

        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        item.append(start)
        number_format = OxmlElement("w:numFmt")
        number_format.set(qn("w:val"), "decimal")
        item.append(number_format)
        paragraph_style = OxmlElement("w:pStyle")
        paragraph_style.set(qn("w:val"), style.style_id)
        item.append(paragraph_style)
        level_text = OxmlElement("w:lvlText")
        level_text.set(
            qn("w:val"),
            ".".join(f"%{index}" for index in range(1, level + 2)),
        )
        item.append(level_text)
        level_justification = OxmlElement("w:lvlJc")
        level_justification.set(qn("w:val"), "left")
        item.append(level_justification)

        paragraph_properties = OxmlElement("w:pPr")
        indentation = OxmlElement("w:ind")
        if level < 2:
            indentation.set(qn("w:left"), "0")
            indentation.set(qn("w:firstLine"), "0")
        else:
            authoritative_indent = {
                2: ("567", "567"),
                3: ("851", "851"),
            }[level]
            indentation.set(qn("w:left"), authoritative_indent[0])
            indentation.set(qn("w:hanging"), authoritative_indent[1])
        paragraph_properties.append(indentation)
        item.append(paragraph_properties)

        run_properties = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Times New Roman")
        fonts.set(qn("w:eastAsia"), "宋体")
        fonts.set(qn("w:hAnsi"), "Times New Roman")
        fonts.set(qn("w:cs"), "Times New Roman")
        fonts.set(qn("w:hint"), "default")
        run_properties.append(fonts)
        bold = OxmlElement("w:b")
        run_properties.append(bold)
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "auto")
        run_properties.append(color)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "24")
        run_properties.append(size)
        size_cs = OxmlElement("w:szCs")
        size_cs.set(qn("w:val"), "24")
        run_properties.append(size_cs)
        item.append(run_properties)
        abstract.append(item)

    first_num_index = next(
        (
            index
            for index, element in enumerate(numbering)
            if element.tag == qn("w:num")
        ),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)
    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_num_id))
    concrete.append(abstract_reference)
    numbering.append(concrete)

    for level in range(4):
        style = output.styles[f"Heading {level + 1}"]
        paragraph_properties = style.element.find(qn("w:pPr"))
        if paragraph_properties is None:
            paragraph_properties = OxmlElement("w:pPr")
            run_properties = style.element.find(qn("w:rPr"))
            if run_properties is None:
                style.element.append(paragraph_properties)
            else:
                style.element.insert(style.element.index(run_properties), paragraph_properties)
        existing = paragraph_properties.find(qn("w:numPr"))
        if existing is not None:
            paragraph_properties.remove(existing)
        number_properties = OxmlElement("w:numPr")
        if level:
            level_element = OxmlElement("w:ilvl")
            level_element.set(qn("w:val"), str(level))
            number_properties.append(level_element)
        number_id_element = OxmlElement("w:numId")
        number_id_element.set(qn("w:val"), str(num_id))
        number_properties.append(number_id_element)
        table_docx._insert_before_first(
            paragraph_properties,
            number_properties,
            (
                "w:suppressLineNumbers",
                "w:pBdr",
                "w:shd",
                "w:tabs",
                "w:suppressAutoHyphens",
                "w:kinsoku",
                "w:wordWrap",
                "w:overflowPunct",
                "w:topLinePunct",
                "w:autoSpaceDE",
                "w:autoSpaceDN",
                "w:bidi",
                "w:adjustRightInd",
                "w:snapToGrid",
                "w:spacing",
                "w:ind",
                "w:contextualSpacing",
                "w:mirrorIndents",
                "w:suppressOverlap",
                "w:jc",
                "w:textDirection",
                "w:textAlignment",
                "w:textboxTightWrap",
                "w:outlineLvl",
                "w:divId",
                "w:cnfStyle",
                "w:rPr",
                "w:sectPr",
                "w:pPrChange",
            ),
        )
    return num_id


def _apply_native_heading_number(paragraph, parts: Sequence[int]) -> None:
    """Apply an exact editable number without depending on ancestor list state."""

    heading_level = len(parts) - 1
    if heading_level not in range(4):
        raise MedicalWritingDocumentDocxExportError(
            "native heading numbering supports Heading 1-4 only"
        )
    numbering = paragraph.part.numbering_part.element
    style_number_id = paragraph.style.element.find(
        "./w:pPr/w:numPr/w:numId",
        paragraph.style.element.nsmap,
    )
    if style_number_id is None:
        raise MedicalWritingDocumentDocxExportError(
            f"{paragraph.style.name} is not bound to company multilevel numbering"
        )
    base_num_id = style_number_id.get(qn("w:val"))
    base_num = next(
        (
            element
            for element in numbering.findall(qn("w:num"))
            if element.get(qn("w:numId")) == base_num_id
        ),
        None,
    )
    if base_num is None:
        raise MedicalWritingDocumentDocxExportError(
            "company heading numbering instance is missing"
        )
    abstract_reference = base_num.find(qn("w:abstractNumId"))
    base_abstract_id = (
        abstract_reference.get(qn("w:val")) if abstract_reference is not None else None
    )
    base_abstract = next(
        (
            element
            for element in numbering.findall(qn("w:abstractNum"))
            if element.get(qn("w:abstractNumId")) == base_abstract_id
        ),
        None,
    )
    if base_abstract is None:
        raise MedicalWritingDocumentDocxExportError(
            "company heading abstract numbering definition is missing"
        )

    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
        if str(element.get(qn("w:abstractNumId")) or "").isdigit()
    ]
    abstract_num_id = max(abstract_ids, default=-1) + 1
    abstract = copy.deepcopy(base_abstract)
    abstract.set(qn("w:abstractNumId"), str(abstract_num_id))
    numbering_signature = f"{0x627A5780 + abstract_num_id:08X}"[-8:]
    for tag in ("w:nsid", "w:tmpl"):
        signature = abstract.find(qn(tag))
        if signature is not None:
            signature.set(qn("w:val"), numbering_signature)
    multi_level_type = abstract.find(qn("w:multiLevelType"))
    if multi_level_type is not None:
        multi_level_type.set(qn("w:val"), "singleLevel")
    level_definition = abstract.find(qn("w:lvl"))
    if level_definition is None:
        raise MedicalWritingDocumentDocxExportError(
            "company heading numbering has no base level"
        )
    for extra_level in abstract.findall(qn("w:lvl"))[1:]:
        abstract.remove(extra_level)
    level_definition.set(qn("w:ilvl"), "0")
    start = level_definition.find(qn("w:start"))
    level_text = level_definition.find(qn("w:lvlText"))
    if start is None or level_text is None:
        raise MedicalWritingDocumentDocxExportError(
            "company heading numbering base level is incomplete"
        )
    start.set(qn("w:val"), str(parts[-1]))
    level_text.set(
        qn("w:val"),
        ".".join([*(str(part) for part in parts[:-1]), "%1"]),
    )
    paragraph_style = level_definition.find(qn("w:pStyle"))
    if paragraph_style is not None:
        paragraph_style.set(qn("w:val"), paragraph.style.style_id)
    first_num_index = next(
        (
            index
            for index, element in enumerate(numbering)
            if element.tag == qn("w:num")
        ),
        len(numbering),
    )
    numbering.insert(first_num_index, abstract)

    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
        if str(element.get(qn("w:numId")) or "").isdigit()
    ]
    num_id = max(num_ids, default=0) + 1
    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(num_id))
    concrete_reference = OxmlElement("w:abstractNumId")
    concrete_reference.set(qn("w:val"), str(abstract_num_id))
    concrete.append(concrete_reference)
    numbering.append(concrete)

    paragraph_properties = paragraph._p.get_or_add_pPr()
    existing = paragraph_properties.find(qn("w:numPr"))
    if existing is not None:
        paragraph_properties.remove(existing)
    number_properties = OxmlElement("w:numPr")
    level_element = OxmlElement("w:ilvl")
    level_element.set(qn("w:val"), "0")
    number_properties.append(level_element)
    number_id_element = OxmlElement("w:numId")
    number_id_element.set(qn("w:val"), str(num_id))
    number_properties.append(number_id_element)
    table_docx._insert_before_first(
        paragraph_properties,
        number_properties,
        (
            "w:suppressLineNumbers",
            "w:pBdr",
            "w:shd",
            "w:tabs",
            "w:suppressAutoHyphens",
            "w:kinsoku",
            "w:wordWrap",
            "w:overflowPunct",
            "w:topLinePunct",
            "w:autoSpaceDE",
            "w:autoSpaceDN",
            "w:bidi",
            "w:adjustRightInd",
            "w:snapToGrid",
            "w:spacing",
            "w:ind",
            "w:contextualSpacing",
            "w:mirrorIndents",
            "w:suppressOverlap",
            "w:jc",
            "w:textDirection",
            "w:textAlignment",
            "w:textboxTightWrap",
            "w:outlineLvl",
            "w:divId",
            "w:cnfStyle",
            "w:rPr",
            "w:sectPr",
            "w:pPrChange",
        ),
    )


def _set_layout_table_horizontal_rule(
    word_table,
    *,
    edge: str,
    size: str = "6",
    color: str = "000000",
) -> None:
    table_docx._set_table_borders_none(word_table)
    properties = word_table._tbl.tblPr
    borders = properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        properties.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "0")
    border.set(qn("w:color"), color)


def _add_plain_run(
    paragraph,
    text: str,
    *,
    size_pt: float,
    bold: bool = False,
    color: str = "000000",
):
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor.from_string(color)
    table_docx._set_run_fonts(
        run,
        east_asia="宋体",
        latin="Times New Roman",
    )
    return run


def _configure_header(
    section,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
    approval_reference: str,
) -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = Pt(1)
    context = _front_matter_context(document)
    phase = context.study_phase if context is not None else ""
    sponsor = context.sponsor if context is not None else _CMS_SPONSOR_NAME
    document_label = f"{phase}临床试验方案" if phase else "临床试验方案"

    table = header.add_table(rows=1, cols=2, width=Inches(6.3))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_layout_table_horizontal_rule(table, edge="bottom")
    left, right = table.rows[0].cells
    left.width = Inches(3.15)
    right.width = Inches(3.15)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
        cell.paragraphs[0].paragraph_format.space_before = Pt(0)
        cell.paragraphs[0].paragraph_format.space_after = Pt(1.5)

    left_paragraph = left.paragraphs[0]
    _add_plain_run(left_paragraph, document_label, size_pt=9)
    _add_plain_run(
        left_paragraph,
        f"\n{document.protocol_id}",
        size_pt=9,
    )

    right_paragraph = right.paragraphs[0]
    right_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_plain_run(right_paragraph, sponsor, size_pt=9)
    if mode == MedicalWritingDocumentExportMode.DRAFT_PREVIEW:
        _add_plain_run(
            right_paragraph,
            "\n草稿预览 · DRAFT PREVIEW · 不可用于提交",
            size_pt=8,
            color="C00000",
        )
    elif mode == MedicalWritingDocumentExportMode.APPROVED_FINAL:
        _add_plain_run(
            right_paragraph,
            "\n方案终稿",
            size_pt=8,
            color="595959",
        )


def _configure_footer(
    section,
    document: ProtocolDocument,
    mode: MedicalWritingDocumentExportMode,
) -> None:
    """Add company-standard page-number footer to a greenfield section.

    Company templates (D005, D017, RUX, D001) all include page-number
    footers ("第 X 页 共 Y 页") using PAGE/NUMPAGES fields.  Greenfield
    documents rendered from scratch must match this convention rather
    than producing a package with zero footer parts.
    """

    def _add_field_run(paragraph, instruction_text: str, fallback_text: str):
        """Build a properly run-wrapped OOXML field (not dirty).

        Uses ``\\* MERGEFORMAT`` to match the company-template convention
        so LibreOffice and Word both render the live page number.
        """
        run_begin = paragraph.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        run_begin._r.append(fld_begin)
        table_docx._set_run_fonts(run_begin, east_asia="宋体", latin="Times New Roman")
        run_begin.font.size = Pt(9)

        run_instr = paragraph.add_run()
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = f" {instruction_text} \\* MERGEFORMAT "
        run_instr._r.append(instr)
        table_docx._set_run_fonts(run_instr, east_asia="宋体", latin="Times New Roman")
        run_instr.font.size = Pt(9)

        run_sep = paragraph.add_run()
        fld_sep = OxmlElement("w:fldChar")
        fld_sep.set(qn("w:fldCharType"), "separate")
        run_sep._r.append(fld_sep)
        table_docx._set_run_fonts(run_sep, east_asia="宋体", latin="Times New Roman")
        run_sep.font.size = Pt(9)

        run_val = paragraph.add_run(fallback_text)
        run_val.font.size = Pt(9)
        table_docx._set_run_fonts(run_val, east_asia="宋体", latin="Times New Roman")

        run_end = paragraph.add_run()
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run_end._r.append(fld_end)
        table_docx._set_run_fonts(run_end, east_asia="宋体", latin="Times New Roman")
        run_end.font.size = Pt(9)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.clear()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    table = footer.add_table(rows=1, cols=2, width=Inches(6.3))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    _set_layout_table_horizontal_rule(table, edge="top")
    left, right = table.rows[0].cells
    left.width = Inches(3.15)
    right.width = Inches(3.15)
    for cell in (left, right):
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        cell.paragraphs[0].paragraph_format.space_before = Pt(1.5)
        cell.paragraphs[0].paragraph_format.space_after = Pt(0)

    left_paragraph = left.paragraphs[0]
    _add_plain_run(
        left_paragraph,
        f"研究方案.{document.protocol_id}\n版本号/{document.version}",
        size_pt=8,
    )
    if mode == MedicalWritingDocumentExportMode.DRAFT_PREVIEW:
        _add_plain_run(
            left_paragraph,
            "\n草稿预览 · DRAFT PREVIEW · 不可用于提交",
            size_pt=7.5,
            color="C00000",
        )

    page_paragraph = right.paragraphs[0]
    page_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _add_plain_run(page_paragraph, "第 ", size_pt=9)
    _add_field_run(page_paragraph, "PAGE", "1")
    _add_plain_run(page_paragraph, " 页 共 ", size_pt=9)
    _add_field_run(page_paragraph, "NUMPAGES", "1")
    _add_plain_run(page_paragraph, " 页", size_pt=9)


def _looks_like_markdown_table(text: str) -> bool:
    lines = text.splitlines()
    return any(_MARKDOWN_SEPARATOR_RE.match(line) for line in lines)


def _is_table_caption(
    block: Mapping[str, object],
    next_block: Optional[Mapping[str, object]],
) -> bool:
    if not next_block or next_block.get("block_type") != "table":
        return False
    text = str(block.get("text") or "").strip()
    if not text:
        return False
    structure = next_block.get("structured_table")
    structured_title = (
        str(structure.get("title") or "").strip()
        if isinstance(structure, Mapping)
        else ""
    )
    if structured_title and _normalized_visible_text(text) == _normalized_visible_text(
        structured_title
    ):
        return True
    if str(block.get("block_type") or "").strip().lower() == "heading":
        return True
    return bool(re.match(r"^(?:附表|表)\s*[A-Za-z0-9一二三四五六七八九十]*", text))


def _normalized_visible_text(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).casefold()


def _normalized_note_text(value: str) -> str:
    text = str(value or "").strip()
    text = re.sub(r"^(?:附注|备注|注)\s*[:：]?\s*", "", text)
    text = re.sub(r"^[（(]?[A-Za-z0-9]+[）)]?\s*[.、:]?\s+", "", text)
    return _normalized_visible_text(text)


def _deduplicated_notes(
    notes: Sequence[StructuredTableNote],
) -> List[StructuredTableNote]:
    result: List[StructuredTableNote] = []
    seen = set()
    for note in notes:
        key = (note.marker.strip().casefold(), _normalized_note_text(note.text))
        if key in seen:
            continue
        seen.add(key)
        result.append(note)
    return result


def medical_writing_document_export_snapshot_digest(
    document: ProtocolDocument,
    *,
    front_matter_overrides: Optional[Mapping[str, str]] = None,
) -> str:
    normalized_overrides = {
        key: str(value or "").strip()
        for key, value in (front_matter_overrides or {}).items()
        if str(value or "").strip()
    }
    if not normalized_overrides:
        payload = document.model_dump_json(exclude_none=False)
    else:
        payload = json.dumps(
            {
                "document": document.model_dump(mode="json", exclude_none=False),
                "front_matter_overrides": normalized_overrides,
            },
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def _preview_display_units(value: str) -> int:
    """Return a stable, renderer-independent width estimate for preview text."""

    units = 0
    for character in str(value or ""):
        if character.isspace():
            units += 1
        elif unicodedata.east_asian_width(character) in {"W", "F"}:
            units += 2
        else:
            units += 1
    return max(1, units)


def _preview_block_text(block: Mapping[str, object]) -> str:
    text = str(block.get("text") or "").strip()
    if text:
        return text
    block_type = str(block.get("block_type") or "").strip().lower()
    if block_type == "table":
        return str(block.get("title") or block.get("caption") or "结构化表格").strip()
    if block_type in {"figure", "image", "appendix_image"}:
        return str(block.get("title") or block.get("alt_text") or block_type).strip()
    return block_type or "内容块"


def _preview_style_preset(
    block: Mapping[str, object],
    style_presets: Mapping[str, Mapping[str, object]],
) -> Mapping[str, object]:
    block_type = str(block.get("block_type") or "").strip().lower()
    outline = block.get("outline_level")
    level = outline + 1 if isinstance(outline, int) and 0 <= outline <= 8 else 1
    name = f"heading_{min(level, 4)}" if block_type == "heading" else "body"
    return style_presets.get(name, style_presets["body"])


def _preview_estimated_lines(
    block: Mapping[str, object],
    *,
    available_width_inches: float,
    style_presets: Mapping[str, Mapping[str, object]],
) -> int:
    block_type = str(block.get("block_type") or "").strip().lower()
    if block_type == "table":
        table = block.get("table") or block.get("structured_table") or {}
        rows = table.get("rows") if isinstance(table, Mapping) else None
        row_count = len(rows) if isinstance(rows, Sequence) and not isinstance(rows, (str, bytes)) else 1
        return max(2, row_count + 1)
    if block_type in {"figure", "image", "appendix_image"}:
        return max(8, int(block.get("page_count") or 1) * 12)
    preset = _preview_style_preset(block, style_presets)
    font_size = max(8.0, float(preset.get("font_size_pt") or 10.5))
    # 2.0 display units per point is intentionally conservative for CJK body text.
    chars_per_line = max(8, int(max(1.0, available_width_inches) * 72 * 2.0 / font_size))
    lines = str(block.get("text") or "").splitlines() or [""]
    return max(
        1,
        sum(max(1, math.ceil(_preview_display_units(line) / chars_per_line)) for line in lines),
    )


def build_medical_writing_fast_preview(
    document: ProtocolDocument,
    *,
    mode: MedicalWritingDocumentExportMode | str = MedicalWritingDocumentExportMode.DRAFT_PREVIEW,
    front_matter_overrides: Optional[Mapping[str, str]] = None,
) -> MedicalWritingDocumentPreview:
    """Build a deterministic source-bound page estimate without claiming Word layout.

    The preview intentionally shares the exact export snapshot digest and style
    profile resolution with DOCX export, while keeping pagination explicitly
    estimated. It is therefore safe for fast in-app review but never a proof of
    final Microsoft Word pagination. In particular, the estimate does not attempt
    to reproduce Word-generated cover, index/TOC, field-update, section, or table
    pagination; callers must use the explicit page_count_basis contract and a
    Word/PDF receipt for submission-grade page counts.
    """

    export_mode = MedicalWritingDocumentExportMode(mode)
    style_profile = _resolve_style_profile(document)
    snapshot_sha256 = medical_writing_document_export_snapshot_digest(
        document,
        front_matter_overrides=front_matter_overrides,
    )
    ordered_blocks = _ordered_blocks(document)
    visual_plan = _build_visual_layout_plan(
        ordered_blocks,
        style_profile.body_layout,
        style_presets=style_profile.style_presets,
    )

    pages: List[dict[str, object]] = []
    block_records: List[MedicalWritingPreviewBlock] = []
    current_page: Optional[dict[str, object]] = None
    current_used_pt = 0.0
    current_layout = style_profile.body_layout

    def ensure_page(layout: tuple[object, ...]) -> None:
        nonlocal current_page, current_used_pt, current_layout
        if current_page is None:
            orientation, (width, height), _margins = layout
            current_page = {
                "page_number": len(pages) + 1,
                "orientation": str(orientation),
                "width_twips": int(width),
                "height_twips": int(height),
                "section_ids": [],
                "block_ids": [],
                "estimated": True,
                "layout_note": "页面尺寸与分页为 StyleProfile 同源快速估算",
            }
            pages.append(current_page)
            current_used_pt = 0.0
            current_layout = layout

    def close_page() -> None:
        nonlocal current_page, current_used_pt
        current_page = None
        current_used_pt = 0.0

    for ordered_index, (section_index, _block_index, block) in enumerate(ordered_blocks):
        section = document.sections[section_index]
        layout = visual_plan.target_layout_by_index.get(
            ordered_index,
            style_profile.body_layout,
        )
        if current_page is not None and layout != current_layout:
            close_page()
        ensure_page(layout)
        _orientation, (page_width, page_height), margins = layout
        available_width_inches = max(
            1.0,
            (int(page_width) - int(margins[0]) - int(margins[1])) / 1440,
        )
        available_height_pt = max(
            72.0,
            (int(page_height) - int(margins[2]) - int(margins[3])) / 1440 * 72,
        )
        block_type = str(block.get("block_type") or "paragraph").strip().lower()
        block_id = str(block.get("block_id") or f"block-{ordered_index + 1}")
        lines = _preview_estimated_lines(
            block,
            available_width_inches=available_width_inches,
            style_presets=style_profile.style_presets,
        )
        preset = _preview_style_preset(block, style_profile.style_presets)
        line_height_pt = max(
            8.0,
            float(preset.get("font_size_pt") or 10.5)
            * float(preset.get("line_height") or 1.0),
        )
        spacing_pt = float(preset.get("spacing_before_pt") or 0.0) + float(
            preset.get("spacing_after_pt") or 0.0
        )
        if block_type in {"figure", "image", "appendix_image"}:
            block_height_pt = max(96.0, lines * line_height_pt)
            layout_note = "图像/量表块按独立视觉页估算；最终分页须经 Word 验证"
        elif block_type == "table":
            block_height_pt = max(36.0, lines * max(11.0, line_height_pt * 0.9))
            layout_note = "表格行高和跨页为快速估算；最终分页须经 Word 验证"
        else:
            block_height_pt = max(12.0, lines * line_height_pt + spacing_pt)
            layout_note = "文字换行与行高为快速估算"
        page_break_before = bool(block.get("page_break_before"))
        if page_break_before and current_page and current_page["block_ids"]:
            close_page()
            ensure_page(layout)
        page_start = int(current_page["page_number"]) if current_page else len(pages) + 1
        remaining_pt = max(1.0, available_height_pt - current_used_pt)
        page_span = max(1, math.ceil(block_height_pt / max(1.0, available_height_pt)))
        if current_page and current_used_pt > 0 and block_height_pt > remaining_pt:
            close_page()
            ensure_page(layout)
            page_start = int(current_page["page_number"])
        page_end = page_start + page_span - 1
        # Keep a deterministic page record for a block that spans multiple pages.
        for page_number in range(page_start, page_end + 1):
            while len(pages) < page_number:
                close_page()
                ensure_page(layout)
            page = pages[page_number - 1]
            if section.section_id not in page["section_ids"]:
                page["section_ids"].append(section.section_id)
            if block_id not in page["block_ids"]:
                page["block_ids"].append(block_id)
        current_page = pages[page_end - 1]
        current_layout = layout
        current_used_pt = block_height_pt % available_height_pt
        if current_used_pt == 0:
            current_used_pt = available_height_pt
        block_records.append(
            MedicalWritingPreviewBlock(
                block_id=block_id,
                section_id=section.section_id,
                block_type=block_type or "paragraph",
                page_start=page_start,
                page_end=page_end,
                estimated_lines=lines,
                text_preview=_preview_block_text(block)[:600],
                layout_note=layout_note,
            )
        )

    if not pages:
        ensure_page(style_profile.body_layout)
    return MedicalWritingDocumentPreview(
        project_id=document.project_id,
        document_id=document.document_id,
        mode=export_mode.value,
        preview_status="fast_preview",
        snapshot_sha256=snapshot_sha256,
        style_profile_id=style_profile.style_profile_id,
        style_profile_version=style_profile.style_profile_version,
        page_count_basis="style_profile_estimate",
        page_count=len(pages),
        pages=[MedicalWritingPreviewPage(**page) for page in pages],
        blocks=block_records,
        warning=(
            "当前页数是 StyleProfile 快速估算，不等同于 DOCX/ Microsoft Word 原生页数；"
            "封面、目录/索引、域更新、分页、表格跨页和版式节均可能造成差异。"
            "需要申报级页数与最终版式时，必须对同一快照执行 Word verification。"
        ),
    )


def build_medical_writing_word_verified_preview(
    document: ProtocolDocument,
    receipt: MedicalWritingWordVerificationReceipt,
    *,
    expected_docx_sha256: str,
    mode: MedicalWritingDocumentExportMode | str = MedicalWritingDocumentExportMode.DRAFT_PREVIEW,
    front_matter_overrides: Optional[Mapping[str, str]] = None,
) -> MedicalWritingDocumentPreview:
    """Promote a preview only when an external Word/PDF receipt is exact.

    This function does not run Word or PDF rendering. It verifies the receipt's
    source/document identity against the current immutable snapshot and returns
    an explicit non-estimated page contract only for already-produced evidence.
    """

    export_mode = MedicalWritingDocumentExportMode(mode)
    if receipt.project_id != document.project_id or receipt.document_id != document.document_id:
        raise MedicalWritingDocumentDocxExportError(
            "Word verification receipt does not belong to the current document"
        )
    snapshot_sha256 = medical_writing_document_export_snapshot_digest(
        document,
        front_matter_overrides=front_matter_overrides,
    )
    if receipt.source_snapshot_sha256 != snapshot_sha256:
        raise MedicalWritingDocumentDocxExportError(
            "Word verification receipt is stale for the current document snapshot"
        )
    if not re.fullmatch(r"[0-9a-f]{64}", str(expected_docx_sha256 or "")):
        raise MedicalWritingDocumentDocxExportError(
            "expected DOCX verification hash must be lowercase SHA-256"
        )
    if receipt.docx_sha256 != expected_docx_sha256:
        raise MedicalWritingDocumentDocxExportError(
            "Word verification receipt DOCX hash does not match the exported document"
        )
    style_profile = _resolve_style_profile(document)
    pages = [
        MedicalWritingPreviewPage(
            page_number=evidence.page_number,
            orientation=evidence.orientation,
            width_twips=evidence.width_twips,
            height_twips=evidence.height_twips,
            estimated=False,
            layout_note=(
                "Microsoft Word 打开、更新域、保存、重开并导出 PDF 后的页面证据；"
                f"PDF 页 hash={evidence.pdf_page_sha256}"
            ),
        )
        for evidence in receipt.page_evidence
    ]
    return MedicalWritingDocumentPreview(
        project_id=document.project_id,
        document_id=document.document_id,
        mode=export_mode.value,
        preview_status="word_verified",
        snapshot_sha256=snapshot_sha256,
        word_verified_snapshot_sha256=receipt.source_snapshot_sha256,
        style_profile_id=style_profile.style_profile_id,
        style_profile_version=style_profile.style_profile_version,
        page_count_basis="microsoft_word_receipt",
        page_count=receipt.page_count,
        pages=pages,
        blocks=[],
        warning=(
            "当前页数和版式证据来自 Microsoft Word/PDF receipt；"
            "若源快照或 DOCX hash 变化，必须重新验证。"
        ),
    )


def _document_snapshot_digest(document: ProtocolDocument) -> str:
    return medical_writing_document_export_snapshot_digest(document)


def _fill_source_grid_holes(
    table: StructuredTable,
    *,
    expected_holes: set[tuple[int, int]] | None,
) -> StructuredTable:
    candidate = table.model_copy(deep=True)
    columns = sorted(candidate.columns, key=lambda item: (item.order, item.column_id))
    rows = sorted(candidate.rows, key=lambda item: (item.order, item.row_id))
    column_index = {column.column_id: index for index, column in enumerate(columns)}
    occupancy: List[List[bool]] = [[False for _ in columns] for _ in rows]

    for row_index, row in enumerate(rows):
        for cell in row.cells:
            if table_docx._is_covered_merge_continuation(cell):
                continue
            start_column = column_index[cell.column_id]
            row_end = row_index + cell.row_span
            column_end = start_column + cell.column_span
            if row_end > len(rows) or column_end > len(columns):
                raise table_docx.StructuredTableDocxExportError(
                    f"cell {cell.cell_id} merge range exceeds the table grid"
                )
            for occupied_row in range(row_index, row_end):
                for occupied_column in range(start_column, column_end):
                    if occupancy[occupied_row][occupied_column]:
                        raise table_docx.StructuredTableDocxExportError(
                            f"cell {cell.cell_id} overlaps another cell at row "
                            f"{occupied_row}, column {occupied_column}"
                        )
                    occupancy[occupied_row][occupied_column] = True

    actual_holes = {
        (row_index, column_index)
        for row_index, occupied_row in enumerate(occupancy)
        for column_index, occupied in enumerate(occupied_row)
        if not occupied
    }
    if expected_holes is not None and actual_holes != expected_holes:
        raise table_docx.StructuredTableDocxExportError(
            "structured table grid holes differ from the preserved original DOCX grid"
        )

    for row_index, missing_column in sorted(actual_holes):
        row = rows[row_index]
        if not occupancy[row_index][missing_column]:
            column = columns[missing_column]
            token = hashlib.sha256(
                f"{candidate.table_id}|{row.row_id}|{column.column_id}|export-gap".encode(
                    "utf-8"
                )
            ).hexdigest()[:16]
            row.cells.append(
                StructuredTableCell(
                    cell_id=f"mwcell_export_gap_{token}",
                    row_id=row.row_id,
                    column_id=column.column_id,
                    text="",
                    semantic_value={
                        "_medical_writing_table": {
                            "generated": True,
                            "export_only_grid_gap": True,
                        }
                    },
                    style_role=row.style_role,
                )
            )
            occupancy[row_index][missing_column] = True
    return StructuredTable.model_validate(candidate.model_dump(mode="python"))


def _preserved_source_grid_holes(
    block: Mapping[str, object],
) -> set[tuple[int, int]] | None:
    if block.get("source_kind") != "original_protocol_docx":
        return None
    structure = block.get("structured_table")
    if not isinstance(structure, Mapping):
        return None
    raw = structure.get("source_grid_holes")
    if not isinstance(raw, list):
        return None
    holes: set[tuple[int, int]] = set()
    for position in raw:
        if (
            not isinstance(position, list)
            or len(position) != 2
            or any(
                not isinstance(value, int)
                or isinstance(value, bool)
                or value < 0
                for value in position
            )
        ):
            return None
        holes.add((position[0], position[1]))
    return holes


def _is_original_source_version_zero(block: Mapping[str, object]) -> bool:
    if block.get("source_kind") != "original_protocol_docx":
        return False
    structure = block.get("structured_table")
    if not isinstance(structure, Mapping):
        return True
    version = structure.get("version", 0)
    return isinstance(version, int) and not isinstance(version, bool) and version == 0
