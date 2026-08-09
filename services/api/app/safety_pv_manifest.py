from __future__ import annotations

import re
from copy import deepcopy
from collections import Counter, defaultdict
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from threading import Lock
from typing import Dict, Iterable, List, Optional, Tuple

from openpyxl import load_workbook

from packages.contracts.workbench_contracts import (
    RiskSeverity,
    SafetyListingDomainSummary,
    SafetyPvManifestResult,
    SafetyQualityGate,
    SafetySignalCandidate,
    SafetySourceDocumentSummary,
    SafetySourcePackageSummary,
)


MY009_LISTING = Path(
    "/Users/smkzw/Documents/朗来项目资料/MY009治疗UC/S1安全性评价-202604/MY009-UC-2-01-MM Listing_20260410_Comparison.xlsx"
)
MY009_PREVIOUS_LISTING = Path(
    "/Users/smkzw/Documents/朗来项目资料/MY009治疗UC/S1安全性评价-202604/MY009-UC-2-01-MM Listing_20260408(已自动还原).xlsx"
)
MY009_S1_ROOT = Path("/Users/smkzw/Documents/朗来项目资料/MY009治疗UC/S1安全性评价-202604")
MY009_DSUR_DOC = Path("/Users/smkzw/Documents/朗来项目资料/MY009治疗UC/DSUR/附件1：MY009_DSUR#4_资料收集 to CPM、RA、医学-MM.docx")
RUX_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/"
    "RUX-03-002_列表_数据集_Excel_20250612_处理后.xlsx"
)
RUX_PV_ROOT = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/2-RUX-03-002-现场核查项目层面文件目录-20260424/22.PV计划"
)
RUX_274_ROOT = Path("/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/RA/2.7.4/2-7-4临床安全性总结")


@dataclass(frozen=True)
class SafetyPackageConfig:
    package_id: str
    package_label: str
    project_code: str
    source_root_label: str
    package_role: str
    listing_file: Optional[Path]
    document_paths: Tuple[Path, ...]


DEFAULT_SAFETY_PACKAGES = [
    SafetyPackageConfig(
        package_id="my009_uc_s1",
        package_label="MY009 UC S1安全资料与医学复核listing",
        project_code="MY009-UC-2-01",
        source_root_label="MY009 UC S1安全性评价包",
        package_role="安全资料医学复核样本",
        listing_file=MY009_LISTING,
        document_paths=(MY009_S1_ROOT, MY009_DSUR_DOC),
    ),
    SafetyPackageConfig(
        package_id="rux_03_002_pv",
        package_label="RUX-03-002 PV计划与临床安全性总结",
        project_code="RUX-03-002",
        source_root_label="RUX-03-002 PV与2.7.4安全资料",
        package_role="PV边界与监管安全总结样本",
        listing_file=RUX_LISTING,
        document_paths=(RUX_PV_ROOT, RUX_274_ROOT),
    ),
]

SAFETY_PACKAGE_IDS_BY_PROJECT = {
    "proj_my009_uc": {"my009_uc_s1"},
    "proj_rux_03_002": {"rux_03_002_pv"},
}


class SafetyPvManifestService:
    def __init__(self, packages: Optional[List[SafetyPackageConfig]] = None):
        self.packages = packages or DEFAULT_SAFETY_PACKAGES
        self._listing_scan_cache: Dict[
            str,
            tuple[tuple[object, ...], List[SafetyListingDomainSummary], Dict[str, object], List[str]],
        ] = {}
        self._listing_scan_lock = Lock()

    def build_manifest(self, project_id: str, force_refresh: bool = False) -> SafetyPvManifestResult:
        generated_at = datetime.now(timezone.utc)
        package_ids = SAFETY_PACKAGE_IDS_BY_PROJECT.get(project_id)
        if package_ids is None:
            raise KeyError(f"safety_pv not configured for project: {project_id}")
        package_summaries = [
            self._build_package(config, force_refresh=force_refresh)
            for config in self.packages
            if config.package_id in package_ids
        ]
        return SafetyPvManifestResult(
            project_id=project_id,
            generated_at=generated_at,
            package_count=len(package_summaries),
            total_listing_domains=sum(len(package.listing_domains) for package in package_summaries),
            total_documents=sum(len(package.documents) for package in package_summaries),
            total_signal_candidates=sum(len(package.signal_candidates) for package in package_summaries),
            quality_gate_count=sum(len(package.quality_gates) for package in package_summaries),
            packages=package_summaries,
            parser_notes=[
                "P0仅生成安全资料清单、listing域清单、信号候选和PV协同质量门，不替代正式药物警戒系统或监管递交流程。",
                "安全候选均为待医学/PV确认内容；系统不做正式报告性判断。",
                "Listing解析使用确定性字段识别；PDF/PPTX/RTF等资料当前仅登记文件级清单。",
            ],
        )

    def _build_package(
        self,
        config: SafetyPackageConfig,
        *,
        force_refresh: bool = False,
    ) -> SafetySourcePackageSummary:
        warnings: List[str] = []
        listing_domains, listing_metrics, listing_warnings = self._cached_listing_scan(
            config,
            force_refresh=force_refresh,
        )
        warnings.extend(listing_warnings)
        documents = _scan_documents(config, warnings)
        candidates = _build_signal_candidates(config, listing_metrics, listing_domains, documents)
        quality_gates = _build_quality_gates(config, listing_domains, documents, candidates, listing_metrics)
        return SafetySourcePackageSummary(
            package_id=config.package_id,
            package_label=config.package_label,
            project_code=config.project_code,
            source_root_label=config.source_root_label,
            package_role=config.package_role,
            listing_domains=listing_domains,
            documents=documents,
            signal_candidates=candidates,
            quality_gates=quality_gates,
            parser_warnings=warnings,
        )

    def _cached_listing_scan(
        self,
        config: SafetyPackageConfig,
        *,
        force_refresh: bool,
    ) -> tuple[List[SafetyListingDomainSummary], Dict[str, object], List[str]]:
        cache_key = _listing_cache_key(config)
        with self._listing_scan_lock:
            cached = self._listing_scan_cache.get(config.package_id)
            if not force_refresh and cached is not None and cached[0] == cache_key:
                return deepcopy(cached[1]), deepcopy(cached[2]), list(cached[3])
            scan_warnings: List[str] = []
            domains, metrics = _scan_listing_workbook(config, scan_warnings)
            self._listing_scan_cache[config.package_id] = (
                cache_key,
                deepcopy(domains),
                deepcopy(metrics),
                list(scan_warnings),
            )
            return domains, metrics, scan_warnings


def _listing_cache_key(config: SafetyPackageConfig) -> tuple[object, ...]:
    listing_file = config.listing_file
    if listing_file is None:
        return config.package_id, "no_listing"
    resolved = listing_file.expanduser().resolve()
    try:
        stat = resolved.stat()
    except FileNotFoundError:
        return config.package_id, str(resolved), "missing"
    return (
        config.package_id,
        str(resolved),
        stat.st_size,
        stat.st_mtime_ns,
        stat.st_ctime_ns,
    )


def _scan_listing_workbook(
    config: SafetyPackageConfig,
    warnings: List[str],
) -> tuple[List[SafetyListingDomainSummary], Dict[str, object]]:
    metrics: Dict[str, object] = {
        "ae_event_count": 0,
        "ae_related_count": 0,
        "sae_count": 0,
        "lab_cs_count": 0,
        "lab_cs_by_sheet": Counter(),
        "eg_cs_count": 0,
        "eg_ae_link_count": 0,
        "medication_record_count": 0,
    }
    listing_file = config.listing_file
    if listing_file is None:
        return [], metrics
    if not listing_file.exists():
        warnings.append(f"{listing_file.name}不存在，无法读取安全listing。")
        return [], metrics

    try:
        workbook = load_workbook(listing_file, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - defensive for corrupted vendor workbook
        warnings.append(f"{listing_file.name}读取失败：{exc}")
        return [], metrics

    domains: List[SafetyListingDomainSummary] = []
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        headers, data_start_row = _listing_layout(sheet)
        row_count = max(sheet.max_row - data_start_row + 1, 0)
        if row_count <= 0:
            continue
        subject_idx = _column_index_any(headers, "USUBJID", "SUBJID")
        site_idx = _column_index(headers, "SITEID")
        domain = _sheet_domain(sheet_name)
        ae_term_idx = _column_index(headers, "AETERM")
        ae_ser_idx = _column_index(headers, "AESER")
        ae_rel_idx = _column_index(headers, "AEREL")
        lb_sig_idx = _column_index_any(headers, "LBSIG", "LBCLSIGN")
        eg_sig_idx = _column_index_any(headers, "EGSIG", "EGCLSIG")
        eg_ae_idx = _column_index_any(headers, "EGAEMH1", "EGAEMH")
        cm_term_idx = _column_index(headers, "CMTRT")
        subjects = set()
        sites = set()
        for row in sheet.iter_rows(min_row=data_start_row, values_only=True):
            if subject_idx is not None and row[subject_idx] not in (None, ""):
                subjects.add(str(row[subject_idx]).strip())
            if site_idx is not None and row[site_idx] not in (None, ""):
                sites.add(str(row[site_idx]).strip())
            if domain == "AE" and ae_term_idx is not None and _cell_text(row[ae_term_idx]):
                metrics["ae_event_count"] = int(metrics["ae_event_count"]) + 1
                if ae_ser_idx is not None and _is_yes(_cell_text(row[ae_ser_idx])):
                    metrics["sae_count"] = int(metrics["sae_count"]) + 1
                if ae_rel_idx is not None and _is_related(_cell_text(row[ae_rel_idx])):
                    metrics["ae_related_count"] = int(metrics["ae_related_count"]) + 1
            if domain.startswith("LB") and lb_sig_idx is not None and _is_clinically_significant(_cell_text(row[lb_sig_idx])):
                metrics["lab_cs_count"] = int(metrics["lab_cs_count"]) + 1
                metrics["lab_cs_by_sheet"][sheet_name] += 1
            if domain == "EG" and eg_sig_idx is not None and _is_clinically_significant(_cell_text(row[eg_sig_idx])):
                metrics["eg_cs_count"] = int(metrics["eg_cs_count"]) + 1
                if eg_ae_idx is not None and _cell_text(row[eg_ae_idx]):
                    metrics["eg_ae_link_count"] = int(metrics["eg_ae_link_count"]) + 1
            if domain in {"CM", "CM1", "JWYY", "HBYY"} and cm_term_idx is not None and _cell_text(row[cm_term_idx]):
                metrics["medication_record_count"] = int(metrics["medication_record_count"]) + 1
        domain_label, relevance = _domain_label_and_relevance(sheet_name, headers)
        if relevance == "支持性资料" and not _has_any_key_field(headers):
            continue
        domains.append(
            SafetyListingDomainSummary(
                domain_id=f"{config.package_id}:{sheet_name.lower()}",
                sheet_name=sheet_name,
                domain=_domain_code(sheet_name, headers),
                domain_label=domain_label,
                safety_relevance=relevance,
                row_count=row_count,
                subject_count=len(subjects),
                site_count=len(sites),
                key_fields=_key_fields(headers),
                parser_status="parsed",
                parser_notes=_domain_parser_notes(sheet_name, headers, row_count),
                source_package=config.package_id,
            )
        )
    return domains, metrics


def _scan_documents(config: SafetyPackageConfig, warnings: List[str]) -> List[SafetySourceDocumentSummary]:
    source_files: List[tuple[Path, Path]] = []
    for source in config.document_paths:
        if not source.exists():
            warnings.append(f"{source.name}不存在，无法登记。")
            continue
        if source.is_file():
            source_files.append((source, source.parent))
            continue
        source_files.extend(
            (path, source)
            for path in sorted(
                [
                    path
                    for path in source.rglob("*")
                    if path.is_file()
                    and path.suffix.lower() in {".docx", ".pdf", ".pptx", ".xlsx", ".xls", ".rtf", ".csv"}
                    and not path.name.startswith("~$")
                ],
                key=lambda item: str(item).lower(),
            )
        )

    documents: List[SafetySourceDocumentSummary] = []
    seen = set()
    for file_path, source_root in source_files:
        token = str(file_path)
        if token in seen:
            continue
        seen.add(token)
        document_type, role_hint, topics = _document_type(file_path)
        documents.append(
            SafetySourceDocumentSummary(
                document_id=f"{config.package_id}:doc:{_stable_token(file_path)}",
                document_type=document_type,
                project_code=config.project_code,
                public_title=_public_title(file_path),
                source_package=config.package_id,
                relative_path=_public_relative(file_path, source_root),
                file_format=file_path.suffix.lower().lstrip(".").upper(),
                parser_status="inventory_only" if file_path.suffix.lower() != ".docx" else "document_registered",
                size_bytes=file_path.stat().st_size,
                role_hint=role_hint,
                key_topics=topics,
            )
        )
    return documents


def _build_signal_candidates(
    config: SafetyPackageConfig,
    metrics: Dict[str, object],
    domains: List[SafetyListingDomainSummary],
    documents: List[SafetySourceDocumentSummary],
) -> List[SafetySignalCandidate]:
    candidates: List[SafetySignalCandidate] = []
    ae_count = int(metrics.get("ae_event_count") or 0)
    related_count = int(metrics.get("ae_related_count") or 0)
    sae_count = int(metrics.get("sae_count") or 0)
    lab_cs_count = int(metrics.get("lab_cs_count") or 0)
    eg_cs_count = int(metrics.get("eg_cs_count") or 0)
    eg_ae_link_count = int(metrics.get("eg_ae_link_count") or 0)
    medication_count = int(metrics.get("medication_record_count") or 0)

    if ae_count:
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:ae_medical_review",
                signal_type="ae_medical_review_candidate",
                signal_label="不良事件医学复核候选",
                title=f"{config.project_code} 已登记 {ae_count} 条AE，其中 {related_count} 条提示可能相关",
                severity=RiskSeverity.MEDIUM if sae_count == 0 else RiskSeverity.HIGH,
                source_package=config.package_id,
                source_domains=["AE"],
                evidence_locators=["AE sheet / AETERM, AESEV, AEREL, AESER"],
                observation=f"当前解析到SAE字段阳性 {sae_count} 条；该结果只说明listing字段状态，不构成最终安全性结论。",
                medical_pv_boundary="医学经理可复核医学解释和证据链；严重性、预期性、相关性及报告性由PV确认。",
                recommended_next_step="核对AE术语、CTCAE分级、相关性、采取措施和结局，并与安全评估报告/AE TFL交叉确认。",
            )
        )
    if lab_cs_count:
        top_labs = ", ".join(f"{sheet}:{count}" for sheet, count in metrics["lab_cs_by_sheet"].most_common(4))
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:lab_cs_explanation",
                signal_type="lab_abnormality_medical_explanation",
                signal_label="实验室异常医学解释候选",
                title=f"{config.project_code} 实验室异常有临床意义记录 {lab_cs_count} 条",
                severity=RiskSeverity.HIGH if lab_cs_count >= 50 else RiskSeverity.MEDIUM,
                source_package=config.package_id,
                source_domains=[domain.sheet_name for domain in domains if domain.sheet_name.startswith("LB")],
                evidence_locators=[f"{top_labs} / LBSIG或LBCLSIGN"],
                observation="实验室字段出现“异常有临床意义”，需要核对是否已有AE/MH/CM解释，避免把NCS异常误计入候选。",
                medical_pv_boundary="系统只提示解释缺口；是否形成AE、是否升级报告或纳入安全信号由医学/PV确认。",
                recommended_next_step="按受试者核对异常项目、日期、参考范围、AE/MH/CM链接和安全评估报告中的叙述。",
            )
        )
    if eg_cs_count:
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:ecg_cs_linkage",
                signal_type="ecg_abnormality_linkage_review",
                signal_label="心电图异常关联核对候选",
                title=f"{config.project_code} 心电图异常有临床意义 {eg_cs_count} 条，已链接AE {eg_ae_link_count} 条",
                severity=RiskSeverity.MEDIUM,
                source_package=config.package_id,
                source_domains=["EG"],
                evidence_locators=["EG表 / EGSIG或EGCLSIG、EGINDC、EGAEMH、EGAEMH1"],
                observation="部分心电图异常已链接病史或AE，仍需逐条确认链接是否充分。",
                medical_pv_boundary="医学经理可确认临床解释是否充分；正式编码、严重性和报告性由PV确认。",
                recommended_next_step="按EGINDC/EGAEMH/EGAEMH1追踪病史或AE链接，检查未链接的CS异常是否需要补充说明。",
            )
        )
    if medication_count:
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:conmed_indication_linkage",
                signal_type="conmed_indication_linkage_review",
                signal_label="合并用药适应症关联核对候选",
                title=f"{config.project_code} 合并/既往用药记录 {medication_count} 条需与MH/AE解释链对齐",
                severity=RiskSeverity.MEDIUM,
                source_package=config.package_id,
                source_domains=[domain.sheet_name for domain in domains if "用药" in domain.domain_label][:6],
                evidence_locators=["CM/CM1/JWYY/HBYY表 / CMTRT、CMINDC、CMAENO或CMAEMH"],
                observation="用药原因字段存在MH/AE链接信息，可用于发现AE/MH漏报或解释链不完整。",
                medical_pv_boundary="系统不判定禁限用药违背或AE漏报结论，只形成需医学/PV确认的证据链。",
                recommended_next_step="核对用药原因、起止日期和CMAEMH链接，必要时回流医学监查风险账本。",
            )
        )
    document_types = {document.document_type for document in documents}
    if {"dsur_collection", "safety_evaluation_report"} & document_types:
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:dsur_ib_update_sync",
                signal_type="dsur_ib_update_sync",
                signal_label="DSUR/IB安全更新资料同步候选",
                title=f"{config.project_code} 安全评估/DSUR资料包已登记，需同步核对医学叙述",
                severity=RiskSeverity.MEDIUM,
                source_package=config.package_id,
                source_domains=["安全评估报告", "DSUR资料收集"],
                evidence_locators=[doc.public_title for doc in documents if doc.document_type in {"dsur_collection", "safety_evaluation_report"}][:4],
                observation="已识别安全评估报告或DSUR资料收集表，适合形成跨文件安全叙述核对清单。",
                medical_pv_boundary="工作台仅提示叙述一致性和缺口；正式DSUR/IB更新批准流程不在本系统内完成。",
                recommended_next_step="比对AE TFL、安全评估报告和DSUR医学资料收集表，形成待医学/PV确认的更新清单。",
            )
        )
    if {"pv_plan", "clinical_safety_summary"} <= document_types:
        candidates.append(
            SafetySignalCandidate(
                signal_id=f"{config.package_id}:signal:pv_plan_274_alignment",
                signal_type="pv_plan_safety_summary_alignment",
                signal_label="PV计划与安全总结一致性核对候选",
                title=f"{config.project_code} PV计划与2.7.4安全总结均已登记",
                severity=RiskSeverity.MEDIUM,
                source_package=config.package_id,
                source_domains=["PV计划", "临床安全性总结"],
                evidence_locators=[doc.public_title for doc in documents if doc.document_type in {"pv_plan", "clinical_safety_summary"}][:4],
                observation="可用于核对安全管理计划、医学监查关注点和监管安全性总结的边界一致性。",
                medical_pv_boundary="系统不改写PV计划或正式递交文件，仅生成医学/PV协同核对项。",
                recommended_next_step="建立PV计划条款、2.7.4章节和医学监查风险类型之间的证据索引。",
            )
        )
    return candidates


def _build_quality_gates(
    config: SafetyPackageConfig,
    domains: List[SafetyListingDomainSummary],
    documents: List[SafetySourceDocumentSummary],
    candidates: List[SafetySignalCandidate],
    metrics: Dict[str, object],
) -> List[SafetyQualityGate]:
    has_listing = bool(domains)
    has_ae = any(domain.domain == "AE" for domain in domains)
    has_supporting_docs = bool(documents)
    gates = [
        SafetyQualityGate(
            gate_id=f"{config.package_id}:source_traceability",
            gate_label="来源证据链完整性",
            status="ok" if has_supporting_docs and (has_listing or config.listing_file is None) else "warning",
            owner="医学经理",
            detail="已登记安全资料文件和listing域清单；PDF/PPTX当前为文件级清单，后续如需正文级引用需接入OCR/文档解析。",
            source_refs=[document.public_title for document in documents[:5]],
        ),
        SafetyQualityGate(
            gate_id=f"{config.package_id}:listing_mapping",
            gate_label="安全listing字段映射",
            status="ok" if has_ae or config.listing_file is None else "warning",
            owner="医学/数据",
            detail="AE、实验室、心电图、合并用药、病史等域按sheet和字段名确定性识别；未识别域仅作为支持性资料。",
            source_refs=[domain.sheet_name for domain in domains[:10]],
        ),
        SafetyQualityGate(
            gate_id=f"{config.package_id}:pv_boundary",
            gate_label="PV协同确认边界",
            status="warning",
            owner="PV",
            detail="报告性判断、递交流程和药物警戒系统记录均需PV流程确认，本模块不得直接完成最终确认。",
            source_refs=["PV协同边界"],
        ),
        SafetyQualityGate(
            gate_id=f"{config.package_id}:candidate_confirmation",
            gate_label="信号候选确认状态",
            status="warning" if candidates else "ok",
            owner="医学经理/PV",
            detail=f"当前生成 {len(candidates)} 个待医学/PV确认候选；确认前不得进入医学批准或导出为最终安全性结论。",
            source_refs=[candidate.signal_label for candidate in candidates[:6]],
        ),
    ]
    docs_requiring_text = [
        document
        for document in documents
        if document.file_format in {"PDF", "PPTX", "RTF", "DOCX"} and document.parser_status != "parsed_text"
    ]
    if docs_requiring_text:
        gates.append(
            SafetyQualityGate(
                gate_id=f"{config.package_id}:document_text_locator",
                gate_label="正文级证据定位缺口",
                status="warning",
                owner="医学经理/数据",
                detail="部分PDF、PPTX、RTF或DOCX目前仅登记文件级清单；生成正式叙述前需补充页码、表格或段落级证据定位。",
                source_refs=[document.public_title for document in docs_requiring_text[:5]],
            )
        )
    if config.package_id == "my009_uc_s1" and int(metrics.get("ae_event_count") or 0):
        gates.append(
            SafetyQualityGate(
                gate_id=f"{config.package_id}:teae_epoch_unresolved",
                gate_label="AE与TEAE口径待核对",
                status="warning",
                owner="医学/统计/PV",
                detail="Listing AE条目不能自动等同TEAE口径；需与S1安全评估报告和AE TFL表格确认分析集、观察期和事件计数。",
                source_refs=["MY009 20260410 Comparison listing", "MY009 S1安全评估报告", "MY009 AE TFL"],
            )
        )
        gates.append(
            SafetyQualityGate(
                gate_id=f"{config.package_id}:denominator_scope",
                gate_label="安全性分母口径待确认",
                status="warning",
                owner="医学/统计",
                detail="全量listing受试者与S1安全性评价人群不可混用；安全性摘要需保留分析集、数据截止和观察期。",
                source_refs=["MY009 listing", "S1安全性评价报告"],
            )
        )
    if config.package_id == "rux_03_002_pv":
        gates.append(
            SafetyQualityGate(
                gate_id=f"{config.package_id}:rux_population_scope",
                gate_label="RUX安全总结人群口径待确认",
                status="warning",
                owner="医学/统计/PV",
                detail="RUX安全总结中的国内/境外、双盲/开放/全周期、年龄层和JAK关注事件分母必须分开展示，不得合并为单一安全结论。",
                source_refs=["RUX-03-002 PV计划", "RUX-03-002 2.7.4安全总结"],
            )
        )
    return gates


def _listing_layout(sheet) -> tuple[List[str], int]:
    rows = list(sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 2), values_only=True))
    if not rows:
        return [], 2
    first = [_normalize_header(value) for value in rows[0]]
    if len(rows) == 1:
        return first, 2
    second = [_normalize_header(value) for value in rows[1]]
    if _looks_like_field_code_row(second):
        return second, 3
    return first, 2


def _headers(sheet) -> List[str]:
    return _listing_layout(sheet)[0]


def _looks_like_field_code_row(values: List[str]) -> bool:
    exact_tokens = {value.upper() for value in values if value}
    if not ({"USUBJID", "SUBJID"} & exact_tokens):
        return False
    code_like = sum(bool(re.fullmatch(r"[A-Z_][A-Z0-9_]*", value)) for value in values if value)
    return code_like >= 6


def _normalize_header(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _cell_text(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _column_index(headers: List[str], needle: str) -> Optional[int]:
    for index, header in enumerate(headers):
        if needle in header:
            return index
    return None


def _column_index_any(headers: List[str], *needles: str) -> Optional[int]:
    for needle in needles:
        index = _column_index(headers, needle)
        if index is not None:
            return index
    return None


def _has_any_key_field(headers: List[str]) -> bool:
    return bool(_key_fields(headers))


def _key_fields(headers: List[str]) -> List[str]:
    tokens = [
        "USUBJID",
        "SUBJID",
        "SITEID",
        "VISIT",
        "AETERM",
        "AESEV",
        "AEREL",
        "AESER",
        "CMTRT",
        "CMINDC",
        "CMAEMH",
        "MHDAT",
        "LBTEST",
        "LBORRES",
        "LBSIG",
        "LBCLSIGN",
        "EGSIG",
        "EGCLSIG",
        "EGINDC",
        "EGAEMH",
        "VSDAT",
        "DSVER",
    ]
    selected = []
    seen = set()
    for header in headers:
        if any(token in header for token in tokens):
            field_name = _short_field_name(header)
            if field_name not in seen:
                selected.append(field_name)
                seen.add(field_name)
    return selected[:12]


def _short_field_name(header: str) -> str:
    matches = re.findall(r"\(([A-Z0-9_]+)\)", header)
    if matches:
        return matches[-1]
    return header[:36]


def _domain_code(sheet_name: str, headers: List[str]) -> str:
    return _sheet_domain(sheet_name)


def _sheet_domain(sheet_name: str) -> str:
    return sheet_name.split("--", 1)[0].strip().upper()


def _domain_label_and_relevance(sheet_name: str, headers: List[str]) -> tuple[str, str]:
    upper = _sheet_domain(sheet_name)
    if upper == "AE":
        return "不良事件", "核心安全事件"
    if upper in {"CM", "CM1", "JWYY", "HBYY"}:
        return "既往/合并用药（非试验用药）", "非试验用药与解释链"
    if upper in {"EX", "ECB", "ECA", "DA", "DAA", "DAB"}:
        return "试验药物/剂量调整", "试验药物暴露与安全解释"
    if upper == "MH":
        return "病史/UC病史", "病史与安全解释"
    if upper.startswith("LB"):
        return _lab_label(upper), "实验室安全性"
    if upper == "EG":
        return "十二导联心电图", "安全性检查"
    if upper == "VS":
        return "生命体征/身高体重", "安全性检查"
    if upper in {"DM", "DS"}:
        return "受试者状态/知情同意", "分母与状态支持"
    if upper in {"PR", "PR3", "PR4", "JWZL", "HBZL"}:
        return "既往/合并非药物治疗", "治疗暴露与解释链"
    return _page_or_form_label(headers) or sheet_name, "支持性资料"


def _lab_label(sheet_name: str) -> str:
    exact = {
        "LB": "粪便病原体检测",
        "LB1": "血常规",
        "LB2": "血生化",
        "LB3": "eGFR",
        "LB4": "凝血功能",
    }.get(sheet_name)
    if exact:
        return exact
    if "HEMA" in sheet_name:
        return "血常规"
    if "CHEM" in sheet_name:
        return "血生化"
    if "URIN" in sheet_name:
        return "尿常规"
    return "实验室检查"


def _page_or_form_label(headers: List[str]) -> str:
    for header in headers:
        if "PAGE" in header or "FORM" in header:
            return ""
    return ""


def _domain_parser_notes(sheet_name: str, headers: List[str], row_count: int) -> List[str]:
    domain = _sheet_domain(sheet_name)
    notes = [f"{row_count} 条记录"]
    if domain.startswith("LB"):
        notes.append("NCS/CS判定按LBSIG或LBCLSIGN精确区分，避免把异常无临床意义误计为候选。")
    if domain == "AE":
        notes.append("包含AETERM、CTCAE分级、相关性和AESER字段。")
    if not _key_fields(headers):
        notes.append("未识别核心安全字段，仅作为支持性资料。")
    return notes


def _is_yes(value: str) -> bool:
    return value in {"是", "Yes", "YES", "Y", "true", "True"}


def _is_related(value: str) -> bool:
    if not value:
        return False
    if "无关" in value and "可能无关" not in value:
        return False
    return "有关" in value or "相关" in value


def _is_clinically_significant(value: str) -> bool:
    if not value or "异常无临床意义" in value:
        return False
    return "异常有临床意义" in value or value.upper() == "CS"


def _document_type(file_path: Path) -> tuple[str, str, List[str]]:
    name = file_path.name
    lower = name.lower()
    if "listing" in lower or "列表" in name:
        return "listing_workbook", "安全listing/医学复核数据源", ["AE", "实验室", "合并用药", "病史"]
    if "dsur" in lower:
        return "dsur_collection", "DSUR医学资料收集", ["DSUR", "IB安全更新", "医学资料收集"]
    if "pv" in lower or "安全管理计划" in name:
        return "pv_plan", "PV计划/安全管理计划", ["PV协同", "安全管理计划", "质量门"]
    if "2-7-4" in lower or "274" in lower or "安全性总结" in name:
        return "clinical_safety_summary", "临床安全性总结", ["2.7.4", "安全性总结", "监管交付"]
    if "tfl" in lower or "ae" in lower and file_path.suffix.lower() in {".docx", ".pdf"}:
        return "ae_tfl", "AE TFL或安全输出", ["AE", "TFL", "安全分析"]
    if "安全性评估" in name or "安全评估" in name:
        return "safety_evaluation_report", "阶段性安全评估报告", ["安全评估", "AE", "实验室"]
    if file_path.suffix.lower() == ".pptx":
        return "safety_review_presentation", "安全评估汇报材料", ["安全评估", "内部沟通"]
    return "source_document", "安全资料文件", ["安全资料"]


def _public_title(file_path: Path) -> str:
    stem = file_path.stem
    return stem[:120]


def _public_relative(file_path: Path, source_root: Path) -> str:
    try:
        return file_path.relative_to(source_root).as_posix()
    except ValueError:
        return file_path.name


def _stable_token(path: Path) -> str:
    safe = re.sub(r"[^a-zA-Z0-9]+", "_", path.stem).strip("_").lower()
    return safe[:60] or "document"
