"""Render a saved working manuscript into an ordered DOCX on the clean template.

The clean TP-MA-07 template stays the base document (cover, front matter and
its styles).  Each chapter starts with its real template heading, then
paragraphs and structured tables follow in saved document order — merged
cells and table notes included.  This produces a working-draft export: real
field refresh, page-number and native-Word acceptance belong to the unified
verification stage and are explicitly not claimed here.
"""
import hashlib
import json
from typing import Any, Mapping

from docx import Document as open_docx

from app.protocol_workflow.canonical.document import document_revision_hash
from app.protocol_workflow.canonical.hashing import canonical_json

#: Template style-id (from node_tree word_rules) -> python-docx heading style.
#: The clean template numbers its heading styles 3/4/5; Normal carries body text.
STYLE_ID_TO_HEADING = {'3': 1, '4': 2, '5': 3}


def _chapter_titles(template_dir) -> dict[str, dict]:
    from pathlib import Path
    node_tree = json.loads((Path(template_dir) / 'node_tree.json').read_text(encoding='utf-8'))
    titles = {}
    for tree in ('heading_style_tree', 'outlined_tree'):
        for node in node_tree[tree]['nodes']:
            entry = titles.setdefault(node['id'], {})
            if node.get('title_zh'):
                entry['title'] = node['title_zh']
            if node.get('style_id'):
                entry['style_id'] = str(node['style_id'])
            if node.get('is_leaf_heading'):
                entry['heading_leaf'] = True
    return titles


def render_manuscript_docx(template_path, template_dir, document: Mapping[str, Any], output_path) -> dict:
    """Write the ordered working draft DOCX; return export identity hashes.

    Append-style working draft: the clean template's front matter stays
    intact and the saved chapters follow in document order.  Replacing
    template placeholder bodies section-by-section (source-preserving
    export) belongs to the 7R.4 producer productization, not this bridge.
    """
    titles = _chapter_titles(template_dir)
    doc = open_docx(str(template_path))
    doc.add_page_break()
    current_node = None
    for block in document.get('semantic_blocks', []):
        node_id = block.get('semantic_node_id')
        if node_id != current_node:
            current_node = node_id
            info = titles.get(node_id, {})
            heading_level = STYLE_ID_TO_HEADING.get(info.get('style_id', ''), 2)
            doc.add_heading(info.get('title', node_id), level=heading_level)
        kind = block.get('block_kind')
        if kind == 'paragraph':
            doc.add_paragraph(block.get('content') or '')
        elif kind == 'table':
            _add_table(doc, block.get('content') or '')
    doc.save(str(output_path))
    document_sha = document_revision_hash_sha(document)
    output_sha = hashlib.sha256(output_path if isinstance(output_path, bytes) else open(output_path, 'rb').read()).hexdigest()
    return {'document_sha256': document_sha, 'output_sha256': output_sha,
            'export_scope': 'working_draft_docx'}


def document_revision_hash_sha(document: Mapping[str, Any]) -> str:
    """Full canonical revision hash when the input is a complete revision;
    otherwise a plain content hash of the projection."""
    try:
        from packages.contracts.workbench_contracts.protocol_v3 import SemanticDocumentRevision
        return document_revision_hash(SemanticDocumentRevision.model_validate(document))
    except Exception:
        return hashlib.sha256(canonical_json(document).encode()).hexdigest()


def _add_table(doc, content: str) -> None:
    try:
        value = json.loads(content)
    except (TypeError, ValueError):
        doc.add_paragraph('（表格内容需要核对，原记录已保留）')
        return
    table = value.get('table') if isinstance(value, dict) else None
    if not isinstance(table, Mapping) or not table.get('rows'):
        doc.add_paragraph('（表格内容为空，原记录已保留）')
        return
    columns = sorted(table.get('columns', []), key=lambda c: c.get('order', 0))
    column_index = {c['column_id']: i for i, c in enumerate(columns)}
    rows = sorted(table.get('rows', []), key=lambda r: r.get('order', 0))
    header_rows = table.get('header_row_count', 0) or 0
    grid = doc.add_table(rows=len(rows), cols=len(columns))
    grid.style = 'Table Grid'
    occupied = [[False] * len(columns) for _ in rows]
    for r, row in enumerate(rows):
        cells = sorted(row.get('cells', []), key=lambda c: column_index.get(c.get('column_id'), 0))
        cursor = 0
        for cell in cells:
            while cursor < len(columns) and occupied[r][cursor]:
                cursor += 1
            if cursor >= len(columns):
                break
            height = cell.get('row_span', 1) or 1
            width = cell.get('column_span', 1) or 1
            target = grid.cell(r, cursor)
            if height > 1 or width > 1:
                target = target.merge(grid.cell(min(r + height - 1, len(rows) - 1),
                                               min(cursor + width - 1, len(columns) - 1)))
            for rr in range(r, min(r + height, len(rows))):
                for cc in range(cursor, min(cursor + width, len(columns))):
                    if rr < len(occupied) and cc < len(occupied[rr]):
                        occupied[rr][cc] = True
            target.text = cell.get('text', '')
            cursor += width
    if header_rows > 0:
        for c in range(len(columns)):
            for run in grid.rows[0].cells[c].paragraphs[0].runs:
                run.font.bold = True
    for note in table.get('notes', []):
        marker = note.get('marker', '')
        doc.add_paragraph(f'{marker} {note.get("text", "")}', style='Normal')
