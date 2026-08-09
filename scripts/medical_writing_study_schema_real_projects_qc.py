from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path
from typing import Any
from urllib.error import HTTPError
from urllib.parse import urlencode
from urllib.request import Request, urlopen

from docx import Document


PNH_SOURCE = Path(
    "/Users/smkzw/Documents/康哲项目资料/CMS-D017/4.方案/PNH/方案摘要/"
    "CMS-D017-PNH-方案摘要_v0.2.docx"
)
D017_SOURCE = Path(
    "/Users/smkzw/Documents/康哲项目资料/CMS-D017/4.方案/"
    "CMS-D017Ⅰ期方案-v1.1-20260209-clean.docx"
)
CLIENT_CONTRACT = "medical-writing-api-2026-07-17.1"


class ApiError(RuntimeError):
    pass


def file_sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class Api:
    def __init__(self, base_url: str):
        self.base_url = base_url.rstrip("/")

    def request(
        self,
        method: str,
        path: str,
        payload: dict[str, Any] | None = None,
        query: dict[str, str] | None = None,
        *,
        raw: bool = False,
    ) -> Any:
        url = self.base_url + path
        if query:
            url += "?" + urlencode(query)
        body = None
        headers = {"X-Workbench-Api-Contract": CLIENT_CONTRACT}
        if payload is not None:
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        request = Request(url, data=body, headers=headers, method=method)
        try:
            with urlopen(request, timeout=90) as response:
                data = response.read()
                if raw:
                    return data, dict(response.headers)
                return json.loads(data.decode("utf-8"))
        except HTTPError as exc:
            detail = exc.read().decode("utf-8", errors="replace")
            raise ApiError(f"{method} {path} -> {exc.code}: {detail}") from exc

    def get(self, path: str) -> Any:
        return self.request("GET", path)

    def post(self, path: str, payload: dict[str, Any]) -> Any:
        return self.request("POST", path, payload)


def pnh_source_evidence() -> dict[str, str]:
    document = Document(PNH_SOURCE)
    design = " ".join(
        cell.text.strip()
        for row_index in (9, 12, 13)
        for cell in document.tables[1].rows[row_index].cells
    )
    normalized = " ".join(design.split())
    required = [
        "多中心、随机、开放标签、平行、剂量探索II期临床研究",
        "按1:1随机分配至CMS-D017低剂量组或CMS-D017高剂量组",
        "具体剂量和给药频次待I期",
        "安全性随访期",
    ]
    missing = [item for item in required if item not in normalized]
    if missing:
        raise RuntimeError(f"PNH source no longer contains required study-design facts: {missing}")
    return {
        "source_id": f"source_real_pnh_{file_sha256(PNH_SOURCE)[:16]}",
        "locator": "docx:table:1:rows:9,12-13",
        "quote_sha256": hashlib.sha256(normalized.encode("utf-8")).hexdigest(),
        "quote_preview": normalized[:1_000],
    }


def d017_source_evidence() -> dict[str, str]:
    document = Document(D017_SOURCE)
    selected = " ".join(
        " ".join(document.paragraphs[index].text.split())
        for index in range(373, 385)
    )
    required = [
        "Part-1单次给药剂量递增（SAD）",
        "Part-2多次给药剂量递增（MAD）",
        "随机、双盲、安慰剂对照、序贯队列设计",
        "50 mg、100 mg、200 mg、400 mg、800 mg和1200 mg",
        "至少72 h",
        "至少7天",
    ]
    missing = [item for item in required if item not in selected]
    if missing:
        raise RuntimeError(f"D017 I source no longer contains required study-design facts: {missing}")
    return {
        "source_id": f"source_real_d017_i_{file_sha256(D017_SOURCE)[:16]}",
        "locator": "docx:paragraph:373-384",
        "quote_sha256": hashlib.sha256(selected.encode("utf-8")).hexdigest(),
        "quote_preview": selected[:1_000],
    }


def source_binding(evidence: dict[str, str]) -> list[dict[str, str]]:
    return [{
        "study_definition_path": "",
        "source_id": evidence["source_id"],
        "evidence_span_id": evidence["quote_sha256"][:24],
        "locator": evidence["locator"],
    }]


def confirm_protocol_assembly_plan(
    api: Api,
    project_id: str,
    definition: dict[str, Any],
    case_id: str,
    *,
    step: str = "initial",
) -> dict[str, Any]:
    try:
        current_state = api.get(
            f"/api/projects/{project_id}/medical-writing/protocol-assembly-plan"
        )
        current_plan = current_state.get("plan") or {}
        expected_plan_revision = int(current_plan.get("revision") or 0)
    except ApiError as exc:
        if "-> 404:" not in str(exc):
            raise
        expected_plan_revision = 0
    refreshed = api.post(
        f"/api/projects/{project_id}/medical-writing/protocol-assembly-plan/refresh",
        {
            "expected_plan_revision": expected_plan_revision,
            "expected_source_definition_id": definition["definition_id"],
            "expected_source_definition_revision": definition["revision"],
            "expected_source_definition_sha256": definition["state_sha256"],
            "actor": "medical_manager_qc",
            "idempotency_key": (
                f"study-schema-real-{case_id}-plan-refresh-{step}"
            ),
        },
    )
    plan = refreshed["plan"]
    blockers = [
        {
            "module_id": module["module_id"],
            "question_id": question["question_id"],
            "fact_path": question["fact_path"],
            "prompt": question["prompt"],
        }
        for module in plan["modules"]
        for question in module["unresolved_questions"]
        if question["severity"] == "blocker"
    ]
    if blockers:
        raise RuntimeError(
            f"{case_id} protocol assembly plan has unresolved blockers: "
            f"{blockers[:12]}"
        )
    confirmed = api.post(
        f"/api/projects/{project_id}/medical-writing/protocol-assembly-plan/confirm",
        {
            "expected_plan_revision": plan["revision"],
            "expected_plan_sha256": plan["state_sha256"],
            "actor": "medical_manager_qc",
            "idempotency_key": (
                f"study-schema-real-{case_id}-plan-confirm-{step}"
            ),
        },
    )
    return confirmed["plan"]


def create_project(api: Api, case: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    created = api.post("/api/projects", {
        "project_code": case["project_code"],
        "project_name": case["framing"]["document_title"],
        "indication": case["framing"]["indication"],
        "product_name": case["framing"]["investigational_product"],
        "study_phase": case["framing"]["study_phase"],
        "protocol_id": case["framing"]["protocol_id"],
        "protocol_version": case["framing"]["version"],
        "entry_mode": "from_zero",
        "actor": "medical_manager_qc",
        "idempotency_key": f"study-schema-real-{case['case_id']}-project",
    })
    project_id = created["project"]["project_id"]
    journey = created["authoring_journey"]
    framing_preview = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/impact-preview",
        {
            "expected_revision": journey["revision"],
            "stage": "framing",
            "framing": case["framing"],
        },
    )
    journey = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/stages/framing/commit",
        {
            "expected_revision": journey["revision"],
            "stage": "framing",
            "framing": case["framing"],
            "impact_preview_id": framing_preview["preview_id"],
            "actor": "medical_manager_qc",
            "idempotency_key": f"study-schema-real-{case['case_id']}-framing",
        },
    )
    picos_preview = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/impact-preview",
        {
            "expected_revision": journey["revision"],
            "stage": "picos",
            "picos": case["picos"],
        },
    )
    journey = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/stages/picos/commit",
        {
            "expected_revision": journey["revision"],
            "stage": "picos",
            "picos": case["picos"],
            "impact_preview_id": picos_preview["preview_id"],
            "actor": "medical_manager_qc",
            "idempotency_key": f"study-schema-real-{case['case_id']}-picos",
        },
    )
    missing = journey["corpus_gate"]["missing_requirements"]
    journey = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/corpus-gate/override",
        {
            "expected_revision": journey["revision"],
            "reason": "隔离验收保留全部语料缺口，仅验证真实研究设计到M11研究流程图的受治理闭环。",
            "acknowledged_missing_requirements": missing,
            "actor": "medical_manager_qc",
            "idempotency_key": f"study-schema-real-{case['case_id']}-corpus-override",
        },
    )
    template = api.get("/api/medical-writing/protocol-templates/default")
    definition = journey["study_definition"]
    assembly_plan = confirm_protocol_assembly_plan(
        api,
        project_id,
        definition,
        case["case_id"],
    )
    return project_id, {
        "journey": journey,
        "template": template,
        "assembly_plan": assembly_plan,
    }


def node(
    node_id: str,
    part_id: str,
    order: int,
    lane_order: int,
    node_kind: str,
    label: str,
    details: list[str],
    evidence: dict[str, str],
) -> dict[str, Any]:
    return {
        "node_id": node_id,
        "part_id": part_id,
        "order": order,
        "lane_order": lane_order,
        "node_kind": node_kind,
        "label": label,
        "detail_lines": details,
        "fact_status": "confirmed",
        "source_bindings": source_binding(evidence),
    }


def edge(
    edge_id: str,
    start: str,
    end: str,
    kind: str,
    label: str,
    evidence: dict[str, str],
) -> dict[str, Any]:
    return {
        "edge_id": edge_id,
        "from_node_id": start,
        "to_node_id": end,
        "edge_kind": kind,
        "label": label,
        "fact_status": "confirmed",
        "source_bindings": source_binding(evidence),
    }


def pnh_schema(proposal: dict[str, Any], evidence: dict[str, str]) -> dict[str, Any]:
    nodes = [
        node("pnh_screen", "main", 0, 0, "screening", "筛选期", ["D-56至D1给药前"], evidence),
        node("pnh_random", "main", 1, 0, "randomization", "1:1随机分配", ["开放标签、平行、剂量探索"], evidence),
        node("pnh_low", "main", 2, 0, "arm", "CMS-D017低剂量组", ["连续口服治疗12周", "具体剂量/频次待I期结果确认"], evidence),
        node("pnh_high", "main", 2, 1, "arm", "CMS-D017高剂量组", ["连续口服治疗12周", "具体剂量/频次待I期结果确认"], evidence),
        node("pnh_d84", "main", 3, 0, "treatment", "D84主要疗效评价/EOT", ["Hb较基线变化", "综合获益-风险评估"], evidence),
        node("pnh_ole", "main", 4, 0, "end", "转入长期延展研究", ["获益风险适合且知情同意"], evidence),
        node("pnh_follow", "main", 4, 1, "follow_up", "安全性随访", ["末次给药后4周 / D112"], evidence),
    ]
    edges = [
        edge("pnh_e1", "pnh_screen", "pnh_random", "participant_flow", "符合入排标准", evidence),
        edge("pnh_e2", "pnh_random", "pnh_low", "randomization", "1:1", evidence),
        edge("pnh_e3", "pnh_random", "pnh_high", "randomization", "1:1", evidence),
        edge("pnh_e4", "pnh_low", "pnh_d84", "participant_flow", "12周治疗", evidence),
        edge("pnh_e5", "pnh_high", "pnh_d84", "participant_flow", "12周治疗", evidence),
        edge("pnh_e6", "pnh_d84", "pnh_ole", "conditional", "适合且同意", evidence),
        edge("pnh_e7", "pnh_d84", "pnh_follow", "conditional", "不进入OLE", evidence),
    ]
    return {
        **proposal,
        "status": "confirmed",
        "parts": [{
            "part_id": "main",
            "order": 0,
            "label": "CMS-D017 PNH II期研究",
            "flow_direction": "left_to_right",
            "source_bindings": source_binding(evidence),
        }],
        "nodes": nodes,
        "edges": edges,
        "annotations": ["低、高剂量的具体剂量和给药频次待I期SAD/MAD、PK/PD与安全性结果确认。"],
        "state_sha256": "0" * 64,
        "updated_by": "medical_manager_qc",
    }


def d017_schema(proposal: dict[str, Any], evidence: dict[str, str]) -> dict[str, Any]:
    parts = [
        {"part_id": "sad", "order": 0, "label": "Part-1 SAD", "flow_direction": "left_to_right", "source_bindings": source_binding(evidence)},
        {"part_id": "mad", "order": 1, "label": "Part-2 MAD", "flow_direction": "left_to_right", "source_bindings": source_binding(evidence)},
    ]
    nodes = [
        node("sad_screen", "sad", 0, 0, "screening", "筛选/基线", ["D-28至D-1"], evidence),
        node("sad_random", "sad", 1, 0, "randomization", "随机分配", ["双盲、安慰剂对照"], evidence),
        node("sad_follow", "sad", 3, 0, "follow_up", "安全性随访", ["D14±1"], evidence),
        node("mad_screen", "mad", 0, 0, "screening", "筛选/基线", ["D-28至D-1"], evidence),
        node("mad_random", "mad", 1, 0, "randomization", "随机分配", ["双盲、安慰剂对照"], evidence),
        node("mad_follow", "mad", 3, 0, "follow_up", "安全性随访", ["D23±1"], evidence),
    ]
    edges: list[dict[str, Any]] = []
    sad_doses = ["50 mg", "100 mg", "200 mg", "400 mg", "800 mg", "1200 mg（备选）"]
    for index, dose in enumerate(sad_doses):
        cohort_id = f"sad_c{index + 1}"
        nodes.append(node(cohort_id, "sad", 2, index, "dose_cohort", dose, ["每组8例（CMS-D017:安慰剂=6:2）", "先入组2例哨兵"], evidence))
        edges.append(edge(f"sad_rand_{index + 1}", "sad_random", cohort_id, "randomization", "6:2", evidence))
        edges.append(edge(f"sad_follow_{index + 1}", cohort_id, "sad_follow", "follow_up", "D14±1", evidence))
        if index:
            edges.append(edge(f"sad_gate_{index}", f"sad_c{index}", cohort_id, "activation_dependency", "≥72 h安全性评估/SRC", evidence))
    mad_doses = ["50 mg BID", "100 mg BID", "200 mg BID", "400 mg QD（备选）"]
    for index, dose in enumerate(mad_doses):
        cohort_id = f"mad_c{index + 1}"
        nodes.append(node(cohort_id, "mad", 2, index, "dose_cohort", dose, ["连续给药10天", "每组10例（8:2）"], evidence))
        edges.append(edge(f"mad_rand_{index + 1}", "mad_random", cohort_id, "randomization", "8:2", evidence))
        edges.append(edge(f"mad_follow_{index + 1}", cohort_id, "mad_follow", "follow_up", "D23±1", evidence))
        if index:
            edges.append(edge(f"mad_gate_{index}", f"mad_c{index}", cohort_id, "activation_dependency", "末次给药后≥7天/SRC", evidence))
    return {
        **proposal,
        "status": "confirmed",
        "parts": parts,
        "nodes": nodes,
        "edges": edges,
        "annotations": [
            "剂量队列间虚线仅表示SRC审评后的启用依赖，不表示同一参与者跨队列流转。",
            "SAD每个队列先随机入组2例哨兵（CMS-D017与安慰剂各1例）。",
        ],
        "state_sha256": "0" * 64,
        "updated_by": "medical_manager_qc",
    }


def complete_case(api: Api, case: dict[str, Any], output_dir: Path) -> dict[str, Any]:
    project_id, state = create_project(api, case)
    proposal = api.get(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/study-schema/proposal"
    )
    schema = case["schema_builder"](proposal, case["evidence"])
    journey = api.get(f"/api/projects/{project_id}/medical-writing/authoring-journey")
    preview = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/study-schema/impact-preview",
        {"expected_journey_revision": journey["revision"], "study_schema": schema},
    )
    if any(item["severity"] == "blocker" for item in preview["issues"]):
        raise RuntimeError(f"{case['case_id']} schema preview is blocked: {preview['issues']}")
    snapshot = api.post(
        f"/api/projects/{project_id}/medical-writing/authoring-journey/study-schema/commit",
        {
            "expected_journey_revision": journey["revision"],
            "study_schema": schema,
            "impact_preview_id": preview["preview_id"],
            "reason": "医学经理依据当前原始方案逐项核对研究部分、节点、关系和时间条件后确认。",
            "actor": "medical_manager_qc",
            "idempotency_key": f"study-schema-real-{case['case_id']}-commit",
        },
    )
    journey = api.get(
        f"/api/projects/{project_id}/medical-writing/authoring-journey"
    )
    confirm_protocol_assembly_plan(
        api,
        project_id,
        journey["study_definition"],
        case["case_id"],
        step="after-study-schema",
    )
    definition = journey["study_definition"]
    template = state["template"]
    document_result = api.post(
        f"/api/projects/{project_id}/medical-writing/greenfield-document",
        {
            "protocol_id": journey["framing"]["protocol_id"],
            "version": journey["framing"]["version"],
            "document_title": journey["framing"]["document_title"],
            "indication": journey["framing"]["indication"],
            "study_phase": journey["framing"]["study_phase"],
            "source_study_definition_id": definition["definition_id"],
            "source_study_definition_revision": definition["revision"],
            "source_study_definition_sha256": definition["state_sha256"],
            "template_id": template["template_id"],
            "template_version": template["template_version"],
            "actor": "medical_manager_qc",
            "idempotency_key": f"study-schema-real-{case['case_id']}-document",
        },
    )
    document = document_result["document"]
    section_id = next(
        item["section_id"]
        for item in document["sections"]
        if item["section_number"] == "1.2"
    )
    journey = api.get(
        f"/api/projects/{project_id}/medical-writing/authoring-journey"
    )
    working_copy = api.get(
        f"/api/projects/{project_id}/medical-writing/working-copies/{section_id}"
    )
    projected = api.post(
        f"/api/projects/{project_id}/medical-writing/working-copies/{section_id}/study-schema-figure",
        {
            "expected_journey_revision": journey["revision"],
            "expected_schema_revision": snapshot["study_schema"]["revision"],
            "expected_layout_revision": snapshot["presentation"]["layout_revision"],
            "expected_working_copy_revision": working_copy["revision"],
            "actor": "medical_manager_qc",
            "reason": "医学经理确认将已核对的研究流程图插入方案第1.2节。",
            "idempotency_key": f"study-schema-real-{case['case_id']}-projection",
        },
    )
    omitted_figure_rejected = False
    try:
        api.post(
            f"/api/projects/{project_id}/medical-writing/working-copies/{section_id}",
            {
                "document_id": projected["working_copy"]["document_id"],
                "expected_revision": projected["working_copy"]["revision"],
                "content_blocks": [
                    block for block in projected["working_copy"]["content_blocks"]
                    if block.get("figure_kind") != "study_schema"
                ],
                "actor": "medical_manager_qc",
                "idempotency_key": f"study-schema-real-{case['case_id']}-forbidden-removal",
            },
        )
    except ApiError as exc:
        omitted_figure_rejected = "preserve every server-projected" in str(exc)
    if not omitted_figure_rejected:
        raise RuntimeError(f"{case['case_id']} did not reject silent figure removal")
    docx_bytes, headers = api.request(
        "GET",
        f"/api/projects/{project_id}/medical-writing/document.docx",
        query={"mode": "draft_preview"},
        raw=True,
    )
    docx_path = output_dir / f"{case['case_id']}_study_schema.docx"
    docx_path.write_bytes(docx_bytes)
    return {
        "case_id": case["case_id"],
        "project_id": project_id,
        "section_id": section_id,
        "source_path": str(case["source_path"]),
        "source_sha256": file_sha256(case["source_path"]),
        "source_evidence": case["evidence"],
        "schema_revision": snapshot["study_schema"]["revision"],
        "layout_revision": snapshot["presentation"]["layout_revision"],
        "node_count": len(snapshot["study_schema"]["nodes"]),
        "edge_count": len(snapshot["study_schema"]["edges"]),
        "formal_render_allowed": snapshot["formal_render_allowed"],
        "projection_action": projected["projection_action"],
        "silent_removal_rejected": omitted_figure_rejected,
        "docx_path": str(docx_path),
        "docx_sha256": hashlib.sha256(docx_bytes).hexdigest(),
        "export_mode": headers.get("X-Medical-Writing-Export-Mode", "draft_preview"),
    }


def cases() -> list[dict[str, Any]]:
    pnh_evidence = pnh_source_evidence()
    d017_evidence = d017_source_evidence()
    return [
        {
            "case_id": "pnh_phase2",
            "project_code": "QC-PNH-SCHEMA-REAL",
            "source_path": PNH_SOURCE,
            "evidence": pnh_evidence,
            "schema_builder": pnh_schema,
            "framing": {
                "protocol_id": "CMS-D017-PNH-II",
                "version": "V0.2",
                "document_title": "CMS-D017治疗PNH的随机开放剂量探索II期临床研究方案",
                "indication": "阵发性睡眠性血红蛋白尿症",
                "clinicaltrials_condition_term": "Paroxysmal Nocturnal Hemoglobinuria",
                "study_phase": "II期",
                "intrinsic_objectives": ["剂量探索", "概念验证（PoC）", "推荐剂量选择"],
                "investigational_product": "CMS-D017胶囊",
                "product_profile": {
                    "technology_type": "small_molecule",
                    "technology_description": "口服补体因子B小分子抑制剂",
                    "administration_routes": ["口服"],
                    "dosage_forms": ["胶囊"],
                    "exposure_scope": "systemic",
                    "device_dependency": "none",
                    "immunogenicity_relevance": "not_expected",
                    "pharmacology_considerations": ["补体旁路途径抑制"],
                    "safety_considerations": ["荚膜菌感染", "突破性溶血"],
                    "pk_pd_considerations": ["血浆PK", "补体旁路途径PD"],
                },
                "target_mechanism": "补体旁路途径因子B抑制剂",
                "competitor_target_scope": "补体因子B及其他补体旁路抑制剂",
                "development_regions": ["中国"],
                "design_pattern": "多中心、随机、开放标签、平行、低高剂量探索研究",
                "structured_design": {
                    "randomization_mode": "randomized",
                    "randomization_details": "按1:1随机分配至CMS-D017低剂量组或高剂量组。",
                    "blinding_mode": "open_label",
                    "blinding_details": "研究采用开放标签设计。",
                    "comparator_type": "active",
                    "comparator_intervention": "CMS-D017高剂量组",
                    "assignment_model": "平行组",
                    "center_model": "中国多中心",
                    "treatment_switch": {"planned": False},
                    "crossover": {"planned": False},
                    "open_label_extension": {
                        "planned": True,
                        "entry_source": "完成D84治疗结束访视的本研究参与者",
                        "entry_eligibility": "经获益风险评估适合且另行签署知情同意",
                        "treatment_regimen": "按长期延展研究方案接受CMS-D017治疗",
                        "duration": "按长期延展研究方案执行",
                        "blind_break_and_transition": "本研究为开放标签；D84后直接转入长期延展研究",
                        "long_term_objectives": ["评价长期安全性", "评价长期有效性"],
                    },
                    "sample_size_reestimation": {"planned": False},
                    "adaptive_design": {"planned": False},
                    "src_planned": False,
                    "dmc_planned": False,
                    "interim_analysis": {"planned": False},
                    "phase1_parts": [],
                    "arm_or_cohort_kind": "剂量组",
                    "arm_or_cohort_labels": [
                        "CMS-D017低剂量组",
                        "CMS-D017高剂量组",
                    ],
                    "other_design_notes": "D84后按获益风险评估和知情同意决定进入长期延展研究或安全性随访。",
                },
                "population_intent": "既往未接受补体抑制剂且存在活动性溶血的PNH成人患者",
                "key_uncertainties": ["低高剂量的具体剂量和给药频次待I期结果确认"],
            },
            "picos": {
                "design_archetype": "randomized_exploratory",
                "population_summary": "既往未接受补体抑制剂且存在活动性溶血的PNH成人患者。",
                "inclusion_modules": ["PNH克隆水平、Hb和LDH满足方案阈值", "完成荚膜菌疫苗接种或预防性抗生素"],
                "exclusion_modules": ["既往使用补体抑制剂", "活动性感染或明显骨髓衰竭"],
                "washout_rules": ["既往试验药物按28天或5个半衰期完成洗脱"],
                "intervention_summary": "CMS-D017低剂量组口服治疗12周。",
                "intervention_dose_regimen": "具体低剂量、给药频次和进食要求待I期SAD/MAD、PK/PD与安全性结果确认。",
                "allowed_concomitant_rules": ["符合稳定用药要求的营养补充剂、抗凝药及必要抗生素"],
                "required_background_rules": [],
                "prohibited_concomitant_rules": ["其他补体抑制剂及方案禁止的免疫抑制治疗"],
                "assessment_timing_restrictions": ["D84完成主要疗效评价"],
                "intervention_rules": {
                    "schema_version": "medical_writing_intervention_rules_v1",
                    "authority": "structured",
                    "ip_regimens": [
                        {
                            "regimen_id": "cms_d017_low_dose",
                            "product_name": "CMS-D017低剂量",
                            "product_role": "investigational_product",
                            "dose_and_frequency": "具体剂量和给药频次待I期结果确认",
                            "route": "口服",
                            "treatment_period": "连续治疗12周",
                        },
                        {
                            "regimen_id": "cms_d017_high_dose",
                            "product_name": "CMS-D017高剂量",
                            "product_role": "active_comparator",
                            "dose_and_frequency": "具体剂量和给药频次待I期结果确认",
                            "route": "口服",
                            "treatment_period": "连续治疗12周",
                        },
                    ],
                    "ip_adjustment_policy": "no_planned_adjustment",
                    "no_planned_adjustment_statement": "本研究不设置参与者内计划性剂量调整；安全性处置按方案规定的暂停或永久停药规则执行。",
                    "ip_action_rules": [],
                    "non_ip_treatment_rules": [
                        {
                            "rule_id": "allowed_cm_001",
                            "rule_class": "allowed_cm",
                            "policy": "allowed_if_stable",
                            "agent_or_category": "符合稳定用药要求的营养补充剂、抗凝药及必要抗生素",
                        },
                        {
                            "rule_id": "prohibited_cm_001",
                            "rule_class": "prohibited_cm",
                            "policy": "prohibited",
                            "agent_or_category": "其他补体抑制剂及方案禁止的免疫抑制治疗",
                        },
                    ],
                    "cross_object_links": [],
                },
                "comparator_summary": "CMS-D017高剂量组口服治疗12周；为活性剂量比较组，非安慰剂对照。",
                "primary_endpoint": "治疗D84时Hb水平较基线的变化值。",
                "key_secondary_endpoints": ["D14、D28、D56和D84的Hb应答与溶血控制"],
                "other_secondary_endpoints": ["输血需求、FACIT-F、LDH、网织红细胞与胆红素"],
                "exploratory_endpoints": ["PK/PD与PNH克隆相关指标"],
                "safety_endpoints": ["AE、SAE、AESI、停药AE及实验室/心电图等安全性指标"],
                "aesi_definitions": ["荚膜菌感染", "突破性溶血"],
                "study_epochs": ["筛选期", "疗效观察期", "安全性随访期"],
                "visit_strategy": "筛选D-56至D1给药前；治疗D1至D84；不进入OLE者随访至D112。",
                "estimand_strategy": "探索性比较低、高剂量组D84 Hb较基线变化及相关支持性指标。",
                "sample_size_strategy": "计划约24例，按1:1随机至低、高剂量组，每组约12例。",
                "statistical_strategy": "以描述性统计和探索性组间比较支持推荐剂量选择。",
            },
        },
        {
            "case_id": "d017_phase1",
            "project_code": "QC-D017-I-SCHEMA-REAL",
            "source_path": D017_SOURCE,
            "evidence": d017_evidence,
            "schema_builder": d017_schema,
            "framing": {
                "protocol_id": "D017-01-001",
                "version": "V1.1",
                "document_title": "CMS-D017健康参与者SAD/MAD随机双盲安慰剂对照I期临床研究方案",
                "indication": "健康参与者",
                "clinicaltrials_condition_term": "Healthy Participants",
                "study_phase": "I期",
                "intrinsic_objectives": ["首次人体（FIH）", "单次剂量递增（SAD）", "多次剂量递增（MAD）", "PK/PD表征"],
                "investigational_product": "CMS-D017胶囊",
                "product_profile": {
                    "technology_type": "small_molecule",
                    "technology_description": "口服补体因子B小分子抑制剂",
                    "administration_routes": ["口服"],
                    "dosage_forms": ["胶囊"],
                    "exposure_scope": "systemic",
                    "device_dependency": "none",
                    "immunogenicity_relevance": "not_expected",
                    "pharmacology_considerations": ["补体旁路途径抑制"],
                    "safety_considerations": ["荚膜菌感染", "QTc间期延长"],
                    "pk_pd_considerations": ["血浆PK", "尿液排泄", "补体旁路途径PD"],
                },
                "target_mechanism": "补体旁路途径因子B抑制剂",
                "competitor_target_scope": "补体因子B抑制剂健康受试者I期研究",
                "development_regions": ["中国"],
                "design_pattern": "随机、双盲、安慰剂对照、序贯剂量递增，包含Part-1 SAD和Part-2 MAD",
                "structured_design": {
                    "randomization_mode": "randomized",
                    "randomization_details": "各剂量队列内按预设比例随机至CMS-D017或匹配安慰剂。",
                    "blinding_mode": "double_blind",
                    "blinded_roles": ["参与者", "研究者", "申办方研究团队"],
                    "blinding_details": "研究采用双盲设计，SRC按预设程序审阅安全性资料。",
                    "comparator_type": "placebo",
                    "comparator_intervention": "匹配安慰剂",
                    "assignment_model": "序贯剂量队列内平行分配",
                    "center_model": "中国单中心临床研究中心",
                    "treatment_switch": {"planned": False},
                    "crossover": {"planned": False},
                    "open_label_extension": {"planned": False},
                    "sample_size_reestimation": {"planned": False},
                    "adaptive_design": {"planned": False},
                    "src_planned": True,
                    "dmc_planned": False,
                    "interim_analysis": {"planned": False},
                    "phase1_parts": [
                        {
                            "part_code": "SAD",
                            "part_label": "Part-1单次给药剂量递增",
                            "population": "中国健康成年参与者",
                            "cohort_dose": "50、100、200、400、800和1200 mg（备选）6个序贯队列",
                            "pk_pd": "单次给药PK、尿液排泄及补体旁路途径PD",
                            "safety": "给药后至少72 h安全性和耐受性评估",
                            "stopping_rules": "每队列经SRC审评后方可启动下一队列",
                            "soa_summary": "筛选D-28至D-2、基线D-1、给药D1、住院观察至D7、D14±1随访",
                            "transition_dependencies": "完成至少72 h安全性评估并经SRC同意后递增",
                        },
                        {
                            "part_code": "MAD",
                            "part_label": "Part-2多次给药剂量递增",
                            "population": "中国健康成年参与者",
                            "cohort_dose": "50 mg BID、100 mg BID、200 mg BID和400 mg QD（备选）4个序贯队列",
                            "pk_pd": "多次给药PK、蓄积特征及补体旁路途径PD",
                            "safety": "连续给药10天并完成末次给药后至少7天安全性评估",
                            "stopping_rules": "每队列经SRC审评后方可启动下一队列",
                            "soa_summary": "筛选/基线、连续给药D1至D10、D23±1随访",
                            "transition_dependencies": "末次给药后至少7天安全性资料经SRC审评后递增",
                        },
                    ],
                    "phase1_sequence": "先开展Part-1 SAD，再结合SAD安全性、PK/PD结果启动Part-2 MAD。",
                    "arm_or_cohort_kind": "序贯剂量队列",
                    "arm_or_cohort_labels": ["Part-1 SAD", "Part-2 MAD"],
                    "other_design_notes": "剂量队列间为启用依赖，不代表同一参与者跨队列流转。",
                },
                "population_intent": "中国健康成年参与者",
                "key_uncertainties": ["后续剂量组及PK/PD采样可根据SRC审评调整"],
            },
            "picos": {
                "design_archetype": "randomized_exploratory",
                "field_applicability": {
                    "estimand_strategy": {
                        "status": "not_applicable",
                        "reason": "首次人体SAD/MAD研究以描述性安全性、耐受性和PK/PD表征为主，不设置正式估计目标。",
                        "confirmed_by_medical_manager": True,
                    }
                },
                "population_summary": "18至55周岁的中国健康成年参与者。",
                "inclusion_modules": ["年龄、BMI、体重和避孕要求符合方案", "能够理解并遵守住院与随访要求"],
                "exclusion_modules": ["具有临床意义的病史、检查异常或感染风险", "近期使用禁用药物或参加其他临床研究"],
                "washout_rules": ["处方/非处方药按2周或5个半衰期洗脱；特定酶诱导剂/抑制剂按4周洗脱"],
                "intervention_summary": "Part-1 SAD单次给药6个序贯剂量队列；Part-2 MAD连续10天给药4个序贯剂量队列。",
                "intervention_dose_regimen": "SAD：50、100、200、400、800、1200 mg（备选）；MAD：50 mg BID、100 mg BID、200 mg BID、400 mg QD（备选）。",
                "allowed_concomitant_rules": ["仅允许治疗AE或医疗紧急情况所需的药物/治疗"],
                "required_background_rules": [],
                "prohibited_concomitant_rules": ["除试验用药品和允许急救用药外的其他药物及方案列明禁用治疗"],
                "assessment_timing_restrictions": ["SAD至少72 h安全性评估后SRC决定下一队列", "MAD末次给药后至少7天SRC决定下一队列"],
                "intervention_rules": {
                    "schema_version": "medical_writing_intervention_rules_v1",
                    "authority": "structured",
                    "ip_regimens": [
                        {
                            "regimen_id": "cms_d017_sad_mad",
                            "product_name": "CMS-D017",
                            "product_role": "investigational_product",
                            "dose_and_frequency": "SAD 50至1200 mg；MAD 50 mg BID至400 mg QD",
                            "route": "口服",
                            "treatment_period": "SAD单次给药；MAD连续给药10天",
                        },
                        {
                            "regimen_id": "matched_placebo",
                            "product_name": "匹配安慰剂",
                            "product_role": "placebo",
                            "dose_and_frequency": "按各队列随机比例匹配给药",
                            "route": "口服",
                            "treatment_period": "与相应CMS-D017队列一致",
                        },
                    ],
                    "ip_adjustment_policy": "no_planned_adjustment",
                    "no_planned_adjustment_statement": "同一参与者不进行计划性剂量调整或跨剂量队列递增；队列递增由SRC基于累积资料决定。",
                    "ip_action_rules": [],
                    "non_ip_treatment_rules": [
                        {
                            "rule_id": "allowed_cm_001",
                            "rule_class": "allowed_cm",
                            "policy": "allowed_with_approval",
                            "agent_or_category": "治疗AE或医疗紧急情况所必需的药物或治疗",
                        },
                        {
                            "rule_id": "prohibited_cm_001",
                            "rule_class": "prohibited_cm",
                            "policy": "prohibited",
                            "agent_or_category": "除试验用药品和允许急救用药外的其他药物及方案列明禁用治疗",
                        },
                    ],
                    "cross_object_links": [],
                },
                "comparator_summary": "各队列使用匹配安慰剂；SAD为CMS-D017:安慰剂=6:2，MAD为8:2。",
                "primary_endpoint": "单次及多次口服CMS-D017的安全性和耐受性。",
                "key_secondary_endpoints": ["CMS-D017血浆PK参数", "补体旁路途径抑制率等PD指标"],
                "other_secondary_endpoints": [],
                "exploratory_endpoints": ["浓度-QTc、PK/PD相关性和尿液排泄"],
                "safety_endpoints": ["AE、安全性实验室检查、心电图、生命体征和体格检查"],
                "aesi_definitions": ["荚膜菌感染", "QTc间期延长"],
                "study_epochs": ["筛选/基线期", "Part-1 SAD", "Part-2 MAD", "安全性随访期"],
                "visit_strategy": "SAD随访至D14±1；MAD连续给药10天并随访至D23±1。",
                "sample_size_strategy": "计划88例；SAD 6组各8例共48例，MAD 4组各10例共40例。",
                "statistical_strategy": "以描述性统计评价安全性、耐受性、PK与PD，不设正式统计假设。",
            },
        },
    ]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--api-url", required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    api = Api(args.api_url)
    results = [complete_case(api, case, args.output_dir) for case in cases()]
    report = {
        "passed": all(
            item["formal_render_allowed"]
            and item["projection_action"] == "inserted"
            and item["silent_removal_rejected"]
            for item in results
        ),
        "cases": results,
    }
    report_path = args.output_dir / "real_projects_api_qc.json"
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"report": str(report_path), **report}, ensure_ascii=False, indent=2))
    if not report["passed"]:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
