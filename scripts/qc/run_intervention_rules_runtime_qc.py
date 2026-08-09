#!/usr/bin/env python3
"""Run real-project intervention-rule API, approval, and DOCX QC.

The project manifest must have been created by
``scripts/seed_intervention_rules_qc_projects.py`` against an isolated runtime.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import urllib.error
import urllib.parse
import urllib.request
from copy import deepcopy
from pathlib import Path
from typing import Any

from docx import Document


EXPECTED = {
    "QC-RUX-IR": {
        "6.4": {
            "include": ["开放治疗期IGA为0分", "总BSA超过20%", "个体化中断给药"],
            "exclude": ["合并用药剂量变化", "替代治疗或支持治疗"],
        },
        "6.9": None,
        "6.10": None,
    },
    "QC-D001-IR": {
        "6.4": {
            "include": ["立即暂停研究药物治疗", "5个半衰期"],
            "exclude": ["合并用药剂量变化仍作为CM记录", "补救治疗自身的剂量变化"],
        },
        "6.9": {
            "include": ["系统性糖皮质激素", "第24周访视"],
            "exclude": ["方案允许的稳定剂量合并用药", "合并用药剂量变化仍作为CM记录"],
        },
        "6.10": {
            "include": ["方案允许的稳定剂量合并用药", "合并用药剂量变化仍作为CM记录"],
            "exclude": ["系统性糖皮质激素", "立即暂停研究药物治疗"],
        },
    },
    "QC-PNH-IR": {
        "6.4": {
            "include": ["没有计划调整剂量", "永久停药", "停药前递减", "停药后随访"],
            "exclude": ["替代治疗或支持治疗"],
        },
        "6.9": {
            "include": ["替代治疗或支持治疗", "永久停药或严重溶血时"],
            "exclude": ["没有计划调整剂量", "100 mg每晚1次连续7天"],
        },
        "6.10": None,
    },
}


class ApiError(RuntimeError):
    def __init__(self, status: int, payload: Any):
        super().__init__(f"HTTP {status}: {payload}")
        self.status = status
        self.payload = payload


def request(base_url: str, method: str, path: str, payload: Any = None) -> Any:
    body = None
    headers = {}
    if payload is not None:
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        headers["Content-Type"] = "application/json"
    req = urllib.request.Request(
        base_url.rstrip("/") + path,
        data=body,
        headers=headers,
        method=method,
    )
    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            content = response.read()
            if response.headers.get_content_type() == "application/json":
                return json.loads(content)
            return {"content": content, "headers": dict(response.headers)}
    except urllib.error.HTTPError as exc:
        raw = exc.read()
        try:
            detail = json.loads(raw)
        except json.JSONDecodeError:
            detail = raw.decode("utf-8", errors="replace")
        raise ApiError(exc.code, detail) from exc


def projection_payload(
    journey: dict,
    working_copy: dict,
    *,
    key: str,
    overwrite: bool = False,
) -> dict:
    return {
        "expected_journey_revision": journey["revision"],
        "expected_intervention_rules_sha256": journey["intervention_rules_sha256"],
        "expected_working_copy_revision": working_copy["revision"],
        "overwrite_medical_edits": overwrite,
        "actor": "medical_manager_qc",
        "reason": (
            "医学经理确认以最新结构化规则覆盖本章隔离QC中的人工修订。"
            if overwrite
            else "医学经理确认将真实方案结构化干预规则应用到当前章节。"
        ),
        "idempotency_key": key,
    }


def assert_text(text: str, expectation: dict) -> None:
    for phrase in expectation["include"]:
        assert phrase in text, f"missing expected phrase: {phrase}\n{text}"
    for phrase in expectation["exclude"]:
        assert phrase not in text, f"unexpected cross-boundary phrase: {phrase}\n{text}"


def assert_includes(text: str, expectation: dict) -> None:
    for phrase in expectation["include"]:
        assert phrase in text, f"missing expected DOCX phrase: {phrase}"


def edited_projection_blocks(working_copy: dict, block_id: str, marker: str) -> list[dict]:
    blocks = deepcopy(working_copy["content_blocks"])
    block = next(item for item in blocks if item.get("block_id") == block_id)
    block["text"] = f"{block['text']}\n{marker}"
    block["rich_text"] = {
        "type": "paragraph",
        "attrs": {"stylePreset": "body"},
        "content": [{"type": "text", "text": block["text"]}],
    }
    return blocks


def docx_text(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def run(base_url: str, manifest_path: Path, output_dir: Path) -> dict:
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    output_dir.mkdir(parents=True, exist_ok=True)
    report = {"api_url": base_url, "projects": []}

    for project in manifest["projects"]:
        project_id = project["project_id"]
        project_code = project["project_code"]
        expectations = EXPECTED[project_code]
        project_result = {
            "project_id": project_id,
            "project_code": project_code,
            "source_path": project["source_path"],
            "sections": {},
        }
        journey = request(
            base_url,
            "GET",
            f"/api/projects/{project_id}/medical-writing/authoring-journey",
        )

        for section_number, section_id in project["sections"].items():
            route = (
                f"/api/projects/{project_id}/medical-writing/working-copies/"
                f"{section_id}"
            )
            working_copy = request(base_url, "GET", route)
            expectation = expectations[section_number]
            key_prefix = f"runtime-qc-{project_code}-{section_number.replace('.', '_')}"

            if expectation is None:
                try:
                    request(
                        base_url,
                        "POST",
                        route + "/intervention-rules-projection",
                        projection_payload(
                            journey,
                            working_copy,
                            key=key_prefix + "-empty",
                        ),
                    )
                except ApiError as exc:
                    assert exc.status == 422, exc
                    project_result["sections"][section_number] = {
                        "result": "correctly_rejected_empty_rule_panel",
                        "status": exc.status,
                    }
                    continue
                raise AssertionError(
                    f"empty rule panel unexpectedly projected: {project_code} {section_number}"
                )

            projected = request(
                base_url,
                "POST",
                route + "/intervention-rules-projection",
                projection_payload(
                    journey,
                    working_copy,
                    key=f"{key_prefix}-project-r{working_copy['revision']}",
                ),
            )
            projected_copy = projected["working_copy"]
            projection_block = projected["projection_block"]
            assert_text(projection_block["text"], expectation)

            reloaded = request(base_url, "GET", route)
            assert reloaded["revision"] == projected_copy["revision"]
            assert reloaded["content_blocks"] == projected_copy["content_blocks"]

            conflict_verified = False
            if section_number == "6.4":
                marker = f"医学经理人工修订保护验证：{project_code}。"
                manually_saved = request(
                    base_url,
                    "POST",
                    route,
                    {
                        "document_id": reloaded["document_id"],
                        "expected_revision": reloaded["revision"],
                        "content_blocks": edited_projection_blocks(
                            reloaded, projection_block["block_id"], marker
                        ),
                        "actor": "medical_manager_qc",
                        "idempotency_key": f"{key_prefix}-manual-r{reloaded['revision']}",
                    },
                )
                try:
                    request(
                        base_url,
                        "POST",
                        route + "/intervention-rules-projection",
                        projection_payload(
                            journey,
                            manually_saved,
                            key=f"{key_prefix}-protect-r{manually_saved['revision']}",
                        ),
                    )
                except ApiError as exc:
                    assert exc.status == 409, exc
                    assert "medical edits" in json.dumps(exc.payload, ensure_ascii=False)
                else:
                    raise AssertionError("manual medical edits were overwritten without confirmation")

                projected = request(
                    base_url,
                    "POST",
                    route + "/intervention-rules-projection",
                    projection_payload(
                        journey,
                        manually_saved,
                        key=f"{key_prefix}-overwrite-r{manually_saved['revision']}",
                        overwrite=True,
                    ),
                )
                projected_copy = projected["working_copy"]
                projection_block = projected["projection_block"]
                assert marker not in projection_block["text"]
                assert_text(projection_block["text"], expectation)
                conflict_verified = True

            project_result["sections"][section_number] = {
                "result": "projected",
                "working_copy_revision": projected_copy["revision"],
                "projection_action": projected["projection_action"],
                "projection_block_id": projection_block["block_id"],
                "projection_text": projection_block["text"],
                "manual_edit_conflict_verified": conflict_verified,
            }

        draft = request(
            base_url,
            "GET",
            f"/api/projects/{project_id}/medical-writing/document.docx?mode=draft_preview",
        )
        draft_content = draft["content"]
        draft_path = output_dir / f"{project_code}_draft_preview.docx"
        draft_path.write_bytes(draft_content)
        rendered_draft = docx_text(draft_content)
        for section_number, expectation in expectations.items():
            if expectation is not None:
                assert_includes(rendered_draft, expectation)

        approvals = []
        document_session = request(
            base_url,
            "GET",
            f"/api/projects/{project_id}/medical-writing/document-session",
        )
        section_number_by_id = {
            section_id: section_number
            for section_number, section_id in project["sections"].items()
        }
        for section in document_session["sections"]:
            section_id = section["section_id"]
            section_number = section_number_by_id.get(
                section_id, section.get("section_number") or "front-matter"
            )
            working_route = (
                f"/api/projects/{project_id}/medical-writing/working-copies/{section_id}"
            )
            current_copy = request(base_url, "GET", working_route)
            if current_copy["revision"] == 0:
                current_copy = request(
                    base_url,
                    "POST",
                    working_route,
                    {
                        "document_id": current_copy["document_id"],
                        "expected_revision": 0,
                        "content_blocks": current_copy["content_blocks"],
                        "actor": "medical_manager_qc",
                        "idempotency_key": f"save-full-document-{project_code}-{section_id}",
                    },
                )
            if current_copy["approval_state"] == "medically_approved":
                approvals.append(
                    {
                        "section_number": section_number,
                        "approval_id": "existing",
                        "state": "medically_approved",
                    }
                )
                continue
            gate_path = (
                f"/api/projects/{project_id}/medical-writing/working-copies/"
                f"{section_id}/approval-gate?"
                + urllib.parse.urlencode({"requested_by": "medical_manager_qc"})
            )
            gate = request(base_url, "POST", gate_path)
            approval = request(
                base_url,
                "POST",
                f"/api/projects/{project_id}/approvals/{gate['approval_id']}/actions",
                {
                    "action": "approve",
                    "actor": "medical_director_qc",
                    "comment": "真实方案干预规则章节隔离QC通过，批准当前章节快照。",
                    "idempotency_key": f"approve-{project_code}-{section_id}",
                },
            )
            assert approval["approval"]["state"] == "medically_approved"
            approvals.append(
                {
                    "section_number": section_number,
                    "approval_id": gate["approval_id"],
                    "state": approval["approval"]["state"],
                }
            )

        final = request(
            base_url,
            "GET",
            f"/api/projects/{project_id}/medical-writing/document.docx?mode=approved_final",
        )
        final_content = final["content"]
        final_path = output_dir / f"{project_code}_approved_final.docx"
        final_path.write_bytes(final_content)
        rendered_final = docx_text(final_content)
        for section_number, expectation in expectations.items():
            if expectation is not None:
                assert_includes(rendered_final, expectation)

        project_result["approvals"] = approvals
        project_result["draft_docx"] = {
            "path": str(draft_path),
            "sha256": hashlib.sha256(draft_content).hexdigest(),
        }
        project_result["approved_docx"] = {
            "path": str(final_path),
            "sha256": hashlib.sha256(final_content).hexdigest(),
        }
        report["projects"].append(project_result)

    return report


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--api-url", required=True)
    parser.add_argument("--manifest", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    report = run(args.api_url, args.manifest, args.output_dir)
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(
        json.dumps(
            {
                "project_count": len(report["projects"]),
                "report": str(args.report),
            },
            ensure_ascii=False,
        )
    )


if __name__ == "__main__":
    main()
