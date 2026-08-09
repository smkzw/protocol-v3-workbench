from __future__ import annotations

import argparse
import json
from pathlib import Path
import sys
from urllib.error import HTTPError
from urllib.request import Request, urlopen

from docx import Document

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from packages.contracts.workbench_contracts import (
    MedicalWritingPicosDefinition,
    MedicalWritingStudyFraming,
)
from services.api.app.medical_writing_manifest import (
    D001_PROTOCOL_DOCX,
    MY008_PNH_3_01_PROTOCOL_DOCX,
    RUX_PROTOCOL_DOCX,
)
from services.api.app.medical_writing_protocol_template import (
    TEMPLATE_ID,
    TEMPLATE_VERSION,
)


def _request(base_url: str, method: str, path: str, payload: dict | None = None):
    data = None if payload is None else json.dumps(payload, ensure_ascii=False).encode()
    request = Request(
        f"{base_url.rstrip('/')}{path}",
        data=data,
        method=method,
        headers={"Content-Type": "application/json"} if data is not None else {},
    )
    try:
        with urlopen(request, timeout=60) as response:
            return json.load(response)
    except HTTPError as exc:
        detail = exc.read().decode("utf-8", errors="replace")
        raise RuntimeError(f"{method} {path} failed with HTTP {exc.code}: {detail}") from exc


def _source_locator(path: Path, phrase: str) -> str:
    document = Document(str(path))
    for index, paragraph in enumerate(document.paragraphs):
        if phrase in paragraph.text:
            return f"{path.name} | paragraph:{index + 1} | {phrase}"
    for table_index, table in enumerate(document.tables):
        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                if phrase in cell.text:
                    return (
                        f"{path.name} | table:{table_index + 1} row:{row_index + 1} "
                        f"cell:{cell_index + 1} | {phrase}"
                    )
    raise RuntimeError(f"source phrase not found in {path}: {phrase}")


def _base_picos(**overrides) -> dict:
    payload = {
        "design_archetype": "randomized_confirmatory",
        "field_applicability": {},
        "population_summary": "符合方案诊断与疾病严重程度要求的目标适应症试验参与者。",
        "inclusion_modules": ["筛选期和基线期均满足方案规定的疾病活动度标准"],
        "exclusion_modules": ["存在方案规定的活动性感染或其他重大安全性风险"],
        "washout_rules": ["既往治疗须按方案规定完成洗脱"],
        "intervention_summary": "按方案规定接受试验用药品治疗。",
        "intervention_dose_regimen": "按方案规定的剂量、频次和给药途径使用试验用药品。",
        "allowed_concomitant_rules": [],
        "required_background_rules": [],
        "prohibited_concomitant_rules": [],
        "assessment_timing_restrictions": ["疗效与安全性评价前遵守方案规定的用药时限"],
        "intervention_rules": None,
        "comparator_summary": "匹配安慰剂对照。",
        "primary_endpoint": "在方案规定的主要评价时点评价主要疗效终点。",
        "key_secondary_endpoints": ["在方案规定时点评价关键次要疗效终点"],
        "other_secondary_endpoints": ["评价其他疗效指标"],
        "exploratory_endpoints": ["评价探索性生物标志物或患者报告结局"],
        "safety_endpoints": ["TEAE、SAE及导致停药的不良事件"],
        "aesi_definitions": ["方案规定的特别关注不良事件"],
        "assessment_instruments": [],
        "study_epochs": ["筛选期", "治疗期", "安全性随访期"],
        "visit_strategy": "按研究流程表完成筛选、基线、治疗期和随访期访视。",
        "estimand_strategy": "主要估计目标按方案规定的治疗策略和伴发事件处理规则定义。",
        "sample_size_strategy": "根据主要终点假设、显著性水平、把握度和脱落率估算。",
        "statistical_strategy": "采用与主要终点类型和研究设计匹配的统计模型进行分析。",
    }
    payload.update(overrides)
    return MedicalWritingPicosDefinition.model_validate(payload).model_dump(mode="json")


def _project_definitions() -> list[dict]:
    rux_locator_stop = _source_locator(RUX_PROTOCOL_DOCX, "若治疗期间总BSA超过20%的受试者必须停止使用研究药物")
    rux_locator_restart = _source_locator(RUX_PROTOCOL_DOCX, "若AD复发可重新开始用药")
    rux_locator_interrupt = _source_locator(RUX_PROTOCOL_DOCX, "做出关于个体中断给药的决定")
    d001_locator = _source_locator(D001_PROTOCOL_DOCX, "不能早于系统性补救药物最后一次给药后的5个半衰期")
    pnh_locator_no_plan = _source_locator(MY008_PNH_3_01_PROTOCOL_DOCX, "没有计划调整剂量和（或）中断剂量")
    pnh_locator_taper = _source_locator(MY008_PNH_3_01_PROTOCOL_DOCX, "应考虑在14天内逐渐减少MY008211A片的服药剂量")

    return [
        {
            "code": "QC-RUX-IR",
            "project_name": "RUX-03-002特应性皮炎III期方案干预规则QC",
            "framing": {
                "protocol_id": "RUX-03-002",
                "version": "V1.3",
                "document_title": "磷酸芦可替尼乳膏治疗特应性皮炎的III期临床研究方案",
                "indication": "特应性皮炎",
                "clinicaltrials_condition_term": "Atopic Dermatitis",
                "study_phase": "III期",
                "intrinsic_objectives": ["确证性研究"],
                "investigational_product": "磷酸芦可替尼乳膏",
                "target_mechanism": "JAK1/JAK2抑制剂",
                "competitor_target_scope": "外用JAK抑制剂及其他创新外用治疗",
                "development_regions": ["中国"],
                "design_pattern": "随机、双盲、安慰剂对照、多中心III期研究",
                "population_intent": "轻中度特应性皮炎受试者",
                "key_uncertainties": [],
                "manual_source_ids": [str(RUX_PROTOCOL_DOCX)],
                "terminology_policy": "cde_participant",
            },
            "picos": _base_picos(
                population_summary="符合方案诊断标准的轻中度特应性皮炎受试者。",
                intervention_summary="磷酸芦可替尼乳膏双盲治疗后进入开放治疗期。",
                intervention_dose_regimen="磷酸芦可替尼乳膏每日2次外用，给药间隔至少8小时。",
                primary_endpoint="第8周达到IGA-TS的受试者比例。",
                key_secondary_endpoints=["第8周EASI 75应答率", "第8周Itch NRS 4应答率"],
                intervention_rules={
                    "authority": "structured",
                    "ip_regimens": [{
                        "regimen_id": "rux_ip_main",
                        "product_name": "磷酸芦可替尼乳膏",
                        "product_role": "investigational_product",
                        "dose_and_frequency": "每日2次，间隔至少8小时",
                        "route": "外用",
                        "treatment_period": "双盲治疗期8周，开放治疗期16周",
                        "source_location": rux_locator_restart,
                    }],
                    "ip_adjustment_policy": "protocol_defined",
                    "ip_action_rules": [
                        {
                            "rule_id": "rux_pause_clearance",
                            "action_kind": "planned_on_off",
                            "trigger": "开放治疗期IGA为0分且皮损完全清除3天",
                            "study_product_action": "暂停使用研究药物",
                            "retest_recovery": "AD复发时立即重新开始使用研究药物",
                            "source_location": rux_locator_restart,
                        },
                        {
                            "rule_id": "rux_stop_bsa",
                            "action_kind": "permanent_discontinuation",
                            "trigger": "治疗期间总BSA超过20%",
                            "study_product_action": "停止使用研究药物",
                            "permanent_discontinuation_condition": "总BSA超过20%",
                            "source_location": rux_locator_stop,
                        },
                        {
                            "rule_id": "rux_interrupt_safety",
                            "action_kind": "temporary_interruption",
                            "trigger": "发生需结合AE、基础疾病和研究治疗相关性判断的安全性情况",
                            "confirmation_required": "尽可能与医学监查员协商并由研究者作出临床判断",
                            "study_product_action": "个体化中断给药",
                            "source_location": rux_locator_interrupt,
                        },
                    ],
                    "non_ip_treatment_rules": [],
                    "cross_object_links": [],
                },
            ),
            "source_path": str(RUX_PROTOCOL_DOCX),
        },
        {
            "code": "QC-D001-IR",
            "project_name": "CMS-D001银屑病II/III期方案干预规则QC",
            "framing": {
                "protocol_id": "CMS-D001",
                "version": "V1.0",
                "document_title": "CMS-D001片治疗中重度斑块状银屑病的II/III期临床研究方案",
                "indication": "中重度斑块状银屑病",
                "clinicaltrials_condition_term": "Plaque Psoriasis",
                "study_phase": "II/III期",
                "intrinsic_objectives": ["剂量探索", "确证性研究"],
                "investigational_product": "CMS-D001片",
                "target_mechanism": "创新口服免疫调节剂",
                "competitor_target_scope": "银屑病口服小分子及生物制剂",
                "development_regions": ["中国"],
                "design_pattern": "多中心、随机、双盲、安慰剂对照的II/III期操作无缝适应性设计",
                "population_intent": "中度至重度斑块状银屑病成人患者",
                "key_uncertainties": [],
                "manual_source_ids": [str(D001_PROTOCOL_DOCX)],
                "terminology_policy": "cde_participant",
            },
            "picos": _base_picos(
                population_summary="中度至重度斑块状银屑病成人患者。",
                intervention_summary="CMS-D001片多个剂量组接受口服治疗。",
                intervention_dose_regimen="CMS-D001片按随机分组每日口服给药。",
                intervention_rules={
                    "authority": "structured",
                    "ip_regimens": [{
                        "regimen_id": "d001_ip_main",
                        "product_name": "CMS-D001片",
                        "product_role": "investigational_product",
                        "dose_and_frequency": "按随机分组每日口服给药",
                        "route": "口服",
                        "treatment_period": "方案规定治疗期",
                        "source_location": d001_locator,
                    }],
                    "ip_adjustment_policy": "protocol_defined",
                    "ip_action_rules": [{
                        "rule_id": "d001_hold_for_rescue",
                        "action_kind": "temporary_interruption",
                        "trigger": "开始全身系统性补救治疗",
                        "study_product_action": "立即暂停研究药物治疗",
                        "retest_recovery": "研究者和申办者认为可恢复时恢复",
                        "wait_period": "不得早于系统性补救药物末次给药后5个半衰期",
                        "approvers": ["研究者", "申办者"],
                        "linked_non_ip_rule_ids": ["d001_systemic_rescue"],
                        "source_location": d001_locator,
                    }],
                    "non_ip_treatment_rules": [
                        {
                            "rule_id": "d001_systemic_rescue",
                            "rule_class": "rescue",
                            "policy": "rescue_policy",
                            "agent_or_category": "系统性糖皮质激素或其他系统性免疫抑制/调节治疗",
                            "phase_applicability": "III期临床研究阶段",
                            "timing_restrictions": ["原则上不早于第24周访视", "外用补救治疗14天无改善后方可升级"],
                            "cm_dose_rule": "补救治疗自身的剂量变化记录为非试验用药变化",
                            "source_location": d001_locator,
                        },
                        {
                            "rule_id": "d001_allowed_cm",
                            "rule_class": "allowed_cm",
                            "policy": "allowed_if_stable",
                            "agent_or_category": "方案允许的稳定剂量合并用药",
                            "cm_dose_rule": "合并用药剂量变化仍作为CM记录",
                            "source_location": d001_locator,
                        },
                    ],
                    "cross_object_links": [{
                        "link_id": "d001_rescue_hold_link",
                        "source_rule_id": "d001_systemic_rescue",
                        "source_kind": "rescue",
                        "target_kind": "ip",
                        "target_rule_id": "d001_hold_for_rescue",
                        "action": "hold",
                        "resume_condition": "末次系统性补救治疗后至少5个半衰期，且经研究者和申办者确认",
                        "source_location": d001_locator,
                    }],
                },
            ),
            "source_path": str(D001_PROTOCOL_DOCX),
        },
        {
            "code": "QC-PNH-IR",
            "project_name": "MY008211A-PNH III期方案干预规则QC",
            "framing": {
                "protocol_id": "MY008211A-PNH-3-01",
                "version": "V1.1",
                "document_title": "MY008211A片治疗PNH的多中心单臂开放标签III期临床研究方案",
                "indication": "阵发性睡眠性血红蛋白尿症",
                "clinicaltrials_condition_term": "Paroxysmal Nocturnal Hemoglobinuria",
                "study_phase": "III期",
                "intrinsic_objectives": ["确证性研究"],
                "investigational_product": "MY008211A片",
                "target_mechanism": "补体通路抑制剂",
                "competitor_target_scope": "补体C5抑制剂及补体替代通路抑制剂",
                "development_regions": ["中国"],
                "design_pattern": "多中心、单臂、开放标签III期研究",
                "population_intent": "C5单抗疗效不佳的PNH患者",
                "key_uncertainties": [],
                "manual_source_ids": [str(MY008_PNH_3_01_PROTOCOL_DOCX)],
                "terminology_policy": "cde_participant",
            },
            "picos": _base_picos(
                design_archetype="other",
                field_applicability={
                    "comparator_summary": {
                        "status": "not_applicable",
                        "reason": "本研究为多中心、单臂、开放标签III期研究，不设置同期对照组。",
                        "confirmed_by_medical_manager": True,
                    }
                },
                population_summary="C5单抗治疗后疗效不佳的PNH患者。",
                intervention_summary="MY008211A片400 mg口服，每日2次。",
                intervention_dose_regimen="MY008211A片400 mg口服BID。",
                comparator_summary="",
                primary_endpoint="在方案规定时点评价血红蛋白改善及溶血控制。",
                intervention_rules={
                    "authority": "structured",
                    "ip_regimens": [{
                        "regimen_id": "pnh_ip_main",
                        "product_name": "MY008211A片",
                        "product_role": "investigational_product",
                        "dose_and_frequency": "400 mg，每日2次",
                        "route": "口服",
                        "treatment_period": "疗效观察期24周",
                        "source_location": pnh_locator_no_plan,
                    }],
                    "ip_adjustment_policy": "no_planned_adjustment",
                    "no_planned_adjustment_statement": "本研究没有计划调整剂量和（或）中断剂量，但方案规定的停止/提前终止、递减和停药后随访除外。",
                    "ip_action_rules": [
                        {
                            "rule_id": "pnh_permanent_stop",
                            "action_kind": "permanent_discontinuation",
                            "trigger": "发生重大安全性风险，需立即停止或提前终止治疗",
                            "study_product_action": "永久停止MY008211A片治疗",
                            "permanent_discontinuation_condition": "研究者判断继续治疗的风险超过获益",
                            "source_location": pnh_locator_taper,
                        },
                        {
                            "rule_id": "pnh_taper",
                            "action_kind": "discontinuation_taper",
                            "trigger": "无需立即停止或提前终止时",
                            "taper_steps": ["100 mg每晚1次连续7天", "50 mg每晚1次连续7天", "完成递减后停药"],
                            "source_location": pnh_locator_taper,
                        },
                        {
                            "rule_id": "pnh_follow_up",
                            "action_kind": "post_discontinuation_follow_up",
                            "trigger": "永久停药或完成剂量递减后",
                            "study_product_action": "监测溶血、实验室检查、不良事件及PNH症状体征",
                            "retest_recovery": "永久停药后一周完成规定实验室评估",
                            "source_location": pnh_locator_taper,
                        },
                    ],
                    "non_ip_treatment_rules": [{
                        "rule_id": "pnh_supportive_rescue",
                        "rule_class": "rescue",
                        "policy": "rescue_policy",
                        "agent_or_category": "替代治疗或支持治疗",
                        "phase_applicability": "永久停药或严重溶血时",
                        "notes": "可包括补体C5抑制剂、输血、糖皮质激素、抗凝及研究者判断的其他支持治疗",
                        "source_location": pnh_locator_taper,
                    }],
                    "cross_object_links": [],
                },
            ),
            "source_path": str(MY008_PNH_3_01_PROTOCOL_DOCX),
        },
    ]


def _commit_stage(base_url: str, project_id: str, stage: str, payload: dict) -> dict:
    journey = _request(base_url, "GET", f"/api/projects/{project_id}/medical-writing/authoring-journey")
    preview = _request(
        base_url,
        "POST",
        f"/api/projects/{project_id}/medical-writing/authoring-journey/impact-preview",
        {"expected_revision": journey["revision"], "stage": stage, stage: payload},
    )
    return _request(
        base_url,
        "POST",
        f"/api/projects/{project_id}/medical-writing/authoring-journey/stages/{stage}/commit",
        {
            "expected_revision": journey["revision"],
            "stage": stage,
            stage: payload,
            "impact_preview_id": preview["preview_id"],
            "actor": "medical_manager_qc",
            "idempotency_key": f"qc-{project_id}-{stage}-r{journey['revision']}",
        },
    )


def seed(base_url: str) -> list[dict]:
    results = []
    for definition in _project_definitions():
        framing = MedicalWritingStudyFraming.model_validate(definition["framing"]).model_dump(mode="json")
        created = _request(
            base_url,
            "POST",
            "/api/projects",
            {
                "project_code": definition["code"],
                "project_name": definition["project_name"],
                "indication": framing["indication"],
                "product_name": framing["investigational_product"],
                "study_phase": framing["study_phase"],
                "protocol_id": framing["protocol_id"],
                "protocol_version": framing["version"],
                "entry_mode": "from_zero",
                "actor": "medical_manager_qc",
                "idempotency_key": f"seed-{definition['code']}-20260717",
            },
        )
        project_id = created["project"]["project_id"]
        _commit_stage(base_url, project_id, "framing", framing)
        _commit_stage(base_url, project_id, "picos", definition["picos"])
        gate_state = _request(
            base_url,
            "POST",
            f"/api/projects/{project_id}/medical-writing/authoring-journey/corpus-gate/recalculate",
            {},
        )
        missing = gate_state["corpus_gate"]["missing_requirements"]
        if missing:
            gate_state = _request(
                base_url,
                "POST",
                f"/api/projects/{project_id}/medical-writing/authoring-journey/corpus-gate/override",
                {
                    "expected_revision": gate_state["revision"],
                    "reason": "仅用于隔离运行时真实项目功能与浏览器QC；不构成正式语料充分性结论。",
                    "acknowledged_missing_requirements": missing,
                    "actor": "medical_manager_qc",
                    "idempotency_key": f"qc-corpus-override-{project_id}-r{gate_state['revision']}",
                },
            )
        study_definition = gate_state["study_definition"]
        document_result = _request(
            base_url,
            "POST",
            f"/api/projects/{project_id}/medical-writing/greenfield-document",
            {
                "protocol_id": framing["protocol_id"],
                "version": framing["version"],
                "document_title": framing["document_title"],
                "indication": framing["indication"],
                "study_phase": framing["study_phase"],
                "source_study_definition_id": study_definition["definition_id"],
                "source_study_definition_revision": study_definition["revision"],
                "source_study_definition_sha256": study_definition["state_sha256"],
                "template_id": TEMPLATE_ID,
                "template_version": TEMPLATE_VERSION,
                "actor": "medical_manager_qc",
                "idempotency_key": f"qc-m11-document-{project_id}-{study_definition['revision']}",
            },
        )
        document = document_result["document"]
        sections = {
            item["section_number"]: item["section_id"]
            for item in document["sections"]
            if item["section_number"] in {"6.4", "6.9", "6.10"}
        }
        if set(sections) != {"6.4", "6.9", "6.10"}:
            raise RuntimeError(f"M11 target sections missing for {project_id}: {sections}")
        journey = _request(base_url, "GET", f"/api/projects/{project_id}/medical-writing/authoring-journey")
        results.append(
            {
                "project_id": project_id,
                "project_code": definition["code"],
                "source_path": definition["source_path"],
                "journey_revision": journey["revision"],
                "intervention_rules_sha256": journey["intervention_rules_sha256"],
                "document_id": document["document_id"],
                "sections": sections,
            }
        )
    return results


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--api-url", required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    projects = seed(args.api_url)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(
        json.dumps({"api_url": args.api_url, "projects": projects}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(json.dumps({"project_count": len(projects), "output": str(args.output)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
