#!/usr/bin/env python3
"""Build a self-contained company protocol corpus snapshot for the workbench."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
DEFAULT_BASE = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/玄关入排系统构建/"
    "protocol_corpus/corpus/protocol_clause_library.jsonl"
)
DEFAULT_OUTPUT = (
    ROOT
    / "services/api/assets/medical_writing_corpus"
    / "cms_cn_protocol_corpus_20260715_v1.jsonl"
)

SUPPLEMENTAL_SOURCES = (
    {
        "path": Path(
            "/Users/smkzw/Documents/康哲项目资料/CMS-D017/4.方案/PNH/方案摘要/"
            "CMS-D017-PNH-方案摘要_v0.2.docx"
        ),
        "authority_priority": 0,
        "source_role": "primary_synopsis_reference",
        "authority_scope": "方案摘要、中文表达、首页、概要表格、研究流程表及附注的最高参照",
    },
    {
        "path": Path(
            "/Users/smkzw/Documents/康哲项目资料/模版/方案模版/"
            "CMS-D017Ⅰ期方案-v1.1-20260209-clean.docx"
        ),
        "authority_priority": 2,
        "source_role": "full_protocol_style_and_clause_reference",
        "authority_scope": "完整方案正文、多级标题、目录和全文条款补充",
    },
    {
        "path": Path(
            "/Users/smkzw/Documents/朗来项目资料/MY004/RA/"
            "MY004-RA-2b 研究方案摘要_V0.3-with Comments to ABBV.docx"
        ),
        "authority_priority": 3,
        "source_role": "cross_project_synopsis_reference",
        "authority_scope": "RA IIb、阳性药/安慰剂对照、量表和研究流程的项目化参考",
    },
    {
        "path": Path(
            "/Users/smkzw/Documents/朗来项目资料/MY004/AD-CSU-PN三合一整合/"
            "MY004567片-炎症性皮肤病-方案摘要-V0.4-clean.docx"
        ),
        "authority_priority": 3,
        "source_role": "cross_project_synopsis_reference",
        "authority_scope": "多适应症PoC、量表、PD和皮肤活检的项目化参考",
    },
)

_CATEGORY_PATTERNS = {
    "ethics_consent_privacy": r"伦理|知情同意|GCP|隐私|保密",
    "reproductive_pregnancy_contraception": r"妊娠|怀孕|哺乳|避孕|育龄|hCG|FSH",
    "infection_screening": r"感染|结核|TB|乙肝|丙肝|HIV|疫苗|脑膜炎球菌",
    "prior_concomitant_treatment": r"合并用药|合并治疗|禁用|限制用药|洗脱|既往治疗|救援治疗",
    "study_design_structure": r"随机|盲|开放标签|安慰剂|阳性对照|筛选期|治疗期|随访期|剂量探索",
    "visit_eos_eot_withdrawal": r"EOT|EOS|研究结束|治疗结束|提前终止|退出研究|失访",
    "soa_procedure_framework": r"研究流程表|访视|窗口|生命体征|体格检查|实验室|心电图|PK|PD|量表",
    "safety_reporting": r"不良事件|严重不良事件|AESI|AE|SAE|安全性|过量",
    "endpoint_objective_stats": r"研究目的|主要终点|次要终点|探索性终点|样本量|分析集|统计分析|估计目标",
    "eligibility_disease_specific": r"入选标准|排除标准|诊断|疾病活动|FACIT|ACR|DAS28|EASI|UAS7|PNRS",
}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def normalized(value: str) -> str:
    return re.sub(r"\s+", " ", value).strip()


def looks_like_heading(style_name: str, text: str) -> bool:
    if style_name.casefold().startswith(("heading", "标题")):
        return True
    return bool(re.match(r"^\d+(?:\.\d+){0,4}\s+\S+", text)) and len(text) <= 160


def iter_docx_blocks(path: Path) -> Iterable[tuple[str, str, str]]:
    document = Document(path)
    section = "文档开头"
    for paragraph in document.paragraphs:
        text = normalized(paragraph.text)
        if not text:
            continue
        style_name = paragraph.style.name if paragraph.style else ""
        if looks_like_heading(style_name, text):
            section = text[:160]
        yield section, "paragraph", text
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            cells = [normalized(cell.text) for cell in row.cells]
            if any(cells):
                yield (
                    f"表{table_index}",
                    "table_row",
                    " | ".join(cells),
                )


def split_text(text: str, limit: int = 850) -> list[str]:
    if len(text) <= limit:
        return [text]
    units = [item.strip() for item in re.split(r"(?<=[。；;])", text) if item.strip()]
    chunks: list[str] = []
    current = ""
    for unit in units:
        if len(current) + len(unit) <= limit:
            current += unit
            continue
        if current:
            chunks.append(current)
        while len(unit) > limit:
            chunks.append(unit[:limit])
            unit = unit[limit:]
        current = unit
    if current:
        chunks.append(current)
    return chunks


def categories_for(text: str) -> list[str]:
    return [
        category
        for category, pattern in _CATEGORY_PATTERNS.items()
        if re.search(pattern, text, re.IGNORECASE)
    ]


def supplemental_entries(source: dict[str, object]) -> list[dict[str, object]]:
    path = source["path"]
    assert isinstance(path, Path)
    source_hash = file_sha256(path)
    document_id = hashlib.sha1(
        f"{path.name}|{source_hash}".encode("utf-8")
    ).hexdigest()[:16]
    result: list[dict[str, object]] = []
    for block_number, (section, block_type, text) in enumerate(iter_docx_blocks(path), start=1):
        if len(text) < 12:
            continue
        for chunk_number, chunk in enumerate(split_text(text), start=1):
            categories = categories_for(f"{section}\n{chunk}")
            entry_id = hashlib.sha1(
                f"{document_id}|{section}|{block_number}|{chunk_number}|{chunk}".encode("utf-8")
            ).hexdigest()[:14]
            result.append(
                {
                    "entry_id": entry_id,
                    "source_doc_id": document_id,
                    "source_role": source["source_role"],
                    "source_file": path.name,
                    "source_sha256": source_hash,
                    "authority_priority": source["authority_priority"],
                    "authority_scope": source["authority_scope"],
                    "section": section,
                    "block_no": block_number,
                    "chunk_no": chunk_number,
                    "block_type": block_type,
                    "categories": categories,
                    "reuse_level": (
                        "disease_specific_reference"
                        if "eligibility_disease_specific" in categories
                        else "adapt_with_project"
                    ),
                    "text_hash": hashlib.sha1(normalized(chunk).encode("utf-8")).hexdigest()[:16],
                    "text": chunk,
                    "review_note": "仅按当前项目事实、适应症、研究分期和设计适用性审阅后复用。",
                    "canonical_entry_id": entry_id,
                    "near_duplicate_of": "",
                    "near_duplicate_similarity": 1.0,
                    "duplicate_group_size": 1,
                }
            )
    return result


def base_authority(entry: dict[str, object]) -> tuple[int, str]:
    name = str(entry.get("source_file") or "")
    if name.startswith("CMS-D005-减重II期") and "KZYY0525-clean" in name:
        return 1, "secondary_synopsis_reference"
    if name.startswith(("CMS-D017", "1-3-4-1-2", "MG-K10", "CMD-D001")):
        return 2, "full_protocol_style_and_clause_reference"
    return 4, str(entry.get("source_role") or "legacy_company_reference")


def build(base_path: Path, output_path: Path) -> dict[str, object]:
    if not base_path.is_file():
        raise FileNotFoundError(base_path)
    rows: list[dict[str, object]] = []
    with base_path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip():
                continue
            entry = json.loads(line)
            priority, source_role = base_authority(entry)
            entry["authority_priority"] = priority
            entry["source_role"] = source_role
            rows.append(entry)

    source_manifest: list[dict[str, object]] = []
    for source in SUPPLEMENTAL_SOURCES:
        path = source["path"]
        assert isinstance(path, Path)
        if not path.is_file():
            raise FileNotFoundError(path)
        additions = supplemental_entries(source)
        rows.extend(additions)
        source_manifest.append(
            {
                "file_name": path.name,
                "sha256": file_sha256(path),
                "authority_priority": source["authority_priority"],
                "source_role": source["source_role"],
                "authority_scope": source["authority_scope"],
                "entry_count": len(additions),
            }
        )

    rows.sort(
        key=lambda item: (
            int(item.get("authority_priority", 99)),
            str(item.get("source_file") or ""),
            int(item.get("block_no") or 0),
            int(item.get("chunk_no") or 0),
            str(item.get("entry_id") or ""),
        )
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with output_path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n")
    snapshot_hash = file_sha256(output_path)
    manifest = {
        "schema_version": "cms_medical_writing_corpus_snapshot_v1",
        "snapshot_id": "cms_cn_protocol_corpus",
        "snapshot_version": "cms_cn_protocol_corpus_20260715_v1",
        "snapshot_sha256": snapshot_hash,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "base_library": {
            "file_name": base_path.name,
            "sha256": file_sha256(base_path),
            "entry_count": sum(1 for _ in base_path.open(encoding="utf-8")),
        },
        "supplemental_sources": source_manifest,
        "entry_count": len(rows),
    }
    manifest_path = output_path.with_suffix(".manifest.json")
    manifest_path.write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--base", type=Path, default=DEFAULT_BASE)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    print(json.dumps(build(args.base, args.output), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
