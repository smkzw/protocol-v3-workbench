from __future__ import annotations

import argparse
import io
import json
import os
import sys
import tempfile
from pathlib import Path

import docx

ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = ROOT.parents[1]
sys.path.insert(0, str(ROOT))

from packages.contracts.workbench_contracts import AiTaskFromRegistryRequest
from services.api.app.ai_execution_policy import AiExecutionPolicyResolver
from services.api.app.ai_gateway import (
    DIRECT_DEEPSEEK_BASE_URL,
    DIRECT_DEEPSEEK_MODEL,
    DisabledAiProvider,
    configured_ai_provider_from_env,
)
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore, public_ai_run
from services.api.app.demo_repository import DemoRepository
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore


PROJECT_ID = "proj_mgk10_sar_demo"


def _synthetic_protocol_bytes() -> bytes:
    document = docx.Document()
    document.add_heading("非受试者数据 AI Provider 验证方案", level=1)
    document.add_paragraph("本文件只用于验证模型网关，不包含真实受试者、中心或项目机密数据。")
    document.add_paragraph("IN-01：受试者应在任何研究程序前签署知情同意书。")
    document.add_paragraph("EX-01：研究者判断存在无法安全参加研究的情况时不得入组。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _required_env(name: str) -> str:
    value = os.environ.get(name, "").strip()
    if not value:
        raise RuntimeError(f"missing required environment variable: {name}")
    return value


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    model_name = _required_env("WORKBENCH_AI_MODEL")
    provider_name = os.environ.get("WORKBENCH_AI_PROVIDER", "openai_compatible").strip()
    transport = os.environ.get("WORKBENCH_AI_TRANSPORT", "openai_compatible").strip()
    base_url = os.environ.get("WORKBENCH_AI_BASE_URL", DIRECT_DEEPSEEK_BASE_URL).rstrip("/")
    deployment_profile = _required_env("WORKBENCH_AI_DEPLOYMENT_PROFILE")
    if provider_name != "deepseek":
        raise RuntimeError("product AI provider must be direct DeepSeek")
    if transport != "openai_compatible":
        raise RuntimeError("product AI transport must not invoke Hermes")
    if base_url != DIRECT_DEEPSEEK_BASE_URL:
        raise RuntimeError("product AI base URL must be the direct DeepSeek API")
    if model_name != DIRECT_DEEPSEEK_MODEL:
        raise RuntimeError(f"product AI model must be {DIRECT_DEEPSEEK_MODEL}")

    with tempfile.TemporaryDirectory() as tmp:
        root = Path(tmp)
        registry = SourceRegistryService(SourceRegistryStore(root / "sources.jsonl"))
        registered = registry.register_protocol_docx(
            PROJECT_ID,
            "synthetic_non_patient_protocol.docx",
            _synthetic_protocol_bytes(),
            module="eligibility_review",
        )
        provider = configured_ai_provider_from_env()
        if isinstance(provider, DisabledAiProvider):
            raise RuntimeError("configured AI provider transport is unavailable")
        policy = AiExecutionPolicyResolver(
            deployment_profile=deployment_profile,
            provider_name=provider_name,
            model_name=model_name,
        )
        runner = AiTaskRunner(
            DemoRepository(PROJECT_ROOT / "demo_data" / "workbench_demo_v0_1.json"),
            AiTaskStore(root / "ai_runs.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=policy,
        )
        request = AiTaskFromRegistryRequest(
            module="eligibility_review",
            task_type="protocol_rule_extraction",
            expected_prompt_version="protocol_rule_extraction_v0_1",
            source_ids=[span.source_id for span in registered.spans[:4]],
            forbidden_source_ids=["previous_ai_summary", "legacy_eligibility_report"],
            user_instruction=(
                "仅基于登记的合成非受试者方案片段，抽取 IN-01 和 EX-01 候选规则；"
                "证据不足时明确标记 uncertainty，所有内容均待医学确认。"
            ),
        )
        run = runner.submit_registered(PROJECT_ID, request, registry)
        public = public_ai_run(run)
        public["validation"] = {
            "registered_source_kind": registered.entry.source_kind,
            "registered_span_count": registered.entry.span_count,
            "contains_real_subject_data": False,
            "provider_invoked_outside_codex": True,
            "provider_invoked_outside_hermes": True,
            "response_model_identity_required": True,
        }

    serialized = json.dumps(public, ensure_ascii=False, indent=2)
    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(serialized + "\n", encoding="utf-8")
    print(serialized)
    return 0 if public["status"] == "completed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
