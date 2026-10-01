#!/usr/bin/env python3
"""Generate a 2-3 page NSCLC first-line Phase III protocol synopsis for T2 E2E testing."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

out = Path(__file__).with_name("T2_NSCLC_1L_PhaseIII_Synopsis.docx")
doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.paragraph_format.line_spacing = 1.15
style.paragraph_format.space_after = Pt(6)


def set_run_font(run, size=11, bold=False, east="宋体"):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east)


def add_heading_cn(text, size=16):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, east="黑体")
    return p


def add_section(title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(title)
    set_run_font(run, size=13, bold=True, east="黑体")
    return p


def add_body(text):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    set_run_font(run, size=11, bold=False)
    return p


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=11, bold=False)
    return p


add_heading_cn("临床试验方案摘要", 18)
add_heading_cn("XS-201联合含铂双药化疗对比安慰剂联合含铂双药化疗", 13)
add_heading_cn("用于一线治疗驱动基因阴性晚期非小细胞肺癌的III期随机双盲研究", 12)

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("方案编号：XS-201-NSCLC-301    版本：草案 0.1    日期：2026-09-28")
set_run_font(r, size=10)

add_section("1. 研究背景与目的")
add_body(
    "非小细胞肺癌（NSCLC）是肺癌的主要组织学类型。对于无EGFR敏感突变、ALK/ROS1融合等驱动基因改变的IV期NSCLC患者，"
    "含铂双药化疗联合PD-1通路阻断已成为一线标准策略之一。XS-201是一种人源化抗PD-1单克隆抗体，拟在一线场景评估其"
    "与含铂双药化疗联合对比安慰剂联合含铂双药化疗的疗效与安全性。"
)
add_body("主要目的：比较XS-201联合含铂双药化疗与安慰剂联合含铂双药化疗在一线驱动基因阴性IV期NSCLC患者中的无进展生存期（PFS）。")
add_body("次要目的：比较总生存期（OS）、客观缓解率（ORR）、缓解持续时间（DoR）及安全性；探索PD-L1表达与疗效的关系。")

add_section("2. 研究设计")
add_body(
    "本研究为多中心、随机、双盲、安慰剂对照的III期临床试验。计划入组约500例受试者，按1:1随机分配至试验组或对照组。"
    "随机分层因素包括：组织学类型（鳞状 vs 非鳞状）、PD-L1肿瘤比例评分（TPS：<1% vs 1%–49% vs ≥50%）、"
    "东部肿瘤协作组体能状态（ECOG PS 0 vs 1）及吸烟史（从不吸烟 vs 既往/目前吸烟）。"
)
add_body(
    "盲法：受试者、研究者、申办方医学监查人员对分组保持盲态。XS-201与安慰剂在外观、包装和给药操作上保持一致。"
    "研究分为筛选期（最长28天）、治疗期及随访期。治疗持续至疾病进展、不可耐受毒性、受试者撤回知情同意或研究结束，以先发生者为准。"
)
add_bullet("试验组：XS-201 200 mg静脉输注，每3周一次（Q3W），联合含铂双药化疗（最多4个周期），随后XS-201维持治疗。")
add_bullet("对照组：XS-201安慰剂 200 mg静脉输注 Q3W，联合同样的含铂双药化疗（最多4个周期），随后安慰剂维持治疗。")
add_bullet("非鳞状NSCLC化疗方案：培美曲塞 500 mg/m² + 顺铂 75 mg/m² 或卡铂 AUC 5，Q3W，最多4个周期；允许培美曲塞维持。")
add_bullet("鳞状NSCLC化疗方案：紫杉醇 175 mg/m² 或白蛋白结合型紫杉醇 100 mg/m²（d1、d8、d15）+ 卡铂 AUC 5–6，Q3W，最多4个周期。")

add_section("3. 研究人群")
add_body("目标人群：组织学或细胞学确诊的IV期（AJCC第8版）NSCLC，一线治疗，驱动基因阴性。计划随机约500例。")
add_body("主要入选标准：")
add_bullet("年龄≥18岁，签署书面知情同意书。")
add_bullet("组织学或细胞学确诊的IV期NSCLC，既往未接受过针对晚期/转移性疾病的系统性抗肿瘤治疗。")
add_bullet("无EGFR敏感突变、ALK融合、ROS1融合（非鳞状受试者须提供上述检测结果；鳞状受试者按方案规定执行）。")
add_bullet("至少有一处符合RECIST 1.1的可测量病灶。")
add_bullet("ECOG体能状态评分0或1；预期生存≥12周。")
add_bullet("器官功能满足方案规定的实验室标准（包括骨髓、肝、肾功能）。")
add_bullet("有生育能力的受试者须同意在研究期间及末次给药后按规定期限采取高效避孕措施。")
add_body("主要排除标准：")
add_bullet("小细胞肺癌成分或神经内分泌癌；已知有症状性中枢神经系统转移且未经治疗或未稳定。")
add_bullet("既往接受过抗PD-1/PD-L1/CTLA-4等免疫检查点抑制剂治疗。")
add_bullet("活动性自身免疫性疾病需系统免疫抑制治疗；需长期使用相当于泼尼松>10 mg/日的糖皮质激素。")
add_bullet("活动性乙型肝炎、丙型肝炎或活动性肺结核；人类免疫缺陷病毒感染。")
add_bullet("间质性肺病病史或当前需要类固醇治疗的非感染性肺炎。")
add_bullet("妊娠或哺乳期妇女；研究者判断不适合参加本研究的其他严重合并症。")

add_section("4. 给药方案")
add_body(
    "XS-201或安慰剂：200 mg，静脉输注，每3周给药一次。首次输注时间建议不少于60分钟，若耐受良好，后续输注可缩短至30分钟。"
    "含铂双药化疗在XS-201/安慰剂输注完成后当日给药，具体水化、止吐和预处理按各中心标准及所用化疗药物说明书执行。"
)
add_body(
    "剂量调整：XS-201不允许减量。出现免疫相关不良事件时，按方案规定暂停给药、给予糖皮质激素或其他免疫调节治疗，"
    "恢复后可继续原剂量。含铂化疗可按血液学或非血液学毒性进行减量或延迟。治疗最长持续时间建议不超过24个月，除非研究者与申办方评估后认为继续治疗对受试者有明确获益。"
)

add_section("5. 疗效终点")
add_body(
    "主要终点：由独立影像学评审委员会（BICR）按RECIST 1.1评估的无进展生存期（PFS），定义为自随机至疾病进展或任何原因死亡的时间，以先发生者为准。"
)
add_body("关键次要终点：总生存期（OS）。")
add_body("其他次要终点：")
add_bullet("研究者评估的PFS。")
add_bullet("BICR和研究者评估的客观缓解率（ORR）及缓解持续时间（DoR）。")
add_bullet("疾病控制率（DCR）；至缓解时间（TTR）。")
add_bullet("安全性与耐受性：不良事件、实验室检查、生命体征、心电图等。")
add_body("探索性终点：基线肿瘤组织PD-L1 TPS与PFS/OS/ORR的关系；血浆循环肿瘤DNA动态变化（如可行）。影像学评估频率：随机后每6周一次，直至第48周，此后每9周一次。")

add_section("6. 统计考虑")
add_body(
    "样本量：假设对照组中位PFS约为6个月，试验组可将中位PFS延长至约8.5–9个月（风险比HR=0.70）。"
    "在双侧α=0.05、把握度80%–85%、1:1随机、约10%脱落的前提下，计划随机约500例受试者，预计观察到约350件PFS事件时进行主要分析。"
)
add_body(
    "分析集：全分析集（FAS，所有随机受试者，按随机分组分析）用于主要疗效分析；符合方案集（PPS）作敏感性分析；"
    "安全性分析集包括至少接受过一次研究药物治疗的受试者。"
)
add_body(
    "主要分析方法：采用分层Log-rank检验比较两组PFS；用分层Cox比例风险模型估计HR及其95%置信区间；"
    "用Kaplan-Meier法估计中位PFS及生存曲线。OS分析采用同样框架。ORR比较采用分层Cochran-Mantel-Haenszel检验。"
    "期中分析：可在约70% PFS事件时进行一次疗效/无效性期中分析，α消耗采用Lan-DeMets（O’Brien-Fleming型）边界。"
    "多重性：主终点PFS与关键次要终点OS按预设层级检验控制I类错误。"
)

add_section("7. 安全性")
add_body(
    "不良事件按NCI-CTCAE 5.0分级，MedDRA编码。特别关注的不良事件（AESI）包括：免疫相关肺炎、肝炎、结肠炎、内分泌疾病"
    "（甲状腺功能异常、垂体炎、肾上腺功能不全、1型糖尿病）、重症皮肤反应、肾炎、心肌炎、输液反应等。"
)
add_body(
    "严重不良事件须在获知后24小时内报告申办方。免疫相关≥3级不良事件按方案给予糖皮质激素并暂停或永久停用XS-201。"
    "独立数据监查委员会（IDMC）将定期审阅安全性数据，必要时可建议调整研究。"
)
add_body(
    "随访：末次给药后30天进行安全性随访，其后每12周进行生存随访，直至死亡、失访或研究结束。妊娠事件按方案即时报告。"
)

add_section("8. 伦理与质量管理")
add_body(
    "本研究遵循《药物临床试验质量管理规范》（GCP）、赫尔辛基宣言及适用法规。所有中心须经伦理委员会批准后方可入组。"
    "受试者在任何研究相关操作前须签署知情同意书。数据和源文件按方案及监查计划进行核查。本摘要仅用于方案起草起点，正式方案以伦理批准版本为准。"
)

footer = doc.add_paragraph()
footer.paragraph_format.space_before = Pt(12)
r = footer.add_run("—— 方案摘要结束（供医学写作工作台导入测试使用）——")
set_run_font(r, size=9)
r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.save(out)
print(f"saved {out} bytes={out.stat().st_size}")
