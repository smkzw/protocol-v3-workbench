"""Isolated actual Office renderer + product snapshot API, synthetic data only."""
import base64
import io
import json
import os
from pathlib import Path
from datetime import datetime, timezone
from urllib.parse import urlencode, quote
from fastapi import FastAPI
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.testclient import TestClient
from docx import Document
import uvicorn
import integration_shared as shared
from test_mounted_api_integration import PROJECT, SD_ID, _create_body
from test_manuscript_edit_control import seed_working_document
from app.protocol_workflow.api.composition import ProtocolWorkflowMountConfig, mount_protocol_workflow_router
from app.protocol_workflow.application.manuscript_documents import ManuscriptDocumentService
from app.protocol_workflow.artifacts.local_store import LocalArtifactStore
from app.protocol_workflow.storage.sqlite import build_unit_of_work_factory
ROOT = Path.cwd()
RUN = Path(os.environ.get(
    'OFFICE_FIXTURE_RUN',
    ROOT / 'runs/requirements_v2_20260919/t17_round11/office_browser_isolated',
))
RUN.mkdir(parents=True, exist_ok=True)
db = RUN / 'product.sqlite'
app = FastAPI()
shared.admit(db, PROJECT)
mount_protocol_workflow_router(app, ProtocolWorkflowMountConfig(enabled=True, db_path=db))
service = ManuscriptDocumentService(build_unit_of_work_factory({'backend':'sqlite','path':str(db)}),
    clock=lambda: datetime.now(timezone.utc), office_store=LocalArtifactStore(str(db)+'.office-artifacts'))
base = f'/api/projects/{quote(PROJECT,safe="")}/protocol-workflow/study-definitions/{quote(SD_ID,safe="")}/manuscript-draft'
with TestClient(app) as client:
    if not (RUN/'seed.json').exists():
        result = client.post(f'/api/projects/{PROJECT}/protocol-workflow/study-definitions', json=_create_body())
        assert result.status_code == 200, result.text
        seed_working_document(service, {})
        current = service.current(PROJECT, SD_ID)
        doc = Document()
        doc.sections[0].header.paragraphs[0].text = '隔离验证 · 页眉保留'
        doc.sections[0].footer.paragraphs[0].text = '隔离验证 · 页脚保留'
        doc.add_heading('研究方案编辑验证', 0)
        doc.add_heading('一、研究设计', 1)
        doc.add_paragraph('这是合成验证文档，不用于申报。')
        doc.add_paragraph('保存前的原文。')
        table = doc.add_table(rows=2, cols=2)
        for cell,text in zip([table.cell(0,0),table.cell(0,1),table.cell(1,0),table.cell(1,1)], ['访视','检查','筛选期','合成检查项目']): cell.text=text
        buffer=io.BytesIO();doc.save(buffer)
        intent={'operation_id':'office:isolated:seed','actor_id':'user:example',
            'expected_revision':current['expected_revision'],'expected_document_sha256':current['expected_document_sha256'],
            'content_base64':base64.b64encode(buffer.getvalue()).decode(), 'base_artifact_revision':None,
            'opened_study_revision_sha256':'a'*64}
        response=client.post(base+'/office-draft/snapshots',json=intent)
        assert response.status_code == 201,response.text
        (RUN/'seed.json').write_text(json.dumps({'receipt':response.json(),'current':current},ensure_ascii=False))
        (RUN/'before.docx').write_bytes(buffer.getvalue())

app.mount('/genoffice', StaticFiles(directory=ROOT/'frontend/public/genoffice'),name='genoffice')
@app.get('/cover-fixed.docx')
def cover_fixture():
    return FileResponse(RUN/'cover-fixed.docx')

@app.get('/header-layout-aligned.docx')
def latest_layout_fixture():
    return FileResponse(RUN/'header-layout-aligned.docx')

@app.get('/api/r11-fixture')
def react_fixture_metadata():
    current = service.current(PROJECT, SD_ID)
    return {'projectId': PROJECT, 'studyDefinitionId': SD_ID,
        'actorId': 'user:example', 'savedDocument': {
            'document_sha256': current['expected_document_sha256'],
            'document': {'revision': current['expected_revision'],
                         'study_definition_sha256': 'a' * 64}}}

@app.get('/', response_class=HTMLResponse)
def page():
    latest=service.latest_office_snapshot(PROJECT,SD_ID)
    current=service.current(PROJECT,SD_ID)
    url=base+'/office-draft/snapshots/'+quote(latest['operation_id'],safe='')+'/content'
    query=urlencode({'docUrl':url,'docName':'隔离验证.docx','saveUrl':base+'/office-draft/snapshots',
        'rev':current['expected_revision'],'sha':current['expected_document_sha256'],'openedStudySha':'a'*64,'actor':'user:example'})
    return '<html lang="zh"><meta charset="utf-8"><title>隔离文档验证</title><body style="margin:0">'+f'<a id="download" href="{url}" download>下载已保存版本</a><span id="status">隔离合成文档</span><iframe id="office" title="隔离Office" style="width:100%;height:95vh;border:0" src="/genoffice/index.html?{query}"></iframe>'+'''<script>
window.addEventListener('message', e=>{if(e.origin!==location.origin||e.source!==document.querySelector('#office').contentWindow||e.data?.type!=='protocol-office:saved')return;const r=e.data.receipt;document.querySelector('#status').textContent='已保存版本 '+r.artifact_revision;document.querySelector('#download').href='''+json.dumps(base+'/office-draft/snapshots/')+'''+encodeURIComponent(r.operation_id)+'/content';});
</script></body></html>'''

if __name__=='__main__':
    uvicorn.run(app, host='127.0.0.1',
                port=int(os.environ.get('OFFICE_FIXTURE_PORT', '5293')),
                log_level='warning')
