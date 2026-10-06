#!/usr/bin/env python3
"""R3-D 第3波自备《方案摘要》：难治性慢性咳嗽 II 期 PX-590。"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).with_name("R3D-W3-慢性咳嗽-PX590-摘要.docx")

SECTIONS = [
    (
        "立项背景",
        "难治性慢性咳嗽（RCC）指咳嗽超过八周、常规检查未明病因、且对经验性治疗反应不足。"
        "气道感觉神经元 P2X3 受体介导咳嗽反射亢进，是近年被验证的干预靶点。"
        "现有止咳药对 RCC 覆盖不足。本摘要拟启动口服选择性 P2X3 拮抗剂 PX-590 的 II 期剂量探索。",
    ),
    (
        "研究目的",
        "主要目的：评价 PX-590 10 mg 与 30 mg 每日两次口服相较安慰剂，在第 8 周降低 24 小时客观咳嗽次数的疗效。"
        "次要目的：评价 LCQ、咳嗽 VAS 及应答率，并描述安全性，为 III 期剂量选择提供依据。",
    ),
    (
        "总体设计",
        "多中心、随机、双盲、安慰剂对照的剂量探索试验。三臂 1:1:1 分配：PX-590 10 mg bid、PX-590 30 mg bid、匹配安慰剂。"
        "双盲治疗期 8 周，随后 2 周安全性随访。片剂口服，每日两次。随机化分层可按中心实施。",
    ),
    (
        "研究人群",
        "入选 40–75 岁、咳嗽病程不少于 12 个月、胸片无明显异常、符合 RCC 定义的受试者。"
        "基线 24 小时客观咳嗽次数不少于 20 次。ACEI 诱发咳嗽、活动性呼吸道感染、明显肝肾功能不全及 P2X3 拮抗剂过敏者排除。",
    ),
    (
        "主要终点",
        "第 8 周 24 小时客观咳嗽次数相对基线的变化。咳嗽次数由经验证的咳嗽监测仪连续采集，分析集以全分析集为主，并报告符合方案集。",
    ),
    (
        "次要终点",
        "第 8 周 LCQ 总分与咳嗽 VAS 较基线变化；应答率（24 小时咳嗽次数下降不少于 30%）；"
        "治疗及随访期内不良事件、实验室检查和心电图（含 QT 间期）的发生与分级。",
    ),
    (
        "统计考虑",
        "1:1:1 随机。主要终点用混合效应模型或协方差分析估计各剂量相对安慰剂的差值及 95% 置信区间，不设正式期中分析。"
        "样本量兼顾主要终点估计精度与三臂安全性观察，例数在正式方案确定。缺失数据按预先规则处理。",
    ),
    (
        "安全性关注",
        "P2X3 拮抗剂类特征性不良事件为味觉障碍，另需关注口干、咽炎、肝酶升高及 QT 间期延长。"
        "试验期间定期采集生命体征、血生化、血常规及心电图；出现 3 级及以上肝酶升高或有临床意义的 QT 改变时按方案暂停或退出。",
    ),
]


def set_east_asia(run) -> None:
    run.font.name = "宋体"
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "宋体")


def main() -> None:
    doc = Document()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("PX-590 治疗难治性慢性咳嗽 II 期研究方案摘要")
    run.bold = True
    run.font.size = Pt(16)
    set_east_asia(run)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("项目代号：R3D-W3-慢性咳嗽-PX590")
    set_east_asia(r2)

    body_chars = 0
    for heading, text in SECTIONS:
        h = doc.add_paragraph()
        hr = h.add_run(heading)
        hr.bold = True
        hr.font.size = Pt(14)
        set_east_asia(hr)
        p = doc.add_paragraph()
        pr = p.add_run(text)
        set_east_asia(pr)
        body_chars += len(heading) + len(text)

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"section_body_chars={body_chars}")


if __name__ == "__main__":
    main()
