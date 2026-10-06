#!/usr/bin/env python3
"""Generate the R2-D wave-2 protocol synopsis (CKD-aP / KOR-228)."""
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).with_name("R2D-W2-CKDaP-KOR228-摘要.docx")

SECTIONS = [
    (
        "一、背景与立题依据",
        "慢性肾病相关瘙痒（CKD-aP）是维持性血液透析患者常见并发症，常导致睡眠障碍和生活质量下降。抗组胺药、润肤剂等对症手段对相当一部分患者效果有限。KOR-228为外周选择性κ阿片受体激动剂，拟抑制外周瘙痒信号传导，不激动μ受体，理论上无呼吸抑制与依赖风险。本试验拟在该人群评价疗效与安全性，为确证性研究提供剂量依据。",
    ),
    (
        "二、研究目的",
        "主要目的：评价KOR-228（2.5 mg bid及5 mg bid）较安慰剂在第8周降低周平均最重瘙痒数字评分（WI-NRS）的疗效。次要目的：评价WI-NRS应答率（改善≥4分）、5-D瘙痒量表、Skindex-16、PSQI睡眠及透析中不适，并描述安全性。探索剂量-效应趋势，为III期剂量选择提供依据。",
    ),
    (
        "三、总体设计",
        "多中心、随机、双盲、安慰剂对照、剂量探索的II期试验。约180例按1:1:1随机至KOR-228 2.5 mg bid、5 mg bid或匹配安慰剂，每日两次口服。8周双盲治疗期后4周安全性随访。随机分层拟包括基线周平均WI-NRS及透析龄。试验期间维持原透析处方及稳定对症用药。",
    ),
    (
        "四、研究人群",
        "拟纳入18–80岁、维持性血液透析≥3个月且每周透析3次者；入组前4周周平均最重瘙痒WI-NRS≥5，且常规抗组胺药和润肤剂无效或不耐受。主要排除：活动性恶性肿瘤、原发性皮肤病所致显著瘙痒、阿片类药物过敏、严重肝功能异常、妊娠或哺乳。允许稳定的透析常规合并用药。",
    ),
    (
        "五、主要与次要终点",
        "主要终点：第8周周平均WI-NRS较基线的变化。关键次要终点：第8周WI-NRS应答率（改善≥4分）。其他次要终点：5-D瘙痒量表、Skindex-16、PSQI、透析中不适评分的变化。安全性终点：不良事件、严重不良事件、停药事件，以及肝酶、电解质和透析中低血压。",
    ),
    (
        "六、统计考虑",
        "全分析集作主要分析。主要终点采用混合效应模型重复测量（MMRM）。II期不设严格多重性校正，报告点估计与95%置信区间。按第8周WI-NRS较安慰剂多下降约1.0–1.5分估算，每组约60例。不设正式有效性期中分析。",
    ),
    (
        "七、安全性关注",
        "重点关注头晕、嗜睡、口干、恶心，以及透析中低血压、电解质紊乱和肝酶升高。出现不能耐受的嗜睡、反复低血压或2级及以上肝酶升高时按方案减量、暂停或停药。不良事件按CTCAE分级。随访至末次给药后4周。",
    ),
]


def set_cjk_font(run, name: str = "Songti SC") -> None:
    run.font.name = name
    run.font.size = Pt(12)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), name)


def main() -> None:
    doc = Document()
    title = doc.add_heading(
        "KOR-228治疗维持性血液透析患者慢性肾病相关瘙痒（CKD-aP）II期研究方案摘要",
        level=0,
    )
    for run in title.runs:
        set_cjk_font(run, "Heiti SC")
        run.font.size = Pt(16)

    meta = doc.add_paragraph()
    r = meta.add_run(
        "试验分期：II期　研究药物：KOR-228（口服κ阿片受体激动剂）　适应症：慢性肾病相关瘙痒（维持性血液透析人群）"
    )
    set_cjk_font(r)

    body_chars = 0
    for heading, text in SECTIONS:
        h = doc.add_heading(heading, level=1)
        for run in h.runs:
            set_cjk_font(run, "Heiti SC")
            run.font.size = Pt(14)
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_cjk_font(run)
        body_chars += len(text)

    doc.save(OUT)
    print(f"wrote {OUT}")
    print(f"body_chars={body_chars} (target 600-900)")
    print(f"with_headings={body_chars + sum(len(h) for h, _ in SECTIONS)}")


if __name__ == "__main__":
    main()
