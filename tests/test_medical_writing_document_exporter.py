from __future__ import annotations

import base64
import hashlib
import io
import re
import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml.ns import qn

from packages.contracts.workbench_contracts import (
    ApprovalState,
    MedicalWritingGreenfieldCreateRequest,
    MedicalWritingGreenfieldSectionSeed,
    ProtocolDocument,
    ProtocolSection,
)
from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_document_exporter import (
    MedicalWritingDocumentDocxExportError,
    document_index_catalog,
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_greenfield import (
    _build_greenfield_document,
    _greenfield_company_synopsis_table,
)
from services.api.app.medical_writing_manifest import (
    DEFAULT_WRITING_PACKAGES,
    WRITING_PACKAGE_IDS_BY_PROJECT,
)
from services.api.app.medical_writing_tables import MedicalWritingTableService
from services.api.app.medical_writing_study_schema import (
    render_study_schema_png,
    render_study_schema_svg,
)
from services.api.app.medical_writing_style_profile import (
    MedicalWritingStyleProfileService,
)
from tests.test_medical_writing_study_schema import _pnh_schema


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WORD_NS}
SOA_CASES = (
    ("proj_rux_03_002", "docx:table:2"),
    ("proj_d001", "docx:table:6"),
    ("proj_my008_pnh_3_01", "docx:table:14"),
)


def test_protocol_synopsis_objectives_and_endpoints_export_as_numbered_nested_table():
    def ordered_list(*items: str) -> dict:
        return {
            "type": "doc",
            "content": [
                {
                    "type": "orderedList",
                    "attrs": {"numberingFormat": "decimal_half_paren"},
                    "content": [
                        {
                            "type": "listItem",
                            "content": [
                                {
                                    "type": "paragraph",
                                    "content": [{"type": "text", "text": item}],
                                }
                            ],
                        }
                        for item in items
                    ],
                }
            ],
        }

    group_id = "objectives_endpoints"
    synopsis = {
        "columns": ["项目", "目的", "相应的研究终点"],
        "nested_groups": [
            {
                "group_id": group_id,
                "label": "目的与估计目标/终点",
                "columns": ["目的", "相应的研究终点"],
                "sections": [],
            }
        ],
        "rows": [
            {
                "label": "研究题目",
                "values": ["一项代表性III期临床试验"],
                "merge_content": True,
            },
            {
                "label": "目的与估计目标/终点",
                "values": ["主要目的", "相应的研究终点"],
                "style_role": "section_header",
                "nested_group_id": group_id,
                "nested_group_start": True,
                "nested_group_row_span": 4,
                "nested_section_type": "primary",
            },
            {
                "label": "目的与估计目标/终点",
                "values": ["评价诱导治疗有效性。", "第12周临床缓解比例。"],
                "rich_values": [
                    ordered_list("评价诱导治疗有效性。"),
                    ordered_list("第12周临床缓解比例。"),
                ],
                "nested_group_id": group_id,
                "nested_group_continuation": True,
                "nested_section_type": "primary",
            },
            {
                "label": "目的与估计目标/终点",
                "values": ["次要目的", "相应的研究终点"],
                "style_role": "section_header",
                "nested_group_id": group_id,
                "nested_group_continuation": True,
                "nested_section_type": "secondary",
            },
            {
                "label": "目的与估计目标/终点",
                "values": [
                    "评价持续有效性。\n评价安全性。",
                    "第52周临床缓解比例。\nAE发生率。",
                ],
                "rich_values": [
                    ordered_list("评价持续有效性。", "评价安全性。"),
                    ordered_list("第52周临床缓解比例。", "AE发生率。"),
                ],
                "nested_group_id": group_id,
                "nested_group_continuation": True,
                "nested_section_type": "secondary",
            },
        ],
    }
    table_block = _greenfield_company_synopsis_table(
        section_id="section_synopsis_nested",
        document_id="doc_synopsis_nested",
        body_order=1,
        synopsis=synopsis,
        source_fact_ids=[],
    )
    document = ProtocolDocument(
        document_id="doc_synopsis_nested",
        project_id="proj_synopsis_nested",
        protocol_id="CMS-SYNOPSIS-NESTED",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_synopsis_nested",
                document_id="doc_synopsis_nested",
                heading="方案摘要",
                section_number="1.1",
                node_kind="protocol_synopsis",
                content_blocks=[table_block],
            )
        ],
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")

    assert result.metadata["tables"][0]["nested_synopsis_table_count"] == 1
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
        numbering_xml = archive.read("word/numbering.xml").decode("utf-8")
    outer_table = root.find("./w:body/w:tbl", NS)
    assert outer_table is not None
    nested_tables = outer_table.findall(".//w:tc/w:tbl", NS)
    assert len(nested_tables) == 1
    outer_rows = outer_table.findall("./w:tr", NS)
    assert all(
        row.find("./w:trPr/w:cantSplit", NS) is None
        for row in outer_rows[1:5]
    )
    nested_rows = nested_tables[0].findall("./w:tr", NS)
    assert len(nested_rows) == 4
    assert all(
        row.find("./w:trPr/w:cantSplit", NS) is not None
        for row in nested_rows
    )
    assert all(len(row.findall("./w:tc", NS)) == 2 for row in nested_rows)
    assert len(nested_tables[0].findall(".//w:numPr", NS)) == 6
    assert "%1)" in numbering_xml
    assert "%1）" in numbering_xml


def test_governed_study_schema_exports_svg_png_caption_bookmark_and_figure_metadata():
    schema = _pnh_schema()
    svg = render_study_schema_svg(schema)
    png = render_study_schema_png(svg)
    figure = {
        "block_id": "mwgenerated_figure_schema_pnh",
        "block_type": "figure",
        "figure_id": "mwfigure_schema_pnh",
        "figure_kind": "study_schema",
        "title": "研究设计概况",
        "alt_text": "研究流程图：研究设计概况",
        "body_order": 2,
        "source_kind": "medical_writing_study_schema",
        "source_locator": "generated:medical_writing_study_schema:schema_pnh:r1:layout0",
        "editable": False,
        "schema_id": schema.schema_id,
        "schema_revision": schema.revision,
        "schema_state_sha256": schema.state_sha256,
        "layout_revision": 0,
        "svg": svg,
        "svg_sha256": hashlib.sha256(svg.encode("utf-8")).hexdigest(),
        "png_base64": base64.b64encode(png).decode("ascii"),
        "png_sha256": hashlib.sha256(png).hexdigest(),
        "width_inches": 6.45,
        "bookmark_name": "mwfig_schema_pnh",
    }
    document = ProtocolDocument(
        document_id="doc_study_schema",
        project_id="proj_study_schema",
        protocol_id="CMS-STUDY-SCHEMA",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_1_2",
                document_id="doc_study_schema",
                heading="试验示意图",
                section_number="1.2",
                node_kind="study_schema",
                content_blocks=[
                    {"block_id": "heading_1_2", "block_type": "heading", "text": "1.2 试验示意图", "body_order": 1},
                    figure,
                ],
            )
        ],
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    figure_bookmark = document_index_catalog(document)["figures"][0]["bookmark_name"]

    assert result.metadata["figure_count"] == 1
    assert result.metadata["figures"][0]["caption"] == "图 1 研究设计概况"
    assert result.metadata["figures"][0]["vector_primary"] is True
    assert result.metadata["figures"][0]["orientation"] == "portrait"
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        names = archive.namelist()
        assert any(name.startswith("word/media/") and name.endswith(".svg") for name in names)
        assert any(name.startswith("word/media/") and name.endswith(".png") for name in names)
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_xml = archive.read("word/settings.xml").decode("utf-8")
    assert "asvg:svgBlip" in document_xml
    assert "SEQ 图" in document_xml
    assert 'TOC \\o "1-4" \\h \\z \\u' in document_xml
    assert 'TOC \\h \\z \\c "图"' in document_xml
    assert 'TOC \\h \\z \\c "表"' not in document_xml
    assert "w:updateFields" in settings_xml and 'w:val="1"' in settings_xml
    assert f'w:name="{figure_bookmark}"' in document_xml
    exported = Document(io.BytesIO(result.content))
    image_paragraph = next(
        paragraph for paragraph in exported.paragraphs if paragraph._p.xpath(".//w:drawing")
    )
    caption_paragraph = next(
        paragraph for paragraph in exported.paragraphs if "研究设计概况" in paragraph.text
    )
    assert caption_paragraph._p is image_paragraph._p
    assert image_paragraph.paragraph_format.keep_together is True
    assert image_paragraph.paragraph_format.keep_with_next is False
    assert image_paragraph.paragraph_format.line_spacing == 1.0
    assert image_paragraph.paragraph_format.space_before.pt == 0
    assert image_paragraph.paragraph_format.space_after.pt == 0
    assert image_paragraph._p.xpath("./w:r/w:br")
    assert "SEQ 图" in image_paragraph._p.xml

    wide_svg = re.sub(
        r'(<svg\b[^>]*\bwidth=")[0-9.]+"',
        r'\g<1>2288"',
        svg,
        count=1,
    )
    wide_svg = re.sub(
        r'(<svg\b[^>]*\bheight=")[0-9.]+"',
        r'\g<1>1564"',
        wide_svg,
        count=1,
    )
    wide_png = render_study_schema_png(wide_svg)
    wide_figure = {
        **figure,
        "svg": wide_svg,
        "svg_sha256": hashlib.sha256(wide_svg.encode("utf-8")).hexdigest(),
        "png_base64": base64.b64encode(wide_png).decode("ascii"),
        "png_sha256": hashlib.sha256(wide_png).hexdigest(),
    }
    wide_document = document.model_copy(deep=True)
    style_profile = MedicalWritingStyleProfileService().definition()
    wide_document = wide_document.model_copy(
        update={
            "style_profile_id": style_profile.style_profile_id,
            "style_profile_version": style_profile.style_profile_version,
            "style_profile_definition_sha256": style_profile.definition_sha256,
        },
        deep=True,
    )
    wide_document.sections[0] = wide_document.sections[0].model_copy(
        update={
            "content_blocks": [
                {"block_id": "wide_heading", "block_type": "heading", "text": "1.2 试验示意图", "body_order": 1},
                {
                    "block_id": "wide_intro",
                    "block_type": "paragraph",
                    "text": "研究示意图概括筛选、随机、治疗与随访流程。",
                    "body_order": 2,
                },
                {**wide_figure, "body_order": 3},
                {"block_id": "after_figure", "block_type": "heading", "text": "1.3 活动时间表", "body_order": 4},
            ]
        },
        deep=True,
    )
    wide_result = export_medical_writing_document_docx(wide_document, mode="draft_preview")
    wide_exported = Document(io.BytesIO(wide_result.content))
    assert wide_result.metadata["figures"][0]["orientation"] == "landscape"
    assert wide_result.metadata["figures"][0]["width_inches"] >= 8.0
    assert wide_exported.inline_shapes[0].width.inches >= 8.0
    assert [section.orientation for section in wide_exported.sections] == [
        0,
        1,
        0,
    ]
    wide_heading = next(
        paragraph
        for paragraph in wide_exported.paragraphs
        if paragraph.text == "1.2 试验示意图"
    )
    wide_intro = next(
        paragraph
        for paragraph in wide_exported.paragraphs
        if paragraph.text == "研究示意图概括筛选、随机、治疗与随访流程。"
    )
    assert wide_heading.paragraph_format.keep_together is True
    assert wide_heading.paragraph_format.keep_with_next is True
    assert wide_heading.paragraph_format.line_spacing == 1.0
    assert wide_heading.paragraph_format.space_before.pt == 0
    assert wide_heading.paragraph_format.space_after.pt == 0
    assert wide_intro.paragraph_format.keep_together is True
    assert wide_intro.paragraph_format.keep_with_next is True
    assert wide_intro.paragraph_format.line_spacing == 1.0
    assert wide_intro.paragraph_format.space_before.pt == 0
    assert wide_intro.paragraph_format.space_after.pt == 0
    landscape = wide_exported.sections[1]
    assert (
        landscape.top_margin.twips
        >= landscape.header_distance.twips + 790
    )
    with zipfile.ZipFile(io.BytesIO(wide_result.content)) as archive:
        wide_root = ElementTree.fromstring(archive.read("word/document.xml"))
    wide_body_items = list(wide_root.find("./w:body", NS))
    figure_index = next(
        index
        for index, element in enumerate(wide_body_items)
        if element.find(".//w:drawing", NS) is not None
    )
    figure_paragraph = wide_body_items[figure_index]
    assert figure_paragraph.find(".//w:br", NS) is not None
    assert "SEQ 图" in "".join(figure_paragraph.itertext())
    assert "研究设计概况" in "".join(figure_paragraph.itertext())
    assert figure_paragraph.find("./w:pPr/w:keepLines", NS) is not None
    assert figure_paragraph.find("./w:pPr/w:sectPr", NS) is None
    section_break_paragraph = wide_body_items[figure_index + 1]
    assert section_break_paragraph.find("./w:pPr/w:sectPr", NS) is not None
    assert section_break_paragraph.find(".//w:drawing", NS) is None
    assert "研究设计概况" not in "".join(section_break_paragraph.itertext())


def test_greenfield_cover_uses_upright_company_label_without_repeating_protocol_id():
    request = MedicalWritingGreenfieldCreateRequest(
        protocol_id="CMS-COVER-001",
        version="V1.0",
        document_title="一项代表性临床试验方案",
        indication="代表性适应症",
        study_phase="III期",
        investigational_product="CMS-COVER",
        protocol_date="2026年07月27日",
        sponsor="康哲药业",
        sections=[
            MedicalWritingGreenfieldSectionSeed(
                section_key="front_matter",
                heading="方案首页与版本信息",
                node_kind="front_matter",
                interaction_types=["front_matter_editor"],
                title_locked=True,
            )
        ],
        actor="medical_manager_test",
        idempotency_key="cover-style-regression",
    )
    document = _build_greenfield_document(
        "proj_cover_style",
        request,
        "c" * 64,
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    cover_label = next(
        paragraph for paragraph in exported.paragraphs if paragraph.text == "研究方案"
    )

    assert all(run.font.italic is not True for run in cover_label.runs)
    assert all(run.font.color.rgb is None or str(run.font.color.rgb) == "000000" for run in cover_label.runs)
    assert not any(
        paragraph.text.startswith(("研究方案.", "研究方案·"))
        for paragraph in exported.paragraphs
    )
    assert "CMS-COVER-001" not in cover_label.text


def test_real_document_index_catalog_excludes_layout_synopsis_and_unnumbered_source_tables():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    catalog = document_index_catalog(document)

    indexed_table_ids = {item["object_id"] for item in catalog["tables"]}
    tables_by_locator = {
        block.get("source_locator"): block
        for section in document.sections
        for block in section.content_blocks
        if block.get("block_type") == "table"
    }
    assert tables_by_locator["docx:table:0"]["table_id"] not in indexed_table_ids
    assert tables_by_locator["docx:table:1"]["table_id"] not in indexed_table_ids
    assert tables_by_locator["docx:table:2"]["table_id"] not in indexed_table_ids
    assert tables_by_locator["docx:table:3"]["table_id"] not in indexed_table_ids
    assert tables_by_locator["docx:table:4"]["table_id"] in indexed_table_ids
    assert [item["number"] for item in catalog["tables"]] == list(
        range(1, len(catalog["tables"]) + 1)
    )


def test_explicit_layout_role_overrides_numbered_title_for_table_catalog():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target = next(
        block
        for section in document.sections
        for block in section.content_blocks
        if block.get("source_locator") == "docx:table:4"
    )
    target["structured_table"]["role"] = "layout"

    indexed_ids = {
        item["object_id"] for item in document_index_catalog(document)["tables"]
    }

    assert target["table_id"] not in indexed_ids


def test_rux_static_table_list_is_replaced_by_dynamic_toc_and_seq_bookmarks():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    result = export_medical_writing_document_docx(document, mode="draft_preview")

    assert result.metadata["toc_field"] is True
    assert result.metadata["table_index_field"] is True
    assert result.metadata["skipped_source_index_block_count"] >= 9
    assert result.metadata["indexed_table_count"] == 8
    source_image = next(
        block
        for section in document.sections
        for block in section.content_blocks
        if block.get("figure_kind") == "source_docx_image"
        and block.get("title") == "表 8 SCORAD-主观症状评分"
    )
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml_bytes = archive.read("word/document.xml")
        document_xml = document_xml_bytes.decode("utf-8")
        exported_media_hashes = {
            hashlib.sha256(archive.read(name)).hexdigest()
            for name in archive.namelist()
            if name.startswith("word/media/") and not name.endswith("/")
        }
    assert 'TOC \\h \\z \\c "表"' in document_xml
    assert document_xml.count("SEQ 表") == 8
    assert "表 8 SCORAD-主观症状评分87" not in document_xml
    assert document_xml.count('w:name="_MWTAB_') == 8
    assert source_image["image_sha256"] in exported_media_hashes
    assert result.metadata["figures"][-1]["indexed_as"] == "table"

    root = ElementTree.fromstring(document_xml_bytes)
    parents = {
        child: parent
        for parent in root.iter()
        for child in parent
    }
    field_children = list(root.iter(f"{{{WORD_NS}}}fldChar")) + list(
        root.iter(f"{{{WORD_NS}}}instrText")
    )
    assert field_children
    assert all(
        parents[element].tag == f"{{{WORD_NS}}}r"
        for element in field_children
    )


def test_cross_reference_mark_exports_server_resolved_ref_field():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target = document_index_catalog(document)["tables"][0]
    block = next(
        block
        for section in document.sections
        if section.heading not in {"目录", "表格列表", "参考文献"}
        for block in section.content_blocks
        if block.get("block_type") == "paragraph" and block.get("text")
    )
    block["text"] = f"见表 {target['number']}。"
    block["rich_text"] = {
        "type": "paragraph",
        "content": [
            {"type": "text", "text": "见"},
            {
                "type": "text",
                "text": f"表 {target['number']}",
                "marks": [
                    {
                        "type": "crossReference",
                        "attrs": {
                            "targetKind": "table",
                            "targetId": target["object_id"],
                        },
                    }
                ],
            },
            {"type": "text", "text": "。"},
        ],
    }

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert f"REF {target['bookmark_name']} \\h" in document_xml
    assert f'w:name="{target["bookmark_name"]}"' in document_xml


def test_cross_reference_to_missing_target_fails_closed():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    block = next(
        block
        for section in document.sections
        if section.heading not in {"目录", "表格列表", "参考文献"}
        for block in section.content_blocks
        if block.get("block_type") == "paragraph" and block.get("text")
    )
    block["text"] = "见表 99。"
    block["rich_text"] = {
        "type": "paragraph",
        "content": [
            {"type": "text", "text": "见"},
            {
                "type": "text",
                "text": "表 99",
                "marks": [
                    {
                        "type": "crossReference",
                        "attrs": {
                            "targetKind": "table",
                            "targetId": "ptbl_missing_target",
                        },
                    }
                ],
            },
            {"type": "text", "text": "。"},
        ],
    }

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="cross-reference target is not available",
    ):
        export_medical_writing_document_docx(document, mode="draft_preview")


def test_cross_reference_renumbers_after_target_reorder_without_changing_bookmark():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    original_catalog = document_index_catalog(document)
    original_target = next(
        item
        for item in original_catalog["tables"]
        if item["title"] == "SCORAD-主观症状评分"
    )
    section = next(
        item
        for item in document.sections
        if any(
            block.get("figure_kind") == "source_docx_image"
            for block in item.content_blocks
        )
    )
    image_block = next(
        block
        for block in section.content_blocks
        if block.get("figure_kind") == "source_docx_image"
    )
    native_table = next(
        block
        for block in section.content_blocks
        if block.get("block_type") == "table"
        and block.get("title") == "表 7 SCORAD-皮损严重程度评分"
    )
    image_block["body_order"], native_table["body_order"] = (
        native_table["body_order"],
        image_block["body_order"],
    )
    paragraph = next(
        block
        for block in section.content_blocks
        if block.get("block_type") == "paragraph" and block.get("text")
    )
    paragraph["text"] = "重排后的目标见表 8。"
    paragraph["rich_text"] = {
        "type": "paragraph",
        "content": [
            {"type": "text", "text": "重排后的目标见"},
            {
                "type": "text",
                "text": "表 8",
                "marks": [
                    {
                        "type": "crossReference",
                        "attrs": {
                            "targetKind": "table",
                            "targetId": original_target["object_id"],
                        },
                    }
                ],
            },
            {"type": "text", "text": "。"},
        ],
    }

    reordered_catalog = document_index_catalog(document)
    reordered_target = next(
        item
        for item in reordered_catalog["tables"]
        if item["object_id"] == original_target["object_id"]
    )
    assert original_target["number"] == 8
    assert reordered_target["number"] == 7
    assert reordered_target["bookmark_name"] == original_target["bookmark_name"]

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        root = ElementTree.fromstring(archive.read("word/document.xml"))
    ref_instruction = f"REF {original_target['bookmark_name']} \\h"
    ref_paragraph = next(
        node
        for node in root.findall(".//w:p", NS)
        if ref_instruction in "".join(node.itertext())
    )
    assert "".join(ref_paragraph.findall(".//w:t", NS)[-2].itertext()) == "7"
    assert original_target["bookmark_name"] in {
        node.get(qn("w:name"))
        for node in root.findall(".//w:bookmarkStart", NS)
    }


@pytest.mark.parametrize(("project_id", "table_locator"), SOA_CASES)
def test_real_protocols_export_paragraph_table_paragraph_order_merges_headers_and_notes(
    project_id,
    table_locator,
):
    source_path = _source_path(project_id)
    source_hash = _file_hash(source_path)
    document, context = _document_with_structured_table(
        project_id,
        table_locator,
        repeat_header_rows=3,
    )
    approved = _approved(document)
    snapshot = approved.model_dump_json(exclude_none=False)
    approval_reference = f"approval-{project_id}-r1"

    result = export_medical_writing_document_docx(
        approved,
        mode="approved_final",
        approval_reference=approval_reference,
    )

    assert approved.model_dump_json(exclude_none=False) == snapshot
    assert _file_hash(source_path) == source_hash
    assert result.metadata["mode"] == "approved_final"
    assert result.metadata["approval_reference"] == approval_reference
    assert (
        result.metadata["source_snapshot_sha256"]
        == hashlib.sha256(snapshot.encode("utf-8")).hexdigest()
    )

    exported = Document(io.BytesIO(result.content))
    items = _body_items(exported)
    table_meta = next(
        table
        for table in result.metadata["tables"]
        if table["table_id"] == context["table_id"]
    )
    expected_caption = table_meta["caption"] or context["caption"]
    caption_index = _paragraph_index(items, expected_caption)
    next_paragraph_index = _paragraph_index(items, context["next_text"])
    table_index, table_element = next(
        (index, element)
        for index, (kind, _, element) in enumerate(items)
        if kind == "table" and caption_index < index < next_paragraph_index
    )
    note_index = _paragraph_containing_index(
        items,
        context["note_text"],
        start_after=table_index,
    )
    assert caption_index < table_index < note_index < next_paragraph_index
    caption_element = items[caption_index][2]
    assert caption_element.find("w:pPr/w:sectPr", NS) is None

    note_region = [
        text
        for kind, text, _ in items[table_index + 1 : next_paragraph_index]
        if kind == "paragraph" and text
    ]
    assert note_region[0] == "附注"
    assert len(note_region) - 1 == context["deduplicated_note_count"]

    paragraph_texts = [text for kind, text, _ in items if kind == "paragraph"]
    assert paragraph_texts.count(expected_caption) == 1
    assert (
        sum(
            context["note_text"] in text
            for kind, text, _ in items[note_index:next_paragraph_index]
            if kind == "paragraph"
        )
        == 1
    )
    assert "附注" in paragraph_texts
    assert not any("|---" in text for text in paragraph_texts)

    rows = table_element.findall("w:tr", NS)
    assert all(
        rows[index].find("w:trPr/w:tblHeader", NS) is not None for index in range(3)
    )
    assert table_element.findall(".//w:gridSpan", NS) or table_element.findall(
        ".//w:vMerge", NS
    )

    assert table_meta["repeat_header_rows"] == 3
    assert table_meta["merged_ranges"]
    assert context["note_marker"] in table_meta["note_markers"]
    assert len(table_meta["note_markers"]) == context["deduplicated_note_count"]
    assert len(exported.tables) == result.metadata["table_count"]

    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        headers = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )
    assert "方案终稿" in headers
    assert "DRAFT PREVIEW" not in headers
    assert approval_reference not in headers
    assert all(
        section.top_margin.twips >= section.header_distance.twips + 790
        for section in exported.sections
    )


def test_real_generic_table_uses_same_inline_renderer_without_nested_docx():
    document, context = _document_with_structured_table(
        "proj_rux_03_002",
        "docx:table:4",
        repeat_header_rows=1,
        keep_only_target_section=True,
    )

    result = export_medical_writing_document_docx(
        _approved(document),
        mode="approved_final",
        approval_reference="approval-rux-generic-r1",
    )

    exported = Document(io.BytesIO(result.content))
    items = _body_items(exported)
    caption_index = _paragraph_index(items, context["caption"])
    note_index = _paragraph_containing_index(items, context["note_text"])
    table_element = next(
        element
        for index, (kind, _, element) in enumerate(items)
        if kind == "table" and caption_index < index < note_index
    )
    assert caption_index < note_index
    assert table_element.find("w:tr/w:trPr/w:tblHeader", NS) is not None
    assert table_element.findall(".//w:gridSpan", NS) or table_element.findall(
        ".//w:vMerge", NS
    )
    table_meta = next(
        table
        for table in result.metadata["tables"]
        if table["table_id"] == context["table_id"]
    )
    assert table_meta["orientation"] == "portrait"
    assert context["note_marker"] in table_meta["note_markers"]
    assert len(exported.tables) == result.metadata["table_count"]


def test_draft_preview_is_visibly_marked_and_final_export_is_approval_gated():
    source = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")

    draft = export_medical_writing_document_docx(source, mode="draft_preview")
    with zipfile.ZipFile(io.BytesIO(draft.content)) as archive:
        headers = "".join(
            archive.read(name).decode("utf-8")
            for name in archive.namelist()
            if name.startswith("word/header") and name.endswith(".xml")
        )
    assert "草稿预览" in headers
    assert "待医学批准" not in headers
    assert "DRAFT PREVIEW" in headers
    assert draft.metadata["approval_reference"] == ""
    assert (
        export_medical_writing_document_docx(
            source,
            mode="draft_preview",
        ).content
        == draft.content
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="requires an approval reference",
    ):
        export_medical_writing_document_docx(source, mode="approved_final")
    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="contains unapproved sections",
    ):
        export_medical_writing_document_docx(
            source,
            mode="approved_final",
            approval_reference="approval-unapproved-r1",
        )


def test_rich_text_paragraph_attrs_and_inline_marks_are_exported_to_word():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target = _replace_first_text_block(
        document,
        text="一级标题 正文 H2O 2 黄色 蓝色",
        rich_text={
            "type": "paragraph",
            "attrs": {
                "stylePreset": "body",
                "textAlign": "justify",
                "lineHeight": 1.5,
                "spacingBeforePt": 6,
                "spacingAfterPt": 8,
                "leftIndentChars": 1,
                "rightIndentChars": 0.5,
                "firstLineIndentChars": 2,
            },
            "content": [
                {
                    "type": "text",
                    "text": "一级标题",
                    "marks": [
                        {"type": "bold"},
                        {"type": "underline"},
                        {
                            "type": "textStyle",
                            "attrs": {
                                "fontFamily": "黑体",
                                "fontSize": "16pt",
                                "color": "#C00000",
                            },
                        },
                    ],
                },
                {"type": "text", "text": " 正文 H"},
                {"type": "text", "text": "2", "marks": [{"type": "subscript"}]},
                {"type": "text", "text": "O "},
                {"type": "text", "text": "2", "marks": [{"type": "superscript"}]},
                {"type": "text", "text": " 黄色", "marks": [{"type": "highlight", "attrs": {"color": "#FFFF00"}}]},
                {"type": "text", "text": " 蓝色", "marks": [{"type": "textStyle", "attrs": {"color": "#4472C4"}}]},
            ],
        },
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    paragraph = next(item for item in exported.paragraphs if item.text == target["text"])

    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paragraph.paragraph_format.line_spacing == 1.5
    assert paragraph.paragraph_format.space_before.pt == pytest.approx(6)
    assert paragraph.paragraph_format.space_after.pt == pytest.approx(8)
    assert paragraph.paragraph_format.left_indent.pt == pytest.approx(10.5)
    assert paragraph.paragraph_format.right_indent.pt == pytest.approx(5.25)
    assert paragraph.paragraph_format.first_line_indent.pt == pytest.approx(21)
    assert paragraph._p.pPr.ind.get(qn("w:firstLineChars")) == "200"

    runs = paragraph.runs
    assert runs[0].bold is True
    assert runs[0].underline is True
    assert runs[0].font.size.pt == pytest.approx(16)
    assert str(runs[0].font.color.rgb) == "C00000"
    assert runs[2].font.subscript is True
    assert paragraph.runs[4].font.superscript is True
    assert paragraph.runs[5]._element.rPr.highlight.val == WD_COLOR_INDEX.YELLOW
    assert str(paragraph.runs[6].font.color.rgb) == "4472C4"
    for run in (run for run in runs if run.text):
        assert run.font.name == "Times New Roman"
        assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
        assert run._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
        assert run._element.rPr.rFonts.get(qn("w:hAnsi")) == "Times New Roman"


def test_rich_text_heading_style_and_list_semantics_are_exported():
    document = MedicalWritingDocumentService().document_for_revision("proj_d001")
    _replace_first_text_block(
        document,
        text="研究目的",
        rich_text={
            "type": "heading",
            "attrs": {
                "level": 3,
                "stylePreset": "heading_3",
                "textAlign": "left",
                "lineHeight": 1.0,
                "spacingBeforePt": 10,
                "spacingAfterPt": 4,
                "leftIndentChars": 0,
                "rightIndentChars": 0,
                "firstLineIndentChars": 0,
            },
            "content": [{"type": "text", "text": "研究目的"}],
        },
    )
    _replace_next_text_block(
        document,
        after_text="研究目的",
        text="主要目的次要目的",
        rich_text={
            "type": "bulletList",
            "content": [
                {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "主要目的"}]}]},
                {"type": "listItem", "content": [{"type": "paragraph", "content": [{"type": "text", "text": "次要目的"}]}]},
            ],
        },
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    heading = next(item for item in exported.paragraphs if item.text == "研究目的")
    list_items = [
        item
        for item in exported.paragraphs
        if item.text in {"主要目的", "次要目的"} and item.style.name == "List Bullet"
    ]

    assert heading.style.name == "Heading 3"
    assert heading.paragraph_format.keep_with_next is True
    assert heading._p.pPr.ind.get(qn("w:firstLineChars")) == "0"
    assert [item.style.name for item in list_items] == ["List Bullet", "List Bullet"]


def test_rich_text_document_root_exports_enter_as_separate_word_paragraphs():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    _replace_first_text_block(
        document,
        text="第一段。\n第二段。",
        rich_text={
            "type": "doc",
            "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": "第一段。"}]},
                {"type": "paragraph", "content": [{"type": "text", "text": "第二段。"}]},
            ],
        },
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    paragraph_text = [paragraph.text for paragraph in exported.paragraphs]

    assert "第一段。" in paragraph_text
    assert "第二段。" in paragraph_text
    assert paragraph_text.index("第二段。") == paragraph_text.index("第一段。") + 1


def test_rich_text_paragraph_with_newline_does_not_use_plain_text_split_path():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target = _replace_first_text_block(
        document,
        text="第一行富文本。\n第二行富文本。",
        rich_text={
            "type": "paragraph",
            "content": [
                {
                    "type": "text",
                    "text": "第一行富文本。\n第二行富文本。",
                    "marks": [{"type": "bold"}],
                }
            ],
        },
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    matching = [
        paragraph
        for paragraph in exported.paragraphs
        if "第一行富文本。" in paragraph.text
        or "第二行富文本。" in paragraph.text
    ]

    assert len(matching) == 1
    assert matching[0].text == "第一行富文本。\n第二行富文本。"
    assert all(run.bold is True for run in matching[0].runs if run.text)
    assert target["text"] == "第一行富文本。\n第二行富文本。"


def test_plain_greenfield_multiline_body_exports_as_indented_word_paragraphs():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target = _replace_first_text_block(
        document,
        text="第一段正文。\n第二段正文。\n\n第三段正文。",
        rich_text=None,
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    paragraph_text = [paragraph.text for paragraph in exported.paragraphs]
    first_index = paragraph_text.index("第一段正文。")
    emitted = exported.paragraphs[first_index : first_index + 3]

    assert [paragraph.text for paragraph in emitted] == [
        "第一段正文。",
        "第二段正文。",
        "第三段正文。",
    ]
    for paragraph in emitted:
        assert paragraph.style.name == "Normal"
        assert paragraph._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
        assert int(paragraph._p.pPr.ind.get(qn("w:firstLine"))) > 0
    assert target["text"] == "第一段正文。\n第二段正文。\n\n第三段正文。"


def test_rich_text_unsupported_mark_fails_closed_instead_of_silent_format_loss():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    _replace_first_text_block(
        document,
        text="不允许链接",
        rich_text={
            "type": "paragraph",
            "content": [
                {
                    "type": "text",
                    "text": "不允许链接",
                    "marks": [{"type": "link", "attrs": {"href": "https://example.com"}}],
                }
            ],
        },
    )

    with pytest.raises(MedicalWritingDocumentDocxExportError, match="unsupported rich-text mark"):
        export_medical_writing_document_docx(document, mode="draft_preview")


def test_markdown_table_in_paragraph_is_rejected_instead_of_exported_raw():
    section = ProtocolSection(
        section_id="section_markdown",
        document_id="document_markdown",
        heading="研究设计",
        approval_state=ApprovalState.MEDICALLY_APPROVED,
        content_blocks=[
            {
                "block_id": "block_markdown",
                "block_type": "paragraph",
                "body_order": 1,
                "text": "|访视|D1|\n|---|---|\n|血常规|X|",
            }
        ],
    )
    document = ProtocolDocument(
        document_id="document_markdown",
        project_id="project_markdown",
        protocol_id="TEST-001",
        version="V1.0",
        status="medically_approved",
        sections=[section],
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="unrendered Markdown table",
    ):
        export_medical_writing_document_docx(
            document,
            mode="approved_final",
            approval_reference="approval-markdown-r1",
        )


def test_legacy_greenfield_objective_marker_is_resolved_without_touching_body_text():
    section = ProtocolSection(
        section_id="section_primary_objective",
        document_id="document_objective",
        heading="主要目的 <#>",
        template_node_id="ich_m11_3_1_1",
        section_number="3.1.1",
        repeatable=True,
        title_locked=True,
        content_blocks=[
            {
                "block_id": "block_primary_objective_heading",
                "block_type": "heading",
                "body_order": 1,
                "text": "主要目的 <#>",
            },
            {
                "block_id": "block_primary_objective_body",
                "block_type": "paragraph",
                "body_order": 2,
                "text": "入选标准条目 <#> 由结构化编辑器另行实例化。",
            },
        ],
    )
    document = ProtocolDocument(
        document_id="document_objective",
        project_id="project_objective",
        protocol_id="TEST-OBJECTIVE-001",
        version="V0.1",
        sections=[section],
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    texts = [paragraph.text for paragraph in exported.paragraphs]

    assert "主要目的" in texts
    assert "主要目的 <#>" not in texts
    assert "入选标准条目 <#> 由结构化编辑器另行实例化。" in texts


def test_custom_repeatable_objective_heading_is_not_overwritten_on_export():
    section = ProtocolSection(
        section_id="section_custom_objective",
        document_id="document_custom_objective",
        heading="主要目的：维持临床缓解",
        template_node_id="ich_m11_3_1_1",
        section_number="3.1.1",
        repeatable=True,
        title_locked=True,
        content_blocks=[
            {
                "block_id": "block_custom_objective_heading",
                "block_type": "heading",
                "body_order": 1,
                "text": "主要目的：维持临床缓解",
            }
        ],
    )
    document = ProtocolDocument(
        document_id="document_custom_objective",
        project_id="project_custom_objective",
        protocol_id="TEST-OBJECTIVE-002",
        version="V0.1",
        sections=[section],
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))

    assert "主要目的：维持临床缓解" in [
        paragraph.text for paragraph in exported.paragraphs
    ]


def test_edited_structured_table_with_grid_hole_is_not_silently_repaired():
    document = MedicalWritingDocumentService().document_for_revision("proj_rux_03_002")
    target_locator = "docx:table:3"
    found = False
    for section_index, section in enumerate(document.sections):
        updated_blocks = list(section.content_blocks)
        for block_index, block in enumerate(updated_blocks):
            if block.get("source_locator") != target_locator:
                continue
            changed = dict(block)
            structure = dict(changed.get("structured_table") or {})
            structure["version"] = 1
            changed["structured_table"] = structure
            updated_blocks[block_index] = changed
            document.sections[section_index] = section.model_copy(
                update={"content_blocks": updated_blocks},
                deep=True,
            )
            found = True
            break
        if found:
            break
    assert found

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="uncovered grid positions",
    ):
        export_medical_writing_document_docx(document, mode="draft_preview")


def _document_with_structured_table(
    project_id: str,
    table_locator: str,
    *,
    repeat_header_rows: int,
    keep_only_target_section: bool = False,
):
    document = MedicalWritingDocumentService().document_for_revision(project_id)
    table_service = MedicalWritingTableService()
    target_section_index = -1
    context = {}
    globally_ordered_blocks = sorted(
        (block for section in document.sections for block in section.content_blocks),
        key=lambda block: block.get("body_order", 10**9),
    )

    for section_index, section in enumerate(document.sections):
        for block_index, block in enumerate(section.content_blocks):
            if block.get("source_locator") != table_locator:
                continue
            structured = table_service.from_table_block(block)
            caption = structured.title.strip()
            assert caption
            assert structured.notes
            note = structured.notes[0]
            note_text = note.text.strip()
            assert note_text
            structured.header_row_count = min(repeat_header_rows, len(structured.rows))
            structured.word_layout.update(
                {
                    "orientation": "auto",
                    "repeat_header_rows": structured.header_row_count,
                    "long_table_split_strategy": "repeat_header_keep_rows_intact",
                }
            )
            for row in structured.rows[: structured.header_row_count]:
                row.style_role = "header"
                for cell in row.cells:
                    cell.style_role = "header"
            updated_block = table_service.to_table_block(structured)
            updated_blocks = list(section.content_blocks)
            updated_blocks[block_index] = updated_block
            document.sections[section_index] = section.model_copy(
                update={"content_blocks": updated_blocks},
                deep=True,
            )
            target_section_index = section_index
            target_global_index = next(
                index
                for index, candidate in enumerate(globally_ordered_blocks)
                if candidate.get("source_locator") == table_locator
            )
            structured_note_texts = {
                _normalized_test_note(item.text) for item in structured.notes
            }
            next_text = next(
                str(candidate.get("text") or "").strip()
                for candidate in globally_ordered_blocks[target_global_index + 1 :]
                if candidate.get("block_type") in {"paragraph", "heading"}
                and str(candidate.get("text") or "").strip()
                and _strip_note_prefix(str(candidate.get("text") or "").strip())
                and _normalized_test_note(str(candidate.get("text") or ""))
                not in structured_note_texts
            )
            deduplicated_note_keys = {
                (item.marker.strip().casefold(), _normalized_test_note(item.text))
                for item in structured.notes
            }
            context = {
                "caption": caption,
                "note_text": note_text,
                "next_text": next_text,
                "table_id": structured.table_id,
                "note_marker": note.marker.strip(),
                "deduplicated_note_count": len(deduplicated_note_keys),
            }
            break
        if context:
            break

    assert context, f"real table not found: {project_id}/{table_locator}"
    if keep_only_target_section:
        document = document.model_copy(
            update={"sections": [document.sections[target_section_index]]},
            deep=True,
        )
    return document, context


def _replace_first_text_block(document, *, text: str, rich_text: dict):
    for section_index, section in enumerate(document.sections):
        for block_index, block in enumerate(section.content_blocks):
            if block.get("block_type") not in {"paragraph", "heading"}:
                continue
            changed = dict(block)
            changed["text"] = text
            changed["rich_text"] = rich_text
            blocks = list(section.content_blocks)
            blocks[block_index] = changed
            document.sections[section_index] = section.model_copy(
                update={"content_blocks": blocks},
                deep=True,
            )
            return changed
    raise AssertionError("real protocol contains no text block")


def _replace_next_text_block(document, *, after_text: str, text: str, rich_text: dict):
    found_previous = False
    for section_index, section in enumerate(document.sections):
        blocks = list(section.content_blocks)
        for block_index, block in enumerate(blocks):
            if block.get("text") == after_text:
                found_previous = True
                continue
            if not found_previous or block.get("block_type") not in {"paragraph", "heading"}:
                continue
            changed = dict(block)
            changed["text"] = text
            changed["rich_text"] = rich_text
            blocks[block_index] = changed
            document.sections[section_index] = section.model_copy(
                update={"content_blocks": blocks},
                deep=True,
            )
            return changed
    raise AssertionError("real protocol contains no second text block")


def _approved(document: ProtocolDocument) -> ProtocolDocument:
    return document.model_copy(
        update={
            "status": "medically_approved",
            "sections": [
                section.model_copy(
                    update={"approval_state": ApprovalState.MEDICALLY_APPROVED},
                    deep=True,
                )
                for section in document.sections
            ],
        },
        deep=True,
    )


def _body_items(document: Document):
    items = []
    for element in document.element.body.iterchildren():
        if element.tag == qn("w:p"):
            kind = "paragraph"
        elif element.tag == qn("w:tbl"):
            kind = "table"
        else:
            continue
        text = "".join(node.text or "" for node in element.findall(".//w:t", NS))
        items.append((kind, text, element))
    return items


def _paragraph_index(items, expected_text: str) -> int:
    return next(
        index
        for index, (kind, text, _) in enumerate(items)
        if kind == "paragraph" and text == expected_text
    )


def _paragraph_containing_index(
    items,
    expected_text: str,
    *,
    start_after: int = -1,
) -> int:
    return next(
        index
        for index, (kind, text, _) in enumerate(items)
        if index > start_after and kind == "paragraph" and expected_text in text
    )


def _strip_note_prefix(value: str) -> str:
    for prefix in ("附注：", "附注:", "备注：", "备注:", "注：", "注:"):
        if value.startswith(prefix):
            return value[len(prefix) :].strip()
    return value.strip()


def _normalized_test_note(value: str) -> str:
    return "".join(_strip_note_prefix(value).split()).casefold()


def _source_path(project_id: str):
    package_ids = WRITING_PACKAGE_IDS_BY_PROJECT[project_id]
    return next(
        package.protocol_path
        for package in DEFAULT_WRITING_PACKAGES
        if package.package_id in package_ids
    )


def _file_hash(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()
