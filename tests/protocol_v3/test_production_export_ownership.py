"""Production export serializes the current version faithfully (R4/T13).

Export owns layout only: template front matter, bookmarks, orientation,
placeholder real values.  It must not re-project the synopsis, swap the SOA
table, silently replace terms, delete marker-hit paragraphs, or write
"不适用" based on block kind (B01/A13/A14/A19).
"""
import json

from app.protocol_workflow.agent3.word_export_production import render_production_docx
from app.protocol_workflow.registries.template_runtime import default_template_root


def _template_paths():
    root = default_template_root()
    meta = json.loads((root / 'template.json').read_text(encoding='utf-8'))
    return meta['source']['path'], root


def _document(blocks):
    return {'study_definition_id': 'study:v3:export', 'revision': 1,
            'semantic_blocks': blocks}


def _para(node, text):
    return {'semantic_node_id': node, 'block_kind': 'paragraph', 'content': text}


def _table(node, cell_text):
    payload = {'schema_version': 'semantic-structured-table.v1', 'table': {
        'columns': [{'column_id': 'c1', 'order': 1}],
        'header_row_count': 1,
        'rows': [{'order': 1, 'cells': [{'column_id': 'c1', 'order': 1, 'text': '访视'}]},
                 {'order': 2, 'cells': [{'column_id': 'c1', 'order': 2, 'text': cell_text}]}],
    }}
    return {'semantic_node_id': node, 'block_kind': 'table',
            'content': json.dumps(payload, ensure_ascii=False)}


def _all_text(doc):
    parts = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(para.text for para in cell.paragraphs)
    return '\n'.join(parts)


def test_bookmark_preserves_paragraph_schema_order_and_stable_identity():
    import hashlib
    from docx import Document
    from docx.oxml.ns import qn
    from app.protocol_workflow.agent3.word_export_production import _bookmark

    paragraph = Document().add_heading('稳定标题', level=1)
    _bookmark(paragraph, 'chap_v2_n_1_1')

    children = list(paragraph._p)
    assert children[0].tag == qn('w:pPr')
    assert children[1].tag == qn('w:bookmarkStart')
    expected_id = str(
        int(hashlib.sha256(b'chap_v2_n_1_1').hexdigest()[:8], 16)
        % 2_147_483_647
    )
    assert children[1].get(qn('w:id')) == expected_id
    assert children[-1].tag == qn('w:bookmarkEnd')
    assert children[-1].get(qn('w:id')) == expected_id


def test_table_first_chapter_never_gets_not_applicable_text(tmp_path):
    template_path, template_dir = _template_paths()
    blocks = [_para('v2_n_11_1_1', '第一章正文。'),
              _table('v2_n_11_1_2', '表格首块数据')]
    out = tmp_path / 'out.docx'
    render_production_docx(template_path, template_dir, _document(blocks),
                           out, {'research.input_context': {}})
    from docx import Document as open_docx
    text = _all_text(open_docx(str(out)))
    assert '本节不适用于本研究' not in text
    assert '表格首块数据' in text


def test_export_preserves_user_content_without_reprojection_or_rewriting(tmp_path):
    template_path, template_dir = _template_paths()
    synopsis_node = 'v2_n_11_1_1'
    blocks = [
        _para(synopsis_node, '用户手改过的概要段落，含受试者一词与计划入组约240名受试者。'),
        _para(synopsis_node, '缩略语按用户自己的写法排列。'),
        _para('v2_n_12_x', '流程表章用户文字，含受试者与随访安排。'),
        _table('v2_n_12_x', '用户改过的访视安排'),
        _para('v2_n_13_x', '本候选不作合规断言。'),
    ]
    out = tmp_path / 'out.docx'
    render_production_docx(template_path, template_dir, _document(blocks),
                           out, {'research.input_context': {'user_brief': '合成'}})
    from docx import Document as open_docx
    doc = open_docx(str(out))
    text = _all_text(doc)
    # R4: the user's current version ships verbatim.
    assert '用户手改过的概要段落' in text
    assert '用户改过的访视安排' in text
    # No silent term replacement at export (R4/A13).
    assert text.count('受试者') == 3
    # No export-time synopsis projection: the projected abbreviation line
    # would differ from the user's own wording.
    assert '缩略语：AE = 不良事件' not in text
    # No marker-hit paragraph deletion: cleaning is a visible candidate, not
    # a silent export behavior (R-C04).
    assert '本候选不作合规断言' in text
    # Layout-only tuning still runs: wide tables stay readable.
    receipt = render_production_docx(template_path, template_dir,
                                     _document(blocks), tmp_path / 'out2.docx',
                                     {'research.input_context': {}})
    assert receipt['export_scope'] == 'production_docx'


def test_template_orphan_media_never_ships(tmp_path):
    """The clean template's flow-diagram image is referenced only by its
    example body; after body replacement it must not ship as dead weight
    (T07 spike finding), while real content parts stay."""
    import zipfile
    template_path, template_dir = _template_paths()
    with zipfile.ZipFile(template_path) as template_zip:
        template_media = [name for name in template_zip.namelist() if 'word/media/' in name]
    assert template_media, 'precondition: the template ships media parts'
    blocks = [_para('v2_n_11_1_1', '第一章正文。')]
    out = tmp_path / 'out.docx'
    receipt = render_production_docx(template_path, template_dir, _document(blocks),
                                     out, {'research.input_context': {}})
    with zipfile.ZipFile(out) as exported:
        exported_media = [name for name in exported.namelist() if 'word/media/' in name]
    assert exported_media == []
    assert [name.split('/')[-1] for name in receipt['orphaned_media_pruned']] == \
        [name.split('/')[-1] for name in template_media]
    # Header/footer/style parts are untouched by the pruning.
    with zipfile.ZipFile(out) as exported:
        kept = [name for name in exported.namelist()
                if name.startswith('word/header') or name.startswith('word/footer')
                or name == 'word/styles.xml']
    assert kept, 'headers, footers and styles must survive pruning'


def test_template_instruction_page_does_not_leave_a_blank_page_before_cover(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    template_path, template_dir = _template_paths()
    out = tmp_path / 'cover.docx'
    render_production_docx(template_path, template_dir,
        _document([_para('v2_n_11_1_1', '保留的研究正文。')]), out,
        {'framing.document_title': '用于核对封面的合成研究方案'})
    doc = Document(out)
    title_index = next(i for i, p in enumerate(doc.paragraphs)
        if p.text == '用于核对封面的合成研究方案')
    prefix = doc.paragraphs[:title_index]
    assert not any(b.get(qn('w:type')) == 'page'
        for p in prefix for b in p._p.findall('.//' + qn('w:br')))
    assert title_index <= 3, 'only the cover spacing may precede the title'
    assert '保留的研究正文。' in _all_text(doc)


def test_generated_table_of_contents_has_balanced_fields_and_no_template_page_cache(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    template_path, template_dir = _template_paths()
    out = tmp_path / 'toc.docx'
    render_production_docx(template_path, template_dir,
        _document([_para('v2_n_11_1_1','独立研究正文。')]),out,{})
    doc=Document(out)
    depth=0
    for field in doc.element.body.iter(qn('w:fldChar')):
        kind=field.get(qn('w:fldCharType'))
        if kind=='begin': depth+=1
        elif kind=='end': depth-=1
        assert depth>=0
    assert depth==0, 'template TOC must not be cut midway through an open field'
    assert '方案修订历史记录\t5' not in _all_text(doc)
    toc = next(p for p in doc.paragraphs if p._p.findall('.//' + qn('w:instrText')))
    assert '方案修订历史记录' in toc.text
    assert '主要研究者签字页' in toc.text
    assert '独立研究正文。' in _all_text(doc)
    assert len(doc.settings.element.findall(qn('w:updateFields'))) == 1
    assert doc.settings.element.find(qn('w:trackRevisions')) is None


def test_template_textbox_alternate_branches_are_not_concatenated(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    template_path, template_dir = _template_paths()
    out = tmp_path / 'textbox.docx'
    render_production_docx(template_path,template_dir,
        _document([_para('v2_n_11_1_1','用户正文仍包含所有版本都应当具有版本号和日期。')]),out,
        {'framing.sponsor':'合成申办者'})
    doc=Document(out)
    branches = list(doc.element.body.iter(qn('w:txbxContent')))
    template_branches = [branch for branch in Document(template_path).element.body.iter(qn('w:txbxContent'))
        if '本文件包含重要的保密性商业信息' in ''.join(n.text or '' for n in branch.iter(qn('w:t')))]
    assert len(branches) == len(template_branches) == 2
    texts = [''.join(n.text or '' for n in branch.iter(qn('w:t'))) for branch in branches]
    assert texts[0] == texts[1]
    assert all(text.count('本文件包含重要的保密性商业信息') == 1 for text in texts)
    assert all('合成申办者' in text for text in texts)
    assert '用户正文仍包含所有版本都应当具有版本号和日期。' in _all_text(doc)


def test_placeholder_replacement_preserves_runs_breaks_and_tabs():
    from docx import Document
    from app.protocol_workflow.agent3.word_export_production import _replace_in_paragraph
    doc = Document()
    p = doc.add_paragraph()
    p.add_run('前文<申办')
    p.add_run('者名称>').bold = True
    tail = p.add_run()
    tail.add_break()
    tail.add_text('后文')
    tail.add_tab()
    tail.add_text('尾部')
    _replace_in_paragraph(p, {'<申办者名称>': '合成单位'})
    assert p.text == '前文合成单位\n后文\t尾部'
    assert p.runs[1].bold is True
    assert p.runs[2].text == '\n后文\t尾部'


def test_template_header_uses_width_aware_tabs_and_retains_cover_section(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn
    template_path, template_dir = _template_paths()
    out = tmp_path / 'header.docx'
    render_production_docx(template_path, template_dir,
        _document([_para('v2_n_1_1', '合成正文。')]), out,
        {'framing.protocol_date': '2026-09-21', 'framing.protocol_id': 'TEST-20260921'})
    doc = Document(out)
    header = next(p for p in doc.sections[0].header.paragraphs if '版本日期：' in p.text)
    assert '\t' in header.text
    assert header.text.count('\t') == 2
    version = next(p for p in doc.sections[0].header.paragraphs if '版本号：' in p.text)
    assert version.text.count('\t') == 1
    assert '2026-09-21' in header.text
    from docx.enum.text import WD_TAB_ALIGNMENT
    active_tabs = [t for t in header.paragraph_format.tab_stops if t.alignment != WD_TAB_ALIGNMENT.CLEAR]
    assert [t.alignment for t in active_tabs] == [WD_TAB_ALIGNMENT.CENTER, WD_TAB_ALIGNMENT.RIGHT]
    cleared = {t.position for t in header.paragraph_format.tab_stops if t.alignment == WD_TAB_ALIGNMENT.CLEAR}
    assert {t.position for t in header.style.paragraph_format.tab_stops}.issubset(cleared)
    source = Document(template_path)
    assert doc.sections[0]._sectPr.xml == source.sections[0]._sectPr.xml
    assert len(list(doc.element.body.iter(qn('w:sectPr')))) >= 2


def test_front_matter_has_clean_version_sponsor_and_signature_text(tmp_path):
    from docx import Document
    from docx.oxml.ns import qn

    template_path, template_dir = _template_paths()
    out = tmp_path / 'front-matter.docx'
    render_production_docx(
        template_path,
        template_dir,
        _document([_para('v2_n_11_1_1', '试验参与者完成筛选后进入研究。')]),
        out,
        {
            'framing.document_title': '合成研究方案',
            'framing.version': 'v0.1',
            'framing.sponsor': '',
        },
    )
    doc = Document(out)
    text = _all_text(doc)
    all_xml_text = ''.join(node.text or '' for node in doc.element.body.iter(qn('w:t')))
    assert 'vV0.1' not in text
    assert 'v0.1' in text
    assert '(申办者名称)' not in text
    assert '()' not in text
    assert '此信息属申办者所有' in all_xml_text
    assert '[我已阅读' not in text
    assert '年     月     日]' not in text
