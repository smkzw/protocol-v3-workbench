from __future__ import annotations

import base64
import io
import zipfile

import pytest
from PIL import Image, ImageDraw
from docx import Document

from packages.contracts.workbench_contracts import ProtocolDocument, ProtocolSection

from services.api.app.medical_writing_document_exporter import (
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_instrument_appendix import (
    MedicalWritingInstrumentAppendixError,
    render_instrument_pdf_appendix,
    validate_instrument_appendix_block,
)


def _two_page_pdf() -> bytes:
    pages = []
    for page_number in (1, 2):
        image = Image.new("RGB", (595, 842), "white")
        draw = ImageDraw.Draw(image)
        draw.rectangle((32, 32, 563, 810), outline="black", width=2)
        draw.text((64, 64), f"Instrument page {page_number}", fill="black")
        pages.append(image)
    output = io.BytesIO()
    pages[0].save(
        output,
        format="PDF",
        save_all=True,
        append_images=pages[1:],
        resolution=72,
    )
    return output.getvalue()


def test_pdf_instrument_appendix_renders_each_page_above_200_dpi():
    rendered = render_instrument_pdf_appendix(
        _two_page_pdf(),
        instrument_id="instrument_test",
        title="测试量表",
        original_filename="test-scale.pdf",
        dpi=220,
    )
    blocks = rendered.content_blocks(body_order_start=20)

    assert len(blocks) == 2
    assert [block["page_number"] for block in blocks] == [1, 2]
    assert [block["body_order"] for block in blocks] == [20, 21]
    assert all(block["render_dpi"] == 220 for block in blocks)
    assert all(block["indexed"] is False for block in blocks)
    assert all(block["block_type"] == "appendix_image" for block in blocks)
    for block in blocks:
        payload = base64.b64decode(block["image_base64"])
        image = Image.open(io.BytesIO(payload))
        assert image.width >= 1800
        assert image.height >= 2500
        validate_instrument_appendix_block(block)


def test_pdf_instrument_appendix_rejects_low_dpi_and_non_pdf():
    with pytest.raises(MedicalWritingInstrumentAppendixError, match="DPI"):
        render_instrument_pdf_appendix(
            _two_page_pdf(),
            instrument_id="instrument_test",
            title="测试量表",
            original_filename="test-scale.pdf",
            dpi=150,
        )
    with pytest.raises(MedicalWritingInstrumentAppendixError, match="readable PDF"):
        render_instrument_pdf_appendix(
            b"not a PDF",
            instrument_id="instrument_test",
            title="测试量表",
            original_filename="test-scale.pdf",
        )


def test_pdf_instrument_appendix_hash_tampering_is_rejected():
    rendered = render_instrument_pdf_appendix(
        _two_page_pdf(),
        instrument_id="instrument_test",
        title="测试量表",
        original_filename="test-scale.pdf",
    )
    block = rendered.content_blocks(body_order_start=1)[0]
    block["image_sha256"] = "0" * 64
    with pytest.raises(MedicalWritingInstrumentAppendixError, match="hash"):
        validate_instrument_appendix_block(block)


def test_pdf_instrument_appendix_exports_as_unindexed_full_page_word_images():
    rendered = render_instrument_pdf_appendix(
        _two_page_pdf(),
        instrument_id="instrument_ibdq",
        title="炎症性肠病问卷",
        original_filename="IBDQ.pdf",
        dpi=220,
    )
    document = ProtocolDocument(
        document_id="doc_instrument_appendix",
        project_id="proj_instrument_appendix",
        protocol_id="CMS-INSTRUMENT",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_appendix",
                document_id="doc_instrument_appendix",
                heading="量表与评估工具",
                section_number="14.2",
                node_kind="assessment_instrument_appendix",
                content_blocks=[
                    {
                        "block_id": "appendix_heading",
                        "block_type": "heading",
                        "text": "14.2 量表与评估工具",
                        "body_order": 1,
                    },
                    *rendered.content_blocks(body_order_start=2),
                ],
            )
        ],
    )

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
    )

    assert result.metadata["figure_count"] == 0
    assert result.metadata["indexed_figure_count"] == 0
    assert result.metadata["appendix_image_page_count"] == 2
    assert [
        item["page_number"] for item in result.metadata["appendix_image_pages"]
    ] == [1, 2]
    exported = Document(io.BytesIO(result.content))
    drawing_paragraphs = [
        paragraph
        for paragraph in exported.paragraphs
        if paragraph._p.xpath(".//w:drawing")
    ]
    assert len(drawing_paragraphs) == 2
    assert len(exported.inline_shapes) == 2
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        media = [
            name
            for name in archive.namelist()
            if name.startswith("word/media/") and name.endswith(".png")
        ]
    assert len(media) == 2
    assert "炎症性肠病问卷原始附件第1页，共2页" in document_xml
    assert "炎症性肠病问卷原始附件第2页，共2页" in document_xml
    assert "SEQ 图" not in document_xml
    assert 'TOC \\h \\z \\c "图"' not in document_xml
