from __future__ import annotations

import io
import tempfile
import unittest
from copy import deepcopy
from pathlib import Path

from docx import Document

from packages.contracts.workbench_contracts import (
    MedicalWritingRevisionRequest,
    MedicalWritingWorkingCopySaveRequest,
    RevisionActionRequest,
)
from packages.contracts.workbench_contracts.models import (
    MedicalWritingSectionFreezeRequest,
)
from services.api.app.ai_execution_policy import AiExecutionPolicyResolver
from services.api.app.ai_gateway import AiPromptEnvelope
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore
from services.api.app.medical_writing import MedicalWritingRevisionService
from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_document_exporter import (
    document_index_catalog,
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_manifest import (
    D001_PROTOCOL_DOCX,
    MY008_PNH_3_01_PROTOCOL_DOCX,
    RUX_PROTOCOL_DOCX,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository
from services.api.app.medical_writing_table_templates import (
    MedicalWritingTableTemplateService,
)
from services.api.app.medical_writing_tables import MedicalWritingTableService
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore
from services.api.app.sqlite_runtime_store import SqliteRuntimeStore


class RealProjectWritingProvider:
    provider_name = "buddy"
    model_name = "glm-5.2"

    def __init__(self):
        self.envelopes = []

    def run(self, envelope: AiPromptEnvelope):
        self.envelopes.append(envelope)
        source = envelope.payload["allowed_sources"][0]
        source_id = source["source_id"]
        selection = source["text_preview"]
        return {
            "task_id": envelope.task_id,
            "task_type": envelope.task_type.value,
            "provider": self.provider_name,
            "model": self.model_name,
            "prompt_version": envelope.prompt_version,
            "input_source_ids": [source_id],
            "forbidden_source_ids": envelope.payload["forbidden_source_ids"],
            "findings": [
                {
                    "finding_id": f"finding_{len(self.envelopes):03d}",
                    "status": "supported",
                    "title": "方案正文修订候选",
                    "source_id": source_id,
                    "evidence_span_ids": [f"span_{len(self.envelopes):03d}"],
                }
            ],
            "evidence_spans": [
                {
                    "span_id": f"span_{len(self.envelopes):03d}",
                    "source_id": source_id,
                    "locator": source["locator"],
                    "quote": selection,
                }
            ],
            "uncertainties": [
                {"level": "medical_review", "description": "需医学经理确认修订与方案全局一致性。"}
            ],
            "needs_medical_confirmation": True,
            "schema_version": "ai_task_output_v0_1",
            "revision": {
                "proposal_text": f"{selection}（修订候选）",
                "diff_patch": f"- {selection}\n+ {selection}（修订候选）",
                "rationale": "在保留原始医学含义的前提下提高表述清晰度。",
                "evidence_span_ids": [f"span_{len(self.envelopes):03d}"],
                "alternatives": [
                    {
                        "proposal_text": f"{selection}（备选修订一）",
                        "diff_patch": f"- {selection}\n+ {selection}（备选修订一）",
                        "rationale": "提供一个保持原意的备选表述。",
                        "evidence_span_ids": [f"span_{len(self.envelopes):03d}"],
                    },
                    {
                        "proposal_text": f"{selection}（备选修订二）",
                        "diff_patch": f"- {selection}\n+ {selection}（备选修订二）",
                        "rationale": "提供另一个保持原意的备选表述。",
                        "evidence_span_ids": [f"span_{len(self.envelopes):03d}"],
                    },
                ],
            },
        }


@unittest.skipUnless(
    RUX_PROTOCOL_DOCX.exists() and D001_PROTOCOL_DOCX.exists() and MY008_PNH_3_01_PROTOCOL_DOCX.exists(),
    "real protocol DOCX fixtures are unavailable",
)
class MedicalWritingRealProjectFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        root = Path(self.tmpdir.name)
        self.document_service = MedicalWritingDocumentService()
        self.runtime_store = SqliteRuntimeStore(root / "runtime.sqlite3")
        self.repository = MedicalWritingRuntimeRepository(
            self.document_service,
            self.runtime_store,
        )
        self.provider = RealProjectWritingProvider()
        self.source_registry = SourceRegistryService(
            SourceRegistryStore(root / "source_registry.jsonl"),
            allowed_roots=[
                RUX_PROTOCOL_DOCX.parent,
                D001_PROTOCOL_DOCX.parent,
                MY008_PNH_3_01_PROTOCOL_DOCX.parent,
            ],
        )
        self.runner = AiTaskRunner(
            self.repository,
            AiTaskStore(root / "ai_runs.jsonl"),
            provider_factory=lambda resolution: self.provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name="buddy",
                model_name="glm-5.2",
                test_only_provider_injection=True,
            ),
        )
        self.service = MedicalWritingRevisionService(
            self.repository,
            self.runner,
            source_registry=self.source_registry,
            protocol_source_paths={
                "proj_rux_03_002": RUX_PROTOCOL_DOCX,
                "proj_d001": D001_PROTOCOL_DOCX,
                "proj_my008_pnh_3_01": MY008_PNH_3_01_PROTOCOL_DOCX,
            },
        )

    def tearDown(self):
        self.tmpdir.cleanup()

    def _selection(self, project_id: str):
        session = self.document_service.document_session(project_id)
        for summary in session.sections:
            section = self.document_service.section(project_id, summary.section_id)
            block = next(
                (
                    item
                    for item in section.content_blocks
                    if item.get("block_type") in {"paragraph", "heading"}
                    and str(item.get("text") or "").strip()
                ),
                None,
            )
            if block is not None:
                return summary, block
        self.fail(f"no editable paragraph found for {project_id}")

    def _submit(self, project_id: str):
        section, block = self._selection(project_id)
        return self.service.submit_revision(
            project_id,
            MedicalWritingRevisionRequest(
                document_id=section.document_id,
                section_id=section.section_id,
                anchor_path=block["source_locator"],
                selected_text=block["text"],
                user_instruction="请保持原始医学含义，优化为可供医学作者确认的方案正文候选。",
                requested_by="medical_manager_test",
            ),
        )

    def test_real_front_matter_and_synopsis_roles_survive_save_restart_and_word_export(self):
        cases = (
            ("proj_d001", "概要", "docx:table:4"),
            ("proj_my008_pnh_3_01", "方案摘要", "docx:table:13"),
        )
        table_service = MedicalWritingTableService()
        for project_id, heading, locator in cases:
            with self.subTest(project_id=project_id):
                session = self.document_service.document_session(project_id)
                section_summary = next(item for item in session.sections if item.heading == heading)
                source_section = self.document_service.section(project_id, section_summary.section_id)
                source_table = next(
                    block
                    for block in source_section.content_blocks
                    if block.get("source_locator") == locator
                )
                self.assertEqual(
                    "protocol_synopsis",
                    source_table["structured_table"]["role"],
                )
                initial = self.repository.working_copy(project_id, section_summary.section_id)
                first = self.repository.save_working_copy(
                    project_id,
                    section_summary.section_id,
                    MedicalWritingWorkingCopySaveRequest(
                        document_id=session.document_id,
                        expected_revision=0,
                        content_blocks=initial.content_blocks,
                        actor="medical_manager_object_role_test",
                        idempotency_key=f"save-{project_id}-synopsis-role-r1",
                    ),
                )
                edited_blocks = deepcopy(first.content_blocks)
                table_index = next(
                    index
                    for index, block in enumerate(edited_blocks)
                    if block.get("source_locator") == locator
                )
                table = table_service.from_table_block(edited_blocks[table_index])
                edited_cell = next(
                    cell
                    for row in table.rows
                    if row.order >= table.header_row_count
                    for cell in row.cells
                    if str(cell.text or "").strip()
                )
                table = table_service.apply_operations(
                    table,
                    [
                        {
                            "op": "edit_cell",
                            "cell_id": edited_cell.cell_id,
                            "text": f"{edited_cell.text}（对象角色重载测试）",
                        }
                    ],
                    expected_version=table.version,
                )
                edited_blocks[table_index] = table_service.to_table_block(table)
                saved = self.repository.save_working_copy(
                    project_id,
                    section_summary.section_id,
                    MedicalWritingWorkingCopySaveRequest(
                        document_id=session.document_id,
                        expected_revision=first.revision,
                        content_blocks=edited_blocks,
                        actor="medical_manager_object_role_test",
                        idempotency_key=f"save-{project_id}-synopsis-role-r2",
                    ),
                )
                restarted = MedicalWritingRuntimeRepository(
                    MedicalWritingDocumentService(),
                    SqliteRuntimeStore(self.runtime_store.db_path),
                )
                reloaded = restarted.working_copy(project_id, section_summary.section_id)
                reloaded_table_block = next(
                    block
                    for block in reloaded.content_blocks
                    if block.get("source_locator") == locator
                )
                reloaded_table = table_service.from_table_block(reloaded_table_block)
                self.assertEqual("protocol_synopsis", reloaded_table.role.value)
                self.assertTrue(
                    any(
                        "对象角色重载测试" in cell.text
                        for row in reloaded_table.rows
                        for cell in row.cells
                    )
                )
                assembled = restarted.assemble_document_for_export(project_id, "draft_preview")
                self.assertNotIn(
                    reloaded_table.table_id,
                    {item["object_id"] for item in document_index_catalog(assembled)["tables"]},
                )
                exported = export_medical_writing_document_docx(
                    assembled,
                    mode="draft_preview",
                )
                word = Document(io.BytesIO(exported.content))
                self.assertTrue(
                    any(
                        "对象角色重载测试" in cell.text
                        for word_table in word.tables
                        for row in word_table.rows
                        for cell in row.cells
                    )
                )
                self.assertEqual(2, saved.revision)

        rux = self.document_service.document_for_revision("proj_rux_03_002")
        rux_front = next(
            block
            for section in rux.sections
            for block in section.content_blocks
            if block.get("source_locator") == "docx:table:0"
        )
        self.assertEqual("layout", rux_front["structured_table"]["role"])

    def test_three_original_protocols_run_two_turn_independent_persistent_revision_workflows(self):
        submitted = {
            project_id: self._submit(project_id)
            for project_id in ("proj_rux_03_002", "proj_d001", "proj_my008_pnh_3_01")
        }
        rux_result = submitted["proj_rux_03_002"]
        d001_result = submitted["proj_d001"]
        pnh_result = submitted["proj_my008_pnh_3_01"]

        self.assertNotEqual(rux_result.thread.document_id, d001_result.thread.document_id)
        self.assertNotEqual(rux_result.thread.section_id, d001_result.thread.section_id)
        self.assertEqual(3, len({item.thread.document_id for item in submitted.values()}))
        self.assertTrue(all(item.thread.status == "candidate_ready" for item in submitted.values()))

        rewritten = {}
        for project_id, item in submitted.items():
            rewritten[project_id] = self.service.apply_action(
                project_id,
                item.thread.thread_id,
                RevisionActionRequest(
                    action="request_rewrite",
                    suggestion_id=item.suggestion.suggestion_id,
                    actor="medical_manager_test",
                    comment=f"{project_id}：保留原始医学事实，减少模板化表述。",
                    rewrite_instruction="请结合上一轮候选进一步压缩句式，不新增原始方案未支持的事实。",
                ),
            )
        self.assertEqual(6, len(self.provider.envelopes))
        for project_id, item in rewritten.items():
            first_round = [suggestion for suggestion in item.thread.suggestions if suggestion.turn_number == 1]
            second_round = [suggestion for suggestion in item.thread.suggestions if suggestion.turn_number == 2]
            first_selected = next(
                suggestion
                for suggestion in first_round
                if suggestion.suggestion_id == submitted[project_id].suggestion.suggestion_id
            )
            self.assertEqual(6, len(item.thread.suggestions), project_id)
            self.assertEqual(3, len(first_round), project_id)
            self.assertEqual(3, len(second_round), project_id)
            self.assertEqual(
                {submitted[project_id].suggestion.suggestion_id},
                {suggestion.parent_suggestion_id for suggestion in second_round},
            )
            self.assertEqual(
                "rewrite_requested",
                first_selected.user_decision,
            )
            self.assertTrue(
                all(
                    suggestion.user_decision == "not_selected"
                    for suggestion in first_round
                    if suggestion.suggestion_id
                    != submitted[project_id].suggestion.suggestion_id
                )
            )
            self.assertTrue(all(suggestion.user_decision == "pending" for suggestion in second_round))
            self.assertTrue(all(project_id in suggestion.user_comment for suggestion in second_round))

        accepted = self.service.apply_action(
            "proj_rux_03_002",
            rux_result.thread.thread_id,
            RevisionActionRequest(
                action="accept",
                suggestion_id=rewritten["proj_rux_03_002"].suggestion.suggestion_id,
                actor="medical_manager_test",
            ),
        )
        rejected = self.service.apply_action(
            "proj_d001",
            d001_result.thread.thread_id,
            RevisionActionRequest(
                action="reject",
                suggestion_id=rewritten["proj_d001"].suggestion.suggestion_id,
                actor="medical_manager_test",
                comment="当前版本不采用该候选。",
            ),
        )
        # Accepting a revision candidate is author selection, not medical ApprovalGate.
        self.assertEqual("author_selected", accepted.thread.status)
        self.assertEqual("rejected", rejected.thread.status)

        restarted_store = SqliteRuntimeStore(self.runtime_store.db_path)
        restarted_repository = MedicalWritingRuntimeRepository(
            MedicalWritingDocumentService(),
            restarted_store,
        )
        self.assertEqual(
            "author_selected",
            restarted_repository.revision_thread(
                "proj_rux_03_002", rux_result.thread.thread_id
            ).status,
        )
        self.assertEqual(
            "rejected",
            restarted_repository.revision_thread("proj_d001", d001_result.thread.thread_id).status,
        )
        pnh_restarted = restarted_repository.revision_thread(
            "proj_my008_pnh_3_01", pnh_result.thread.thread_id
        )
        self.assertEqual("candidate_ready", pnh_restarted.status)
        self.assertEqual(6, len(pnh_restarted.suggestions))
        self.assertEqual(2, pnh_restarted.suggestions[-1].turn_number)
        self.assertEqual(
            {pnh_result.suggestion.suggestion_id},
            {
                suggestion.parent_suggestion_id
                for suggestion in pnh_restarted.suggestions
                if suggestion.turn_number == 2
            },
        )
        self.assertTrue(pnh_restarted.suggestions[-1].ai_run_id)
        self.assertIsNotNone(pnh_restarted.suggestions[-1].created_at)
        self.assertEqual(0, len(restarted_store.gates("proj_rux_03_002")))
        self.assertEqual(0, len(restarted_store.gates("proj_d001")))
        self.assertEqual(0, len(restarted_store.gates("proj_my008_pnh_3_01")))
        self.assertEqual([], restarted_store.verify_audit_chain("proj_rux_03_002"))
        self.assertEqual([], restarted_store.verify_audit_chain("proj_d001"))
        self.assertEqual([], restarted_store.verify_audit_chain("proj_my008_pnh_3_01"))

    def test_unregistered_selected_text_is_denied_before_provider_call(self):
        section, block = self._selection("proj_rux_03_002")
        with self.assertRaisesRegex(ValueError, "does not match"):
            self.service.submit_revision(
                "proj_rux_03_002",
                MedicalWritingRevisionRequest(
                    document_id=section.document_id,
                    section_id=section.section_id,
                    anchor_path=block["source_locator"],
                    selected_text="这段文字并不存在于原始研究方案。",
                    user_instruction="请修订。",
                ),
            )
        self.assertEqual([], self.provider.envelopes)
        self.assertEqual([], self.repository.revision_threads("proj_rux_03_002"))

    def test_three_real_projects_keep_all_table_types_isolated_through_approval_and_word(self):
        templates = MedicalWritingTableTemplateService()
        table_service = MedicalWritingTableService()
        template_ids = (
            "schedule_of_activities",
            "objectives_endpoints",
            "sample_size_assumptions",
            "treatment_dose",
            "dose_modification",
            "stopping_rules",
            "ae_management",
            "laboratory_panel",
            "pk_immunogenicity_schedule",
            "analysis_sets",
            "version_history",
        )
        profiled_template_ids = {
            "objectives_endpoints",
            "sample_size_assumptions",
            "treatment_dose",
            "dose_modification",
            "laboratory_panel",
            "pk_immunogenicity_schedule",
            "analysis_sets",
            "version_history",
        }
        expected_header_labels = [
            next(item for item in templates.catalog() if item["template_id"] == template_id)[
                "columns"
            ][0]
            for template_id in template_ids
        ]
        for project_id in (
            "proj_rux_03_002",
            "proj_d001",
            "proj_my008_pnh_3_01",
        ):
            session = self.document_service.document_session(project_id)
            section_summary = session.sections[0]
            source_section = self.document_service.section(
                project_id,
                section_summary.section_id,
            )
            created = self.repository.save_working_copy(
                project_id,
                section_summary.section_id,
                MedicalWritingWorkingCopySaveRequest(
                    document_id=session.document_id,
                    expected_revision=0,
                    content_blocks=source_section.content_blocks,
                    actor="medical_manager_real_project_test",
                    idempotency_key=f"create-profile-copy-{project_id}",
                ),
            )
            generated = [
                templates.instantiate(
                    template_id,
                    instance_id=f"{project_id}-{template_id}",
                )
                for template_id in template_ids
            ]
            saved = self.repository.save_working_copy(
                project_id,
                section_summary.section_id,
                MedicalWritingWorkingCopySaveRequest(
                    document_id=session.document_id,
                    expected_revision=created.revision,
                    content_blocks=[*created.content_blocks, *generated],
                    actor="medical_manager_real_project_test",
                    idempotency_key=f"save-profile-tables-{project_id}",
                ),
            )

            candidate_blocks = deepcopy(saved.content_blocks)
            objectives_index = next(
                index
                for index, block in enumerate(candidate_blocks)
                if block.get("template_id") == "objectives_endpoints"
            )
            sibling_snapshots = {
                block["template_id"]: deepcopy(block["structured_table"])
                for block in candidate_blocks
                if block.get("source_kind") == "medical_writing_template"
                and block.get("template_id") != "objectives_endpoints"
            }
            objectives_block = candidate_blocks[objectives_index]
            objectives_table = table_service.from_table_block(objectives_block)
            objectives_table.word_layout["domain_profile"].update(
                {
                    "mapping_status": "confirmed_by_user",
                    "confirmed_by_user": True,
                }
            )
            objectives_table.rows[1].cells[1].text = f"{project_id} 主要研究目的"
            updated_objectives = table_service.to_table_block(objectives_table)
            updated_objectives.update(
                {
                    key: objectives_block[key]
                    for key in (
                        "source_kind",
                        "template_id",
                        "template_instance_id",
                        "editable",
                    )
                }
            )
            candidate_blocks[objectives_index] = updated_objectives
            edited = self.repository.save_working_copy(
                project_id,
                section_summary.section_id,
                MedicalWritingWorkingCopySaveRequest(
                    document_id=session.document_id,
                    expected_revision=saved.revision,
                    content_blocks=candidate_blocks,
                    actor="medical_manager_real_project_test",
                    idempotency_key=f"edit-one-profile-table-{project_id}",
                ),
            )
            restarted = MedicalWritingRuntimeRepository(
                MedicalWritingDocumentService(),
                SqliteRuntimeStore(self.runtime_store.db_path),
            ).working_copy(project_id, section_summary.section_id)
            self.assertEqual(
                created.content_blocks,
                restarted.content_blocks[: len(created.content_blocks)],
                project_id,
            )
            reloaded_generated = [
                block
                for block in restarted.content_blocks
                if block.get("source_kind") == "medical_writing_template"
            ]
            self.assertEqual(len(template_ids), len(reloaded_generated), project_id)
            self.assertEqual(
                list(template_ids),
                [block["template_id"] for block in reloaded_generated],
                project_id,
            )
            for block in reloaded_generated:
                table = table_service.from_table_block(block)
                if block["template_id"] not in profiled_template_ids:
                    self.assertNotIn("domain_profile", table.word_layout, block["template_id"])
                    self.assertTrue(
                        all(not column.semantic_role for column in table.columns),
                        block["template_id"],
                    )
                    continue
                self.assertTrue(all(column.semantic_role for column in table.columns))
                if block["template_id"] == "objectives_endpoints":
                    self.assertEqual(
                        "confirmed_by_user",
                        table.word_layout["domain_profile"]["mapping_status"],
                    )
                    self.assertEqual(f"{project_id} 主要研究目的", table.rows[1].cells[1].text)
                else:
                    self.assertEqual(
                        sibling_snapshots[block["template_id"]],
                        block["structured_table"],
                        block["template_id"],
                    )

            # Domain-profile isolation: only the edited objectives table is
            # author-confirmed; sibling profiled tables keep non-confirmed status.
            for block in reloaded_generated:
                if block["template_id"] not in profiled_template_ids:
                    continue
                table = table_service.from_table_block(block)
                status = table.word_layout.get("domain_profile", {}).get("mapping_status")
                if block["template_id"] == "objectives_endpoints":
                    self.assertEqual("confirmed_by_user", status, project_id)
                else:
                    self.assertNotEqual("confirmed_by_user", status, block["template_id"])

            # Section freeze is author version confirmation (not medical ApprovalGate).
            frozen = self.repository.freeze_current_version(
                project_id,
                section_summary.section_id,
                MedicalWritingSectionFreezeRequest(
                    document_id=session.document_id,
                    expected_working_copy_revision=restarted.revision,
                    reason="医学作者已核对本章结构化表格并确认冻结当前版本。",
                    actor="medical_manager_real_project_test",
                    idempotency_key=f"freeze-profile-tables-{project_id}",
                ),
            )
            self.assertEqual("frozen", frozen.working_copy.freeze_status, project_id)
            history = self.repository.section_freeze_history(
                project_id, section_summary.section_id
            )
            self.assertEqual(1, len(history), project_id)
            self.assertTrue(history[0].is_current, project_id)

            assembled = self.repository.assemble_document_for_export(
                project_id,
                "draft_preview",
            )
            exported = export_medical_writing_document_docx(
                assembled,
                mode="draft_preview",
            )
            word = Document(io.BytesIO(exported.content))
            self.assertGreaterEqual(len(word.tables), len(template_ids), project_id)
            self.assertEqual(
                expected_header_labels,
                [table.rows[0].cells[0].text for table in word.tables[-len(template_ids) :]],
                project_id,
            )
            self.assertNotIn("|---", "\n".join(paragraph.text for paragraph in word.paragraphs))
            self.assertEqual(edited.revision, restarted.revision)


if __name__ == "__main__":
    unittest.main()
