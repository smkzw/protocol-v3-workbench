from __future__ import annotations

import io

from docx import Document

from packages.contracts.workbench_contracts import ProtocolDocument, ProtocolSection
from services.api.app.medical_writing_document import (
    _blocks_for_range,
    _serialize_source_block,
)
from services.api.app.medical_writing_legacy_reference_index import (
    build_legacy_reference_index,
)
from services.api.app.protocol_text_extractor import parse_protocol_docx


def test_imported_table_cell_rich_text_preserves_superscript_and_plain_cells():
    source = Document()
    table = source.add_table(rows=1, cols=2)

    cited = table.cell(0, 0).paragraphs[0]
    cited.add_run("疗效结果")
    citation = cited.add_run("[1]")
    citation.font.superscript = True

    ordinary = table.cell(0, 1).paragraphs[0]
    ordinary.add_run("普通单元格")

    buffer = io.BytesIO()
    source.save(buffer)
    parsed = parse_protocol_docx("synthetic-rich-table.docx", buffer.getvalue())

    source_block = next(
        block
        for block in _blocks_for_range(parsed, 0, 1)
        if block.table is not None
    )
    payload = _serialize_source_block("mwsec_rich_table", source_block, 0)

    cited_cell = payload["rows"][0][0]
    ordinary_cell = payload["rows"][0][1]
    assert _plain_text(cited_cell["rich_text"]) == cited_cell["text"]
    assert _plain_text(ordinary_cell["rich_text"]) == ordinary_cell["text"]
    citation_node = cited_cell["rich_text"]["content"][1]
    assert citation_node["text"] == "[1]"
    assert {mark["type"] for mark in citation_node["marks"]} >= {"superscript"}
    assert ordinary_cell["rich_text"]["content"] == [
        {"type": "text", "text": "普通单元格"}
    ]

    document_id = "mwdoc_imported_table_rich_text"
    indexed = build_legacy_reference_index(
        ProtocolDocument(
            document_id=document_id,
            project_id="proj_imported_table_rich_text",
            protocol_id="RICH-TABLE-TEST",
            version="V1.0",
            sections=[
                ProtocolSection(
                    section_id="body",
                    document_id=document_id,
                    heading="研究设计",
                    content_blocks=[payload],
                ),
                ProtocolSection(
                    section_id="references",
                    document_id=document_id,
                    heading="参考文献",
                    content_blocks=[
                        {
                            "block_id": "reference_1",
                            "block_type": "paragraph",
                            "text": "[1] Imported table reference.",
                            "source_locator": "docx:paragraph:10",
                        }
                    ],
                ),
            ],
        )
    )
    assert len(indexed.occurrences) == 1
    assert indexed.occurrences[0].source_numbers == (1,)
    assert indexed.occurrences[0].source_locator == cited_cell["source_locator"]


def test_imported_multi_paragraph_table_cell_uses_document_rich_text_root():
    source = Document()
    cell = source.add_table(rows=1, cols=1).cell(0, 0)
    cell.paragraphs[0].add_run("第一段")
    second = cell.add_paragraph()
    second.add_run("第二段")

    buffer = io.BytesIO()
    source.save(buffer)
    parsed = parse_protocol_docx("synthetic-multi-paragraph-table.docx", buffer.getvalue())

    source_block = next(
        block
        for block in _blocks_for_range(parsed, 0, 1)
        if block.table is not None
    )
    payload = _serialize_source_block("mwsec_multi_paragraph", source_block, 0)
    serialized_cell = payload["rows"][0][0]

    assert serialized_cell["text"] == "第一段\n第二段"
    assert serialized_cell["rich_text"]["type"] == "doc"
    assert [
        _plain_text(paragraph)
        for paragraph in serialized_cell["rich_text"]["content"]
    ] == ["第一段", "第二段"]


def _plain_text(node: dict) -> str:
    if node.get("type") == "text":
        return str(node.get("text") or "")
    return "".join(
        _plain_text(child)
        for child in node.get("content", [])
        if isinstance(child, dict)
    )
