"""R8 片X（P0-26+P1-07）：Word导出渲染降级自愈 + 片C′清扫根修。

现场（R7-A loop27-r7-a，截图w7-10-导出渲染失败.png）：占位门知悉后
导出必报 `MedicalWritingDocumentDocxExportError: block mwblock_mwsec_
greenfield_proj_user_3084cb2bd954_… rich text does not match its text
projection`，三次尝试（含完全删除编辑）均同一块失败。

根因（本轮亲证，workbench_runtime.sqlite3 只读取证）：R7-A 存储块的
rich_text 是嵌套树 doc>paragraph>text，『医学审核注』在第二个嵌套段落；
第6轮片C′清扫只遍历 rich_text.content 一层（paragraph 节点无 text 键
未被触及）——导出副本上 text 被剥离而 rich_text 原样，投影不再相等，
导出器双表示不变量（exporter 两处 fail-closed raise）拒导出。存储行
本身投影一致（复算 mismatch=0），故用户编辑怎么改都不能解除。

契约（红先修后）：
- T1 根修：清扫递归改写 rich_text 全部 text 节点（任意深度），清扫后
  重算校验 text==投影；
- T2/T3 自愈：渲染路径投影不一致不再整单失败——剥离该块 rich_text
  按纯文本渲染，导出报告记降级注记（块ID+首20字）；
- T4 人话化：注记映射为『第X章「标题」第N段的富文本与正文不一致，
  已自动降级为纯文本导出』。
"""

from __future__ import annotations

import unittest

from packages.contracts.workbench_contracts import (
    ProtocolDocument,
    ProtocolSection,
)

PROSE = (
    "本研究为一项随机、双盲、安慰剂对照、两臂平行组、多中心III期确证性临床研究，"
    "旨在评价R7A-BD502口服治疗良性前列腺增生伴中重度下尿路症状的疗效与安全性。"
    "研究设两臂按每日一次口服给药，双盲核心期持续12周。"
)
ANNOTATION = "（医学审核注：本章与已确认PICOS事实一致；共同主要终点α分配建议保持同一口径。）"


def _nested_rich_block() -> dict:
    """R7-A 现场形态：doc>paragraph>text 嵌套树，批注在第二个段落。"""
    text = PROSE + "\n" + ANNOTATION
    return {
        "block_id": "mwblock_r8_degrade_test",
        "block_type": "paragraph",
        "text": text,
        "rich_text": {
            "type": "doc",
            "content": [
                {"type": "paragraph", "content": [{"type": "text", "text": PROSE}]},
                {"type": "paragraph", "content": [{"type": "text", "text": ANNOTATION}]},
            ],
        },
    }


def _document_with(block: dict) -> ProtocolDocument:
    return ProtocolDocument(
        document_id="doc_r8_degrade",
        project_id="proj_r8_degrade",
        protocol_id="CMS-R8-T",
        version="1.0",
        sections=[
            ProtocolSection(
                section_id="sec_overall",
                document_id="doc_r8_degrade",
                heading="总体设计",
                section_number="4.1",
                content_blocks=[block],
            )
        ],
    )


def _projection(rich) -> str:
    def walk(node):
        node_type = str(node.get("type") or "")
        if node_type == "text":
            return node.get("text")
        if node_type == "hardBreak":
            return "\n"
        parts = [walk(child) for child in node.get("content", [])]
        return "\n".join(parts) if node_type == "doc" else "".join(parts)

    return walk(rich)


class SanitizerNestedTreeTests(unittest.TestCase):
    def test_nested_annotation_stripped_in_both_representations(self) -> None:
        from services.api.app.main import _apply_export_annotation_sanitization

        block = _nested_rich_block()
        document = _document_with(block)
        notes = _apply_export_annotation_sanitization(document)
        # pydantic 会深拷贝 content_blocks——断言必须读文档内的副本。
        stored = document.sections[0].content_blocks[0]
        self.assertNotIn("医学审核注", stored["text"])
        self.assertNotIn("医学审核注", _projection(stored["rich_text"]))
        # 根修判据：清扫后双表示仍一致（导出器不变量不再被清扫破坏）。
        self.assertEqual(
            _projection(stored["rich_text"]),
            stored["text"],
            "清扫后 text 与 rich_text 投影必须一致（R7-A 死因）。",
        )
        self.assertTrue(notes)


class RenderDegradeSelfHealTests(unittest.TestCase):
    def test_render_paragraph_degrades_to_plain_text_with_note(self) -> None:
        from services.api.app.medical_writing_document_exporter import (
            reset_render_degradation_notes,
            take_render_degradation_notes,
            _render_paragraph,
        )
        from docx import Document

        block = _nested_rich_block()
        # 制造与清扫无关的投影不一致（任一写入路径只改其一的现场形态）。
        block["text"] = PROSE
        reset_render_degradation_notes()
        output = Document()
        # 旧契约在此 raise（P0-26 现场）；新契约降级渲染。
        count = _render_paragraph(output, block, "paragraph")
        self.assertGreaterEqual(count, 1)
        rendered = "\n".join(p.text for p in output.paragraphs)
        self.assertIn(PROSE[:20], rendered)
        notes = take_render_degradation_notes()
        self.assertEqual(1, len(notes))
        self.assertEqual(block["block_id"], notes[0]["block_id"])
        self.assertTrue(notes[0].get("preview"))

    def test_source_patch_paragraphs_degrades_instead_of_raising(self) -> None:
        from services.api.app.medical_writing_document_exporter import (
            reset_render_degradation_notes,
            take_render_degradation_notes,
            _render_source_patch_paragraphs,
        )

        block = _nested_rich_block()
        block["text"] = PROSE  # 投影不一致
        reset_render_degradation_notes()
        paragraphs = _render_source_patch_paragraphs(block)
        self.assertTrue(paragraphs)
        self.assertGreaterEqual(len(take_render_degradation_notes()), 1)

    def test_consistent_rich_text_still_renders_rich_and_notes_nothing(self) -> None:
        from services.api.app.medical_writing_document_exporter import (
            reset_render_degradation_notes,
            take_render_degradation_notes,
            _render_paragraph,
        )
        from docx import Document

        block = _nested_rich_block()  # text==投影（一致）
        reset_render_degradation_notes()
        output = Document()
        _render_paragraph(output, block, "paragraph")
        self.assertEqual(0, len(take_render_degradation_notes()))


class HumanizedDegradationTests(unittest.TestCase):
    def test_note_maps_to_section_ordinal_copy(self) -> None:
        from services.api.app.main import _humanize_render_degradations

        block = _nested_rich_block()
        document = _document_with(block)
        entries = _humanize_render_degradations(
            document,
            [{"block_id": block["block_id"], "preview": block["text"][:20]}],
        )
        self.assertEqual(1, len(entries))
        self.assertIn("4.1", entries[0])
        self.assertIn("总体设计", entries[0])
        self.assertIn("已自动降级为纯文本导出", entries[0])
        self.assertIn("富文本与正文不一致", entries[0])


if __name__ == "__main__":
    unittest.main()
