from __future__ import annotations

import io
import json
import sys
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import Mock, patch

import docx
import openpyxl
from fastapi.testclient import TestClient

ROOT = Path(__file__).resolve().parents[1]
PROJECT_ROOT = ROOT.parents[1]
sys.path.insert(0, str(ROOT))

from packages.contracts.workbench_contracts import AiTaskFromRegistryRequest  # noqa: E402
from services.api.app.ai_execution_policy import AiExecutionPolicyResolver  # noqa: E402
from services.api.app.ai_gateway import AiPromptEnvelope, DisabledAiProvider  # noqa: E402
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore  # noqa: E402
from services.api.app.demo_repository import DemoRepository  # noqa: E402
from services.api.app.main import app  # noqa: E402
from services.api.app import main as main_module  # noqa: E402
from services.api.app.monitoring_identity_authorization import MonitoringRole  # noqa: E402
from services.api.app.monitoring_runtime_principal import MonitoringAuthenticatedPrincipal  # noqa: E402
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore  # noqa: E402
from services.api.app.source_content_validation import (  # noqa: E402
    SourceContentValidationService,
    SourceContentValidationStore,
    SourceExpectedContext,
)


RUX_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/RUX-03-002_列表_数据集_Excel_20250612_处理后.xlsx"
)
RUX_PROTOCOL = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/RUX-03-002-自查文件包-20260107/10-临床试验重要文件/1-临床试验方案/V1.3版-2024.8.14/磷酸芦可替尼乳膏-AD3期临床研究方案V1.3-clean-20240814.docx"
)
D001_RAW_SUBJECT = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/全量-入组/白铭江/0316-SA07007-T-SPOT阳性导致筛败"
)


class EchoProvider:
    provider_name = "buddy"
    model_name = "deepseek-v4-pro"

    def run(self, envelope: AiPromptEnvelope):
        source_id = envelope.payload["allowed_sources"][0]["source_id"]
        return {
            "task_id": envelope.task_id,
            "task_type": envelope.task_type.value,
            "provider": self.provider_name,
            "model": self.model_name,
            "prompt_version": envelope.prompt_version,
            "input_source_ids": [source["source_id"] for source in envelope.payload["allowed_sources"]],
            "forbidden_source_ids": envelope.payload["forbidden_source_ids"],
            "findings": [
                {
                    "finding_id": "finding_001",
                    "status": "supported",
                    "title": "字段语义需要医学经理确认",
                    "source_id": source_id,
                    "evidence_span_ids": ["span_001"],
                }
            ],
            "evidence_spans": [
                {
                    "span_id": "span_001",
                    "source_id": source_id,
                    "locator": envelope.payload["allowed_sources"][0]["locator"],
                    "quote": envelope.payload["allowed_sources"][0]["text_preview"][:80],
                }
            ],
            "uncertainties": [{"level": "medical_review", "description": "示例 provider，仅用于边界测试。"}],
            "needs_medical_confirmation": True,
            "schema_version": "ai_task_output_v0_1",
        }


class SourceRegistryTests(unittest.TestCase):
    class _PrincipalMiddleware:
        def __init__(self, application, principal):
            self.application = application
            self.principal = principal

        async def __call__(self, scope, receive, send):
            if scope["type"] == "http":
                scope.setdefault("state", {})["monitoring_principal"] = self.principal
            await self.application(scope, receive, send)

    @staticmethod
    def _principal() -> MonitoringAuthenticatedPrincipal:
        return MonitoringAuthenticatedPrincipal(
            principal_id="source-registry-test",
            tenant_id="tenant-kangzhe",
            roles=(MonitoringRole.MEDICAL_MANAGER,),
            project_scope=("proj_mgk10_sar_demo",),
            issued_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
            expires_at=datetime(2099, 1, 1, tzinfo=timezone.utc),
            authenticated=True,
            authn_method="test-server-session",
            session_id="source-registry-test-session",
            directory_revision="source-registry-test-v1",
            verification_ref_sha256="4" * 64,
        )

    def _client(self):
        return TestClient(self._PrincipalMiddleware(app, self._principal()))

    def setUp(self):
        self.repo = DemoRepository(PROJECT_ROOT / "demo_data" / "workbench_demo_v0_1.json")
        self.ai_policy = AiExecutionPolicyResolver(
            deployment_profile="local_private_clinical",
            provider_name="buddy",
            model_name="deepseek-v4-pro",
        )
        self.disabled_ai_policy = AiExecutionPolicyResolver(
            deployment_profile="local_private_clinical",
            provider_name="disabled",
            model_name="not_configured",
        )

    def test_protocol_and_listing_registration_public_payload_hides_absolute_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))

            protocol_result = service.register_protocol_docx(
                "proj_mgk10_sar_demo",
                "/private/raw/CMS-D001 临床方案.docx",
                _minimal_docx_bytes(),
                module="eligibility_review",
            )
            listing_result = service.register_listing_file(
                "proj_mgk10_sar_demo",
                "/private/raw/edc_listing.xlsx",
                _minimal_xlsx_bytes(),
                module="medical_monitoring",
            )

            public_payload = json.dumps(
                {
                    "protocol": protocol_result.model_dump(mode="json"),
                    "listing": listing_result.model_dump(mode="json"),
                },
                ensure_ascii=False,
            )
            self.assertNotIn("/private/raw", public_payload)
            self.assertNotIn("server_path", public_payload)
            self.assertGreaterEqual(protocol_result.entry.span_count, 2)
            self.assertEqual(1, listing_result.entry.span_count)
            self.assertIn("SUBJID", listing_result.spans[0].text_preview)
            self.assertIn("sample_rows", listing_result.spans[0].text_preview)

    def test_raw_subject_bundle_registry_is_metadata_only_and_sanitized(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "0316-SA07007-T-SPOT阳性导致筛败"
            root.mkdir()
            (root / "合格邮件.pdf").write_bytes(b"%PDF-1.4 synthetic")
            (root / "检查照片.jpg").write_bytes(b"synthetic image")

            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                allowed_roots=[Path(tmp)],
            )
            result = service.register_raw_subject_bundle("proj_mgk10_sar_demo", root, module="eligibility_review")

            payload = json.dumps(result.model_dump(mode="json"), ensure_ascii=False)
            for leaked_label in ["合格", "筛败", "T-SPOT", str(root)]:
                self.assertNotIn(leaked_label, payload)
            self.assertEqual("inventory_only", result.entry.parser_status)
            self.assertIn("metadata_only_pending_ocr_vlm", result.entry.metadata["evidence_status"])
            self.assertTrue(any("needs_ocr_vlm=True" in span.text_preview for span in result.spans))

    def test_api_from_registered_sources_is_blocked_until_named_execution_action(self):
        with tempfile.TemporaryDirectory() as tmp:
            source_service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            registered = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "edc_listing.xlsx",
                _minimal_xlsx_bytes(),
                module="medical_monitoring",
            )
            runner = AiTaskRunner(
                self.repo,
                AiTaskStore(Path(tmp) / "ai_runs.jsonl"),
                provider_factory=lambda resolution: EchoProvider(),
                policy_resolver=self.ai_policy,
            )
            request = AiTaskFromRegistryRequest(
                module="medical_monitoring",
                task_type="listing_semantic_mapping",
                expected_prompt_version="listing_semantic_mapping_v0_1",
                source_ids=[registered.spans[0].source_id],
                forbidden_source_ids=["legacy_timeline_json"],
                user_instruction="仅基于服务端 registry 中的原始 listing sheet 建议字段映射。",
            )

            client = self._client()
            with patch("services.api.app.main.source_registry", source_service), patch(
                "services.api.app.main.ai_task_runner", runner
            ):
                response = client.post(
                    "/api/projects/proj_mgk10_sar_demo/ai-runs/from-sources",
                    json=request.model_dump(mode="json"),
                )

        self.assertEqual(403, response.status_code, response.text)
        self.assertEqual(
            "monitoring_write_action_unconfigured",
            response.json()["detail"]["code"],
        )
        self.assertEqual([], runner.list_runs("proj_mgk10_sar_demo"))

    def test_from_sources_accepts_evidence_design_data_analysis_tfl_safety_pv_task_types_when_provider_disabled(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            runner = AiTaskRunner(
                self.repo,
                AiTaskStore(Path(tmp) / "ai_runs.jsonl"),
                provider_factory=lambda resolution: DisabledAiProvider(),
                policy_resolver=self.disabled_ai_policy,
            )
            cases = [
                ("evidence_design", "competitive_intelligence"),
                ("data_analysis_tfl", "tfl_generation_assist"),
                ("safety_pv", "safety_case_medical_review"),
            ]
            for module, task_type in cases:
                with self.subTest(module=module, task_type=task_type):
                    result = service.register_listing_file(
                        "proj_mgk10_sar_demo",
                        f"{module}_{task_type}.xlsx",
                        _minimal_xlsx_bytes(),
                        module=module,
                    )
                    request = AiTaskFromRegistryRequest(
                        module=module,
                        task_type=task_type,
                        expected_prompt_version=f"{task_type}_v0_1",
                        source_ids=[result.spans[0].source_id],
                        forbidden_source_ids=["legacy_summary"],
                    )
                    run = runner.submit_registered("proj_mgk10_sar_demo", request, service)
                    self.assertEqual("blocked", run.status)
                    self.assertEqual(module, run.module)
                    self.assertEqual(task_type, run.task_type)
                    self.assertFalse(run.codex_runtime_dependency)

    def test_local_file_registration_uses_allowlist_and_hides_server_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_file = root / "edc_listing.xlsx"
            source_file.write_bytes(_minimal_xlsx_bytes())
            service = SourceRegistryService(
                SourceRegistryStore(root / "sources.jsonl"),
                allowed_roots=[root],
            )

            result = service.register_local_file("proj_mgk10_sar_demo", source_file, module="data_analysis_tfl")
            payload = json.dumps(result.model_dump(mode="json"), ensure_ascii=False)

            self.assertEqual("data_analysis_tfl", result.entry.module)
            self.assertEqual("listing_file", result.entry.source_kind)
            self.assertNotIn(str(source_file), payload)
            self.assertNotIn("server_path", payload)

    def test_operational_ai_rejects_superseded_source_entry(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                content_validation_service=SourceContentValidationService(
                    SourceContentValidationStore(Path(tmp) / "validations.sqlite3")
                ),
                expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                    expected_file_role="clinical_data_file",
                ),
            )
            first = service.register_listing_file(
                "proj_mgk10_sar_demo",
                "listing.xlsx",
                _minimal_xlsx_bytes("10001"),
                module="medical_monitoring",
                expected_file_role="clinical_data_file",
            )
            service.register_listing_file(
                "proj_mgk10_sar_demo",
                "listing.xlsx",
                _minimal_xlsx_bytes("10002"),
                module="medical_monitoring",
                expected_file_role="clinical_data_file",
            )

            with self.assertRaisesRegex(ValueError, "superseded registered source"):
                service.ai_task_request_from_registry(
                    "proj_mgk10_sar_demo",
                    AiTaskFromRegistryRequest(
                        module="medical_monitoring",
                        task_type="monitoring_risk_interpretation",
                        expected_prompt_version="monitoring_risk_interpretation_v0_1",
                        source_ids=[first.spans[0].source_id],
                    ),
                    self.ai_policy,
                )

    def test_operational_ai_rejects_stale_validator_decision(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                content_validation_service=SourceContentValidationService(
                    SourceContentValidationStore(Path(tmp) / "validations.sqlite3")
                ),
                expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                    expected_file_role="clinical_data_file",
                ),
            )
            result = service.register_listing_file(
                "proj_mgk10_sar_demo",
                "listing.xlsx",
                _minimal_xlsx_bytes(),
                module="safety_pv",
                expected_file_role="clinical_data_file",
            )

            with patch("services.api.app.source_intake.VALIDATOR_VERSION", "future_validator_v3"):
                with self.assertRaisesRegex(ValueError, "validation is stale"):
                    service.ai_task_request_from_registry(
                        "proj_mgk10_sar_demo",
                        AiTaskFromRegistryRequest(
                            module="safety_pv",
                            task_type="safety_case_medical_review",
                            expected_prompt_version="safety_case_medical_review_v0_1",
                            source_ids=[result.spans[0].source_id],
                        ),
                        self.ai_policy,
                    )

    def test_listing_preview_redacts_local_paths_and_public_api_bounds_preview(self):
        with tempfile.TemporaryDirectory() as tmp:
            source_service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            result = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "competitor_index.xlsx",
                _xlsx_with_local_path_cell(),
                module="evidence_design",
            )
            request = source_service.ai_task_request_from_registry(
                "proj_mgk10_sar_demo",
                AiTaskFromRegistryRequest(
                    module="evidence_design",
                    task_type="competitive_intelligence",
                    expected_prompt_version="competitive_intelligence_v0_1",
                    source_ids=[result.spans[0].source_id],
                ),
                self.ai_policy,
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                response = client.get("/api/projects/proj_mgk10_sar_demo/sources")

        combined_ai_payload = json.dumps(request.model_dump(mode="json"), ensure_ascii=False)
        self.assertNotIn("/Users/", combined_ai_payload)
        self.assertIn("[local_path_redacted]", combined_ai_payload)
        self.assertEqual(200, response.status_code)
        public_payload = response.json()
        self.assertEqual("proj_mgk10_sar_demo", public_payload["project_id"])
        self.assertNotIn("/Users/", json.dumps(public_payload, ensure_ascii=False))
        self.assertNotIn("content_hash", json.dumps(public_payload, ensure_ascii=False))
        self.assertNotIn("preview_hash", json.dumps(public_payload, ensure_ascii=False))
        self.assertLessEqual(len(public_payload["spans"][0]["text_preview"]), 614)

    def test_source_api_exposes_sanitized_validation_history_for_override_audit(self):
        with tempfile.TemporaryDirectory() as tmp:
            from services.api.app.source_content_validation import (
                SourceContentValidationService,
                SourceContentValidationStore,
                SourceExpectedContext,
            )
            from packages.contracts.workbench_contracts import SourceContentValidationConfirmationRequest

            validation_service = SourceContentValidationService(
                SourceContentValidationStore(Path(tmp) / "validations.sqlite3")
            )
            source_service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                content_validation_service=validation_service,
                expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                    project_identifiers=("STUDY-01",),
                    expected_file_role="edc_data_listing",
                ),
            )
            result = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "listing.xlsx",
                _minimal_xlsx_bytes(),
                module="medical_monitoring",
            )
            current = source_service.current_content_validation(
                "proj_mgk10_sar_demo", result.entry.entry_id
            )
            unresolved = [
                check.check_code
                for check in current.checks
                if check.outcome in {"warning", "mismatch"}
            ]
            source_service.confirm_content_validation(
                "proj_mgk10_sar_demo",
                result.entry.entry_id,
                SourceContentValidationConfirmationRequest(
                    reason="医学经理已对照当前项目资料，确认本次继续沿用。",
                    acknowledged_check_codes=unresolved,
                    expected_revision=current.revision,
                    idempotency_key="source-registry-history-001",
                ),
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                response = client.get("/api/projects/proj_mgk10_sar_demo/sources")

        self.assertEqual(200, response.status_code)
        payload = response.json()
        history = payload["content_validation_histories"][result.entry.entry_id]
        self.assertEqual([2, 1], [record["revision"] for record in history])
        self.assertEqual("confirmed_after_warning", history[0]["use_status"])
        self.assertTrue(history[0]["confirmation_reason"])
        serialized = json.dumps(payload, ensure_ascii=False)
        self.assertNotIn("file_sha256", serialized)
        self.assertNotIn("expected_context_hash", serialized)

    def test_source_api_marks_only_latest_stable_role_as_current(self):
        with tempfile.TemporaryDirectory() as tmp:
            source_service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            old = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "safety_listing.xlsx",
                _minimal_xlsx_bytes("10001"),
                module="safety_pv",
            )
            current = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "safety_listing.xlsx",
                _minimal_xlsx_bytes("10002"),
                module="safety_pv",
                expected_file_role="safety_medical_review_listing",
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                response = client.get("/api/projects/proj_mgk10_sar_demo/sources")

        by_id = {entry["entry_id"]: entry for entry in response.json()["entries"]}
        self.assertFalse(by_id[old.entry.entry_id]["is_current"])
        self.assertTrue(by_id[current.entry.entry_id]["is_current"])

    def test_source_api_replaces_numeric_document_title_with_safe_filename(self):
        with tempfile.TemporaryDirectory() as tmp:
            source_service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            result = source_service.register_protocol_docx(
                "proj_mgk10_sar_demo",
                "MY009_Document.docx",
                _minimal_docx_bytes(),
                module="safety_pv",
            )
            stored = result.entry.model_copy(update={"public_title": "18"})
            source_service.store.jsonl_path.write_text(
                json.dumps(result.model_copy(update={"entry": stored}).model_dump(mode="json"), ensure_ascii=False) + "\n",
                encoding="utf-8",
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                response = client.get("/api/projects/proj_mgk10_sar_demo/sources")

        self.assertEqual("MY009_Document.docx", response.json()["entries"][0]["public_title"])

    def test_public_source_api_hides_registry_hash_fields_and_raw_hash_prefixes(self):
        with tempfile.TemporaryDirectory() as tmp:
            source_service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))
            internal = source_service.register_listing_file(
                "proj_mgk10_sar_demo",
                "edc_listing.xlsx",
                _minimal_xlsx_bytes(),
                module="medical_monitoring",
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                response = client.get("/api/projects/proj_mgk10_sar_demo/sources")

        self.assertEqual(200, response.status_code)
        serialized = response.text
        self.assertNotIn("content_hash", serialized)
        self.assertNotIn("preview_hash", serialized)
        self.assertNotIn(internal.entry.content_hash, serialized)
        self.assertNotIn(internal.entry.content_hash[:12], serialized)
        self.assertNotIn(internal.spans[0].preview_hash, serialized)

    def test_public_source_registration_routes_fail_closed_before_registry_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            raw = root / "0316-SA07005"
            raw.mkdir()
            (raw / "合格邮件.pdf").write_bytes(b"%PDF-1.4 synthetic")
            source_service = SourceRegistryService(
                SourceRegistryStore(root / "sources.jsonl"),
                allowed_roots=[root],
            )
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service):
                listing_response = client.post(
                    "/api/projects/proj_mgk10_sar_demo/sources/listing-file",
                    params={"filename": "edc_listing.xlsx", "module": "medical_monitoring"},
                    content=_minimal_xlsx_bytes(),
                )
                raw_response = client.post(
                    "/api/projects/proj_mgk10_sar_demo/sources/raw-subject-bundle",
                    params={"root_path": str(raw), "module": "eligibility_review"},
                )

        for response in (listing_response, raw_response):
            self.assertEqual(403, response.status_code, response.text)
            self.assertEqual(
                "monitoring_write_action_unconfigured",
                response.json()["detail"]["code"],
            )
        self.assertFalse((root / "sources.jsonl").exists())

    def test_local_candidate_route_uses_server_side_mapping_and_hides_internal_path(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_file = root / "candidate_listing.xlsx"
            source_file.write_bytes(_minimal_xlsx_bytes())
            source_service = SourceRegistryService(
                SourceRegistryStore(root / "sources.jsonl"),
                allowed_roots=[root],
            )
            candidate_map = {
                "unit-candidate-listing": {
                    "kind": "local-file",
                    "module": "data_analysis_tfl",
                    "project_ids": {"proj_mgk10_sar_demo"},
                    "path": source_file,
                }
            }
            client = self._client()
            with patch("services.api.app.main.source_registry", source_service), patch.dict(
                "services.api.app.main.SOURCE_REGISTRY_CANDIDATES",
                candidate_map,
                clear=True,
            ):
                success = client.post(
                    "/api/projects/proj_mgk10_sar_demo/sources/local-candidate",
                    params={"candidate_id": "unit-candidate-listing", "module": "data_analysis_tfl"},
                )
                wrong_module = client.post(
                    "/api/projects/proj_mgk10_sar_demo/sources/local-candidate",
                    params={"candidate_id": "unit-candidate-listing", "module": "safety_pv"},
                )
                unknown = client.post(
                    "/api/projects/proj_mgk10_sar_demo/sources/local-candidate",
                    params={"candidate_id": "missing-candidate", "module": "data_analysis_tfl"},
                )

        for response in (success, wrong_module, unknown):
            self.assertEqual(403, response.status_code, response.text)
            self.assertEqual(
                "monitoring_write_action_unconfigured",
                response.json()["detail"]["code"],
            )
        self.assertFalse((root / "sources.jsonl").exists())

    def test_local_file_candidate_registration_reuses_unchanged_parse_and_invalidates_on_change(self):
        from types import SimpleNamespace

        with tempfile.TemporaryDirectory() as tmp:
            source_file = Path(tmp) / "candidate.xlsx"
            source_file.write_bytes(_minimal_xlsx_bytes())
            candidate = {
                "kind": "local-file",
                "module": "safety_pv",
                "project_ids": {"proj_test"},
                "path": source_file,
                "source_kind": "safety_analysis_listing",
            }
            fake_registry = Mock()
            fake_registry.reusable_local_file_registration.return_value = None
            fake_registry.register_local_file.side_effect = [
                SimpleNamespace(marker="first"),
                SimpleNamespace(marker="changed"),
            ]
            cache_id = ("proj_test", "safety_pv", "test-listing")
            main_module._SOURCE_CANDIDATE_REGISTRATION_CACHE.pop(cache_id, None)
            with patch.object(main_module, "source_registry", fake_registry), patch.dict(
                main_module.SOURCE_REGISTRY_CANDIDATES,
                {"test-listing": candidate},
                clear=True,
            ):
                first = main_module._register_source_candidate("proj_test", "safety_pv", "test-listing")
                replay = main_module._register_source_candidate("proj_test", "safety_pv", "test-listing")
                source_file.write_bytes(_minimal_xlsx_bytes() + b"changed")
                changed = main_module._register_source_candidate("proj_test", "safety_pv", "test-listing")

            self.assertIs(first, replay)
            self.assertNotEqual(first.marker, changed.marker)
            self.assertEqual(2, fake_registry.register_local_file.call_count)
            main_module._SOURCE_CANDIDATE_REGISTRATION_CACHE.pop(cache_id, None)

    def test_registry_recovers_persisted_local_file_registration_without_reparsing(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source_file = root / "safety.xlsx"
            source_file.write_bytes(_minimal_xlsx_bytes())
            service = SourceRegistryService(
                SourceRegistryStore(root / "sources.jsonl"),
                allowed_roots=[root],
            )
            registered = service.register_local_file(
                "proj_test",
                source_file,
                module="safety_pv",
                expected_file_role="safety_analysis_listing",
            )

            reusable = service.reusable_local_file_registration(
                "proj_test",
                source_file,
                module="safety_pv",
                expected_file_role="safety_analysis_listing",
            )
            source_file.write_bytes(source_file.read_bytes() + b"changed")
            changed = service.reusable_local_file_registration(
                "proj_test",
                source_file,
                module="safety_pv",
                expected_file_role="safety_analysis_listing",
            )

            self.assertIsNotNone(reusable)
            self.assertEqual(registered.entry.entry_id, reusable.entry.entry_id)
            self.assertIsNone(changed)

    def test_frontend_source_registry_candidates_do_not_embed_local_paths(self):
        frontend_src = ROOT / "frontend" / "src"
        combined = "\n".join(
            path.read_text(encoding="utf-8", errors="replace")
            for path in sorted(frontend_src.rglob("*"))
            if path.is_file()
        )

        self.assertNotIn("/Users/", combined)
        app_source = (frontend_src / "App.jsx").read_text(encoding="utf-8")
        self.assertNotIn("candidate.path", app_source)
        self.assertNotIn('params.set("file_path"', app_source)
        self.assertNotIn('params.set("root_path"', app_source)
        self.assertIn("/sources/local-candidate", app_source)

    def test_local_file_registration_rejects_outside_allowlist(self):
        with tempfile.TemporaryDirectory() as tmp, tempfile.TemporaryDirectory() as outside:
            source_file = Path(outside) / "edc_listing.xlsx"
            source_file.write_bytes(_minimal_xlsx_bytes())
            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                allowed_roots=[Path(tmp)],
            )

            with self.assertRaises(ValueError):
                service.register_local_file("proj_mgk10_sar_demo", source_file, module="data_analysis_tfl")

    def test_local_directory_registration_creates_generic_inventory_without_paths(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "tfl_package"
            root.mkdir()
            (root / "ae.xpt").write_bytes(b"synthetic xpt")
            (root / "define.xml").write_text("<ODM />", encoding="utf-8")
            (root / "csr.pdf").write_bytes(b"%PDF-1.4 synthetic")
            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                allowed_roots=[Path(tmp)],
            )

            result = service.register_local_directory(
                "proj_mgk10_sar_demo",
                root,
                module="data_analysis_tfl",
                source_kind="tfl_dataset_package_inventory",
            )
            payload = json.dumps(result.model_dump(mode="json"), ensure_ascii=False)

            self.assertEqual("tfl_dataset_package_inventory", result.entry.source_kind)
            self.assertEqual("inventory_only", result.entry.parser_status)
            self.assertEqual(3, result.entry.span_count)
            self.assertIn("metadata_only_pending_parser", result.entry.metadata["evidence_status"])
            self.assertNotIn(str(root), payload)
            self.assertNotIn("ae.xpt", payload)
            self.assertTrue(any(span.source_type == "file_bundle_pdf" for span in result.spans))

    @unittest.skipUnless(RUX_LISTING.exists() and RUX_PROTOCOL.exists(), "RUX-03-002 raw files not available")
    def test_real_rux_listing_and_protocol_register_from_original_files(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(SourceRegistryStore(Path(tmp) / "sources.jsonl"))

            listing_result = service.register_listing_file(
                "proj_mgk10_sar_demo",
                RUX_LISTING.name,
                RUX_LISTING.read_bytes(),
                module="medical_monitoring",
            )
            protocol_result = service.register_protocol_docx(
                "proj_mgk10_sar_demo",
                RUX_PROTOCOL.name,
                RUX_PROTOCOL.read_bytes(),
                module="medical_monitoring",
            )
            request = service.ai_task_request_from_registry(
                "proj_mgk10_sar_demo",
                AiTaskFromRegistryRequest(
                    module="medical_monitoring",
                    task_type="monitoring_risk_interpretation",
                    expected_prompt_version="monitoring_risk_interpretation_v0_1",
                    source_ids=[listing_result.spans[0].source_id, protocol_result.spans[0].source_id],
                    forbidden_source_ids=["existing_ae_mh_report", "existing_patient_profile_html"],
                ),
                self.ai_policy,
            )

        self.assertGreaterEqual(listing_result.entry.span_count, 10)
        self.assertGreaterEqual(protocol_result.entry.span_count, 20)
        self.assertEqual(2, len(request.allowed_sources))
        self.assertIn("existing_ae_mh_report", request.forbidden_source_ids)

    @unittest.skipUnless(D001_RAW_SUBJECT.exists(), "D001 raw subject folder not available")
    def test_real_d001_raw_subject_bundle_does_not_leak_conclusion_labels(self):
        with tempfile.TemporaryDirectory() as tmp:
            service = SourceRegistryService(
                SourceRegistryStore(Path(tmp) / "sources.jsonl"),
                allowed_roots=[D001_RAW_SUBJECT.parents[2]],
            )
            result = service.register_raw_subject_bundle("proj_mgk10_sar_demo", D001_RAW_SUBJECT, module="eligibility_review")

        combined = json.dumps(result.model_dump(mode="json"), ensure_ascii=False)
        for leaked_label in ["筛败", "T-SPOT", str(D001_RAW_SUBJECT)]:
            self.assertNotIn(leaked_label, combined)
        self.assertGreater(result.entry.span_count, 0)


def _minimal_docx_bytes() -> bytes:
    document = docx.Document()
    document.add_paragraph("CMS-D001 研究方案")
    document.add_paragraph("IN-01 受试者需签署知情同意。")
    document.add_paragraph("EX-01 活动性感染者不得入组。")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _minimal_xlsx_bytes(subject_id: str = "10001") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AE"
    sheet.append(["受试者编号", "中心编号", "访视名称", "AE术语"])
    sheet.append([subject_id, "01", "W1", "头痛"])
    sheet.append(["10002", "02", "W2", "瘙痒"])
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _xlsx_with_local_path_cell() -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Evidence"
    sheet.append(["trial_id", "source_path", "note"])
    sheet.append(
        [
            "TRIAL-001",
            "/Users/smkzw/Documents/康哲项目资料/竞品调研/CRSwNP/03_Protocols/source.pdf",
            "long evidence note " * 200,
        ]
    )
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


if __name__ == "__main__":
    unittest.main()
