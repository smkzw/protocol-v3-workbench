from __future__ import annotations

import base64
import hashlib
import io
import re
import shutil
import subprocess
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from packages.contracts.workbench_contracts import ProtocolDocument, ProtocolSection
from services.api.app import medical_writing_document_exporter as exporter
from services.api.app.medical_writing_document_exporter import (
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_instrument_appendix import (
    render_instrument_pdf_appendix,
)
from services.api.app.medical_writing_study_schema import (
    render_study_schema_png,
    render_study_schema_svg,
)
from services.api.app.medical_writing_style_profile import (
    MedicalWritingStyleProfileService,
)
from scripts.render_docx_visual_qc import render_docx
from tests.test_medical_writing_instrument_appendix import _two_page_pdf
from tests.test_medical_writing_study_schema import _pnh_schema


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WORD_NS}


def _appendix_document() -> tuple[ProtocolDocument, list[dict[str, object]]]:
    rendered = render_instrument_pdf_appendix(
        _two_page_pdf(),
        instrument_id="instrument_pagination",
        title="测试量表",
        original_filename="pagination-scale.pdf",
        dpi=220,
    )
    pages = rendered.content_blocks(body_order_start=2)
    document = ProtocolDocument(
        document_id="doc_appendix_pagination",
        project_id="proj_appendix_pagination",
        protocol_id="CMS-APPENDIX-PAGINATION",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_appendix",
                document_id="doc_appendix_pagination",
                heading="量表与评估工具",
                section_number="14.2",
                node_kind="assessment_instrument_appendix",
                content_blocks=[
                    {
                        "block_id": "appendix_heading",
                        "block_type": "heading",
                        "text": "14.2 量表与评估工具",
                        "body_order": 1,
                    },
                    *pages,
                ],
            )
        ],
    )
    return document, pages


def _pdf_with_page_sizes(page_sizes: list[tuple[int, int]]) -> bytes:
    pages: list[Image.Image] = []
    for page_number, (width, height) in enumerate(page_sizes, start=1):
        image = Image.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(image)
        inset = max(12, min(width, height) // 24)
        draw.rectangle(
            (inset, inset, width - inset, height - inset),
            outline="black",
            width=max(2, inset // 8),
        )
        draw.rectangle(
            (inset * 2, inset * 3, width - inset * 2, inset * 5),
            fill=(222, 236, 226),
            outline=(0, 92, 73),
            width=max(2, inset // 10),
        )
        draw.text(
            (inset * 2, inset * 3 + 4),
            f"Source appendix page {page_number}",
            fill="black",
        )
        pages.append(image)
    output = io.BytesIO()
    pages[0].save(
        output,
        format="PDF",
        save_all=True,
        append_images=pages[1:],
        resolution=72,
    )
    return output.getvalue()


def _appendix_document_from_pdf(
    pdf_bytes: bytes,
    *,
    case_name: str,
) -> tuple[ProtocolDocument, list[dict[str, object]]]:
    rendered = render_instrument_pdf_appendix(
        pdf_bytes,
        instrument_id=f"instrument_{case_name}",
        title=f"{case_name}测试量表",
        original_filename=f"{case_name}.pdf",
        dpi=220,
    )
    pages = rendered.content_blocks(body_order_start=2)
    document = ProtocolDocument(
        document_id=f"doc_{case_name}",
        project_id=f"proj_{case_name}",
        protocol_id=f"CMS-{case_name.upper()}",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id=f"section_{case_name}",
                document_id=f"doc_{case_name}",
                heading="量表与评估工具",
                section_number="14.2",
                node_kind="assessment_instrument_appendix",
                content_blocks=[
                    {
                        "block_id": f"heading_{case_name}",
                        "block_type": "heading",
                        "text": "14.2 量表与评估工具",
                        "body_order": 1,
                    },
                    *pages,
                ],
            ),
            ProtocolSection(
                section_id=f"section_{case_name}_after",
                document_id=f"doc_{case_name}",
                heading="附录后正文",
                section_number="14.3",
                content_blocks=[
                    {
                        "block_id": f"after_{case_name}",
                        "block_type": "heading",
                        "text": "14.3 附录后正文",
                        "body_order": len(pages) + 2,
                    }
                ],
            ),
        ],
    )
    return document, pages


def _render_docx_with_libreoffice(
    tmp_path: Path,
    *,
    case_name: str,
    content: bytes,
) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is not available")
    case_dir = tmp_path / case_name
    output_dir = case_dir / "pdf"
    output_dir.mkdir(parents=True)
    docx_path = case_dir / f"{case_name}.docx"
    pdf_path = output_dir / f"{case_name}.pdf"
    assert not pdf_path.exists()
    docx_path.write_bytes(content)
    docx_written_at = time.time_ns()
    report = render_docx(
        docx_path,
        output_dir,
        soffice=Path(soffice),
        max_preview_pages=1,
    )
    assert report["passed"] is True
    assert report["expected_text"] == ""
    assert pdf_path.is_file()
    assert pdf_path.stat().st_size > 0
    assert pdf_path.stat().st_mtime_ns >= docx_written_at
    return pdf_path


def _assert_rendered_appendix_page_assignment(
    pdf_path: Path,
    *,
    page_count: int,
    source_aspect_ratio: float,
    expect_landscape: bool,
) -> None:
    fitz = pytest.importorskip("fitz")
    rendered = fitz.open(pdf_path)
    appendix_pages: list[int] = []
    image_pages: list[int] = []
    following_content_pages: list[int] = []
    try:
        for page_index, page in enumerate(rendered):
            normalized_text = re.sub(r"\s+", "", page.get_text("text"))
            if "14.3附录后正文" in normalized_text:
                following_content_pages.append(page_index)
            title_numbers = [
                page_number
                for page_number in range(1, page_count + 1)
                if f"原始附件第{page_number}/{page_count}页" in normalized_text
            ]
            images = page.get_image_info(xrefs=True)
            if title_numbers:
                assert len(title_numbers) == 1
                appendix_pages.append(page_index)
                assert len(images) == 1
                image_pages.append(page_index)
                image_box = fitz.Rect(images[0]["bbox"])
                title_boxes = page.search_for("原始附件")
                assert title_boxes
                assert image_box.y0 > max(box.y1 for box in title_boxes)
                assert image_box.x0 >= 0 and image_box.y0 >= 0
                assert image_box.x1 <= page.rect.x1 and image_box.y1 <= page.rect.y1
                rendered_ratio = image_box.width / image_box.height
                assert rendered_ratio == pytest.approx(source_aspect_ratio, rel=0.025)
                pixmap = page.get_pixmap(
                    matrix=fitz.Matrix(1.5, 1.5),
                    clip=image_box,
                    colorspace=fitz.csGRAY,
                    alpha=False,
                )
                assert pixmap.samples
                assert sum(value < 235 for value in pixmap.samples) / len(
                    pixmap.samples
                ) > 0.002
                assert (page.rect.width > page.rect.height) is expect_landscape
            elif images:
                image_pages.append(page_index)

        assert len(appendix_pages) == page_count
        assert image_pages == appendix_pages
        assert appendix_pages == list(
            range(appendix_pages[0], appendix_pages[0] + page_count)
        )
        assert following_content_pages
        assert min(following_content_pages) > appendix_pages[-1]
        for page_number in range(1, page_count + 1):
            token = f"原始附件第{page_number}/{page_count}页"
            assert sum(
                token in re.sub(r"\s+", "", page.get_text("text"))
                for page in rendered
            ) == 1
    finally:
        rendered.close()


def _wide_study_schema_figure(*, body_order: int) -> dict[str, object]:
    schema = _pnh_schema()
    svg = render_study_schema_svg(schema)
    wide_svg = re.sub(r'(<svg\b[^>]*\bwidth=")[0-9.]+"', r'\g<1>1972"', svg, count=1)
    png = render_study_schema_png(wide_svg)
    return {
        "block_id": "wide_study_schema",
        "block_type": "figure",
        "figure_id": "mwfigure_wide_pagination",
        "figure_kind": "study_schema",
        "title": "研究设计概况",
        "alt_text": "研究流程图：研究设计概况",
        "body_order": body_order,
        "source_kind": "medical_writing_study_schema",
        "source_locator": "generated:medical_writing_study_schema:pagination:r1",
        "editable": False,
        "schema_id": schema.schema_id,
        "schema_revision": schema.revision,
        "schema_state_sha256": schema.state_sha256,
        "layout_revision": 0,
        "svg": wide_svg,
        "svg_sha256": hashlib.sha256(wide_svg.encode("utf-8")).hexdigest(),
        "png_base64": base64.b64encode(png).decode("ascii"),
        "png_sha256": hashlib.sha256(png).hexdigest(),
        "width_inches": 6.45,
    }


def _paragraph_section_indices(content: bytes) -> dict[str, int]:
    root = ElementTree.fromstring(content)
    body = root if root.tag == qn("w:body") else root.find("w:body", NS)
    assert body is not None
    section_index = 0
    result: dict[str, int] = {}
    for child in body:
        if child.tag != qn("w:p"):
            continue
        text = "".join(node.text or "" for node in child.findall(".//w:t", NS))
        if text:
            result[text] = section_index
        if child.find(".//w:drawing", NS) is not None:
            result.setdefault("<drawing>", section_index)
        if child.find("w:pPr/w:sectPr", NS) is not None:
            section_index += 1
    return result


def test_appendix_pages_use_title_page_breaks_without_separator_paragraphs():
    document, pages = _appendix_document()
    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))

    heading = next(paragraph for paragraph in exported.paragraphs if paragraph.text == "14.2 量表与评估工具")
    page_titles = [
        paragraph
        for paragraph in exported.paragraphs
        if "原始附件第" in paragraph.text
    ]
    drawing_indices = [
        index
        for index, paragraph in enumerate(exported.paragraphs)
        if paragraph._p.xpath(".//w:drawing")
    ]
    assert len(page_titles) == len(drawing_indices) == 2
    assert heading.paragraph_format.keep_with_next is True
    assert page_titles[0].paragraph_format.page_break_before is not True
    assert page_titles[0].paragraph_format.keep_with_next is True
    assert page_titles[1].paragraph_format.page_break_before is not True

    first_drawing_index = drawing_indices[0]
    second_title_index = next(
        index
        for index, paragraph in enumerate(exported.paragraphs)
        if paragraph.text == page_titles[1].text
    )
    assert second_title_index == first_drawing_index + 1
    assert not page_titles[1]._p.xpath('.//w:br[@w:type="page"]')
    assert exported.paragraphs[first_drawing_index]._p.xpath(
        './/w:br[@w:type="page"]'
    )
    assert not any(
        not paragraph.text
        and paragraph._p.xpath('.//w:br[@w:type="page"]')
        and not paragraph._p.xpath(".//w:drawing")
        for paragraph in exported.paragraphs[first_drawing_index:second_title_index]
    )

    ordered_blocks = exporter._ordered_blocks(document)
    layout_plan = exporter._build_visual_layout_plan(
        ordered_blocks,
        exporter._BODY_LAYOUT,
    )
    appendix_layouts = [
        layout_plan.visual_blocks_by_index[index]
        for index, (_, _, block) in enumerate(ordered_blocks)
        if block.get("block_type") == "appendix_image"
    ]
    assert appendix_layouts[0].context_indices
    assert (
        appendix_layouts[0].effective_width_inches
        < appendix_layouts[1].effective_width_inches
    )

    _, width_inches, _ = exporter._instrument_appendix_page_word_layout(
        pages[1], exporter._BODY_LAYOUT
    )
    _, (page_width_twips, page_height_twips), margins = exporter._BODY_LAYOUT
    del page_width_twips
    available_height = (page_height_twips - margins[2] - margins[3]) / 1440
    image_height = width_inches * pages[1]["pixel_height"] / pages[1]["pixel_width"]
    assert available_height - image_height >= 1.30


@pytest.mark.parametrize(
    ("case_name", "page_sizes", "expect_landscape"),
    [
        ("near_square_1_page", [(760, 800)], False),
        ("standard_portrait_2_pages", [(595, 842)] * 2, False),
        ("wide_landscape_1_page", [(1120, 560)], True),
        ("standard_portrait_9_pages", [(595, 842)] * 9, False),
    ],
)
def test_libreoffice_appendix_title_and_source_image_share_exactly_one_page(
    tmp_path,
    case_name,
    page_sizes,
    expect_landscape,
):
    document, source_pages = _appendix_document_from_pdf(
        _pdf_with_page_sizes(page_sizes),
        case_name=case_name,
    )
    assert len(source_pages) == len(page_sizes)
    assert all(int(page["render_dpi"]) >= 200 for page in source_pages)
    assert all(
        int(page["pixel_width"]) / int(page["pixel_height"])
        == pytest.approx(page_sizes[0][0] / page_sizes[0][1], rel=0.01)
        for page in source_pages
    )

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
    )
    exported = Document(io.BytesIO(result.content))
    assert exported.sections[-1].orientation == WD_ORIENT.PORTRAIT
    assert [
        int(page["page_number"])
        for page in result.metadata["appendix_image_pages"]
    ] == list(range(1, len(page_sizes) + 1))

    pdf_path = _render_docx_with_libreoffice(
        tmp_path,
        case_name=case_name,
        content=result.content,
    )
    _assert_rendered_appendix_page_assignment(
        pdf_path,
        page_count=len(page_sizes),
        source_aspect_ratio=page_sizes[0][0] / page_sizes[0][1],
        expect_landscape=expect_landscape,
    )


def test_wide_flowchart_context_and_caption_share_one_landscape_layout_group():
    figure = _wide_study_schema_figure(body_order=4)
    document = ProtocolDocument(
        document_id="doc_flowchart_pagination",
        project_id="proj_flowchart_pagination",
        protocol_id="CMS-FLOWCHART-PAGINATION",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_intro",
                document_id="doc_flowchart_pagination",
                heading="研究背景",
                section_number="1.1",
                content_blocks=[
                    {
                        "block_id": "intro",
                        "block_type": "paragraph",
                        "text": "前置正文。",
                        "body_order": 1,
                    }
                ],
            ),
            ProtocolSection(
                section_id="section_schema",
                document_id="doc_flowchart_pagination",
                heading="研究流程图",
                section_number="1.2",
                node_kind="study_schema",
                content_blocks=[
                    {
                        "block_id": "schema_heading",
                        "block_type": "heading",
                        "text": "1.2 研究流程图",
                        "body_order": 2,
                    },
                    {
                        "block_id": "schema_context",
                        "block_type": "paragraph",
                        "text": "下图展示筛选、治疗和随访阶段。",
                        "body_order": 3,
                    },
                    figure,
                ],
            ),
            ProtocolSection(
                section_id="section_after",
                document_id="doc_flowchart_pagination",
                heading="后续正文",
                section_number="1.3",
                content_blocks=[
                    {
                        "block_id": "after_heading",
                        "block_type": "heading",
                        "text": "1.3 后续正文",
                        "body_order": 5,
                    }
                ],
            ),
        ],
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    section_indices = _paragraph_section_indices(exported.element.body.xml.encode("utf-8"))

    assert [section.orientation for section in exported.sections] == [
        WD_ORIENT.PORTRAIT,
        WD_ORIENT.LANDSCAPE,
        WD_ORIENT.PORTRAIT,
    ]
    assert section_indices["1.2 研究流程图"] == 1
    assert section_indices["下图展示筛选、治疗和随访阶段。"] == 1
    assert section_indices["<drawing>"] == 1
    assert section_indices["图 1 研究设计概况"] == 1
    assert section_indices["1.3 后续正文"] == 2

    heading = next(paragraph for paragraph in exported.paragraphs if paragraph.text == "1.2 研究流程图")
    context = next(
        paragraph
        for paragraph in exported.paragraphs
        if paragraph.text == "下图展示筛选、治疗和随访阶段。"
    )
    assert heading.paragraph_format.keep_with_next is True
    assert context.paragraph_format.keep_with_next is True
    assert context.paragraph_format.keep_together is True


def test_document_indexes_follow_explicit_semantics_and_stable_front_matter_boundary():
    explicit = ProtocolDocument(
        document_id="doc_explicit_indexes",
        project_id="proj_explicit_indexes",
        protocol_id="CMS-EXPLICIT-INDEXES",
        version="V0.1",
        sections=[
            _text_section("front", "首页", "首页内容", 1, node_kind="front_matter"),
            _text_section("version", "版本记录", "版本记录内容", 2, node_kind="document_control"),
            _text_section(
                "indexes",
                "索引占位",
                "索引占位不应导出",
                3,
                node_kind="document_control",
                template_node_id="cms_indexes",
            ),
            _text_section("body", "正文", "1 正文", 4),
        ],
    )
    explicit_export = Document(
        io.BytesIO(export_medical_writing_document_docx(explicit, mode="draft_preview").content)
    )
    explicit_text = [paragraph.text for paragraph in explicit_export.paragraphs]
    assert explicit_text.index("版本记录内容") < explicit_text.index("目录")
    assert explicit_text.index("目录") < explicit_text.index("1 正文")
    assert "索引占位不应导出" not in explicit_text

    no_explicit = ProtocolDocument(
        document_id="doc_default_indexes",
        project_id="proj_default_indexes",
        protocol_id="CMS-DEFAULT-INDEXES",
        version="V0.1",
        sections=[
            _text_section("body_1", "第一章", "1 第一章", 1),
            _text_section("body_2", "第二章", "2 第二章", 2),
            _text_section("references", "参考文献", "参考文献条目", 3),
        ],
    )
    default_export = Document(
        io.BytesIO(export_medical_writing_document_docx(no_explicit, mode="draft_preview").content)
    )
    default_text = [paragraph.text for paragraph in default_export.paragraphs]
    assert default_text.index("目录") < default_text.index("1 第一章")
    assert default_text.index("1 第一章") < default_text.index("2 第二章")
    assert default_text.index("2 第二章") < default_text.index("参考文献条目")
    toc_heading = next(paragraph for paragraph in default_export.paragraphs if paragraph.text == "目录")
    assert toc_heading.style.name == "TOC Heading"
    assert toc_heading._p.find("w:pPr/w:numPr", NS) is None

    inferred_boundary = ProtocolDocument(
        document_id="doc_inferred_index_boundary",
        project_id="proj_inferred_index_boundary",
        protocol_id="CMS-INFERRED-INDEX-BOUNDARY",
        version="V0.1",
        sections=[
            _text_section("cover", "封面", "封面内容", 1, node_kind="front_matter"),
            _text_section("summary", "摘要", "摘要内容", 2, node_kind="document_control"),
            _text_section("body", "正文", "1 正文", 3),
        ],
    )
    inferred_export = Document(
        io.BytesIO(
            export_medical_writing_document_docx(
                inferred_boundary,
                mode="draft_preview",
            ).content
        )
    )
    inferred_text = [paragraph.text for paragraph in inferred_export.paragraphs]
    assert inferred_text.index("封面内容") < inferred_text.index("摘要内容")
    assert inferred_text.index("摘要内容") < inferred_text.index("目录")
    assert inferred_text.index("目录") < inferred_text.index("1 正文")


def test_company_heading_number_is_visible_once_in_editable_heading_text():
    document = _numbering_regression_document()
    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))

    expected = [
        "1.2 Study Flow Diagram",
        "14.2 Assessment Instruments",
        "14.3 Post-appendix Body",
    ]
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")

    for title in expected:
        paragraph = next(item for item in exported.paragraphs if item.text == title)
        assert paragraph.style.name == "Heading 2"
        num_id_node = paragraph._p.find("w:pPr/w:numPr/w:numId", NS)
        assert num_id_node is not None
        assert num_id_node.get(qn("w:val")) == "0"
        assert document_xml.count(title) == 1

    source_text = [
        str(block.get("text") or "")
        for section in document.sections
        for block in section.content_blocks
    ]
    assert source_text == [
        "1.2 Study Flow Diagram",
        "14.2 Assessment Instruments",
        "14.3 Post-appendix Body",
    ]
    assert 'TOC \\o "1-4" \\h \\z \\u' in document_xml


def test_libreoffice_rendered_heading_has_complete_single_composite_number(tmp_path):
    soffice = shutil.which("soffice")
    if not soffice:
        pytest.skip("LibreOffice is not available")
    fitz = pytest.importorskip("fitz")
    result = export_medical_writing_document_docx(
        _numbering_regression_document(),
        mode="draft_preview",
    )
    docx_path = tmp_path / "numbering-regression.docx"
    output_dir = tmp_path / "pdf"
    profile_dir = tmp_path / "libreoffice-profile"
    output_dir.mkdir()
    docx_path.write_bytes(result.content)
    pdf_path = output_dir / "numbering-regression.pdf"
    assert not pdf_path.exists()
    subprocess.run(
        [
            soffice,
            "--headless",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "pdf",
            "--outdir",
            str(output_dir),
            str(docx_path),
        ],
        check=True,
        capture_output=True,
        text=True,
        timeout=60,
    )
    assert pdf_path.is_file()
    assert pdf_path.stat().st_mtime_ns >= docx_path.stat().st_mtime_ns
    rendered = fitz.open(pdf_path)
    pdf_text = "\n".join(page.get_text("text") for page in rendered)
    normalized = re.sub(r"\s+", " ", pdf_text)

    assert "1.2 Study Flow Diagram" in normalized
    assert "14.2 Assessment Instruments" in normalized
    assert "14.3 Post-appendix Body" in normalized
    assert not re.search(r"1\.1\s+1\.2\s+Study Flow Diagram", normalized)
    assert not re.search(r"1\.2\s+14\.2\s+Assessment Instruments", normalized)
    assert not re.search(r"1\.3\s+14\.3\s+Post-appendix Body", normalized)

    target_page = next(
        page for page in rendered if "Post-appendix Body" in page.get_text("text")
    )
    number_boxes = target_page.search_for("14.3")
    assert len(number_boxes) == 1
    number_box = number_boxes[0]
    prefix_box = fitz.Rect(
        number_box.x0,
        number_box.y0,
        number_box.x0 + number_box.width * 0.72,
        number_box.y1,
    )
    suffix_box = fitz.Rect(
        number_box.x0 + number_box.width * 0.72,
        number_box.y0,
        number_box.x1,
        number_box.y1,
    )
    assert _rendered_ink_ratio(target_page, prefix_box, fitz) >= 0.015
    assert _rendered_ink_ratio(target_page, suffix_box, fitz) >= 0.015
    rendered.close()


def _rendered_ink_ratio(page, clip, fitz_module) -> float:
    pixmap = page.get_pixmap(
        matrix=fitz_module.Matrix(2, 2),
        clip=clip,
        colorspace=fitz_module.csGRAY,
        alpha=False,
    )
    if not pixmap.samples:
        return 0.0
    return sum(value < 220 for value in pixmap.samples) / len(pixmap.samples)


def _text_section(
    section_id: str,
    heading: str,
    text: str,
    body_order: int,
    *,
    node_kind: str = "section",
    template_node_id: str = "",
) -> ProtocolSection:
    return ProtocolSection(
        section_id=section_id,
        document_id="unused",
        heading=heading,
        node_kind=node_kind,
        template_node_id=template_node_id,
        content_blocks=[
            {
                "block_id": f"{section_id}_block",
                "block_type": "paragraph",
                "text": text,
                "body_order": body_order,
            }
        ],
    )


def _numbering_regression_document() -> ProtocolDocument:
    definition = MedicalWritingStyleProfileService().definition()
    headings = (
        ("flow", "1.2", "Study Flow Diagram"),
        ("appendix", "14.2", "Assessment Instruments"),
        ("after", "14.3", "Post-appendix Body"),
    )
    return ProtocolDocument(
        document_id="doc_heading_numbering_regression",
        project_id="proj_heading_numbering_regression",
        protocol_id="CMS-HEADING-NUMBERING",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id=f"section_{key}",
                document_id="doc_heading_numbering_regression",
                heading=title,
                section_number=number,
                content_blocks=[
                    {
                        "block_id": f"heading_{key}",
                        "block_type": "heading",
                        "text": f"{number} {title}",
                        "body_order": index,
                        "outline_level": 1,
                        **(
                            {
                                "rich_text": {
                                    "type": "heading",
                                    "attrs": {"level": 2},
                                    "content": [
                                        {"type": "text", "text": "1."},
                                        {"type": "text", "text": "2 "},
                                        {
                                            "type": "text",
                                            "text": title,
                                            "marks": [{"type": "bold"}],
                                        },
                                    ],
                                }
                            }
                            if key == "flow"
                            else {}
                        ),
                    }
                ],
            )
            for index, (key, number, title) in enumerate(headings, start=1)
        ],
        style_profile_id=definition.style_profile_id,
        style_profile_version=definition.style_profile_version,
        style_profile_definition_sha256=definition.definition_sha256,
    )
