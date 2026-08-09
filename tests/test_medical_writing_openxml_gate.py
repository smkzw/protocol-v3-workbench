from __future__ import annotations

import json
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document

from deploy.medical_writing_local.build_release_bundle import (
    validate_word_validation_report,
)


WORKBENCH_ROOT = Path(__file__).resolve().parents[1]
VALIDATOR = (
    WORKBENCH_ROOT
    / "tools/openxml_docx_validator/bin/osx-arm64/openxml-docx-validator"
)
GATE = WORKBENCH_ROOT / "scripts/qc/run_medical_writing_openxml_gate.py"


def _create_docx(path: Path, text: str) -> None:
    document = Document()
    document.add_heading("临床试验方案", level=1)
    document.add_paragraph(text)
    document.save(path)


def _replace_document_xml(source: Path, target: Path, before: bytes, after: bytes) -> None:
    with ZipFile(source, "r") as input_zip, ZipFile(
        target, "w", compression=ZIP_DEFLATED
    ) as output_zip:
        replaced = False
        for item in input_zip.infolist():
            payload = input_zip.read(item.filename)
            if item.filename == "word/document.xml":
                if before not in payload:
                    raise AssertionError(f"expected marker is missing: {before!r}")
                payload = payload.replace(before, after, 1)
                replaced = True
            output_zip.writestr(item, payload)
        if not replaced:
            raise AssertionError("word/document.xml was not replaced")


def _gate_report(
    generated: Path,
    source_a: Path,
    passthrough_a: Path,
    edited_a: Path,
    source_b: Path,
    passthrough_b: Path,
    edited_b: Path,
    output: Path,
) -> subprocess.CompletedProcess[str]:
    return subprocess.run(
        [
            "python3",
            str(GATE),
            "--validator",
            str(VALIDATOR),
            "--generated",
            f"greenfield={generated}",
            "--passthrough",
            f"project-a={source_a}:{passthrough_a}",
            "--passthrough",
            f"project-b={source_b}:{passthrough_b}",
            "--imported-edit",
            f"project-a-paragraph={source_a}:{edited_a}",
            "--imported-edit",
            f"project-b-paragraph={source_b}:{edited_b}",
            "--output",
            str(output),
        ],
        check=False,
        capture_output=True,
        text=True,
    )


@unittest.skipUnless(VALIDATOR.is_file(), "OpenXML validator binary is unavailable")
class MedicalWritingOpenXmlGateTests(unittest.TestCase):
    def test_gate_accepts_greenfield_and_two_distinct_imported_projects(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            generated = root / "generated.docx"
            source_a = root / "source_a.docx"
            source_b = root / "source_b.docx"
            _create_docx(generated, "从零生成的研究方案。")
            _create_docx(source_a, "项目A来源文本。")
            _create_docx(source_b, "项目B来源文本。")
            passthrough_a = root / "passthrough_a.docx"
            passthrough_b = root / "passthrough_b.docx"
            shutil.copyfile(source_a, passthrough_a)
            shutil.copyfile(source_b, passthrough_b)
            edited_a = root / "edited_a.docx"
            edited_b = root / "edited_b.docx"
            _replace_document_xml(
                source_a,
                edited_a,
                "项目A来源文本。".encode(),
                "项目A修订文本。".encode(),
            )
            _replace_document_xml(
                source_b,
                edited_b,
                "项目B来源文本。".encode(),
                "项目B修订文本。".encode(),
            )

            output = root / "report.json"
            result = _gate_report(
                generated,
                source_a,
                passthrough_a,
                edited_a,
                source_b,
                passthrough_b,
                edited_b,
                output,
            )

            self.assertEqual(0, result.returncode, result.stderr)
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual("pass", report["status"])
            self.assertEqual(5, report["case_count"])
            summary = validate_word_validation_report(report)
            self.assertEqual(2, summary["distinct_imported_project_count"])

    def test_gate_rejects_a_new_openxml_error(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            valid = root / "valid.docx"
            invalid = root / "invalid.docx"
            _create_docx(valid, "结构有效。")
            _replace_document_xml(
                valid,
                invalid,
                b"</w:body>",
                b"<w:tblHeader w:val=\"true\"/></w:body>",
            )
            output = root / "negative.json"
            result = subprocess.run(
                [
                    "python3",
                    str(GATE),
                    "--validator",
                    str(VALIDATOR),
                    "--generated",
                    f"invalid={invalid}",
                    "--output",
                    str(output),
                ],
                check=False,
                capture_output=True,
                text=True,
            )

            self.assertEqual(1, result.returncode)
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual("fail", report["status"])
            self.assertGreater(report["results"][0]["error_count"], 0)


class WordValidationReleaseContractTests(unittest.TestCase):
    def _report(self) -> dict:
        def imported(name: str, contract: str, source_sha: str) -> dict:
            return {
                "name": name,
                "contract": contract,
                "passed": True,
                "source": {"sha256": source_sha},
                "candidate": {"sha256": source_sha},
                "sha256_identical": True,
                "new_error_count": 0,
            }

        results = [
            {
                "name": "greenfield",
                "contract": "generated_zero_errors",
                "passed": True,
                "error_count": 0,
            },
            imported("project-a-pass", "imported_passthrough_identical", "a" * 64),
            imported("project-b-pass", "imported_passthrough_identical", "b" * 64),
            imported("project-a-edit", "imported_edit_no_new_errors", "a" * 64),
            imported("project-b-edit", "imported_edit_no_new_errors", "b" * 64),
        ]
        return {
            "schema_version": "medical_writing_openxml_gate_v1",
            "status": "pass",
            "case_count": len(results),
            "passed_count": len(results),
            "failed_count": 0,
            "validator": {
                "sha256": "c" * 64,
                "license": "MIT",
                "document_format_openxml_version": "3.5.1",
                "file_format_version": "Microsoft365",
            },
            "results": results,
        }

    def test_release_contract_accepts_complete_free_tool_report(self):
        summary = validate_word_validation_report(self._report())
        self.assertEqual("MIT", summary["validator_license"])
        self.assertEqual(2, summary["distinct_imported_project_count"])

    def test_release_contract_rejects_commercial_validator(self):
        report = self._report()
        report["validator"]["license"] = "commercial"
        with self.assertRaisesRegex(RuntimeError, "MIT"):
            validate_word_validation_report(report)

    def test_release_contract_rejects_one_imported_project_disguised_as_two(self):
        report = self._report()
        for item in report["results"]:
            if item["contract"].startswith("imported_"):
                item["source"]["sha256"] = "a" * 64
        with self.assertRaisesRegex(RuntimeError, "two distinct"):
            validate_word_validation_report(report)

    def test_release_contract_rejects_unknown_contract(self):
        report = self._report()
        report["results"].append(
            {"name": "unclassified", "contract": "unknown", "passed": True}
        )
        report["case_count"] += 1
        report["passed_count"] += 1
        with self.assertRaisesRegex(RuntimeError, "unknown contract"):
            validate_word_validation_report(report)


if __name__ == "__main__":
    unittest.main()
