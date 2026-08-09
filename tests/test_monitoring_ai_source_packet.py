from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

import pytest
from docx import Document

from packages.contracts.workbench_contracts import (
    SourceRegistrationResult,
    SourceRegistryEntry,
    SourceRegistrySpan,
)
from services.api.app.monitoring_ai_source_packet import (
    MAX_PROTOCOL_CANDIDATE_EVIDENCE_IDS,
    MonitoringAiSourcePacketResolver,
    PROTOCOL_EVIDENCE_PACKET_VERSION,
    PROTOCOL_STRUCTURAL_REPAIR_VERSION,
    ProtocolStructuralRepairError,
    expand_protocol_evidence_spans,
    focus_protocol_provider_evidence,
    repair_protocol_structural_bundles,
)
from services.api.app.source_intake import (
    MAX_PROTOCOL_REGISTRY_SPANS,
    SourceRegistryService,
    SourceRegistryStore,
    protocol_document_to_ai_sources,
)
from services.api.app.protocol_text_extractor import (
    ProtocolTextDocument,
    ProtocolTextSpan,
)


def _registry(tmp_path: Path) -> SourceRegistryService:
    store = SourceRegistryStore(tmp_path / "sources.jsonl")
    now = datetime(2026, 7, 29, tzinfo=timezone.utc)
    entry = SourceRegistryEntry(
        entry_id="src_protocol_001",
        project_id="project-alpha",
        module="medical_monitoring",
        source_kind="protocol_docx",
        public_title="研究方案",
        content_hash="a" * 64,
        created_at=now,
    )
    spans = [
        SourceRegistrySpan(
            source_id=f"src_protocol_001_span_{index}",
            entry_id=entry.entry_id,
            project_id=entry.project_id,
            module=entry.module,
            source_type="protocol_docx_paragraph",
            title="研究方案",
            locator=f"docx:paragraph:{index}",
            text_preview=f"方案原文 {index}",
            created_at=now,
        )
        for index in (1, 2)
    ]
    store.append(SourceRegistrationResult(entry=entry, spans=spans))
    return SourceRegistryService(store)


def test_source_packet_is_stable_source_bound_and_server_materialized(
    tmp_path: Path,
) -> None:
    resolver = MonitoringAiSourcePacketResolver(_registry(tmp_path))

    first = resolver.resolve(
        "project-alpha",
        ["src_protocol_001_span_2", "src_protocol_001_span_1"],
    )
    second = resolver.resolve(
        "project-alpha",
        ["src_protocol_001_span_2", "src_protocol_001_span_1"],
    )

    assert first == second
    assert first.source_ids == (
        "src_protocol_001_span_2",
        "src_protocol_001_span_1",
    )
    assert len(first.input_revision.sources) == 1
    assert first.input_revision.sources[0].source_entry_id == "src_protocol_001"
    assert {
        item["source_entry_id"] for item in first.evidence_packet
    } == {"src_protocol_001"}
    assert [item["quote"] for item in first.evidence_packet] == [
        "方案原文 2",
        "方案原文 1",
    ]


def test_source_packet_rejects_missing_empty_or_wrong_project_sources(
    tmp_path: Path,
) -> None:
    resolver = MonitoringAiSourcePacketResolver(_registry(tmp_path))

    with pytest.raises(ValueError, match="non-empty"):
        resolver.resolve("project-alpha", [])
    with pytest.raises(KeyError, match="not registered"):
        resolver.resolve("project-alpha", ["missing"])
    with pytest.raises(KeyError, match="not registered"):
        resolver.resolve("project-beta", ["src_protocol_001_span_1"])


def _protocol_bytes(paragraph_count: int) -> bytes:
    document = Document()
    for index in range(1, paragraph_count + 1):
        suffix = " 后半部禁用药证据" if index == paragraph_count else ""
        document.add_paragraph(f"方案段落 {index}{suffix} 共同检索词")
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def test_protocol_registration_preserves_tail_spans_and_duplicate_is_compatible(
    tmp_path: Path,
) -> None:
    store = SourceRegistryStore(tmp_path / "sources.jsonl")
    registry = SourceRegistryService(store)
    content = _protocol_bytes(260)

    first = registry.register_protocol_docx(
        "project-tail",
        "tail-protocol.docx",
        content,
        module="medical_monitoring",
    )
    second = registry.register_protocol_docx(
        "project-tail",
        "tail-protocol.docx",
        content,
        module="medical_monitoring",
    )
    resolver = MonitoringAiSourcePacketResolver(registry)
    results = resolver.search(
        "project-tail",
        first.entry.entry_id,
        ["后半部", "禁用药"],
    )

    assert first.entry.span_count >= 260
    assert second.entry.entry_id == first.entry.entry_id
    assert len(store.list_results("project-tail")) == 1
    assert len(results) == 1
    assert "后半部禁用药证据" in results[0]["text"]
    assert results[0]["source_id"] == first.spans[-1].source_id
    assert results[0]["matched_keywords"] == ["后半部", "禁用药"]
    assert "命中关键词" in results[0]["match_reason"]


def test_evidence_span_search_is_project_scoped_empty_and_bounded(
    tmp_path: Path,
) -> None:
    registry = SourceRegistryService(SourceRegistryStore(tmp_path / "sources.jsonl"))
    registered = registry.register_protocol_docx(
        "project-alpha",
        "protocol.docx",
        _protocol_bytes(260),
        module="medical_monitoring",
    )
    resolver = MonitoringAiSourcePacketResolver(registry)

    assert resolver.search(
        "project-alpha",
        registered.entry.entry_id,
        ["不存在的关键词"],
    ) == ()
    limited = resolver.search(
        "project-alpha",
        registered.entry.entry_id,
        ["共同检索词"],
        limit=200,
    )
    assert len(limited) == 200
    assert all(item["source_entry_id"] == registered.entry.entry_id for item in limited)
    with pytest.raises(ValueError, match="between 1 and 200"):
        resolver.search(
            "project-alpha",
            registered.entry.entry_id,
            ["共同检索词"],
            limit=201,
        )
    with pytest.raises(KeyError, match="not registered"):
        resolver.search(
            "project-beta",
            registered.entry.entry_id,
            ["共同检索词"],
        )


def test_protocol_registry_hard_limit_rejects_abnormal_document() -> None:
    span = ProtocolTextSpan(
        span_id="p1",
        kind="paragraph",
        text="方案正文",
        source_locator="docx:paragraph:1",
    )
    document = ProtocolTextDocument(
        filename="abnormal.docx",
        title="异常方案",
        paragraphs=[],
        tables=[],
        spans=[span] * (MAX_PROTOCOL_REGISTRY_SPANS + 1),
    )

    with pytest.raises(ValueError, match="registry hard limit"):
        protocol_document_to_ai_sources(document, "abnormal")


def test_evidence_span_search_rejects_wrong_module_and_superseded_protocol(
    tmp_path: Path,
) -> None:
    store = SourceRegistryStore(tmp_path / "sources.jsonl")
    registry = SourceRegistryService(store)
    old = registry.register_protocol_docx(
        "project-boundary",
        "old.docx",
        _protocol_bytes(2),
        module="medical_monitoring",
    )
    current = registry.register_protocol_docx(
        "project-boundary",
        "current.docx",
        _protocol_bytes(3),
        module="medical_monitoring",
    )
    writing = registry.register_protocol_docx(
        "project-boundary",
        "writing.docx",
        _protocol_bytes(4),
        module="medical_writing",
    )
    resolver = MonitoringAiSourcePacketResolver(registry)

    assert resolver.search(
        "project-boundary",
        current.entry.entry_id,
        ["共同检索词"],
    )
    with pytest.raises(ValueError, match="superseded"):
        resolver.search(
            "project-boundary",
            old.entry.entry_id,
            ["共同检索词"],
        )
    with pytest.raises(ValueError, match="module boundary"):
        resolver.search(
            "project-boundary",
            writing.entry.entry_id,
            ["共同检索词"],
        )


def _shape_span(
    source_id: str,
    locator: str,
    text: str,
    *,
    source_type: str = "protocol_docx_paragraph",
) -> SourceRegistrySpan:
    return SourceRegistrySpan(
        source_id=source_id,
        entry_id="shape-entry",
        project_id="shape-project",
        module="medical_monitoring",
        source_type=source_type,
        title="跨项目真实形状方案",
        locator=locator,
        text_preview=text,
        created_at=datetime(2026, 7, 30, tzinfo=timezone.utc),
    )


def test_protocol_evidence_v2_expands_table_header_same_row_and_caption() -> None:
    spans = [
        _shape_span("caption", "docx:paragraph:40", "表4 暂停和停药标准"),
        _shape_span(
            "header-condition",
            "docx:table:4:row:0:cell:0:paragraph:41",
            "触发条件",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "header-action",
            "docx:table:4:row:0:cell:1:paragraph:42",
            "处理措施",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "row-threshold",
            "docx:table:4:row:2:cell:0:paragraph:43",
            "ALT或AST>3×ULN",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "row-action",
            "docx:table:4:row:2:cell:1:paragraph:44",
            "48小时内复查并暂停研究药物",
            source_type="protocol_docx_table_cell",
        ),
    ]
    expanded = expand_protocol_evidence_spans(
        spans,
        [
            {
                "source_id": "row-action",
                "source_entry_id": "shape-entry",
                "locator": spans[-1].locator,
                "text": spans[-1].text_preview,
                "score": 1001,
                "matched_keywords": ["暂停"],
                "match_reason": "命中关键词：暂停",
            }
        ],
    )

    by_id = {item["source_id"]: item for item in expanded}
    assert set(by_id) == {
        "caption",
        "header-condition",
        "header-action",
        "row-threshold",
        "row-action",
    }
    assert "table_caption" in by_id["caption"]["evidence_context"]["roles"]
    assert "table_header" in by_id["header-action"]["evidence_context"]["roles"]
    assert (
        "table_row_context"
        in by_id["row-threshold"]["evidence_context"]["roles"]
    )
    assert by_id["row-action"]["evidence_context"]["packet_version"] == (
        PROTOCOL_EVIDENCE_PACKET_VERSION
    )


def test_protocol_evidence_v2_carries_prohibited_list_and_sae_exceptions() -> None:
    spans = [
        _shape_span("cm-heading", "docx:paragraph:100", "禁止使用以下药物或治疗："),
        _shape_span("cm-1", "docx:paragraph:101", "（1）其他试验用药；"),
        _shape_span("cm-2", "docx:paragraph:102", "（2）外用糖皮质激素；"),
        _shape_span("cm-3", "docx:paragraph:103", "（3）免疫抑制剂及生物疗法。"),
        _shape_span("sae", "docx:paragraph:200", "住院或延长住院属于严重不良事件。"),
        _shape_span(
            "sae-exception-title",
            "docx:paragraph:201",
            "以下住院情形可由研究者综合判断不作为SAE：",
        ),
        _shape_span("sae-ex-1", "docx:paragraph:202", "（1）不超过24小时的留观；"),
        _shape_span("sae-ex-2", "docx:paragraph:203", "（2）预先计划的住院；"),
        _shape_span("sae-ex-3", "docx:paragraph:204", "（3）方案计划的程序。"),
    ]
    matches = [
        {
            "source_id": "cm-heading",
            "source_entry_id": "shape-entry",
            "locator": spans[0].locator,
            "text": spans[0].text_preview,
            "score": 1001,
            "matched_keywords": ["禁止使用"],
            "match_reason": "命中关键词：禁止使用",
        },
        {
            "source_id": "sae",
            "source_entry_id": "shape-entry",
            "locator": spans[4].locator,
            "text": spans[4].text_preview,
            "score": 1001,
            "matched_keywords": ["严重不良事件"],
            "match_reason": "命中关键词：严重不良事件",
        },
    ]

    expanded = expand_protocol_evidence_spans(spans, matches)
    by_id = {item["source_id"]: item for item in expanded}

    assert {"cm-heading", "cm-1", "cm-2", "cm-3"}.issubset(by_id)
    assert {
        "sae",
        "sae-exception-title",
        "sae-ex-1",
        "sae-ex-2",
        "sae-ex-3",
    }.issubset(by_id)
    assert "list_title" in by_id["cm-heading"]["evidence_context"]["roles"]
    assert "list_item" in by_id["cm-3"]["evidence_context"]["roles"]
    assert (
        "list_title"
        in by_id["sae-exception-title"]["evidence_context"]["roles"]
    )


def test_protocol_evidence_v2_recovers_preceding_title_for_item_only_match() -> None:
    spans = [
        _shape_span(
            "section-heading",
            "docx:paragraph:8",
            "6.3.5 受试者终止标准",
        ),
        _shape_span(
            "list-title",
            "docx:paragraph:10",
            "符合以下任一情况时，应提前退出研究：",
        ),
        _shape_span(
            "list-item-1",
            "docx:paragraph:11",
            "（1）受试者撤回知情同意；",
        ),
        _shape_span(
            "list-item-2",
            "docx:paragraph:12",
            "（2）研究者判断继续参加不符合受试者利益；",
        ),
        _shape_span(
            "after-list",
            "docx:paragraph:13",
            "退出后应尽量完成安全性随访。",
        ),
    ]
    matches = [
        {
            "source_id": "list-item-2",
            "source_entry_id": "shape-entry",
            "locator": spans[3].locator,
            "text": spans[3].text_preview,
            "score": 1001,
            "matched_keywords": ["研究者判断"],
            "match_reason": "命中关键词：研究者判断",
        }
    ]

    expanded = expand_protocol_evidence_spans(spans, matches)
    by_id = {item["source_id"]: item for item in expanded}

    assert "list_title" in by_id["list-title"]["evidence_context"]["roles"]
    assert "list_item" in by_id["list-item-1"]["evidence_context"]["roles"]
    assert {
        "primary_match",
        "list_item",
    }.issubset(by_id["list-item-2"]["evidence_context"]["roles"])
    assert "list_item" not in by_id["after-list"]["evidence_context"]["roles"]
    assert by_id["after-list"]["evidence_context"].get(
        "list_bundle_id"
    ) is None
    assert by_id["list-item-2"]["evidence_context"]["list_bundle_id"] == (
        by_id["list-title"]["evidence_context"]["list_bundle_id"]
    )
    assert by_id["list-title"]["evidence_context"]["list_bundle_id"] == (
        "list-title"
    )


def test_protocol_evidence_v2_binds_visit_table_window_to_visit_name() -> None:
    spans = [
        _shape_span("visit-caption", "docx:paragraph:20", "表2 研究流程表"),
        _shape_span(
            "visit-header",
            "docx:table:2:row:0:cell:0:paragraph:21",
            "访视",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "window-header",
            "docx:table:2:row:0:cell:1:paragraph:22",
            "允许时间窗",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "visit-name",
            "docx:table:2:row:1:cell:0:paragraph:23",
            "D15、D29、D57",
            source_type="protocol_docx_table_cell",
        ),
        _shape_span(
            "visit-window",
            "docx:table:2:row:1:cell:1:paragraph:24",
            "±3天",
            source_type="protocol_docx_table_cell",
        ),
    ]
    expanded = expand_protocol_evidence_spans(
        spans,
        [
            {
                "source_id": "visit-window",
                "source_entry_id": "shape-entry",
                "locator": spans[-1].locator,
                "text": spans[-1].text_preview,
                "score": 1001,
                "matched_keywords": ["时间窗"],
                "match_reason": "命中关键词：时间窗",
            }
        ],
    )
    by_id = {item["source_id"]: item for item in expanded}

    assert {"visit-header", "window-header", "visit-name", "visit-window"}.issubset(
        by_id
    )
    assert (
        by_id["visit-name"]["evidence_context"]["structure"]["row_index"]
        == by_id["visit-window"]["evidence_context"]["structure"]["row_index"]
        == 1
    )


def _payload_evidence(
    evidence_id: str,
    locator: str,
    quote: str,
    roles: tuple[str, ...],
    structure: dict[str, Any],
    parents: tuple[str, ...] = (),
    bundle_id: str | None = None,
) -> dict[str, Any]:
    protocol_context: dict[str, Any] = {
        "packet_version": PROTOCOL_EVIDENCE_PACKET_VERSION,
        "primary_match": "primary_match" in roles,
        "roles": list(roles),
        "parent_match_source_ids": list(parents),
        "structure": structure,
    }
    if bundle_id is not None:
        protocol_context["list_bundle_id"] = bundle_id
    return {
        "evidence_id": evidence_id,
        "source_entry_id": "shape-entry",
        "source_content_sha256": "a" * 64,
        "locator": locator,
        "quote": quote,
        "raw_fields": {
            "source_id": f"src-{evidence_id}",
            "source_type": "protocol_docx_paragraph",
            "protocol_context": protocol_context,
        },
    }


def test_item_only_match_expansion_remains_closed_after_provider_focus() -> None:
    spans = [
        _shape_span(
            "list-title",
            "docx:paragraph:20",
            "研究期间禁止使用以下治疗：",
        ),
        _shape_span(
            "list-item-1",
            "docx:paragraph:21",
            "（1）系统性糖皮质激素；",
        ),
        _shape_span(
            "list-item-2",
            "docx:paragraph:22",
            "（2）其他免疫抑制剂。",
        ),
    ]
    matches = [
        {
            "source_id": "list-item-2",
            "source_entry_id": "shape-entry",
            "locator": spans[2].locator,
            "text": spans[2].text_preview,
            "score": 1001,
            "matched_keywords": ["免疫抑制剂"],
            "match_reason": "命中关键词：免疫抑制剂",
        }
    ]
    expanded = expand_protocol_evidence_spans(spans, matches)
    packet = [
        _payload_evidence(
            f"evidence-{item['source_id']}",
            item["locator"],
            item["text"],
            tuple(item["evidence_context"]["roles"]),
            item["evidence_context"]["structure"],
            parents=tuple(
                item["evidence_context"]["parent_match_source_ids"]
            ),
            bundle_id=str(
                item["evidence_context"].get("list_bundle_id") or ""
            )
            or None,
        )
        for item in expanded
    ]

    focused = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in focused] == [
        "evidence-list-title",
        "evidence-list-item-2",
    ]
    assert "evidence-list-item-1" not in {
        item["evidence_id"] for item in focused
    }


def test_protocol_evidence_focus_keeps_primary_and_conflict_ids_only() -> None:
    packet = [
        _payload_evidence(
            "evidence-primary",
            "docx:paragraph:10",
            "主要命中条款。",
            ("primary_match",),
            {"kind": "paragraph", "paragraph_index": 10},
        ),
        _payload_evidence(
            "evidence-conflict",
            "docx:paragraph:20",
            "冲突另一侧条款。",
            ("adjacent_paragraph",),
            {"kind": "paragraph", "paragraph_index": 20},
        ),
        _payload_evidence(
            "evidence-adjacent",
            "docx:paragraph:11",
            "非必要的邻接段落。",
            ("adjacent_paragraph",),
            {"kind": "paragraph", "paragraph_index": 11},
        ),
        _payload_evidence(
            "evidence-caption",
            "docx:paragraph:12",
            "表4 暂停标准",
            ("table_caption",),
            {"kind": "paragraph", "paragraph_index": 12},
        ),
    ]
    focused = focus_protocol_provider_evidence(
        packet,
        conflict_evidence_ids=("evidence-conflict",),
    )

    assert [item["evidence_id"] for item in focused] == [
        "evidence-primary",
        "evidence-conflict",
    ]
    assert packet[0]["raw_fields"]["protocol_context"]["roles"] == [
        "primary_match"
    ]


def test_protocol_evidence_focus_keeps_same_row_and_header_bundle() -> None:
    packet = [
        _payload_evidence(
            "evidence-header-condition",
            "docx:table:4:row:0:cell:0:paragraph:1",
            "触发条件",
            ("table_header",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 0,
                "cell_index": 0,
            },
        ),
        _payload_evidence(
            "evidence-header-action",
            "docx:table:4:row:0:cell:1:paragraph:2",
            "处理措施",
            ("table_header",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 0,
                "cell_index": 1,
            },
        ),
        _payload_evidence(
            "evidence-threshold",
            "docx:table:4:row:2:cell:0:paragraph:3",
            "ALT或AST>3×ULN",
            ("table_row_context",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 2,
                "cell_index": 0,
            },
        ),
        _payload_evidence(
            "evidence-action",
            "docx:table:4:row:2:cell:1:paragraph:4",
            "48小时内复查并暂停研究药物",
            ("primary_match", "table_row_context"),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 2,
                "cell_index": 1,
            },
        ),
        _payload_evidence(
            "evidence-other-row",
            "docx:table:4:row:3:cell:0:paragraph:5",
            "未触碰行的阈值",
            ("table_row_context",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 3,
                "cell_index": 0,
            },
        ),
    ]
    focused = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in focused] == [
        "evidence-header-condition",
        "evidence-header-action",
        "evidence-threshold",
        "evidence-action",
    ]


def test_protocol_evidence_focus_title_only_stays_title_only() -> None:
    packet = [
        _payload_evidence(
            "evidence-list-title",
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("primary_match", "list_title"),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=("matched-1",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:31",
            "（1）其他试验用药；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-1",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-2",
            "docx:paragraph:32",
            "（2）免疫抑制剂；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 32},
            parents=("matched-1",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-unrelated",
            "docx:paragraph:40",
            "其他章节内容。",
            ("adjacent_paragraph",),
            {"kind": "paragraph", "paragraph_index": 40},
        ),
    ]
    focused = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in focused] == [
        "evidence-list-title",
    ]


def test_protocol_evidence_focus_pulls_unique_title_not_siblings() -> None:
    packet = [
        _payload_evidence(
            "evidence-list-title",
            "docx:paragraph:30",
            "允许使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=("matched-2",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:31",
            "（1）外用糖皮质激素；",
            ("primary_match", "list_item"),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-2",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-2",
            "docx:paragraph:32",
            "（2）支持治疗；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 32},
            parents=("matched-2",),
            bundle_id="src-list-title",
        ),
    ]
    focused = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in focused] == [
        "evidence-list-title",
        "evidence-item-1",
    ]


def test_protocol_evidence_focus_keeps_section_heading_and_is_deterministic() -> None:
    packet = [
        _payload_evidence(
            "evidence-heading",
            "docx:paragraph:5",
            "入选标准",
            ("section_heading",),
            {"kind": "paragraph", "paragraph_index": 5},
            parents=("src-evidence-primary",),
        ),
        _payload_evidence(
            "evidence-primary",
            "docx:paragraph:6",
            "年龄18至75岁。",
            ("primary_match",),
            {"kind": "paragraph", "paragraph_index": 6},
        ),
    ]
    first = focus_protocol_provider_evidence(packet)
    second = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in first] == [
        "evidence-heading",
        "evidence-primary",
    ]
    assert first == second


def test_protocol_evidence_focus_is_passthrough_without_v2_context() -> None:
    bare = {
        "evidence_id": "evidence-legacy",
        "source_entry_id": "shape-entry",
        "source_content_sha256": "a" * 64,
        "locator": "docx:paragraph:1",
        "quote": "旧合同证据。",
        "raw_fields": {"source_id": "src-legacy"},
    }
    partial = _payload_evidence(
        "evidence-partial",
        "docx:paragraph:2",
        "缺少roles的证据。",
        (),
        {"kind": "paragraph"},
    )
    partial["raw_fields"]["protocol_context"].pop("roles", None)

    assert focus_protocol_provider_evidence((bare,)) == (bare,)
    assert [item["evidence_id"] for item in focus_protocol_provider_evidence(
        (partial,)
    )] == ["evidence-partial"]


def test_protocol_evidence_focus_limit_never_splits_root_bundle() -> None:
    packet = [
        _payload_evidence(
            "evidence-heading",
            "docx:paragraph:2",
            "安全性评估",
            ("section_heading",),
            {"kind": "paragraph", "paragraph_index": 2},
            parents=("src-evidence-action",),
        ),
        _payload_evidence(
            "evidence-header",
            "docx:table:4:row:0:cell:0:paragraph:3",
            "处理措施",
            ("table_header",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 0,
                "cell_index": 0,
            },
        ),
        _payload_evidence(
            "evidence-threshold",
            "docx:table:4:row:2:cell:0:paragraph:4",
            "ALT>3×ULN",
            ("table_row_context",),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 2,
                "cell_index": 0,
            },
        ),
        _payload_evidence(
            "evidence-action",
            "docx:table:4:row:2:cell:1:paragraph:5",
            "48小时内复查",
            ("primary_match", "table_row_context"),
            {
                "kind": "table_cell",
                "table_index": 4,
                "row_index": 2,
                "cell_index": 1,
            },
        ),
    ]
    focused = focus_protocol_provider_evidence(packet, limit=2)

    assert {item["evidence_id"] for item in focused} == {
        "evidence-header",
        "evidence-threshold",
        "evidence-action",
    }


def test_protocol_evidence_focus_rejects_invalid_limit_and_does_not_mutate() -> None:
    packet = [
        _payload_evidence(
            "evidence-primary",
            "docx:paragraph:1",
            "主要命中。",
            ("primary_match",),
            {"kind": "paragraph", "paragraph_index": 1},
        )
    ]
    snapshot = [dict(item) for item in packet]

    with pytest.raises(ValueError, match="between 1 and 200"):
        focus_protocol_provider_evidence(packet, limit=201)

    assert focus_protocol_provider_evidence(packet)[0] == packet[0]
    assert packet == snapshot
    assert MAX_PROTOCOL_CANDIDATE_EVIDENCE_IDS == 50


def _table_packet(
    table_index: int = 4,
    id_prefix: str = "",
    rows: tuple[tuple[int, tuple[tuple[str, str, int], ...]], ...] = (
        (
            2,
            (
                ("evidence-threshold", "ALT或AST>3×ULN。", 0),
                ("evidence-action", "48小时内复查并暂停研究药物。", 1),
            ),
        ),
    ),
    headers: tuple[tuple[str, str, int], ...] = (
        ("evidence-header-condition", "触发条件", 0),
        ("evidence-header-action", "处理措施", 1),
    ),
) -> list[dict[str, Any]]:
    packet: list[dict[str, Any]] = []
    for evidence_id, quote, cell_index in headers:
        packet.append(
            _payload_evidence(
                f"{id_prefix}{evidence_id}",
                f"docx:table:{table_index}:row:0:cell:{cell_index}:paragraph:1",
                quote,
                ("table_header",),
                {
                    "kind": "table_cell",
                    "table_index": table_index,
                    "row_index": 0,
                    "cell_index": cell_index,
                },
            )
        )
    for row_index, row_cells in rows:
        for cell_index, (evidence_id, quote, _cell_index) in enumerate(
            row_cells
        ):
            packet.append(
                _payload_evidence(
                    f"{id_prefix}{evidence_id}",
                    (
                        f"docx:table:{table_index}:row:{row_index}:"
                        f"cell:{cell_index}:paragraph:{row_index + 1}"
                    ),
                    quote,
                    ("table_row_context",),
                    {
                        "kind": "table_cell",
                        "table_index": table_index,
                        "row_index": row_index,
                        "cell_index": cell_index,
                    },
                )
            )
    return packet


def test_structural_repair_table_expands_same_row_and_header_path() -> None:
    packet = _table_packet()

    result = repair_protocol_structural_bundles(
        ["evidence-action"],
        packet,
    )

    assert result.schema_version == PROTOCOL_STRUCTURAL_REPAIR_VERSION
    assert result.original_evidence_ids == ("evidence-action",)
    assert result.added_structural_context_ids == (
        "evidence-header-condition",
        "evidence-header-action",
        "evidence-threshold",
    )
    assert result.expanded_evidence_ids == (
        "evidence-header-condition",
        "evidence-header-action",
        "evidence-threshold",
        "evidence-action",
    )
    binding = result.bundle_bindings[0]
    assert binding["kind"] == "table"
    assert binding["table_index"] == 4
    assert binding["row_index"] == 2
    assert binding["row_evidence_ids"] == [
        "evidence-threshold",
        "evidence-action",
    ]
    assert binding["header_evidence_ids"] == [
        "evidence-header-condition",
        "evidence-header-action",
    ]


def test_structural_repair_is_deterministic_and_idempotent() -> None:
    packet = _table_packet()

    first = repair_protocol_structural_bundles(["evidence-action"], packet)
    second = repair_protocol_structural_bundles(["evidence-action"], packet)

    assert first == second
    assert first.to_lineage() == second.to_lineage()
    reclosed = repair_protocol_structural_bundles(
        list(first.expanded_evidence_ids),
        packet,
    )
    assert reclosed.expanded_evidence_ids == first.expanded_evidence_ids
    assert reclosed.added_structural_context_ids == ()
    assert reclosed.bundle_bindings == first.bundle_bindings


def test_structural_repair_rejects_header_without_data_row() -> None:
    packet = _table_packet()

    with pytest.raises(ProtocolStructuralRepairError, match="header-only"):
        repair_protocol_structural_bundles(
            ["evidence-header-condition"],
            packet,
        )


def test_structural_repair_rejects_cross_row_union() -> None:
    packet = _table_packet(
        rows=(
            (
                2,
                (
                    ("evidence-row2-threshold", "ALT>3×ULN。", 0),
                    ("evidence-row2-action", "暂停研究药物。", 1),
                ),
            ),
            (
                3,
                (
                    ("evidence-row3-threshold", "ALT>5×ULN。", 0),
                    ("evidence-row3-action", "永久停药。", 1),
                ),
            ),
        ),
    )

    with pytest.raises(ProtocolStructuralRepairError, match="across rows"):
        repair_protocol_structural_bundles(
            ["evidence-row2-action", "evidence-row3-threshold"],
            packet,
        )


def test_structural_repair_rejects_different_tables() -> None:
    packet = [
        *_table_packet(table_index=4),
        *_table_packet(table_index=5, id_prefix="t5-"),
    ]

    with pytest.raises(ProtocolStructuralRepairError, match="across tables"):
        repair_protocol_structural_bundles(
            ["evidence-action", "t5-evidence-threshold"],
            packet,
        )


def test_structural_repair_rejects_ambiguous_table_identity() -> None:
    packet = [
        _payload_evidence(
            "evidence-cell-untyped",
            "docx:table:4:row:2:cell:0:paragraph:3",
            "缺少表索引类型的单元格。",
            ("table_row_context",),
            {"kind": "table_cell", "table_index": None, "row_index": 2},
        )
    ]

    with pytest.raises(ProtocolStructuralRepairError, match="typed table identity"):
        repair_protocol_structural_bundles(["evidence-cell-untyped"], packet)


@pytest.mark.parametrize("field_name", ["table_index", "row_index"])
def test_structural_repair_rejects_boolean_table_indices(field_name: str) -> None:
    packet = _table_packet()
    action = next(
        item for item in packet if item["evidence_id"] == "evidence-action"
    )
    action["raw_fields"]["protocol_context"]["structure"][field_name] = True

    with pytest.raises(ProtocolStructuralRepairError, match="typed table identity"):
        repair_protocol_structural_bundles(["evidence-action"], packet)


def _list_packet(
    title_id: str = "evidence-list-title",
    item_ids: tuple[str, ...] = (
        "evidence-item-1",
        "evidence-item-2",
    ),
    parents: tuple[str, ...] = ("matched-1",),
    bundle_id: str | None = None,
) -> list[dict[str, Any]]:
    # The stable bundle identity is the exact ancestor-title source
    # identity; it never derives from the parent-match lineage.
    stable_bundle_id = bundle_id or f"src-{title_id}"
    packet = [
        _payload_evidence(
            title_id,
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=parents,
            bundle_id=stable_bundle_id,
        )
    ]
    for index, item_id in enumerate(item_ids, start=31):
        packet.append(
            _payload_evidence(
                item_id,
                f"docx:paragraph:{index}",
                f"（{index - 30}）条目内容；",
                ("list_item",),
                {"kind": "paragraph", "paragraph_index": index},
                parents=parents,
                bundle_id=stable_bundle_id,
            )
        )
    return packet


def test_structural_repair_item_adds_unique_title_not_siblings() -> None:
    packet = _list_packet()

    result = repair_protocol_structural_bundles(
        ["evidence-item-2"],
        packet,
    )

    assert result.original_evidence_ids == ("evidence-item-2",)
    assert result.added_structural_context_ids == ("evidence-list-title",)
    assert result.expanded_evidence_ids == (
        "evidence-list-title",
        "evidence-item-2",
    )
    binding = result.bundle_bindings[0]
    assert binding["kind"] == "list"
    assert binding["ancestor_title_evidence_id"] == "evidence-list-title"
    assert binding["list_bundle_id"] == "src-evidence-list-title"
    assert binding["item_evidence_ids"] == [
        "evidence-item-1",
        "evidence-item-2",
    ]


def test_structural_repair_title_alone_adds_no_items() -> None:
    packet = _list_packet()

    result = repair_protocol_structural_bundles(
        ["evidence-list-title"],
        packet,
    )

    assert result.expanded_evidence_ids == ("evidence-list-title",)
    assert result.added_structural_context_ids == ()
    assert result.bundle_bindings[0]["kind"] == "list"


def test_structural_repair_rejects_missing_or_ambiguous_list_ancestry() -> None:
    orphaned = [
        _payload_evidence(
            "evidence-item-orphan",
            "docx:paragraph:31",
            "（1）无标题的孤儿条目；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-0",),
            bundle_id="src-title-orphan",
        )
    ]
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="unique list ancestor title",
    ):
        repair_protocol_structural_bundles(["evidence-item-orphan"], orphaned)

    ambiguous = [
        _payload_evidence(
            "evidence-title-a",
            "docx:paragraph:30",
            "允许使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
        _payload_evidence(
            "evidence-title-b",
            "docx:paragraph:31",
            "禁止使用以下治疗：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:32",
            "（1）双重标题下的条目；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 32},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
    ]
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="unique list ancestor title",
    ):
        repair_protocol_structural_bundles(["evidence-item-1"], ambiguous)


def test_structural_repair_rejects_empty_list_ancestry() -> None:
    packet = [
        _payload_evidence(
            "evidence-title-empty",
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
        ),
        _payload_evidence(
            "evidence-item-empty",
            "docx:paragraph:31",
            "（1）无父身份的条目；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 31},
        ),
    ]

    with pytest.raises(
        ProtocolStructuralRepairError,
        match="typed list identity",
    ):
        repair_protocol_structural_bundles(["evidence-item-empty"], packet)
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="typed list identity",
    ):
        repair_protocol_structural_bundles(["evidence-title-empty"], packet)


def test_protocol_evidence_focus_omits_item_with_empty_ancestry() -> None:
    packet = [
        _payload_evidence(
            "evidence-list-title",
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:31",
            "（1）其他试验用药；",
            ("primary_match", "list_item"),
            {"kind": "paragraph", "paragraph_index": 31},
        ),
    ]

    focused = focus_protocol_provider_evidence(packet)

    assert [item["evidence_id"] for item in focused] == []


def test_structural_repair_rejects_title_from_other_list() -> None:
    packet = [
        *_list_packet(
            title_id="evidence-title-a",
            item_ids=("evidence-item-a",),
            parents=("matched-a",),
        ),
        *_list_packet(
            title_id="evidence-title-b",
            item_ids=("evidence-item-b",),
            parents=("matched-b",),
        ),
    ]

    with pytest.raises(
        ProtocolStructuralRepairError,
        match="from another list",
    ):
        repair_protocol_structural_bundles(
            ["evidence-title-a", "evidence-item-b"],
            packet,
        )


def test_structural_repair_rejects_cross_list_item_union() -> None:
    packet = [
        *_list_packet(
            title_id="evidence-title-a",
            item_ids=("evidence-item-a",),
            parents=("matched-a",),
        ),
        *_list_packet(
            title_id="evidence-title-b",
            item_ids=("evidence-item-b",),
            parents=("matched-b",),
        ),
    ]

    with pytest.raises(ProtocolStructuralRepairError, match="across lists"):
        repair_protocol_structural_bundles(
            ["evidence-item-a", "evidence-item-b"],
            packet,
        )


def test_structural_repair_rejects_out_of_packet_evidence() -> None:
    packet = _list_packet()

    with pytest.raises(ProtocolStructuralRepairError, match="frozen packet"):
        repair_protocol_structural_bundles(["evidence-not-in-packet"], packet)


def test_structural_repair_rejects_expansion_over_limit() -> None:
    row_cells = tuple(
        (f"evidence-cell-{index}", f"单元格 {index}。", 0)
        for index in range(52)
    )
    packet = _table_packet(rows=((2, row_cells),))

    with pytest.raises(ProtocolStructuralRepairError, match="exceed 50"):
        repair_protocol_structural_bundles(["evidence-cell-0"], packet)


def test_structural_repair_plain_paragraph_evidence_is_passthrough() -> None:
    packet = [
        _payload_evidence(
            "evidence-paragraph",
            "docx:paragraph:10",
            "普通段落条款。",
            ("primary_match",),
            {"kind": "paragraph", "paragraph_index": 10},
        )
    ]

    result = repair_protocol_structural_bundles(
        ["evidence-paragraph"],
        packet,
    )

    assert result.expanded_evidence_ids == ("evidence-paragraph",)
    assert result.added_structural_context_ids == ()
    assert result.bundle_bindings == ()


def test_v3_list_bundle_binds_title_and_items_across_different_parent_windows() -> None:
    """N1/N3: one physical list keeps one identity even when adjacent members
    carry different causal keyword-match windows; a title and its item bind
    even though their ``parent_match_source_ids`` differ.
    """

    packet = [
        _payload_evidence(
            "evidence-list-title",
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=("matched-1",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:31",
            "（1）其他试验用药；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-1", "matched-2"),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-2",
            "docx:paragraph:32",
            "（2）免疫抑制剂；",
            ("list_item",),
            {"kind": "paragraph", "paragraph_index": 32},
            parents=("matched-1", "matched-3", "matched-4"),
            bundle_id="src-list-title",
        ),
    ]

    result = repair_protocol_structural_bundles(
        ["evidence-item-2"],
        packet,
    )

    assert result.expanded_evidence_ids == (
        "evidence-list-title",
        "evidence-item-2",
    )
    binding = result.bundle_bindings[0]
    assert binding["list_bundle_id"] == "src-list-title"
    assert binding["ancestor_title_evidence_id"] == "evidence-list-title"
    assert binding["item_evidence_ids"] == [
        "evidence-item-1",
        "evidence-item-2",
    ]


def test_v3_list_bundle_identity_ignores_unrelated_keyword_hits() -> None:
    """N2: adding an unrelated keyword hit can never change the stable
    bundle identity of an existing list.
    """

    def packet_with_parents(parents: tuple[str, ...]) -> list[dict[str, Any]]:
        return _list_packet(parents=parents, bundle_id="src-list-title")

    base = repair_protocol_structural_bundles(
        ["evidence-item-2"],
        packet_with_parents(("matched-1",)),
    )
    unrelated = repair_protocol_structural_bundles(
        ["evidence-item-2"],
        packet_with_parents(("matched-1", "unrelated-hit")),
    )

    assert base.bundle_bindings[0]["list_bundle_id"] == "src-list-title"
    assert unrelated.bundle_bindings[0]["list_bundle_id"] == "src-list-title"
    assert unrelated.bundle_bindings[0] == base.bundle_bindings[0]


def test_v3_identical_parent_sets_keep_physical_lists_separate() -> None:
    """N4: genuinely different physical lists with identical match-parent
    sets must stay separate bundles and reject a cross-list union.
    """

    packet = [
        *_list_packet(
            title_id="evidence-title-a",
            item_ids=("evidence-item-a",),
            parents=("matched-9",),
            bundle_id="src-title-a",
        ),
        *_list_packet(
            title_id="evidence-title-b",
            item_ids=("evidence-item-b",),
            parents=("matched-9",),
            bundle_id="src-title-b",
        ),
    ]

    with pytest.raises(ProtocolStructuralRepairError, match="across lists"):
        repair_protocol_structural_bundles(
            ["evidence-item-a", "evidence-item-b"],
            packet,
        )
    first = repair_protocol_structural_bundles(["evidence-item-a"], packet)
    second = repair_protocol_structural_bundles(["evidence-item-b"], packet)
    assert first.bundle_bindings[0]["list_bundle_id"] == "src-title-a"
    assert second.bundle_bindings[0]["list_bundle_id"] == "src-title-b"


def test_v3_heading_body_pair_is_not_two_list_items() -> None:
    """N5: a colon-ending heading and its unmarked body (paragraphs
    1171-1172 shape) must not become two list items; explicit item syntax
    after a later title still binds to that title only.
    """

    spans = [
        _shape_span(
            "unscheduled-heading",
            "docx:paragraph:1171",
            "计划外访视：",
        ),
        _shape_span(
            "unscheduled-body",
            "docx:paragraph:1172",
            "研究者应根据受试者情况安排计划外访视。",
        ),
        _shape_span(
            "unscheduled-title",
            "docx:paragraph:1175",
            "计划外访视安排如下：",
        ),
        _shape_span(
            "unscheduled-item",
            "docx:paragraph:1176",
            "（1）计划外访视应在24小时内安排；",
        ),
        _shape_span(
            "after-list",
            "docx:paragraph:1177",
            "计划外访视记录应归档。",
        ),
    ]
    matches = [
        {
            "source_id": "unscheduled-heading",
            "source_entry_id": "shape-entry",
            "locator": spans[0].locator,
            "text": spans[0].text_preview,
            "score": 1001,
            "matched_keywords": ["计划外访视"],
            "match_reason": "命中关键词：计划外访视",
        },
        {
            "source_id": "unscheduled-title",
            "source_entry_id": "shape-entry",
            "locator": spans[2].locator,
            "text": spans[2].text_preview,
            "score": 1001,
            "matched_keywords": ["计划外访视"],
            "match_reason": "命中关键词：计划外访视",
        },
    ]

    expanded = expand_protocol_evidence_spans(spans, matches)
    by_id = {item["source_id"]: item for item in expanded}

    assert "list_title" in by_id["unscheduled-heading"]["evidence_context"][
        "roles"
    ]
    assert by_id["unscheduled-heading"]["evidence_context"][
        "list_bundle_id"
    ] == "unscheduled-heading"
    assert "list_item" not in by_id["unscheduled-body"]["evidence_context"][
        "roles"
    ]
    assert by_id["unscheduled-body"]["evidence_context"].get(
        "list_bundle_id"
    ) is None
    assert "list_item" in by_id["unscheduled-item"]["evidence_context"][
        "roles"
    ]
    assert by_id["unscheduled-item"]["evidence_context"]["list_bundle_id"] == (
        "unscheduled-title"
    )
    assert "list_item" not in by_id["after-list"]["evidence_context"]["roles"]
    assert by_id["after-list"]["evidence_context"].get(
        "list_bundle_id"
    ) is None


def test_v3_promis_prose_cannot_be_list_title() -> None:
    """N6: term-based prose ending in ``。`` (PROMIS paragraph 486 shape)
    must never become a list ancestor title.
    """

    spans = [
        _shape_span(
            "promis-description",
            "docx:paragraph:486",
            "受试者使用患者报告结局测量信息系统（PROMIS）完成评估，"
            "包括疲劳、疼痛等维度。",
        ),
        _shape_span(
            "first-dose-rule",
            "docx:paragraph:487",
            "访视当天应完成首次给药。",
        ),
    ]
    matches = [
        {
            "source_id": "first-dose-rule",
            "source_entry_id": "shape-entry",
            "locator": spans[1].locator,
            "text": spans[1].text_preview,
            "score": 1001,
            "matched_keywords": ["首次给药"],
            "match_reason": "命中关键词：首次给药",
        }
    ]

    expanded = expand_protocol_evidence_spans(spans, matches)
    by_id = {item["source_id"]: item for item in expanded}

    promis_roles = set(by_id["promis-description"]["evidence_context"]["roles"])
    assert "list_title" not in promis_roles
    assert "list_item" not in promis_roles
    assert by_id["promis-description"]["evidence_context"].get(
        "list_bundle_id"
    ) is None


def test_v3_focus_keeps_stable_ancestor_and_orphan_gets_no_ancestor() -> None:
    """N7: provider focus must never retain a selected list item while
    omitting its stable ancestor; a primary list item without a stable
    bundle, or without exactly one typed ancestor title in its bundle, is
    omitted from the provider-focused packet.
    """

    packet = [
        _payload_evidence(
            "evidence-list-title",
            "docx:paragraph:30",
            "禁止使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 30},
            parents=("matched-1",),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-1",
            "docx:paragraph:31",
            "（1）其他试验用药；",
            ("primary_match", "list_item"),
            {"kind": "paragraph", "paragraph_index": 31},
            parents=("matched-1", "matched-5"),
            bundle_id="src-list-title",
        ),
        _payload_evidence(
            "evidence-item-orphan",
            "docx:paragraph:32",
            "（2）无稳定身份的条目；",
            ("primary_match", "list_item"),
            {"kind": "paragraph", "paragraph_index": 32},
            parents=("matched-5",),
        ),
    ]

    focused = focus_protocol_provider_evidence(packet)
    focused_ids = [item["evidence_id"] for item in focused]

    assert focused_ids == ["evidence-list-title", "evidence-item-1"]
    assert "evidence-item-orphan" not in focused_ids
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="typed list identity",
    ):
        repair_protocol_structural_bundles(["evidence-item-orphan"], packet)

    ambiguous = [
        _payload_evidence(
            "evidence-title-a",
            "docx:paragraph:40",
            "允许使用以下药物：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 40},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
        _payload_evidence(
            "evidence-title-b",
            "docx:paragraph:41",
            "禁止使用以下治疗：",
            ("list_title",),
            {"kind": "paragraph", "paragraph_index": 41},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
        _payload_evidence(
            "evidence-item-ambiguous",
            "docx:paragraph:42",
            "（1）双重标题下的条目；",
            ("primary_match", "list_item"),
            {"kind": "paragraph", "paragraph_index": 42},
            parents=("matched-3",),
            bundle_id="src-double-title",
        ),
    ]

    focused = focus_protocol_provider_evidence(ambiguous)
    assert [item["evidence_id"] for item in focused] == []
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="unique list ancestor title",
    ):
        repair_protocol_structural_bundles(
            ["evidence-item-ambiguous"],
            ambiguous,
        )


def test_v3_genuine_cross_list_selection_fails_closed() -> None:
    """N8: paragraph 1169 (withdrawal block) plus paragraph 1405 (PK list)
    must fail as a genuinely cross-list selection.
    """

    packet = [
        *_list_packet(
            title_id="evidence-title-withdrawal",
            item_ids=("evidence-item-1169",),
            parents=("matched-1164", "matched-1169", "matched-1170", "matched-1171"),
            bundle_id="src-title-withdrawal",
        ),
        *_list_packet(
            title_id="evidence-title-pk",
            item_ids=("evidence-item-1405",),
            parents=("matched-1400", "matched-1405"),
            bundle_id="src-title-pk",
        ),
    ]

    with pytest.raises(ProtocolStructuralRepairError, match="across lists"):
        repair_protocol_structural_bundles(
            ["evidence-item-1169", "evidence-item-1405"],
            packet,
        )


def test_v3_ambiguous_multi_bundle_membership_fails_closed() -> None:
    """N4/N8 corollary: a member bound to more than one ancestor title at
    expansion time gets no stable identity and fails closed when cited.
    """

    spans = [
        _shape_span(
            "parent-title",
            "docx:paragraph:10",
            "禁止使用以下药物：",
        ),
        _shape_span(
            "nested-title",
            "docx:paragraph:11",
            "（1）允许使用以下治疗：",
        ),
        _shape_span(
            "nested-item",
            "docx:paragraph:12",
            "（2）外用糖皮质激素；",
        ),
    ]
    matches = [
        {
            "source_id": "parent-title",
            "source_entry_id": "shape-entry",
            "locator": spans[0].locator,
            "text": spans[0].text_preview,
            "score": 1001,
            "matched_keywords": ["禁止使用"],
            "match_reason": "命中关键词：禁止使用",
        },
        {
            "source_id": "nested-title",
            "source_entry_id": "shape-entry",
            "locator": spans[1].locator,
            "text": spans[1].text_preview,
            "score": 1001,
            "matched_keywords": ["允许使用"],
            "match_reason": "命中关键词：允许使用",
        },
    ]

    expanded = expand_protocol_evidence_spans(spans, matches)
    by_id = {item["source_id"]: item for item in expanded}

    # The nested title is simultaneously an item of the parent list and the
    # title of its own list: ambiguous membership strips its bundle identity.
    assert by_id["nested-title"]["evidence_context"].get(
        "list_bundle_id"
    ) is None
    packet = [
        _payload_evidence(
            f"evidence-{item['source_id']}",
            item["locator"],
            item["text"],
            tuple(item["evidence_context"]["roles"]),
            item["evidence_context"]["structure"],
            parents=tuple(
                item["evidence_context"]["parent_match_source_ids"]
            ),
            bundle_id=str(
                item["evidence_context"].get("list_bundle_id") or ""
            )
            or None,
        )
        for item in expanded
    ]
    with pytest.raises(
        ProtocolStructuralRepairError,
        match="typed list identity",
    ):
        repair_protocol_structural_bundles(["evidence-nested-title"], packet)
