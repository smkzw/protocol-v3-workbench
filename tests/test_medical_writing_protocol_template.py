from __future__ import annotations

import io
import copy
import unittest
import zipfile
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from xml.etree import ElementTree

from docx import Document

from packages.contracts.workbench_contracts import (
    MedicalWritingGreenfieldCreateRequest,
    MedicalWritingGreenfieldDecision,
    MedicalWritingGreenfieldSectionSeed,
    MedicalWritingPhase1Part,
    MedicalWritingProtocolModuleResolution,
    MedicalWritingProtocolModuleResolutionApplyRequest,
    MedicalWritingStructuredStudyDesign,
    MedicalWritingStudyFactState,
)
from services.api.app.medical_writing_greenfield import (
    GreenfieldMedicalWritingConflictError,
    GreenfieldMedicalWritingDocumentService,
)
from services.api.app.medical_writing_protocol_template import (
    DRAFT_TEMPLATE_VERSION,
    M11_TEMPLATE_ID,
    M11_TEMPLATE_VERSION,
    TEMPLATE_ID,
    TEMPLATE_VERSION,
    MedicalWritingProtocolTemplateService,
)
from services.api.app.medical_writing_style_profile import (
    MedicalWritingStyleProfileService,
)
from services.api.app.medical_writing_company_corpus import (
    MedicalWritingCompanyCorpusService,
)
from services.api.app.medical_writing_document_exporter import (
    export_medical_writing_document_docx,
)


class MedicalWritingProtocolTemplateTests(unittest.TestCase):
    def setUp(self):
        self.service = MedicalWritingProtocolTemplateService()

    def test_company_template_is_default_and_traces_to_supplied_authorities(self):
        template = self.service.definition()
        self.assertEqual(TEMPLATE_ID, template.template_id)
        self.assertEqual(TEMPLATE_VERSION, template.template_version)
        self.assertEqual(130, len(template.nodes))
        self.assertEqual("final", template.lifecycle_status)
        self.assertEqual("effective", template.china_regulatory_status)
        self.assertEqual("2026-07-19", template.effective_date)
        self.assertIn("公司权威方案优先", template.china_applicability_rule)
        self.assertEqual(7, len(template.source_documents))
        self.assertEqual(
            "CMS-D017-PNH-方案摘要_v0.2.docx", template.source_documents[0]["title"]
        )
        self.assertEqual("cms_front_matter", template.nodes[0].node_id)
        by_semantic = {node.semantic_node_id: node for node in template.nodes}
        self.assertEqual("protocol_synopsis", by_semantic["synopsis.summary"].node_kind)
        self.assertEqual("研究示意图", by_semantic["synopsis.schema"].title_zh)
        self.assertEqual("研究治疗/干预", by_semantic["intervention"].title_zh)
        self.assertEqual(
            [
                "cms_d017_pnh_synopsis_v0_2",
                "cms_d005_obesity_synopsis_v0_3",
                "my004_ra_synopsis_v0_3",
                "my004_dermatology_synopsis_v0_4",
            ],
            by_semantic["synopsis.summary"].authority_source_ids,
        )
        self.assertEqual(["1.1"], by_semantic["synopsis.summary"].m11_coverage_anchors)
        self.assertNotEqual(
            by_semantic["intervention.dose_modification"].interaction_types[0],
            by_semantic["intervention.concomitant"].interaction_types[0],
        )

    def test_current_chinese_m11_tree_remains_queryable_for_replay(self):
        template = self.service.definition(M11_TEMPLATE_ID, M11_TEMPLATE_VERSION)
        self.assertEqual(160, len(template.nodes))
        self.assertEqual("public_consultation", template.china_regulatory_status)
        by_number = {node.section_number: node for node in template.nodes}
        self.assertEqual("临床试验干预描述", by_number["6.1"].title_zh)
        self.assertEqual("附加附录", by_number["12.X"].title_zh)
        self.assertEqual("ich_m11_12", by_number["12.X"].parent_node_id)

    def test_historical_draft_remains_queryable_without_becoming_default(self):
        current = self.service.definition()
        draft = self.service.definition(M11_TEMPLATE_ID, DRAFT_TEMPLATE_VERSION)
        self.assertEqual("draft", draft.lifecycle_status)
        self.assertEqual("historical_draft", draft.china_regulatory_status)
        self.assertEqual(159, len(draft.nodes))
        self.assertNotEqual(current.definition_sha256, draft.definition_sha256)
        draft_by_number = {node.section_number: node for node in draft.nodes}
        self.assertEqual("试验方案", draft_by_number["1.2"].title_zh)
        self.assertEqual("研究性干预描述", draft_by_number["6.1"].title_zh)
        self.assertNotIn("12.X", draft_by_number)

    def test_unknown_template_version_fails_closed(self):
        with self.assertRaises(KeyError):
            self.service.definition(TEMPLATE_ID, "unregistered")

    def test_server_projects_synopsis_and_design_into_canonical_nodes(self):
        definition = self._definition()
        seeds = self.service.section_seeds(definition)
        synopsis = next(
            item for item in seeds if item.section_key == "cms_synopsis_summary"
        )
        design = next(
            item for item in seeds if item.section_key == "cms_study_design_overall"
        )
        self.assertEqual("结构化研究摘要", synopsis.initial_text)
        self.assertEqual(5, len(synopsis.source_fact_ids))
        self.assertEqual(
            "cms_protocol_synopsis_v1", synopsis.initial_data["schema_version"]
        )
        self.assertEqual(18, len(synopsis.initial_data["rows"]))
        self.assertEqual(
            ["项目", "目的/内容", "相应的研究终点"],
            synopsis.initial_data["columns"],
        )
        self.assertIn("随机化", design.initial_text)
        self.assertIn("双盲", design.initial_text)
        self.assertIn("安慰剂对照", design.initial_text)
        self.assertTrue(
            all(item.template_node_id == item.section_key for item in seeds)
        )
        by_key = {item.section_key: item for item in seeds}
        self.assertEqual(
            "主要目的和主要终点", by_key["cms_objectives_endpoints_primary"].heading
        )
        self.assertEqual(
            "次要目的和次要终点", by_key["cms_objectives_endpoints_secondary"].heading
        )
        self.assertEqual("研究流程表", by_key["cms_synopsis_schedule"].heading)

    def test_applicable_sections_have_typed_draft_or_actionable_blocker(self):
        definitions = []
        for phase in ("I期", "II期", "III期"):
            definition = copy.deepcopy(self._definition())
            definition.framing.study_phase = phase
            definitions.append(definition)

        for definition in definitions:
            with self.subTest(phase=definition.framing.study_phase):
                seeds = self.service.section_seeds(definition)
                parent_ids = {
                    seed.parent_key for seed in seeds if seed.parent_key
                }
                applicable = [
                    seed
                    for seed in seeds
                    if seed.applicability_status == "applicable"
                ]
                self.assertTrue(applicable)
                self.assertNotIn(
                    "unclassified",
                    {seed.drafting_status for seed in applicable},
                )
                blockers = [
                    seed
                    for seed in applicable
                    if seed.drafting_status == "actionable_blocker"
                ]
                self.assertTrue(blockers)
                for seed in blockers:
                    self.assertEqual("", seed.initial_text)
                    self.assertTrue(seed.drafting_blocker_code)
                    self.assertTrue(seed.drafting_blocker_reason)
                    self.assertTrue(seed.drafting_missing_inputs)
                    self.assertTrue(seed.drafting_resolution_actions)
                for seed in applicable:
                    if seed.drafting_status == "structural_container":
                        self.assertIn(seed.section_key, parent_ids)

    def test_typed_blocker_survives_materialization_without_polluting_body(self):
        definition = self._definition()
        seeds = self.service.section_seeds(definition)
        blocked_seed = next(
            seed
            for seed in seeds
            if seed.drafting_status == "actionable_blocker"
        )
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-READINESS",
            version="V0.1",
            document_title="起草就绪状态测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="typed-readiness-materialization",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            first = store.create(
                "project_typed_readiness",
                request,
                resolved_sections=seeds,
            )
            second = store.create(
                "project_typed_readiness",
                request,
                resolved_sections=seeds,
            )
            revision_document = store.document_for_revision(
                "project_typed_readiness"
            )
        self.assertEqual(
            first.document.model_dump(mode="json"),
            second.document.model_dump(mode="json"),
        )
        materialized = next(
            section
            for section in revision_document.sections
            if section.template_node_id == blocked_seed.template_node_id
        )
        self.assertEqual("actionable_blocker", materialized.drafting_status)
        self.assertEqual("blocked_missing_inputs", materialized.completion_status)
        self.assertEqual(
            blocked_seed.drafting_blocker_code,
            materialized.drafting_blocker_code,
        )
        self.assertEqual(2, len(materialized.content_blocks))
        self.assertEqual("", materialized.content_blocks[1]["text"])

    def test_create_idempotency_hash_includes_server_resolved_sections(self):
        definition = self._definition()
        seeds = self.service.section_seeds(definition)
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-SEMANTIC-HASH",
            version="V0.1",
            document_title="服务端解析内容哈希测试",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="server-resolved-sections-hash",
        )
        changed_seeds = list(seeds)
        changed_seeds[0] = changed_seeds[0].model_copy(
            update={"heading": f"{changed_seeds[0].heading}（变更）"},
            deep=True,
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            store.create(
                "project_semantic_hash",
                request,
                resolved_sections=seeds,
            )
            with self.assertRaisesRegex(
                GreenfieldMedicalWritingConflictError,
                "baseline already exists",
            ):
                store.create(
                    "project_semantic_hash",
                    request,
                    resolved_sections=changed_seeds,
                )

    def test_unknown_modules_create_one_document_level_readiness_blocker(self):
        definition = self._definition()
        sections = self.service.section_seeds(definition)
        resolutions = self.service.module_resolutions(definition)
        unresolved = [
            item
            for item in resolutions
            if item.status in {"unknown", "deferred"}
        ]
        self.assertTrue(unresolved)
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-MODULE-READINESS",
            version="V0.1",
            document_title="动态章节就绪测试",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="module-readiness-create",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            created = store.create(
                "project_module_readiness",
                request,
                resolved_sections=sections,
                resolved_module_resolutions=resolutions,
            )
            blockers = store.approval_blockers("project_module_readiness")
            state = store.baseline_state("project_module_readiness")
        module_gate = next(
            gate
            for gate in created.document.quality_gates
            if gate["label"] == "动态章节适用性"
        )
        module_blockers = [
            blocker
            for blocker in blockers
            if blocker.blocker_type
            == "greenfield_unresolved_module_applicability"
        ]
        self.assertEqual("blocked", module_gate["status"])
        self.assertIn(str(len(unresolved)), module_gate["detail"])
        self.assertEqual(1, len(module_blockers))
        self.assertEqual(len(unresolved), state["unresolved_module_count"])
        self.assertEqual(1, state["approval_blocker_count"])

    def test_company_synopsis_materializes_the_p0_three_column_contract(self):
        definition = self._definition()
        definition.picos.primary_objectives = [
            "保持来源限定条件完整的主要目的。"
        ]
        definition.picos.secondary_objectives = [
            "保持来源顺序的第一次要目的。",
            "保持来源顺序的第二次要目的。",
        ]
        definition.picos.exploratory_objectives = [
            "保持来源顺序的第一探索性目的。",
            "保持来源顺序的第二探索性目的。",
            "保持来源顺序的第三探索性目的。",
        ]
        definition.picos.key_secondary_endpoints = [
            "次要有效性终点：\n第12周症状评分较基线的变化值。"
        ]
        definition.picos.safety_endpoints = ["安全性终点：\nAE、SAE和AESI的发生率。"]
        definition.picos.exploratory_endpoints = [
            "探索性终点：\n研究药物的PK特征。\n研究药物的PD特征。"
        ]
        definition.framing.product_profile.pk_pd_considerations = ["暴露-效应分析"]
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-001",
            version="V0.1",
            document_title="测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="company-synopsis-p0-contract",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            store.create(
                "project_company_synopsis",
                request,
                resolved_sections=self.service.section_seeds(definition),
            )
            synopsis = next(
                section
                for section in store.document_for_revision(
                    "project_company_synopsis"
                ).sections
                if section.template_node_id == "cms_synopsis_summary"
            )
            by_template_node = {
                section.template_node_id: section
                for section in store.document_for_revision(
                    "project_company_synopsis"
                ).sections
            }
            exported = export_medical_writing_document_docx(
                store.document_for_revision("project_company_synopsis"),
                mode="draft_preview",
            )
        table = synopsis.content_blocks[1]["structured_table"]
        self.assertTrue(
            by_template_node["cms_confidentiality"].content_blocks[0][
                "page_break_before"
            ]
        )
        self.assertTrue(
            by_template_node["cms_signatures"].content_blocks[0]["page_break_before"]
        )
        self.assertTrue(
            by_template_node["cms_version_history"].content_blocks[0][
                "page_break_before"
            ]
        )
        self.assertTrue(
            by_template_node["cms_synopsis_summary"].content_blocks[0][
                "page_break_before"
            ]
        )
        self.assertFalse(
            by_template_node["cms_indexes"].content_blocks[0]["page_break_before"]
        )
        signature_table = by_template_node["cms_signatures"].content_blocks[1][
            "structured_table"
        ]
        self.assertEqual("document_control", signature_table["role"])
        self.assertEqual(
            ["签署角色", "姓名/职务", "签名", "日期"],
            [column["label"] for column in signature_table["columns"]],
        )
        self.assertEqual(
            ["申办者代表", "主要研究者"],
            [row["cells"][0]["text"] for row in signature_table["rows"][1:]],
        )
        version_history = by_template_node["cms_version_history"].content_blocks[1][
            "structured_table"
        ]
        self.assertEqual("version_history", version_history["domain"])
        self.assertEqual(
            ["V0.1", "", "", "", "", ""],
            [cell["text"] for cell in version_history["rows"][1]["cells"]],
        )
        word_document = Document(io.BytesIO(exported.content))
        self.assertNotIn(
            "保密声明",
            {
                paragraph.text.strip()
                for paragraph in word_document.paragraphs
                if paragraph.text.strip()
            },
        )
        self.assertIn(
            "保密声明：",
            "\n".join(paragraph.text for paragraph in word_document.paragraphs),
        )
        self.assertTrue(
            any(
                [cell.text for cell in word_table.rows[0].cells]
                == ["签署角色", "姓名/职务", "签名", "日期"]
                for word_table in word_document.tables
            )
        )
        self.assertTrue(
            any(
                [cell.text for cell in word_table.rows[0].cells]
                == [
                    "版本号",
                    "版本日期",
                    "变更范围",
                    "变更说明",
                    "变更理由",
                    "批准状态",
                ]
                for word_table in word_document.tables
            )
        )
        with zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
            document_root = ElementTree.fromstring(archive.read("word/document.xml"))
        word_ns = {
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        }
        body = document_root.find("w:body", word_ns)
        self.assertIsNotNone(body)
        body_children = list(body)
        toc_heading_index = next(
            index
            for index, element in enumerate(body_children)
            if element.tag.endswith("}p")
            and "".join(
                node.text or "" for node in element.findall(".//w:t", word_ns)
            ).strip()
            == "目录"
        )
        preceding_window = body_children[max(0, toc_heading_index - 2) : toc_heading_index]
        has_section_break = any(
            element.find(".//w:sectPr", word_ns) is not None
            for element in preceding_window
        )
        has_page_break = any(
            element.find('.//w:br[@w:type="page"]', word_ns) is not None
            for element in preceding_window
        )
        self.assertFalse(
            has_section_break and has_page_break,
            "the TOC must be preceded by one pagination mechanism only; "
            "a section break plus a page break creates a blank page",
        )
        self.assertEqual("protocol_synopsis", table["role"])
        self.assertEqual(3, len(table["columns"]))
        self.assertEqual(
            [1554, 2993, 4514],
            [column["width_twips"] for column in table["columns"]],
        )
        self.assertEqual(18, len(table["rows"]))
        self.assertEqual("pct", table["word_layout"]["table_width_type"])
        self.assertEqual(4999, table["word_layout"]["table_width_value"])
        self.assertEqual(
            "repeat_header_allow_row_split",
            table["word_layout"]["long_table_split_strategy"],
        )
        self.assertEqual(
            {
                str(index): height
                for index, height in {
                    3: 335,
                    4: 90,
                    5: 335,
                    6: 335,
                    7: 335,
                    8: 335,
                }.items()
            },
            {
                str(index): height
                for index, height in table["word_layout"][
                    "row_min_heights_twips"
                ].items()
            },
        )
        self.assertEqual(
            [
                "研究题目",
                "试验分期",
                "研究对象",
                "目的与估计目标/终点",
            ],
            [row["label"] for row in table["rows"][:4]],
        )
        self.assertEqual(2, table["rows"][0]["cells"][1]["column_span"])
        objective_rows = [
            row for row in table["rows"] if row["label"] == "目的与估计目标/终点"
        ]
        self.assertEqual(6, len(objective_rows))
        self.assertEqual("相应的研究终点", objective_rows[0]["cells"][2]["text"])
        self.assertEqual(6, objective_rows[0]["cells"][0]["row_span"])
        self.assertTrue(all(row["cells"][0]["hidden"] for row in objective_rows[1:]))
        self.assertEqual(
            "objectives_endpoints",
            table["word_layout"]["nested_groups"][0]["group_id"],
        )
        primary_content = objective_rows[1]["cells"][1:]
        self.assertTrue(
            all(
                block["type"] == "paragraph"
                for cell in primary_content
                for block in cell["rich_text"]["content"]
            )
        )
        self.assertTrue(all("• " not in cell["text"] for cell in primary_content))
        for row in (objective_rows[3], objective_rows[5]):
            objective_cell, endpoint_cell = row["cells"][1:]
            objective_lists = [
                block
                for block in objective_cell["rich_text"]["content"]
                if block["type"] == "orderedList"
            ]
            endpoint_lists = [
                block
                for block in endpoint_cell["rich_text"]["content"]
                if block["type"] == "orderedList"
            ]
            self.assertEqual(
                ["decimal_half_paren"],
                [block["attrs"]["numberingFormat"] for block in objective_lists],
            )
            self.assertTrue(endpoint_lists)
            self.assertEqual(
                {"decimal_fullwidth_paren"},
                {block["attrs"]["numberingFormat"] for block in endpoint_lists},
            )
            self.assertIn("1) ", objective_cell["text"])
            self.assertIn("1）", endpoint_cell["text"])
        self.assertEqual(
            2,
            objective_rows[3]["cells"][1]["text"].count(") "),
        )
        self.assertEqual(
            "保持来源限定条件完整的主要目的。",
            objective_rows[1]["cells"][1]["text"],
        )
        self.assertEqual(
            2,
            objective_rows[3]["cells"][1]["text"].count(") "),
        )
        self.assertEqual(
            3,
            objective_rows[5]["cells"][1]["text"].count(") "),
        )
        self.assertIn(
            "3) 保持来源顺序的第三探索性目的。",
            objective_rows[5]["cells"][1]["text"],
        )
        self.assertEqual(
            3,
            objective_rows[5]["cells"][1]["text"].count(") "),
        )

    def test_glossary_prefills_only_abbreviations_supported_by_project_text(self):
        definition = self._definition()
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="D017-02-001",
            version="0.1",
            document_title=(
                "在阵发性睡眠性血红蛋白尿症（PNH）患者中评价CMS-D017的"
                "有效性和安全性"
            ),
            indication="阵发性睡眠性血红蛋白尿症（PNH）",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="glossary-prefill",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            store.create(
                "project_glossary_prefill",
                request,
                resolved_sections=self.service.section_seeds(definition),
            )
            glossary = next(
                section
                for section in store.document_for_revision(
                    "project_glossary_prefill"
                ).sections
                if section.template_node_id == "cms_glossary"
            ).content_blocks[1]["structured_table"]
        rows = [
            [cell["text"] for cell in row["cells"]]
            for row in glossary["rows"][1:]
        ]
        self.assertEqual(
            [["PNH", "阵发性睡眠性血红蛋白尿症"]],
            rows,
        )
        self.assertNotIn("CMS-D017", {row[0] for row in rows})

    def test_negative_interim_analysis_language_does_not_add_a_module(self):
        definition = self._definition()
        definition.picos.statistical_strategy = (
            "本研究不设置正式期中疗效分析。可开展内部安全性和剂量评估，"
            "但不用于正式疗效假设检验。"
        )
        seeds = self.service.section_seeds(definition)
        synopsis = next(
            item for item in seeds if item.section_key == "cms_synopsis_summary"
        )
        self.assertEqual(18, len(synopsis.initial_data["rows"]))
        self.assertNotIn(
            "cms_statistics_interim",
            {item.section_key for item in seeds},
        )

    def test_post_creation_interim_resolution_propagates_and_quarantines(self):
        definition = self._definition()
        template = self.service.definition()
        initial_sections = self.service.section_seeds(definition, template)
        initial_resolutions = self.service.module_resolutions(definition, template)
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-DYNAMIC-001",
            version="V0.1",
            document_title="动态章节传播测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="dynamic-module-create",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            created = store.create(
                "project_dynamic_module",
                request,
                resolved_sections=initial_sections,
                resolved_module_resolutions=initial_resolutions,
            )
            self.assertNotIn(
                "cms_statistics_interim",
                {section.template_node_id for section in created.document.sections},
            )

            interim_on = MedicalWritingProtocolModuleResolution(
                template_node_id="cms_statistics_interim",
                semantic_node_id="statistics.interim",
                status="applicable",
                render_action="retain_full",
                resolution_source="user_override",
                rationale="医学经理确认在达到50%信息量时开展一次正式期中分析。",
                source_fact_ids=["medical_manager:interim_decision:v1"],
                user_override=True,
            )
            definition.framing.indication = "测试适应症"
            definition.field_states["framing.indication"] = (
                MedicalWritingStudyFactState(status="confirmed")
            )
            definition.module_resolutions = {"statistics.interim": interim_on}
            on_sections = self.service.section_seeds(definition, template)
            target_background = next(
                section
                for section in on_sections
                if section.template_node_id == "cms_background_disease"
            )
            self.assertEqual(
                "substantive_draft",
                target_background.drafting_status,
            )
            on_resolutions = self.service.module_resolutions(definition, template)
            applied_on = store.apply_module_resolution(
                "project_dynamic_module",
                MedicalWritingProtocolModuleResolutionApplyRequest(
                    expected_baseline_revision=created.baseline_revision,
                    expected_baseline_sha256=created.baseline_sha256,
                    semantic_node_id="statistics.interim",
                    status="applicable",
                    render_action="retain_full",
                    rationale=interim_on.rationale,
                    source_refs=interim_on.source_fact_ids,
                    actor="medical_manager",
                    idempotency_key="dynamic-module-interim-on",
                ),
                resolved_sections=on_sections,
                resolved_module_resolutions=on_resolutions,
            )
            self.assertTrue(applied_on.added_section_ids)
            after_on = store.document_for_revision("project_dynamic_module")
            self.assertIn(
                "cms_statistics_interim",
                {section.template_node_id for section in after_on.sections},
            )
            preserved_background = next(
                section
                for section in after_on.sections
                if section.template_node_id == "cms_background_disease"
            )
            self.assertEqual(
                "actionable_blocker",
                preserved_background.drafting_status,
            )
            self.assertEqual(
                "blocked_missing_inputs",
                preserved_background.completion_status,
            )
            self.assertEqual("", preserved_background.content_blocks[1]["text"])
            synopsis_on = next(
                section
                for section in after_on.sections
                if section.template_node_id == "cms_synopsis_summary"
            )
            self.assertIn(
                "期中分析",
                [
                    row["label"]
                    for row in synopsis_on.content_blocks[1]["structured_table"]["rows"]
                ],
            )

            interim_off = interim_on.model_copy(
                update={
                    "status": "not_applicable",
                    "render_action": "omit",
                    "rationale": "医学经理确认本研究不设置正式期中分析。",
                    "source_fact_ids": ["medical_manager:interim_decision:v2"],
                },
                deep=True,
            )
            definition.module_resolutions = {"statistics.interim": interim_off}
            off_sections = self.service.section_seeds(definition, template)
            off_resolutions = self.service.module_resolutions(definition, template)
            applied_off = store.apply_module_resolution(
                "project_dynamic_module",
                MedicalWritingProtocolModuleResolutionApplyRequest(
                    expected_baseline_revision=applied_on.baseline_revision,
                    expected_baseline_sha256=applied_on.baseline_sha256,
                    semantic_node_id="statistics.interim",
                    status="not_applicable",
                    render_action="omit",
                    rationale=interim_off.rationale,
                    source_refs=interim_off.source_fact_ids,
                    actor="medical_manager",
                    idempotency_key="dynamic-module-interim-off",
                ),
                resolved_sections=off_sections,
                resolved_module_resolutions=off_resolutions,
            )
            self.assertTrue(applied_off.removed_section_ids)
            self.assertTrue(
                set(applied_off.removed_section_ids).issubset(
                    applied_off.quarantined_section_ids
                )
            )
            after_off = store.document_for_revision("project_dynamic_module")
            self.assertNotIn(
                "cms_statistics_interim",
                {section.template_node_id for section in after_off.sections},
            )
            synopsis_off = next(
                section
                for section in after_off.sections
                if section.template_node_id == "cms_synopsis_summary"
            )
            self.assertNotIn(
                "期中分析",
                [
                    row["label"]
                    for row in synopsis_off.content_blocks[1]["structured_table"][
                        "rows"
                    ]
                ],
            )

    def test_phase1_explicitly_excluding_sad_does_not_materialize_sad(self):
        definition = self._dynamic_project_definition(
            definition_id="mwdef_phase1_mad_only_oral_autoimmune",
            study_phase="I期",
            document_title="口服小分子在自身免疫病受试者中的I期研究方案",
            investigational_product="CMS-P1-ORAL",
            population_intent="自身免疫病受试者",
            design_pattern="随机、双盲、安慰剂对照、多次给药剂量递增",
            intrinsic_objectives=[
                "首次患者研究",
                "明确不开展单次给药部分",
                "仅开展多次给药剂量递增研究",
            ],
            phase1_parts=[
                MedicalWritingPhase1Part(
                    part_code="mad",
                    part_label="多次给药剂量递增",
                    unresolved=False,
                )
            ],
            technology_type="small_molecule",
            administration_routes=["口服"],
            immunogenicity_relevance="not_expected",
        )

        keys = {seed.section_key for seed in self.service.section_seeds(definition)}

        self.assertNotIn("cms_study_design_phase1_sad", keys)
        self.assertIn("cms_study_design_phase1_mad", keys)

    def test_phase1_free_text_objectives_do_not_invent_sad_mad_facets(self):
        """Dual-authority guard: prose intrinsic_objectives must not drive Parts."""
        definition = self._dynamic_project_definition(
            definition_id="mwdef_phase1_prose_only_no_parts",
            study_phase="I期",
            document_title="无typed Parts的I期方案",
            investigational_product="CMS-P1-PROSE",
            population_intent="健康成人受试者",
            design_pattern="随机、双盲、安慰剂对照、剂量递增",
            intrinsic_objectives=["FIH", "SAD", "MAD", "单次给药", "多次给药"],
            phase1_parts=[],
            technology_type="small_molecule",
            administration_routes=["口服"],
            immunogenicity_relevance="not_expected",
        )
        keys = {seed.section_key for seed in self.service.section_seeds(definition)}
        self.assertNotIn("cms_study_design_phase1_sad", keys)
        self.assertNotIn("cms_study_design_phase1_mad", keys)

    def test_phase1_sad_mad_biologic_materializes_both_parts_and_immunogenicity(self):
        definition = self._dynamic_project_definition(
            definition_id="mwdef_phase1_sad_mad_biologic",
            study_phase="I期",
            document_title="注射用生物制品首次人体SAD和MAD研究方案",
            investigational_product="CMS-P1-BIO",
            population_intent="健康成人受试者及目标疾病受试者",
            design_pattern="随机、双盲、安慰剂对照、剂量递增",
            intrinsic_objectives=["FIH", "SAD", "MAD"],
            phase1_parts=[
                MedicalWritingPhase1Part(
                    part_code="sad", part_label="SAD", unresolved=False
                ),
                MedicalWritingPhase1Part(
                    part_code="mad", part_label="MAD", unresolved=False
                ),
            ],
            technology_type="other_biologic",
            administration_routes=["皮下注射"],
            immunogenicity_relevance="expected",
        )

        keys = {seed.section_key for seed in self.service.section_seeds(definition)}

        self.assertIn("cms_study_design_phase1_sad", keys)
        self.assertIn("cms_study_design_phase1_mad", keys)
        self.assertIn("cms_study_design_starting_dose", keys)
        self.assertIn("cms_procedures_assessments_immunogenicity", keys)
        self.assertIn("cms_study_design_safety_committee", keys)

    def test_phase2_randomized_open_label_has_randomization_without_global_blinding(
        self,
    ):
        definition = self._dynamic_project_definition(
            definition_id="mwdef_phase2_open_label_randomized_active_control",
            study_phase="II期",
            document_title="炎症性疾病随机开放标签II期研究方案",
            investigational_product="CMS-P2-TAB",
            population_intent="中重度炎症性疾病成人受试者",
            design_pattern="多中心、随机、开放标签、阳性药平行对照",
            intrinsic_objectives=["概念验证", "剂量探索"],
            technology_type="small_molecule",
            administration_routes=["口服"],
            immunogenicity_relevance="not_expected",
        )

        keys = {seed.section_key for seed in self.service.section_seeds(definition)}

        self.assertIn("cms_study_design_randomization", keys)
        self.assertNotIn("cms_study_design_blinding", keys)
        self.assertNotIn("cms_study_design_phase1_parts", keys)

    def test_phase3_monoclonal_antibody_with_planned_interim_analysis_propagates_modules(
        self,
    ):
        definition = self._dynamic_project_definition(
            definition_id="mwdef_phase3_mab_planned_interim",
            study_phase="III期",
            document_title="单克隆抗体确证性III期研究方案",
            investigational_product="CMS-P3-MAB",
            population_intent="活动性自身免疫病成人受试者",
            design_pattern="多中心、随机、双盲、安慰剂对照；计划开展期中分析",
            intrinsic_objectives=["确证性研究"],
            technology_type="monoclonal_antibody",
            administration_routes=["静脉输注"],
            immunogenicity_relevance="expected",
            statistical_strategy=(
                "计划在约50%的受试者完成主要终点评估后开展一次正式期中分析，"
                "并采用预设的α消耗方法控制总体I类错误。"
            ),
        )

        seeds = self.service.section_seeds(definition)
        keys = {seed.section_key for seed in seeds}
        synopsis = next(
            seed for seed in seeds if seed.section_key == "cms_synopsis_summary"
        )

        self.assertIn("cms_statistics_interim", keys)
        self.assertIn("cms_statistics_multiplicity", keys)
        self.assertNotIn("cms_statistics_subgroup", keys)
        self.assertIn("cms_procedures_assessments_immunogenicity", keys)
        self.assertIn(
            "期中分析",
            [row["label"] for row in synopsis.initial_data["rows"]],
        )

    def test_aesi_undecided_and_explicitly_none_have_distinct_observable_states(self):
        undecided = self._dynamic_project_definition(
            definition_id="mwdef_phase2_aesi_undecided",
            study_phase="II期",
            document_title="吸入制剂II期研究方案",
            investigational_product="CMS-P2-INHALE",
            population_intent="慢性呼吸系统疾病成人受试者",
            design_pattern="多中心、随机、双盲、安慰剂对照",
            intrinsic_objectives=["概念验证"],
            technology_type="small_molecule",
            administration_routes=["吸入"],
            immunogenicity_relevance="not_expected",
        )
        explicitly_none = self._dynamic_project_definition(
            definition_id="mwdef_phase2_aesi_explicit_none",
            study_phase="II期",
            document_title="外用制剂II期研究方案",
            investigational_product="CMS-P2-TOPICAL",
            population_intent="炎症性皮肤病成人受试者",
            design_pattern="多中心、随机、双盲、赋形剂对照",
            intrinsic_objectives=["剂量探索"],
            technology_type="small_molecule",
            administration_routes=["外用"],
            immunogenicity_relevance="not_expected",
        )
        explicitly_none.field_states["picos.aesi_definitions"] = (
            MedicalWritingStudyFactState(status="not_applicable")
        )

        undecided_resolution = next(
            resolution
            for resolution in self.service.module_resolutions(undecided)
            if resolution.semantic_node_id == "safety.aesi"
        )
        explicitly_none_resolution = next(
            resolution
            for resolution in self.service.module_resolutions(explicitly_none)
            if resolution.semantic_node_id == "safety.aesi"
        )

        self.assertEqual("unknown", undecided_resolution.status)
        self.assertEqual("omit", undecided_resolution.render_action)
        self.assertEqual("not_applicable", explicitly_none_resolution.status)
        self.assertEqual(
            "retain_not_applicable",
            explicitly_none_resolution.render_action,
        )
        self.assertNotEqual(
            undecided_resolution.model_dump(mode="json"),
            explicitly_none_resolution.model_dump(mode="json"),
        )

    def test_dynamic_project_section_numbers_are_unique_after_conditional_selection(
        self,
    ):
        fixtures = {
            "phase1_mad_only": self._dynamic_project_definition(
                definition_id="mwdef_numbering_phase1_mad_only",
                study_phase="I期",
                document_title="I期MAD研究方案",
                investigational_product="CMS-NUM-P1",
                population_intent="健康成人受试者",
                design_pattern="随机、双盲、安慰剂对照、多次给药剂量递增",
                intrinsic_objectives=["MAD"],
                technology_type="small_molecule",
                administration_routes=["口服"],
                immunogenicity_relevance="not_expected",
            ),
            "phase2_open_label": self._dynamic_project_definition(
                definition_id="mwdef_numbering_phase2_open_label",
                study_phase="II期",
                document_title="II期随机开放标签研究方案",
                investigational_product="CMS-NUM-P2",
                population_intent="慢性炎症性疾病成人受试者",
                design_pattern="随机、开放标签、阳性药平行对照",
                intrinsic_objectives=["概念验证"],
                technology_type="small_molecule",
                administration_routes=["口服"],
                immunogenicity_relevance="not_expected",
            ),
            "phase3_mab_interim": self._dynamic_project_definition(
                definition_id="mwdef_numbering_phase3_mab_interim",
                study_phase="III期",
                document_title="III期单抗确证性研究方案",
                investigational_product="CMS-NUM-P3",
                population_intent="活动性自身免疫病成人受试者",
                design_pattern="随机、双盲、安慰剂对照；计划开展期中分析",
                intrinsic_objectives=["确证性研究"],
                technology_type="monoclonal_antibody",
                administration_routes=["皮下注射"],
                immunogenicity_relevance="expected",
                statistical_strategy="计划开展一次正式期中分析。",
            ),
        }

        for fixture_name, definition in fixtures.items():
            with self.subTest(fixture=fixture_name):
                numbers = [
                    seed.section_number
                    for seed in self.service.section_seeds(definition)
                    if seed.section_number
                ]
                self.assertEqual(len(numbers), len(set(numbers)))

    def test_section_seeds_follow_the_explicit_registered_version(self):
        draft = self.service.definition(M11_TEMPLATE_ID, DRAFT_TEMPLATE_VERSION)
        draft_seeds = self.service.section_seeds(self._definition(), draft)
        by_number = {item.section_number: item for item in draft_seeds}
        self.assertEqual(159, len(draft_seeds))
        self.assertEqual("试验方案", by_number["1.2"].heading)
        self.assertNotIn("12.X", by_number)

    def test_confirmed_instruments_seed_efficacy_methods_without_internal_rights_log(
        self,
    ):
        definition = self._definition()
        definition.picos.assessment_instruments = [
            SimpleNamespace(
                confirmation_status="confirmed",
                acronym="DLQI",
                canonical_name_zh="皮肤病生活质量指数",
                respondent="试验参与者",
                administration_mode="访视现场填写",
                recall_period="过去7天",
                visit_labels=["基线（D1）", "第8周（D57±3天）"],
                study_purpose="评价生活质量较基线的变化",
                scoring_range="0-30",
                scoring_direction="分值越高表示生活质量受影响程度越大",
                protocol_modified=False,
            )
        ]
        definition.field_states["picos.assessment_instruments"] = (
            MedicalWritingStudyFactState(status="confirmed")
        )
        efficacy = next(
            item
            for item in self.service.section_seeds(definition)
            if item.section_key == "cms_procedures_assessments_efficacy"
        )
        self.assertIn("DLQI（皮肤病生活质量指数）", efficacy.initial_text)
        self.assertIn("基线（D1）", efficacy.initial_text)
        self.assertIn("0-30", efficacy.initial_text)
        self.assertNotIn("permission_required", efficacy.initial_text)
        self.assertNotIn("Cardiff", efficacy.initial_text)
        self.assertEqual(
            ["study_definition:mwdef_test:r3:picos.assessment_instruments"],
            efficacy.source_fact_ids,
        )

    def test_confirmed_instrument_methods_reach_docx_as_protocol_text(self):
        definition = self._definition()
        definition.picos.assessment_instruments = [
            SimpleNamespace(
                confirmation_status="confirmed",
                acronym="EASI",
                canonical_name_zh="湿疹面积和严重程度指数",
                respondent="研究者",
                administration_mode="访视现场评估",
                recall_period="",
                visit_labels=["基线（D1）", "第8周（D57±3天）"],
                study_purpose="评价特应性皮炎体征较基线的变化",
                scoring_range="0-72",
                scoring_direction="分值越高表示疾病严重程度越高",
                protocol_modified=False,
            )
        ]
        definition.field_states["picos.assessment_instruments"] = (
            MedicalWritingStudyFactState(status="confirmed")
        )
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-001",
            version="V0.1",
            document_title="测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="instrument-docx-projection",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            result = store.create(
                "project_instrument_docx",
                request,
                resolved_sections=self.service.section_seeds(definition),
            )
            self.assertEqual(
                len(self.service.section_seeds(definition)),
                len(result.document.sections),
            )
            exported = export_medical_writing_document_docx(
                store.document_for_revision("project_instrument_docx"),
                mode="draft_preview",
            )
        text = "\n".join(
            paragraph.text
            for paragraph in Document(io.BytesIO(exported.content)).paragraphs
        )
        self.assertIn("有效性评估", text)
        self.assertIn("EASI（湿疹面积和严重程度指数）", text)
        self.assertIn("第8周（D57±3天）", text)
        self.assertNotIn("permission_required", text)
        self.assertNotIn("量表题项", text)

    def test_server_template_request_rejects_client_sections_and_decisions(self):
        common = dict(
            protocol_id="P-001",
            version="V0.1",
            document_title="测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="template-request",
        )
        with self.assertRaisesRegex(ValueError, "client section seeds"):
            MedicalWritingGreenfieldCreateRequest(
                **common,
                sections=[
                    MedicalWritingGreenfieldSectionSeed(section_key="x", heading="x")
                ],
            )
        with self.assertRaisesRegex(ValueError, "client approval decisions"):
            MedicalWritingGreenfieldCreateRequest(
                **common,
                decisions=[
                    MedicalWritingGreenfieldDecision(decision_id="d", label="d")
                ],
            )

    def test_actionable_blocker_contract_fails_closed_when_incomplete(self):
        with self.assertRaisesRegex(
            ValueError,
            "actionable drafting blocker requires",
        ):
            MedicalWritingGreenfieldSectionSeed(
                section_key="blocked",
                heading="证据不足章节",
                drafting_status="actionable_blocker",
                drafting_blocker_code="project_evidence_missing",
            )
        with self.assertRaisesRegex(
            ValueError,
            "must not contain initial body text",
        ):
            MedicalWritingGreenfieldSectionSeed(
                section_key="polluted",
                heading="带正文的阻塞章节",
                initial_text="不应在阻塞状态写入的临床正文。",
                drafting_status="actionable_blocker",
                drafting_blocker_code="project_evidence_missing",
                drafting_blocker_reason="证据不足。",
                drafting_missing_inputs=["picos.primary_endpoint"],
                drafting_resolution_actions=["由系统生成候选。"],
            )
        with self.assertRaisesRegex(
            ValueError,
            "substantive drafting seed requires",
        ):
            MedicalWritingGreenfieldSectionSeed(
                section_key="empty-substantive",
                heading="空实质章节",
                drafting_status="substantive_draft",
            )

    def test_materialized_document_persists_exact_template_identity(self):
        template = self.service.definition()
        style_profile = MedicalWritingStyleProfileService().definition()
        corpus_snapshot = MedicalWritingCompanyCorpusService().definition()
        request = MedicalWritingGreenfieldCreateRequest(
            protocol_id="P-001",
            version="V0.1",
            document_title="测试方案",
            indication="测试适应症",
            study_phase="II期",
            template_id=TEMPLATE_ID,
            template_version=TEMPLATE_VERSION,
            actor="medical_manager",
            idempotency_key="materialized-template",
        )
        with TemporaryDirectory() as directory:
            store = GreenfieldMedicalWritingDocumentService(
                Path(directory) / "greenfield.sqlite3"
            )
            result = store.create(
                "project_template",
                request,
                resolved_sections=self.service.section_seeds(self._definition()),
                template_definition_sha256=template.definition_sha256,
                style_profile_id=style_profile.style_profile_id,
                style_profile_version=style_profile.style_profile_version,
                style_profile_definition_sha256=style_profile.definition_sha256,
                corpus_snapshot_id=corpus_snapshot.snapshot_id,
                corpus_snapshot_version=corpus_snapshot.snapshot_version,
                corpus_snapshot_sha256=corpus_snapshot.snapshot_sha256,
            )
            self.assertEqual(
                len(self.service.section_seeds(self._definition())),
                len(result.document.sections),
            )
            self.assertEqual(TEMPLATE_ID, result.document.template_id)
            self.assertEqual(TEMPLATE_VERSION, result.document.template_version)
            self.assertEqual(
                template.definition_sha256, result.document.template_definition_sha256
            )
            self.assertEqual(
                style_profile.style_profile_id, result.document.style_profile_id
            )
            self.assertEqual(
                style_profile.style_profile_version,
                result.document.style_profile_version,
            )
            self.assertEqual(
                style_profile.definition_sha256,
                result.document.style_profile_definition_sha256,
            )
            self.assertEqual(
                corpus_snapshot.snapshot_id, result.document.corpus_snapshot_id
            )
            self.assertEqual(
                corpus_snapshot.snapshot_version,
                result.document.corpus_snapshot_version,
            )
            self.assertEqual(
                corpus_snapshot.snapshot_sha256,
                result.document.corpus_snapshot_sha256,
            )
            self.assertNotIn(
                "<#>",
                "\n".join(
                    str(block.get("text") or "")
                    for section in result.document.sections
                    for block in section.content_blocks
                ),
            )
            self.assertEqual(
                result.document.model_dump(mode="json"),
                store.document_session("project_template").model_dump(mode="json"),
            )

    def test_phase_design_and_modality_plugins_do_not_leak_between_projects(self):
        phase1 = copy.deepcopy(self._definition())
        phase1.framing.study_phase = "I期"
        phase1.framing.design_pattern = "随机、双盲、安慰剂对照、剂量递增"
        phase1.framing.intrinsic_objectives = ["FIH", "SAD", "MAD"]
        phase1.framing.structured_design.phase1_parts = [
            MedicalWritingPhase1Part(
                part_code="sad", part_label="SAD", unresolved=False
            ),
            MedicalWritingPhase1Part(
                part_code="mad", part_label="MAD", unresolved=False
            ),
        ]
        phase1_keys = {seed.section_key for seed in self.service.section_seeds(phase1)}
        phase1_synopsis = next(
            seed
            for seed in self.service.section_seeds(phase1)
            if seed.section_key == "cms_synopsis_summary"
        )
        self.assertIn("cms_study_design_phase1_sad", phase1_keys)
        self.assertIn("cms_study_design_phase1_mad", phase1_keys)
        self.assertIn("cms_study_design_starting_dose", phase1_keys)
        self.assertNotIn("cms_statistics_multiplicity", phase1_keys)
        self.assertNotIn("cms_procedures_assessments_immunogenicity", phase1_keys)
        self.assertGreater(len(phase1_synopsis.initial_data["rows"]), 18)
        self.assertIn(
            "I期研究组成",
            [row["label"] for row in phase1_synopsis.initial_data["rows"]],
        )

        phase3 = copy.deepcopy(self._definition())
        phase3.framing.study_phase = "III期"
        phase3.framing.design_pattern = "多中心、随机、双盲、安慰剂对照"
        phase3.framing.product_profile.technology_type = "monoclonal_antibody"
        phase3.framing.product_profile.administration_routes = ["皮下注射"]
        phase3.framing.product_profile.immunogenicity_relevance = "expected"
        phase3_keys = {seed.section_key for seed in self.service.section_seeds(phase3)}
        self.assertNotIn("cms_statistics_multiplicity", phase3_keys)
        self.assertNotIn("cms_statistics_subgroup", phase3_keys)
        self.assertIn("cms_procedures_assessments_immunogenicity", phase3_keys)
        self.assertNotIn("cms_study_design_starting_dose", phase3_keys)

        phase3.picos.statistical_strategy = (
            "主要和关键次要终点采用预先规定的层级检验控制多重性，"
            "并按基线疾病严重程度开展预设亚组分析。"
        )
        phase3_confirmed_keys = {
            seed.section_key for seed in self.service.section_seeds(phase3)
        }
        self.assertIn("cms_statistics_multiplicity", phase3_confirmed_keys)
        self.assertIn("cms_statistics_subgroup", phase3_confirmed_keys)

        phase3.picos.statistical_strategy = (
            "本研究不进行多重性调整，也不开展预设亚组分析。"
        )
        phase3_excluded = {
            resolution.semantic_node_id: resolution
            for resolution in self.service.module_resolutions(phase3)
        }
        self.assertEqual(
            "not_applicable",
            phase3_excluded["statistics.multiplicity"].status,
        )
        self.assertEqual(
            "not_applicable",
            phase3_excluded["statistics.subgroup"].status,
        )

        open_label = copy.deepcopy(self._definition())
        open_label.framing.design_pattern = "多中心、随机、开放标签、平行对照"
        open_label_keys = {
            seed.section_key for seed in self.service.section_seeds(open_label)
        }
        self.assertIn("cms_study_design_randomization", open_label_keys)
        self.assertNotIn("cms_study_design_blinding", open_label_keys)

        for seeds in (
            self.service.section_seeds(phase1),
            self.service.section_seeds(phase3),
            self.service.section_seeds(open_label),
        ):
            numbers = [seed.section_number for seed in seeds if seed.section_number]
            self.assertEqual(len(numbers), len(set(numbers)))

    @staticmethod
    def _definition():
        return SimpleNamespace(
            definition_id="mwdef_test",
            revision=3,
            state_sha256="test_definition_state_sha256",
            synopsis_text="结构化研究摘要",
            framing=SimpleNamespace(
                study_phase="II期",
                document_title="测试方案",
                investigational_product="CMS-TEST",
                design_pattern="随机、双盲、安慰剂对照",
                population_intent="目标研究人群",
                intrinsic_objectives=["剂量探索"],
                structured_design=MedicalWritingStructuredStudyDesign(),
                product_profile=SimpleNamespace(
                    technology_type="small_molecule",
                    administration_routes=["口服"],
                    dosage_forms=["片剂"],
                    immunogenicity_relevance="not_expected",
                    safety_considerations=[],
                    pk_pd_considerations=[],
                    pharmacology_considerations=[],
                ),
            ),
            picos=SimpleNamespace(
                population_summary="目标研究人群",
                design_archetype="",
                inclusion_modules=[],
                exclusion_modules=[],
                intervention_summary="试验药物方案",
                intervention_dose_regimen="每日一次口服给药",
                allowed_concomitant_rules=[],
                required_background_rules=[],
                prohibited_concomitant_rules=[],
                comparator_summary="安慰剂",
                primary_endpoint="主要终点及评价时间",
                key_secondary_endpoints=[],
                other_secondary_endpoints=[],
                exploratory_endpoints=[],
                safety_endpoints=[],
                aesi_definitions=[],
                assessment_instruments=[],
                sample_size_strategy="计划入组约120例参与者。",
                statistical_strategy="采用预先规定的统计分析集进行分析。",
            ),
            field_states={
                "framing.design_pattern": MedicalWritingStudyFactState(
                    status="confirmed"
                ),
                "framing.population_intent": MedicalWritingStudyFactState(
                    status="confirmed"
                ),
                "picos.intervention_summary": MedicalWritingStudyFactState(
                    status="confirmed"
                ),
                "picos.comparator_summary": MedicalWritingStudyFactState(
                    status="confirmed"
                ),
                "picos.primary_endpoint": MedicalWritingStudyFactState(
                    status="confirmed"
                ),
            },
        )

    @classmethod
    def _dynamic_project_definition(
        cls,
        *,
        definition_id: str,
        study_phase: str,
        document_title: str,
        investigational_product: str,
        population_intent: str,
        design_pattern: str,
        intrinsic_objectives: list[str],
        technology_type: str,
        administration_routes: list[str],
        immunogenicity_relevance: str,
        statistical_strategy: str = "采用预先规定的统计分析集进行分析。",
        phase1_parts: list[MedicalWritingPhase1Part] | None = None,
    ):
        definition = copy.deepcopy(cls._definition())
        definition.definition_id = definition_id
        definition.framing.study_phase = study_phase
        definition.framing.document_title = document_title
        definition.framing.investigational_product = investigational_product
        definition.framing.population_intent = population_intent
        definition.framing.design_pattern = design_pattern
        definition.framing.intrinsic_objectives = list(intrinsic_objectives)
        if phase1_parts is not None:
            definition.framing.structured_design.phase1_parts = list(phase1_parts)
        definition.framing.product_profile.technology_type = technology_type
        definition.framing.product_profile.administration_routes = list(
            administration_routes
        )
        definition.framing.product_profile.immunogenicity_relevance = (
            immunogenicity_relevance
        )
        definition.picos.population_summary = population_intent
        definition.picos.statistical_strategy = statistical_strategy
        return definition


if __name__ == "__main__":
    unittest.main()
