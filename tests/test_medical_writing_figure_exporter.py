from __future__ import annotations

import binascii
import io
import struct
import zipfile
import zlib
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.shared import Inches

from services.api.app.medical_writing_figure_exporter import (
    ASVG_EXTENSION_URI,
    ASVG_NAMESPACE,
    MedicalWritingFigureExportError,
    add_svg_figure_with_png_fallback,
)


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
PIC_NS = "http://schemas.openxmlformats.org/drawingml/2006/picture"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
IMAGE_REL_TYPE = f"{R_NS}/image"
NS = {
    "a": A_NS,
    "asvg": ASVG_NAMESPACE,
    "pic": PIC_NS,
    "r": R_NS,
    "wp": WP_NS,
}


def test_document_target_embeds_svg_with_png_fallback_and_correct_ooxml_package():
    svg = _svg_bytes()
    png = _png_bytes(width=2, height=1)
    document = Document()

    shape = add_svg_figure_with_png_fallback(
        document,
        svg_bytes=svg,
        png_bytes=png,
        figure_id="mw-flow-primary-endpoint",
        alt_text="主要终点医学写作研究流程图",
        width=Inches(4),
    )

    assert shape.width == Inches(4)
    assert shape.height == Inches(2)

    output = io.BytesIO()
    document.save(output)
    with zipfile.ZipFile(io.BytesIO(output.getvalue())) as archive:
        names = archive.namelist()
        svg_names = [
            name
            for name in names
            if name.startswith("word/media/") and name.endswith(".svg")
        ]
        png_names = [
            name
            for name in names
            if name.startswith("word/media/") and name.endswith(".png")
        ]
        assert len(svg_names) == 1
        assert len(png_names) == 1
        assert archive.read(svg_names[0]) == svg
        assert archive.read(png_names[0]) == png

        document_root = ElementTree.fromstring(archive.read("word/document.xml"))
        relationships_root = ElementTree.fromstring(
            archive.read("word/_rels/document.xml.rels")
        )
        content_types_root = ElementTree.fromstring(archive.read("[Content_Types].xml"))

    fallback_blip = document_root.find(".//a:blip", NS)
    svg_blip = fallback_blip.find("./a:extLst/a:ext/asvg:svgBlip", NS)
    extension = fallback_blip.find("./a:extLst/a:ext", NS)
    assert fallback_blip.get(f"{{{R_NS}}}embed")
    assert svg_blip.get(f"{{{R_NS}}}embed")
    assert extension.get("uri") == ASVG_EXTENSION_URI

    relationship_by_id = {
        relationship.get("Id"): relationship
        for relationship in relationships_root.findall(f"{{{PACKAGE_REL_NS}}}Relationship")
    }
    fallback_relationship = relationship_by_id[fallback_blip.get(f"{{{R_NS}}}embed")]
    svg_relationship = relationship_by_id[svg_blip.get(f"{{{R_NS}}}embed")]
    assert fallback_relationship.get("Type") == IMAGE_REL_TYPE
    assert svg_relationship.get("Type") == IMAGE_REL_TYPE
    assert fallback_relationship.get("Target") == png_names[0].removeprefix("word/")
    assert svg_relationship.get("Target") == svg_names[0].removeprefix("word/")
    assert fallback_relationship.get("TargetMode") is None
    assert svg_relationship.get("TargetMode") is None

    svg_override = next(
        element
        for element in content_types_root.findall(f"{{{CONTENT_TYPES_NS}}}Override")
        if element.get("PartName") == f"/{svg_names[0]}"
    )
    assert svg_override.get("ContentType") == "image/svg+xml"

    doc_properties = document_root.find(".//wp:docPr", NS)
    picture_properties = document_root.find(".//pic:cNvPr", NS)
    for properties in (doc_properties, picture_properties):
        assert properties.get("name") == "mw-flow-primary-endpoint"
        assert properties.get("title") == "mw-flow-primary-endpoint"
        assert properties.get("descr") == "主要终点医学写作研究流程图"


def test_existing_paragraph_is_used_without_adding_another_paragraph():
    document = Document()
    paragraph = document.add_paragraph("图前文本")

    add_svg_figure_with_png_fallback(
        paragraph,
        svg_bytes=_svg_bytes(),
        png_bytes=_png_bytes(width=1, height=1),
        figure_id="mw-flow-paragraph",
        alt_text="段落内研究流程图",
        width=Inches(1),
    )

    assert len(document.paragraphs) == 1
    assert paragraph.text == "图前文本"
    assert paragraph._element.find(".//w:drawing", paragraph._element.nsmap) is not None


@pytest.mark.parametrize(
    ("unsafe_svg", "message"),
    [
        (
            b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>',
            "script",
        ),
        (
            b'<svg xmlns="http://www.w3.org/2000/svg"><foreignObject/></svg>',
            "foreignobject",
        ),
        (
            b'<svg xmlns="http://www.w3.org/2000/svg"><image '
            b'href="https://example.test/a.png"/></svg>',
            "external SVG references",
        ),
        (
            b'<svg xmlns="http://www.w3.org/2000/svg" '
            b'xmlns:xlink="http://www.w3.org/1999/xlink"><use '
            b'xlink:href="//example.test/a.svg#icon"/></svg>',
            "external SVG references",
        ),
    ],
)
def test_unsafe_svg_is_rejected_before_document_mutation(unsafe_svg, message):
    document = Document()
    existing_relationship_ids = set(document.part.rels)

    with pytest.raises(MedicalWritingFigureExportError, match=message):
        add_svg_figure_with_png_fallback(
            document,
            svg_bytes=unsafe_svg,
            png_bytes=_png_bytes(width=1, height=1),
            figure_id="unsafe-figure",
            alt_text="不安全流程图",
            width=Inches(1),
        )

    assert document.paragraphs == []
    assert set(document.part.rels) == existing_relationship_ids


def _svg_bytes() -> bytes:
    return (
        b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100" '
        b'viewBox="0 0 200 100"><defs><linearGradient id="g"><stop '
        b'offset="0" stop-color="#4472C4"/></linearGradient></defs><rect '
        b'width="200" height="100" fill="url(#g)"/></svg>'
    )


def _png_bytes(*, width: int, height: int) -> bytes:
    signature = b"\x89PNG\r\n\x1a\n"
    header = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    rows = b"".join(b"\x00" + (b"\x44\x72\xC4" * width) for _ in range(height))
    return signature + _png_chunk(b"IHDR", header) + _png_chunk(
        b"IDAT", zlib.compress(rows)
    ) + _png_chunk(b"IEND", b"")


def _png_chunk(chunk_type: bytes, payload: bytes) -> bytes:
    checksum = binascii.crc32(chunk_type + payload) & 0xFFFFFFFF
    return struct.pack(">I", len(payload)) + chunk_type + payload + struct.pack(">I", checksum)
