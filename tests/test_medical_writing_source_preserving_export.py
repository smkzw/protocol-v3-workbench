from __future__ import annotations

import base64
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone

import pytest
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
from packages.contracts.workbench_contracts import (
    ApprovalState,
    MedicalWritingLiteratureLibrary,
    MedicalWritingProjectReference,
)

from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_document_exporter import (
    MedicalWritingDocumentDocxExportError,
    document_index_catalog,
    export_source_preserving_medical_writing_document_docx,
)
from services.api.app.medical_writing_table_templates import (
    MedicalWritingTableTemplateService,
)
from services.api.app.medical_writing_study_schema import (
    render_study_schema_png,
    render_study_schema_svg,
)
from tests.test_medical_writing_study_schema import _pnh_schema


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
WORD = {"w": WORD_NS}


def _assert_required_bilingual_fonts(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for name in archive.namelist():
            if not (name.startswith("word/") and name.endswith(".xml")):
                continue
            root = etree.fromstring(archive.read(name))
            for run in root.xpath(".//w:r", namespaces=WORD):
                fonts = run.find(
                    f"./{{{WORD_NS}}}rPr/{{{WORD_NS}}}rFonts"
                )
                assert fonts is not None, (name, "missing rFonts")
                assert fonts.get(f"{{{WORD_NS}}}eastAsia") == "宋体"
                assert fonts.get(f"{{{WORD_NS}}}ascii") == "Times New Roman"
                assert fonts.get(f"{{{WORD_NS}}}hAnsi") == "Times New Roman"


def _assert_only_font_normalizable_parts_changed(source_zip, output_zip) -> None:
    assert set(source_zip.namelist()) == set(output_zip.namelist())
    for name in source_zip.namelist():
        if source_zip.read(name) == output_zip.read(name):
            continue
        assert (
            name.startswith("word/") and name.endswith(".xml")
        ) or (
            name.startswith("word/media/") and name.lower().endswith(".svg")
        ), name


def _strip_font_properties(element) -> None:
    for fonts in element.findall(".//" + qn("w:rFonts")):
        parent = fonts.getparent()
        parent.remove(fonts)
        if parent.tag == qn("w:rPr") and len(parent) == 0 and not parent.attrib:
            parent.getparent().remove(parent)


def test_imported_document_without_content_changes_preserves_content_and_normalizes_fonts():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        _assert_only_font_normalizable_parts_changed(source_zip, output_zip)
    assert "\n".join(
        paragraph.text for paragraph in Document(source_path).paragraphs
    ) == "\n".join(
        paragraph.text
        for paragraph in Document(io.BytesIO(result.content)).paragraphs
    )
    _assert_required_bilingual_fonts(result.content)
    assert (
        result.metadata["source_preservation_mode"]
        == "original_package_font_normalized"
    )
    assert result.metadata["changed_source_block_count"] == 0
    reference_status = result.metadata["source_reference_reindex"]
    assert reference_status["status"] == "blocked"
    assert reference_status["action"] == "block"
    assert reference_status["user_action_required"] is True
    assert "Word" in reference_status["warning_message"]
    assert (
        "刷新" in reference_status["warning_message"]
        or "重新绑定" in reference_status["warning_message"]
    )
    assert "重新导入" in reference_status["warning_message"]
    json.dumps(result.metadata)


def test_original_protocol_path_is_available_only_for_configured_imported_project():
    service = MedicalWritingDocumentService()

    assert service.original_protocol_path("proj_d001").name.endswith(".docx")

    try:
        service.original_protocol_path("proj_not_configured")
    except KeyError:
        pass
    else:
        raise AssertionError("unconfigured project unexpectedly resolved an original DOCX")


def test_approved_unchanged_imported_document_is_blocked_until_external_reindex():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    approved = baseline.model_copy(
        update={
            "status": "medically_approved",
            "sections": [
                section.model_copy(
                    update={"approval_state": ApprovalState.MEDICALLY_APPROVED},
                    deep=True,
                )
                for section in baseline.sections
            ],
        },
        deep=True,
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="approved_final export blocked",
    ):
        export_source_preserving_medical_writing_document_docx(
            source_path,
            baseline,
            approved,
            mode="approved_final",
            approval_reference="approval-rux-source-preserving-r1",
        )


def test_imported_paragraph_edit_changes_only_main_document_part_and_keeps_source_layout():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    target = next(
        block
        for section in assembled.sections
        for block in section.content_blocks
        if block.get("source_locator") == "docx:paragraph:25"
    )
    replacement = "本文件中的信息仅供本研究医学审核使用。"
    target["text"] = replacement
    target["rich_text"]["content"] = [
        {
            "type": "text",
            "text": replacement,
            "marks": target["rich_text"]["content"][0]["marks"],
        }
    ]
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        _assert_only_font_normalizable_parts_changed(source_zip, output_zip)
        assert replacement in "\n".join(
            paragraph.text
            for paragraph in Document(io.BytesIO(result.content)).paragraphs
        )

    assert (
        result.metadata["source_preservation_mode"]
        == "source_locator_patch_font_normalized"
    )
    assert result.metadata["changed_source_block_count"] == 1


def test_text_only_imported_paragraph_edit_preserves_original_runs_fields_and_properties():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    target = next(
        block
        for section in assembled.sections
        for block in section.content_blocks
        if block.get("source_locator") == "docx:paragraph:714"
    )
    target["text"] = str(target["text"]).replace("受试者", "参与者", 1)
    rich_node = next(
        node
        for node in target["rich_text"]["content"]
        if "受试者" in str(node.get("text") or "")
    )
    rich_node["text"] = str(rich_node["text"]).replace("受试者", "参与者", 1)
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        source_root = etree.fromstring(source_zip.read("word/document.xml"))
        output_root = etree.fromstring(output_zip.read("word/document.xml"))
    body_order = int(target["body_order"])
    source_paragraph = list(source_root.find(qn("w:body")))[body_order]
    output_paragraph = list(output_root.find(qn("w:body")))[body_order]
    assert "参与者" in "".join(
        str(node.text or "") for node in output_paragraph.findall(".//" + qn("w:t"))
    )
    for paragraph in (source_paragraph, output_paragraph):
        for node in paragraph.findall(".//" + qn("w:t")):
            node.text = ""
        _strip_font_properties(paragraph)
    assert etree.tostring(source_paragraph, method="c14n") == etree.tostring(
        output_paragraph,
        method="c14n",
    )


def test_imported_table_cell_edit_preserves_table_geometry_and_all_other_package_parts():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    table = next(
        block
        for section in assembled.sections
        for block in section.content_blocks
        if block.get("source_locator") == "docx:table:1"
    )
    cell = next(
        cell
        for row in table["rows"]
        for cell in row
        if cell.get("text")
    )
    replacement = f"{cell['text']}（医学复核）"
    cell["text"] = replacement
    rich_text = cell.get("rich_text")
    if rich_text and rich_text.get("type") == "doc":
        target_paragraph = rich_text["content"][-1]
        target_paragraph["content"][-1]["text"] += "（医学复核）"
    elif rich_text:
        rich_text["content"][-1]["text"] += "（医学复核）"
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        _assert_only_font_normalizable_parts_changed(source_zip, output_zip)
        assert replacement.encode("utf-8") not in source_zip.read("word/document.xml")

    output_document = Document(io.BytesIO(result.content))
    assert replacement in [
        cell.text
        for row in output_document.tables[1].rows
        for cell in row.cells
    ]

    assert result.metadata["changed_source_block_count"] == 1


def test_governed_generated_table_is_inserted_at_imported_section_end_without_rebuilding_source_tables():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    target_section = next(
        section for section in assembled.sections if section.heading == "试验研究计划"
    )
    generated = MedicalWritingTableTemplateService().instantiate(
        "analysis_sets",
        title="分析集定义（医学复核）",
        instance_id="source_preserving_export_test",
    )
    generated.update(
        {
            "source_kind": "medical_writing_template",
            "template_id": "analysis_sets",
            "editable": True,
        }
    )
    target_section.content_blocks.append(generated)
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    source_document = Document(source_path)
    output_document = Document(io.BytesIO(result.content))
    assert len(output_document.tables) == len(source_document.tables) + 1
    assert "分析集定义（医学复核）" in "\n".join(
        paragraph.text for paragraph in output_document.paragraphs
    )
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert "SEQ 表" in document_xml
    assert 'w:name="_MWTAB_' in document_xml
    assert result.metadata["changed_source_block_count"] == 1


def test_imported_paragraph_cross_reference_resolves_against_source_document_index():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    target = document_index_catalog(assembled)["tables"][0]
    paragraph = next(
        block
        for section in assembled.sections
        for block in section.content_blocks
        if block.get("block_type") == "paragraph"
        and block.get("source_kind") == "original_protocol_docx"
        and block.get("text")
    )
    paragraph["text"] = f"见表 {target['number']}。"
    paragraph["rich_text"] = {
        "type": "paragraph",
        "content": [
            {"type": "text", "text": "见"},
            {
                "type": "text",
                "text": f"表 {target['number']}",
                "marks": [
                    {
                        "type": "crossReference",
                        "attrs": {
                            "targetKind": "table",
                            "targetId": target["object_id"],
                        },
                    }
                ],
            },
            {"type": "text", "text": "。"},
        ],
    }
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
    assert f"REF {target['bookmark_name']} \\h" in document_xml
    assert f'w:name="{target["bookmark_name"]}"' in document_xml


def test_managed_citation_in_imported_protocol_gets_hyperlink_and_reindexed_reference_bookmark():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    target = next(
        block
        for section in assembled.sections
        for block in section.content_blocks
        if block.get("source_locator") == "docx:paragraph:25"
    )
    reference_id = "ref_source_preserving_test"
    target["text"] = "既往研究支持该设计[1]"
    target["rich_text"] = {
        "type": "paragraph",
        "attrs": target["rich_text"].get("attrs", {}),
        "content": [
            {"type": "text", "text": "既往研究支持该设计"},
            {
                "type": "text",
                "text": "[1]",
                "marks": [
                    {"type": "citation", "attrs": {"referenceId": reference_id}}
                ],
            },
        ],
    }
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)
    now = datetime(2026, 7, 17, tzinfo=timezone.utc)
    library = MedicalWritingLiteratureLibrary(
        project_id=project_id,
        citation_style="gbt_7714_2015_numeric",
        references=[
            MedicalWritingProjectReference(
                reference_id=reference_id,
                project_id=project_id,
                canonical_key="doi:10.1000/source-preserving-test",
                source_kind="doi",
                source_input="10.1000/source-preserving-test",
                title="Source preserving evidence",
                authors=["Zhang S", "Li S"],
                journal="Clinical Evidence",
                year="2026",
                volume="1",
                issue="1",
                pages="1-8",
                doi="10.1000/source-preserving-test",
                validation_status="confirmed",
                created_at=now,
                updated_at=now,
            )
        ],
    )

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
        literature_library=library,
    )

    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_payload = archive.read("word/document.xml")
        document_xml = document_payload.decode("utf-8")
        document_root = etree.fromstring(document_payload)
    assert "w:hyperlink" in document_xml
    assert "_MWREF_" in document_xml
    bookmark_ids = document_root.xpath(
        ".//w:bookmarkStart/@w:id",
        namespaces=WORD,
    )
    assert len(bookmark_ids) == len(set(bookmark_ids))
    assert result.metadata["citation_rewritten_block_count"] > 0
    reindex_metadata = result.metadata["source_reference_reindex"]
    assert reindex_metadata["status"] == "applied"
    assert reindex_metadata["action"] == "managed_citation_plan"
    assert reindex_metadata["changed"] is True
    assert reindex_metadata["citation_count"] > 0
    assert reindex_metadata["reference_count"] > 0
    assert reindex_metadata["mapping"]
    assert reindex_metadata["failure_code"] == ""
    assert "Source preserving evidence" in "\n".join(
        paragraph.text for paragraph in Document(io.BytesIO(result.content)).paragraphs
    )


def test_governed_study_schema_figure_adds_only_relationship_content_type_and_media_parts():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    schema = _pnh_schema()
    svg = render_study_schema_svg(schema)
    png = render_study_schema_png(svg)
    figure = {
        "block_id": "mwgenerated_figure_source_preserving",
        "block_type": "figure",
        "figure_id": "mwfigure_source_preserving",
        "figure_kind": "study_schema",
        "title": "研究设计概况（医学复核）",
        "alt_text": "研究流程图：研究设计概况",
        "source_kind": "medical_writing_study_schema",
        "source_locator": "generated:medical_writing_study_schema:source_preserving:r1",
        "editable": False,
        "schema_id": schema.schema_id,
        "schema_revision": schema.revision,
        "schema_state_sha256": schema.state_sha256,
        "layout_revision": 0,
        "svg": svg,
        "svg_sha256": hashlib.sha256(svg.encode("utf-8")).hexdigest(),
        "png_base64": base64.b64encode(png).decode("ascii"),
        "png_sha256": hashlib.sha256(png).hexdigest(),
        "width_inches": 6.45,
    }
    target_section = next(
        section for section in assembled.sections if section.heading == "试验研究计划"
    )
    target_section.content_blocks.append(figure)
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        source_names = set(source_zip.namelist())
        output_names = set(output_zip.namelist())
        new_names = output_names - source_names
        assert new_names
        assert all(name.startswith("word/media/") for name in new_names)
        changed_existing = {
            name
            for name in source_names
            if source_zip.read(name) != output_zip.read(name)
        }
        required_changes = {
            "[Content_Types].xml",
            "word/_rels/document.xml.rels",
            "word/document.xml",
        }
        assert required_changes.issubset(changed_existing)
        assert all(
            name in required_changes
            or (name.startswith("word/") and name.endswith(".xml"))
            or (name.startswith("word/media/") and name.lower().endswith(".svg"))
            for name in changed_existing
        )
        document_xml = output_zip.read("word/document.xml").decode("utf-8")
        assert "asvg:svgBlip" in document_xml
        assert "研究设计概况（医学复核）" in document_xml
        document_root = etree.fromstring(output_zip.read("word/document.xml"))
        body = document_root.find(qn("w:body"))
        assert body is not None
        assert list(body)[-1].tag == qn("w:sectPr")
        assert sum(child.tag == qn("w:sectPr") for child in body) == 1
        figure_paragraphs = [
            paragraph
            for paragraph in body.findall(qn("w:p"))
            if "研究设计概况（医学复核）"
            in "".join(paragraph.xpath(".//w:t/text()", namespaces=WORD))
        ]
        assert len(figure_paragraphs) == 1
        figure_paragraph = figure_paragraphs[0]
        assert len(figure_paragraph.xpath(".//w:drawing", namespaces=WORD)) == 1
        assert len(figure_paragraph.xpath(".//w:br", namespaces=WORD)) == 1
        assert figure_paragraph.xpath(".//w:bookmarkStart", namespaces=WORD)
        assert "SEQ 图" in "".join(
            figure_paragraph.xpath(".//w:instrText/text()", namespaces=WORD)
        )


def test_multiple_governed_figures_receive_distinct_svg_parts_and_may_share_identical_png_fallback():
    service = MedicalWritingDocumentService()
    project_id = "proj_rux_03_002"
    source_path = service.original_protocol_path(project_id)
    baseline = service.document_for_revision(project_id)
    assembled = baseline.model_copy(deep=True)
    schema = _pnh_schema()
    svg = render_study_schema_svg(schema)
    png = render_study_schema_png(svg)

    def figure(index: int):
        return {
            "block_id": f"mwgenerated_figure_source_preserving_{index}",
            "block_type": "figure",
            "figure_id": f"mwfigure_source_preserving_{index}",
            "figure_kind": "study_schema",
            "title": f"研究设计概况 {index}",
            "alt_text": f"研究流程图 {index}",
            "source_kind": "medical_writing_study_schema",
            "source_locator": f"generated:medical_writing_study_schema:source_preserving:{index}",
            "editable": False,
            "schema_id": schema.schema_id,
            "schema_revision": schema.revision,
            "schema_state_sha256": schema.state_sha256,
            "layout_revision": index,
            "svg": svg,
            "svg_sha256": hashlib.sha256(svg.encode("utf-8")).hexdigest(),
            "png_base64": base64.b64encode(png).decode("ascii"),
            "png_sha256": hashlib.sha256(png).hexdigest(),
            "width_inches": 6.45,
        }

    target_section = next(
        section for section in assembled.sections if section.heading == "试验研究计划"
    )
    target_section.content_blocks.extend([figure(1), figure(2)])
    assembled = assembled.model_copy(update={"status": "draft_preview"}, deep=True)

    result = export_source_preserving_medical_writing_document_docx(
        source_path,
        baseline,
        assembled,
        mode="draft_preview",
    )

    with zipfile.ZipFile(source_path) as source_zip, zipfile.ZipFile(
        io.BytesIO(result.content)
    ) as output_zip:
        new_media = {
            name
            for name in output_zip.namelist()
            if name not in source_zip.namelist() and name.startswith("word/media/")
        }
        # python-docx de-duplicates identical binary image payloads. The two
        # SVGs are distinct OOXML parts while the identical PNG fallback may
        # be shared safely by both drawings.
        assert len({name for name in new_media if name.endswith(".svg")}) == 2
        assert len({name for name in new_media if name.endswith(".png")}) >= 1
        document_xml = output_zip.read("word/document.xml").decode("utf-8")
        assert document_xml.count("asvg:svgBlip") >= 2
        assert "研究设计概况 1" in document_xml
        assert "研究设计概况 2" in document_xml
