from __future__ import annotations

import io
import json
import sqlite3
import tempfile
import time
from copy import deepcopy
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from docx import Document
from fastapi.testclient import TestClient

from packages.contracts.workbench_contracts import (
    ApprovalState,
    MedicalWritingWorkingCopySaveRequest,
    ProtocolDocument,
    ProtocolSection,
)
from packages.contracts.workbench_contracts.models import (
    InterventionRulesAuthority,
    InterventionRulesIpAdjustmentPolicy,
    InterventionRulesProductRole,
    MedicalWritingInterimAnalysisDesign,
    MedicalWritingInterventionIpRegimen,
    MedicalWritingInterventionRules,
    MedicalWritingPhase1Part,
    MedicalWritingPicosDefinition,
    MedicalWritingProtocolAssemblyPlanConfirmRequest,
    MedicalWritingProtocolAssemblyPlanRefreshRequest,
    MedicalWritingSectionFreezeRequest,
    MedicalWritingStructuredStudyDesign,
    MedicalWritingStudyDefinition,
    MedicalWritingStudyFraming,
)
from services.api.app import main as app_main
from services.api.app.medical_writing_plan_consumption import (
    MedicalWritingPlanConsumptionHelper,
)
from services.api.app.medical_writing_document_export_jobs import (
    MedicalWritingDocumentExportJobService,
)
from services.api.app.medical_writing_durable_jobs import (
    DurableJobStore,
    DurableJobWorker,
)
from services.api.app.medical_writing_protocol_assembly_plan import (
    MedicalWritingProtocolAssemblyPlanService,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository
from services.api.app.medical_writing_table_templates import (
    MedicalWritingTableTemplateService,
)
from services.api.app.sqlite_runtime_store import RuntimeStoreError, SqliteRuntimeStore


PROJECT_ID = "proj_rux_03_002"


def _confirm_managed_export_plan(tmpdir: Path, project_id: str = PROJECT_ID):
    """Create/confirm a ProtocolAssemblyPlan for managed greenfield-path export tests."""
    from datetime import datetime, timezone

    now = datetime(2026, 7, 22, tzinfo=timezone.utc)
    design = MedicalWritingStructuredStudyDesign(
        randomization_mode="randomized",
        blinding_mode="double_blind",
        comparator_type="placebo",
        assignment_model="parallel_group",
        adaptive_design_enabled=False,
        sample_size_reestimation_planned=False,
        treatment_switch_planned=False,
        crossover_planned=False,
        open_label_extension_planned=False,
        src_planned=False,
        dmc_planned=True,
        phase1_parts=[],
        interim_analysis=MedicalWritingInterimAnalysisDesign(planned=False),
    )
    rules = MedicalWritingInterventionRules(
        authority=InterventionRulesAuthority.STRUCTURED,
        ip_regimens=[
            MedicalWritingInterventionIpRegimen(
                regimen_id="ip-1",
                product_name="RUX",
                product_role=InterventionRulesProductRole.INVESTIGATIONAL_PRODUCT,
                dose_and_frequency="test dose",
                route="外用",
                treatment_period="研究期",
            ),
            MedicalWritingInterventionIpRegimen(
                regimen_id="pbo-1",
                product_name="Placebo",
                product_role=InterventionRulesProductRole.PLACEBO,
                dose_and_frequency="matching placebo",
                route="外用",
                treatment_period="研究期",
            ),
        ],
        ip_adjustment_policy=InterventionRulesIpAdjustmentPolicy.NO_PLANNED_ADJUSTMENT,
        no_planned_adjustment_statement="无计划剂量调整。",
        non_ip_treatment_rules=[],
    )
    definition = MedicalWritingStudyDefinition(
        definition_id=f"def-{project_id}-export-api",
        project_id=project_id,
        revision=1,
        origin="guided_greenfield",
        framing=MedicalWritingStudyFraming(
            protocol_id="RUX-03-002",
            document_title="RUX export API plan fixture",
            indication="特应性皮炎",
            clinicaltrials_condition_term="Atopic Dermatitis",
            study_phase="III期",
            investigational_product="RUX",
            structured_design=design,
            product_profile={
                "technology_type": "small_molecule",
                "administration_routes": ["外用"],
                "dosage_forms": ["乳膏"],
                "exposure_scope": "local",
            },
        ),
        picos=MedicalWritingPicosDefinition(
            population_summary="AD患者",
            primary_endpoint="EASI",
            intervention_rules=rules,
        ),
        state_sha256="b" * 64,
        created_at=now,
        updated_at=now,
        updated_by="export_api_test",
    )
    store = {project_id: definition}
    plan_service = MedicalWritingProtocolAssemblyPlanService(
        tmpdir / f"plan_{project_id}.sqlite3",
        store.__getitem__,
    )
    refreshed = plan_service.refresh(
        project_id,
        MedicalWritingProtocolAssemblyPlanRefreshRequest(
            expected_plan_revision=0,
            expected_source_definition_id=definition.definition_id,
            expected_source_definition_revision=definition.revision,
            expected_source_definition_sha256=definition.state_sha256,
            actor="export_api_test",
            idempotency_key=f"refresh-{project_id}",
        ),
    )
    confirmed = plan_service.confirm(
        project_id,
        MedicalWritingProtocolAssemblyPlanConfirmRequest(
            expected_plan_revision=refreshed.plan.revision,
            expected_plan_sha256=refreshed.plan.state_sha256,
            actor="export_api_test",
            idempotency_key=f"confirm-{project_id}",
        ),
    )
    helper = MedicalWritingPlanConsumptionHelper(plan_service)
    return plan_service, helper, confirmed.plan


class TwoSectionDocumentService:
    def __init__(self):
        document_id = "mwdoc_export_api_test"
        table = MedicalWritingTableTemplateService().instantiate(
            "objectives_endpoints",
            title="研究目的与终点",
            instance_id="source-objectives",
        )
        table.update(
            {
                "source_kind": "original_protocol_docx",
                "source_locator": "docx:table:1",
                "body_order": 2,
                "table_caption": {
                    "caption_id": "caption_source_objectives",
                    "text": "研究目的与终点",
                    "source_locator": "docx:paragraph:5",
                    "paragraph_index": 5,
                    "body_order": 2,
                    "review_status": "pending_human_review",
                    "association_reason": "表格前相邻表题。",
                },
            }
        )
        table["structured_table"]["source_lineage"] = ["docx:table:1"]
        table["structured_table"]["source_caption"] = deepcopy(
            table["table_caption"]
        )
        table["structured_table"]["source_note_metadata"] = [
            {
                "note_id": "source_note_objectives_1",
                "source_locator": "docx:paragraph:6",
                "target_scope": "cell",
                "target_row_index": 1,
                "target_cell_id": table["rows"][1][0]["cell_id"],
                "source_kind": "post_table_note",
                "review_status": "pending_human_review",
                "association_reason": "表后相邻附注。",
            }
        ]
        self.document = ProtocolDocument(
            document_id=document_id,
            project_id=PROJECT_ID,
            protocol_id="RUX-03-002测试方案",
            version="V1.3",
            status="source_imported_unverified",
            sections=[
                ProtocolSection(
                    section_id="section_objectives",
                    document_id=document_id,
                    heading="研究目的",
                    approval_state=ApprovalState.IN_MEDICAL_REVIEW,
                    content_blocks=[
                        {
                            "block_id": "source_heading_1",
                            "block_type": "heading",
                            "body_order": 0,
                            "text": "研究目的",
                            "source_kind": "original_protocol_docx",
                            "source_locator": "docx:paragraph:1",
                        },
                        {
                            "block_id": "source_paragraph_1",
                            "block_type": "paragraph",
                            "body_order": 1,
                            "text": "第一章节原始正文。",
                            "source_kind": "original_protocol_docx",
                            "source_locator": "docx:paragraph:2",
                        },
                        table,
                    ],
                ),
                ProtocolSection(
                    section_id="section_design",
                    document_id=document_id,
                    heading="研究设计",
                    approval_state=ApprovalState.IN_MEDICAL_REVIEW,
                    content_blocks=[
                        {
                            "block_id": "source_heading_2",
                            "block_type": "heading",
                            "body_order": 3,
                            "text": "研究设计",
                            "source_kind": "original_protocol_docx",
                            "source_locator": "docx:paragraph:3",
                        },
                        {
                            "block_id": "source_paragraph_2",
                            "block_type": "paragraph",
                            "body_order": 4,
                            "text": "第二章节保留原始正文。",
                            "source_kind": "original_protocol_docx",
                            "source_locator": "docx:paragraph:4",
                        },
                    ],
                ),
            ],
        )

    def document_session(self, project_id: str) -> ProtocolDocument:
        assert project_id == PROJECT_ID
        return self.document.model_copy(deep=True)

    def document_for_revision(self, project_id: str) -> ProtocolDocument:
        return self.document_session(project_id)

    def section(self, project_id: str, section_id: str) -> ProtocolSection:
        return next(
            section.model_copy(deep=True)
            for section in self.document_session(project_id).sections
            if section.section_id == section_id
        )


class TestMedicalWritingDocumentExportApi:
    def setup_method(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmpdir.name) / "runtime.sqlite3"
        self.documents = TwoSectionDocumentService()
        self.store = SqliteRuntimeStore(self.db_path)
        self.repository = MedicalWritingRuntimeRepository(self.documents, self.store)
        self.repository_patch = patch(
            "services.api.app.main.medical_writing_runtime_repository",
            self.repository,
        )
        self.repository_patch.start()
        # The synthetic document service does not have a real source DOCX,
        # so route exports through the greenfield template path instead of
        # the source-preserving baseline-matching path.
        self.source_mode_patch = patch(
            "services.api.app.main.medical_writing_document_service.source_mode",
            return_value="greenfield",
        )
        self.source_mode_patch.start()
        # Managed exporter path requires a confirmed ProtocolAssemblyPlan pin.
        self.plan_service, self.plan_helper, self.plan = _confirm_managed_export_plan(
            Path(self.tmpdir.name),
            PROJECT_ID,
        )
        self.plan_service_patch = patch(
            "services.api.app.main.medical_writing_protocol_assembly_plan_service",
            self.plan_service,
        )
        self.plan_helper_patch = patch(
            "services.api.app.main.medical_writing_plan_consumption_helper",
            self.plan_helper,
        )
        self.plan_service_patch.start()
        self.plan_helper_patch.start()
        self.client = TestClient(app_main.app)

    def teardown_method(self):
        self.plan_helper_patch.stop()
        self.plan_service_patch.stop()
        self.repository_patch.stop()
        self.source_mode_patch.stop()
        self.tmpdir.cleanup()

    def _save_section(self, section_id: str, text: str):
        source = self.documents.section(PROJECT_ID, section_id)
        blocks = deepcopy(source.content_blocks)
        paragraph = next(
            block for block in blocks if block.get("block_type") == "paragraph"
        )
        paragraph["text"] = text
        return self.repository.save_working_copy(
            PROJECT_ID,
            section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=source.document_id,
                expected_revision=0,
                content_blocks=blocks,
                actor="medical_manager_test",
                idempotency_key=f"save-{section_id}",
            ),
        )

    def _freeze(self, working_copy):
        return self.repository.freeze_current_version(
            PROJECT_ID,
            working_copy.section_id,
            MedicalWritingSectionFreezeRequest(
                document_id=working_copy.document_id,
                expected_working_copy_revision=working_copy.revision,
                actor="medical_manager_test",
                reason="医学作者已核对本章节并确认冻结当前版本。",
                idempotency_key=f"freeze-{working_copy.section_id}",
            ),
        )

    def test_draft_preview_mixes_saved_working_copy_and_source_sections_in_real_docx(self):
        self._save_section("section_objectives", "第一章节当前工作副本正文。")

        response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
            params={"mode": "draft_preview"},
        )

        assert response.status_code == 200, response.text
        assert response.content.startswith(b"PK")
        assert response.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
        disposition = response.headers["content-disposition"]
        assert "filename*=UTF-8''" in disposition
        assert "%E8%8D%89%E7%A8%BF%E9%A2%84%E8%A7%88" in disposition
        assert response.headers["x-medical-writing-export-mode"] == "draft_preview"
        assert len(response.headers["x-medical-writing-docx-sha256"]) == 64

        exported = Document(io.BytesIO(response.content))
        text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
        assert text.index("第一章节当前工作副本正文。") < text.index(
            "第二章节保留原始正文。"
        )
        assert "第一章节当前工作副本正文。" in text
        assert "第一章节原始正文。" not in text
        assert "第二章节保留原始正文。" in text
        assert len(exported.tables) == 1
        assert "|---" not in text

    def test_async_document_export_start_adapts_to_durable_service(self):
        start_response = SimpleNamespace(
            model_dump=lambda mode="json": {
                "job_id": "job-export-001",
                "project_id": PROJECT_ID,
                "job_type": "medical_writing_document_export",
                "status": "queued",
                "reused": False,
                "request_hash": "a" * 64,
            }
        )
        with patch.object(
            app_main.medical_writing_document_export_job_service,
            "start",
            return_value=start_response,
        ) as start:
            response = self.client.post(
                f"/api/projects/{PROJECT_ID}/medical-writing/document-exports",
                json={
                    "mode": "draft_preview",
                    "actor": "medical_manager_test",
                    "idempotency_key": "export-start-001",
                },
            )

        assert response.status_code == 200, response.text
        assert response.json()["job_id"] == "job-export-001"
        start.assert_called_once_with(
            PROJECT_ID,
            mode="draft_preview",
            idempotency_key="export-start-001",
            actor="medical_manager_test",
            payload={},
        )

    def test_async_document_export_download_preserves_docx_metadata_headers(self):
        source_response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
            params={"mode": "draft_preview"},
        )
        assert source_response.status_code == 200, source_response.text
        artifact_path = Path(self.tmpdir.name) / "async-export.docx"
        artifact_path.write_bytes(source_response.content)
        metadata = {
            "mode": "draft_preview",
            "document_id": self.documents.document.document_id,
            "source_snapshot_sha256": source_response.headers[
                "x-medical-writing-source-snapshot-sha256"
            ],
            "docx_sha256": source_response.headers[
                "x-medical-writing-docx-sha256"
            ],
            "section_count": 2,
            "table_count": 1,
            "freeze_reference": "",
            "source_reference_status": "",
            "user_action_required": False,
            "warning_message": "",
        }
        artifact = SimpleNamespace(
            path=artifact_path,
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            filename="RUX-03-002_V1.3_草稿预览.docx",
            metadata=metadata,
        )
        with patch.object(
            app_main.medical_writing_document_export_job_service,
            "read_artifact",
            return_value=artifact,
        ) as read_artifact:
            response = self.client.get(
                f"/api/projects/{PROJECT_ID}/medical-writing/"
                "document-exports/job-export-001/download"
            )

        assert response.status_code == 200, response.text
        assert response.content.startswith(b"PK")
        assert (
            response.headers["x-medical-writing-export-mode"]
            == "draft_preview"
        )
        assert response.headers["x-medical-writing-section-count"] == "2"
        assert response.headers["x-medical-writing-table-count"] == "1"
        read_artifact.assert_called_once_with(PROJECT_ID, "job-export-001")

    def test_async_document_export_runs_through_shared_status_and_download_routes(self):
        durable_store = DurableJobStore(
            Path(self.tmpdir.name) / "document-export-jobs.sqlite3"
        )
        durable_worker = DurableJobWorker(durable_store)
        export_service = MedicalWritingDocumentExportJobService(
            store=durable_store,
            worker=durable_worker,
            artifact_root=Path(self.tmpdir.name) / "document-export-artifacts",
            callbacks=app_main._medical_writing_document_export_callbacks(),
        )
        try:
            with (
                patch.object(app_main, "mw_durable_store", durable_store),
                patch.object(app_main, "mw_durable_worker", durable_worker),
                patch.object(
                    app_main,
                    "medical_writing_document_export_job_service",
                    export_service,
                ),
            ):
                started = self.client.post(
                    f"/api/projects/{PROJECT_ID}/medical-writing/document-exports",
                    json={
                        "mode": "draft_preview",
                        "actor": "medical_manager_test",
                        "idempotency_key": "document-export-real-loop-001",
                    },
                )
                assert started.status_code == 200, started.text
                job_id = started.json()["job_id"]
                statuses = []
                progress_steps = []
                for _ in range(200):
                    status_response = self.client.get(
                        f"/api/projects/{PROJECT_ID}/medical-writing/jobs/{job_id}"
                    )
                    assert status_response.status_code == 200, status_response.text
                    payload = status_response.json()
                    statuses.append(payload["status"])
                    progress_steps.append(payload["progress"]["step"])
                    if payload["status"] in {"completed", "failed", "cancelled"}:
                        break
                    time.sleep(0.01)

                assert statuses[-1] == "completed", payload
                assert progress_steps[-1] == 5
                assert payload["progress"]["percent"] == 1.0
                result = self.client.get(
                    f"/api/projects/{PROJECT_ID}/medical-writing/jobs/{job_id}/result"
                )
                assert result.status_code == 200, result.text
                assert result.json()["artifact"]["job_id"] == job_id

                download = self.client.get(
                    f"/api/projects/{PROJECT_ID}/medical-writing/"
                    f"document-exports/{job_id}/download"
                )
                assert download.status_code == 200, download.text
                assert download.content.startswith(b"PK")
                assert (
                    download.headers["x-medical-writing-export-mode"]
                    == "draft_preview"
                )
                assert len(
                    download.headers["x-medical-writing-docx-sha256"]
                ) == 64
        finally:
            durable_worker.shutdown(timeout=5.0)

    def test_greenfield_export_passes_authoritative_cover_facts_to_exporter(self):
        real_exporter = app_main.export_medical_writing_document_docx
        with (
            patch.object(
                app_main.medical_writing_authoring_journey_service,
                "has_project",
                return_value=True,
            ),
            patch.object(
                app_main.medical_writing_authoring_journey_service,
                "get",
                return_value=SimpleNamespace(
                    study_definition=SimpleNamespace(
                        framing=SimpleNamespace(
                            investigational_product="CMS-D017胶囊"
                        )
                    )
                ),
            ),
            patch.object(
                app_main.user_project_store,
                "get",
                return_value=SimpleNamespace(
                    product_name="CMS-D017",
                    protocol_date="2026年07月20日",
                ),
            ),
            patch(
                "services.api.app.main.export_medical_writing_document_docx",
                wraps=real_exporter,
            ) as exporter,
        ):
            response = self.client.get(
                f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
                params={"mode": "draft_preview"},
            )

        assert response.status_code == 200, response.text
        assert exporter.call_args.kwargs["front_matter_overrides"] == {
            "investigational_product": "CMS-D017胶囊",
            "protocol_date": "2026年07月20日",
        }

    def test_document_index_route_returns_server_resolved_targets(self):
        table = self.documents.document.sections[0].content_blocks[2]
        table["title"] = "表 1 研究目的与终点"
        table["structured_table"]["title"] = "表 1 研究目的与终点"

        response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document-index"
        )

        assert response.status_code == 200, response.text
        payload = response.json()
        assert payload["schema_version"] == "medical_writing_document_index_v1"
        assert payload["project_id"] == PROJECT_ID
        assert payload["figures"] == []
        assert len(payload["tables"]) == 1
        assert payload["tables"][0]["object_id"] == table["table_id"]
        assert payload["tables"][0]["number"] == 1
        assert payload["tables"][0]["bookmark_name"].startswith("_MWTAB_")

    def test_final_export_rejects_any_applicable_section_without_current_freeze(self):
        first = self._save_section("section_objectives", "第一章节定稿正文。")
        self._freeze(first)

        response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
            params={"mode": "approved_final"},
        )

        assert response.status_code == 409
        assert "section_design" in response.json()["detail"]

    def test_final_export_uses_freeze_snapshots_and_rejects_payload_tampering(self):
        first = self._save_section("section_objectives", "第一章节冻结快照正文。")
        second = self._save_section("section_design", "第二章节冻结快照正文。")
        self._freeze(first)
        self._freeze(second)

        exported_response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
            params={"mode": "approved_final"},
        )
        assert exported_response.status_code == 200, exported_response.text
        assert exported_response.headers[
            "x-medical-writing-freeze-reference"
        ].startswith("author-frozen-document:")
        assert (
            exported_response.headers[
                "x-medical-writing-source-reference-status"
            ]
            == ""
        )
        assert (
            exported_response.headers[
                "x-medical-writing-user-action-required"
            ]
            == "false"
        )
        assert (
            exported_response.headers["x-medical-writing-warning-message"]
            == ""
        )
        exported = Document(io.BytesIO(exported_response.content))
        text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
        assert "第一章节冻结快照正文。" in text
        assert "第二章节冻结快照正文。" in text
        assert len(exported.tables) == 1

        current = self.repository.working_copy(PROJECT_ID, "section_objectives")
        tampered = current.model_copy(deep=True)
        tampered.content_blocks[1]["text"] = "当前可变记录中的错误正文。"
        with sqlite3.connect(self.db_path) as connection:
            connection.execute(
                """
                UPDATE medical_writing_working_copies
                SET payload_json = ?
                WHERE project_id = ? AND working_copy_id = ?
                """,
                (
                    json.dumps(tampered.model_dump(mode="json"), ensure_ascii=False),
                    PROJECT_ID,
                    current.working_copy_id,
                ),
            )
            connection.commit()

        tampered_response = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/document.docx",
            params={"mode": "approved_final"},
        )
        assert tampered_response.status_code == 409
        assert "inconsistent" in tampered_response.json()["detail"]

    def test_template_catalog_and_instantiate_status_boundaries(self):
        catalog = self.client.get("/api/medical-writing/table-templates")
        assert catalog.status_code == 200
        assert "dose_adjustment" not in {item["template_id"] for item in catalog.json()}
        assert "dose_modification" in {item["template_id"] for item in catalog.json()}
        assert "generic_table" in {item["template_id"] for item in catalog.json()}

        profile_catalog = self.client.get("/api/medical-writing/table-domain-profiles")
        assert profile_catalog.status_code == 200
        assert {item["domain"] for item in profile_catalog.json()} == {
            "objectives_endpoints",
            "sample_size_assumptions",
            "treatment_dose",
            "dose_modification",
            "laboratory_panel",
            "pk_immunogenicity_schedule",
            "analysis_sets",
            "version_history",
        }

        route = (
            f"/api/projects/{PROJECT_ID}/medical-writing/working-copies/"
            "section_objectives/table-templates/dose_modification/instantiate"
        )
        unsaved = self.client.post(route)
        assert unsaved.status_code == 409

        self._save_section("section_objectives", "保存后允许插入表格。")
        inserted = self.client.post(
            route,
            params={"title": "试验用药剂量调整与变更", "actor": "medical_manager_test"},
        )
        assert inserted.status_code == 200, inserted.text
        payload = inserted.json()
        assert payload["working_copy"]["revision"] == 2
        assert payload["table_block"]["template_id"] == "dose_modification"
        assert payload["table_block"]["source_kind"] == "medical_writing_template"

        unknown = self.client.post(route.replace("dose_modification", "unknown"))
        assert unknown.status_code == 404

        custom = self.client.post(
            route.replace("dose_modification", "generic_table"),
            params={
                "title": "自定义安全性汇总表",
                "row_count": 8,
                "column_count": 5,
                "header_row_count": 2,
                "orientation": "portrait",
                "notes_area": True,
                "actor": "medical_manager_test",
            },
        )
        assert custom.status_code == 200, custom.text
        custom_block = custom.json()["table_block"]
        assert custom_block["title"] == "自定义安全性汇总表"
        assert len(custom_block["rows"]) == 8
        assert len(custom_block["rows"][0]) == 5
        assert custom_block["header_row_count"] == 2
        assert custom_block["structured_table"]["word_layout"]["notes_area_enabled"] is True

        with patch.object(
            app_main.medical_writing_table_template_service,
            "instantiate",
            side_effect=ValueError("invalid template payload"),
        ):
            invalid = self.client.post(route)
        assert invalid.status_code == 422

        current = self.repository.working_copy(PROJECT_ID, "section_objectives")
        frozen = self._freeze(current)
        assert frozen.working_copy.freeze_status == "frozen"
        edited_after_freeze = self.client.post(
            route, params={"allow_duplicate": True}
        )
        assert edited_after_freeze.status_code == 200, edited_after_freeze.text
        assert edited_after_freeze.json()["working_copy"]["freeze_status"] == "editable"

    def test_new_promoted_templates_instantiate_through_http_and_persist(self):
        self._save_section("section_objectives", "保存后逐一插入新增结构化表格。")
        expected = {
            "sample_size_assumptions": {
                "label": "样本量估算假设",
                "headers": ["参数", "基本假设", "依据/来源", "敏感性情景", "待确认项"],
                "starter_rows": ["主要终点效应量", "变异/事件率", "检验水准与把握度", "脱落/不可评价比例"],
            },
            "analysis_sets": {
                "label": "分析集定义",
                "headers": ["分析集", "定义", "纳入规则", "排除规则", "主要用途", "偏离/敏感性处理"],
                "starter_rows": ["全分析集（FAS）", "符合方案集（PPS）", "安全性集（SS）"],
            },
            "pk_immunogenicity_schedule": {
                "label": "PK/PD/免疫原性采样计划",
                "headers": ["评估类型", "访视/研究日", "相对给药时间", "允许时间窗", "样本类型", "处理与保存", "实际时间记录"],
                "starter_rows": ["PK", "PD", "免疫原性"],
            },
        }

        inserted_block_ids = []
        for template_id, contract in expected.items():
            route = (
                f"/api/projects/{PROJECT_ID}/medical-writing/working-copies/"
                f"section_objectives/table-templates/{template_id}/instantiate"
            )
            response = self.client.post(
                route,
                params={"actor": "medical_manager_http_template_test"},
            )
            assert response.status_code == 200, response.text
            payload = response.json()
            block = payload["table_block"]
            inserted_block_ids.append(block["block_id"])
            assert block["template_id"] == template_id
            assert block["structured_table"]["domain"] == template_id
            assert block["structured_table"]["title"] == contract["label"]
            assert [cell["text"] for cell in block["rows"][0]] == contract["headers"]
            assert [row[0]["text"] for row in block["rows"][1:]] == contract["starter_rows"]
            structured_rows = block["structured_table"]["rows"]
            assert [cell["text"] for cell in structured_rows[0]["cells"]] == contract["headers"]
            assert block in payload["working_copy"]["content_blocks"]

        current = self.client.get(
            f"/api/projects/{PROJECT_ID}/medical-writing/working-copies/section_objectives"
        )
        assert current.status_code == 200
        persisted = {
            block["block_id"]: block
            for block in current.json()["content_blocks"]
            if block.get("block_id") in inserted_block_ids
        }
        assert set(persisted) == set(inserted_block_ids)
        assert {block["template_id"] for block in persisted.values()} == set(expected)

    def test_structured_table_content_is_editable_but_source_identity_and_caption_are_not(self):
        source = self.documents.section(PROJECT_ID, "section_objectives").content_blocks[2]
        inconsistent = deepcopy(source)
        inconsistent["rows"][1][1]["text"] = "仅修改顶层表示"
        with pytest.raises(
            RuntimeStoreError,
            match="top-level and structured table cell representations are inconsistent",
        ):
            MedicalWritingRuntimeRepository._validate_working_copy_blocks(
                [source],
                [inconsistent],
            )

        generated = MedicalWritingTableTemplateService().instantiate(
            "objectives_endpoints",
            title="新增研究目的与终点",
            instance_id="dual-representation-regression",
        )
        generated["rows"][1][1]["text"] = "仅修改模板表顶层表示"
        with pytest.raises(
            RuntimeStoreError,
            match="top-level and structured table cell representations are inconsistent",
        ):
            MedicalWritingRuntimeRepository._validate_working_copy_blocks(
                [source],
                [source, generated],
            )

        candidate = deepcopy(source)
        candidate["rows"][1][1]["text"] = "合法单元格修订"
        edited_cell_id = candidate["rows"][1][1]["cell_id"]
        for row in candidate["structured_table"]["rows"]:
            for cell in row["cells"]:
                if cell["cell_id"] == edited_cell_id:
                    cell["text"] = "合法单元格修订"
        candidate["title"] = "修订后的研究目的与终点"
        candidate["structured_table"]["title"] = "修订后的研究目的与终点"
        candidate["table_caption"]["review_status"] = "confirmed"
        candidate["structured_table"]["source_caption"][
            "review_status"
        ] = "confirmed"
        candidate["structured_table"]["source_note_metadata"][0][
            "review_status"
        ] = "confirmed"
        MedicalWritingRuntimeRepository._validate_working_copy_blocks(
            [source],
            [candidate],
        )

        for identity_field in (
            "block_id",
            "table_id",
            "source_locator",
            "source_kind",
            "body_order",
        ):
            changed = deepcopy(candidate)
            changed[identity_field] = f"changed-{identity_field}"
            try:
                MedicalWritingRuntimeRepository._validate_working_copy_blocks(
                    [source],
                    [changed],
                )
            except ValueError:
                pass
            else:
                raise AssertionError(f"source identity must be immutable: {identity_field}")

        changed_caption = deepcopy(candidate)
        changed_caption["table_caption"]["text"] = "被改写的来源表题"
        try:
            MedicalWritingRuntimeRepository._validate_working_copy_blocks(
                [source],
                [changed_caption],
            )
        except ValueError as exc:
            assert "caption text" in str(exc)
        else:
            raise AssertionError("source caption text must be immutable")

        for note_field in ("note_id", "source_locator", "target_cell_id"):
            changed_note = deepcopy(candidate)
            changed_note["structured_table"]["source_note_metadata"][0][
                note_field
            ] = f"changed-{note_field}"
            try:
                MedicalWritingRuntimeRepository._validate_working_copy_blocks(
                    [source],
                    [changed_note],
                )
            except ValueError as exc:
                assert f"note {note_field}" in str(exc)
            else:
                raise AssertionError(
                    f"source note metadata must be immutable: {note_field}"
                )

    def test_source_linked_text_and_matching_rich_text_are_jointly_editable(self):
        source = deepcopy(next(
            block
            for block in self.documents.section(PROJECT_ID, "section_objectives").content_blocks
            if block.get("block_type") != "table"
        ))
        source["rich_text"] = {
            "type": "paragraph",
            "attrs": {"stylePreset": "body"},
            "content": [{"type": "text", "text": source["text"]}],
        }
        candidate = deepcopy(source)
        candidate["text"] = "加粗后的研究目的"
        candidate["rich_text"] = {
            "type": "paragraph",
            "attrs": {"stylePreset": "body", "textAlign": "justify"},
            "content": [{
                "type": "text",
                "text": candidate["text"],
                "marks": [{"type": "bold"}],
            }],
        }

        MedicalWritingRuntimeRepository._validate_working_copy_blocks([source], [candidate])

        mismatched = deepcopy(candidate)
        mismatched["text"] = "与富文本不一致"
        with pytest.raises(ValueError, match="rich text must match block text"):
            MedicalWritingRuntimeRepository._validate_working_copy_blocks([source], [mismatched])

        changed_identity = deepcopy(candidate)
        changed_identity["source_locator"] = "docx:paragraph:changed"
        with pytest.raises(ValueError, match="source-linked source_locator"):
            MedicalWritingRuntimeRepository._validate_working_copy_blocks([source], [changed_identity])

    def test_legacy_working_copy_formatting_is_hydrated_without_changing_source_identity(self):
        source = deepcopy(next(
            block
            for block in self.documents.section(PROJECT_ID, "section_objectives").content_blocks
            if block.get("block_type") != "table"
        ))
        source["rich_text"] = {
            "type": "paragraph",
            "attrs": {"stylePreset": "body", "textAlign": "justify"},
            "content": [{
                "type": "text",
                "text": source["text"],
                "marks": [{"type": "bold"}],
            }],
        }
        unchanged_legacy = deepcopy(source)
        unchanged_legacy.pop("rich_text")
        changed_legacy = deepcopy(unchanged_legacy)
        changed_legacy["text"] = "旧工作副本已修订文本"

        unchanged = MedicalWritingRuntimeRepository._hydrate_legacy_working_copy_formatting(
            [source], [unchanged_legacy]
        )[0]
        changed = MedicalWritingRuntimeRepository._hydrate_legacy_working_copy_formatting(
            [source], [changed_legacy]
        )[0]

        assert unchanged["rich_text"] == source["rich_text"]
        assert changed["rich_text"]["attrs"] == source["rich_text"]["attrs"]
        assert changed["rich_text"]["content"] == [
            {"type": "text", "text": changed_legacy["text"]}
        ]
        assert changed["source_locator"] == source["source_locator"]
        MedicalWritingRuntimeRepository._validate_working_copy_blocks([source], [unchanged])
        MedicalWritingRuntimeRepository._validate_working_copy_blocks([source], [changed])
