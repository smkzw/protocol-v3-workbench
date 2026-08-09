from __future__ import annotations

import io

from docx import Document
import pytest

from packages.contracts.workbench_contracts.models import (
    MedicalWritingInterventionCrossObjectLink,
    MedicalWritingInterventionIpActionRule,
    MedicalWritingInterventionNonIpTreatmentRule,
    MedicalWritingInterventionRules,
    ProtocolDocument,
    ProtocolSection,
)
from services.api.app.medical_writing_document_exporter import (
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_intervention_rules_projection import (
    build_intervention_rules_projection_block,
    merge_intervention_rules_projection_block,
    render_intervention_rules_panel,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository


def _d001_like_rules() -> MedicalWritingInterventionRules:
    return MedicalWritingInterventionRules(
        authority="structured",
        ip_adjustment_policy="protocol_defined",
        ip_action_rules=[
            MedicalWritingInterventionIpActionRule(
                rule_id="ip_hold_for_systemic_rescue",
                action_kind="temporary_interruption",
                trigger="受试者接受全身系统性补救治疗",
                study_product_action="立即暂停试验用药品",
                retest_recovery="研究者与申办者确认可恢复给药",
                wait_period="补救治疗结束后至少 5 个半衰期",
                approvers=["研究者", "申办者"],
                linked_non_ip_rule_ids=["rescue_systemic"],
            )
        ],
        non_ip_treatment_rules=[
            MedicalWritingInterventionNonIpTreatmentRule(
                rule_id="rescue_systemic",
                rule_class="rescue",
                policy="rescue_policy",
                agent_or_category="全身系统性糖皮质激素",
                phase_applicability="治疗期",
                cm_dose_rule="按临床需要调整补救治疗剂量",
            ),
            MedicalWritingInterventionNonIpTreatmentRule(
                rule_id="cm_folic_acid",
                rule_class="allowed_cm",
                policy="allowed_if_stable",
                agent_or_category="叶酸",
                cm_dose_rule="1 mg/日；剂量变化记录为 CM 变化",
            ),
        ],
        cross_object_links=[
            MedicalWritingInterventionCrossObjectLink(
                link_id="link_rescue_hold",
                source_rule_id="rescue_systemic",
                source_kind="rescue",
                target_rule_id="ip_hold_for_systemic_rescue",
                action="hold",
                resume_condition="补救治疗结束后至少 5 个半衰期，且经研究者和申办者确认",
            )
        ],
    )


def _pnh_like_rules() -> MedicalWritingInterventionRules:
    return MedicalWritingInterventionRules(
        authority="structured",
        ip_adjustment_policy="no_planned_adjustment",
        no_planned_adjustment_statement="本研究没有计划调整剂量和/或中断剂量",
        ip_action_rules=[
            MedicalWritingInterventionIpActionRule(
                rule_id="ip_permanent_stop",
                action_kind="permanent_discontinuation",
                trigger="发生方案规定的永久停药条件",
                study_product_action="永久停止试验用药品",
                permanent_discontinuation_condition="继续给药的风险超过获益",
            ),
            MedicalWritingInterventionIpActionRule(
                rule_id="ip_taper",
                action_kind="discontinuation_taper",
                trigger="决定永久停药后",
                taper_steps=["按方案规定逐步递减", "完成递减后停止给药"],
            ),
            MedicalWritingInterventionIpActionRule(
                rule_id="ip_follow_up",
                action_kind="post_discontinuation_follow_up",
                trigger="永久停药后",
                study_product_action="按方案继续安全性随访",
            ),
        ],
    )


def test_d001_like_rescue_link_is_visible_but_cm_dose_does_not_enter_section_6_4():
    rules = _d001_like_rules()
    ip_text = render_intervention_rules_panel(rules, "6.4")
    rescue_text = render_intervention_rules_panel(rules, "6.9")
    cm_text = render_intervention_rules_panel(rules, "6.10")

    assert "立即暂停试验用药品" in ip_text
    assert "至少 5 个半衰期" in ip_text
    assert "全身系统性糖皮质激素" in ip_text
    assert "1 mg/日" not in ip_text
    assert "全身系统性糖皮质激素" in rescue_text
    assert "关联试验用药品处置：暂停试验用药品" in rescue_text
    assert "叶酸" not in rescue_text
    assert "叶酸" in cm_text
    assert "剂量变化记录为 CM 变化" in cm_text
    assert "立即暂停试验用药品" not in cm_text


def test_no_planned_adjustment_coexists_with_stop_taper_and_follow_up():
    text = render_intervention_rules_panel(_pnh_like_rules(), "6.4")
    assert text.startswith("本研究没有计划调整剂量和/或中断剂量。")
    assert "永久停药" in text
    assert "停药前递减" in text
    assert "停药后随访" in text


def test_empty_panel_is_not_silently_projected():
    with pytest.raises(ValueError, match="no projectable content"):
        render_intervention_rules_panel(_pnh_like_rules(), "6.10")


def test_projection_block_is_editable_auditable_and_repository_valid():
    rules = _pnh_like_rules()
    block = build_intervention_rules_projection_block(
        rules=rules,
        section_number="6.4",
        journey_id="journey_d001",
        journey_revision=8,
        body_order=30,
    )
    assert block["projection_panel"] == "ip_actions"
    assert block["source_journey_revision"] == 8
    assert len(block["source_intervention_rules_sha256"]) == 64
    assert len(block["projected_text_sha256"]) == 64
    assert len(block["projected_content_sha256"]) == 64
    assert block["rich_text"]["type"] == "doc"
    assert block["text"] == "\n".join(
        item["content"][0]["text"] for item in block["rich_text"]["content"]
    )
    source = {
        "block_id": "source_heading_6_4",
        "block_type": "heading",
        "text": "6.4 试验用药品剂量调整",
    }
    MedicalWritingRuntimeRepository._validate_working_copy_blocks(
        [source],
        [source, block],
        existing_blocks=[source],
        authorized_generated_blocks={block["block_id"]: block},
    )


def test_reprojection_preserves_manual_text_or_formatting_until_explicit_override():
    original = build_intervention_rules_projection_block(
        rules=_pnh_like_rules(),
        section_number="6.4",
        journey_id="journey_pnh",
        journey_revision=3,
        body_order=30,
    )
    manually_formatted = {**original, "rich_text": dict(original["rich_text"])}
    manually_formatted["rich_text"] = {
        **manually_formatted["rich_text"],
        "attrs": {"stylePreset": "note"},
    }
    refreshed = build_intervention_rules_projection_block(
        rules=_pnh_like_rules(),
        section_number="6.4",
        journey_id="journey_pnh",
        journey_revision=4,
        body_order=40,
    )
    with pytest.raises(ValueError, match="medical edits"):
        merge_intervention_rules_projection_block(
            [manually_formatted],
            refreshed,
        )
    merged, action = merge_intervention_rules_projection_block(
        [manually_formatted],
        refreshed,
        overwrite_medical_edits=True,
    )
    assert action == "updated"
    assert merged[0]["source_journey_revision"] == 4
    assert merged[0]["body_order"] == 30


def test_projected_no_adjustment_stop_and_follow_up_render_in_docx():
    block = build_intervention_rules_projection_block(
        rules=_pnh_like_rules(),
        section_number="6.4",
        journey_id="journey_pnh_docx",
        journey_revision=5,
        body_order=2,
    )
    document = ProtocolDocument(
        document_id="doc_pnh_projection",
        project_id="proj_pnh_projection",
        protocol_id="CMS-PNH",
        version="V0.1",
        sections=[
            ProtocolSection(
                section_id="section_6_4",
                document_id="doc_pnh_projection",
                heading="试验用药品剂量调整",
                section_number="6.4",
                content_blocks=[
                    {
                        "block_id": "heading_6_4",
                        "block_type": "heading",
                        "text": "6.4 试验用药品剂量调整",
                        "body_order": 1,
                    },
                    block,
                ],
            )
        ],
    )
    exported = export_medical_writing_document_docx(document, mode="draft_preview")
    rendered = Document(io.BytesIO(exported.content))
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    assert "本研究没有计划调整剂量和/或中断剂量" in text
    assert "永久停止试验用药品" in text
    assert "按方案规定逐步递减" in text
    assert "按方案继续安全性随访" in text
    assert "ip_permanent_stop" not in text
