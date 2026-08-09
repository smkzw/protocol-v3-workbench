from __future__ import annotations

import json
import sys
import tempfile
import time
import unittest
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook as openpyxl_load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.api.app.main import app  # noqa: E402
from services.api.app.safety_pv_manifest import (  # noqa: E402
    MY009_LISTING,
    RUX_274_ROOT,
    RUX_PV_ROOT,
    SafetyPackageConfig,
    SafetyPvManifestService,
)

RUX_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/"
    "RUX-03-002_列表_数据集_Excel_20250612_处理后.xlsx"
)


class SafetyPvManifestCacheTests(unittest.TestCase):
    def _workbook(self, path: Path, extra_subject: str = "") -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "AE"
        sheet.append(["USUBJID", "SITEID", "AETERM", "AESER", "AEREL"])
        sheet.append(["S001", "01", "头痛", "否", "可能相关"])
        if extra_subject:
            sheet.append([extra_subject, "02", "恶心", "否", "可能无关"])
        workbook.save(path)

    def test_listing_scan_cache_is_single_flight_and_invalidates_on_source_change(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            listing = Path(tmp) / "safety.xlsx"
            self._workbook(listing)
            config = SafetyPackageConfig(
                package_id="my009_uc_s1",
                package_label="test",
                project_code="TEST",
                source_root_label="test",
                package_role="test",
                listing_file=listing,
                document_paths=(),
            )
            service = SafetyPvManifestService(packages=[config])

            def counted_load(*args, **kwargs):
                time.sleep(0.02)
                return openpyxl_load_workbook(*args, **kwargs)

            with patch("services.api.app.safety_pv_manifest.load_workbook", side_effect=counted_load) as loader:
                with ThreadPoolExecutor(max_workers=2) as executor:
                    manifests = list(executor.map(lambda _: service.build_manifest("proj_my009_uc"), range(2)))
                self.assertEqual(1, loader.call_count)
                self.assertEqual(
                    manifests[0].packages[0].listing_domains,
                    manifests[1].packages[0].listing_domains,
                )

                self._workbook(listing, extra_subject="S002")
                changed = service.build_manifest("proj_my009_uc")
                self.assertEqual(2, loader.call_count)
                self.assertEqual(2, changed.packages[0].listing_domains[0].subject_count)

                service.build_manifest("proj_my009_uc", force_refresh=True)
                self.assertEqual(3, loader.call_count)

@unittest.skipUnless(MY009_LISTING.exists(), "MY009 safety listing path is unavailable")
class SafetyPvManifestTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        service = SafetyPvManifestService()
        cls.manifests = {
            "my009_uc_s1": service.build_manifest("proj_my009_uc"),
            "rux_03_002_pv": service.build_manifest("proj_rux_03_002"),
        }

    def package(self, package_id: str):
        return self.manifests[package_id].packages[0]

    def test_my009_manifest_reads_primary_comparison_listing_and_keeps_teae_boundary(self):
        package = self.package("my009_uc_s1")

        self.assertGreaterEqual(len(package.listing_domains), 50)
        self.assertGreaterEqual(len(package.documents), 7)
        ae_domain = next(domain for domain in package.listing_domains if domain.sheet_name == "AE")
        self.assertEqual(27, ae_domain.row_count)
        self.assertEqual("不良事件", ae_domain.domain_label)
        self.assertIn("AETERM", ae_domain.key_fields)
        self.assertIn("AESER", ae_domain.key_fields)

        ae_candidate = next(candidate for candidate in package.signal_candidates if candidate.signal_type == "ae_medical_review_candidate")
        self.assertIn("8 条AE", ae_candidate.title)
        self.assertIn("4 条提示可能相关", ae_candidate.title)
        self.assertEqual("待医学/PV确认", ae_candidate.confirmation_status)
        self.assertIn("不构成最终安全性结论", ae_candidate.observation)

        gate_labels = {gate.gate_label for gate in package.quality_gates}
        self.assertIn("AE与TEAE口径待核对", gate_labels)
        self.assertIn("安全性分母口径待确认", gate_labels)

    @unittest.skipUnless(
        RUX_LISTING.exists() and RUX_PV_ROOT.exists() and RUX_274_ROOT.exists(),
        "RUX safety source paths are unavailable",
    )
    def test_rux_manifest_integrates_real_listing_pv_plan_and_274(self):
        package = self.package("rux_03_002_pv")

        self.assertGreaterEqual(len(package.listing_domains), 40)
        self.assertTrue(any(document.document_type == "pv_plan" for document in package.documents))
        self.assertTrue(any(document.document_type == "clinical_safety_summary" for document in package.documents))
        domains = {domain.domain: domain for domain in package.listing_domains}
        self.assertIn("AE", domains)
        self.assertGreater(domains["AE"].subject_count, 0)
        self.assertIn("AETERM", domains["AE"].key_fields)
        self.assertIn("CM", domains)
        self.assertEqual("既往/合并用药（非试验用药）", domains["CM"].domain_label)
        self.assertIn("ECB", domains)
        self.assertEqual("试验药物/剂量调整", domains["ECB"].domain_label)

        candidates = {candidate.signal_type: candidate for candidate in package.signal_candidates}
        self.assertEqual(
            {
                "ae_medical_review_candidate",
                "lab_abnormality_medical_explanation",
                "ecg_abnormality_linkage_review",
                "conmed_indication_linkage_review",
                "pv_plan_safety_summary_alignment",
            },
            set(candidates),
        )
        self.assertIn("356 条AE", candidates["ae_medical_review_candidate"].title)
        self.assertIn("110 条提示可能相关", candidates["ae_medical_review_candidate"].title)
        self.assertIn("SAE字段阳性 5 条", candidates["ae_medical_review_candidate"].observation)
        self.assertIn("1293 条", candidates["lab_abnormality_medical_explanation"].title)
        self.assertIn("22 条", candidates["ecg_abnormality_linkage_review"].title)
        self.assertIn("3061 条", candidates["conmed_indication_linkage_review"].title)
        self.assertTrue(any(gate.gate_label == "RUX安全总结人群口径待确认" for gate in package.quality_gates))

    def test_safety_pv_manifest_endpoint_is_desensitized_and_does_not_claim_formal_pv_actions(self):
        response = TestClient(app).get("/api/projects/proj_my009_uc/safety-pv/manifest")
        self.assertEqual(200, response.status_code)
        payload = response.json()
        serialized = json.dumps(payload, ensure_ascii=False)

        self.assertEqual("安全信号与PV协同", payload["module_label"])
        self.assertEqual(1, payload["package_count"])
        self.assertGreaterEqual(payload["total_signal_candidates"], 4)
        for forbidden in [
            "/Users/",
            "第9环节",
            "阶段9",
            "Stage 9",
            "E2B",
            "监管clock",
            "case intake",
            "PV数据库写入",
            "正式PV判定",
            "final_seriousness",
            "final_expectedness",
            "final_causality",
            "pv_database_record_id",
        ]:
            self.assertNotIn(forbidden, serialized)


if __name__ == "__main__":
    unittest.main()
