from __future__ import annotations

import tempfile
import unittest
import io
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

from docx import Document

from packages.contracts.workbench_contracts import (
    AuditEvent,
    AiTaskRequest,
    AiTaskSourceRef,
    MedicalWritingRevisionApplyRequest,
    MedicalWritingRevisionRequest,
    MedicalWritingTableCellAnchor,
    MedicalWritingWorkingCopySaveRequest,
    ProtocolDocument,
    ProtocolSection,
    RevisionSuggestion,
    RevisionThread,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository
from services.api.app.medical_writing import MedicalWritingRevisionService
from services.api.app.medical_writing_revision_prompts import revision_task_context
from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_manifest import (
    D001_PROTOCOL_DOCX,
    RUX_PROTOCOL_DOCX,
)
from services.api.app.medical_writing_document_exporter import (
    export_medical_writing_document_docx,
)
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore
from services.api.app.medical_writing_tables import MedicalWritingTableService
from services.api.app.medical_writing_table_templates import (
    MedicalWritingTableTemplateService,
)
from services.api.app.ai_execution_policy import (
    AiExecutionPolicyDenied,
    AiExecutionPolicyResolver,
)
from services.api.app.ai_gateway import AiSourceRef, AiTaskSpec, PromptRegistry
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore
from services.api.app.sqlite_runtime_store import RuntimeStoreError, SqliteRuntimeStore


NOW = datetime(2026, 7, 14, 2, 0, tzinfo=timezone.utc)
PROJECTS = ("proj_table_ai_a", "proj_table_ai_b")


def table_block(project_id: str) -> dict:
    suffix = project_id.rsplit("_", 1)[-1]
    block_id = f"block_{suffix}_table"
    table_id = f"table_{suffix}_objectives"
    row_ids = [f"row_{suffix}_header", f"row_{suffix}_primary"]
    column_ids = [f"column_{suffix}_type", f"column_{suffix}_endpoint"]
    rows = []
    values = [["终点类型", "终点名称"], ["主要", f"{suffix.upper()}项目主要终点原文"]]
    for row_index, values_row in enumerate(values):
        rows.append(
            [
                {
                    "cell_id": f"cell_{suffix}_{row_index}_{column_index}",
                    "text": value,
                    "grid_column_index": column_index,
                    "row_span": 1,
                    "column_span": 1,
                    "hidden": False,
                    "style_role": "header" if row_index == 0 else "body",
                    "source_locator": f"docx:table:1:r{row_index}:c{column_index}",
                    "structure_row_id": row_ids[row_index],
                    "structure_column_id": column_ids[column_index],
                }
                for column_index, value in enumerate(values_row)
            ]
        )
    return {
        "block_id": block_id,
        "block_type": "table",
        "table_id": table_id,
        "title": "研究目的与终点",
        "source_locator": "docx:table:1",
        "header_row_count": 1,
        "column_count": 2,
        "rows": rows,
        "structured_table": {
            "schema_version": "structured_table_v1",
            "version": 0,
            "domain": "objectives_endpoints",
            "title": "研究目的与终点",
            "review_state": "ai_draft",
            "row_ids": row_ids,
            "row_labels": ["表头", "主要终点"],
            "row_style_roles": ["header", "body"],
            "column_ids": column_ids,
            "column_labels": ["终点类型", "终点名称"],
            "column_style_roles": ["header", "header"],
            "column_width_twips": [1800, 4800],
            "column_source_locators": ["docx:table:1:c0", "docx:table:1:c1"],
            "notes": [],
        },
    }


class TableDocumentService:
    def __init__(self):
        self.documents = {}
        for project_id in PROJECTS:
            suffix = project_id.rsplit("_", 1)[-1]
            document_id = f"document_{suffix}"
            section = ProtocolSection(
                section_id=f"section_{suffix}",
                document_id=document_id,
                heading="研究目的与终点",
                content_blocks=[
                    {
                        "block_id": f"block_{suffix}_paragraph",
                        "block_type": "paragraph",
                        "text": f"{suffix.upper()}项目原始方案研究目的。",
                        "source_locator": "docx:paragraph:1",
                    },
                    table_block(project_id),
                ],
            )
            self.documents[project_id] = ProtocolDocument(
                document_id=document_id,
                project_id=project_id,
                protocol_id=f"PROTOCOL-{suffix.upper()}",
                version="V1.0",
                sections=[section],
            )

    def document_session(self, project_id):
        return self.documents[project_id]

    def document_for_revision(self, project_id):
        return self.documents[project_id]

    def section(self, project_id, section_id):
        return next(
            section
            for section in self.documents[project_id].sections
            if section.section_id == section_id
        )


class RevisionRepositoryAdapter:
    """Expose the runtime revision API without the real-source-registry marker."""

    def __init__(self, repository):
        self.repository = repository

    def project(self, project_id):
        return self.repository.project(project_id)

    def protocol(self, project_id):
        return self.repository.protocol(project_id)

    def normalize_table_cell_revision(self, *args):
        return self.repository.normalize_table_cell_revision(*args)

    def commit_revision_submission(self, thread, audit_event):
        return self.repository.commit_revision_submission(thread, audit_event)


class TableRevisionProvider:
    provider_name = "buddy"
    model_name = "deepseek-v4-pro"

    def __init__(self):
        self.envelopes = []

    def run(self, envelope):
        self.envelopes.append(envelope)
        source = envelope.payload["allowed_sources"][0]
        source_text = str(source["text_preview"]).strip()
        punctuation_base = source_text.rstrip("。；，")
        return {
            "task_id": envelope.task_id,
            "task_type": envelope.task_type.value,
            "provider": self.provider_name,
            "model": self.model_name,
            "prompt_version": envelope.prompt_version,
            "input_source_ids": [
                item["source_id"] for item in envelope.payload["allowed_sources"]
            ],
            "forbidden_source_ids": envelope.payload["forbidden_source_ids"],
            "findings": [
                {
                    "finding_id": "finding_table_cell",
                    "status": "needs_medical_confirmation",
                    "title": "目标单元格可优化",
                    "source_id": source["source_id"],
                    "evidence_span_ids": ["span_table_cell"],
                }
            ],
            "evidence_spans": [
                {
                    "span_id": "span_table_cell",
                    "source_id": source["source_id"],
                    "locator": source["locator"],
                    "quote": source["text_preview"],
                }
            ],
            "uncertainties": [
                {"level": "medical_review", "description": "需医学经理确认。"}
            ],
            "needs_medical_confirmation": True,
            "schema_version": "ai_task_output_v0_1",
            "revision": {
                "proposal_text": source_text,
                "diff_patch": "cell replacement",
                "rationale": "保留原单元格事实，仅提供可审阅的格式候选。",
                "evidence_span_ids": ["span_table_cell"],
                "alternatives": [
                    {
                        "proposal_text": punctuation_base + "。",
                        "diff_patch": "alternative cell replacement 1",
                        "rationale": "保留单元格语义并统一句末标点。",
                        "evidence_span_ids": ["span_table_cell"],
                    },
                    {
                        "proposal_text": punctuation_base + "；",
                        "diff_patch": "alternative cell replacement 2",
                        "rationale": "保留单元格语义并提供分号衔接版本。",
                        "evidence_span_ids": ["span_table_cell"],
                    },
                ],
            },
        }


class MedicalWritingTableAiRevisionTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.documents = TableDocumentService()
        self.store = SqliteRuntimeStore(Path(self.tmpdir.name) / "runtime.sqlite3")
        self.repository = MedicalWritingRuntimeRepository(self.documents, self.store)

    def tearDown(self):
        self.tmpdir.cleanup()

    def normalize(self, project_id: str) -> dict:
        section = self.documents.documents[project_id].sections[0]
        block = section.content_blocks[1]
        selected = block["rows"][1][1]
        return self.repository.normalize_table_cell_revision(
            project_id,
            section.section_id,
            MedicalWritingTableCellAnchor(
                working_copy_revision=0,
                table_version=0,
                block_id=block["block_id"],
                table_id=block["table_id"],
                row_id=selected["structure_row_id"],
                column_id=selected["structure_column_id"],
                cell_id=selected["cell_id"],
                captured_text=selected["text"],
            ),
            selected["text"],
        )

    def accepted_thread(
        self,
        project_id: str,
        normalized: dict,
        *,
        token: str = "",
        proposal: str = "",
    ) -> RevisionThread:
        section = self.documents.documents[project_id].sections[0]
        identity = token or project_id
        proposal = proposal or f"{project_id} 医学批准后的主要终点表述"
        pending = RevisionThread(
            thread_id=f"thread_{identity}",
            project_id=project_id,
            document_id=section.document_id,
            section_id=section.section_id,
            anchor_type="table_cell",
            anchor_path=normalized["anchor_path"],
            selected_text=normalized["selected_text"],
            user_instruction="请优化该单元格的方案正文表述。",
            intent="medical_writing_revision",
            ai_run_id=f"run_{identity}",
            source_locator=normalized["source_locator"],
            table_cell_anchor=normalized["table_cell_anchor"],
            suggestions=[
                RevisionSuggestion(
                    suggestion_id=f"suggestion_{identity}",
                    proposal_text=proposal,
                    diff_patch="cell diff",
                    rationale="提高表述清晰度。",
                    evidence_span_ids=[f"span_{project_id}"],
                )
            ],
            status="candidate_ready",
            created_at=NOW,
        )
        self.store.commit_medical_writing_revision_submission(
            pending,
            AuditEvent(
                audit_id=f"audit_submit_{identity}",
                project_id=project_id,
                actor="medical_manager_test",
                action="medical_writing_revision_submitted",
                target_type="revision_thread",
                target_id=pending.thread_id,
                created_at=NOW,
            ),
        )
        accepted = pending.model_copy(deep=True)
        accepted.suggestions[0].user_decision = "accepted"
        accepted.status = "medically_approved"
        self.store.commit_medical_writing_revision_action(
            pending,
            accepted,
            AuditEvent(
                audit_id=f"audit_accept_{identity}",
                project_id=project_id,
                actor="medical_manager_test",
                action="medical_writing_revision_accept",
                target_type="revision_thread",
                target_id=pending.thread_id,
                created_at=NOW,
            ),
            None,
        )
        return self.repository.revision_thread(project_id, pending.thread_id)

    def test_normalization_builds_bounded_working_copy_context_not_evidence(self):
        for project_id in PROJECTS:
            with self.subTest(project_id=project_id):
                normalized = self.normalize(project_id)
                context = normalized["task_context"]
                self.assertEqual("working_copy_table_cell", context["context_type"])
                self.assertTrue(context["is_source_linked"])
                self.assertEqual(normalized["selected_text"], context["selected_text"])
                self.assertLessEqual(len(context["column_window"]), 12)
                self.assertIn("column_semantic_role", context)
                self.assertTrue(
                    all("semantic_role" in item for item in context["column_window"])
                )
                self.assertEqual(64, len(context["block_hash"]))
                self.assertTrue(normalized["source_locator"].startswith("docx:"))

    def test_ai_gateway_accepts_table_context_but_keeps_it_out_of_evidence_sources(
        self,
    ):
        project_id = PROJECTS[0]
        context = revision_task_context(
            "medical_writing_revision",
            self.normalize(project_id)["task_context"],
        )
        source = AiTaskSourceRef(
            source_id="source_original_protocol_cell",
            source_type="protocol_docx_selection",
            title="原始方案表格单元格",
            locator="docx:table:1:r1:c1",
            text_preview="A项目主要终点原文",
            project_id=project_id,
            module="medical_writing",
        )
        resolver = AiExecutionPolicyResolver(
            deployment_profile="approved_private_documents",
            provider_name="buddy",
            model_name="deepseek-v4-pro",
            test_only_provider_injection=True,
        )
        resolution = resolver.resolve_internal(
            project_id,
            AiTaskRequest(
                module="medical_writing",
                task_type="medical_writing_revision",
                prompt_version="medical_writing_revision_v1_4",
                allowed_sources=[source],
                user_instruction="优化目标单元格。",
                task_context=context,
            ),
        )
        envelope = PromptRegistry().build(
            AiTaskSpec(
                task_id="task_table_cell_revision",
                task_type=resolution.task_type,
                prompt_version=resolution.prompt_version,
                allowed_sources=[
                    AiSourceRef(
                        **source.model_dump(
                            exclude={"project_id", "module", "source_entry_id"}
                        )
                    )
                ],
                user_instruction=resolution.user_instruction,
                provider_name=resolution.provider_name,
                model_name=resolution.model_name,
                task_context=resolution.task_context,
            )
        )
        self.assertEqual(context, envelope.payload["task_context"])
        self.assertEqual(1, len(envelope.payload["allowed_sources"]))
        self.assertNotIn("working_copy_table_cell", envelope.payload["allowed_sources"])
        self.assertIn("不是原始方案或临床事实证据", envelope.system_prompt)

        invalid = dict(context)
        invalid["client_freeform_path"] = "/another/project/table"
        with self.assertRaisesRegex(AiExecutionPolicyDenied, "keys must match"):
            resolver.resolve_internal(
                project_id,
                AiTaskRequest(
                    module="medical_writing",
                    task_type="medical_writing_revision",
                    prompt_version="medical_writing_revision_v1_4",
                    allowed_sources=[source],
                    user_instruction="优化目标单元格。",
                    task_context=invalid,
                ),
            )

    def test_revision_service_submits_server_normalized_cell_context_to_independent_ai(
        self,
    ):
        project_id = PROJECTS[0]
        section = self.documents.documents[project_id].sections[0]
        block = section.content_blocks[1]
        selected = block["rows"][1][1]
        provider = TableRevisionProvider()
        runner = AiTaskRunner(
            self.repository,
            AiTaskStore(Path(self.tmpdir.name) / "table-ai-runs.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name="buddy",
                model_name="deepseek-v4-pro",
                test_only_provider_injection=True,
            ),
        )
        service = MedicalWritingRevisionService(
            RevisionRepositoryAdapter(self.repository),
            runner,
        )
        result = service.submit_revision(
            project_id,
            MedicalWritingRevisionRequest(
                document_id=section.document_id,
                section_id=section.section_id,
                anchor_type="table_cell",
                selected_text=selected["text"],
                table_cell_anchor=MedicalWritingTableCellAnchor(
                    working_copy_revision=0,
                    table_version=0,
                    block_id=block["block_id"],
                    table_id=block["table_id"],
                    row_id=selected["structure_row_id"],
                    column_id=selected["structure_column_id"],
                    cell_id=selected["cell_id"],
                    captured_text=selected["text"],
                ),
                user_instruction="优化目标单元格。",
            ),
        )
        self.assertEqual("table_cell", result.thread.anchor_type)
        self.assertEqual(selected["cell_id"], result.thread.table_cell_anchor.cell_id)
        context = provider.envelopes[0].payload["task_context"]
        self.assertEqual("working_copy_table_cell", context["context_type"])
        self.assertEqual(selected["cell_id"], context["cell_id"])
        self.assertEqual(
            "protocol_section_selection",
            provider.envelopes[0].payload["allowed_sources"][0]["source_type"],
        )
        self.assertEqual(
            context,
            runner.get(project_id, result.thread.ai_run_id).task_context_summary,
        )

    def test_approved_cell_revision_updates_both_representations_and_is_idempotent(
        self,
    ):
        for project_id in PROJECTS:
            with self.subTest(project_id=project_id):
                normalized = self.normalize(project_id)
                thread = self.accepted_thread(project_id, normalized)
                request = MedicalWritingRevisionApplyRequest(
                    expected_working_copy_revision=0,
                    actor="medical_manager_test",
                    idempotency_key=f"apply-{project_id}",
                )
                result = self.repository.apply_approved_revision_to_working_copy(
                    project_id,
                    thread.section_id,
                    thread.thread_id,
                    request,
                )
                self.assertEqual(1, result.working_copy.revision)
                block = next(
                    item
                    for item in result.working_copy.content_blocks
                    if item.get("block_id") == normalized["table_cell_anchor"].block_id
                )
                cell_id = normalized["table_cell_anchor"].cell_id
                top_text = next(
                    cell["text"]
                    for row in block["rows"]
                    for cell in row
                    if cell["cell_id"] == cell_id
                )
                structured_text = next(
                    cell["text"]
                    for row in block["structured_table"]["rows"]
                    for cell in row["cells"]
                    if cell["cell_id"] == cell_id
                )
                self.assertEqual(thread.suggestions[0].proposal_text, top_text)
                self.assertEqual(top_text, structured_text)
                self.assertEqual(1, block["structured_table"]["version"])
                self.assertEqual("table_cell", result.audit_event.detail["anchor_type"])
                self.assertEqual(cell_id, result.audit_event.detail["cell_id"])
                replay = self.repository.apply_approved_revision_to_working_copy(
                    project_id,
                    thread.section_id,
                    thread.thread_id,
                    request,
                )
                self.assertEqual(result.working_copy, replay.working_copy)
                source_block = self.documents.section(
                    project_id, thread.section_id
                ).content_blocks[1]
                self.assertEqual(
                    normalized["selected_text"], source_block["rows"][1][1]["text"]
                )

    def test_promoted_generated_templates_complete_sequential_ai_cell_revision_cycles(
        self,
    ):
        project_id = PROJECTS[0]
        section = self.documents.documents[project_id].sections[0]
        template_service = MedicalWritingTableTemplateService()
        template_ids = (
            "sample_size_assumptions",
            "analysis_sets",
            "pk_immunogenicity_schedule",
        )
        generated_blocks = [
            template_service.instantiate(
                template_id,
                instance_id=f"ai-revision-{template_id}",
            )
            for template_id in template_ids
        ]
        initial = self.repository.working_copy(project_id, section.section_id)
        saved = self.repository.save_working_copy(
            project_id,
            section.section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=section.document_id,
                expected_revision=initial.revision,
                content_blocks=[*initial.content_blocks, *generated_blocks],
                actor="medical_manager_test",
                idempotency_key="save-promoted-generated-tables",
            ),
        )
        self.assertEqual(1, saved.revision)

        expected_text_by_cell = {}
        untouched_cells = {}
        current_revision = saved.revision
        for template_id, original_block in zip(template_ids, generated_blocks):
            current = self.repository.working_copy(project_id, section.section_id)
            block = next(
                item
                for item in current.content_blocks
                if item.get("block_id") == original_block["block_id"]
            )
            table = MedicalWritingTableService().from_table_block(block)
            target = next(
                cell
                for row in table.rows[1:]
                for cell in row.cells
                if str(cell.text).strip()
            )
            row = next(item for item in table.rows if item.row_id == target.row_id)
            sibling = next(
                cell
                for candidate_row in table.rows
                for cell in candidate_row.cells
                if cell.cell_id != target.cell_id and str(cell.text).strip()
            )
            untouched_cells[sibling.cell_id] = sibling.text
            normalized = self.repository.normalize_table_cell_revision(
                project_id,
                section.section_id,
                MedicalWritingTableCellAnchor(
                    working_copy_revision=current_revision,
                    table_version=table.version,
                    block_id=table.block_id,
                    table_id=table.table_id,
                    row_id=row.row_id,
                    column_id=target.column_id,
                    cell_id=target.cell_id,
                    captured_text=target.text,
                ),
                target.text,
            )
            proposal = f"{target.text}（{template_id}医学批准修订）"
            thread = self.accepted_thread(
                project_id,
                normalized,
                token=f"generated_{template_id}",
                proposal=proposal,
            )
            result = self.repository.apply_approved_revision_to_working_copy(
                project_id,
                section.section_id,
                thread.thread_id,
                MedicalWritingRevisionApplyRequest(
                    expected_working_copy_revision=current_revision,
                    actor="medical_manager_test",
                    idempotency_key=f"apply-generated-{template_id}",
                ),
            )
            current_revision += 1
            self.assertEqual(current_revision, result.working_copy.revision)
            changed_block = next(
                item
                for item in result.working_copy.content_blocks
                if item.get("block_id") == table.block_id
            )
            top_text = next(
                cell["text"]
                for values in changed_block["rows"]
                for cell in values
                if cell["cell_id"] == target.cell_id
            )
            structured_text = next(
                cell["text"]
                for values in changed_block["structured_table"]["rows"]
                for cell in values["cells"]
                if cell["cell_id"] == target.cell_id
            )
            self.assertEqual(proposal, top_text)
            self.assertEqual(top_text, structured_text)
            self.assertEqual(1, changed_block["structured_table"]["version"])
            expected_text_by_cell[target.cell_id] = proposal

        reloaded = self.repository.working_copy(project_id, section.section_id)
        all_cells = {
            cell["cell_id"]: cell["text"]
            for block in reloaded.content_blocks
            if block.get("block_type") == "table"
            for row in block.get("rows", [])
            for cell in row
        }
        for cell_id, expected in expected_text_by_cell.items():
            self.assertEqual(expected, all_cells[cell_id])
        for cell_id, expected in untouched_cells.items():
            self.assertEqual(expected, all_cells[cell_id])

        assembled = self.repository.assemble_document_for_export(
            project_id, "draft_preview"
        )
        exported = export_medical_writing_document_docx(assembled, mode="draft_preview")
        exported_docx = Document(io.BytesIO(exported.content))
        exported_text = "\n".join(
            cell.text
            for table in exported_docx.tables
            for row in table.rows
            for cell in row.cells
        )
        for expected in expected_text_by_cell.values():
            self.assertIn(expected, exported_text)
        self.assertEqual([], self.store.verify_audit_chain(project_id))

    def test_changed_table_block_is_rejected_even_with_current_global_revision(self):
        project_id = PROJECTS[0]
        normalized = self.normalize(project_id)
        thread = self.accepted_thread(project_id, normalized)
        current = self.repository.working_copy(project_id, thread.section_id)
        blocks = deepcopy(current.content_blocks)
        table_index = next(
            index
            for index, block in enumerate(blocks)
            if block.get("block_type") == "table"
        )
        table_service = MedicalWritingTableService()
        table = table_service.from_table_block(blocks[table_index])
        changed = table_service.apply_operations(
            table,
            [
                {
                    "op": "edit_cell",
                    "cell_id": normalized["table_cell_anchor"].cell_id,
                    "text": "并发修改",
                }
            ],
            expected_version=0,
        )
        blocks[table_index] = table_service.to_table_block(changed)
        saved = self.repository.save_working_copy(
            project_id,
            thread.section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=thread.document_id,
                expected_revision=0,
                content_blocks=blocks,
                actor="medical_manager_other",
                idempotency_key="concurrent-table-save-a",
            ),
        )
        with self.assertRaisesRegex(
            RuntimeStoreError, "changed after AI revision submission"
        ):
            self.repository.apply_approved_revision_to_working_copy(
                project_id,
                thread.section_id,
                thread.thread_id,
                MedicalWritingRevisionApplyRequest(
                    expected_working_copy_revision=saved.revision,
                    actor="medical_manager_test",
                    idempotency_key="apply-stale-table-a",
                ),
            )

    def test_wrong_hierarchy_and_cross_project_access_are_rejected(self):
        normalized = self.normalize(PROJECTS[0])
        bad_anchor = normalized["table_cell_anchor"].model_copy(
            update={"column_id": "column_from_another_table"}
        )
        section = self.documents.documents[PROJECTS[0]].sections[0]
        with self.assertRaisesRegex(ValueError, "does not resolve exactly once"):
            self.repository.normalize_table_cell_revision(
                PROJECTS[0],
                section.section_id,
                bad_anchor,
                normalized["selected_text"],
            )
        thread = self.accepted_thread(PROJECTS[0], normalized)
        with self.assertRaises(KeyError):
            self.repository.apply_approved_revision_to_working_copy(
                PROJECTS[1],
                self.documents.documents[PROJECTS[1]].sections[0].section_id,
                thread.thread_id,
                MedicalWritingRevisionApplyRequest(
                    expected_working_copy_revision=0,
                    actor="medical_manager_test",
                    idempotency_key="cross-project-apply",
                ),
            )


@unittest.skipUnless(
    RUX_PROTOCOL_DOCX.exists() and D001_PROTOCOL_DOCX.exists(),
    "real protocol DOCX fixtures are unavailable",
)
class RealMedicalWritingTableAiRevisionTests(unittest.TestCase):
    def test_rux_and_d001_table_cells_reach_independent_ai_with_registered_original_evidence(
        self,
    ):
        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            documents = MedicalWritingDocumentService()
            repository = MedicalWritingRuntimeRepository(
                documents,
                SqliteRuntimeStore(root / "real-table-ai-service.sqlite3"),
            )
            registry = SourceRegistryService(
                SourceRegistryStore(root / "source-registry.jsonl"),
                allowed_roots=[RUX_PROTOCOL_DOCX.parent, D001_PROTOCOL_DOCX.parent],
            )
            provider = TableRevisionProvider()
            runner = AiTaskRunner(
                repository,
                AiTaskStore(root / "real-table-ai-runs.jsonl"),
                provider_factory=lambda resolution: provider,
                policy_resolver=AiExecutionPolicyResolver(
                    deployment_profile="approved_private_documents",
                    provider_name="buddy",
                    model_name="deepseek-v4-pro",
                    test_only_provider_injection=True,
                ),
            )
            service = MedicalWritingRevisionService(
                repository,
                runner,
                source_registry=registry,
                protocol_source_paths={
                    "proj_rux_03_002": RUX_PROTOCOL_DOCX,
                    "proj_d001": D001_PROTOCOL_DOCX,
                },
            )

            for project_id in ("proj_rux_03_002", "proj_d001"):
                session = documents.document_session(project_id)
                section = next(
                    documents.section(project_id, summary.section_id)
                    for summary in session.sections
                    if any(
                        block.get("block_type") == "table"
                        for block in documents.section(
                            project_id, summary.section_id
                        ).content_blocks
                    )
                )
                block = next(
                    item
                    for item in section.content_blocks
                    if item.get("block_type") == "table"
                )
                table = MedicalWritingTableService().from_table_block(block)
                cell = next(
                    candidate
                    for row in table.rows
                    for candidate in row.cells
                    if str(candidate.text).strip() and candidate.source_locator
                )
                row = next(item for item in table.rows if item.row_id == cell.row_id)
                result = service.submit_revision(
                    project_id,
                    MedicalWritingRevisionRequest(
                        document_id=section.document_id,
                        section_id=section.section_id,
                        anchor_type="table_cell",
                        selected_text=cell.text,
                        table_cell_anchor=MedicalWritingTableCellAnchor(
                            working_copy_revision=0,
                            table_version=table.version,
                            block_id=table.block_id,
                            table_id=table.table_id,
                            row_id=row.row_id,
                            column_id=cell.column_id,
                            cell_id=cell.cell_id,
                            captured_text=cell.text,
                        ),
                        user_instruction="基于原始方案证据优化目标单元格。",
                    ),
                )
                envelope = provider.envelopes[-1]
                self.assertEqual(
                    "working_copy_table_cell",
                    envelope.payload["task_context"]["context_type"],
                )
                self.assertEqual(
                    cell.cell_id, envelope.payload["task_context"]["cell_id"]
                )
                self.assertIn(
                    envelope.payload["allowed_sources"][0]["source_type"],
                    {"protocol_docx_selection", "protocol_docx_table_cell_selection"},
                )
                self.assertTrue(result.thread.source_entry_id.startswith("src_"))
                self.assertEqual("table_cell", result.thread.anchor_type)

            self.assertEqual(2, len(provider.envelopes))

    def test_rux_and_d001_raw_protocol_tables_complete_the_same_cell_revision_cycle(
        self,
    ):
        with tempfile.TemporaryDirectory() as tmpdir:
            documents = MedicalWritingDocumentService()
            repository = MedicalWritingRuntimeRepository(
                documents,
                SqliteRuntimeStore(Path(tmpdir) / "real-table-ai.sqlite3"),
            )
            applied = {}
            for project_id in ("proj_rux_03_002", "proj_d001"):
                session = documents.document_session(project_id)
                section = next(
                    documents.section(project_id, summary.section_id)
                    for summary in session.sections
                    if any(
                        block.get("block_type") == "table"
                        for block in documents.section(
                            project_id, summary.section_id
                        ).content_blocks
                    )
                )
                block = next(
                    item
                    for item in section.content_blocks
                    if item.get("block_type") == "table"
                )
                table = MedicalWritingTableService().from_table_block(block)
                cell = next(
                    candidate
                    for row in table.rows
                    for candidate in row.cells
                    if str(candidate.text).strip()
                    and not any(
                        raw.get("cell_id") == candidate.cell_id and raw.get("hidden")
                        for raw_row in block["rows"]
                        for raw in raw_row
                    )
                )
                row = next(item for item in table.rows if item.row_id == cell.row_id)
                normalized = repository.normalize_table_cell_revision(
                    project_id,
                    section.section_id,
                    MedicalWritingTableCellAnchor(
                        working_copy_revision=0,
                        table_version=table.version,
                        block_id=table.block_id,
                        table_id=table.table_id,
                        row_id=row.row_id,
                        column_id=cell.column_id,
                        cell_id=cell.cell_id,
                        captured_text=cell.text,
                    ),
                    cell.text,
                )
                thread_id = f"thread_real_table_{project_id}"
                pending = RevisionThread(
                    thread_id=thread_id,
                    project_id=project_id,
                    document_id=section.document_id,
                    section_id=section.section_id,
                    anchor_type="table_cell",
                    anchor_path=normalized["anchor_path"],
                    selected_text=normalized["selected_text"],
                    user_instruction="真实方案表格单元格修订测试。",
                    intent="medical_writing_revision",
                    ai_run_id=f"run_real_{project_id}",
                    source_locator=normalized["source_locator"],
                    table_cell_anchor=normalized["table_cell_anchor"],
                    suggestions=[
                        RevisionSuggestion(
                            suggestion_id=f"suggestion_real_{project_id}",
                            proposal_text=f"{cell.text}（医学修订候选）",
                            diff_patch="real cell replacement",
                            rationale="验证真实原始方案表格闭环。",
                            evidence_span_ids=[f"span_real_{project_id}"],
                        )
                    ],
                    status="candidate_ready",
                    created_at=NOW,
                )
                repository.runtime_store.commit_medical_writing_revision_submission(
                    pending,
                    AuditEvent(
                        audit_id=f"audit_real_submit_{project_id}",
                        project_id=project_id,
                        actor="medical_manager_test",
                        action="medical_writing_revision_submitted",
                        target_type="revision_thread",
                        target_id=thread_id,
                        created_at=NOW,
                    ),
                )
                accepted = pending.model_copy(deep=True)
                accepted.suggestions[0].user_decision = "accepted"
                accepted.status = "medically_approved"
                repository.runtime_store.commit_medical_writing_revision_action(
                    pending,
                    accepted,
                    AuditEvent(
                        audit_id=f"audit_real_accept_{project_id}",
                        project_id=project_id,
                        actor="medical_manager_test",
                        action="medical_writing_revision_accept",
                        target_type="revision_thread",
                        target_id=thread_id,
                        created_at=NOW,
                    ),
                    None,
                )
                approved = repository.revision_thread(project_id, thread_id)
                result = repository.apply_approved_revision_to_working_copy(
                    project_id,
                    section.section_id,
                    thread_id,
                    MedicalWritingRevisionApplyRequest(
                        expected_working_copy_revision=0,
                        actor="medical_manager_test",
                        idempotency_key=f"apply-real-{project_id}",
                    ),
                )
                changed_block = next(
                    item
                    for item in result.working_copy.content_blocks
                    if item.get("block_id") == table.block_id
                )
                changed_text = next(
                    item["text"]
                    for values in changed_block["rows"]
                    for item in values
                    if item["cell_id"] == cell.cell_id
                )
                self.assertEqual(approved.suggestions[0].proposal_text, changed_text)
                self.assertEqual(1, result.working_copy.revision)
                self.assertEqual(
                    [], repository.runtime_store.verify_audit_chain(project_id)
                )
                assembled = repository.assemble_document_for_export(
                    project_id, "draft_preview"
                )
                exported = export_medical_writing_document_docx(
                    assembled, mode="draft_preview"
                )
                exported_docx = Document(io.BytesIO(exported.content))
                exported_text = "\n".join(
                    cell.text
                    for exported_table in exported_docx.tables
                    for exported_row in exported_table.rows
                    for cell in exported_row.cells
                )
                self.assertIn(changed_text, exported_text)
                applied[project_id] = (
                    result.working_copy.working_copy_id,
                    changed_block["table_id"],
                )

            self.assertNotEqual(applied["proj_rux_03_002"], applied["proj_d001"])


if __name__ == "__main__":
    unittest.main()
