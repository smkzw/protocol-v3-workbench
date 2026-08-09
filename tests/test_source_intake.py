from __future__ import annotations

import io
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.api.app.ai_gateway import (  # noqa: E402
    AiGatewayConfigurationError,
    AiTaskSpec,
    AiTaskType,
    PromptRegistry,
)
from services.api.app.protocol_text_extractor import parse_protocol_docx  # noqa: E402
from services.api.app.raw_subject_bundle import inventory_subject_bundle  # noqa: E402
from services.api.app.source_intake import (  # noqa: E402
    SourceRegistryService,
    SourceRegistryStore,
    protocol_document_to_ai_sources,
    raw_subject_inventory_to_ai_sources,
)


MGK10_PROTOCOL = Path(
    "/Users/smkzw/Documents/康哲项目资料/MG-K10/SAR/4. Protocol/MG-K10-SAR-001_临床研究方案_ V2.1_20250919_clean版 .docx"
)
D001_SUBJECT_DIR = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/全量-入组/白铭江/0316-SA07005"
)


class SourceIntakeTests(unittest.TestCase):
    def test_investigator_brochure_is_registered_with_role_and_indication_evidence(self):
        document = Document()
        document.add_heading("研究者手册", level=1)
        document.add_paragraph("试验药物 CMS-RA-001 拟用于类风湿关节炎。")
        document.add_heading("非临床研究", level=1)
        document.add_paragraph("已完成药理、毒理和安全药理研究。")
        document.add_heading("临床试验经验", level=1)
        document.add_paragraph("现有安全性信息和药代动力学结果支持后续开发。")
        buffer = io.BytesIO()
        document.save(buffer)

        with tempfile.TemporaryDirectory() as tmpdir:
            root = Path(tmpdir)
            service = SourceRegistryService(
                SourceRegistryStore(root / "sources.jsonl"),
                artifact_root=root / "artifacts",
            )
            result = service.register_medical_writing_document(
                "project_ra",
                "CMS-RA-001_IB_V1.0.docx",
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                buffer.getvalue(),
                document_role="investigator_brochure",
                expected_indication="类风湿关节炎",
            )

            validation = result.entry.metadata["content_validation"]
            self.assertEqual("investigator_brochure", result.entry.source_kind)
            self.assertEqual("matched", validation["role_status"])
            self.assertEqual("matched", validation["indication_status"])
            self.assertTrue(validation["usable_without_override"])
            self.assertTrue(result.spans)
            self.assertTrue(
                any("安全性信息" in span.text_preview for span in result.spans)
            )
            self.assertTrue(list((root / "artifacts").rglob("*.docx")))

    @unittest.skipUnless(MGK10_PROTOCOL.exists(), "MG-K10 raw protocol not available on this machine")
    def test_protocol_sources_feed_protocol_rule_extraction_prompt_without_using_reference_outputs(self):
        document = parse_protocol_docx(MGK10_PROTOCOL.name, MGK10_PROTOCOL.read_bytes())
        sources = protocol_document_to_ai_sources(document, "mgk10_v21", max_spans=12)

        spec = AiTaskSpec(
            task_id="task_protocol_rules_mgk10_v21",
            task_type=AiTaskType.PROTOCOL_RULE_EXTRACTION,
            prompt_version="protocol_rule_extraction_v0_1",
            allowed_sources=sources,
            forbidden_source_ids=["criteria_rules_md", "timeline_json_reference", "patient_profile_v10_dataset"],
        )
        envelope = PromptRegistry().build(spec)

        self.assertEqual(12, len(envelope.payload["allowed_sources"]))
        self.assertTrue(all(source["source_id"].startswith("mgk10_v21_") for source in envelope.payload["allowed_sources"]))
        self.assertIn("criteria_rules_md", envelope.payload["forbidden_source_ids"])
        self.assertTrue(all(source["source_type"].startswith("protocol_docx_") for source in envelope.payload["allowed_sources"]))
        self.assertTrue(any("MG-K10" in source["text_preview"] for source in envelope.payload["allowed_sources"]))

    @unittest.skipUnless(D001_SUBJECT_DIR.exists(), "D001 raw subject directory not available on this machine")
    def test_raw_subject_inventory_metadata_cannot_bypass_trusted_eligibility_context(self):
        inventory = inventory_subject_bundle(D001_SUBJECT_DIR)
        sources = raw_subject_inventory_to_ai_sources(inventory, "d001_sa07005")

        spec = AiTaskSpec(
            task_id="task_eligibility_d001_sa07005",
            task_type=AiTaskType.ELIGIBILITY_RULE_REVIEW,
            prompt_version="eligibility_rule_review_v0_1",
            allowed_sources=sources,
            forbidden_source_ids=["legacy_evidence_bundle", "legacy_review_report"],
        )
        with self.assertRaisesRegex(
            AiGatewayConfigurationError,
            "eligibility task_context keys must match",
        ):
            PromptRegistry().build(spec)

        self.assertEqual(inventory.total_files, len(sources))
        self.assertTrue(any(source.source_type == "raw_subject_pdf" for source in sources))
        self.assertTrue(any("needs_ocr_vlm=True" in source.text_preview for source in sources))
        combined_prompt_text = "\n".join(
            f"{source.title}\n{source.locator}\n{source.text_preview}"
            for source in sources
        )
        for leaked_label in ["合格", "筛败", "T-SPOT"]:
            self.assertNotIn(leaked_label, combined_prompt_text)
        self.assertIn("legacy_evidence_bundle", spec.forbidden_source_ids)


if __name__ == "__main__":
    unittest.main()
