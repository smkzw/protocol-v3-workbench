from __future__ import annotations

import io
import sys
import unittest
import zipfile
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from services.api.app.listing_file_parser import parse_listing_file  # noqa: E402


MGK10_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/MG-K10/SAR/13. CFDI核查/自查/评分SDV/【锁库后Data Listing】MG-K10-SAR-001_FormExcelAllVersion_202601201126.xlsx"
)
RUX_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/RUX-03-002_列表_数据集_Excel_20250612_处理后.xlsx"
)
RUX_SUBJECT_REPORT = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/RUX-03-002_受试者报表_20250613.xls"
)
RUX_RAW_DIMENSION_DEFECT = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/"
    "RUX-03-002-自查文件包-20260107/1-RUX-03-002_列表_数据集_压缩包_20250612/"
    "RUX-03-002_列表_数据集_Excel_20250612.xlsx"
)


class ListingFileParserTests(unittest.TestCase):
    def test_parse_xlsx_keeps_sheets_dates_and_duplicate_headers(self):
        workbook = openpyxl.Workbook()
        ae = workbook.active
        ae.title = "AE"
        ae.append(["SUBJID", "VISIT", "AETERM", "AESTDTC", "SUBJID"])
        ae.append(["10008", "W1", "头痛", date(2026, 7, 13), "010-10008"])
        ae.append([None, None, None, None, None])
        cm = workbook.create_sheet("CM")
        cm.append(["SUBJID", "CMTRT", "CMSTDTC", "CMENDTC"])
        cm.append(["10008", "氯雷他定", "2026-06-30", "2026-07-05"])

        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        sheets = parse_listing_file("edc_listing_batch_004.xlsx", buffer.getvalue())

        self.assertEqual(["AE", "CM"], [sheet.sheet_name for sheet in sheets])
        self.assertEqual(1, len(sheets[0].rows))
        self.assertEqual("2026-07-13", sheets[0].rows[0]["AESTDTC"])
        self.assertEqual("010-10008", sheets[0].rows[0]["SUBJID__2"])
        self.assertEqual("氯雷他定", sheets[1].rows[0]["CMTRT"])
        self.assertEqual([2], sheets[0].row_numbers)
        self.assertEqual([2], sheets[1].row_numbers)
        self.assertEqual(
            ["SUBJID", "VISIT", "AETERM", "AESTDTC", "SUBJID__2"],
            sheets[0].headers,
        )

    def test_parse_xlsx_preserves_header_only_domain_schema(self):
        workbook = openpyxl.Workbook()
        ae = workbook.active
        ae.title = "AE"
        ae.append(["SUBJID", "AETERM", "AESTDTC"])
        dm = workbook.create_sheet("DM")
        dm.append(["SUBJID", "SEX"])
        dm.append(["S01001", "女"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        sheets = parse_listing_file("header_only_domain.xlsx", buffer.getvalue())
        sheet_index = {sheet.sheet_name: sheet for sheet in sheets}

        self.assertEqual([], sheet_index["AE"].rows)
        self.assertEqual(["SUBJID", "AETERM", "AESTDTC"], sheet_index["AE"].headers)
        self.assertEqual(1, len(sheet_index["DM"].rows))

    def test_parse_csv_uses_filename_as_sheet_name_and_utf8_sig(self):
        content = "SUBJID,VISIT,AETERM,SUBJID\n06021,W2,轻度鼻出血,010-06021\n".encode("utf-8-sig")

        sheets = parse_listing_file("AE导出.csv", content)

        self.assertEqual(1, len(sheets))
        self.assertEqual("AE导出", sheets[0].sheet_name)
        self.assertEqual("06021", sheets[0].rows[0]["SUBJID"])
        self.assertEqual("010-06021", sheets[0].rows[0]["SUBJID__2"])
        self.assertEqual("轻度鼻出血", sheets[0].rows[0]["AETERM"])
        self.assertEqual([2], sheets[0].row_numbers)

    def test_parse_csv_prefers_semantic_code_in_double_parenthesized_headers(self):
        content = (
            "受试者编号(SUBJID)(RAW),开始日期(AESTDAT)(RAW),结束日期(AEENDAT)(RAW)\n"
            "S01003,2026-03-01,2026-03-03\n"
        ).encode("utf-8-sig")

        row = parse_listing_file("AE.csv", content)[0].rows[0]

        self.assertEqual("S01003", row["SUBJID"])
        self.assertEqual("2026-03-01", row["AESTDAT"])
        self.assertEqual("2026-03-03", row["AEENDAT"])
        self.assertNotIn("RAW", row)

    def test_parse_csv_deduplicates_double_parenthesized_semantic_codes(self):
        content = (
            "受试者编号(SUBJID)(RAW),开始日期(AESTDAT)(RAW),复核开始日期(AESTDAT)(RAW)\n"
            "S01003,2026-03-01,2026-03-02\n"
        ).encode("utf-8-sig")

        row = parse_listing_file("AE.csv", content)[0].rows[0]

        self.assertEqual("2026-03-01", row["AESTDAT"])
        self.assertEqual("2026-03-02", row["AESTDAT__2"])

    def test_parse_csv_preserves_distinct_semantic_field_names_before_raw_qualifier(self):
        content = (
            "受试者编号(SUBJID)(RAW),不良事件名称(AETERM)(RAW),严重程度(AESEV)(RAW)\n"
            "S01003,头痛,1级\n"
        ).encode("utf-8-sig")

        row = parse_listing_file("AE.csv", content)[0].rows[0]

        self.assertEqual("头痛", row["AETERM"])
        self.assertEqual("1级", row["AESEV"])
        self.assertNotIn("RAW", row)
        self.assertNotIn("RAW__2", row)

    def test_parse_csv_does_not_treat_ascii_heavy_first_data_row_as_variable_header(self):
        content = (
            "状态,__STUDYOID,__STUDYEVENTOID,__STUDYEVENTREPEATKEY,DOMAIN,"
            "USUBJID,PAGE,FORM,LINE,AETERM\n"
            "Changed,STUDY-1,V1,1,AE,S01001,AE,AE,1,头痛\n"
        ).encode("utf-8")

        rows = parse_listing_file("Comparison.csv", content)[0].rows

        self.assertEqual(1, len(rows))
        self.assertEqual("Changed", rows[0]["状态"])
        self.assertEqual("头痛", rows[0]["AETERM"])

    def test_parse_xlsx_uses_edc_variable_row_when_present(self):
        workbook = openpyxl.Workbook()
        ae = workbook.active
        ae.title = "AE--不良事件"
        ae.append(["研究中心", "中心编号", "受试者", "是否发生不良事件？", "不良事件名称"])
        ae.append(["SITENM", "SITEID", "SUBJID", "AEYN", "AETERM"])
        ae.append(["上海市皮肤病医院", "01", "S01001", "是", "用药部位灼烧感"])

        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        sheets = parse_listing_file("rux_listing.xlsx", buffer.getvalue())

        self.assertEqual("S01001", sheets[0].rows[0]["SUBJID"])
        self.assertEqual("用药部位灼烧感", sheets[0].rows[0]["AETERM"])
        self.assertNotIn("受试者", sheets[0].rows[0])
        self.assertEqual([3], sheets[0].row_numbers)

    def test_parse_xlsx_ignores_invalid_autofilter_metadata_without_changing_values(self):
        workbook = openpyxl.Workbook()
        ae = workbook.active
        ae.title = "AE"
        ae.append(["SUBJID", "SITEID", "AETERM"])
        ae.append(["S01001", "01", "头痛"])
        buffer = io.BytesIO()
        workbook.save(buffer)
        workbook.close()

        source = zipfile.ZipFile(io.BytesIO(buffer.getvalue()), "r")
        patched = io.BytesIO()
        with source, zipfile.ZipFile(patched, "w") as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    payload = payload.replace(
                        b"</sheetData>",
                        b'</sheetData><autoFilter ref="1:1"/>',
                    )
                target.writestr(entry, payload)

        sheets = parse_listing_file("apache_poi_export.xlsx", patched.getvalue())

        self.assertEqual(["AE"], [sheet.sheet_name for sheet in sheets])
        self.assertEqual("S01001", sheets[0].rows[0]["SUBJID"])
        self.assertEqual("头痛", sheets[0].rows[0]["AETERM"])
        self.assertIn(
            "worksheet_autofilter_ignored: invalid range metadata ignored; cell values preserved",
            sheets[0].parser_warnings,
        )

    def test_unsupported_listing_type_raises_clear_error(self):
        with self.assertRaisesRegex(ValueError, "unsupported listing file type"):
            parse_listing_file("edc_listing.txt", b"SUBJID")

    @unittest.skipUnless(MGK10_LISTING.exists(), "MG-K10 raw listing not available on this machine")
    def test_parse_real_mgk10_listing_extracts_codes_from_chinese_headers(self):
        sheets = parse_listing_file(MGK10_LISTING.name, MGK10_LISTING.read_bytes())
        sheet_index = {sheet.sheet_name: sheet for sheet in sheets}

        self.assertIn("AE", sheet_index)
        self.assertIn("MH", sheet_index)
        ae_row = next(row for row in sheet_index["AE"].rows if row.get("AEYN") == "是" and row.get("AETERM"))
        mh_row = next(row for row in sheet_index["MH"].rows if row.get("MHYN") == "是" and row.get("MHTERM"))

        self.assertEqual("流鼻涕", ae_row["AETERM"])
        self.assertTrue(ae_row["SUBJID"])
        self.assertTrue(mh_row["MHTERM"])
        self.assertIn("SITEID", ae_row)

    @unittest.skipUnless(RUX_LISTING.exists(), "RUX raw listing not available on this machine")
    def test_parse_real_rux_listing_uses_second_variable_header(self):
        sheets = parse_listing_file(RUX_LISTING.name, RUX_LISTING.read_bytes())
        sheet_index = {sheet.sheet_name: sheet for sheet in sheets}

        self.assertIn("AE--不良事件", sheet_index)
        self.assertIn("CM--既往及合并用药治疗", sheet_index)
        ae_row = sheet_index["AE--不良事件"].rows[0]
        cm_row = sheet_index["CM--既往及合并用药治疗"].rows[0]

        self.assertEqual("S01001", ae_row["SUBJID"])
        self.assertEqual("用药部位灼烧感", ae_row["AETERM"])
        self.assertEqual("盐酸西替利嗪分散片", cm_row["CMTRT"])
        self.assertNotIn("受试者", ae_row)

    @unittest.skipUnless(RUX_SUBJECT_REPORT.exists(), "RUX subject report xls not available on this machine")
    def test_parse_real_rux_xls_subject_report_finds_header_after_metadata(self):
        sheets = parse_listing_file(RUX_SUBJECT_REPORT.name, RUX_SUBJECT_REPORT.read_bytes())

        self.assertEqual(["受试者报表"], [sheet.sheet_name for sheet in sheets])
        self.assertGreaterEqual(len(sheets[0].rows), 190)
        first = sheets[0].rows[0]
        self.assertEqual("S08004", first["SUBJID"])
        self.assertEqual("安慰剂组Y", first["ARM"])
        self.assertIn("RANDDTC", first)

    @unittest.skipUnless(
        RUX_RAW_DIMENSION_DEFECT.exists(),
        "RUX raw dimension-defective listing not available on this machine",
    )
    def test_parse_real_rux_raw_listing_recovers_incorrect_a1_dimensions(self):
        sheets = parse_listing_file(
            RUX_RAW_DIMENSION_DEFECT.name,
            RUX_RAW_DIMENSION_DEFECT.read_bytes(),
        )
        sheet_index = {sheet.sheet_name: sheet for sheet in sheets}

        self.assertEqual(241, len(sheet_index["SUBJ"].rows))
        self.assertGreater(len(sheet_index["LBCHEM"].rows), 20_000)
        self.assertIn(
            "worksheet_dimension_reset: declared A1:A1; parsed from worksheet XML bounds",
            sheet_index["SUBJ"].parser_warnings,
        )


if __name__ == "__main__":
    unittest.main()
