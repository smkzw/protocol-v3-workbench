from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document
from fastapi.testclient import TestClient

from packages.contracts.workbench_contracts import (
    ApprovalAction,
    ApprovalActionRequest,
    MedicalWritingGreenfieldCreateRequest,
    MedicalWritingGreenfieldDecision,
    MedicalWritingGreenfieldDecisionResolveRequest,
    MedicalWritingRevisionApplyRequest,
    MedicalWritingGreenfieldSectionSeed,
    MedicalWritingRevisionRequest,
    MedicalWritingWorkingCopySaveRequest,
    RevisionActionRequest,
    RevisionThread,
)
from packages.contracts.workbench_contracts.models import (
    InterventionRulesAuthority,
    InterventionRulesIpAdjustmentPolicy,
    InterventionRulesProductRole,
    MedicalWritingInterimAnalysisDesign,
    MedicalWritingInterventionIpRegimen,
    MedicalWritingInterventionRules,
    MedicalWritingPhase1Part,
    MedicalWritingPicosDefinition,
    MedicalWritingProtocolAssemblyPlanConfirmRequest,
    MedicalWritingProtocolAssemblyPlanRefreshRequest,
    MedicalWritingSectionFreezeRequest,
    MedicalWritingStructuredStudyDesign,
    MedicalWritingStudyDefinition,
    MedicalWritingStudyFraming,
)
from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing import MedicalWritingRevisionService
from services.api.app.ai_execution_policy import AiExecutionPolicyResolver
from services.api.app.ai_gateway import AiPromptEnvelope
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore
from services.api.app.medical_writing_document_exporter import (
    document_index_catalog,
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_greenfield import (
    CompositeMedicalWritingDocumentService,
    GreenfieldMedicalWritingConflictError,
    GreenfieldMedicalWritingDocumentService,
)
from services.api.app.medical_writing_manifest import RUX_PROTOCOL_DOCX
from services.api.app.medical_writing_company_corpus import (
    MedicalWritingCompanyCorpusService,
)
from services.api.app.medical_writing_protocol_template import (
    TEMPLATE_ID,
    TEMPLATE_VERSION,
    MedicalWritingProtocolTemplateService,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository
from services.api.app.medical_writing_plan_consumption import (
    MedicalWritingPlanConsumptionHelper,
)
from services.api.app.medical_writing_protocol_assembly_plan import (
    MedicalWritingProtocolAssemblyPlanService,
)
from services.api.app.sqlite_runtime_store import RuntimeStoreError, SqliteRuntimeStore
from services.api.app import main as app_main


def _study_definition_for_project(
    project_id: str,
    *,
    study_phase: str = "I期",
    indication: str = "慢性鼻窦炎伴鼻息肉",
    protocol_id: str = "MG-K10-CRSwNP-GREENFIELD",
    title: str = "CRSwNP研究方案绿地结构化表格测试",
) -> MedicalWritingStudyDefinition:
    """StudyDefinition with enough decided facts to confirm an AssemblyPlan."""
    from datetime import datetime, timezone

    now = datetime(2026, 7, 22, tzinfo=timezone.utc)
    design = MedicalWritingStructuredStudyDesign(
        randomization_mode="non_randomized",
        blinding_mode="open_label",
        comparator_type="none_or_dose_escalation",
        assignment_model="sequential",
        adaptive_design_enabled=False,
        sample_size_reestimation_planned=False,
        treatment_switch_planned=False,
        crossover_planned=False,
        open_label_extension_planned=False,
        src_planned=True,
        dmc_planned=False,
        phase1_parts=[
            MedicalWritingPhase1Part(
                part_code="SAD",
                population="健康志愿者",
                cohort_dose="单次递增",
            ),
            MedicalWritingPhase1Part(
                part_code="MAD",
                population="健康志愿者",
                cohort_dose="多次递增",
            ),
        ],
        interim_analysis=MedicalWritingInterimAnalysisDesign(planned=False),
    )
    rules = MedicalWritingInterventionRules(
        authority=InterventionRulesAuthority.STRUCTURED,
        ip_regimens=[
            MedicalWritingInterventionIpRegimen(
                regimen_id="ip-1",
                product_name="MG-K10",
                product_role=InterventionRulesProductRole.INVESTIGATIONAL_PRODUCT,
                dose_and_frequency="试验剂量",
                route="鼻喷",
                treatment_period="研究期",
            )
        ],
        ip_adjustment_policy=InterventionRulesIpAdjustmentPolicy.NO_PLANNED_ADJUSTMENT,
        no_planned_adjustment_statement="无计划剂量调整。",
        non_ip_treatment_rules=[],
    )
    framing = MedicalWritingStudyFraming(
        protocol_id=protocol_id,
        document_title=title,
        indication=indication,
        clinicaltrials_condition_term=indication,
        study_phase=study_phase,
        investigational_product="MG-K10",
        structured_design=design,
        product_profile={
            "technology_type": "other_biologic",
            "administration_routes": ["鼻喷"],
            "dosage_forms": ["喷雾剂"],
            "exposure_scope": "local",
        },
    )
    picos = MedicalWritingPicosDefinition(
        population_summary="CRSwNP患者",
        primary_endpoint="鼻息肉评分",
        intervention_rules=rules,
    )
    return MedicalWritingStudyDefinition(
        definition_id=f"def-{project_id}",
        project_id=project_id,
        revision=1,
        origin="guided_greenfield",
        framing=framing,
        picos=picos,
        state_sha256="a" * 64,
        created_at=now,
        updated_at=now,
        updated_by="greenfield_test",
    )


def _confirm_assembly_plan_for_project(
    tmpdir: Path,
    definition: MedicalWritingStudyDefinition,
):
    """Create and confirm a ProtocolAssemblyPlan for *definition* via the service."""
    store = {definition.project_id: definition}
    plan_service = MedicalWritingProtocolAssemblyPlanService(
        tmpdir / f"plan_{definition.project_id}.sqlite3",
        store.__getitem__,
    )
    refreshed = plan_service.refresh(
        definition.project_id,
        MedicalWritingProtocolAssemblyPlanRefreshRequest(
            expected_plan_revision=0,
            expected_source_definition_id=definition.definition_id,
            expected_source_definition_revision=definition.revision,
            expected_source_definition_sha256=definition.state_sha256,
            actor="greenfield_test",
            idempotency_key=f"refresh-{definition.project_id}",
        ),
    )
    confirmed = plan_service.confirm(
        definition.project_id,
        MedicalWritingProtocolAssemblyPlanConfirmRequest(
            expected_plan_revision=refreshed.plan.revision,
            expected_plan_sha256=refreshed.plan.state_sha256,
            actor="greenfield_test",
            idempotency_key=f"confirm-{definition.project_id}",
        ),
    )
    helper = MedicalWritingPlanConsumptionHelper(plan_service)
    return plan_service, helper, confirmed.plan


def _greenfield_request(
    *,
    idempotency_key: str = "create-syn-ra-v01",
) -> MedicalWritingGreenfieldCreateRequest:
    return MedicalWritingGreenfieldCreateRequest(
        protocol_id="SYN-RA-201",
        version="V0.1",
        document_title="CMS-RA-101治疗类风湿关节炎的探索性研究方案",
        indication="类风湿关节炎",
        study_phase="II期",
        investigational_product="CMS-RA-101",
        protocol_date="2026年07月20日",
        sponsor="深圳市康哲生物科技有限公司",
        sections=[
            MedicalWritingGreenfieldSectionSeed(
                section_key="study_design",
                heading="研究设计",
                ich_m11_anchor="C.3 Trial Design",
                initial_text="本研究拟采用随机、双盲、安慰剂对照设计。",
                source_fact_ids=[
                    "decision.design.randomized",
                    "decision.design.blinded",
                ],
            )
        ],
        decisions=[
            MedicalWritingGreenfieldDecision(
                decision_id="dose_selection",
                label="试验药物剂量",
                status="unresolved",
                value="",
                rationale="仍需结合非临床与早期临床资料确认。",
                source_refs=["synthetic_sponsor_brief:dose_selection"],
                approval_blocking=True,
            )
        ],
        actor="medical_manager_test",
        idempotency_key=idempotency_key,
    )


class StudyDefinitionGreenfieldBindingTests(unittest.TestCase):
    @staticmethod
    def _definition():
        return SimpleNamespace(
            definition_id="mwstudydef_ra_001",
            revision=3,
            state_sha256="b" * 64,
            synopsis_text="本研究拟评价试验药物治疗类风湿关节炎的有效性和安全性。",
            framing=SimpleNamespace(
                protocol_id="SYN-RA-201",
                version="V0.1",
                document_title="CMS-RA-101治疗类风湿关节炎的探索性研究方案",
                indication="类风湿关节炎",
                study_phase="II期",
                intrinsic_objectives=["剂量探索"],
                investigational_product="CMS-RA-101",
                development_regions=["中国"],
                target_mechanism="靶向炎症通路",
                design_pattern="随机、双盲、安慰剂对照设计",
                population_intent="中重度活动性类风湿关节炎成人患者",
                product_profile=SimpleNamespace(
                    technology_type="small_molecule",
                    administration_routes=["口服"],
                    dosage_forms=["片剂"],
                    immunogenicity_relevance="not_expected",
                    safety_considerations=[],
                    pk_pd_considerations=[],
                ),
                minimum_product_fact_packet=SimpleNamespace(),
            ),
            picos=SimpleNamespace(
                population_summary="中重度活动性类风湿关节炎成人患者",
                inclusion_modules=[],
                exclusion_modules=[],
                washout_rules=[],
                intervention_summary="CMS-RA-101治疗",
                intervention_dose_regimen="每日一次口服给药",
                allowed_concomitant_rules=[],
                required_background_rules=[],
                prohibited_concomitant_rules=[],
                assessment_timing_restrictions=[],
                comparator_summary="安慰剂",
                primary_objectives=[],
                secondary_objectives=[],
                exploratory_objectives=[],
                primary_endpoint="第12周ACR20应答率",
                key_secondary_endpoints=[],
                other_secondary_endpoints=[],
                exploratory_endpoints=[],
                safety_endpoints=[],
                aesi_definitions=[],
                assessment_instruments=[],
                study_epochs=[],
                visit_strategy="",
                estimand_strategy="",
                sample_size_strategy="计划入组约120例参与者。",
                statistical_strategy="采用预先规定的统计分析集进行分析。",
            ),
            field_states={
                "framing.protocol_id": SimpleNamespace(status="confirmed"),
                "framing.version": SimpleNamespace(status="confirmed"),
                "framing.document_title": SimpleNamespace(status="confirmed"),
                "framing.indication": SimpleNamespace(status="confirmed"),
                "framing.study_phase": SimpleNamespace(status="confirmed"),
                "framing.intrinsic_objectives": SimpleNamespace(status="confirmed"),
                "framing.investigational_product": SimpleNamespace(status="confirmed"),
                "framing.development_regions": SimpleNamespace(status="confirmed"),
                "framing.target_mechanism": SimpleNamespace(status="confirmed"),
                "framing.design_pattern": SimpleNamespace(status="confirmed"),
                "framing.population_intent": SimpleNamespace(status="confirmed"),
            },
            module_resolutions={},
        )

    def _request(self, *, protocol_id="SYN-RA-201", initial_text=None):
        definition = self._definition()
        prefix = f"study_definition:{definition.definition_id}:r{definition.revision}:"
        design_text = "随机、双盲、安慰剂对照设计；中重度活动性类风湿关节炎成人患者"
        return MedicalWritingGreenfieldCreateRequest(
            protocol_id=protocol_id,
            version="V0.1",
            document_title="CMS-RA-101治疗类风湿关节炎的探索性研究方案",
            indication="类风湿关节炎",
            study_phase="II期",
            source_study_definition_id=definition.definition_id,
            source_study_definition_revision=definition.revision,
            source_study_definition_sha256="a" * 64,
            sections=[
                MedicalWritingGreenfieldSectionSeed(
                    section_key="study_design",
                    heading="研究设计",
                    initial_text=design_text if initial_text is None else initial_text,
                    source_fact_ids=[
                        prefix + "framing.design_pattern",
                        prefix + "framing.population_intent",
                    ],
                )
            ],
            actor="medical_manager_test",
            idempotency_key="create-from-definition-r3",
        )

    def test_bound_request_accepts_exact_server_projection(self):
        app_main._validate_greenfield_request_against_study_definition(
            self._request(), self._definition()
        )

    def test_bound_request_rejects_client_identity_and_fact_text_mismatch(self):
        with self.assertRaisesRegex(Exception, "identity does not match"):
            app_main._validate_greenfield_request_against_study_definition(
                self._request(protocol_id="WRONG"), self._definition()
            )
        mismatched_product = self._request().model_copy(
            update={"investigational_product": "错误试验药物"}
        )
        with self.assertRaisesRegex(
            Exception,
            "investigational_product",
        ):
            app_main._validate_greenfield_request_against_study_definition(
                mismatched_product,
                self._definition(),
            )
        with self.assertRaisesRegex(Exception, "does not match.*projection"):
            app_main._validate_greenfield_request_against_study_definition(
                self._request(initial_text="与已确认研究事实无关的任意正文。"),
                self._definition(),
            )

    def test_template_bound_request_accepts_all_canonical_dynamic_fact_projections(self):
        definition = self._definition()
        definition.field_states.update(
            {
                "picos.intervention_summary": SimpleNamespace(status="confirmed"),
                "picos.comparator_summary": SimpleNamespace(status="confirmed"),
                "picos.primary_endpoint": SimpleNamespace(status="confirmed"),
            }
        )
        template_service = MedicalWritingProtocolTemplateService()
        resolved_sections = template_service.section_seeds(definition)
        request = self._request().model_copy(
            update={
                "template_id": TEMPLATE_ID,
                "template_version": TEMPLATE_VERSION,
                "sections": [],
            }
        )

        app_main._validate_greenfield_request_against_study_definition(
            request,
            definition,
            resolved_sections,
        )

        sourced_section = next(
            section for section in resolved_sections if section.source_fact_ids
        )
        tampered_section = sourced_section.model_copy(
            update={"initial_text": sourced_section.initial_text + "客户端篡改"}
        )
        tampered_sections = [
            tampered_section
            if section.section_key == sourced_section.section_key
            else section
            for section in resolved_sections
        ]
        with self.assertRaisesRegex(Exception, "does not match.*projection"):
            app_main._validate_greenfield_request_against_study_definition(
                request,
                definition,
                tampered_sections,
            )

    def test_greenfield_request_rejects_whitespace_only_identity_fields(self):
        payload = self._request().model_dump(mode="json")
        payload["protocol_id"] = "   "
        with self.assertRaisesRegex(ValueError, "must not be blank"):
            MedicalWritingGreenfieldCreateRequest.model_validate(payload)


class _GreenfieldWritingProvider:
    provider_name = "buddy"
    model_name = "deepseek-v4-pro"

    def __init__(self):
        self.envelopes: list[AiPromptEnvelope] = []

    def run(self, envelope: AiPromptEnvelope):
        self.envelopes.append(envelope)
        source = envelope.payload["allowed_sources"][0]
        selected_text = source["text_preview"]
        evidence_source = source
        proposal_seed = selected_text
        if not selected_text.strip():
            evidence_source = next(
                item
                for item in envelope.payload["allowed_sources"][1:]
                if str(item.get("text_preview") or "").strip()
            )
            proposal_seed = "本章节正文候选"
        span_id = f"span_greenfield_{len(self.envelopes):03d}"
        return {
            "task_id": envelope.task_id,
            "task_type": envelope.task_type.value,
            "provider": self.provider_name,
            "model": self.model_name,
            "prompt_version": envelope.prompt_version,
            "input_source_ids": [
                item["source_id"] for item in envelope.payload["allowed_sources"]
            ],
            "forbidden_source_ids": envelope.payload["forbidden_source_ids"],
            "findings": [
                {
                    "finding_id": f"finding_greenfield_{len(self.envelopes):03d}",
                    "status": "supported",
                    "title": "绿地方案工作副本修订候选",
                    "source_id": evidence_source["source_id"],
                    "evidence_span_ids": [span_id],
                }
            ],
            "evidence_spans": [
                {
                    "span_id": span_id,
                    "source_id": evidence_source["source_id"],
                    "locator": evidence_source["locator"],
                    "quote": evidence_source["text_preview"],
                }
            ],
            "uncertainties": [
                {
                    "level": "medical_review",
                    "description": "候选文本仍需医学确认，不得补写未决项目事实。",
                }
            ],
            "needs_medical_confirmation": True,
            "schema_version": "ai_task_output_v0_1",
            "revision": {
                "proposal_text": f"{proposal_seed}（AI修订候选）",
                "diff_patch": f"- {selected_text}\n+ {proposal_seed}（AI修订候选）",
                "rationale": "仅优化当前选中文本，不新增项目事实。",
                "evidence_span_ids": [span_id],
                "alternatives": [
                    {
                        "proposal_text": f"{proposal_seed}（AI备选修订一）",
                        "diff_patch": f"- {selected_text}\n+ {proposal_seed}（AI备选修订一）",
                        "rationale": "仅提供当前选中文本的备选表达。",
                        "evidence_span_ids": [span_id],
                    },
                    {
                        "proposal_text": f"{proposal_seed}（AI备选修订二）",
                        "diff_patch": f"- {selected_text}\n+ {proposal_seed}（AI备选修订二）",
                        "rationale": "仅提供当前选中文本的另一备选表达。",
                        "evidence_span_ids": [span_id],
                    },
                ],
            },
        }


class GreenfieldMedicalWritingDocumentServiceTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.db_path = Path(self.tmpdir.name) / "greenfield.sqlite3"
        self.service = GreenfieldMedicalWritingDocumentService(self.db_path)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_create_is_persistent_traceable_and_idempotent(self):
        created = self.service.create("proj_synthetic_ra", _greenfield_request())

        self.assertEqual("proj_synthetic_ra", created.document.project_id)
        self.assertEqual("SYN-RA-201", created.document.protocol_id)
        self.assertEqual("greenfield_candidate", created.document.status)
        self.assertEqual(1, created.baseline_revision)
        self.assertEqual(64, len(created.baseline_sha256))
        self.assertTrue(created.created)
        self.assertNotIn("/Users/", created.model_dump_json())
        serialized_gates = str(created.document.quality_gates)
        self.assertNotIn("待医学批准", serialized_gates)
        self.assertNotIn('"label": "医学批准"', serialized_gates)
        self.assertIn("章节确认与版本冻结", serialized_gates)
        self.assertIn("当前医学经理逐节确认", serialized_gates)

        replay = self.service.create("proj_synthetic_ra", _greenfield_request())
        self.assertFalse(replay.created)
        self.assertEqual(created.document.document_id, replay.document.document_id)
        self.assertEqual(created.baseline_sha256, replay.baseline_sha256)

        restarted = GreenfieldMedicalWritingDocumentService(self.db_path)
        session = restarted.document_session("proj_synthetic_ra")
        section = restarted.section("proj_synthetic_ra", session.sections[0].section_id)
        self.assertFalse(session.sections[0].content_blocks)
        self.assertEqual("研究设计", section.heading)
        self.assertEqual(2, len(section.content_blocks))
        self.assertEqual(
            "greenfield_project_decision", section.content_blocks[1]["source_kind"]
        )
        self.assertEqual(
            "本研究拟采用随机、双盲、安慰剂对照设计。",
            section.content_blocks[1]["text"],
        )

    def test_greenfield_front_matter_and_synopsis_are_real_tables_with_distinct_roles(
        self,
    ):
        request = _greenfield_request().model_copy(
            update={
                "sections": [
                    MedicalWritingGreenfieldSectionSeed(
                        section_key="front_matter",
                        heading="方案首页与版本信息",
                        node_kind="front_matter",
                        interaction_types=["front_matter_editor"],
                        title_locked=True,
                    ),
                    MedicalWritingGreenfieldSectionSeed(
                        section_key="protocol_synopsis",
                        heading="方案摘要",
                        ich_m11_anchor="1.1",
                        section_number="1.1",
                        node_kind="protocol_synopsis",
                        interaction_types=["synopsis_editor", "structured_table"],
                        initial_text=(
                            "研究题目：CMS-RA-101治疗类风湿关节炎的探索性研究方案\n"
                            "研究分期：II期\n"
                            "主要终点：第12周ACR20应答率"
                        ),
                        source_fact_ids=[
                            "study_definition:mwstudydef_ra_001:r3:framing.document_title",
                            "study_definition:mwstudydef_ra_001:r3:framing.study_phase",
                            "study_definition:mwstudydef_ra_001:r3:picos.primary_endpoint",
                        ],
                    ),
                ],
                "idempotency_key": "create-greenfield-document-objects-v1",
            }
        )

        created = self.service.create("proj_synthetic_ra", request)
        sections = {
            item.node_kind: self.service.section("proj_synthetic_ra", item.section_id)
            for item in created.document.sections
        }
        front_table = sections["front_matter"].content_blocks[1]
        synopsis_table = sections["protocol_synopsis"].content_blocks[1]

        self.assertEqual("table", front_table["block_type"])
        self.assertEqual("layout", front_table["structured_table"]["role"])
        self.assertEqual(0, front_table["header_row_count"])
        self.assertNotIn(
            "项目",
            [cell["text"] for row in front_table["rows"] for cell in row],
        )
        self.assertIn(
            "SYN-RA-201",
            [cell["text"] for row in front_table["rows"] for cell in row],
        )
        front_values = {
            row["label"]: row["cells"][1]["text"]
            for row in front_table["structured_table"]["rows"]
        }
        self.assertEqual("CMS-RA-101", front_values["研究药物"])
        self.assertEqual("2026年07月20日", front_values["版本日期"])
        self.assertEqual(
            "深圳市康哲生物科技有限公司",
            front_values["申办者"],
        )
        self.assertEqual("table", synopsis_table["block_type"])
        self.assertEqual(
            "protocol_synopsis",
            synopsis_table["structured_table"]["role"],
        )
        self.assertIn(
            "第12周ACR20应答率",
            [cell["text"] for row in synopsis_table["rows"] for cell in row],
        )
        self.assertEqual(
            [],
            document_index_catalog(
                self.service.document_for_revision("proj_synthetic_ra")
            )["tables"],
        )

    def test_duplicate_project_does_not_silently_replace_a_baseline(self):
        self.service.create("proj_synthetic_ra", _greenfield_request())
        changed = _greenfield_request(idempotency_key="create-syn-ra-v02")
        changed.version = "V0.2"

        with self.assertRaises(GreenfieldMedicalWritingConflictError):
            self.service.create("proj_synthetic_ra", changed)

        self.assertEqual(
            "V0.1",
            self.service.document_session("proj_synthetic_ra").version,
        )

    def test_unresolved_decision_can_be_resolved_with_optimistic_concurrency(self):
        created = self.service.create("proj_synthetic_ra", _greenfield_request())
        self.assertEqual(1, len(self.service.approval_blockers("proj_synthetic_ra")))

        resolved = self.service.resolve_decision(
            "proj_synthetic_ra",
            "dose_selection",
            MedicalWritingGreenfieldDecisionResolveRequest(
                expected_baseline_revision=created.baseline_revision,
                value="剂量A（仅用于系统测试）",
                rationale="已由医学负责人完成测试场景确认。",
                source_refs=["synthetic_sponsor_decision:dose_selection:v2"],
                actor="medical_director_test",
                idempotency_key="resolve-dose-selection-v2",
            ),
        )
        self.assertEqual(2, resolved.baseline_revision)
        self.assertEqual("resolved", resolved.decision.status)
        self.assertEqual([], self.service.approval_blockers("proj_synthetic_ra"))

        with self.assertRaises(GreenfieldMedicalWritingConflictError):
            self.service.resolve_decision(
                "proj_synthetic_ra",
                "dose_selection",
                MedicalWritingGreenfieldDecisionResolveRequest(
                    expected_baseline_revision=1,
                    value="陈旧值",
                    rationale="使用陈旧版本提交。",
                    source_refs=["stale:test"],
                    actor="medical_manager_test",
                    idempotency_key="stale-resolution",
                ),
            )

    @unittest.skipUnless(
        RUX_PROTOCOL_DOCX.exists(), "real RUX protocol fixture is unavailable"
    )
    def test_composite_service_preserves_real_docx_precedence(self):
        greenfield = GreenfieldMedicalWritingDocumentService(self.db_path)
        greenfield.create("proj_rux_03_002", _greenfield_request())
        composite = CompositeMedicalWritingDocumentService(
            MedicalWritingDocumentService(),
            greenfield,
        )

        session = composite.document_session("proj_rux_03_002")

        self.assertEqual("RUX-03-002", session.protocol_id)
        self.assertTrue(session.document_id.startswith("mwdoc_proj_rux_03_002_"))
        self.assertEqual(
            "original_protocol_docx", composite.source_mode("proj_rux_03_002")
        )
        with self.assertRaises(GreenfieldMedicalWritingConflictError):
            composite.create_greenfield("proj_rux_03_002", _greenfield_request())


class GreenfieldMedicalWritingRuntimeFlowTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        root = Path(self.tmpdir.name)
        self.greenfield = GreenfieldMedicalWritingDocumentService(
            root / "greenfield.sqlite3"
        )
        self.greenfield.create("proj_synthetic_ra", _greenfield_request())
        self.documents = CompositeMedicalWritingDocumentService(
            MedicalWritingDocumentService(),
            self.greenfield,
        )
        self.store = SqliteRuntimeStore(root / "runtime.sqlite3")
        self.repository = MedicalWritingRuntimeRepository(self.documents, self.store)

    def tearDown(self):
        self.tmpdir.cleanup()

    def test_heading_only_greenfield_word_export_is_blocked(self):
        with self.assertRaisesRegex(RuntimeStoreError, "缺少实质正文"):
            self.repository.assemble_document_for_export(
                "proj_synthetic_ra", "draft_preview"
            )

    def test_greenfield_reuses_working_copy_approval_and_docx_export(self):
        """Author save + freeze semantics replace retired writing ApprovalGate.

        - Saving working-copy content is the author's decision.
        - Unresolved greenfield project decisions remain visible blockers.
        - After resolving required decisions, freeze the current revision and
          verify freeze history / final readiness / final export.
        """
        project_id = "proj_synthetic_ra"
        session = self.documents.document_session(project_id)
        section_id = session.sections[0].section_id
        initial = self.repository.working_copy(project_id, section_id)
        self.assertEqual(0, initial.revision)

        author_text = (
            "本研究采用随机、双盲、安慰剂对照设计（作者已确认当前章节措辞）。"
            "研究对象、给药安排、主要评价路径及安全性观察均须以已确认的研究事实为依据，"
            "本段正文用于验证完整导出闸门不会把短句或标题骨架误认为可递交内容。"
        )
        candidate_blocks = [dict(block) for block in initial.content_blocks]
        candidate_blocks[1] = dict(candidate_blocks[1])
        candidate_blocks[1]["text"] = author_text
        saved = self.repository.save_working_copy(
            project_id,
            section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=session.document_id,
                expected_revision=0,
                content_blocks=candidate_blocks,
                actor="medical_manager_test",
                idempotency_key="save-greenfield-r1",
            ),
        )
        self.assertEqual(1, saved.revision)
        self.assertEqual("editable", saved.freeze_status)

        # Draft export uses the author-saved working copy (no plan helper here).
        draft = self.repository.assemble_document_for_export(
            project_id, "draft_preview"
        )
        exported = export_medical_writing_document_docx(draft, mode="draft_preview")
        rendered = Document(io.BytesIO(exported.content))
        rendered_text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
        self.assertIn(author_text, rendered_text)
        self.assertNotIn("待医学批准", rendered_text)

        # Unresolved project/design decisions remain document-level blockers.
        unresolved_blockers = self.greenfield.approval_blockers(project_id)
        self.assertTrue(
            any(
                item.blocker_type == "greenfield_unresolved_decision"
                for item in unresolved_blockers
            )
        )
        # Retired writing ApprovalGate must stay retired.
        with self.assertRaisesRegex(RuntimeStoreError, "retired|freeze-current-version"):
            self.repository.ensure_working_copy_approval_gate(
                saved,
                "medical_manager_test",
            )

        # Resolve the required greenfield decision, then freeze the author version.
        self.greenfield.resolve_decision(
            project_id,
            "dose_selection",
            MedicalWritingGreenfieldDecisionResolveRequest(
                expected_baseline_revision=1,
                value="剂量A（仅用于系统测试）",
                rationale="测试场景中完成项目决策。",
                source_refs=["synthetic_sponsor_decision:dose_selection:v2"],
                actor="medical_manager_test",
                idempotency_key="resolve-before-author-freeze",
            ),
        )
        self.assertEqual([], self.greenfield.approval_blockers(project_id))

        frozen = self.repository.freeze_current_version(
            project_id,
            section_id,
            MedicalWritingSectionFreezeRequest(
                document_id=session.document_id,
                expected_working_copy_revision=saved.revision,
                reason="医学作者已完成本章节内容核对，确认冻结当前版本。",
                actor="medical_manager_test",
                idempotency_key="freeze-greenfield-r1",
            ),
        )
        self.assertEqual("frozen", frozen.working_copy.freeze_status)
        self.assertFalse(frozen.replayed)
        history = self.repository.section_freeze_history(project_id, section_id)
        self.assertEqual(1, len(history))
        self.assertTrue(history[0].is_current)
        readiness = self.repository.final_freeze_readiness(project_id)
        self.assertTrue(readiness.ready)
        final = self.repository.assemble_document_for_export(
            project_id, "approved_final"
        )
        self.assertEqual("author_frozen_final", final.status)

    def test_greenfield_ai_revision_uses_versioned_internal_source_and_stales_on_decision_change(
        self,
    ):
        project_id = "proj_synthetic_ra"
        session = self.documents.document_session(project_id)
        section_id = session.sections[0].section_id
        section = self.documents.section(project_id, section_id)
        source_block = section.content_blocks[1]
        provider = _GreenfieldWritingProvider()
        runner = AiTaskRunner(
            self.repository,
            AiTaskStore(Path(self.tmpdir.name) / "ai_runs.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name="buddy",
                model_name="deepseek-v4-pro",
                test_only_provider_injection=True,
            ),
        )
        revision_service = MedicalWritingRevisionService(
            self.repository,
            runner,
            company_corpus_service=MedicalWritingCompanyCorpusService(),
            authoring_journey_service=SimpleNamespace(
                has_project=lambda candidate_project_id: (
                    candidate_project_id == project_id
                ),
                get=lambda candidate_project_id: SimpleNamespace(
                    study_definition=StudyDefinitionGreenfieldBindingTests._definition(),
                    framing=StudyDefinitionGreenfieldBindingTests._definition().framing,
                ),
            ),
        )

        submitted = revision_service.submit_revision(
            project_id,
            MedicalWritingRevisionRequest(
                document_id=session.document_id,
                section_id=section_id,
                anchor_path=source_block["source_locator"],
                selected_text=source_block["text"],
                user_instruction="请只优化当前句式，不新增剂量、终点或样本量。",
                requested_by="medical_manager_test",
            ),
        )
        self.assertTrue(submitted.thread.source_entry_id.startswith("greenfield:"))
        self.assertEqual(
            RevisionThread.selected_text_hash(submitted.thread.selected_text),
            submitted.thread.selected_hash,
        )
        envelope = provider.envelopes[0]
        self.assertEqual(
            "greenfield_working_copy_selection",
            envelope.payload["allowed_sources"][0]["source_type"],
        )
        project_fact_sources = [
            item
            for item in envelope.payload["allowed_sources"]
            if item["source_type"] == "current_project_study_definition"
        ]
        self.assertEqual(1, len(project_fact_sources))
        self.assertIn("试验药物：CMS-RA-101", project_fact_sources[0]["text_preview"])
        self.assertIn("适应症：类风湿关节炎", project_fact_sources[0]["text_preview"])
        self.assertIn("研究分期：II期", project_fact_sources[0]["text_preview"])
        self.assertIn("开发区域：中国", project_fact_sources[0]["text_preview"])
        self.assertLess(
            envelope.payload["allowed_sources"].index(project_fact_sources[0]),
            min(
                envelope.payload["allowed_sources"].index(item)
                for item in envelope.payload["allowed_sources"]
                if item["source_type"] == "company_protocol_reference_corpus"
            ),
        )
        corpus_sources = [
            item
            for item in envelope.payload["allowed_sources"]
            if item["source_type"] == "company_protocol_reference_corpus"
        ]
        self.assertGreaterEqual(len(corpus_sources), 3)
        self.assertLessEqual(len(corpus_sources), 5)
        self.assertTrue(
            all(
                item["source_id"].startswith("company_corpus:")
                for item in corpus_sources
            )
        )
        task_context = envelope.payload["task_context"]
        self.assertEqual("medical_writing_revision", task_context["revision_intent"])
        self.assertEqual("改写", task_context["intent_label"])
        self.assertEqual(4, task_context["candidate_count"])
        self.assertTrue(task_context["directional_goal"])
        self.assertTrue(task_context["preservation_rules"])
        self.assertTrue(
            any(
                "不得据此新增、替换或确认当前项目" in rule
                for rule in task_context["preservation_rules"]
            )
        )
        self.assertTrue(
            any(
                "current_project_study_definition" in rule
                and "优先级高于竞品证据和参考语料" in rule
                for rule in task_context["preservation_rules"]
            )
        )
        self.assertEqual(4, len(task_context["candidate_blueprints"]))
        self.assertIn(
            self.greenfield.baseline_state(project_id)["baseline_sha256"],
            submitted.thread.source_entry_id,
        )

        self.greenfield.resolve_decision(
            project_id,
            "dose_selection",
            MedicalWritingGreenfieldDecisionResolveRequest(
                expected_baseline_revision=1,
                value="剂量A（仅用于系统测试）",
                rationale="测试场景完成项目决策。",
                source_refs=["synthetic_sponsor_decision:dose_selection:v2"],
                actor="medical_director_test",
                idempotency_key="resolve-after-ai-thread",
            ),
        )
        with self.assertRaisesRegex(ValueError, "stale"):
            revision_service.apply_action(
                project_id,
                submitted.thread.thread_id,
                RevisionActionRequest(
                    action="request_rewrite",
                    suggestion_id=submitted.suggestion.suggestion_id,
                    actor="medical_manager_test",
                    comment="项目决策已更新，请按最新基线重写。",
                    rewrite_instruction="保持事实不变并压缩句式。",
                ),
            )
        self.assertEqual(1, len(provider.envelopes))

    def test_blank_greenfield_body_can_be_drafted_and_applied_without_touching_heading(
        self,
    ):
        project_id = "proj_synthetic_ra_blank"
        request = _greenfield_request(idempotency_key="create-syn-ra-blank-v01")
        request.protocol_id = "SYN-RA-BLANK-201"
        request.sections[0].initial_text = ""
        request.sections[0].source_fact_ids = []
        self.greenfield.create(project_id, request)
        session = self.documents.document_session(project_id)
        section = self.documents.section(project_id, session.sections[0].section_id)
        heading_block, body_block = section.content_blocks
        self.assertEqual("", body_block["text"])
        self.assertEqual("greenfield_scaffold", body_block["source_kind"])

        provider = _GreenfieldWritingProvider()
        runner = AiTaskRunner(
            self.repository,
            AiTaskStore(Path(self.tmpdir.name) / "blank_ai_runs.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name="buddy",
                model_name="deepseek-v4-pro",
                test_only_provider_injection=True,
            ),
        )
        revision_service = MedicalWritingRevisionService(
            self.repository,
            runner,
            company_corpus_service=MedicalWritingCompanyCorpusService(),
        )
        initial_copy = self.repository.working_copy(project_id, section.section_id)
        submitted = revision_service.submit_revision(
            project_id,
            MedicalWritingRevisionRequest(
                document_id=session.document_id,
                section_id=section.section_id,
                anchor_path=body_block["source_locator"],
                selected_text="",
                user_instruction="请基于已提供来源起草当前章节正文。",
                requested_by="medical_manager_test",
            ),
        )

        self.assertEqual("", submitted.thread.selected_text)
        self.assertEqual(body_block["source_locator"], submitted.thread.anchor_path)
        self.assertEqual(
            0, self.repository.working_copy(project_id, section.section_id).revision
        )
        task_context = provider.envelopes[0].payload["task_context"]
        self.assertIn("当前空白", task_context["directional_goal"])
        self.assertTrue(
            any("空白正文块" in rule for rule in task_context["preservation_rules"])
        )

        accepted = revision_service.apply_action(
            project_id,
            submitted.thread.thread_id,
            RevisionActionRequest(
                action="accept",
                suggestion_id=submitted.suggestion.suggestion_id,
                actor="medical_manager_test",
            ),
        ).thread
        applied = self.repository.apply_approved_revision_to_working_copy(
            project_id,
            section.section_id,
            accepted.thread_id,
            MedicalWritingRevisionApplyRequest(
                expected_working_copy_revision=initial_copy.revision,
                actor="medical_manager_test",
                idempotency_key="apply-blank-greenfield-draft",
            ),
        )
        changed_heading, changed_body = applied.working_copy.content_blocks
        self.assertEqual(heading_block["text"], changed_heading["text"])
        self.assertEqual(submitted.suggestion.proposal_text, changed_body["text"])
        self.assertEqual(body_block["source_locator"], changed_body["source_locator"])

    def test_blank_greenfield_draft_cannot_overwrite_later_manual_text(self):
        project_id = "proj_synthetic_ra_blank_drift"
        request = _greenfield_request(idempotency_key="create-syn-ra-blank-drift-v01")
        request.protocol_id = "SYN-RA-BLANK-DRIFT-201"
        request.sections[0].initial_text = ""
        request.sections[0].source_fact_ids = []
        self.greenfield.create(project_id, request)
        session = self.documents.document_session(project_id)
        section = self.documents.section(project_id, session.sections[0].section_id)
        body_block = section.content_blocks[1]
        provider = _GreenfieldWritingProvider()
        revision_service = MedicalWritingRevisionService(
            self.repository,
            AiTaskRunner(
                self.repository,
                AiTaskStore(Path(self.tmpdir.name) / "blank_drift_ai_runs.jsonl"),
                provider_factory=lambda resolution: provider,
                policy_resolver=AiExecutionPolicyResolver(
                    deployment_profile="approved_private_documents",
                    provider_name="buddy",
                    model_name="deepseek-v4-pro",
                    test_only_provider_injection=True,
                ),
            ),
            company_corpus_service=MedicalWritingCompanyCorpusService(),
        )
        submitted = revision_service.submit_revision(
            project_id,
            MedicalWritingRevisionRequest(
                document_id=session.document_id,
                section_id=section.section_id,
                anchor_path=body_block["source_locator"],
                selected_text="",
                user_instruction="请起草当前章节正文。",
                requested_by="medical_manager_test",
            ),
        )
        accepted = revision_service.apply_action(
            project_id,
            submitted.thread.thread_id,
            RevisionActionRequest(
                action="accept",
                suggestion_id=submitted.suggestion.suggestion_id,
                actor="medical_manager_test",
            ),
        ).thread
        current = self.repository.working_copy(project_id, section.section_id)
        manual_blocks = [dict(item) for item in current.content_blocks]
        manual_blocks[1]["text"] = "医学经理已手工填写的正文。"
        saved = self.repository.save_working_copy(
            project_id,
            section.section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=session.document_id,
                expected_revision=0,
                content_blocks=manual_blocks,
                actor="medical_manager_test",
                idempotency_key="manual-text-before-ai-apply",
            ),
        )
        with self.assertRaisesRegex(RuntimeStoreError, "selected text"):
            self.repository.apply_approved_revision_to_working_copy(
                project_id,
                section.section_id,
                accepted.thread_id,
                MedicalWritingRevisionApplyRequest(
                    expected_working_copy_revision=saved.revision,
                    actor="medical_manager_test",
                    idempotency_key="reject-stale-blank-greenfield-draft",
                ),
            )

    def test_greenfield_docx_export_has_company_standard_page_footer(self):
        import zipfile as _zipfile

        project_id = "proj_synthetic_ra"
        session = self.documents.document_session(project_id)
        initial = self.repository.working_copy(project_id, session.sections[0].section_id)
        substantive_text = (
            "本研究采用随机、双盲、安慰剂对照设计，按照预先定义的研究人群、给药安排和评价指标"
            "开展受试者管理与安全性观察，所有关键事实均须经医学经理审核后进入正式方案完整正文。"
        )
        body_blocks = [dict(block) for block in initial.content_blocks]
        body_blocks[1]["text"] = substantive_text
        self.repository.save_working_copy(
            project_id,
            session.sections[0].section_id,
            MedicalWritingWorkingCopySaveRequest(
                document_id=session.document_id,
                expected_revision=initial.revision,
                content_blocks=body_blocks,
                actor="medical_manager_test",
                idempotency_key="save-footer-export-body",
            ),
        )
        draft = self.repository.assemble_document_for_export(
            project_id, "draft_preview"
        )
        exported = export_medical_writing_document_docx(draft, mode="draft_preview")
        with _zipfile.ZipFile(io.BytesIO(exported.content)) as archive:
            footer_parts = sorted(
                name for name in archive.namelist() if "footer" in name.lower()
            )
            self.assertTrue(
                footer_parts, "greenfield DOCX must have at least one footer part"
            )
            footer_xml = archive.read(footer_parts[0]).decode("utf-8")
        self.assertIn(" PAGE ", footer_xml)
        self.assertIn(" NUMPAGES ", footer_xml)
        document_xml = io.BytesIO(exported.content)
        with _zipfile.ZipFile(document_xml) as archive:
            doc_xml = archive.read("word/document.xml").decode("utf-8")
        self.assertIn("footerReference", doc_xml)


class GreenfieldMedicalWritingApiTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        root = Path(self.tmpdir.name)
        self.greenfield_db_path = root / "greenfield.sqlite3"
        self.runtime_db_path = root / "runtime.sqlite3"
        self.greenfield = GreenfieldMedicalWritingDocumentService(
            self.greenfield_db_path
        )
        self.documents = CompositeMedicalWritingDocumentService(
            MedicalWritingDocumentService(),
            self.greenfield,
        )
        self.runtime_store = SqliteRuntimeStore(self.runtime_db_path)
        self.repository = MedicalWritingRuntimeRepository(
            self.documents,
            self.runtime_store,
        )
        self.patches = [
            patch(
                "services.api.app.main.medical_writing_greenfield_document_service",
                self.greenfield,
            ),
            patch(
                "services.api.app.main.medical_writing_document_service",
                self.documents,
            ),
            patch(
                "services.api.app.main.medical_writing_runtime_repository",
                self.repository,
            ),
        ]
        for item in self.patches:
            item.start()
        self.client = TestClient(app_main.app)

    def tearDown(self):
        for item in reversed(self.patches):
            item.stop()
        self.tmpdir.cleanup()

    def test_known_project_can_explicitly_create_and_resume_greenfield_session(self):
        project_id = "proj_mgk10_crswnp"
        payload = _greenfield_request().model_dump(mode="json")
        payload["protocol_id"] = "MG-K10-CRSwNP-GREENFIELD"
        payload["document_title"] = "CRSwNP研究方案绿地写作测试"
        payload["indication"] = "慢性鼻窦炎伴鼻息肉"

        created = self.client.post(
            f"/api/projects/{project_id}/medical-writing/greenfield-document",
            json=payload,
        )
        self.assertEqual(200, created.status_code, created.text)
        self.assertTrue(created.json()["created"])

        session = self.client.get(
            f"/api/projects/{project_id}/medical-writing/document-session"
        )
        self.assertEqual(200, session.status_code, session.text)
        self.assertEqual("MG-K10-CRSwNP-GREENFIELD", session.json()["protocol_id"])
        self.assertEqual("greenfield_candidate", session.json()["status"])

        state = self.client.get(
            f"/api/projects/{project_id}/medical-writing/greenfield-document"
        )
        self.assertEqual(200, state.status_code, state.text)
        self.assertEqual(1, state.json()["approval_blocker_count"])
        self.assertEqual("greenfield_project_decision", state.json()["source_mode"])

        manifest = self.client.get(
            f"/api/projects/{project_id}/medical-writing/manifest"
        )
        self.assertEqual(200, manifest.status_code, manifest.text)
        self.assertEqual(1, manifest.json()["package_count"])
        self.assertEqual(
            "structured_project_decisions",
            manifest.json()["packages"][0]["documents"][0]["file_format"],
        )

        duplicate = self.client.post(
            f"/api/projects/{project_id}/medical-writing/greenfield-document",
            json={**payload, "idempotency_key": "different-create-request"},
        )
        self.assertEqual(409, duplicate.status_code, duplicate.text)

    def test_pre_document_state_is_a_normal_optional_response(self):
        project_id = "proj_mgk10_crswnp"

        session = self.client.get(
            f"/api/projects/{project_id}/medical-writing/document-session",
            params={"allow_missing": "true"},
        )
        self.assertEqual(200, session.status_code, session.text)
        self.assertEqual({"available": False}, session.json())

        manifest = self.client.get(
            f"/api/projects/{project_id}/medical-writing/manifest",
            params={"allow_missing": "true"},
        )
        self.assertEqual(200, manifest.status_code, manifest.text)
        self.assertEqual({"available": False}, manifest.json())

    def test_decision_resolution_endpoint_preserves_warning_until_explicit_resolution(
        self,
    ):
        project_id = "proj_mgk10_crswnp"
        payload = _greenfield_request().model_dump(mode="json")
        payload.update(
            {
                "protocol_id": "MG-K10-CRSwNP-GREENFIELD",
                "document_title": "CRSwNP研究方案绿地写作测试",
                "indication": "慢性鼻窦炎伴鼻息肉",
            }
        )
        self.assertEqual(
            200,
            self.client.post(
                f"/api/projects/{project_id}/medical-writing/greenfield-document",
                json=payload,
            ).status_code,
        )

        resolved = self.client.post(
            f"/api/projects/{project_id}/medical-writing/greenfield-document/decisions/dose_selection/resolve",
            json={
                "expected_baseline_revision": 1,
                "value": "测试剂量A",
                "rationale": "由医学负责人确认本测试项目决策。",
                "source_refs": ["project_decision:dose_selection:v2"],
                "actor": "medical_director_test",
                "idempotency_key": "resolve-crswnp-dose-v2",
            },
        )
        self.assertEqual(200, resolved.status_code, resolved.text)
        self.assertEqual(2, resolved.json()["baseline_revision"])

        state = self.client.get(
            f"/api/projects/{project_id}/medical-writing/greenfield-document"
        )
        self.assertEqual(0, state.json()["approval_blocker_count"])

    def test_greenfield_structured_table_round_trip_restart_and_export_boundaries(self):
        project_id = "proj_mgk10_crswnp"
        payload = _greenfield_request().model_dump(mode="json")
        payload.update(
            {
                "protocol_id": "MG-K10-CRSwNP-GREENFIELD",
                "document_title": "CRSwNP研究方案绿地结构化表格测试",
                "indication": "慢性鼻窦炎伴鼻息肉",
            }
        )
        created = self.client.post(
            f"/api/projects/{project_id}/medical-writing/greenfield-document",
            json=payload,
        )
        self.assertEqual(200, created.status_code, created.text)
        document = created.json()["document"]
        section_id = document["sections"][0]["section_id"]
        working_copy_route = (
            f"/api/projects/{project_id}/medical-writing/working-copies/{section_id}"
        )

        initial = self.client.get(working_copy_route)
        self.assertEqual(200, initial.status_code, initial.text)
        self.assertEqual(0, initial.json()["revision"])
        saved = self.client.post(
            working_copy_route,
            json={
                "document_id": document["document_id"],
                "expected_revision": 0,
                "content_blocks": initial.json()["content_blocks"],
                "actor": "medical_manager_test",
                "idempotency_key": "greenfield-table-first-save-001",
            },
        )
        self.assertEqual(200, saved.status_code, saved.text)
        self.assertEqual(1, saved.json()["revision"])

        instantiated = self.client.post(
            f"{working_copy_route}/table-templates/objectives_endpoints/instantiate",
            params={
                "actor": "medical_manager_test",
                "idempotency_key": "greenfield-objectives-table-r1",
            },
        )
        self.assertEqual(200, instantiated.status_code, instantiated.text)
        table_block = instantiated.json()["table_block"]
        self.assertEqual(2, instantiated.json()["working_copy"]["revision"])
        self.assertEqual(0, table_block["structured_table"]["version"])

        replayed = self.client.post(
            f"{working_copy_route}/table-templates/objectives_endpoints/instantiate",
            params={
                "actor": "medical_manager_test",
                "idempotency_key": "greenfield-objectives-table-r1",
            },
        )
        self.assertEqual(200, replayed.status_code, replayed.text)
        self.assertEqual(2, replayed.json()["working_copy"]["revision"])
        self.assertEqual(
            table_block["block_id"], replayed.json()["table_block"]["block_id"]
        )
        self.assertEqual(
            1,
            len(
                [
                    block
                    for block in replayed.json()["working_copy"]["content_blocks"]
                    if block.get("block_id") == table_block["block_id"]
                ]
            ),
        )
        conflicting_replay = self.client.post(
            f"{working_copy_route}/table-templates/objectives_endpoints/instantiate",
            params={
                "actor": "medical_manager_test",
                "idempotency_key": "greenfield-objectives-table-r1",
                "title": "同键不同表题",
            },
        )
        self.assertEqual(409, conflicting_replay.status_code, conflicting_replay.text)

        table_route = f"{working_copy_route}/tables/{table_block['block_id']}"
        table = self.client.get(table_route)
        self.assertEqual(200, table.status_code, table.text)
        table_payload = table.json()
        body_rows = table_payload["rows"][table_payload["header_row_count"] :]
        target_cell = next(cell for row in body_rows for cell in row["cells"][1:])
        updated_text = "主要终点：第24周鼻息肉评分较基线变化（待医学批准）"
        updated = self.client.post(
            f"{table_route}/batch-update",
            json={
                "expected_working_copy_revision": 2,
                "expected_table_version": 0,
                "operations": [
                    {
                        "op": "edit_cell",
                        "cell_id": target_cell["cell_id"],
                        "text": updated_text,
                    }
                ],
                "actor": "medical_manager_test",
                "idempotency_key": "greenfield-table-cell-edit-001",
            },
        )
        self.assertEqual(200, updated.status_code, updated.text)
        self.assertEqual(3, updated.json()["working_copy"]["revision"])
        self.assertEqual(1, updated.json()["table"]["version"])

        reread_copy = self.client.get(working_copy_route)
        reread_table = self.client.get(table_route)
        self.assertEqual(3, reread_copy.json()["revision"])
        self.assertEqual(1, reread_table.json()["version"])
        self.assertIn(
            updated_text,
            [
                cell["text"]
                for row in reread_table.json()["rows"]
                for cell in row["cells"]
            ],
        )

        # Managed DOCX export consumes a confirmed ProtocolAssemblyPlan via the
        # production exporter helper. Without a plan the export must fail closed.
        blocked = self.client.get(
            f"/api/projects/{project_id}/medical-writing/document.docx",
            params={"mode": "draft_preview"},
        )
        self.assertEqual(409, blocked.status_code, blocked.text)
        self.assertIn("plan", blocked.json()["detail"].lower())

        definition = _study_definition_for_project(
            project_id,
            indication="慢性鼻窦炎伴鼻息肉",
            protocol_id="MG-K10-CRSwNP-GREENFIELD",
            title="CRSwNP研究方案绿地结构化表格测试",
        )
        plan_service, plan_helper, plan = _confirm_assembly_plan_for_project(
            Path(self.tmpdir.name),
            definition,
        )
        self.assertEqual("author_confirmed", plan.confirmation_status)
        pin = plan_helper.require_confirmed_projection(
            project_id=project_id,
            projection_kind="docx_toc",
        )
        self.assertEqual(plan.plan_id, pin.plan.plan_id)
        self.assertEqual(plan.revision, pin.plan.revision)
        self.assertEqual(plan.state_sha256, pin.plan.state_sha256)

        with (
            patch.object(
                app_main,
                "medical_writing_protocol_assembly_plan_service",
                plan_service,
            ),
            patch.object(
                app_main,
                "medical_writing_plan_consumption_helper",
                plan_helper,
            ),
        ):
            draft = self.client.get(
                f"/api/projects/{project_id}/medical-writing/document.docx",
                params={"mode": "draft_preview"},
            )
            self.assertEqual(200, draft.status_code, draft.text)
            self.assertTrue(draft.content.startswith(b"PK"))
            self.assertEqual(
                "draft_preview", draft.headers["x-medical-writing-export-mode"]
            )
            self.assertEqual("1", draft.headers["x-medical-writing-table-count"])
            self.assertEqual(64, len(draft.headers["x-medical-writing-docx-sha256"]))
            exported = Document(io.BytesIO(draft.content))
            self.assertEqual(1, len(exported.tables))
            exported_table_text = "\n".join(
                cell.text
                for table_item in exported.tables
                for row in table_item.rows
                for cell in row.cells
            )
            self.assertIn(updated_text, exported_table_text)

            formal = self.client.get(
                f"/api/projects/{project_id}/medical-writing/document.docx",
                params={"mode": "approved_final"},
            )
            # Still fail-closed for final: sections not author-frozen.
            self.assertEqual(409, formal.status_code, formal.text)

        restarted_greenfield = GreenfieldMedicalWritingDocumentService(
            self.greenfield_db_path
        )
        restarted_documents = CompositeMedicalWritingDocumentService(
            MedicalWritingDocumentService(),
            restarted_greenfield,
        )
        restarted_repository = MedicalWritingRuntimeRepository(
            restarted_documents,
            SqliteRuntimeStore(self.runtime_db_path),
        )
        recovered = restarted_repository.working_copy(project_id, section_id)
        self.assertEqual(3, recovered.revision)
        recovered_table = next(
            block
            for block in recovered.content_blocks
            if block.get("block_id") == table_block["block_id"]
        )
        self.assertIn(
            updated_text,
            [cell["text"] for row in recovered_table["rows"] for cell in row],
        )


if __name__ == "__main__":
    unittest.main()
