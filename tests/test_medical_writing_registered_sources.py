from __future__ import annotations

import io
import json
import tempfile
import unittest
from pathlib import Path

import docx

from packages.contracts.workbench_contracts import (
    AiTaskFromRegistryRequest,
    MedicalWritingRevisionRequest,
)
from services.api.app.ai_execution_policy import (
    AiExecutionPolicyDenied,
    AiExecutionPolicyResolver,
)
from services.api.app.ai_gateway import AiPromptEnvelope, DisabledAiProvider
from services.api.app.ai_task_runner import AiTaskRunner, AiTaskStore
from services.api.app.medical_writing import MedicalWritingRevisionService
from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_manifest import (
    D001_PROTOCOL_DOCX,
    MY008_PNH_3_01_PROTOCOL_DOCX,
    RUX_PROTOCOL_DOCX,
)
from services.api.app.medical_writing_repository import MedicalWritingRuntimeRepository
from services.api.app.protocol_text_extractor import parse_protocol_docx
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore
from services.api.app.sqlite_runtime_store import SqliteRuntimeStore


PROJECT_SOURCES = {
    "proj_rux_03_002": RUX_PROTOCOL_DOCX,
    "proj_d001": D001_PROTOCOL_DOCX,
    "proj_my008_pnh_3_01": MY008_PNH_3_01_PROTOCOL_DOCX,
}


class RegisteredRevisionProvider:
    provider_name = "buddy"
    model_name = "glm-5.2"

    def __init__(self, evidence_mode: str = "valid"):
        self.evidence_mode = evidence_mode
        self.envelopes: list[AiPromptEnvelope] = []

    def run(self, envelope: AiPromptEnvelope):
        self.envelopes.append(envelope)
        source = envelope.payload["allowed_sources"][0]
        locator = source["locator"]
        quote = source["text_preview"]
        if self.evidence_mode == "wrong_locator":
            locator = "docx:paragraph:999999"
        elif self.evidence_mode == "wrong_quote":
            quote = "这不是已登记来源中的原文。"
        return {
            "task_id": envelope.task_id,
            "task_type": envelope.task_type.value,
            "provider": self.provider_name,
            "model": self.model_name,
            "prompt_version": envelope.prompt_version,
            "input_source_ids": [source["source_id"]],
            "forbidden_source_ids": envelope.payload["forbidden_source_ids"],
            "findings": [
                {
                    "finding_id": "finding_revision_001",
                    "status": "supported",
                    "title": "方案正文修订候选",
                    "source_id": source["source_id"],
                    "evidence_span_ids": ["evidence_revision_001"],
                }
            ],
            "evidence_spans": [
                {
                    "span_id": "evidence_revision_001",
                    "source_id": source["source_id"],
                    "locator": locator,
                    "quote": quote,
                }
            ],
            "uncertainties": [
                {
                    "level": "medical_review",
                    "description": "需医学经理确认表述及全方案一致性。",
                }
            ],
            "needs_medical_confirmation": True,
            "schema_version": "ai_task_output_v0_1",
            "revision": {
                "proposal_text": f"{source['text_preview']}（修订候选）",
                "diff_patch": f"- {source['text_preview']}\n+ {source['text_preview']}（修订候选）",
                "rationale": "保持原始医学含义，仅优化方案正文表达。",
                "evidence_span_ids": ["evidence_revision_001"],
                "alternatives": [
                    {
                        "proposal_text": f"{source['text_preview']}（备选表述一）",
                        "diff_patch": "alternative 1",
                        "rationale": "提供等义的方案正文备选表述。",
                        "evidence_span_ids": ["evidence_revision_001"],
                    },
                    {
                        "proposal_text": f"{source['text_preview']}（备选表述二）",
                        "diff_patch": "alternative 2",
                        "rationale": "提供另一种不新增医学事实的表述。",
                        "evidence_span_ids": ["evidence_revision_001"],
                    },
                ],
            },
        }


@unittest.skipUnless(
    RUX_PROTOCOL_DOCX.exists() and D001_PROTOCOL_DOCX.exists(),
    "real RUX and D001 protocol DOCX sources are unavailable",
)
class MedicalWritingRegisteredSourceTests(unittest.TestCase):
    def setUp(self):
        self.tmpdir = tempfile.TemporaryDirectory()
        self.root = Path(self.tmpdir.name)
        self.registry = SourceRegistryService(
            SourceRegistryStore(self.root / "source_registry.jsonl"),
            allowed_roots=[path.parent for path in PROJECT_SOURCES.values()],
        )

    def tearDown(self):
        self.tmpdir.cleanup()

    def _runner(self, provider: RegisteredRevisionProvider) -> AiTaskRunner:
        return AiTaskRunner(
            MedicalWritingDocumentService(),
            AiTaskStore(self.root / f"ai_runs_{len(provider.envelopes)}.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name=provider.provider_name,
                model_name=provider.model_name,
                test_only_provider_injection=True,
            ),
        )

    def _registered_request(self, registration) -> AiTaskFromRegistryRequest:
        return AiTaskFromRegistryRequest(
            module="medical_writing",
            task_type="medical_writing_revision",
            expected_prompt_version="medical_writing_revision_v1_4",
            source_ids=[registration.spans[0].source_id],
            expected_source_entry_ids=[registration.entry.entry_id],
            user_instruction="仅基于登记的原始段落生成供医学作者确认的修订候选。",
        )

    def _real_repository_and_selection(self, project_id: str):
        document_service = MedicalWritingDocumentService()
        repository = MedicalWritingRuntimeRepository(
            document_service,
            SqliteRuntimeStore(self.root / f"{project_id}_runtime.sqlite3"),
        )
        session = document_service.document_session(project_id)
        section = next(
            document_service.section(project_id, item.section_id)
            for item in session.sections
            if document_service.section(project_id, item.section_id).content_blocks
        )
        block = next(
            item
            for item in section.content_blocks
            if isinstance(item.get("text"), str) and item["text"].strip()
        )
        return repository, section, block

    def test_rux_and_d001_can_register_one_selected_paragraph_after_span_200(self):
        registrations = {}
        for project_id, protocol_path in PROJECT_SOURCES.items():
            with self.subTest(project_id=project_id):
                document = parse_protocol_docx(protocol_path)
                target = next(
                    span
                    for span in document.spans[250:]
                    if span.kind == "paragraph" and len(span.text) >= 20
                )

                result = self.registry.register_protocol_selection_from_file(
                    project_id,
                    protocol_path,
                    locator=target.source_locator,
                    quote=target.text,
                    module="medical_writing",
                )

                self.assertEqual(1, result.entry.span_count)
                self.assertEqual("medical_writing", result.entry.module)
                self.assertEqual("protocol_docx_selection", result.entry.source_kind)
                self.assertEqual(target.source_locator, result.spans[0].locator)
                self.assertEqual(target.text, result.spans[0].text_preview)
                self.assertNotIn(
                    "/Users/",
                    json.dumps(result.model_dump(mode="json"), ensure_ascii=False),
                )
                registrations[project_id] = result

        self.assertNotEqual(
            registrations["proj_rux_03_002"].entry.entry_id,
            registrations["proj_d001"].entry.entry_id,
        )

    def test_real_revision_service_uses_registered_source_and_persists_opaque_policy_metadata(
        self,
    ):
        document_service = MedicalWritingDocumentService()
        repository = MedicalWritingRuntimeRepository(
            document_service,
            SqliteRuntimeStore(self.root / "runtime.sqlite3"),
        )
        provider = RegisteredRevisionProvider()
        runner = self._runner(provider)
        service = MedicalWritingRevisionService(
            repository,
            runner,
            source_registry=self.registry,
            protocol_source_paths=PROJECT_SOURCES,
        )

        for project_id in PROJECT_SOURCES:
            with self.subTest(project_id=project_id):
                session = document_service.document_session(project_id)
                section = next(
                    document_service.section(project_id, item.section_id)
                    for item in session.sections
                    if any(
                        isinstance(block.get("text"), str) and block["text"].strip()
                        for block in document_service.section(
                            project_id, item.section_id
                        ).content_blocks
                    )
                )
                block = next(
                    item
                    for item in section.content_blocks
                    if isinstance(item.get("text"), str) and item["text"].strip()
                )
                result = service.submit_revision(
                    project_id,
                    MedicalWritingRevisionRequest(
                        document_id=section.document_id,
                        section_id=section.section_id,
                        anchor_path=block["source_locator"],
                        selected_text=block["text"],
                        user_instruction="优化为待医学审阅的方案正文候选。",
                    ),
                )
                run = runner.get(project_id, result.thread.ai_run_id)

                self.assertEqual("registered_sources", run.request_origin)
                self.assertEqual("medical_writing", run.input_sources[0].module)
                self.assertEqual(
                    run.input_sources[0].source_id, result.thread.source_id
                )
                self.assertEqual(
                    run.input_sources[0].locator, result.thread.source_locator
                )
                self.assertTrue(result.thread.source_entry_id.startswith("src_"))
                self.assertEqual(
                    run.policy_decision_id, result.thread.ai_policy_decision_id
                )
                self.assertNotIn(
                    "/Users/",
                    json.dumps(
                        result.thread.model_dump(mode="json"), ensure_ascii=False
                    ),
                )

        self.assertEqual(3, len(provider.envelopes))

    def test_provider_evidence_locator_and_quote_must_match_registered_selection(self):
        protocol_path = RUX_PROTOCOL_DOCX
        document = parse_protocol_docx(protocol_path)
        target = next(span for span in document.spans[250:] if len(span.text) >= 20)
        registration = self.registry.register_protocol_selection_from_file(
            "proj_rux_03_002",
            protocol_path,
            locator=target.source_locator,
            quote=target.text,
            module="medical_writing",
        )

        for mode, expected_error in (
            ("wrong_locator", "locator"),
            ("wrong_quote", "quote"),
        ):
            with self.subTest(mode=mode):
                provider = RegisteredRevisionProvider(mode)
                run = self._runner(provider).submit_registered(
                    "proj_rux_03_002",
                    self._registered_request(registration),
                    self.registry,
                )
                self.assertEqual("failed", run.status)
                self.assertTrue(
                    any(expected_error in error for error in run.validation_errors),
                    run.validation_errors,
                )

    def test_multiple_selections_share_one_document_entry_without_losing_spans(self):
        content = _protocol_docx_bytes("同一版本来源")
        document = parse_protocol_docx("same_protocol.docx", content)
        first = document.spans[220]
        second = document.spans[221]

        for span in (first, second):
            self.registry.register_protocol_selection(
                "proj_synthetic_writing",
                "same_protocol.docx",
                content,
                locator=span.source_locator,
                quote=span.text,
                module="medical_writing",
            )

        entries = self.registry.list_entries("proj_synthetic_writing")
        spans = self.registry.list_spans("proj_synthetic_writing")
        self.assertEqual(1, len(entries))
        self.assertEqual(2, entries[0].span_count)
        self.assertEqual(2, len({span.source_id for span in spans}))

    def test_real_revision_missing_registry_fails_before_provider_and_thread_creation(
        self,
    ):
        repository, section, block = self._real_repository_and_selection(
            "proj_rux_03_002"
        )
        provider = RegisteredRevisionProvider()
        service = MedicalWritingRevisionService(repository, self._runner(provider))

        with self.assertRaisesRegex(ValueError, "Source Registry"):
            service.submit_revision(
                "proj_rux_03_002",
                MedicalWritingRevisionRequest(
                    document_id=section.document_id,
                    section_id=section.section_id,
                    anchor_path=block["source_locator"],
                    selected_text=block["text"],
                    user_instruction="请修订。",
                ),
            )

        self.assertEqual([], provider.envelopes)
        self.assertEqual([], repository.revision_threads("proj_rux_03_002"))

    def test_registered_revision_missing_policy_profile_fails_before_provider(self):
        document = parse_protocol_docx(RUX_PROTOCOL_DOCX)
        target = next(span for span in document.spans[250:] if len(span.text) >= 20)
        registration = self.registry.register_protocol_selection_from_file(
            "proj_rux_03_002",
            RUX_PROTOCOL_DOCX,
            locator=target.source_locator,
            quote=target.text,
            module="medical_writing",
        )
        provider = RegisteredRevisionProvider()
        runner = AiTaskRunner(
            MedicalWritingDocumentService(),
            AiTaskStore(self.root / "disabled_policy_runs.jsonl"),
            provider_factory=lambda resolution: provider,
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="disabled",
                provider_name=provider.provider_name,
                model_name=provider.model_name,
            ),
        )

        with self.assertRaisesRegex(AiExecutionPolicyDenied, "disabled or unapproved"):
            runner.submit_registered(
                "proj_rux_03_002",
                self._registered_request(registration),
                self.registry,
            )
        self.assertEqual([], provider.envelopes)

    def test_real_revision_missing_provider_fails_without_thread_creation(self):
        repository, section, block = self._real_repository_and_selection("proj_d001")
        runner = AiTaskRunner(
            repository,
            AiTaskStore(self.root / "disabled_provider_runs.jsonl"),
            provider_factory=lambda resolution: DisabledAiProvider(),
            policy_resolver=AiExecutionPolicyResolver(
                deployment_profile="approved_private_documents",
                provider_name="disabled",
                model_name="not_configured",
            ),
        )
        service = MedicalWritingRevisionService(
            repository,
            runner,
            source_registry=self.registry,
            protocol_source_paths=PROJECT_SOURCES,
        )

        # Unauthorized/non-product provider must fail closed before thread creation.
        with self.assertRaisesRegex(
            ValueError,
            "product-owned approved direct route|独立AI未配置",
        ):
            service.submit_revision(
                "proj_d001",
                MedicalWritingRevisionRequest(
                    document_id=section.document_id,
                    section_id=section.section_id,
                    anchor_path=block["source_locator"],
                    selected_text=block["text"],
                    user_instruction="请修订。",
                ),
            )
        self.assertEqual([], repository.revision_threads("proj_d001"))

    def test_cross_project_and_stale_registered_sources_fail_before_provider_call(self):
        document = parse_protocol_docx(D001_PROTOCOL_DOCX)
        target = next(span for span in document.spans[250:] if len(span.text) >= 20)
        d001_registration = self.registry.register_protocol_selection_from_file(
            "proj_d001",
            D001_PROTOCOL_DOCX,
            locator=target.source_locator,
            quote=target.text,
            module="medical_writing",
        )
        cross_project_provider = RegisteredRevisionProvider()

        with self.assertRaises(AiExecutionPolicyDenied):
            self._runner(cross_project_provider).submit_registered(
                "proj_rux_03_002",
                self._registered_request(d001_registration),
                self.registry,
            )
        self.assertEqual([], cross_project_provider.envelopes)

        source_v1 = _protocol_docx_bytes("版本一附加内容")
        source_v2 = _protocol_docx_bytes("版本二附加内容")
        selection_document = parse_protocol_docx("same_protocol.docx", source_v1)
        selected_span = selection_document.spans[220]
        old_registration = self.registry.register_protocol_selection(
            "proj_synthetic_writing",
            "same_protocol.docx",
            source_v1,
            locator=selected_span.source_locator,
            quote=selected_span.text,
            module="medical_writing",
        )
        self.registry.register_protocol_selection(
            "proj_synthetic_writing",
            "same_protocol.docx",
            source_v2,
            locator=selected_span.source_locator,
            quote=selected_span.text,
            module="medical_writing",
        )
        stale_provider = RegisteredRevisionProvider()

        with self.assertRaisesRegex(AiExecutionPolicyDenied, "stale"):
            self._runner(stale_provider).submit_registered(
                "proj_synthetic_writing",
                self._registered_request(old_registration),
                self.registry,
            )
        self.assertEqual([], stale_provider.envelopes)


def _protocol_docx_bytes(version_marker: str) -> bytes:
    document = docx.Document()
    for index in range(240):
        text = f"稳定的方案正文段落 {index:03d}。"
        if index == 220:
            text = "该段落用于验证按需来源注册和陈旧版本阻断。"
        document.add_paragraph(text)
    document.add_paragraph(version_marker)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
