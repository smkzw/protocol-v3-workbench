from __future__ import annotations

import io
import zipfile
from datetime import datetime, timezone

import pytest
from docx import Document
from lxml import etree

from packages.contracts.workbench_contracts import (
    ApprovalState,
    MedicalWritingLiteratureLibrary,
    MedicalWritingProjectReference,
    ProtocolDocument,
    ProtocolSection,
    StructuredTable,
    StructuredTableCell,
    StructuredTableColumn,
    StructuredTableDomain,
    StructuredTableRow,
)
from services.api.app.medical_writing_document_exporter import (
    MedicalWritingDocumentDocxExportError,
    export_medical_writing_document_docx,
)
from services.api.app import medical_writing_document_exporter
from services.api.app.medical_writing_tables import MedicalWritingTableService


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WORD_NS}


def test_citations_are_numbered_by_first_occurrence_and_link_to_generated_references():
    document = _document(
        _citation_block(
            "block_body_1",
            1,
            [
                ("既往研究", None),
                ("[9]", "ref_b"),
                ("支持该设计，重复引用", None),
                ("[9]", "ref_b"),
                ("。", None),
            ],
        ),
        _citation_block(
            "block_body_2",
            2,
            [("另一项研究", None), ("[4]", "ref_a"), ("提供补充证据。", None)],
        ),
    )
    library = _library(_reference("ref_a"), _reference("ref_b", second=True))

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=library,
    )

    root = _document_xml(result.content)
    hyperlinks = root.xpath(".//w:hyperlink", namespaces=NS)
    assert ["".join(item.xpath(".//w:t/text()", namespaces=NS)) for item in hyperlinks] == [
        "[1]",
        "[1]",
        "[2]",
    ]
    assert all(
        item.xpath(".//w:vertAlign[@w:val='superscript']", namespaces=NS)
        for item in hyperlinks
    )

    bookmark_names = set(
        root.xpath(".//w:bookmarkStart/@w:name", namespaces=NS)
    )
    hyperlink_anchors = [
        item.get(f"{{{WORD_NS}}}anchor") for item in hyperlinks
    ]
    assert set(hyperlink_anchors) <= bookmark_names
    assert hyperlink_anchors[0] == hyperlink_anchors[1]
    assert hyperlink_anchors[0] != hyperlink_anchors[2]

    paragraph_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]
    heading_index = paragraph_text.index("参考文献")
    assert paragraph_text[heading_index + 1].startswith("[1] Hewitt M, Berry M, Provan D, et al.")
    assert "Second evidence title[J]. Journal B, 2022, 7(2): 20-29." in paragraph_text[
        heading_index + 1
    ]
    assert paragraph_text[heading_index + 2].startswith("[2] Zhang San, Li Si.")
    assert "First evidence title[J]. Journal A, 2021, 6(1): 10-19." in paragraph_text[
        heading_index + 2
    ]
    assert result.metadata["citation_style"] == "gbt_7714_2015_numeric"
    assert result.metadata["cited_reference_ids"] == ["ref_b", "ref_a"]


def test_missing_cited_reference_fails_explicitly():
    document = _document(
        _citation_block("block_body", 1, [("证据", None), ("[1]", "ref_missing")])
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="cited reference does not exist.*ref_missing",
    ):
        export_medical_writing_document_docx(
            document,
            mode="draft_preview",
            literature_library=_library(_reference("ref_a")),
        )


def test_multi_reference_mark_keeps_one_visual_group_and_links_each_reference():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "联合证据[8,9]",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "联合证据"},
                    {
                        "type": "text",
                        "text": "[8,9]",
                        "marks": [
                            {
                                "type": "citation",
                                "attrs": {"referenceIds": ["ref_b", "ref_a"]},
                            }
                        ],
                    },
                ],
            },
        }
    )
    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a"), _reference("ref_b", second=True)),
    )
    root = _document_xml(result.content)
    hyperlinks = root.xpath(".//w:hyperlink", namespaces=NS)
    assert [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in hyperlinks
    ] == ["1", "2"]
    paragraph_texts = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]
    assert "联合证据[1,2]" in paragraph_texts
    assert paragraph_texts.index("目录") < paragraph_texts.index("联合证据[1,2]")
    assert hyperlinks[0].get(f"{{{WORD_NS}}}anchor") != hyperlinks[1].get(
        f"{{{WORD_NS}}}anchor"
    )
    assert result.metadata["cited_reference_ids"] == ["ref_b", "ref_a"]


def test_needs_review_cited_reference_fails_explicitly():
    document = _document(
        _citation_block("block_body", 1, [("证据", None), ("[1]", "ref_a")])
    )
    reference = _reference("ref_a").model_copy(
        update={"validation_status": "needs_review"}
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="requires review before export: ref_a",
    ):
        export_medical_writing_document_docx(
            document,
            mode="draft_preview",
            literature_library=_library(reference),
        )


def test_no_citation_export_is_byte_identical_when_an_unused_library_is_provided():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "正文没有引文。",
            "rich_text": {
                "type": "paragraph",
                "content": [{"type": "text", "text": "正文没有引文。"}],
            },
        }
    )

    without_library = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
    )
    with_library = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a")),
    )

    assert with_library.content == without_library.content
    assert with_library.metadata == without_library.metadata
    exported = Document(io.BytesIO(with_library.content))
    assert sum(paragraph.text == "参考文献" for paragraph in exported.paragraphs) == 1


def test_generated_reference_entries_stay_in_reference_section_before_later_content():
    document = _document(
        _citation_block("block_body", 1, [("证据", None), ("[7]", "ref_a")])
    )
    document.sections.append(
        ProtocolSection(
            section_id="section_appendix",
            document_id=document.document_id,
            heading="附录",
            content_blocks=[
                {
                    "block_id": "block_appendix",
                    "block_type": "heading",
                    "body_order": 200,
                    "text": "附录",
                }
            ],
        )
    )

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a")),
    )
    root = _document_xml(result.content)
    paragraph_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]

    reference_heading = paragraph_text.index("参考文献")
    reference_entry = next(
        index for index, text in enumerate(paragraph_text) if text.startswith("[1] ")
    )
    appendix = paragraph_text.index("附录")
    assert reference_heading < reference_entry < appendix


def test_imported_and_managed_citations_share_one_first_occurrence_index():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "旧证据[19]、新增证据[88]、IGA[2，3]及另一旧证据[1]。",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "旧证据"},
                    {
                        "type": "text",
                        "text": "[19]",
                        "marks": [{"type": "superscript"}],
                    },
                    {"type": "text", "text": "、新增证据"},
                    {
                        "type": "text",
                        "text": "[88]",
                        "marks": [
                            {"type": "citation", "attrs": {"referenceId": "ref_a"}}
                        ],
                    },
                    {"type": "text", "text": "、IGA[2，3]及另一旧证据"},
                    {
                        "type": "text",
                        "text": "[1]",
                        "marks": [{"type": "superscript"}],
                    },
                    {"type": "text", "text": "。"},
                ],
            },
        }
    )
    document.sections[1].content_blocks.extend(
        [
            {
                "block_id": "legacy_reference_1",
                "block_type": "paragraph",
                "body_order": 101,
                "text": "[1] Existing reference one[J]. Journal, 2020.",
            },
            {
                "block_id": "legacy_reference_19",
                "block_type": "paragraph",
                "body_order": 119,
                "text": "[19] Existing reference nineteen[J]. Journal, 2024.",
            },
        ]
    )

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a")),
    )
    root = _document_xml(result.content)
    hyperlink_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:hyperlink", namespaces=NS)
    ]
    paragraph_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]

    assert hyperlink_text == ["[1]", "[2]", "[3]"]
    assert any(text.startswith("[1] Existing reference nineteen") for text in paragraph_text)
    assert any(text.startswith("[2] Zhang San, Li Si.") for text in paragraph_text)
    assert any(text.startswith("[3] Existing reference one") for text in paragraph_text)
    assert not any(text.startswith("[19]") for text in paragraph_text)
    assert sum(text == "参考文献" for text in paragraph_text) == 1
    assert result.metadata["citation_count"] == 3
    assert result.metadata["legacy_reference_count"] == 2
    assert result.metadata["uncited_reference_ids"] == []
    assert "number_gap" in result.metadata["reference_index_issue_codes"]


def test_managed_doi_matching_imported_reference_is_deduplicated():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "旧引文[1]与新增引用[9]指向同一文献。",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "旧引文"},
                    {
                        "type": "text",
                        "text": "[1]",
                        "marks": [{"type": "superscript"}],
                    },
                    {"type": "text", "text": "与新增引用"},
                    {
                        "type": "text",
                        "text": "[9]",
                        "marks": [
                            {"type": "citation", "attrs": {"referenceId": "ref_a"}}
                        ],
                    },
                    {"type": "text", "text": "指向同一文献。"},
                ],
            },
        }
    )
    document.sections[1].content_blocks.append(
        {
            "block_id": "legacy_reference_1",
            "block_type": "paragraph",
            "body_order": 101,
            "text": "[1] Existing matching reference[J]. Journal, 2020. DOI: 10.1000/ref_a.",
        }
    )

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a")),
    )
    root = _document_xml(result.content)
    hyperlink_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:hyperlink", namespaces=NS)
    ]
    paragraph_text = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]

    assert hyperlink_text == ["[1]", "[1]"]
    assert sum(text.startswith("[1] Existing matching reference") for text in paragraph_text) == 1
    assert result.metadata["citation_count"] == 1


@pytest.mark.parametrize(
    ("legacy_text", "managed_update"),
    [
        (
            "[1] Alpha. Legacy title[J]. Journal, 2021. PMID: 34567890.",
            {
                "doi": "",
                "pmid": "34567890",
                "canonical_key": "pmid:34567890",
            },
        ),
        (
            "[1] Alpha. Legacy title[J]. Journal, 2021. https://www.example.org/article/42/.",
            {
                "doi": "",
                "url": "https://example.org/article/42",
                "canonical_key": "url:https://example.org/article/42",
            },
        ),
        (
            "[1] Alpha. First Evidence Title[J]. Journal, 2021.",
            {"doi": "", "canonical_key": "title_year:firstevidencetitle2021"},
        ),
    ],
)
def test_managed_reference_uses_deterministic_legacy_identity_fallbacks(
    legacy_text,
    managed_update,
):
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "旧引文[1]与新增引用[待编号]指向同一文献。",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "旧引文"},
                    {
                        "type": "text",
                        "text": "[1]",
                        "marks": [{"type": "superscript"}],
                    },
                    {"type": "text", "text": "与新增引用"},
                    {
                        "type": "text",
                        "text": "[待编号]",
                        "marks": [
                            {"type": "citation", "attrs": {"referenceId": "ref_a"}}
                        ],
                    },
                    {"type": "text", "text": "指向同一文献。"},
                ],
            },
        }
    )
    document.sections[1].content_blocks.append(
        {
            "block_id": "legacy_reference_1",
            "block_type": "paragraph",
            "body_order": 101,
            "text": legacy_text,
        }
    )
    managed = _reference("ref_a").model_copy(update=managed_update)

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(managed),
    )
    root = _document_xml(result.content)

    assert [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:hyperlink", namespaces=NS)
    ] == ["[1]", "[1]"]
    assert result.metadata["citation_count"] == 1
    assert result.metadata["reference_count"] == 1


def test_english_references_heading_is_reused_without_duplicate_section():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "Evidence[1]",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "Evidence"},
                    {
                        "type": "text",
                        "text": "[1]",
                        "marks": [{"type": "superscript"}],
                    },
                ],
            },
        }
    )
    document.sections[1].heading = "14 References"
    document.sections[1].content_blocks[0]["text"] = "References"
    document.sections[1].content_blocks.append(
        {
            "block_id": "legacy_reference_1",
            "block_type": "paragraph",
            "body_order": 101,
            "text": "[1] Alpha. Evidence title[J]. Journal, 2021.",
        }
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    paragraph_text = [paragraph.text for paragraph in Document(io.BytesIO(result.content)).paragraphs]

    assert paragraph_text.count("References") == 1
    assert "参考文献" not in paragraph_text
    heading_index = paragraph_text.index("References")
    assert paragraph_text[heading_index + 1].startswith("[1] Alpha. Evidence title")


def test_weak_managed_identity_matching_conflicting_strong_legacy_ids_is_rejected():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "旧引文[1,2]与新增引用[待编号]。",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "旧引文"},
                    {
                        "type": "text",
                        "text": "[1,2]",
                        "marks": [{"type": "superscript"}],
                    },
                    {"type": "text", "text": "与新增引用"},
                    {
                        "type": "text",
                        "text": "[待编号]",
                        "marks": [
                            {"type": "citation", "attrs": {"referenceId": "ref_a"}}
                        ],
                    },
                    {"type": "text", "text": "。"},
                ],
            },
        }
    )
    document.sections[1].content_blocks.extend(
        [
            {
                "block_id": "legacy_reference_1",
                "block_type": "paragraph",
                "body_order": 101,
                "text": "[1] Alpha. Article[J]. Journal, 2021. DOI: 10.1000/one. PMID: 34567890.",
            },
            {
                "block_id": "legacy_reference_2",
                "block_type": "paragraph",
                "body_order": 102,
                "text": "[2] Alpha. Article[J]. Journal, 2021. DOI: 10.1000/two. PMID: 34567890.",
            },
        ]
    )
    managed = _reference("ref_a").model_copy(
        update={
            "doi": "",
            "pmid": "34567890",
            "canonical_key": "pmid:34567890",
        }
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="managed reference matches multiple imported references: ref_a",
    ):
        export_medical_writing_document_docx(
            document,
            mode="draft_preview",
            literature_library=_library(managed),
        )


def test_unparsed_reference_list_entry_blocks_export_instead_of_mixing_lists():
    document = _document(
        _citation_block(
            "managed_citation",
            1,
            [("新增证据", None), ("[待编号]", "ref_a")],
        )
    )
    document.sections[1].content_blocks.append(
        {
            "block_id": "legacy_unparsed",
            "block_type": "paragraph",
            "body_order": 101,
            "text": "(1) Word auto-numbered reference.",
        }
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="unparsed_entry:legacy_unparsed",
    ):
        export_medical_writing_document_docx(
            document,
            mode="draft_preview",
            literature_library=_library(_reference("ref_a")),
        )


def test_legacy_grouped_aliases_render_one_canonical_number_once():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "同一证据[1,2]",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "同一证据"},
                    {
                        "type": "text",
                        "text": "[1,2]",
                        "marks": [{"type": "superscript"}],
                    },
                ],
            },
        }
    )
    document.sections[1].content_blocks.extend(
        [
            {
                "block_id": "legacy_reference_1",
                "block_type": "paragraph",
                "body_order": 101,
                "text": "[1] Alpha. Same article[J]. Journal, 2021. DOI: 10.1000/same.",
            },
            {
                "block_id": "legacy_reference_2",
                "block_type": "paragraph",
                "body_order": 102,
                "text": "[2] Beta. Other rendering[J]. Journal, 2021. DOI: 10.1000/same.",
            },
        ]
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    root = _document_xml(result.content)

    assert [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:hyperlink", namespaces=NS)
    ] == ["[1]"]
    paragraph_texts = [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in root.xpath(".//w:body/w:p", namespaces=NS)
    ]
    assert "同一证据[1]" in paragraph_texts
    assert paragraph_texts.index("目录") < paragraph_texts.index("同一证据[1]")
    assert result.metadata["reference_count"] == 1


def test_superscript_numbers_without_a_reference_list_are_preserved_not_guessed():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "量表条目[1]",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "量表条目"},
                    {
                        "type": "text",
                        "text": "[1]",
                        "marks": [{"type": "superscript"}],
                    },
                ],
            },
        }
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    root = _document_xml(result.content)

    assert root.xpath(".//w:hyperlink", namespaces=NS) == []
    assert "citation_style" not in result.metadata
    assert "量表条目[1]" in "".join(root.xpath(".//w:t/text()", namespaces=NS))


def test_partial_legacy_reference_list_with_missing_cited_number_blocks_export():
    document = _document(
        {
            "block_id": "block_body",
            "block_type": "paragraph",
            "body_order": 1,
            "text": "证据[2]",
            "rich_text": {
                "type": "paragraph",
                "content": [
                    {"type": "text", "text": "证据"},
                    {
                        "type": "text",
                        "text": "[2]",
                        "marks": [{"type": "superscript"}],
                    },
                ],
            },
        }
    )
    document.sections[1].content_blocks.append(
        {
            "block_id": "legacy_reference_1",
            "block_type": "paragraph",
            "body_order": 101,
            "text": "[1] Available reference[J]. Journal, 2020.",
        }
    )

    with pytest.raises(
        MedicalWritingDocumentDocxExportError,
        match="missing_entry:2",
    ):
        export_medical_writing_document_docx(document, mode="draft_preview")


def test_table_cell_citation_uses_the_same_number_and_bookmark_plan():
    cell = StructuredTableCell(
        cell_id="cell_evidence",
        row_id="row_evidence",
        column_id="column_evidence",
        text="表格证据[待编号]",
        rich_text={
            "type": "doc",
            "content": [
                {
                    "type": "paragraph",
                    "content": [
                        {"type": "text", "text": "表格证据"},
                        {
                            "type": "text",
                            "text": "[待编号]",
                            "marks": [
                                {
                                    "type": "citation",
                                    "attrs": {"referenceIds": ["ref_a"]},
                                }
                            ],
                        },
                    ],
                }
            ],
        },
    )
    table = StructuredTable(
        table_id="table_evidence",
        block_id="block_table_evidence",
        domain=StructuredTableDomain.GENERIC,
        title="",
        columns=[
            StructuredTableColumn(
                column_id="column_evidence",
                order=0,
                label="证据",
                width_twips=3600,
            )
        ],
        rows=[
            StructuredTableRow(
                row_id="row_evidence",
                order=0,
                cells=[cell],
            )
        ],
    )
    block = MedicalWritingTableService().to_table_block(table)
    block["body_order"] = 1
    document = _document(block)

    result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        literature_library=_library(_reference("ref_a")),
    )
    root = _document_xml(result.content)
    hyperlinks = root.xpath(".//w:tbl//w:hyperlink", namespaces=NS)
    assert [
        "".join(item.xpath(".//w:t/text()", namespaces=NS))
        for item in hyperlinks
    ] == ["[1]"]
    anchor = hyperlinks[0].get(f"{{{WORD_NS}}}anchor")
    assert anchor in set(root.xpath(".//w:bookmarkStart/@w:name", namespaces=NS))
    assert "表格证据[1]" in "".join(root.xpath(".//w:tbl//w:t/text()", namespaces=NS))


def _document(*body_blocks: dict) -> ProtocolDocument:
    body = ProtocolSection(
        section_id="section_body",
        document_id="document_citation",
        heading="研究背景",
        approval_state=ApprovalState.IN_MEDICAL_REVIEW,
        content_blocks=list(body_blocks),
    )
    references = ProtocolSection(
        section_id="section_references",
        document_id="document_citation",
        heading="参考文献",
        section_number="14",
        approval_state=ApprovalState.IN_MEDICAL_REVIEW,
        content_blocks=[
            {
                "block_id": "block_references_heading",
                "block_type": "heading",
                "body_order": 100,
                "outline_level": 0,
                "text": "参考文献",
            }
        ],
    )
    return ProtocolDocument(
        document_id="document_citation",
        project_id="project_citation",
        protocol_id="CITATION-001",
        version="V0.1",
        sections=[body, references],
    )


def _citation_block(
    block_id: str,
    body_order: int,
    fragments: list[tuple[str, str | None]],
) -> dict:
    content = []
    for text, reference_id in fragments:
        node = {"type": "text", "text": text}
        if reference_id:
            node["marks"] = [
                {"type": "citation", "attrs": {"referenceId": reference_id}}
            ]
        content.append(node)
    return {
        "block_id": block_id,
        "block_type": "paragraph",
        "body_order": body_order,
        "text": "".join(text for text, _ in fragments),
        "rich_text": {"type": "paragraph", "content": content},
    }


def _library(*references: MedicalWritingProjectReference) -> MedicalWritingLiteratureLibrary:
    return MedicalWritingLiteratureLibrary(
        project_id="project_citation",
        citation_style="gbt_7714_2015_numeric",
        references=list(references),
    )


def _reference(
    reference_id: str,
    *,
    second: bool = False,
) -> MedicalWritingProjectReference:
    now = datetime(2026, 7, 15, tzinfo=timezone.utc)
    return MedicalWritingProjectReference(
        reference_id=reference_id,
        project_id="project_citation",
        canonical_key=f"doi:10.1000/{reference_id}",
        source_kind="doi",
        source_input=f"10.1000/{reference_id}",
        title="Second evidence title" if second else "First evidence title",
        authors=(
            ["Hewitt M", "Berry M", "Provan D", "A fourth author"]
            if second
            else ["Zhang San", "Li Si"]
        ),
        journal="Journal B" if second else "Journal A",
        year="2022" if second else "2021",
        volume="7" if second else "6",
        issue="2" if second else "1",
        pages="20-29" if second else "10-19",
        doi=f"10.1000/{reference_id}",
        validation_status="confirmed",
        created_at=now,
        updated_at=now,
    )


def _document_xml(content: bytes):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return etree.fromstring(archive.read("word/document.xml"))


def _make_minimal_docx_with_citation():
    """Create a minimal DOCX bytes that will have reference groups when scanned."""
    from docx import Document
    from io import BytesIO
    
    # Create a document with a reference section that has explicit entries
    doc = Document()
    doc.add_paragraph("Some text with [1] citation.")
    
    # Add a references section heading
    doc.add_heading("参考文献", level=1)
    doc.add_paragraph("[1] Some reference entry.")
    
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.read()


from unittest import mock


# Worker_01 Round 3: Correct focused tests for _apply_source_reference_reindex_for_export
# Status matrix per product contract:
# - applied: uar=False, final allowed
# - preserved: uar=False, final allowed (preserve_with_notice is NOT pending!)
# - not_applicable: uar=False, final allowed  
# - blocked: uar=True, draft shows actionable warning, final raises error


def test_apply_source_reference_reindex_applied_status():
    """applied status: direct function test, both modes allow."""
    from services.api.app.medical_writing_document_exporter import (
        _apply_source_reference_reindex_for_export,
    )
    
    docx_bytes = _make_minimal_docx_with_citation()
    
    # Create valid manifests with actual reference groups
    from services.api.app.medical_writing_source_reference_reindex import (
        build_source_reference_manifest,
        decide_source_reference_reindex,
    )
    
    initial_manifest = build_source_reference_manifest(docx_bytes)
    current_manifest = build_source_reference_manifest(docx_bytes)
    
    # Patch to return apply decision
    with mock.patch.object(
        medical_writing_document_exporter, 
        'decide_source_reference_reindex',
        return_value=mock.Mock(
            action="apply",
            number_mapping=tuple((n, n) for n in range(1, len(initial_manifest.reference_groups)+1)),
            issues=[],
        )
    ):
        with mock.patch.object(
            medical_writing_document_exporter,
            'apply_source_reference_reindex',
            side_effect=lambda c, m, d: c  # pass-through
        ):
            result, meta = _apply_source_reference_reindex_for_export(
                docx_bytes,
                source_content=docx_bytes,
                managed_citation_plan=None,
                export_mode="draft_preview",
            )
            
            assert meta["status"] == "applied"
            assert meta["action"] == "apply"
            assert meta["user_action_required"] is False
            assert meta["warning_message"] == ""
            
            # Final mode should also work
            result_final, meta_final = _apply_source_reference_reindex_for_export(
                docx_bytes,
                source_content=docx_bytes,
                managed_citation_plan=None,
                export_mode="approved_final",
            )
            assert meta_final["status"] == "applied"
            assert meta_final["user_action_required"] is False


def test_apply_source_reference_reindex_preserved_status():
    """preserved status (preserve_with_notice): final allowed, no user action."""
    from services.api.app.medical_writing_document_exporter import (
        _apply_source_reference_reindex_for_export,
    )
    
    docx_bytes = _make_minimal_docx_with_citation()
    
    from services.api.app.medical_writing_source_reference_reindex import (
        build_source_reference_manifest,
    )
    
    manifest = build_source_reference_manifest(docx_bytes)
    
    with mock.patch.object(
        medical_writing_document_exporter,
        'decide_source_reference_reindex',
        return_value=mock.Mock(
            action="preserve_with_notice",
            number_mapping=tuple((n, n) for n in range(1, len(manifest.reference_groups)+1)),
            issues=[mock.Mock(code="uncited_reference_entry", severity="notice", message="uncited", blocking=False)],
        )
    ):
        result, meta = _apply_source_reference_reindex_for_export(
            docx_bytes,
            source_content=docx_bytes,
            managed_citation_plan=None,
            export_mode="draft_preview",
        )
        
        assert meta["status"] == "preserved"
        assert meta["action"] == "preserve_with_notice"
        assert meta["user_action_required"] is False
        assert meta["warning_message"] == ""  # preserved is not pending
        
        # Crucially: final must also pass
        result_final, meta_final = _apply_source_reference_reindex_for_export(
            docx_bytes,
            source_content=docx_bytes,
            managed_citation_plan=None,
            export_mode="approved_final",
        )
        assert meta_final["status"] == "preserved"
        assert meta_final["user_action_required"] is False


def test_apply_source_reference_reindex_blocked_status_draft_warning_final_error():
    """blocked status: draft returns warning with actionable instruction, final raises."""
    from services.api.app.medical_writing_document_exporter import (
        _apply_source_reference_reindex_for_export,
        MedicalWritingDocumentDocxExportError,
    )
    
    docx_bytes = _make_minimal_docx_with_citation()
    
    from services.api.app.medical_writing_source_reference_reindex import (
        build_source_reference_manifest,
    )
    
    manifest = build_source_reference_manifest(docx_bytes)
    
    # External citation manager issue
    issue = mock.Mock(
        code="citation_manager_field_preserved",
        severity="notice",
        message="manager field preserved",
        blocking=False,
    )
    
    with mock.patch.object(
        medical_writing_document_exporter,
        'decide_source_reference_reindex',
        return_value=mock.Mock(
            action="block",
            number_mapping=tuple((n, n) for n in range(1, len(manifest.reference_groups)+1)),
            issues=(issue,),
        )
    ):
        # Draft must return blocked with actionable Chinese instruction
        result, meta = _apply_source_reference_reindex_for_export(
            docx_bytes,
            source_content=docx_bytes,
            managed_citation_plan=None,
            export_mode="draft_preview",
        )
        
        assert meta["status"] == "blocked"
        assert meta["action"] == "block"
        assert meta["user_action_required"] is True
        assert len(meta["warning_message"]) > 0
        # Must contain actionable Word refresh instruction
        assert "Word" in meta["warning_message"] or "刷新" in meta["warning_message"] or "重新导入" in meta["warning_message"]
        
        # Final must raise exact error
        with pytest.raises(MedicalWritingDocumentDocxExportError) as exc_info:
            _apply_source_reference_reindex_for_export(
                docx_bytes,
                source_content=docx_bytes,
                managed_citation_plan=None,
                export_mode="approved_final",
            )
        assert "blocked" in str(exc_info.value).lower()


def test_apply_source_reference_reindex_not_applicable_status():
    """not_applicable status: both modes allow, no user action."""
    from services.api.app.medical_writing_document_exporter import (
        _apply_source_reference_reindex_for_export,
    )
    
    # Minimal DOCX without citations
    from docx import Document
    from io import BytesIO
    
    doc = Document()
    doc.add_paragraph("No citations here.")
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    docx_bytes = buf.read()
    
    result, meta = _apply_source_reference_reindex_for_export(
        docx_bytes,
        source_content=docx_bytes,
        managed_citation_plan=None,
        export_mode="draft_preview",
    )
    
    assert meta["status"] == "not_applicable"
    assert meta["user_action_required"] is False
    assert meta["warning_message"] == ""
    
    # Final must also pass
    result_final, meta_final = _apply_source_reference_reindex_for_export(
        docx_bytes,
        source_content=docx_bytes,
        managed_citation_plan=None,
        export_mode="approved_final",
    )
    assert meta_final["status"] == "not_applicable"
    assert meta_final["user_action_required"] is False


def test_apply_source_reference_reindex_unknown_action_fails_closed():
    """Unknown actions fail closed as blocked."""
    from services.api.app.medical_writing_document_exporter import (
        _apply_source_reference_reindex_for_export,
        MedicalWritingDocumentDocxExportError,
    )
    
    docx_bytes = _make_minimal_docx_with_citation()
    
    from services.api.app.medical_writing_source_reference_reindex import (
        build_source_reference_manifest,
    )
    
    manifest = build_source_reference_manifest(docx_bytes)
    
    with mock.patch.object(
        medical_writing_document_exporter,
        'decide_source_reference_reindex',
        return_value=mock.Mock(
            action="unknown_xyz_action",
            number_mapping=tuple((n, n) for n in range(1, len(manifest.reference_groups)+1)),
            issues=[],
        )
    ):
        result, meta = _apply_source_reference_reindex_for_export(
            docx_bytes,
            source_content=docx_bytes,
            managed_citation_plan=None,
            export_mode="draft_preview",
        )
        
        assert meta["status"] == "blocked"
        assert meta["user_action_required"] is True
        
        # Final must raise
        with pytest.raises(MedicalWritingDocumentDocxExportError):
            _apply_source_reference_reindex_for_export(
                docx_bytes,
                source_content=docx_bytes,
                managed_citation_plan=None,
                export_mode="approved_final",
            )
