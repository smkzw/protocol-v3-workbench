from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.oxml.ns import qn

from packages.contracts.workbench_contracts import (
    ApprovalState,
    ProtocolDocument,
    ProtocolSection,
    StructuredTableRole,
)
from services.api.app.medical_writing_document_exporter import (
    MedicalWritingDocumentDocxExportError,
    export_medical_writing_document_docx,
)
from services.api.app.medical_writing_greenfield import (
    _greenfield_document_object_table,
)
from services.api.app.medical_writing_style_profile import (
    MedicalWritingStyleProfileService,
)
from services.api.app.medical_writing_company_corpus import MedicalWritingCompanyCorpusService


def _document_with_profile(*, profile=True, digest_override="") -> ProtocolDocument:
    definition = MedicalWritingStyleProfileService().definition()
    section = ProtocolSection(
        section_id="section_1",
        document_id="document_1",
        heading="试验设计",
        approval_state=ApprovalState.AI_DRAFT,
        content_blocks=[
            {
                "block_id": "heading_1",
                "block_type": "heading",
                "text": "试验设计",
                "body_order": 1,
                "outline_level": 0,
            },
            {
                "block_id": "body_1",
                "block_type": "paragraph",
                "text": "本研究采用多中心、随机、开放标签、剂量探索设计。",
                "body_order": 2,
            },
        ],
    )
    return ProtocolDocument(
        document_id="document_1",
        project_id="project_1",
        protocol_id="CMS-D017-PNH",
        version="v0.2",
        sections=[section],
        style_profile_id=definition.style_profile_id if profile else "",
        style_profile_version=definition.style_profile_version if profile else "",
        style_profile_definition_sha256=(
            digest_override or definition.definition_sha256
        )
        if profile
        else "",
    )


def test_style_profile_uses_latest_user_authority_order_and_is_stable():
    service = MedicalWritingStyleProfileService()
    first = service.definition()
    second = service.definition()

    assert first.definition_sha256 == second.definition_sha256
    assert first.source_documents[0]["title"] == "CMS-D017-PNH-方案摘要_v0.2.docx"
    assert first.source_documents[0]["sha256"] == (
        "dcd60942b1a77b2b109583923d0126ae86481a4a12d44c873e26f010463ba4da"
    )
    assert "最高参照" in first.source_documents[0]["authority_role"]
    assert first.source_documents[1]["title"].startswith("CMS-D005")


def test_company_style_profile_is_applied_to_word_export():
    corpus = MedicalWritingCompanyCorpusService().definition()
    source = _document_with_profile().model_copy(
        update={
            "corpus_snapshot_id": corpus.snapshot_id,
            "corpus_snapshot_version": corpus.snapshot_version,
            "corpus_snapshot_sha256": corpus.snapshot_sha256,
        }
    )
    result = export_medical_writing_document_docx(source, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))
    definition = MedicalWritingStyleProfileService().definition()

    assert result.metadata["style_profile_id"] == definition.style_profile_id
    assert result.metadata["style_profile_version"] == definition.style_profile_version
    assert result.metadata["style_profile_definition_sha256"] == definition.definition_sha256
    assert result.metadata["corpus_snapshot_id"] == corpus.snapshot_id
    assert result.metadata["corpus_snapshot_version"] == corpus.snapshot_version
    assert result.metadata["corpus_snapshot_sha256"] == corpus.snapshot_sha256

    section = exported.sections[0]
    assert section.top_margin.twips >= section.header_distance.twips + 790
    assert section.bottom_margin.mm == pytest.approx(20.0, abs=0.05)
    assert section.left_margin.mm == pytest.approx(25.0, abs=0.05)
    assert section.right_margin.mm == pytest.approx(25.0, abs=0.05)

    heading = next(
        paragraph for paragraph in exported.paragraphs if paragraph.text == "试验设计"
    )
    body = next(
        paragraph
        for paragraph in exported.paragraphs
        if paragraph.text == "本研究采用多中心、随机、开放标签、剂量探索设计。"
    )
    assert heading.runs[0].font.size.pt == pytest.approx(12.0)
    assert heading.runs[0].bold is True
    assert heading.runs[0].font.name == "Times New Roman"
    assert heading.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert heading.paragraph_format.line_spacing == pytest.approx(1.5)
    assert body.runs[0].font.size.pt == pytest.approx(12.0)
    assert body.runs[0].font.name == "Times New Roman"
    assert body.runs[0]._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
    assert body.paragraph_format.first_line_indent.pt == pytest.approx(24.0)
    assert body.paragraph_format.line_spacing == pytest.approx(1.5)


def test_greenfield_export_uses_editable_cover_native_heading_styles_and_numbered_toc():
    definition = MedicalWritingStyleProfileService().definition()
    front_section_id = "front_matter"
    front_table = _greenfield_document_object_table(
        section_id=front_section_id,
        document_id="document_greenfield",
        body_order=1,
        title="方案首页信息",
        role=StructuredTableRole.LAYOUT,
        rows=[
            ("方案编号", "CMS-D017-TEST"),
            ("版本", "V0.1"),
            ("方案标题", "一项评价CMS-D017治疗阵发性睡眠性血红蛋白尿症的临床试验"),
            ("适应症", "阵发性睡眠性血红蛋白尿症"),
            ("研究分期", "II期"),
            ("研究药物", "CMS-D017"),
            ("版本日期", "2026年07月20日"),
            ("申办者", "深圳市康哲生物科技有限公司"),
        ],
        source_fact_ids=[],
    )
    sections = [
        ProtocolSection(
            section_id=front_section_id,
            document_id="document_greenfield",
            heading="方案首页与版本信息",
            node_kind="front_matter",
            template_node_id="ich_m11_front_matter",
            content_blocks=[
                {
                    "block_id": "front_heading",
                    "block_type": "heading",
                    "text": "方案首页与版本信息",
                    "body_order": 0,
                    "source_kind": "greenfield_scaffold",
                    "outline_level": 0,
                },
                front_table,
            ],
        ),
        ProtocolSection(
            section_id="document_indexes",
            document_id="document_greenfield",
            heading="目录、表目录和图目录",
            node_kind="document_control",
            template_node_id="cms_indexes",
            content_blocks=[
                {
                    "block_id": "document_indexes_heading",
                    "block_type": "heading",
                    "text": "目录、表目录和图目录",
                    "body_order": 2,
                    "source_kind": "greenfield_scaffold",
                    "outline_level": 0,
                    "include_in_toc": False,
                },
            ],
        ),
    ]
    for index, (number, title) in enumerate(
        (
            ("1", "方案概要"),
            ("1.1", "方案摘要"),
            ("1.1.1", "主要和次要目的及估计目标"),
            ("1.1.1.1", "主要目的"),
        ),
        start=1,
    ):
        level = number.count(".")
        sections.append(
            ProtocolSection(
                section_id=f"section_{index}",
                document_id="document_greenfield",
                heading=title,
                section_number=number,
                content_blocks=[
                    {
                        "block_id": f"heading_{index}",
                        "block_type": "heading",
                        "text": title,
                        "body_order": index * 2,
                        "source_kind": "greenfield_scaffold",
                        "outline_level": level,
                    },
                    {
                        "block_id": f"body_{index}",
                        "block_type": "paragraph",
                        "text": f"{title}正文。",
                        "body_order": index * 2 + 1,
                        "source_kind": "greenfield_scaffold",
                    },
                ],
            )
        )
    document = ProtocolDocument(
        document_id="document_greenfield",
        project_id="project_greenfield",
        protocol_id="CMS-D017-TEST",
        version="V0.1",
        sections=sections,
        style_profile_id=definition.style_profile_id,
        style_profile_version=definition.style_profile_version,
        style_profile_definition_sha256=definition.definition_sha256,
    )

    result = export_medical_writing_document_docx(document, mode="draft_preview")
    exported = Document(io.BytesIO(result.content))

    paragraph_text = [paragraph.text for paragraph in exported.paragraphs]
    assert "方案首页与版本信息" not in paragraph_text
    assert "目录、表目录和图目录" not in paragraph_text
    assert "研究方案" in paragraph_text
    assert "研究方案. CMS-D017-TEST" not in paragraph_text
    assert "研究方案· CMS-D017-TEST" not in paragraph_text
    assert result.metadata["skipped_source_index_block_count"] == 1
    cover_label = next(
        paragraph
        for paragraph in exported.paragraphs
        if paragraph.text == "研究方案"
    )
    assert cover_label.runs[0].font.size.pt == pytest.approx(15.0)
    assert cover_label.runs[0].font.color.rgb == (0, 0, 0)
    assert cover_label.runs[0].font.italic is False
    cover_title = next(
        paragraph
        for paragraph in exported.paragraphs
        if "一项评价CMS-D017" in paragraph.text
    )
    assert cover_title.style.name == "Title"
    assert cover_title.runs[0].font.size.pt == pytest.approx(14.0)
    assert cover_title.runs[0].font.color.rgb == (0, 0, 0)
    metadata_rows = [
        [cell.text for cell in row.cells]
        for row in exported.tables[0].rows
    ]
    assert metadata_rows == [
        ["研究药物：", "CMS-D017"],
        ["研究阶段：", "II期"],
        ["方案/研究编号：", "CMS-D017-TEST"],
        ["版本号：", "V0.1"],
        ["版本日期：", "2026年07月20日"],
        ["申办者：", "深圳市康哲生物科技有限公司"],
    ]
    authoritative_result = export_medical_writing_document_docx(
        document,
        mode="draft_preview",
        front_matter_overrides={
            "investigational_product": "CMS-D017胶囊",
            "protocol_date": "2026年07月21日",
        },
    )
    authoritative_export = Document(io.BytesIO(authoritative_result.content))
    authoritative_rows = [
        [cell.text for cell in row.cells]
        for row in authoritative_export.tables[0].rows
    ]
    assert authoritative_rows[0] == ["研究药物：", "CMS-D017胶囊"]
    assert authoritative_rows[4] == ["版本日期：", "2026年07月21日"]
    assert (
        authoritative_result.metadata["source_snapshot_sha256"]
        != result.metadata["source_snapshot_sha256"]
    )
    assert all(
        run.font.size.pt == pytest.approx(12.0)
        for row in exported.tables[0].rows
        for cell in row.cells
        for paragraph in cell.paragraphs
        for run in paragraph.runs
        if run.text
    )

    # Exporter writes complete authoritative numbers into Heading paragraph text
    # and suppresses paragraph-level inherited numbering (numId=0) to avoid
    # Word/LibreOffice duplicate/truncated numbering.
    expected_headings = (
        "1 方案概要",
        "1.1 方案摘要",
        "1.1.1 主要和次要目的及估计目标",
        "1.1.1.1 主要目的",
    )
    heading_paragraphs = [
        next(paragraph for paragraph in exported.paragraphs if paragraph.text == title)
        for title in expected_headings
    ]
    assert [paragraph.style.name for paragraph in heading_paragraphs] == [
        "Heading 1",
        "Heading 2",
        "Heading 3",
        "Heading 4",
    ]
    for paragraph in heading_paragraphs:
        paragraph_num_id = paragraph._p.find(
            "./w:pPr/w:numPr/w:numId", paragraph._p.nsmap
        )
        assert paragraph_num_id is not None
        assert paragraph_num_id.get(qn("w:val")) == "0"
        assert all(run.font.color.rgb == (0, 0, 0) for run in paragraph.runs if run.text)
    style_num_ids = []
    for level, style in enumerate(
        (exported.styles[f"Heading {index}"] for index in range(1, 5))
    ):
        # Style-level numbering remains for TOC/navigation identity; paragraph
        # inherits are suppressed above so visible numbers are not duplicated.
        num_id = style.element.find("./w:pPr/w:numPr/w:numId", style.element.nsmap)
        assert num_id is not None
        assert style.font.color.rgb == (0, 0, 0)
        assert style.quick_style is True
        assert style.hidden is False
        style_num_ids.append(num_id.get(qn("w:val")))
    assert len(set(style_num_ids)) == 1
    toc_heading_style = exported.styles["TOC Heading"]
    assert toc_heading_style.base_style is not None
    assert toc_heading_style.base_style.name == "Normal"
    assert toc_heading_style.element.find(
        "./w:pPr/w:numPr",
        toc_heading_style.element.nsmap,
    ) is None

    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        document_xml = archive.read("word/document.xml").decode("utf-8")
        settings_root = ElementTree.fromstring(archive.read("word/settings.xml"))
        numbering_root = ElementTree.fromstring(archive.read("word/numbering.xml"))
    assert 'TOC \\o "1-4" \\h \\z \\u' in document_xml
    compatibility_mode = next(
        item
        for item in settings_root.findall(
            ".//w:compatSetting",
            {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )
        if item.get(qn("w:name")) == "compatibilityMode"
    )
    assert compatibility_mode.get(qn("w:val")) == "15"
    level_texts = {
        element.get(qn("w:val"))
        for element in numbering_root.findall(
            ".//w:lvlText",
            {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )
    }
    assert {"%1", "%1.%2", "%1.%2.%3", "%1.%2.%3.%4"}.issubset(level_texts)
    company_abstract = next(
        abstract
        for abstract in numbering_root.findall(
            "./w:abstractNum",
            {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
        )
        if {
            element.get(qn("w:val"))
            for element in abstract.findall(
                "./w:lvl/w:lvlText",
                {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
            )
        }
        == {"%1", "%1.%2", "%1.%2.%3", "%1.%2.%3.%4"}
    )
    word_ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    assert company_abstract.find("./w:nsid", word_ns) is not None
    assert company_abstract.find("./w:tmpl", word_ns) is not None
    assert (
        company_abstract.get(
            "{http://schemas.microsoft.com/office/word/2012/wordml}"
            "restartNumberingAfterBreak"
        )
        is None
    )
    expected_indents = {
        "0": ("0", None),
        "1": ("0", None),
        "2": ("567", "567"),
        "3": ("851", "851"),
    }
    for level in company_abstract.findall("./w:lvl", word_ns):
        level_index = level.get(qn("w:ilvl"))
        indentation = level.find("./w:pPr/w:ind", word_ns)
        assert indentation is not None
        expected_left, expected_hanging = expected_indents[level_index]
        assert indentation.get(qn("w:left")) == expected_left
        assert indentation.get(qn("w:hanging")) == expected_hanging


def test_legacy_document_keeps_legacy_export_style():
    result = export_medical_writing_document_docx(
        _document_with_profile(profile=False),
        mode="draft_preview",
    )
    exported = Document(io.BytesIO(result.content))

    assert result.metadata["style_profile_id"] == ""
    assert (
        exported.sections[0].top_margin.twips
        >= exported.sections[0].header_distance.twips + 790
    )
    assert exported.paragraphs[0].runs[0].font.size.pt == pytest.approx(16.0)


def test_unknown_or_mutated_style_profile_fails_closed():
    with pytest.raises(MedicalWritingDocumentDocxExportError, match="hash"):
        export_medical_writing_document_docx(
            _document_with_profile(digest_override="0" * 64),
            mode="draft_preview",
        )

    incomplete = _document_with_profile(profile=False).model_copy(
        update={"style_profile_id": "cms_cn_clinical_protocol"}
    )
    with pytest.raises(MedicalWritingDocumentDocxExportError, match="incomplete"):
        export_medical_writing_document_docx(incomplete, mode="draft_preview")

    incomplete_corpus = _document_with_profile().model_copy(
        update={"corpus_snapshot_id": "cms_cn_protocol_corpus"}
    )
    with pytest.raises(MedicalWritingDocumentDocxExportError, match="corpus-snapshot"):
        export_medical_writing_document_docx(
            incomplete_corpus,
            mode="draft_preview",
        )
