from __future__ import annotations

import io
import zipfile
from xml.etree import ElementTree

import pytest
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn

from packages.contracts.workbench_contracts import (
    StructuredTable,
    StructuredTableCell,
    StructuredTableColumn,
    StructuredTableDomain,
    StructuredTableNote,
    StructuredTableNoteType,
    StructuredTableRole,
    StructuredTableRow,
)
from services.api.app.medical_writing_table_exporter import (
    StructuredTableDocxExportError,
    export_structured_table_docx,
)


WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS = {"w": WORD_NS}


def test_objectives_endpoints_table_exports_deterministically_with_readable_notes_and_merges():
    table = _objectives_endpoints_table()

    first = export_structured_table_docx(table)
    second = export_structured_table_docx(table)

    assert first.content == second.content
    assert first.metadata["docx_sha256"] == second.metadata["docx_sha256"]
    assert first.metadata["orientation"] == "portrait"
    assert first.metadata["repeat_header_rows"] == 2
    assert (
        first.metadata["long_table_split_strategy"] == "repeat_header_keep_rows_intact"
    )
    assert first.metadata["note_markers"] == ["1", "2"]

    document = Document(io.BytesIO(first.content))
    assert len(document.tables) == 1
    exported = document.tables[0]
    assert len(exported.rows) == 4
    assert len(exported.columns) == 3
    assert exported.cell(0, 0)._tc is exported.cell(0, 2)._tc
    assert exported.cell(0, 0).text == "研究目的与终点"
    assert exported.cell(2, 1).text == "第16周EASI-75应答率"

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[-3:] == [
        "附注",
        "1 主要终点按治疗策略估计目标分析。",
        "2 所有终点定义均需在SAP中预先规定。",
    ]
    assert not any("|---" in text or "| 研究目的" in text for text in paragraphs)

    root = _document_xml(first.content)
    rows = root.findall(".//w:tbl/w:tr", NS)
    assert rows[0].find("w:trPr/w:tblHeader", NS) is not None
    assert rows[1].find("w:trPr/w:tblHeader", NS) is not None
    assert rows[2].find("w:trPr/w:tblHeader", NS) is None
    assert all(row.find("w:trPr/w:cantSplit", NS) is not None for row in rows)
    for row in rows:
        header = row.find("w:trPr/w:tblHeader", NS)
        cant_split = row.find("w:trPr/w:cantSplit", NS)
        if header is not None:
            assert header.get(f"{{{WORD_NS}}}val") is None
        assert cant_split is not None
        assert cant_split.get(f"{{{WORD_NS}}}val") is None
    assert root.find(".//w:sectPr/w:pgSz", NS).get(f"{{{WORD_NS}}}orient") is None
    grid_spans = root.findall(".//w:tbl/w:tr/w:tc/w:tcPr/w:gridSpan", NS)
    assert any(span.get(f"{{{WORD_NS}}}val") == "3" for span in grid_spans)


def test_layout_table_exports_without_visible_title_header_or_grid_borders():
    table = StructuredTable(
        table_id="layout_front_matter",
        block_id="layout_front_matter_block",
        role=StructuredTableRole.LAYOUT,
        title="方案首页信息",
        header_row_count=0,
        columns=[
            StructuredTableColumn(column_id="field", order=0),
            StructuredTableColumn(column_id="value", order=1),
        ],
        rows=[
            StructuredTableRow(
                row_id="row_protocol",
                order=0,
                cells=[
                    _cell("row_protocol", "field", "field_protocol", "方案编号"),
                    _cell("row_protocol", "value", "value_protocol", "CMS-RA-201"),
                ],
            )
        ],
    )

    result = export_structured_table_docx(table)
    document = Document(io.BytesIO(result.content))

    assert [paragraph.text for paragraph in document.paragraphs] == []
    assert document.tables[0].cell(0, 0).text == "方案编号"
    assert document.tables[0].cell(0, 1).text == "CMS-RA-201"
    root = _document_xml(result.content)
    borders = root.find(".//w:tbl/w:tblPr/w:tblBorders", NS)
    assert borders is not None
    assert {border.get(f"{{{WORD_NS}}}val") for border in list(borders)} == {"nil"}
    table_properties = root.find(".//w:tbl/w:tblPr", NS)
    property_names = [item.tag.rsplit("}", 1)[-1] for item in table_properties]
    assert property_names.index("tblBorders") < property_names.index("tblLayout")
    assert property_names.index("tblBorders") < property_names.index("tblLook")


def test_note_sort_accepts_unicode_circled_digits():
    table = _objectives_endpoints_table()
    table.notes = [
        StructuredTableNote(
            note_id="note_circled_three",
            marker="③",
            note_type=StructuredTableNoteType.OPERATIONAL,
            text="第三项说明",
            target_ids=[table.table_id],
        ),
        StructuredTableNote(
            note_id="note_circled_one",
            marker="①",
            note_type=StructuredTableNoteType.OPERATIONAL,
            text="第一项说明",
            target_ids=[table.table_id],
        ),
    ]

    result = export_structured_table_docx(table)

    assert result.metadata["note_markers"] == ["①", "③"]


def test_table_cell_rich_text_exports_inline_and_paragraph_formatting():
    rich_cell = _cell("r0", "c0", "rich", "剂量10 mg/L")
    rich_cell.rich_text = {
        "type": "doc",
        "content": [
            {
                "type": "paragraph",
                "attrs": {
                    "stylePreset": "body",
                    "textAlign": "right",
                    "lineHeight": 1.5,
                    "spacingBeforePt": 2,
                    "spacingAfterPt": 3,
                    "leftIndentChars": 1,
                    "rightIndentChars": 0,
                    "firstLineIndentChars": 0,
                },
                "content": [
                    {"type": "text", "text": "剂量", "marks": [{"type": "bold"}]},
                    {"type": "text", "text": "10", "marks": [{"type": "superscript"}]},
                    {
                        "type": "text",
                        "text": " mg/L",
                        "marks": [
                            {"type": "underline"},
                            {"type": "highlight", "attrs": {"color": "#FFFF00"}},
                            {
                                "type": "textStyle",
                                "attrs": {
                                    "fontFamily": "Arial",
                                    "fontSize": "12pt",
                                    "color": "#C00000",
                                },
                            },
                        ],
                    },
                ],
            }
        ],
    }
    table = _invalid_table(cells=[rich_cell, _cell("r0", "c1", "plain", "对照")])

    result = export_structured_table_docx(table)
    document = Document(io.BytesIO(result.content))
    paragraph = document.tables[0].cell(0, 0).paragraphs[0]

    assert paragraph.text == "剂量10 mg/L"
    assert paragraph.alignment == 2
    assert paragraph.paragraph_format.line_spacing == 1.5
    content_runs = [run for run in paragraph.runs if run.text]
    assert content_runs[0].bold is True
    assert content_runs[1].font.superscript is True
    assert content_runs[2].underline is True
    assert content_runs[2]._element.rPr.highlight.val == WD_COLOR_INDEX.YELLOW
    assert str(content_runs[2].font.color.rgb) == "C00000"
    for run in content_runs:
        assert run.font.name == "Times New Roman"
        assert run._element.rPr.rFonts.get(qn("w:eastAsia")) == "宋体"
        assert run._element.rPr.rFonts.get(qn("w:ascii")) == "Times New Roman"
        assert run._element.rPr.rFonts.get(qn("w:hAnsi")) == "Times New Roman"


def test_42_by_19_soa_exports_landscape_repeating_headers_and_complex_merges():
    table = _dense_soa_table()

    result = export_structured_table_docx(table)

    assert result.metadata["row_count"] == 42
    assert result.metadata["column_count"] == 19
    assert result.metadata["orientation"] == "landscape"
    assert result.metadata["repeat_header_rows"] == 3
    assert (
        result.metadata["long_table_split_strategy"] == "repeat_header_keep_rows_intact"
    )
    assert len(result.metadata["merged_ranges"]) == 6
    assert result.metadata["note_markers"] == ["1", "2", "10"]

    document = Document(io.BytesIO(result.content))
    exported = document.tables[0]
    assert len(exported.rows) == 42
    assert len(exported.columns) == 19
    assert exported.cell(0, 1)._tc is exported.cell(0, 5)._tc
    assert exported.cell(0, 6)._tc is exported.cell(0, 15)._tc
    assert exported.cell(4, 0)._tc is exported.cell(5, 0)._tc
    assert exported.cell(6, 5)._tc is exported.cell(6, 6)._tc
    assert exported.cell(41, 18).text == "X"

    root = _document_xml(result.content)
    page_size = root.find(".//w:sectPr/w:pgSz", NS)
    assert page_size.get(f"{{{WORD_NS}}}orient") == "landscape"
    rows = root.findall(".//w:tbl/w:tr", NS)
    assert len(rows) == 42
    assert all(
        rows[index].find("w:trPr/w:tblHeader", NS) is not None for index in range(3)
    )
    assert rows[3].find("w:trPr/w:tblHeader", NS) is None
    assert all(row.find("w:trPr/w:cantSplit", NS) is not None for row in rows)
    grid_columns = root.findall(".//w:tbl/w:tblGrid/w:gridCol", NS)
    assert [
        int(column.get(f"{{{WORD_NS}}}w")) for column in grid_columns
    ] == result.metadata["column_widths_twips"]
    assert (
        root.find(".//w:tbl/w:tblPr/w:tblLayout", NS).get(f"{{{WORD_NS}}}type")
        == "fixed"
    )
    assert root.findall(".//w:vMerge", NS)
    assert root.findall(".//w:gridSpan", NS)
    header_shading = rows[0].find("w:tc/w:tcPr/w:shd", NS)
    assert header_shading.get(f"{{{WORD_NS}}}fill") == "FFFFFF"
    section_shading = rows[3].find("w:tc/w:tcPr/w:shd", NS)
    assert section_shading.get(f"{{{WORD_NS}}}fill") == "EEECE1"
    serialized = ElementTree.tostring(root, encoding="unicode")
    assert "D9E2F3" not in serialized
    assert "E2F0D9" not in serialized

    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs[-4:] == [
        "附注",
        "1 筛选期血常规包括血红蛋白、白细胞计数、血小板计数。",
        "2 EOT访视适用于提前终止试验治疗的受试者。",
        "10 计划外访视仅在医学需要时实施并记录原因。",
    ]


def test_allow_row_split_strategy_keeps_only_headers_and_sections_intact():
    table = _objectives_endpoints_table()
    table.word_layout["long_table_split_strategy"] = "repeat_header_allow_row_split"

    result = export_structured_table_docx(table)

    root = _document_xml(result.content)
    rows = root.findall(".//w:tbl/w:tr", NS)
    assert (
        result.metadata["long_table_split_strategy"] == "repeat_header_allow_row_split"
    )
    assert rows[0].find("w:trPr/w:cantSplit", NS) is not None
    assert rows[1].find("w:trPr/w:cantSplit", NS) is not None
    assert rows[2].find("w:trPr/w:cantSplit", NS) is None
    assert rows[3].find("w:trPr/w:cantSplit", NS) is None


def test_d017_synopsis_profile_emits_physical_ooxml_fidelity_contract():
    table = _d017_synopsis_table()

    result = export_structured_table_docx(table)

    root = _document_xml(result.content)
    word_table = root.find(".//w:tbl", NS)
    assert word_table.find("w:tblPr/w:tblStyle", NS) is None
    table_width = word_table.find("w:tblPr/w:tblW", NS)
    assert table_width.get(f"{{{WORD_NS}}}type") == "pct"
    assert table_width.get(f"{{{WORD_NS}}}w") == "4999"
    assert [
        int(column.get(f"{{{WORD_NS}}}w"))
        for column in word_table.findall("w:tblGrid/w:gridCol", NS)
    ] == [1554, 2993, 4514]

    borders = word_table.find("w:tblPr/w:tblBorders", NS)
    assert borders is not None
    assert {
        (
            border.tag.rsplit("}", 1)[-1],
            border.get(f"{{{WORD_NS}}}val"),
            border.get(f"{{{WORD_NS}}}sz"),
            border.get(f"{{{WORD_NS}}}space"),
            border.get(f"{{{WORD_NS}}}color"),
        )
        for border in list(borders)
    } == {
        (edge, "single", "4", "0", "auto")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV")
    }

    rows = word_table.findall("w:tr", NS)
    assert [row.find("w:trPr/w:cantSplit", NS) is not None for row in rows] == [
        False,
        True,
        False,
        True,
        False,
    ]
    assert [
        (
            row.find("w:trPr/w:trHeight", NS).get(f"{{{WORD_NS}}}val")
            if row.find("w:trPr/w:trHeight", NS) is not None
            else None
        )
        for row in rows
    ] == [None, "335", "90", "335", "335"]
    assert all(
        row.find("w:trPr/w:trHeight", NS).get(f"{{{WORD_NS}}}hRule") is None
        for row in rows[1:]
    )

    first_row_cells = rows[0].findall("w:tc", NS)
    assert [
        (
            cell.find("w:tcPr/w:tcW", NS).get(f"{{{WORD_NS}}}type"),
            cell.find("w:tcPr/w:tcW", NS).get(f"{{{WORD_NS}}}w"),
        )
        for cell in first_row_cells
    ] == [("pct", "857"), ("pct", "4142")]
    assert first_row_cells[0].find("w:tcPr/w:vAlign", NS) is None
    assert (
        first_row_cells[1].find("w:tcPr/w:vAlign", NS).get(f"{{{WORD_NS}}}val") == "top"
    )
    section_cells = rows[1].findall("w:tc", NS)
    assert all(cell.find("w:tcPr/w:vAlign", NS) is None for cell in section_cells)
    body_cells = rows[2].findall("w:tc", NS)[1:]
    assert all(
        cell.find("w:tcPr/w:vAlign", NS).get(f"{{{WORD_NS}}}val") == "top"
        for cell in body_cells
    )
    assert all(
        shading.get(f"{{{WORD_NS}}}fill") == "FFFFFF"
        for shading in word_table.findall(".//w:tcPr/w:shd", NS)
    )
    assert word_table.find(".//w:tcPr/w:tcMar", NS) is None

    nonempty_paragraphs = [
        paragraph
        for paragraph in word_table.findall(".//w:p", NS)
        if "".join(paragraph.itertext()).strip()
        or paragraph.find("w:pPr/w:numPr", NS) is not None
    ]
    assert nonempty_paragraphs
    for paragraph in nonempty_paragraphs:
        spacing = paragraph.find("w:pPr/w:spacing", NS)
        assert spacing is not None
        assert (
            spacing.get(f"{{{WORD_NS}}}before"),
            spacing.get(f"{{{WORD_NS}}}after"),
            spacing.get(f"{{{WORD_NS}}}line"),
            spacing.get(f"{{{WORD_NS}}}lineRule"),
        ) == ("120", "60", "400", "exact")

    primary_paragraphs = rows[2].findall(".//w:p", NS)
    assert all(
        paragraph.find("w:pPr/w:numPr", NS) is None for paragraph in primary_paragraphs
    )
    numbered_paragraphs = rows[4].findall(".//w:p", NS)
    assert (
        sum(
            paragraph.find("w:pPr/w:numPr", NS) is not None
            for paragraph in numbered_paragraphs
        )
        == 4
    )
    with zipfile.ZipFile(io.BytesIO(result.content)) as archive:
        numbering = ElementTree.fromstring(archive.read("word/numbering.xml"))
    level_texts = {
        item.get(f"{{{WORD_NS}}}val") for item in numbering.findall(".//w:lvlText", NS)
    }
    assert "%1)" in level_texts
    assert "%1）" in level_texts
    assert "• " not in "".join(root.itertext())

    section = root.find(".//w:sectPr", NS)
    margins = section.find("w:pgMar", NS)
    assert (
        margins.get(f"{{{WORD_NS}}}left"),
        margins.get(f"{{{WORD_NS}}}right"),
        margins.get(f"{{{WORD_NS}}}top"),
        margins.get(f"{{{WORD_NS}}}bottom"),
    ) == ("1417", "1417", "1417", "1134")


@pytest.mark.parametrize(
    ("factory_name", "message"),
    [
        pytest.param("_overlapping_table", "overlaps", id="conflicting-merge"),
        pytest.param("_table_with_hole", "uncovered grid positions", id="grid-hole"),
        pytest.param(
            "_out_of_bounds_merge_table", "exceeds the table grid", id="out-of-bounds"
        ),
    ],
)
def test_irregular_or_conflicting_merge_geometry_fails_before_docx_generation(
    factory_name, message
):
    table = globals()[factory_name]()
    with pytest.raises(StructuredTableDocxExportError, match=message):
        export_structured_table_docx(table)


def _objectives_endpoints_table() -> StructuredTable:
    columns = [
        StructuredTableColumn(
            column_id="objective", order=0, label="研究目的", width_twips=2400
        ),
        StructuredTableColumn(
            column_id="endpoint", order=1, label="终点", width_twips=4200
        ),
        StructuredTableColumn(
            column_id="estimand", order=2, label="估计目标", width_twips=3200
        ),
    ]
    rows = [
        StructuredTableRow(
            row_id="title",
            order=0,
            style_role="header",
            cells=[
                _cell(
                    "title",
                    "objective",
                    "table_title",
                    "研究目的与终点",
                    column_span=3,
                    role="header",
                )
            ],
        ),
        StructuredTableRow(
            row_id="headers",
            order=1,
            style_role="header",
            cells=[
                _cell(
                    "headers",
                    column.column_id,
                    f"header_{column.column_id}",
                    column.label,
                    role="header",
                )
                for column in columns
            ],
        ),
        StructuredTableRow(
            row_id="primary",
            order=2,
            cells=[
                _cell(
                    "primary",
                    "objective",
                    "primary_objective",
                    "评价研究药物治疗中重度特应性皮炎的有效性",
                ),
                _cell(
                    "primary",
                    "endpoint",
                    "primary_endpoint",
                    "第16周EASI-75应答率",
                    note_refs=["note_1"],
                ),
                _cell("primary", "estimand", "primary_estimand", "治疗策略估计目标"),
            ],
        ),
        StructuredTableRow(
            row_id="secondary",
            order=3,
            cells=[
                _cell("secondary", "objective", "secondary_objective", "评价瘙痒改善"),
                _cell(
                    "secondary",
                    "endpoint",
                    "secondary_endpoint",
                    "第4周PP-NRS较基线变化",
                ),
                _cell(
                    "secondary",
                    "estimand",
                    "secondary_estimand",
                    "假想策略估计目标",
                    note_refs=["note_2"],
                ),
            ],
        ),
    ]
    notes = [
        StructuredTableNote(
            note_id="note_2",
            marker="2",
            note_type=StructuredTableNoteType.DEFINITION,
            text="所有终点定义均需在SAP中预先规定。",
            target_ids=["secondary_estimand"],
        ),
        StructuredTableNote(
            note_id="note_1",
            marker="1",
            note_type=StructuredTableNoteType.OPERATIONAL,
            text="主要终点按治疗策略估计目标分析。",
            target_ids=["primary_endpoint"],
        ),
    ]
    return StructuredTable(
        table_id="tbl_objectives_endpoints",
        block_id="block_objectives_endpoints",
        domain=StructuredTableDomain.OBJECTIVES_ENDPOINTS,
        title="表2 研究目的与终点",
        header_row_count=2,
        columns=columns,
        rows=rows,
        notes=notes,
        word_layout={
            "orientation": "portrait",
            "repeat_header_rows": 2,
            "long_table_split_strategy": "repeat_header_keep_rows_intact",
        },
    )


def _dense_soa_table() -> StructuredTable:
    columns = [
        StructuredTableColumn(
            column_id=f"c{index}",
            order=index,
            label="评估项目" if index == 0 else f"V{index}",
            width_twips=1800 if index == 0 else 500,
        )
        for index in range(19)
    ]
    rows = [
        StructuredTableRow(
            row_id="r0",
            order=0,
            style_role="header",
            cells=[
                _cell("r0", "c0", "r0c0", "评估项目", role="header"),
                _cell("r0", "c1", "r0c1", "筛选期", column_span=5, role="header"),
                _cell("r0", "c6", "r0c6", "治疗期", column_span=10, role="header"),
                _cell("r0", "c16", "r0c16", "随访期", column_span=3, role="header"),
            ],
        ),
        StructuredTableRow(
            row_id="r1",
            order=1,
            style_role="header",
            cells=[
                _cell(
                    "r1",
                    f"c{index}",
                    f"r1c{index}",
                    "项目" if index == 0 else f"V{index}",
                    role="header",
                )
                for index in range(19)
            ],
        ),
        StructuredTableRow(
            row_id="r2",
            order=2,
            style_role="header",
            cells=[
                _cell(
                    "r2",
                    f"c{index}",
                    f"r2c{index}",
                    "研究日" if index == 0 else f"D{index * 7}",
                    role="header",
                )
                for index in range(19)
            ],
        ),
        StructuredTableRow(
            row_id="r3",
            order=3,
            style_role="section",
            cells=[_cell("r3", "c0", "r3c0", "知情同意与一般评估", column_span=19)],
        ),
    ]
    for row_index in range(4, 42):
        cells = []
        for column_index in range(19):
            if row_index == 5 and column_index == 0:
                continue
            if row_index == 6 and column_index == 6:
                continue
            row_span = 2 if row_index == 4 and column_index == 0 else 1
            column_span = 2 if row_index == 6 and column_index == 5 else 1
            text = (
                f"评估项目{row_index - 3}"
                if column_index == 0
                else ("X" if (row_index + column_index) % 3 == 0 else "")
            )
            note_refs = []
            if row_index == 4 and column_index == 1:
                note_refs = ["soa_note_1"]
            elif row_index == 10 and column_index == 18:
                note_refs = ["soa_note_2", "soa_note_10"]
            cells.append(
                _cell(
                    f"r{row_index}",
                    f"c{column_index}",
                    f"r{row_index}c{column_index}",
                    text,
                    row_span=row_span,
                    column_span=column_span,
                    note_refs=note_refs,
                )
            )
        rows.append(
            StructuredTableRow(row_id=f"r{row_index}", order=row_index, cells=cells)
        )
    rows[-1].cells[-1].text = "X"

    notes = [
        StructuredTableNote(
            note_id="soa_note_10",
            marker="10",
            note_type=StructuredTableNoteType.CONDITION,
            text="计划外访视仅在医学需要时实施并记录原因。",
            target_ids=["r10c18"],
        ),
        StructuredTableNote(
            note_id="soa_note_2",
            marker="2",
            note_type=StructuredTableNoteType.TIMING_RULE,
            text="EOT访视适用于提前终止试验治疗的受试者。",
            target_ids=["r10c18"],
        ),
        StructuredTableNote(
            note_id="soa_note_1",
            marker="1",
            note_type=StructuredTableNoteType.ITEM_SET,
            text="筛选期血常规包括血红蛋白、白细胞计数、血小板计数。",
            target_ids=["r4c1"],
        ),
    ]
    return StructuredTable(
        table_id="tbl_soa_42x19",
        block_id="block_soa_42x19",
        domain=StructuredTableDomain.SCHEDULE_OF_ACTIVITIES,
        title="表1 研究流程表",
        header_row_count=3,
        columns=columns,
        rows=rows,
        notes=notes,
        word_layout={
            "orientation": "landscape",
            "repeat_header_rows": 3,
            "long_table_split_strategy": "repeat_header_keep_rows_intact",
            "margins_twips": {"left": 540, "right": 540, "top": 540, "bottom": 540},
            "font_size_pt": 6.5,
        },
    )


def _overlapping_table() -> StructuredTable:
    return _invalid_table(
        cells=[
            _cell("r0", "c0", "a", "A", column_span=2),
            _cell("r0", "c1", "b", "B"),
        ]
    )


def _table_with_hole() -> StructuredTable:
    return _invalid_table(cells=[_cell("r0", "c0", "a", "A")])


def _out_of_bounds_merge_table() -> StructuredTable:
    return _invalid_table(cells=[_cell("r0", "c0", "a", "A", column_span=3)])


def _d017_synopsis_table() -> StructuredTable:
    columns = [
        StructuredTableColumn(
            column_id="label",
            order=0,
            label="项目",
            width_twips=1554,
        ),
        StructuredTableColumn(
            column_id="objective",
            order=1,
            label="目的/内容",
            width_twips=2993,
        ),
        StructuredTableColumn(
            column_id="endpoint",
            order=2,
            label="相应的研究终点",
            width_twips=4514,
        ),
    ]
    primary_objective = _cell(
        "primary_body",
        "objective",
        "primary_objective",
        "评价主要有效性。",
    )
    primary_objective.rich_text = {
        "type": "doc",
        "content": [
            {
                "type": "paragraph",
                "attrs": {"stylePreset": "synopsis_body"},
                "content": [{"type": "text", "text": "评价主要有效性。"}],
            }
        ],
    }
    primary_endpoint = _cell(
        "primary_body",
        "endpoint",
        "primary_endpoint",
        "治疗D84时Hb较基线变化。",
    )
    primary_endpoint.rich_text = {
        "type": "doc",
        "content": [
            {
                "type": "paragraph",
                "attrs": {"stylePreset": "synopsis_body"},
                "content": [
                    {
                        "type": "text",
                        "text": "主要有效性终点：",
                        "marks": [{"type": "bold"}],
                    }
                ],
            },
            {
                "type": "paragraph",
                "attrs": {"stylePreset": "synopsis_body"},
                "content": [{"type": "text", "text": "治疗D84时Hb较基线变化。"}],
            },
        ],
    }
    secondary_objective = _cell(
        "secondary_body",
        "objective",
        "secondary_objective",
        "评价次要有效性。评价安全性。",
    )
    secondary_objective.rich_text = _numbered_rich_text(
        ["评价次要有效性。", "评价安全性。"],
        "decimal_half_paren",
    )
    secondary_endpoint = _cell(
        "secondary_body",
        "endpoint",
        "secondary_endpoint",
        "第12周应答率。AE发生率。",
    )
    secondary_endpoint.rich_text = _numbered_rich_text(
        ["第12周应答率。", "AE发生率。"],
        "decimal_fullwidth_paren",
        heading="次要有效性终点：",
    )
    rows = [
        StructuredTableRow(
            row_id="title",
            order=0,
            cells=[
                _cell("title", "label", "title_label", "研究题目", role="label"),
                _cell(
                    "title",
                    "objective",
                    "title_value",
                    "CMS-D017 PNH临床试验",
                    column_span=2,
                ),
            ],
        ),
        StructuredTableRow(
            row_id="primary_header",
            order=1,
            style_role="section",
            cells=[
                _cell(
                    "primary_header",
                    "label",
                    "objective_group",
                    "目的与估计目标/终点",
                    row_span=4,
                    role="section",
                ),
                _cell(
                    "primary_header",
                    "objective",
                    "primary_header_objective",
                    "主要目的",
                    role="section",
                ),
                _cell(
                    "primary_header",
                    "endpoint",
                    "primary_header_endpoint",
                    "相应的研究终点",
                    role="section",
                ),
            ],
        ),
        StructuredTableRow(
            row_id="primary_body",
            order=2,
            cells=[
                _hidden_merge_cell(
                    "primary_body",
                    "label",
                    "primary_group_continuation",
                ),
                primary_objective,
                primary_endpoint,
            ],
        ),
        StructuredTableRow(
            row_id="secondary_header",
            order=3,
            style_role="section",
            cells=[
                _hidden_merge_cell(
                    "secondary_header",
                    "label",
                    "secondary_group_continuation",
                ),
                _cell(
                    "secondary_header",
                    "objective",
                    "secondary_header_objective",
                    "次要目的",
                    role="section",
                ),
                _cell(
                    "secondary_header",
                    "endpoint",
                    "secondary_header_endpoint",
                    "相应的研究终点",
                    role="section",
                ),
            ],
        ),
        StructuredTableRow(
            row_id="secondary_body",
            order=4,
            cells=[
                _hidden_merge_cell(
                    "secondary_body",
                    "label",
                    "secondary_body_group_continuation",
                ),
                secondary_objective,
                secondary_endpoint,
            ],
        ),
    ]
    return StructuredTable(
        table_id="d017_synopsis",
        block_id="d017_synopsis_block",
        domain=StructuredTableDomain.GENERIC,
        role=StructuredTableRole.PROTOCOL_SYNOPSIS,
        title="",
        header_row_count=0,
        columns=columns,
        rows=rows,
        word_layout={
            "orientation": "portrait",
            "fit_to_page": True,
            "margins_twips": {
                "left": 1417,
                "right": 1417,
                "top": 1417,
                "bottom": 1134,
            },
            "table_style": "",
            "table_width_type": "pct",
            "table_width_value": 4999,
            "explicit_black_borders": True,
            "long_table_split_strategy": "repeat_header_allow_row_split",
            "row_min_heights_twips": {1: 335, 2: 90, 3: 335, 4: 335},
            "font_size_pt": 11,
            "use_table_font_size_for_body": True,
            "bold_labels": True,
            "paragraph_spacing_before_twips": 120,
            "paragraph_spacing_after_twips": 60,
            "paragraph_line_twips": 400,
            "paragraph_alignment_by_role": {
                "label": "justify",
                "body": "justify",
                "header": "justify",
                "section": "justify",
            },
            "vertical_alignment_by_role": {
                "label": "inherit",
                "body": "top",
                "header": "inherit",
                "section": "inherit",
            },
            "cell_fill_by_role": {
                "body": "FFFFFF",
                "header": "FFFFFF",
                "section": "FFFFFF",
            },
            "cell_margins_twips": None,
        },
    )


def _numbered_rich_text(items, numbering_format, *, heading=None):
    content = []
    if heading:
        content.append(
            {
                "type": "paragraph",
                "attrs": {"stylePreset": "synopsis_body"},
                "content": [
                    {"type": "text", "text": heading, "marks": [{"type": "bold"}]}
                ],
            }
        )
    content.append(
        {
            "type": "orderedList",
            "attrs": {"numberingFormat": numbering_format},
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        {
                            "type": "paragraph",
                            "attrs": {"stylePreset": "synopsis_body"},
                            "content": [{"type": "text", "text": item}],
                        }
                    ],
                }
                for item in items
            ],
        }
    )
    return {"type": "doc", "content": content}


def _hidden_merge_cell(row_id, column_id, cell_id):
    cell = _cell(row_id, column_id, cell_id, "")
    cell.semantic_value = {
        "_medical_writing_table": {
            "hidden": True,
            "merge_parent_cell_id": "objective_group",
        }
    }
    return cell


def _invalid_table(cells) -> StructuredTable:
    return StructuredTable(
        table_id="invalid_table",
        block_id="invalid_block",
        columns=[
            StructuredTableColumn(column_id="c0", order=0),
            StructuredTableColumn(column_id="c1", order=1),
        ],
        rows=[StructuredTableRow(row_id="r0", order=0, cells=cells)],
    )


def _cell(
    row_id,
    column_id,
    cell_id,
    text,
    *,
    row_span=1,
    column_span=1,
    role="body",
    note_refs=None,
):
    return StructuredTableCell(
        cell_id=cell_id,
        row_id=row_id,
        column_id=column_id,
        text=text,
        row_span=row_span,
        column_span=column_span,
        style_role=role,
        note_refs=note_refs or [],
    )


def _document_xml(content: bytes):
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return ElementTree.fromstring(archive.read("word/document.xml"))
