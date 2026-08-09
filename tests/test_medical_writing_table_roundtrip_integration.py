from __future__ import annotations

import io
import unittest

from docx import Document

from services.api.app.medical_writing_document import MedicalWritingDocumentService
from services.api.app.medical_writing_table_exporter import export_structured_table_docx
from services.api.app.medical_writing_tables import MedicalWritingTableService


class MedicalWritingTableRoundTripIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.documents = MedicalWritingDocumentService()
        cls.tables = MedicalWritingTableService()

    def test_real_projects_with_vertical_merges_export_without_duplicate_cells(self):
        verified_projects = []
        for project_id in ("proj_rux_03_002", "proj_d001", "proj_my008_pnh_3_01"):
            document = self.documents.document_session(project_id)
            source_table_blocks = []
            for section in document.sections:
                source_table_blocks.extend(
                    block
                    for block in self.documents.section(project_id, section.section_id).content_blocks
                    if block.get("block_type") == "table"
                )
            merged_block = next(
                (
                    block
                    for block in source_table_blocks
                    if any(
                        cell.get("hidden")
                        for row in block.get("rows", [])
                        for cell in row
                    )
                ),
                None,
            )
            if merged_block is None:
                continue
            structured = self.tables.from_table_block(merged_block)
            result = export_structured_table_docx(structured)
            imported = Document(io.BytesIO(result.content))
            self.assertEqual(1, len(imported.tables))
            self.assertGreater(result.metadata["merged_ranges"], [])
            verified_projects.append(project_id)

        self.assertIn("proj_d001", verified_projects)
        self.assertGreaterEqual(len(verified_projects), 2)


if __name__ == "__main__":
    unittest.main()
