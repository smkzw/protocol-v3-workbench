from __future__ import annotations

import io
import tempfile
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest.mock import Mock, patch

import openpyxl
import docx
from fastapi.testclient import TestClient

from packages.contracts.workbench_contracts import (
    AiTaskFromRegistryRequest,
    ListingSheetPayload,
    SourceContentValidationCheck,
    SourceContentValidationConfirmationRequest,
    SourceContentValidationRecord,
)
from services.api.app.source_content_validation import (
    SourceContentValidationConflict,
    SourceContentValidationService,
    SourceContentValidationStore,
    SourceExpectedContext,
)
from services.api.app.main import app
from services.api.app import main as main_module
from services.api.app.monitoring_identity_authorization import MonitoringRole
from services.api.app.monitoring_runtime_principal import MonitoringAuthenticatedPrincipal
from services.api.app import eligibility
from services.api.app.safety_pv_manifest import MY009_LISTING as MY009_SAFETY_LISTING
from services.api.app.ai_execution_policy import AiExecutionPolicyResolver
from services.api.app.protocol_text_extractor import parse_protocol_docx
from services.api.app.source_intake import SourceRegistryService, SourceRegistryStore


RUX_LISTING = Path(
    "/Users/smkzw/Documents/康哲项目资料/Ruxolitinib-AD/CFDI Inspection/准备阶段/"
    "RUX-03-002_列表_数据集_Excel_20250612_处理后.xlsx"
)
MY009_LISTING = Path(
    "/Users/smkzw/Documents/朗来项目资料/MY009治疗UC/S1安全性评价-202604/"
    "MY009-UC-2-01-MM Listing_20260408(已自动还原).xlsx"
)
D001_PROTOCOL = Path(
    "/Users/smkzw/Documents/康哲项目资料/AI/入排/test-D001项目/"
    "CMS-D001 银屑病2、3期临床方案 v1.0-2025.12.21.docx"
)


class SourceContentValidationTests(unittest.TestCase):
    class _PrincipalMiddleware:
        def __init__(self, application, principal):
            self.application = application
            self.principal = principal

        async def __call__(self, scope, receive, send):
            if scope["type"] == "http":
                scope.setdefault("state", {})["monitoring_principal"] = self.principal
            await self.application(scope, receive, send)

    @staticmethod
    def _principal() -> MonitoringAuthenticatedPrincipal:
        return MonitoringAuthenticatedPrincipal(
            principal_id="source-validation-test",
            tenant_id="tenant-kangzhe",
            roles=(MonitoringRole.DATA_MANAGEMENT, MonitoringRole.MEDICAL_MANAGER),
            project_scope=(
                "proj_d001",
                "proj_mgk10_sar_demo",
                "proj_my009_uc",
                "proj_rux_03_002",
            ),
            issued_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
            expires_at=datetime(2099, 1, 1, tzinfo=timezone.utc),
            authenticated=True,
            authn_method="test-server-session",
            session_id="source-validation-test-session",
            directory_revision="source-validation-test-v1",
            verification_ref_sha256="3" * 64,
        )

    def _client(self):
        return TestClient(self._PrincipalMiddleware(app, self._principal()))

    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.store = SourceContentValidationStore(Path(self.tmp.name) / "source_validations.sqlite3")
        self.service = SourceContentValidationService(self.store)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_listing_match_is_allowed_without_confirmation(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_rux_03_002",
            source_entry_id="src_rux_listing",
            module="medical_monitoring",
            filename="RUX-03-002_listing.xlsx",
            file_sha256="a" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="AE",
                    rows=[{"STUDYID": "RUX-03-002", "SUBJID": "S01001", "AETERM": "用药部位灼烧感"}],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="edc_data_listing",
            ),
            actor="system_validator",
        )

        self.assertEqual("matched", record.content_status)
        self.assertEqual("allowed", record.use_status)
        self.assertEqual(1, record.revision)

    def test_listing_project_check_accepts_common_edc_pstudyid_alias(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_rux_03_002",
            source_entry_id="src_rux_pstudyid",
            module="medical_monitoring",
            filename="RUX-03-002_listing.xlsx",
            file_sha256="b" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="AE",
                    rows=[
                        {
                            "PSTUDYID": "RUX-03-002",
                            "SUBJID": "S01001",
                            "AETERM": "用药部位灼烧感",
                        }
                    ],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="edc_data_listing",
            ),
            actor="system_validator",
        )

        self.assertEqual("matched", record.content_status)
        self.assertEqual("allowed", record.use_status)

    def test_safety_medical_review_listing_keeps_cm_and_study_treatment_changes_separate(self) -> None:
        sheets = [
            ListingSheetPayload(
                sheet_name="AE",
                rows=[{"STUDYID": "STUDY-01", "USUBJID": "01-001", "AETERM": "头痛", "AESTDAT": "2026-01-02", "AESER": "否"}],
            ),
            ListingSheetPayload(sheet_name="MH", rows=[{"USUBJID": "01-001", "MHTERM": "高血压"}]),
            ListingSheetPayload(sheet_name="CM", rows=[{"USUBJID": "01-001", "CMTRT": "氨氯地平"}]),
            ListingSheetPayload(
                sheet_name="EX2",
                rows=[{"USUBJID": "01-001", "EX2DOSE": "5", "EX2ADJ": "减量"}],
            ),
        ]
        record = self.service.assess_listing(
            project_id="proj_study_01",
            source_entry_id="src_safety_listing",
            module="safety_pv",
            filename="safety_review_listing.xlsx",
            file_sha256="d" * 64,
            sheets=sheets,
            expected=SourceExpectedContext(
                project_identifiers=("STUDY-01",),
                expected_file_role="safety_medical_review_listing",
            ),
            actor="system_validator",
        )

        checks = {check.check_code: check for check in record.checks}
        self.assertEqual("match", checks["file_role"].outcome)
        self.assertEqual("match", checks["safety_review_event_structure"].outcome)
        self.assertEqual("match", checks["safety_review_data_scope"].outcome)
        self.assertIn("非试验用药（CM）", checks["safety_review_data_scope"].observed_value)
        self.assertIn("试验用药/剂量调整", checks["safety_review_data_scope"].observed_value)
        self.assertNotIn("TEAE", checks["safety_review_data_scope"].observed_value)
        self.assertEqual("matched", record.content_status)

    def test_ae_only_file_is_role_matched_but_scope_requires_confirmation(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_study_01",
            source_entry_id="src_ae_only",
            module="safety_pv",
            filename="ae_listing.xlsx",
            file_sha256="e" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="AE",
                    rows=[{"STUDYID": "STUDY-01", "USUBJID": "01-001", "AETERM": "头痛", "AESTDAT": "2026-01-02", "AESEV": "轻度"}],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("STUDY-01",),
                expected_file_role="safety_medical_review_listing",
            ),
            actor="system_validator",
        )

        checks = {check.check_code: check for check in record.checks}
        self.assertEqual("mismatch", checks["file_role"].outcome)
        self.assertEqual("match", checks["safety_review_event_structure"].outcome)
        self.assertEqual("mismatch", checks["safety_review_data_scope"].outcome)
        self.assertEqual("mismatch", record.content_status)
        self.assertEqual("requires_confirmation", record.use_status)

    def test_safety_role_requires_subject_and_event_term_on_the_same_sheet(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_study_01",
            source_entry_id="src_cross_sheet_false_positive",
            module="safety_pv",
            filename="split_fields.xlsx",
            file_sha256="1" * 64,
            sheets=[
                ListingSheetPayload(sheet_name="DM", rows=[{"STUDYID": "STUDY-01", "USUBJID": "01-001"}]),
                ListingSheetPayload(sheet_name="AE", rows=[{"AETERM": "头痛", "AESTDAT": "2026-01-02", "AESEV": "轻度"}]),
                ListingSheetPayload(sheet_name="VS", rows=[{"USUBJID": "01-001", "CM": "incidental", "VSDAT": "2026-01-02"}]),
            ],
            expected=SourceExpectedContext(
                project_identifiers=("STUDY-01",),
                expected_file_role="safety_medical_review_listing",
            ),
            actor="system_validator",
        )

        checks = {check.check_code: check for check in record.checks}
        self.assertEqual("mismatch", checks["safety_review_event_structure"].outcome)
        self.assertNotIn("非试验用药（CM）", checks["safety_review_data_scope"].observed_value)
        self.assertEqual("mismatch", checks["file_role"].outcome)

    def test_non_safety_listing_does_not_pass_safety_file_role(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_study_01",
            source_entry_id="src_non_safety",
            module="safety_pv",
            filename="efficacy_listing.xlsx",
            file_sha256="f" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="QS",
                    rows=[{"STUDYID": "STUDY-01", "USUBJID": "01-001", "QSTEST": "评分"}],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("STUDY-01",),
                expected_file_role="safety_medical_review_listing",
            ),
            actor="system_validator",
        )

        checks = {check.check_code: check for check in record.checks}
        self.assertEqual("mismatch", checks["file_role"].outcome)
        self.assertEqual("mismatch", checks["safety_review_data_scope"].outcome)
        self.assertEqual("requires_confirmation", record.use_status)

    def test_same_size_directory_file_replacement_invalidates_prior_confirmation(self) -> None:
        root = Path(self.tmp.name) / "unlabeled_subject_bundle"
        root.mkdir()
        source_file = root / "source.pdf"
        source_file.write_bytes(b"AAAA")
        registry = SourceRegistryService(
            SourceRegistryStore(Path(self.tmp.name) / "bundle_sources.jsonl"),
            content_validation_service=self.service,
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("CMS-D001",),
                expected_file_role="raw_subject_bundle_inventory",
            ),
        )

        first = registry.register_raw_subject_bundle("proj_d001", root)
        first_validation = registry.current_content_validation(
            "proj_d001", first.entry.entry_id
        )
        unresolved = [
            check.check_code
            for check in first_validation.checks
            if check.outcome in {"warning", "mismatch"}
        ]
        confirmed = registry.confirm_content_validation(
            "proj_d001",
            first.entry.entry_id,
            SourceContentValidationConfirmationRequest(
                reason="医学经理已核对资料包目录，确认本项目沿用。",
                acknowledged_check_codes=unresolved,
                actor="medical_manager",
                expected_revision=first_validation.revision,
                idempotency_key="bundle-confirmation-before-replacement",
            ),
        )
        self.assertEqual("confirmed_after_warning", confirmed.use_status)

        source_file.write_bytes(b"BBBB")
        second = registry.register_raw_subject_bundle("proj_d001", root)
        second_validation = registry.current_content_validation(
            "proj_d001", second.entry.entry_id
        )

        self.assertNotEqual(first.entry.entry_id, second.entry.entry_id)
        self.assertEqual("requires_confirmation", second_validation.use_status)

    def test_mismatch_can_be_confirmed_without_changing_mismatch_fact(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_d001",
            source_entry_id="src_wrong_listing",
            module="medical_monitoring",
            filename="RUX-03-002_listing.xlsx",
            file_sha256="b" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="AE",
                    rows=[{"STUDYID": "RUX-03-002", "SUBJID": "S01001", "AETERM": "头痛"}],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("CMS-D001", "D001"),
                expected_file_role="edc_data_listing",
            ),
            actor="system_validator",
        )
        unresolved = [check.check_code for check in record.checks if check.outcome in {"warning", "mismatch"}]

        confirmed = self.service.confirm_after_warning(
            "proj_d001",
            "src_wrong_listing",
            SourceContentValidationConfirmationRequest(
                reason="已核对来源，确认该文件仅用于跨项目结构映射测试，不用于本项目医学结论。",
                acknowledged_check_codes=unresolved,
                actor="medical_manager",
                expected_revision=record.revision,
                idempotency_key="confirm-wrong-listing-001",
            ),
        )

        self.assertEqual("mismatch", confirmed.content_status)
        self.assertEqual("confirmed_after_warning", confirmed.use_status)
        self.assertEqual(record.checks, confirmed.checks)
        self.assertEqual(record.revision + 1, confirmed.revision)

    def test_confirmation_requires_every_current_warning_and_current_revision(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_d001",
            source_entry_id="src_missing_identity",
            module="medical_monitoring",
            filename="listing.xlsx",
            file_sha256="c" * 64,
            sheets=[ListingSheetPayload(sheet_name="AE", rows=[{"SUBJID": "SA07001", "AETERM": "头痛"}])],
            expected=SourceExpectedContext(
                project_identifiers=("CMS-D001",),
                expected_file_role="edc_data_listing",
            ),
            actor="system_validator",
        )

        with self.assertRaisesRegex(ValueError, "all current warning or mismatch checks"):
            self.service.confirm_after_warning(
                "proj_d001",
                "src_missing_identity",
                SourceContentValidationConfirmationRequest(
                    reason="已人工核对原始导出记录，确认继续使用该批次进行医学复核。",
                    acknowledged_check_codes=[],
                    expected_revision=record.revision,
                    idempotency_key="partial-ack-001",
                ),
            )

        unresolved = [check.check_code for check in record.checks if check.outcome in {"warning", "mismatch"}]
        self.service.confirm_after_warning(
            "proj_d001",
            "src_missing_identity",
            SourceContentValidationConfirmationRequest(
                reason="已人工核对原始导出记录，确认继续使用该批次进行医学复核。",
                acknowledged_check_codes=unresolved,
                expected_revision=record.revision,
                idempotency_key="complete-ack-001",
            ),
        )
        with self.assertRaises(SourceContentValidationConflict):
            self.service.confirm_after_warning(
                "proj_d001",
                "src_missing_identity",
                SourceContentValidationConfirmationRequest(
                    reason="再次提交旧版本确认，应由版本冲突阻止，不能覆盖当前记录。",
                    acknowledged_check_codes=unresolved,
                    expected_revision=record.revision,
                    idempotency_key="stale-ack-001",
                ),
            )

    def test_changed_file_hash_creates_new_unconfirmed_revision(self) -> None:
        expected = SourceExpectedContext(
            project_identifiers=("CMS-D001",),
            expected_file_role="edc_data_listing",
        )
        first = self.service.assess_listing(
            project_id="proj_d001",
            source_entry_id="src_versioned",
            module="medical_monitoring",
            filename="listing.xlsx",
            file_sha256="d" * 64,
            sheets=[ListingSheetPayload(sheet_name="AE", rows=[{"SUBJID": "SA07001", "AETERM": "头痛"}])],
            expected=expected,
            actor="system_validator",
        )
        unresolved = [check.check_code for check in first.checks if check.outcome in {"warning", "mismatch"}]
        confirmed = self.service.confirm_after_warning(
            "proj_d001",
            "src_versioned",
            SourceContentValidationConfirmationRequest(
                reason="已与数据管理员核对本次导出批次，确认无研究编号字段但来源正确。",
                acknowledged_check_codes=unresolved,
                expected_revision=first.revision,
                idempotency_key="confirm-version-1",
            ),
        )

        changed = self.service.assess_listing(
            project_id="proj_d001",
            source_entry_id="src_versioned",
            module="medical_monitoring",
            filename="listing.xlsx",
            file_sha256="e" * 64,
            sheets=[ListingSheetPayload(sheet_name="AE", rows=[{"SUBJID": "SA07002", "AETERM": "恶心"}])],
            expected=expected,
            actor="system_validator",
        )

        self.assertEqual(confirmed.revision + 1, changed.revision)
        self.assertEqual("requires_confirmation", changed.use_status)
        self.assertEqual("", changed.confirmation_reason)

    def test_technical_failure_cannot_be_confirmed(self) -> None:
        failed = SourceContentValidationRecord(
            validation_id="val_failed_1",
            project_id="proj_d001",
            source_entry_id="src_failed",
            module="eligibility_review",
            revision=1,
            technical_status="failed",
            content_status="not_assessed",
            use_status="blocked_technical_failure",
            file_sha256="f" * 64,
            expected_context_hash="context-hash",
            checks=[
                SourceContentValidationCheck(
                    check_code="technical_readability",
                    label="文件技术可读性",
                    observed_value="DOCX结构损坏",
                    outcome="mismatch",
                    overridable=False,
                )
            ],
            summary="文件无法解析，未执行内容一致性核验。",
            actor="system_validator",
            created_at=datetime.now(timezone.utc),
        )
        self.store.save_assessment(failed)

        with self.assertRaisesRegex(ValueError, "technical failure"):
            self.service.confirm_after_warning(
                "proj_d001",
                "src_failed",
                SourceContentValidationConfirmationRequest(
                    reason="即使用户确认，也不能将损坏且无法解析的文件送入后续工作流。",
                    acknowledged_check_codes=["technical_readability"],
                    expected_revision=1,
                    idempotency_key="technical-failure-001",
                ),
            )

    def test_source_registry_api_returns_validation_and_confirms_warning(self) -> None:
        root = Path(self.tmp.name)
        validation_service = SourceContentValidationService(
            SourceContentValidationStore(root / "api_validations.sqlite3")
        )
        registry = SourceRegistryService(
            SourceRegistryStore(root / "sources.jsonl"),
            content_validation_service=validation_service,
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("MG-K10-SAR-III",),
                expected_file_role="edc_data_listing",
            ),
        )
        client = self._client()
        with patch("services.api.app.main.source_registry", registry):
            direct_registration = registry.register_listing_file(
                "proj_mgk10_sar_demo",
                "medical_listing.xlsx",
                _listing_bytes_without_study_id(),
                module="medical_monitoring",
            )
            blocked_registration = client.post(
                "/api/projects/proj_mgk10_sar_demo/sources/listing-file",
                params={"filename": "medical_listing.xlsx", "module": "medical_monitoring"},
                content=_listing_bytes_without_study_id(),
            )
            self.assertEqual(403, blocked_registration.status_code)
            self.assertEqual(
                "monitoring_write_action_unconfigured",
                blocked_registration.json()["detail"]["code"],
            )
            validation_record = registry.current_content_validation(
                "proj_mgk10_sar_demo",
                direct_registration.entry.entry_id,
            )
            self.assertIsNotNone(validation_record)
            validation = validation_record.model_dump(mode="json")
            self.assertEqual("warning", validation["content_status"])
            self.assertEqual("requires_confirmation", validation["use_status"])
            self.assertNotIn("file_sha256", blocked_registration.text)
            self.assertNotIn("expected_context_hash", blocked_registration.text)

            ai_request = AiTaskFromRegistryRequest(
                module="medical_monitoring",
                task_type="listing_semantic_mapping",
                expected_prompt_version="listing_semantic_mapping_v0_1",
                source_ids=[direct_registration.spans[0].source_id],
            )
            with self.assertRaisesRegex(ValueError, "requires content-consistency confirmation"):
                registry.ai_task_request_from_registry(
                    "proj_mgk10_sar_demo",
                    ai_request,
                    AiExecutionPolicyResolver(
                        deployment_profile="local_private_clinical",
                        provider_name="buddy",
                        model_name="deepseek-v4-pro",
                    ),
                )

            unresolved = [
                check["check_code"]
                for check in validation["checks"]
                if check["outcome"] in {"warning", "mismatch"}
            ]
            confirmed = client.post(
                f"/api/projects/proj_mgk10_sar_demo/sources/{direct_registration.entry.entry_id}/content-validation/confirm",
                json={
                    "reason": "已与数据管理员核对原始导出批次，确认虽无研究编号字段但来源属于当前项目。",
                    "acknowledged_check_codes": unresolved,
                    "actor": "client-spoof",
                    "expected_revision": validation["revision"],
                    "idempotency_key": "api-confirm-warning-001",
                },
            )

        self.assertEqual(200, confirmed.status_code)
        self.assertEqual("warning", confirmed.json()["content_status"])
        self.assertEqual("confirmed_after_warning", confirmed.json()["use_status"])
        self.assertEqual(
            "source-validation-test",
            validation_service.store.history(
                "proj_mgk10_sar_demo",
                direct_registration.entry.entry_id,
            )[0].actor,
        )
        resolved = registry.ai_task_request_from_registry(
            "proj_mgk10_sar_demo",
            ai_request,
            AiExecutionPolicyResolver(
                deployment_profile="local_private_clinical",
                provider_name="buddy",
                model_name="deepseek-v4-pro",
            ),
        )
        self.assertEqual("listing_semantic_mapping", resolved.task_type)

        monitoring_content = _listing_bytes_without_study_id(subject_id="10002")
        with patch("services.api.app.main.source_registry", registry):
            blocked = client.post(
                "/api/projects/proj_mgk10_sar_demo/monitoring/intake/file",
                params={
                    "filename": "batch_without_study_id.xlsx",
                    "extract_date": "2026-07-13",
                    "batch_label": "EDC listing Batch validation test",
                },
                content=monitoring_content,
            )
            self.assertEqual(409, blocked.status_code)
            detail = blocked.json()["detail"]
            self.assertEqual("source_content_confirmation_required", detail["code"])
            unresolved = [
                check["check_code"]
                for check in detail["validation"]["checks"]
                if check["outcome"] in {"warning", "mismatch"}
            ]
            approved = client.post(
                f"/api/projects/proj_mgk10_sar_demo/sources/{detail['source_entry_id']}/content-validation/confirm",
                json={
                    "reason": "已核对该EDC导出批次来源，确认缺少研究编号字段但属于当前项目。",
                    "acknowledged_check_codes": unresolved,
                    "expected_revision": detail["validation"]["revision"],
                    "idempotency_key": "monitoring-intake-confirm-001",
                },
            )
            self.assertEqual(200, approved.status_code)
            retried = client.post(
                "/api/projects/proj_mgk10_sar_demo/monitoring/intake/file",
                params={
                    "filename": "batch_without_study_id.xlsx",
                    "extract_date": "2026-07-13",
                    "batch_label": "EDC listing Batch validation test",
                },
                content=monitoring_content,
            )
        self.assertEqual(200, retried.status_code, retried.text)
        self.assertEqual(
            "confirmed_after_warning",
            retried.json()["content_validation"]["use_status"],
        )

    def test_eligibility_source_admission_refresh_fails_closed_before_registry(self) -> None:
        registry = Mock()
        with patch("services.api.app.main.source_registry", registry):
            no_principal = TestClient(app).post(
                "/api/projects/proj_d001/eligibility/source-admission/refresh"
            )
            scoped_principal = self._client().post(
                "/api/projects/proj_d001/eligibility/source-admission/refresh"
            )

        self.assertEqual(503, no_principal.status_code, no_principal.text)
        self.assertEqual(
            "monitoring_principal_unavailable",
            no_principal.json()["detail"]["code"],
        )
        self.assertEqual(403, scoped_principal.status_code, scoped_principal.text)
        self.assertEqual(
            "monitoring_write_action_unconfigured",
            scoped_principal.json()["detail"]["code"],
        )
        registry.register_local_file.assert_not_called()
        registry.register_raw_subject_bundle.assert_not_called()

    def test_store_is_restart_durable_and_record_table_is_immutable(self) -> None:
        record = self.service.assess_listing(
            project_id="proj_rux_03_002",
            source_entry_id="src_restart",
            module="medical_monitoring",
            filename="listing.xlsx",
            file_sha256="8" * 64,
            sheets=[
                ListingSheetPayload(
                    sheet_name="AE",
                    rows=[{"STUDYID": "RUX-03-002", "SUBJID": "S01001", "AETERM": "头痛"}],
                )
            ],
            expected=SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="edc_data_listing",
            ),
            actor="system_validator",
        )
        reopened = SourceContentValidationStore(self.store.db_path)
        self.assertEqual(record, reopened.current("proj_rux_03_002", "src_restart"))
        with reopened._connect() as connection:
            with self.assertRaisesRegex(Exception, "immutable"):
                connection.execute(
                    "UPDATE source_content_validation_records SET operation='tampered' WHERE validation_id=?",
                    (record.validation_id,),
                )

    def test_tfl_inventory_role_is_checked_and_project_clue_remains_warning(self) -> None:
        package = Path(self.tmp.name) / "sdtm_package"
        package.mkdir()
        (package / "dm.xpt").write_bytes(b"SAS transport placeholder")
        (package / "define.xml").write_text("<ODM />", encoding="utf-8")
        registry = SourceRegistryService(
            SourceRegistryStore(Path(self.tmp.name) / "bundle_sources.jsonl"),
            allowed_roots=[Path(self.tmp.name)],
            content_validation_service=self.service,
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="tfl_dataset_package_inventory",
            ),
        )

        result = registry.register_local_directory(
            "proj_rux_03_002",
            package,
            module="data_analysis_tfl",
            source_kind="tfl_dataset_package_inventory",
        )
        validation = registry.current_content_validation(
            "proj_rux_03_002", result.entry.entry_id
        )
        checks = {check.check_code: check for check in validation.checks}

        self.assertEqual("match", checks["file_role"].outcome)
        self.assertEqual("warning", checks["project_identity"].outcome)
        self.assertEqual("requires_confirmation", validation.use_status)

    def test_empty_inventory_is_non_overridable_technical_failure(self) -> None:
        package = Path(self.tmp.name) / "empty_package"
        package.mkdir()
        (package / "empty.xpt").write_bytes(b"")
        registry = SourceRegistryService(
            SourceRegistryStore(Path(self.tmp.name) / "empty_sources.jsonl"),
            allowed_roots=[Path(self.tmp.name)],
            content_validation_service=self.service,
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="tfl_dataset_package_inventory",
            ),
        )

        result = registry.register_local_directory(
            "proj_rux_03_002",
            package,
            module="data_analysis_tfl",
            source_kind="tfl_dataset_package_inventory",
        )
        validation = registry.current_content_validation(
            "proj_rux_03_002", result.entry.entry_id
        )

        self.assertEqual("failed", validation.technical_status)
        self.assertEqual("blocked_technical_failure", validation.use_status)

    def test_protocol_role_and_project_identity_are_checked_from_content(self) -> None:
        content = _protocol_bytes(
            "RUX-03-002 临床研究方案",
            "研究设计",
            "入选标准",
            "排除标准",
        )
        document = parse_protocol_docx("RUX-03-002_protocol.docx", content)
        record = self.service.assess_protocol(
            project_id="proj_d001",
            source_entry_id="src_wrong_protocol",
            module="eligibility_review",
            filename="RUX-03-002_protocol.docx",
            file_sha256="9" * 64,
            document=document,
            expected=SourceExpectedContext(
                project_identifiers=("CMS-D001",),
                indication_terms=("银屑病",),
                expected_file_role="protocol_docx",
                expected_protocol_version="V1.0",
            ),
            actor="system_validator",
        )

        checks = {check.check_code: check for check in record.checks}
        self.assertEqual("match", checks["file_role"].outcome)
        self.assertEqual("mismatch", checks["project_identity"].outcome)
        self.assertEqual("warning", checks["indication"].outcome)
        self.assertEqual("mismatch", record.content_status)
        self.assertEqual("requires_confirmation", record.use_status)

    @unittest.skipUnless(
        RUX_LISTING.exists() and MY009_LISTING.exists() and D001_PROTOCOL.exists(),
        "real RUX, MY009, and D001 source files are required",
    )
    def test_two_real_projects_do_not_share_hard_coded_identity_rules(self) -> None:
        root = Path(self.tmp.name)
        contexts = {
            "proj_rux_03_002": SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="edc_data_listing",
            ),
            "proj_my009_uc": SourceExpectedContext(
                project_identifiers=("MY009-UC-2-01",),
                expected_file_role="edc_data_listing",
            ),
            "proj_d001": SourceExpectedContext(
                project_identifiers=("CMS-D001",),
                indication_terms=("银屑病",),
                expected_file_role="protocol_docx",
                expected_protocol_version="V1.0",
            ),
        }
        registry = SourceRegistryService(
            SourceRegistryStore(root / "real_sources.jsonl"),
            content_validation_service=SourceContentValidationService(
                SourceContentValidationStore(root / "real_validations.sqlite3")
            ),
            expected_context_resolver=lambda project_id, module, source_kind: contexts[project_id],
        )

        rux = registry.register_listing_file(
            "proj_rux_03_002",
            RUX_LISTING.name,
            RUX_LISTING.read_bytes(),
            module="medical_monitoring",
        )
        my009 = registry.register_listing_file(
            "proj_my009_uc",
            MY009_LISTING.name,
            MY009_LISTING.read_bytes(),
            module="medical_monitoring",
        )
        d001 = registry.register_protocol_docx(
            "proj_d001",
            D001_PROTOCOL.name,
            D001_PROTOCOL.read_bytes(),
            module="eligibility_review",
        )

        for project_id, result in (
            ("proj_rux_03_002", rux),
            ("proj_my009_uc", my009),
            ("proj_d001", d001),
        ):
            validation = registry.current_content_validation(project_id, result.entry.entry_id)
            checks = {check.check_code: check for check in validation.checks}
            self.assertEqual("match", checks["file_role"].outcome)
            self.assertNotEqual("mismatch", validation.content_status)
        self.assertEqual(
            "warning",
            registry.current_content_validation(
                "proj_rux_03_002", rux.entry.entry_id
            ).checks[2].outcome,
        )
        self.assertEqual(
            "match",
            registry.current_content_validation(
                "proj_my009_uc", my009.entry.entry_id
            ).checks[2].outcome,
        )
        self.assertEqual(
            "match",
            registry.current_content_validation(
                "proj_d001", d001.entry.entry_id
            ).checks[2].outcome,
        )

    @unittest.skipUnless(MY009_SAFETY_LISTING.exists(), "real MY009 safety listing is required")
    def test_real_my009_safety_listing_has_generic_medical_review_scope(self) -> None:
        registry = SourceRegistryService(
            SourceRegistryStore(Path(self.tmp.name) / "real_safety_sources.jsonl"),
            content_validation_service=self.service,
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("MY009-UC-2-01",),
                expected_file_role="safety_medical_review_listing",
            ),
        )
        result = registry.register_listing_file(
            "proj_my009_uc",
            MY009_SAFETY_LISTING.name,
            MY009_SAFETY_LISTING.read_bytes(),
            module="safety_pv",
            expected_file_role="safety_medical_review_listing",
        )
        validation = registry.current_content_validation("proj_my009_uc", result.entry.entry_id)
        checks = {check.check_code: check for check in validation.checks}

        self.assertEqual("match", checks["file_role"].outcome)
        self.assertEqual("match", checks["safety_review_event_structure"].outcome)
        self.assertEqual("match", checks["safety_review_data_scope"].outcome)
        self.assertIn("病史", checks["safety_review_data_scope"].observed_value)
        self.assertIn("非试验用药（CM）", checks["safety_review_data_scope"].observed_value)
        self.assertIn("试验用药/剂量调整", checks["safety_review_data_scope"].observed_value)
        self.assertEqual("matched", validation.content_status)

    @unittest.skipUnless(RUX_LISTING.exists(), "real RUX listing is required")
    def test_real_rux_listing_also_satisfies_generic_safety_review_structure(self) -> None:
        from services.api.app.listing_file_parser import parse_listing_file

        record = self.service.assess_listing(
            project_id="proj_rux_03_002",
            source_entry_id="src_rux_safety_structure",
            module="safety_pv",
            filename=RUX_LISTING.name,
            file_sha256="2" * 64,
            sheets=parse_listing_file(RUX_LISTING.name, RUX_LISTING.read_bytes()),
            expected=SourceExpectedContext(
                project_identifiers=("RUX-03-002",),
                expected_file_role="safety_medical_review_listing",
            ),
            actor="system_validator",
        )
        checks = {check.check_code: check for check in record.checks}

        self.assertEqual("match", checks["safety_review_event_structure"].outcome)
        self.assertEqual("match", checks["safety_review_data_scope"].outcome)
        self.assertEqual("match", checks["file_role"].outcome)

    @unittest.skipUnless(
        D001_PROTOCOL.exists() and MY009_LISTING.exists(),
        "real D001 and MY009 project roots are required",
    )
    def test_two_real_eligibility_projects_refresh_protocol_and_subject_bundle_admission(self) -> None:
        root = Path(self.tmp.name)
        registry = SourceRegistryService(
            SourceRegistryStore(root / "eligibility_sources.jsonl"),
            allowed_roots=[
                Path("/Users/smkzw/Documents/康哲项目资料"),
                Path("/Users/smkzw/Documents/朗来项目资料"),
            ],
            content_validation_service=SourceContentValidationService(
                SourceContentValidationStore(root / "eligibility_validations.sqlite3")
            ),
            expected_context_resolver=lambda project_id, module, source_kind: {
                "proj_d001": SourceExpectedContext(
                    project_identifiers=("CMS-D001",),
                    indication_terms=("银屑病",),
                    expected_file_role=(
                        "protocol_docx"
                        if source_kind == "protocol_docx"
                        else "raw_subject_bundle_inventory"
                    ),
                    expected_protocol_version="V1.0",
                ),
                "proj_my009_uc": SourceExpectedContext(
                    project_identifiers=("MY009-UC-2-01",),
                    indication_terms=("溃疡性结肠炎",),
                    expected_file_role=(
                        "protocol_docx"
                        if source_kind == "protocol_docx"
                        else "raw_subject_bundle_inventory"
                    ),
                    expected_protocol_version="V3.0",
                ),
            }[project_id],
        )
        client = self._client()
        with patch("services.api.app.main.source_registry", registry):
            responses = {
                project_id: client.post(
                    f"/api/projects/{project_id}/eligibility/source-admission/refresh"
                )
                for project_id in ("proj_d001", "proj_my009_uc")
            }

        for project_id, response in responses.items():
            self.assertEqual(403, response.status_code, response.text)
            self.assertEqual(
                "monitoring_write_action_unconfigured",
                response.json()["detail"]["code"],
            )

        with patch("services.api.app.main.source_registry", registry):
            states = {
                project_id: main_module._refresh_eligibility_source_admission(project_id)
                for project_id in ("proj_d001", "proj_my009_uc")
            }
        for project_id, payload in states.items():
            self.assertEqual(2, len(payload["sources"]))
            by_kind = {item["entry"]["source_kind"]: item for item in payload["sources"]}
            protocol_validation = by_kind["protocol_docx"]["content_validation"]
            bundle_validation = by_kind["raw_subject_bundle_inventory"]["content_validation"]
            self.assertEqual("matched", protocol_validation["content_status"])
            self.assertEqual("allowed", protocol_validation["use_status"])
            self.assertEqual("warning", bundle_validation["content_status"])
            self.assertEqual("requires_confirmation", bundle_validation["use_status"])
            self.assertFalse(payload["ready_for_use"])
            self.assertNotIn("/Users/", str(payload))

    @unittest.skipUnless(D001_PROTOCOL.exists(), "real D001 project root is required")
    def test_eligibility_write_guard_opens_only_after_all_real_sources_are_usable(self) -> None:
        root = Path(self.tmp.name)
        registry = SourceRegistryService(
            SourceRegistryStore(root / "eligibility_guard_sources.jsonl"),
            allowed_roots=[Path("/Users/smkzw/Documents/康哲项目资料")],
            content_validation_service=SourceContentValidationService(
                SourceContentValidationStore(root / "eligibility_guard_validations.sqlite3")
            ),
            expected_context_resolver=lambda project_id, module, source_kind: SourceExpectedContext(
                project_identifiers=("CMS-D001",),
                indication_terms=("银屑病",),
                expected_file_role=(
                    "protocol_docx"
                    if source_kind == "protocol_docx"
                    else "raw_subject_bundle_inventory"
                ),
                expected_protocol_version="V1.0",
            ),
        )
        client = self._client()
        with patch("services.api.app.main.source_registry", registry):
            blocked = client.post(
                "/api/projects/proj_d001/eligibility/source-admission/refresh"
            )
            self.assertEqual(403, blocked.status_code, blocked.text)
            self.assertEqual(
                "monitoring_write_action_unconfigured",
                blocked.json()["detail"]["code"],
            )
            admission = main_module._refresh_eligibility_source_admission("proj_d001")
            self.assertFalse(admission["ready_for_use"])
            with self.assertRaises(Exception):
                eligibility._require_eligibility_source_admission("proj_d001")

            bundle = next(
                item
                for item in admission["sources"]
                if item["entry"]["source_kind"] == "raw_subject_bundle_inventory"
            )
            validation = bundle["content_validation"]
            unresolved = [
                check["check_code"]
                for check in validation["checks"]
                if check["outcome"] in {"warning", "mismatch"}
            ]
            confirmation = client.post(
                f"/api/projects/proj_d001/sources/{bundle['entry']['entry_id']}/content-validation/confirm",
                json={
                    "reason": "医学经理已核对资料目录及文件构成，确认本项目沿用。",
                    "acknowledged_check_codes": unresolved,
                    "actor": "medical_manager",
                    "expected_revision": validation["revision"],
                    "idempotency_key": "eligibility-guard-real-d001",
                },
            )
            self.assertEqual(200, confirmation.status_code, confirmation.text)
            refreshed = main_module._refresh_eligibility_source_admission("proj_d001")
            self.assertTrue(refreshed["ready_for_use"])
            self.assertTrue(
                eligibility._require_eligibility_source_admission("proj_d001")["ready_for_use"]
            )


def _listing_bytes_without_study_id(subject_id: str = "10001") -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "AE"
    sheet.append(["SUBJID", "AETERM"])
    sheet.append([subject_id, "头痛"])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _protocol_bytes(*paragraphs: str) -> bytes:
    document = docx.Document()
    for paragraph in paragraphs:
        document.add_paragraph(paragraph)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


if __name__ == "__main__":
    unittest.main()
